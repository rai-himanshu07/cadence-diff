"""Cross-feature integration for the structural hardening contracts."""

from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table

from qc_tool.config.profile import DeliverableProfile
from qc_tool.coverage import CoverageState, QCRunMode
from qc_tool.engine import run_qc
from qc_tool.findings import FindingClass
from qc_tool.io.loader import load_workbook_snapshot


def _save_cycle_workbook(path: Path, *, current: bool) -> None:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Data"
    headers = ["Metric", "Jan-26", "Feb-26", "Mar-26", "Apr-26", "May-26"]
    values: list[str | None] = ["Revenue", "=1", "=2", "=3", "=4", None if current else "=5"]
    if current:
        headers.append("Jun-26")
        values.append(None)
    sheet.append(headers)
    sheet.append(values)
    sheet.add_table(
        Table(
            displayName="Metrics",
            ref=f"A1:{'G' if current else 'F'}2",
        )
    )
    if not current:
        validation = DataValidation(
            type="whole",
            operator="between",
            formula1="0",
            formula2="100",
        )
        validation.add("B2:E2")
        sheet.add_data_validation(validation)
    sheet.conditional_formatting.add(
        "B2:E2",
        CellIsRule(
            operator="greaterThan",
            formula=["0"],
            fill=PatternFill("solid", fgColor="FF00FF00"),
        ),
    )
    workbook.save(path)


def test_structured_workbook_surfaces_integrate_without_false_degradation(
    tmp_path: Path,
) -> None:
    path = tmp_path / "integrated.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Data"
    sheet.append(["Metric", "Jan-26", "Feb-26", "Mar-26", "Apr-26", "May-26", "Total"])
    sheet.append(["Revenue", "=1", "=2", "=3", "=4", None, "=SUM(Metrics[Jan-26])"])
    sheet.add_table(Table(displayName="Metrics", ref="A1:F2"))

    validation = DataValidation(type="whole", operator="between", formula1="0", formula2="100")
    validation.add("B2:E2")
    sheet.add_data_validation(validation)
    sheet.conditional_formatting.add(
        "B2:E2",
        CellIsRule(
            operator="greaterThan",
            formula=["0"],
            fill=PatternFill("solid", fgColor="FF00FF00"),
        ),
    )

    chart = LineChart()
    chart.title = "Revenue trend"
    chart.add_data(Reference(sheet, min_col=2, max_col=5, min_row=2, max_row=2))
    chart.set_categories(Reference(sheet, min_col=2, max_col=5, min_row=1, max_row=1))
    chart.anchor = "I2"
    sheet.add_chart(chart)
    workbook.save(path)

    profile = DeliverableProfile.model_validate(
        {
            "name": "integrated",
            "excel": {
                "sheets": {
                    "Data": {
                        "availability_rules": [
                            {
                                "range": "B2:F2",
                                "periods": "B1:F1",
                                "required_through": "Apr-26",
                            }
                        ]
                    }
                }
            },
        }
    )
    snapshot = load_workbook_snapshot(path)
    result = run_qc(
        current_excel=path,
        profile=profile,
        mode=QCRunMode.CURRENT_FILE_PREFLIGHT,
    )

    assert len(snapshot.tables) == 1
    assert len(snapshot.charts) == 1
    assert len(snapshot.data_validations) == 1
    assert len(snapshot.conditional_formats) == 1
    assert not any(
        finding.finding_class is FindingClass.FORMULA_MISSING
        and finding.location == "F2"
        for finding in result.findings
    )
    coverage = {item.check_id: item for item in result.coverage}
    assert coverage["excel-intrinsic-structure"].state is CoverageState.CHECKED
    assert coverage["excel-interaction-rules"].state is CoverageState.CHECKED
    assert coverage["excel-conditional-format-styles"].state is CoverageState.CHECKED
    assert coverage["excel-dependencies"].state is CoverageState.CHECKED
    assert coverage["excel-availability"].state is CoverageState.CHECKED


def test_cycle_structural_features_interact_without_false_formula_findings(
    tmp_path: Path,
) -> None:
    baseline = tmp_path / "baseline.xlsx"
    current = tmp_path / "current.xlsx"
    _save_cycle_workbook(baseline, current=False)
    _save_cycle_workbook(current, current=True)
    profile = DeliverableProfile.model_validate(
        {
            "name": "integrated-cycle",
            "excel": {
                "sheets": {
                    "Data": {
                        "availability_rules": [
                            {
                                "range": "B2:G2",
                                "periods": "B1:G1",
                                "required_through": "Apr-26",
                            }
                        ]
                    }
                }
            },
        }
    )

    result = run_qc(
        baseline_excel=baseline,
        current_excel=current,
        profile=profile,
    )

    assert any(
        finding.finding_class is FindingClass.TABLE_STRUCTURE_CHANGED
        for finding in result.findings
    )
    assert any(
        finding.finding_class is FindingClass.DATA_VALIDATION_CHANGED
        and "removed" in finding.message
        for finding in result.findings
    )
    assert not any(
        finding.finding_class is FindingClass.FORMULA_REMOVED
        and finding.location == "F2"
        for finding in result.findings
    )
    assert not any(
        finding.finding_class is FindingClass.FORMULA_NOT_EXTENDED
        and finding.location == "G2"
        for finding in result.findings
    )
    coverage = {item.check_id: item for item in result.coverage}
    assert coverage["excel-interaction-rules"].state is CoverageState.CHECKED
    assert coverage["excel-availability"].state is CoverageState.CHECKED
