"""Eight-member final-package memory probe for 16 GB target machines.

Inputs are synthetic and generated under ``tests/fixtures/generated``. The
child runs the complete ``perform_run`` path, including history recording,
under an optional address-space limit and emits aggregate metrics only.
"""

from __future__ import annotations

import argparse
import hashlib
import importlib
import json
import subprocess
import sys
import time
from pathlib import Path

import psutil
from openpyxl import Workbook
from pptx import Presentation

_CHILD = "child"


def _process_memory_bytes() -> tuple[int, int]:
    current = peak = 0
    for line in Path("/proc/self/status").read_text(encoding="ascii").splitlines():
        if line.startswith("VmRSS:"):
            current = int(line.split()[1]) * 1024
        elif line.startswith("VmHWM:"):
            peak = int(line.split()[1]) * 1024
    return current, peak


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1 << 20), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _write_workbook(
    path: Path,
    *,
    member_index: int,
    rows: int,
    columns: int,
    finding_heavy: bool,
) -> None:
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet("Data")
    sheet.append(["Key", *(f"Metric {index}" for index in range(1, columns))])
    for row in range(1, rows + 1):
        if finding_heavy:
            values: list[object] = [
                f"member-{member_index}-row-{row}",
                "#REF!",
                *(member_index * 1_000_000 + row + column for column in range(2, columns)),
            ]
        else:
            values = [
                f"member-{member_index}-row-{row}",
                *(
                    member_index * 1_000_000 + row * columns + column
                    for column in range(1, columns)
                ),
            ]
        sheet.append(values)
    workbook.save(path)


def _write_deck(path: Path) -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    title = slide.shapes.title
    if title is None:
        raise RuntimeError("PowerPoint fixture layout has no title")
    title.text = "Synthetic final package"
    presentation.save(str(path))


def _generate_inputs(
    directory: Path,
    *,
    case: str,
    rows: int,
    columns: int,
) -> dict[str, Path]:
    directory.mkdir(parents=True, exist_ok=True)
    files: dict[str, Path] = {}
    for index in range(8):
        path = directory / f"{case}-member-{index}.xlsx"
        _write_workbook(
            path,
            member_index=index,
            rows=rows,
            columns=columns,
            finding_heavy=case == "findings",
        )
        files[f"current_excel:wb{index}"] = path
    deck = directory / f"{case}-deck.pptx"
    _write_deck(deck)
    files["current_ppt"] = deck
    return files


def _child(arguments: argparse.Namespace) -> int:
    if arguments.limit_gb > 0:
        resource = importlib.import_module("resource")
        limit = int(arguments.limit_gb * 1024**3)
        resource.setrlimit(resource.RLIMIT_AS, (limit, limit))

    from qc_tool.config.profile import DeliverableProfile
    from qc_tool.coverage import QCRunMode
    from qc_tool.package import PackageManifest
    from qc_tool.progress import ProgressEvent
    from qc_tool.run_service import perform_run

    files = {
        role: Path(path)
        for role, path in json.loads(arguments.files_json).items()
    }
    before = {role: _sha256(path) for role, path in files.items()}
    phases: list[dict[str, object]] = []
    started = time.perf_counter()

    def on_progress(event: ProgressEvent) -> None:
        current, peak = _process_memory_bytes()
        phases.append(
            {
                "phase": event.phase.value,
                "processed": event.processed,
                "total": event.total,
                "current_rss_bytes": current,
                "peak_rss_bytes": peak,
            }
        )

    artifacts = perform_run(
        Path(arguments.work_dir),
        files,
        {},
        DeliverableProfile(name=f"package-memory-{arguments.case}"),
        mode=QCRunMode.FINAL_PACKAGE,
        package_manifest=PackageManifest.from_role_files(files),
        on_progress=on_progress,
    )
    current, peak = _process_memory_bytes()
    after = {role: _sha256(path) for role, path in files.items()}
    print(
        json.dumps(
            {
                "case": arguments.case,
                "findings": len(artifacts.result.findings),
                "elapsed_seconds": round(time.perf_counter() - started, 2),
                "current_rss_bytes": current,
                "peak_rss_bytes": peak,
                "source_hashes_unchanged": before == after,
                "run_id": artifacts.run_id,
                "phases": phases,
            }
        )
    )
    return 0


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--case", choices=("snapshot", "findings"), required=True)
    parser.add_argument("--rows", type=int, required=True)
    parser.add_argument("--columns", type=int, required=True)
    parser.add_argument("--limit-gb", type=float, default=6.0)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--resident-pid", type=int, default=0)
    parser.add_argument("--role", default="", help=argparse.SUPPRESS)
    parser.add_argument("--files-json", default="", help=argparse.SUPPRESS)
    parser.add_argument("--work-dir", default="", help=argparse.SUPPRESS)
    arguments = parser.parse_args(argv)
    if arguments.role == _CHILD:
        return _child(arguments)

    fixture_dir = Path("tests/fixtures/generated/package-memory")
    files = _generate_inputs(
        fixture_dir,
        case=arguments.case,
        rows=arguments.rows,
        columns=arguments.columns,
    )
    arguments.output.mkdir(parents=True, exist_ok=True)
    work_dir = arguments.output / f"work-{arguments.case}"
    command = [
        sys.executable,
        "-m",
        "tests.package_memory_probe",
        "--role",
        _CHILD,
        "--case",
        arguments.case,
        "--rows",
        str(arguments.rows),
        "--columns",
        str(arguments.columns),
        "--limit-gb",
        str(arguments.limit_gb),
        "--output",
        str(arguments.output),
        "--files-json",
        json.dumps({role: str(path) for role, path in files.items()}),
        "--work-dir",
        str(work_dir),
    ]
    started = time.perf_counter()
    process = psutil.Popen(
        command,
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
        text=True,
    )
    resident = (
        psutil.Process(arguments.resident_pid)
        if arguments.resident_pid > 0
        else None
    )
    worker_peak_rss = 0
    resident_peak_rss = 0
    combined_peak_rss = 0
    while process.poll() is None:
        try:
            worker_rss = process.memory_info().rss
        except (psutil.AccessDenied, psutil.NoSuchProcess):
            worker_rss = 0
        try:
            resident_rss = resident.memory_info().rss if resident is not None else 0
        except (psutil.AccessDenied, psutil.NoSuchProcess):
            resident_rss = 0
        worker_peak_rss = max(worker_peak_rss, worker_rss)
        resident_peak_rss = max(resident_peak_rss, resident_rss)
        combined_peak_rss = max(combined_peak_rss, worker_rss + resident_rss)
        time.sleep(0.02)
    stdout, stderr = process.communicate()
    record: dict[str, object] = {
        "case": arguments.case,
        "rows_per_member": arguments.rows,
        "columns": arguments.columns,
        "members": 8,
        "limit_gb": arguments.limit_gb,
        "exit_code": process.returncode,
        "wall_seconds": round(time.perf_counter() - started, 2),
        "observed_process_rss": {
            "worker_peak_bytes": worker_peak_rss,
            "resident_peak_bytes": resident_peak_rss,
            "combined_peak_bytes": combined_peak_rss,
        },
    }
    if process.returncode == 0 and stdout.strip():
        record["result"] = json.loads(stdout.strip().splitlines()[-1])
    else:
        record["stderr_tail"] = stderr.strip().splitlines()[-12:]
    target = arguments.output / f"{arguments.case}-r{arguments.rows}.json"
    target.write_text(json.dumps(record, indent=2), encoding="utf-8")
    print(json.dumps(record, indent=2))
    return 0 if process.returncode == 0 else 1


if __name__ == "__main__":
    raise SystemExit(main())
