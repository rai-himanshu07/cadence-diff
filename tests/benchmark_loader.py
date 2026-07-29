"""Reproducible peak-RSS and elapsed-time benchmark for OOXML loaders.

Generated workbooks and JSONL results belong under ignored ``artifacts/``.
Run from the repository root, for example::

    conda run -n py311 python tests/benchmark_loader.py run \
        --output artifacts/loader-benchmark --rows 20000 --repeats 3
"""

from __future__ import annotations

import argparse
import contextlib
import hashlib
import json
import statistics
import subprocess
import sys
import time
import zipfile
from datetime import UTC, datetime
from pathlib import Path
from typing import Any, Literal
from xml.etree import ElementTree

import psutil
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill

LoaderName = Literal["oracle", "streaming"]
CASES = ("dense", "mixed", "formula", "style", "sst", "sparse")


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _rewrite_package(path: Path, replacements: dict[str, bytes]) -> None:
    with zipfile.ZipFile(path) as source:
        parts = {name: source.read(name) for name in source.namelist()}
    parts.update(replacements)
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as target:
        for name, content in parts.items():
            target.writestr(name, content)


def _convert_to_shared_strings(path: Path) -> None:
    with zipfile.ZipFile(path) as archive:
        worksheet = ElementTree.fromstring(archive.read("xl/worksheets/sheet1.xml"))
        content_types = ElementTree.fromstring(archive.read("[Content_Types].xml"))
        relationships = ElementTree.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
    strings: list[str] = []
    indexes: dict[str, int] = {}
    references = 0
    namespace = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    for cell in worksheet.iter(f"{{{namespace}}}c"):
        if cell.get("t") != "inlineStr":
            continue
        text = "".join(element.text or "" for element in cell.iter(f"{{{namespace}}}t"))
        index = indexes.get(text)
        if index is None:
            index = len(strings)
            indexes[text] = index
            strings.append(text)
        for child in list(cell):
            cell.remove(child)
        cell.set("t", "s")
        ElementTree.SubElement(cell, f"{{{namespace}}}v").text = str(index)
        references += 1

    shared = ElementTree.Element(
        f"{{{namespace}}}sst",
        count=str(references),
        uniqueCount=str(len(strings)),
    )
    for text in strings:
        item = ElementTree.SubElement(shared, f"{{{namespace}}}si")
        text_element = ElementTree.SubElement(item, f"{{{namespace}}}t")
        if text != text.strip():
            text_element.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        text_element.text = text

    content_type_namespace = "http://schemas.openxmlformats.org/package/2006/content-types"
    ElementTree.SubElement(
        content_types,
        f"{{{content_type_namespace}}}Override",
        PartName="/xl/sharedStrings.xml",
        ContentType=(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"
        ),
    )
    relationship_namespace = "http://schemas.openxmlformats.org/package/2006/relationships"
    ElementTree.SubElement(
        relationships,
        f"{{{relationship_namespace}}}Relationship",
        Id="rIdSharedStrings",
        Type=("http://schemas.openxmlformats.org/officeDocument/2006/relationships/sharedStrings"),
        Target="sharedStrings.xml",
    )
    _rewrite_package(
        path,
        {
            "[Content_Types].xml": ElementTree.tostring(
                content_types, encoding="utf-8", xml_declaration=True
            ),
            "xl/_rels/workbook.xml.rels": ElementTree.tostring(
                relationships, encoding="utf-8", xml_declaration=True
            ),
            "xl/worksheets/sheet1.xml": ElementTree.tostring(
                worksheet, encoding="utf-8", xml_declaration=True
            ),
            "xl/sharedStrings.xml": ElementTree.tostring(
                shared, encoding="utf-8", xml_declaration=True
            ),
        },
    )


def _build_case(path: Path, *, case: str, rows: int, columns: int) -> None:
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet("Data")
    sheet.append([f"Column {column}" for column in range(1, columns + 1)])
    for row_index in range(1, rows + 1):
        if case == "dense":
            row: list[object] = [float(row_index * columns + column) for column in range(columns)]
        elif case == "mixed":
            row = [
                f"2026-{row_index % 12 + 1:02d}",
                f"Region {row_index % 8}",
                *[
                    float(row_index * max(columns - 2, 1) + column)
                    for column in range(max(columns - 2, 0))
                ],
            ][:columns]
        elif case == "formula":
            row = [float(row_index), float(row_index * 2)]
            for column in range(3, columns + 1):
                row.append(f"=A{row_index + 1}+B{row_index + 1}+{column}")
        elif case == "style":
            row = []
            for column in range(columns):
                cell = WriteOnlyCell(sheet, value=float(row_index + column))
                color = f"FF{row_index % 16:02X}{column % 16:02X}80"
                cell.fill = PatternFill("solid", fgColor=color)
                cell.font = Font(bold=(row_index + column) % 2 == 0)
                cell.number_format = f"0.{(row_index + column) % 4 * '0'}"
                row.append(cell)
        elif case == "sst":
            row = [f" repeated value {(row_index + column) % 128} " for column in range(columns)]
        elif case == "sparse":
            row = [None] * columns
            row[row_index % columns] = float(row_index)
        else:
            raise ValueError(f"unknown benchmark case {case!r}")
        sheet.append(row)
    workbook.save(path)
    if case == "sst":
        _convert_to_shared_strings(path)


def _worker(path: Path, loader_name: LoaderName) -> int:
    from qc_tool.io.loader import _load_ooxml_oracle, _load_ooxml_streaming

    data = path.read_bytes()
    loader = _load_ooxml_oracle if loader_name == "oracle" else _load_ooxml_streaming
    started = time.perf_counter()
    snapshot = loader(
        data,
        source_name=path.name,
        file_format="xlsx",
        allow_large_workbook=True,
    )
    elapsed = time.perf_counter() - started
    populated_cells = sum(len(sheet.cells) for sheet in snapshot.sheets)
    print(
        json.dumps(
            {
                "elapsed_seconds": elapsed,
                "populated_cells": populated_cells,
                "sheet_count": len(snapshot.sheets),
            },
            sort_keys=True,
        )
    )
    return 0


def _sample_worker(path: Path, loader_name: LoaderName) -> dict[str, Any]:
    command = [
        sys.executable,
        "-m",
        "tests.benchmark_loader",
        "worker",
        "--input",
        str(path),
        "--loader",
        loader_name,
    ]
    process = psutil.Popen(
        command,
        cwd=Path(__file__).resolve().parents[1],
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
        text=True,
    )
    peak_rss = 0
    while process.poll() is None:
        with contextlib.suppress(psutil.NoSuchProcess, psutil.AccessDenied):
            peak_rss = max(peak_rss, process.memory_info().rss)
        time.sleep(0.01)
    stdout, stderr = process.communicate()
    if process.returncode != 0:
        raise RuntimeError(
            f"{loader_name} worker failed with {process.returncode}: {stderr.strip()}"
        )
    lines = [line for line in stdout.splitlines() if line.strip()]
    if not lines:
        raise RuntimeError(f"{loader_name} worker returned no metrics")
    metrics = json.loads(lines[-1])
    metrics["peak_rss_bytes"] = peak_rss
    return metrics


def _source_fingerprint() -> dict[str, str]:
    root = Path(__file__).resolve().parents[1]
    files = [
        root / "qc_tool/io/loader.py",
        root / "qc_tool/io/ooxml_chart.py",
        root / "qc_tool/io/ooxml_interaction.py",
        root / "qc_tool/io/ooxml_worksheet.py",
        Path(__file__).resolve(),
    ]
    return {str(path.relative_to(root)): _sha256(path) for path in files}


def _median_summary(records: list[dict[str, Any]]) -> list[dict[str, Any]]:
    summaries: list[dict[str, Any]] = []
    keys = sorted({(record["case"], record["loader"]) for record in records})
    for case, loader_name in keys:
        selected = [
            record
            for record in records
            if record["case"] == case and record["loader"] == loader_name
        ]
        summaries.append(
            {
                "case": case,
                "loader": loader_name,
                "median_elapsed_seconds": statistics.median(
                    record["elapsed_seconds"] for record in selected
                ),
                "median_peak_rss_bytes": statistics.median(
                    record["peak_rss_bytes"] for record in selected
                ),
                "repeats": len(selected),
            }
        )
    return summaries


def _comparisons(summaries: list[dict[str, Any]]) -> list[dict[str, Any]]:
    by_key = {(summary["case"], summary["loader"]): summary for summary in summaries}
    comparisons: list[dict[str, Any]] = []
    for case in sorted({summary["case"] for summary in summaries}):
        oracle = by_key[(case, "oracle")]
        streaming = by_key[(case, "streaming")]
        comparisons.append(
            {
                "case": case,
                "memory_reduction_percent": 100
                * (1 - streaming["median_peak_rss_bytes"] / oracle["median_peak_rss_bytes"]),
                "elapsed_change_percent": 100
                * (streaming["median_elapsed_seconds"] / oracle["median_elapsed_seconds"] - 1),
            }
        )
    return comparisons


def _run(args: argparse.Namespace) -> int:
    output = args.output.resolve()
    input_dir = output / "inputs"
    input_dir.mkdir(parents=True, exist_ok=True)
    selected_cases = tuple(item.strip() for item in args.cases.split(",") if item.strip())
    invalid = sorted(set(selected_cases).difference(CASES))
    if invalid:
        raise ValueError(f"unknown benchmark cases: {invalid}")
    config = {
        "created_at": datetime.now(UTC).isoformat(),
        "rows": args.rows,
        "columns": args.columns,
        "repeats": args.repeats,
        "cases": selected_cases,
        "seed": 0,
        "baseline": "oracle",
        "candidate": "streaming",
        "python": sys.version,
        "platform": sys.platform,
        "source_hashes": _source_fingerprint(),
    }
    records: list[dict[str, Any]] = []
    inputs: dict[str, dict[str, Any]] = {}
    for case in selected_cases:
        path = input_dir / f"{case}.xlsx"
        _build_case(path, case=case, rows=args.rows, columns=args.columns)
        inputs[case] = {"path": str(path), "sha256": _sha256(path), "bytes": path.stat().st_size}
        for loader_name in ("oracle", "streaming"):
            for repeat in range(args.repeats):
                metrics = _sample_worker(path, loader_name)
                record = {
                    "case": case,
                    "loader": loader_name,
                    "repeat": repeat,
                    **metrics,
                }
                records.append(record)
                with (output / "metrics.jsonl").open("a", encoding="utf-8") as stream:
                    stream.write(json.dumps(record, sort_keys=True) + "\n")
    summaries = _median_summary(records)
    comparisons = _comparisons(summaries)
    result = {
        "config": config,
        "inputs": inputs,
        "summaries": summaries,
        "comparisons": comparisons,
    }
    (output / "result.json").write_text(
        json.dumps(result, indent=2, sort_keys=True) + "\n",
        encoding="utf-8",
    )
    print(json.dumps(result, indent=2, sort_keys=True))
    if args.enforce and any(
        comparison["memory_reduction_percent"] < args.min_memory_reduction
        or comparison["elapsed_change_percent"] > args.max_elapsed_regression
        for comparison in comparisons
    ):
        return 1
    return 0


def _parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description=__doc__)
    subparsers = parser.add_subparsers(dest="command", required=True)
    worker = subparsers.add_parser("worker")
    worker.add_argument("--input", type=Path, required=True)
    worker.add_argument("--loader", choices=("oracle", "streaming"), required=True)
    run = subparsers.add_parser("run")
    run.add_argument("--output", type=Path, required=True)
    run.add_argument("--rows", type=int, default=20_000)
    run.add_argument("--columns", type=int, default=10)
    run.add_argument("--repeats", type=int, default=3)
    run.add_argument("--cases", default=",".join(CASES))
    run.add_argument("--enforce", action="store_true")
    run.add_argument("--min-memory-reduction", type=float, default=50.0)
    run.add_argument("--max-elapsed-regression", type=float, default=25.0)
    return parser


def main() -> int:
    args = _parser().parse_args()
    if args.command == "worker":
        return _worker(args.input, args.loader)
    if args.rows < 1 or args.columns < 2 or args.repeats < 1:
        raise ValueError("rows and repeats must be positive; columns must be at least 2")
    return _run(args)


if __name__ == "__main__":
    raise SystemExit(main())
