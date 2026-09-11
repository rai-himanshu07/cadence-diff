"""IO layer tests: snapshot loading across formats, encryption, read-only guarantee."""

import datetime as dt
import hashlib
import io
import shutil
import struct
import zipfile
from copy import copy
from pathlib import Path
from typing import Any
from xml.etree import ElementTree

import pytest
from openpyxl import Workbook
from openpyxl.utils.datetime import to_excel
from openpyxl.worksheet.table import Table
from pyxlsb import biff12

import qc_tool.io.loader as loader_module
from qc_tool.io.decrypt import (
    InvalidPasswordError,
    PasswordRequiredError,
    is_encrypted,
)
from qc_tool.io.formula_enrichment import ExtractedDefinedName, FormulaExtraction
from qc_tool.io.loader import (
    OOXMLWorkloadError,
    UnsupportedFormatError,
    XLSBWorkloadError,
    _assess_ooxml_workload,
    _assess_xlsb_workload,
    _load_ooxml_oracle,
    _load_ooxml_streaming,
    _parse_pivots,
    _parse_pivots_bounded,
    load_workbook_snapshot,
)
from qc_tool.io.model import display_cell_value, serialize_cell_value
from qc_tool.io.native_formula import native_formula_available
from qc_tool.io.native_kernel import native_kernel_available
from qc_tool.io.ooxml_worksheet import (
    OOXMLMetadataError,
    WorkbookMetadata,
    WorksheetMetadata,
    parse_ooxml_worksheet_metadata,
)
from qc_tool.io.opc import (
    InvalidOfficePackageError,
    UnsafeRelationshipTargetError,
    resolve_internal_relationship_target,
    validate_office_package,
)
from qc_tool.io.xlsb_formula import (
    XlsbFormulaScan,
    XlsbFormulaScanError,
    XlsbWorksheetMetrics,
)
from tests.fixtures.manifest_schema import FixtureManifest
from tests.fixtures.xlsb_writer import StyledCell, _record, write_xlsb


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _rewrite_package(path: Path, replacements: dict[str, bytes]) -> None:
    with zipfile.ZipFile(path) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    parts.update(replacements)
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, content in parts.items():
            archive.writestr(name, content)


def _add_shared_strings(path: Path) -> None:
    with zipfile.ZipFile(path) as archive:
        content_types = ElementTree.fromstring(archive.read("[Content_Types].xml"))
        relationships = ElementTree.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
    content_type_namespace = "http://schemas.openxmlformats.org/package/2006/content-types"
    relationship_namespace = "http://schemas.openxmlformats.org/package/2006/relationships"
    ElementTree.SubElement(
        content_types,
        f"{{{content_type_namespace}}}Override",
        PartName="/xl/sharedStrings.xml",
        ContentType=(
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"
        ),
    )
    ElementTree.SubElement(
        relationships,
        f"{{{relationship_namespace}}}Relationship",
        Id="rIdSharedStrings",
        Type=("http://schemas.openxmlformats.org/officeDocument/2006/relationships/sharedStrings"),
        Target="sharedStrings.xml",
    )
    worksheet = b"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">
  <dimension ref="A1:A5"/>
  <sheetData>
    <row r="1"><c r="A1" t="s"><v>0</v></c></row>
    <row r="2"><c r="A2" t="s"><v>1</v></c></row>
    <row r="3"><c r="A3" t="s"><v>2</v></c></row>
    <row r="4"><c r="A4" t="s"><v>0</v></c></row>
    <row r="5"><c r="A5" t="s"><v>3</v></c></row>
  </sheetData>
</worksheet>"""
    shared_strings = b"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"
     count="5" uniqueCount="4">
  <si><t>plain</t></si>
  <si><t/></si>
  <si><t xml:space="preserve"> spaced </t></si>
  <si><r><t>Rich</t></r><r><t xml:space="preserve"> text</t></r></si>
</sst>"""
    _rewrite_package(
        path,
        {
            "[Content_Types].xml": ElementTree.tostring(
                content_types,
                encoding="utf-8",
                xml_declaration=True,
            ),
            "xl/_rels/workbook.xml.rels": ElementTree.tostring(
                relationships,
                encoding="utf-8",
                xml_declaration=True,
            ),
            "xl/worksheets/sheet1.xml": worksheet,
            "xl/sharedStrings.xml": shared_strings,
        },
    )


def _streaming_snapshot(path: Path):
    return _load_ooxml_streaming(
        path.read_bytes(),
        source_name=path.name,
        file_format=path.suffix.lstrip("."),
    )


def _oracle_snapshot(path: Path):
    return _load_ooxml_oracle(
        path.read_bytes(),
        source_name=path.name,
        file_format=path.suffix.lstrip("."),
    )


def _workload_metadata(
    *,
    cell_count: int,
    max_row: int = 1,
    max_column: int = 1,
) -> WorkbookMetadata:
    return WorkbookMetadata(
        sheets=[
            WorksheetMetadata(
                name="Data",
                visibility="visible",
                part="xl/worksheets/sheet1.xml",
                max_row=max_row,
                max_column=max_column,
                cell_count=cell_count,
                xml_bytes=1_000,
            )
        ],
        shared_string_bytes=0,
        styles_bytes=1_000,
        style_count=2,
    )


def test_ooxml_workload_preflight_warns_without_truncating() -> None:
    workload = _assess_ooxml_workload(
        _workload_metadata(cell_count=1_000_000),
        source_name="large.xlsx",
        allow_large_workbook=False,
    )

    assert workload.degraded
    assert not workload.override_used
    assert "warning limit" in workload.detail


def test_ooxml_workload_preflight_refuses_and_requires_explicit_override() -> None:
    metadata = _workload_metadata(cell_count=5_000_000)

    with pytest.raises(OOXMLWorkloadError, match="allow_large_workbook"):
        _assess_ooxml_workload(
            metadata,
            source_name="too-large.xlsx",
            allow_large_workbook=False,
        )

    workload = _assess_ooxml_workload(
        metadata,
        source_name="too-large.xlsx",
        allow_large_workbook=True,
    )
    assert workload.degraded
    assert workload.override_used
    assert "override accepted" in workload.detail


def test_ooxml_workload_preflight_rejects_sparse_pathological_area() -> None:
    metadata = _workload_metadata(
        cell_count=2,
        max_row=1_048_576,
        max_column=16_384,
    )

    with pytest.raises(OOXMLWorkloadError, match="largest sheet area"):
        _assess_ooxml_workload(
            metadata,
            source_name="sparse.xlsx",
            allow_large_workbook=False,
        )


def test_xlsb_workload_preflight_warns_without_requiring_override() -> None:
    scan = XlsbFormulaScan(
        formula_cells={"Data": frozenset()},
        worksheet_metrics={
            "Data": XlsbWorksheetMetrics(
                cell_count=2_500_000,
                max_row=250_000,
                max_column=10,
                binary_bytes=100 * 1024 * 1024,
            )
        },
        sheet_count=1,
    )

    workload = _assess_xlsb_workload(
        scan,
        source_name="large.xlsb",
        allow_large_workbook=False,
    )

    assert workload.degraded
    assert not workload.override_used
    assert workload.warning_reasons == (
        "retained cells 2,500,000 >= warning limit 2,500,000",
    )


def test_sources_unmodified_by_loading(fixture_dir: Path, manifest: FixtureManifest) -> None:
    targets = {
        "baseline.xlsx": None,
        "current.xlsx": None,
        "baseline.xlsb": None,
        "current.xlsb": None,
        "current_encrypted.xlsx": manifest.password,
    }
    before = {name: _sha256(fixture_dir / name) for name in targets}
    for name, password in targets.items():
        load_workbook_snapshot(fixture_dir / name, password=password)
    after = {name: _sha256(fixture_dir / name) for name in targets}
    assert before == after


def test_xlsx_snapshot_content(fixture_dir: Path, manifest: FixtureManifest) -> None:
    snap = load_workbook_snapshot(fixture_dir / "current.xlsx")
    assert snap.file_format == "xlsx"
    assert snap.formulas_available and snap.styles_available
    assert snap.sheet_names == [
        "Long_Monthly",
        "Wide_Weekly",
        "Dashboard",
        "Summary",
        "New_Analysis",
        "Params",
    ]

    long_monthly = snap.sheet("Long_Monthly")
    assert long_monthly.max_row == 25

    e01 = manifest.defect("E01")
    cell = long_monthly.cell("C7")
    assert cell is not None and cell.value == float(e01.current or "")

    hardcoded = long_monthly.cell("E10")  # E02
    assert hardcoded is not None
    assert hardcoded.formula is None and isinstance(hardcoded.value, float)

    logic_changed = long_monthly.cell("E14")  # E03
    assert logic_changed is not None and logic_changed.formula == "=C14-D14*1.1"
    assert logic_changed.has_formula
    assert logic_changed.value is None  # openpyxl-authored file has no cached results

    assert long_monthly.cell("E25") is None  # E04: not extended

    summary = snap.sheet("Summary")
    error_const = summary.cell("C5")  # E05
    assert error_const is not None and error_const.is_error and error_const.value == "#REF!"
    error_formula = summary.cell("D5")  # E13
    assert error_formula is not None and error_formula.formula == "=SUM(#REF!)"

    assert snap.sheet("Params").visibility == "hidden"  # E11


def test_xlsx_snapshot_preserves_typed_dates(tmp_path: Path) -> None:
    path = tmp_path / "typed-date.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet["A1"] = dt.date(2026, 4, 1)
    sheet["A1"].number_format = "mmm-yy"
    workbook.save(path)

    snapshot = load_workbook_snapshot(path)

    cell = snapshot.sheet(sheet.title).cell("A1")
    assert cell is not None
    assert cell.value == dt.datetime(2026, 4, 1)
    assert display_cell_value(cell.value) == "2026-04-01T00:00:00"
    assert serialize_cell_value(cell.value) == "2026-04-01T00:00:00"


def test_xlsx_snapshot_preserves_formula_cached_dates(tmp_path: Path) -> None:
    path = tmp_path / "formula-date.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet["A1"] = "=DATE(2026,4,1)"
    sheet["A1"].number_format = "mmm-yy"
    workbook.save(path)
    with zipfile.ZipFile(path) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    worksheet_part = "xl/worksheets/sheet1.xml"
    xml = parts[worksheet_part].decode("utf-8")
    parts[worksheet_part] = xml.replace(
        "<v></v>",
        f"<v>{to_excel(dt.datetime(2026, 4, 1))}</v>",
    ).encode("utf-8")
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, data in parts.items():
            archive.writestr(name, data)

    snapshot = load_workbook_snapshot(path)

    cell = snapshot.sheet(sheet.title).cell("A1")
    assert cell is not None
    assert cell.formula == "=DATE(2026,4,1)"
    assert cell.value == dt.datetime(2026, 4, 1)


def test_xlsx_snapshot_ignores_false_small_dimension(tmp_path: Path) -> None:
    path = tmp_path / "false-dimension.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet["A1"] = "head"
    sheet["Z40"] = 42
    workbook.save(path)
    with zipfile.ZipFile(path) as archive:
        worksheet = archive.read("xl/worksheets/sheet1.xml")
    _rewrite_package(
        path,
        {
            "xl/worksheets/sheet1.xml": worksheet.replace(
                b'ref="A1:Z40"',
                b'ref="A1:A1"',
            )
        },
    )

    snapshot = load_workbook_snapshot(path)

    loaded = snapshot.sheet(sheet.title)
    assert (loaded.max_row, loaded.max_column) == (40, 26)
    tail = loaded.cell("Z40")
    assert tail is not None
    assert tail.value == 42
    raw = parse_ooxml_worksheet_metadata(path.read_bytes()).sheets[0]
    assert (raw.declared_max_row, raw.declared_max_column) == (1, 1)
    assert (raw.max_row, raw.max_column, raw.cell_count) == (40, 26, 2)
    assert _streaming_snapshot(path) == snapshot


def test_streaming_empty_sheet_bounds_match_oracle(tmp_path: Path) -> None:
    path = tmp_path / "empty.xlsx"
    Workbook().save(path)

    oracle = _oracle_snapshot(path)
    streaming = _streaming_snapshot(path)

    assert streaming == oracle
    assert (streaming.sheets[0].max_row, streaming.sheets[0].max_column) == (1, 1)


def test_xlsx_snapshot_derives_missing_dimension_and_hidden_state(tmp_path: Path) -> None:
    path = tmp_path / "missing-dimension.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet["D8"] = "tail"
    sheet.row_dimensions[8].hidden = True
    sheet.column_dimensions["C"].hidden = True
    workbook.save(path)
    with zipfile.ZipFile(path) as archive:
        root = ElementTree.fromstring(archive.read("xl/worksheets/sheet1.xml"))
    for child in list(root):
        if child.tag.rsplit("}", 1)[-1] == "dimension":
            root.remove(child)
    _rewrite_package(
        path,
        {
            "xl/worksheets/sheet1.xml": ElementTree.tostring(
                root,
                encoding="utf-8",
                xml_declaration=True,
            )
        },
    )

    snapshot = load_workbook_snapshot(path)

    loaded = snapshot.sheet(sheet.title)
    assert (loaded.max_row, loaded.max_column) == (8, 4)
    assert loaded.hidden_rows == frozenset({8})
    assert loaded.hidden_columns == frozenset({3})
    assert loaded.cell("D8") is not None
    raw = parse_ooxml_worksheet_metadata(path.read_bytes()).sheets[0]
    assert (raw.declared_max_row, raw.declared_max_column) == (None, None)
    assert (raw.max_row, raw.max_column) == (8, 4)
    assert raw.hidden_rows == frozenset({8})
    assert raw.hidden_columns == frozenset({3})
    assert _streaming_snapshot(path) == snapshot


def test_xlsx_snapshot_reads_genuine_shared_strings(tmp_path: Path) -> None:
    path = tmp_path / "shared-strings.xlsx"
    workbook = Workbook()
    workbook.save(path)
    _add_shared_strings(path)

    snapshot = load_workbook_snapshot(path)

    sheet = snapshot.sheets[0]
    values = []
    for row in range(1, 6):
        cell = sheet.cell(f"A{row}")
        assert cell is not None
        values.append(cell.value)
    assert values == [
        "plain",
        "",
        " spaced ",
        "plain",
        "Rich text",
    ]
    assert _streaming_snapshot(path) == snapshot


def test_xlsx_snapshot_preserves_shared_formulas_and_cached_values(
    tmp_path: Path,
) -> None:
    path = tmp_path / "shared-formulas.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    for row in range(1, 4):
        sheet.cell(row, 1, f"=ROW()+{row}")
    workbook.save(path)
    with zipfile.ZipFile(path) as archive:
        root = ElementTree.fromstring(archive.read("xl/worksheets/sheet1.xml"))
    cells = [element for element in root.iter() if element.tag.endswith("}c")]
    for index, cell in enumerate(cells):
        formula = next(element for element in cell if element.tag.endswith("}f"))
        value = next(element for element in cell if element.tag.endswith("}v"))
        formula.set("t", "shared")
        formula.set("si", "0")
        if index == 0:
            formula.set("ref", "A1:A3")
            formula.text = "ROW()+1"
        else:
            formula.text = None
        value.text = str(index + 2)
    _rewrite_package(
        path,
        {
            "xl/worksheets/sheet1.xml": ElementTree.tostring(
                root,
                encoding="utf-8",
                xml_declaration=True,
            )
        },
    )

    oracle = _oracle_snapshot(path)
    streaming = _streaming_snapshot(path)

    assert streaming == oracle
    assert [streaming.sheets[0].cells[(row, 1)].formula for row in range(1, 4)] == [
        "=ROW()+1",
        "=ROW()+1",
        "=ROW()+1",
    ]
    assert [streaming.sheets[0].cells[(row, 1)].value for row in range(1, 4)] == [
        2,
        3,
        4,
    ]


def test_xlsx_snapshot_preserves_array_formula_text_and_declared_range(
    tmp_path: Path,
) -> None:
    path = tmp_path / "array-formula.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet["A1"] = 1
    sheet["A2"] = 2
    sheet["B1"] = "=SUM(A1:A2)"
    workbook.save(path)
    with zipfile.ZipFile(path) as archive:
        root = ElementTree.fromstring(archive.read("xl/worksheets/sheet1.xml"))
    formula = next(element for element in root.iter() if element.tag.endswith("}f"))
    formula.set("t", "array")
    formula.set("ref", "B1:B2")
    formula.set("aca", "1")
    _rewrite_package(
        path,
        {
            "xl/worksheets/sheet1.xml": ElementTree.tostring(
                root,
                encoding="utf-8",
                xml_declaration=True,
            )
        },
    )

    oracle = _oracle_snapshot(path)
    streaming = _streaming_snapshot(path)

    assert streaming == oracle
    anchor = streaming.sheets[0].cell("B1")
    assert anchor is not None and anchor.formula == "=SUM(A1:A2)"
    assert len(streaming.formula_ranges) == 1
    descriptor = streaming.formula_ranges[0]
    assert descriptor.sheet == "Sheet"
    assert (descriptor.anchor_row, descriptor.anchor_column) == (1, 2)
    assert descriptor.cell_range == "B1:B2"
    assert descriptor.formula_type == "array"
    assert descriptor.always_calculate is True


def test_xlsx_snapshot_rejects_formula_range_with_wrong_anchor(tmp_path: Path) -> None:
    path = tmp_path / "bad-array-anchor.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet["B2"] = "=SUM(A1:A2)"
    workbook.save(path)
    with zipfile.ZipFile(path) as archive:
        root = ElementTree.fromstring(archive.read("xl/worksheets/sheet1.xml"))
    formula = next(element for element in root.iter() if element.tag.endswith("}f"))
    formula.set("t", "array")
    formula.set("ref", "A1:B2")
    _rewrite_package(
        path,
        {
            "xl/worksheets/sheet1.xml": ElementTree.tostring(
                root,
                encoding="utf-8",
                xml_declaration=True,
            )
        },
    )

    with pytest.raises(OOXMLMetadataError, match="not the top-left"):
        load_workbook_snapshot(path)


def test_xlsx_snapshot_preserves_typed_formula_cached_values(tmp_path: Path) -> None:
    path = tmp_path / "typed-formula-cache.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    formulas = {
        "A1": ("=1+1", "n", "2"),
        "B1": ("=TRUE()", "b", "1"),
        "C1": ('="ready"', "str", "ready"),
        "D1": ("=1/0", "e", "#DIV/0!"),
    }
    for coordinate, (formula, _data_type, _cached) in formulas.items():
        sheet[coordinate] = formula
    workbook.save(path)
    with zipfile.ZipFile(path) as archive:
        root = ElementTree.fromstring(archive.read("xl/worksheets/sheet1.xml"))
    for cell in root.iter():
        if not cell.tag.endswith("}c") or cell.get("r") not in formulas:
            continue
        _formula, data_type, cached = formulas[str(cell.get("r"))]
        cell.set("t", data_type)
        value = next(element for element in cell if element.tag.endswith("}v"))
        value.text = cached
    _rewrite_package(
        path,
        {
            "xl/worksheets/sheet1.xml": ElementTree.tostring(
                root,
                encoding="utf-8",
                xml_declaration=True,
            )
        },
    )

    oracle = _oracle_snapshot(path)
    streaming = _streaming_snapshot(path)

    assert streaming == oracle
    assert [streaming.sheets[0].cells[(1, column)].value for column in range(1, 5)] == [
        2,
        True,
        "ready",
        "#DIV/0!",
    ]


def test_xlsx_snapshot_preserves_shared_string_formula_cache(tmp_path: Path) -> None:
    path = tmp_path / "shared-string-formula-cache.xlsx"
    Workbook().save(path)
    _add_shared_strings(path)
    with zipfile.ZipFile(path) as archive:
        root = ElementTree.fromstring(archive.read("xl/worksheets/sheet1.xml"))
    cell = next(
        element
        for element in root.iter()
        if element.tag.endswith("}c") and element.get("r") == "A1"
    )
    formula = ElementTree.Element(cell.tag.rsplit("}", 1)[0] + "}f")
    formula.text = '"plain"'
    cell.insert(0, formula)
    _rewrite_package(
        path,
        {
            "xl/worksheets/sheet1.xml": ElementTree.tostring(
                root,
                encoding="utf-8",
                xml_declaration=True,
            )
        },
    )

    oracle = _oracle_snapshot(path)
    streaming = _streaming_snapshot(path)

    assert streaming == oracle
    cached = streaming.sheets[0].cells[(1, 1)]
    assert cached.formula == '="plain"'
    assert cached.value == "plain"


def test_streaming_interns_equal_cell_styles(tmp_path: Path) -> None:
    path = tmp_path / "styles.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet["A1"] = 1
    sheet["B1"] = 2
    for cell in (sheet["A1"], sheet["B1"]):
        cell.number_format = "0.00"
        font = copy(cell.font)
        font.bold = True
        cell.font = font
    workbook.save(path)

    snapshot = _streaming_snapshot(path)

    first = snapshot.sheets[0].cells[(1, 1)]
    second = snapshot.sheets[0].cells[(1, 2)]
    assert first.style_key is second.style_key
    assert first.number_format is second.number_format


def test_xlsx_snapshot_expands_hidden_column_spans(tmp_path: Path) -> None:
    path = tmp_path / "hidden-span.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet["A1"] = "head"
    sheet.column_dimensions.group("C", "E", hidden=True)
    workbook.save(path)

    snapshot = load_workbook_snapshot(path)

    assert snapshot.sheets[0].hidden_columns == frozenset({3, 4, 5})
    assert _streaming_snapshot(path) == snapshot


@pytest.mark.parametrize("metadata_kind", ["row", "column"])
def test_raw_metadata_errors_include_sheet_context(
    tmp_path: Path,
    metadata_kind: str,
) -> None:
    path = tmp_path / f"malformed-{metadata_kind}.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Data"
    sheet["C8"] = 1
    sheet.row_dimensions[8].hidden = True
    sheet.column_dimensions["C"].hidden = True
    workbook.save(path)
    with zipfile.ZipFile(path) as archive:
        root = ElementTree.fromstring(archive.read("xl/worksheets/sheet1.xml"))
    if metadata_kind == "row":
        row = next(
            element
            for element in root.iter()
            if element.tag.endswith("}row") and element.get("hidden") == "1"
        )
        row.set("r", "not-a-row")
    else:
        column = next(
            element
            for element in root.iter()
            if element.tag.endswith("}col") and element.get("hidden") == "1"
        )
        column.set("min", "not-a-column")
    _rewrite_package(
        path,
        {
            "xl/worksheets/sheet1.xml": ElementTree.tostring(
                root,
                encoding="utf-8",
                xml_declaration=True,
            )
        },
    )

    with pytest.raises(OOXMLMetadataError, match="Data: invalid hidden"):
        parse_ooxml_worksheet_metadata(path.read_bytes())


@pytest.mark.parametrize("name", ["baseline.xlsx", "current.xlsx"])
def test_streaming_xlsx_snapshot_matches_oracle(
    fixture_dir: Path,
    name: str,
) -> None:
    path = fixture_dir / name

    assert _streaming_snapshot(path) == _oracle_snapshot(path)


def test_streaming_closes_formula_workbook_before_cached_value_failure(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    package = io.BytesIO()
    Workbook().save(package)
    original_load_workbook = loader_module.load_workbook

    class TrackingWorkbook:
        def __init__(self, workbook: Any) -> None:
            self.workbook = workbook
            self.closed = False

        def __getattr__(self, name: str) -> Any:
            return getattr(self.workbook, name)

        def __getitem__(self, name: str) -> Any:
            return self.workbook[name]

        def close(self) -> None:
            self.closed = True
            self.workbook.close()

    formula_workbook: TrackingWorkbook | None = None
    def fake_load_workbook(*args: Any, **kwargs: Any) -> Any:
        nonlocal formula_workbook
        formula_workbook = TrackingWorkbook(original_load_workbook(*args, **kwargs))
        return formula_workbook

    def fail_cached_values(*args: Any, **kwargs: Any) -> None:
        assert formula_workbook is not None
        assert formula_workbook.closed
        raise ValueError("cached-value stream failed")

    monkeypatch.setattr(loader_module, "load_workbook", fake_load_workbook)
    monkeypatch.setattr(
        loader_module,
        "_merge_cached_formula_values",
        fail_cached_values,
    )

    with pytest.raises(ValueError, match="cached-value stream"):
        loader_module._load_ooxml_streaming(
            package.getvalue(),
            source_name="broken.xlsx",
            file_format="xlsx",
        )

    assert formula_workbook is not None
    assert formula_workbook.closed


def test_xlsx_snapshot_structure(fixture_dir: Path, manifest: FixtureManifest) -> None:
    base = load_workbook_snapshot(fixture_dir / "baseline.xlsx")
    curr = load_workbook_snapshot(fixture_dir / "current.xlsx")

    named_base = {n.name: n.target for n in base.named_ranges}
    named_curr = {n.name: n.target for n in curr.named_ranges}
    assert named_base["KPI_Margin"] == "Dashboard!$B$4"
    assert named_curr["KPI_Margin"] == "Dashboard!$B$5"  # E12

    e16 = manifest.defect("E16")
    chart_base = next(c for c in base.charts if c.sheet == "Dashboard")
    chart_curr = next(c for c in curr.charts if c.sheet == "Dashboard")
    assert chart_base.title == "Revenue trend (workbook)"
    assert chart_base.series[0].values_ref == e16.baseline
    assert chart_curr.series[0].values_ref == e16.current

    e17 = manifest.defect("E17")
    assert base.pivots[0].name == "RevenuePivot"
    assert base.pivots[0].source_sheet == "Long_Monthly"
    assert base.pivots[0].source_ref == e17.baseline
    assert curr.pivots[0].source_ref == e17.current

    dash_base = base.sheet("Dashboard")
    dash_curr = curr.sheet("Dashboard")
    b4_base, b4_curr = dash_base.cell("B4"), dash_curr.cell("B4")
    assert b4_base is not None and b4_base.number_format == "0.0%"
    assert b4_curr is not None and b4_curr.number_format == "0.00"  # E09
    a2_base, a2_curr = dash_base.cell("A2"), dash_curr.cell("A2")
    assert a2_base is not None and a2_curr is not None
    assert a2_base.style_key != a2_curr.style_key  # E18


def test_xlsx_snapshot_captures_table_schema(tmp_path: Path) -> None:
    path = tmp_path / "table.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Data"
    sheet.append(["Key", "Amount", "Cost"])
    sheet.append(["A", 10, 4])
    sheet.add_table(Table(displayName="MetricsTable", ref="A1:C2"))
    workbook.save(path)

    snapshot = load_workbook_snapshot(path)

    assert len(snapshot.tables) == 1
    table = snapshot.tables[0]
    assert (table.sheet, table.name, table.display_name) == (
        "Data",
        "MetricsTable",
        "MetricsTable",
    )
    assert table.cell_range == "A1:C2"
    assert table.columns == ["Key", "Amount", "Cost"]
    assert table.header_row_count == 1
    assert table.totals_row_count == 0
    raw = parse_ooxml_worksheet_metadata(path.read_bytes())
    assert raw.sheets[0].tables == snapshot.tables
    assert _streaming_snapshot(path) == snapshot


def test_ooxml_worksheet_relationship_cannot_escape_package(tmp_path: Path) -> None:
    path = tmp_path / "unsafe-relationship.xlsx"
    workbook = Workbook()
    workbook.save(path)
    with zipfile.ZipFile(path) as archive:
        relationships = ElementTree.fromstring(
            archive.read("xl/_rels/workbook.xml.rels")
        )
    for relationship in relationships:
        if (relationship.get("Type") or "").endswith("/worksheet"):
            relationship.set("Target", "../../outside.xml")
            break
    _rewrite_package(
        path,
        {
            "xl/_rels/workbook.xml.rels": ElementTree.tostring(
                relationships,
                encoding="utf-8",
                xml_declaration=True,
            )
        },
    )

    with pytest.raises(OOXMLMetadataError, match="unsafe relationship target"):
        parse_ooxml_worksheet_metadata(path.read_bytes())


def test_opc_relationship_resolution_keeps_valid_parent_targets() -> None:
    assert resolve_internal_relationship_target(
        "ppt/slides/slide1.xml",
        "../slideLayouts/slideLayout1.xml",
    ) == "ppt/slideLayouts/slideLayout1.xml"
    with pytest.raises(UnsafeRelationshipTargetError):
        resolve_internal_relationship_target("xl/workbook.xml", "../../outside.xml")
    with pytest.raises(UnsafeRelationshipTargetError):
        resolve_internal_relationship_target("xl/workbook.xml", "https://example.invalid/a")
    with pytest.raises(UnsafeRelationshipTargetError):
        resolve_internal_relationship_target("xl/workbook.xml", "C:/outside.xml")


def test_office_package_precheck_rejects_escaping_ppt_relationship() -> None:
    payload = io.BytesIO()
    with zipfile.ZipFile(payload, "w") as archive:
        archive.writestr("[Content_Types].xml", b"<Types/>")
        archive.writestr("_rels/.rels", b"<Relationships/>")
        archive.writestr("ppt/presentation.xml", b"<presentation/>")
        archive.writestr(
            "ppt/slides/_rels/slide1.xml.rels",
            b'<Relationships><Relationship Id="rId1" '
            b'Target="../../../outside.xml"/></Relationships>',
        )

    with pytest.raises(
        InvalidOfficePackageError,
        match=r"deck\.pptx: invalid Office package relationship target",
    ):
        validate_office_package(payload.getvalue(), source_name="deck.pptx")


def test_office_package_precheck_rejects_malformed_relationship_metadata() -> None:
    payload = io.BytesIO()
    with zipfile.ZipFile(payload, "w") as archive:
        archive.writestr("[Content_Types].xml", b"<Types/>")
        archive.writestr("_rels/.rels", b"<Relationships/>")
        archive.writestr("xl/workbook.xml", b"<workbook/>")
        archive.writestr("xl/_rels/workbook.xml.rels", b"<not-xml")

    with pytest.raises(
        InvalidOfficePackageError,
        match=r"book\.xlsx: invalid Office package relationship metadata",
    ):
        validate_office_package(payload.getvalue(), source_name="book.xlsx")


def test_pivot_relationship_failure_keeps_filename_at_loader_boundary() -> None:
    payload = io.BytesIO()
    with zipfile.ZipFile(payload, "w") as archive:
        archive.writestr("xl/workbook.xml", b"<workbook/>")
        archive.writestr(
            "xl/_rels/workbook.xml.rels",
            b'<Relationships><Relationship Id="rId1" '
            b'Target="../../../outside.xml"/></Relationships>',
        )

    with pytest.raises(
        InvalidOfficePackageError,
        match=r"book\.xlsx: invalid Office package relationship target",
    ):
        _parse_pivots_bounded(payload.getvalue(), "book.xlsx")


def test_pivots_resolve_shared_and_non_positional_caches() -> None:
    workbook = b"""<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"
        xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">
        <pivotCaches><pivotCache cacheId="1" r:id="rId10"/>
        <pivotCache cacheId="7" r:id="rId11"/></pivotCaches></workbook>"""
    relationships = b"""<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
        <Relationship Id="rId10" Target="pivotCache/pivotCacheDefinition2.xml"/>
        <Relationship Id="rId11" Target="pivotCache/pivotCacheDefinition10.xml"/>
        </Relationships>"""

    def table(name: str, cache_id: int) -> bytes:
        return (
            '<pivotTableDefinition xmlns="http://schemas.openxmlformats.org/'
            f'spreadsheetml/2006/main" name="{name}" cacheId="{cache_id}">'
            '<location ref="A1:B2"/></pivotTableDefinition>'
        ).encode()

    def cache(sheet: str, ref: str) -> bytes:
        return (
            '<pivotCacheDefinition xmlns="http://schemas.openxmlformats.org/'
            'spreadsheetml/2006/main"><cacheSource><worksheetSource '
            f'sheet="{sheet}" ref="{ref}"/></cacheSource></pivotCacheDefinition>'
        ).encode()

    stream = io.BytesIO()
    with zipfile.ZipFile(stream, "w") as archive:
        archive.writestr("xl/workbook.xml", workbook)
        archive.writestr("xl/_rels/workbook.xml.rels", relationships)
        archive.writestr("xl/pivotTables/pivotTable1.xml", table("Other", 7))
        archive.writestr("xl/pivotTables/pivotTable2.xml", table("Shared A", 1))
        archive.writestr("xl/pivotTables/pivotTable10.xml", table("Shared B", 1))
        archive.writestr("xl/pivotCache/pivotCacheDefinition2.xml", cache("Data", "A1:B20"))
        archive.writestr("xl/pivotCache/pivotCacheDefinition10.xml", cache("OtherData", "D1:E8"))

    pivots = {pivot.name: pivot for pivot in _parse_pivots(stream.getvalue())}

    assert (pivots["Shared A"].source_sheet, pivots["Shared A"].source_ref) == (
        "Data",
        "A1:B20",
    )
    assert (pivots["Shared B"].source_sheet, pivots["Shared B"].source_ref) == (
        "Data",
        "A1:B20",
    )
    assert (pivots["Other"].source_sheet, pivots["Other"].source_ref) == (
        "OtherData",
        "D1:E8",
    )


def test_xlsm_loads_identically(fixture_dir: Path, tmp_path: Path) -> None:
    source = fixture_dir / "current.xlsx"
    xlsm = tmp_path / "current.xlsm"
    shutil.copyfile(source, xlsm)

    snap_xlsx = load_workbook_snapshot(source)
    snap_xlsm = load_workbook_snapshot(xlsm)
    assert snap_xlsm.file_format == "xlsm"
    assert snap_xlsm.sheet_names == snap_xlsx.sheet_names
    assert snap_xlsm.sheet("Long_Monthly").cells == snap_xlsx.sheet("Long_Monthly").cells


def test_xlsb_snapshot(fixture_dir: Path, manifest: FixtureManifest) -> None:
    snap = load_workbook_snapshot(fixture_dir / "current.xlsb")
    assert snap.file_format == "xlsb"
    assert not snap.formulas_available and not snap.styles_available
    assert snap.formula_presence_available
    # No xl/styles.bin in this fixture: the scan succeeds but finds nothing to
    # type, so every cell keeps its original cached value unchanged.
    assert snap.number_formats_available is True
    assert snap.sheet_names == ["Long_Monthly"]

    sheet = snap.sheet("Long_Monthly")
    assert sheet.max_row == 25
    header = sheet.cell("A1")
    assert header is not None and header.value == "Period"

    xb01 = manifest.defect("XB01")
    cell = sheet.cell("C7")
    assert cell is not None and cell.value == float(xb01.current or "")
    assert cell.formula is None and cell.style_key is None
    assert cell.is_formula is False and not cell.has_formula
    assert snap.workload.metrics_available
    assert snap.workload.format == "xlsb"
    assert snap.workload.cell_count == len(sheet.cells)
    assert "BIFF12" in snap.workload.detail


@pytest.mark.skipif(
    not native_kernel_available(),
    reason="native/xlsbkernel/ not built in this environment (optional accelerator)",
)
def test_xlsb_native_values_engine_matches_pyxlsb(fixture_dir: Path) -> None:
    """Criterion 10 (plan-20260906): the native kernel's values path must be
    indistinguishable from pyxlsb's for every cell -- type, value AND
    number_format, on a real generated fixture (not just the isolated Rust
    unit-level probe in artifacts/kernel-experiments-20260906/).
    """
    path = fixture_dir / "current.xlsb"
    default_snap = load_workbook_snapshot(path)
    native_snap = load_workbook_snapshot(path, _xlsb_values_engine="native")

    assert native_snap.values_source is not None
    assert native_snap.values_source.startswith("native-biff12:")

    assert native_snap.sheet_names == default_snap.sheet_names
    for sheet_name in default_snap.sheet_names:
        default_sheet = default_snap.sheet(sheet_name)
        native_sheet = native_snap.sheet(sheet_name)
        assert native_sheet.max_row == default_sheet.max_row
        assert native_sheet.max_column == default_sheet.max_column
        assert set(native_sheet.cells) == set(default_sheet.cells)
        for key, default_cell in default_sheet.cells.items():
            native_cell = native_sheet.cells[key]
            assert type(native_cell.value) is type(default_cell.value)
            assert native_cell.value == default_cell.value
            assert native_cell.number_format == default_cell.number_format
            assert native_cell.is_formula == default_cell.is_formula


def test_xlsb_native_values_engine_raises_a_clear_error_when_unavailable(
    monkeypatch: pytest.MonkeyPatch, fixture_dir: Path
) -> None:
    """Deterministic regardless of whether this environment actually has the
    extension built -- forces the unavailable path so the fallback contract
    (never silently guess) is provable without depending on local state.
    """
    import qc_tool.io.native_kernel as native_kernel_module

    monkeypatch.setattr(native_kernel_module, "_xlsbkernel", None)
    with pytest.raises(RuntimeError, match="native xlsbkernel"):
        load_workbook_snapshot(fixture_dir / "current.xlsb", _xlsb_values_engine="native")


@pytest.mark.skipif(
    not native_kernel_available(),
    reason="native/xlsbkernel/ not built in this environment (optional accelerator)",
)
def test_xlsb_values_engine_auto_prefers_native_when_available(fixture_dir: Path) -> None:
    """plan-20260908-phase-b-guest-performance-followup.md: `"auto"` must
    resolve identically to an explicit `"native"` request whenever the kernel
    extension is importable, without changing the shipped `"pyxlsb"` default.
    """
    path = fixture_dir / "current.xlsb"
    auto_snap = load_workbook_snapshot(path, _xlsb_values_engine="auto")
    native_snap = load_workbook_snapshot(path, _xlsb_values_engine="native")

    assert auto_snap.sheet_names == native_snap.sheet_names
    for sheet_name in native_snap.sheet_names:
        auto_sheet = auto_snap.sheet(sheet_name)
        native_sheet = native_snap.sheet(sheet_name)
        assert set(auto_sheet.cells) == set(native_sheet.cells)
        for key, native_cell in native_sheet.cells.items():
            auto_cell = auto_sheet.cells[key]
            assert type(auto_cell.value) is type(native_cell.value)
            assert auto_cell.value == native_cell.value
            assert auto_cell.number_format == native_cell.number_format


def test_xlsb_values_engine_auto_degrades_to_pyxlsb_when_unavailable(
    monkeypatch: pytest.MonkeyPatch, fixture_dir: Path
) -> None:
    """Unlike an explicit `"native"` request (which raises hard when the
    extension is missing, above), `"auto"` must silently degrade to
    `"pyxlsb"` -- deterministic regardless of whether this environment
    actually has the extension built.
    """
    import qc_tool.io.native_kernel as native_kernel_module

    monkeypatch.setattr(native_kernel_module, "_xlsbkernel", None)
    path = fixture_dir / "current.xlsb"

    auto_snap = load_workbook_snapshot(path, _xlsb_values_engine="auto")
    pyxlsb_snap = load_workbook_snapshot(path, _xlsb_values_engine="pyxlsb")

    assert auto_snap.values_source is not None
    assert auto_snap.values_source.startswith("pyxlsb:")

    assert auto_snap.sheet_names == pyxlsb_snap.sheet_names
    for sheet_name in pyxlsb_snap.sheet_names:
        assert set(auto_snap.sheet(sheet_name).cells) == set(
            pyxlsb_snap.sheet(sheet_name).cells
        )


def test_xlsb_values_engine_auto_falls_back_to_pyxlsb_on_a_runtime_failure(
    monkeypatch: pytest.MonkeyPatch, fixture_dir: Path
) -> None:
    """Criterion 19: `auto` must degrade to pyxlsb not only when the kernel
    is absent at resolution time (above) but also when it is present yet
    fails at runtime (a decode bug, or a converted Rust panic) -- the
    fallback must be indistinguishable in outcome from the kernel never
    having been installed, with a fixed, content-free disclosure.
    """
    import qc_tool.io.native_kernel as native_kernel_module

    def _broken_report(data: bytes):
        raise RuntimeError("synthetic native values decode failure")

    monkeypatch.setattr(native_kernel_module, "native_kernel_available", lambda: True)
    monkeypatch.setattr(native_kernel_module, "raw_values_report", _broken_report)
    path = fixture_dir / "current.xlsb"

    auto_snap = load_workbook_snapshot(path, _xlsb_values_engine="auto")
    pyxlsb_snap = load_workbook_snapshot(path, _xlsb_values_engine="pyxlsb")

    assert auto_snap.sheet_names == pyxlsb_snap.sheet_names
    for sheet_name in pyxlsb_snap.sheet_names:
        assert set(auto_snap.sheet(sheet_name).cells) == set(
            pyxlsb_snap.sheet(sheet_name).cells
        )
    assert auto_snap.values_engine_fallback_detail != ""
    assert "RuntimeError" in auto_snap.values_engine_fallback_detail
    assert auto_snap.values_source is not None
    assert auto_snap.values_source.startswith("pyxlsb:")
    assert pyxlsb_snap.values_engine_fallback_detail == ""


def test_xlsb_values_engine_explicit_native_raises_instead_of_falling_back(
    monkeypatch: pytest.MonkeyPatch, fixture_dir: Path
) -> None:
    """Unlike `auto` (above), an explicit `\"native\"` request means the
    caller wants this exact engine or nothing -- a runtime failure must
    fail closed, never silently substitute pyxlsb.
    """
    import qc_tool.io.native_kernel as native_kernel_module

    def _broken_report(data: bytes):
        raise RuntimeError("synthetic native values decode failure")

    monkeypatch.setattr(native_kernel_module, "native_kernel_available", lambda: True)
    monkeypatch.setattr(native_kernel_module, "raw_values_report", _broken_report)

    with pytest.raises(RuntimeError, match="native"):
        load_workbook_snapshot(
            fixture_dir / "current.xlsb", _xlsb_values_engine="native"
        )


# --- B3: XLSB formula-engine selection --------------------------------------


@pytest.mark.skipif(
    not native_formula_available(),
    reason="native/xlsbkernel/ not built in this environment (optional accelerator)",
)
def test_xlsb_native_formula_engine_dispatches_through_load_workbook_snapshot(
    tmp_path: Path,
) -> None:
    """Wiring proof, not a rendering-correctness proof (that is B0/B1/B2's
    real-representative large-workbook-file evidence): a real load with ``formula_engine="native"``
    must reach the native adapter, set `formula_source` to a
    ``native-biff12:`` engine string, and never raise -- even for a cell the
    kernel could not decode (the synthetic writer's formula records carry no
    real Ptg token stream, so this exercises the graceful partial-coverage
    path, not full-text parity).
    """
    path = tmp_path / "native-engine.xlsb"
    write_xlsb(path, {"Data": [[StyledCell(1.0), StyledCell(2.0, is_formula=True)]]})

    snapshot = load_workbook_snapshot(path, formula_engine="native")

    assert snapshot.formula_source is not None
    assert snapshot.formula_source.startswith("native-biff12:")


def test_xlsb_formula_engine_auto_prefers_native_when_available(tmp_path: Path) -> None:
    path = tmp_path / "auto-engine.xlsb"
    write_xlsb(path, {"Data": [[StyledCell(1.0)]]})  # no formulas: engine never runs

    snapshot = load_workbook_snapshot(path, formula_engine="auto")

    # No formula cells at all -- engine choice is moot, but must never raise.
    assert snapshot.formula_presence_available


def test_xlsb_formula_engine_explicit_wrong_platform_degrades_instead_of_raising(
    tmp_path: Path,
) -> None:
    """An engine this platform cannot run (e.g. Excel-COM on Linux) is a
    formula-adapter failure like any other -- it degrades to
    formula-presence-only checks, it never fails the whole load.
    """
    path = tmp_path / "wrong-platform.xlsb"
    write_xlsb(path, {"Data": [[StyledCell(1.0), StyledCell(2.0, is_formula=True)]]})

    snapshot = load_workbook_snapshot(path, formula_engine="excel")

    assert snapshot.formula_source is None
    assert not snapshot.formulas_available
    assert "Formula presence checked" in (snapshot.formula_detail or "")


def test_xlsb_native_formula_engine_unavailable_degrades_gracefully(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    """Unlike `_xlsb_values_engine` (values are required, so an explicit
    ``native`` request that cannot be honoured raises hard), formula TEXT is
    an established best-effort enrichment: an explicit ``formula_engine=
    "native"`` request the environment cannot satisfy must degrade the same
    way a missing Excel/LibreOffice adapter already does, never crash the run.
    """
    import qc_tool.io.native_formula as native_formula_module

    monkeypatch.setattr(native_formula_module, "_xlsbkernel", None)
    path = tmp_path / "native-unavailable.xlsb"
    write_xlsb(path, {"Data": [[StyledCell(1.0), StyledCell(2.0, is_formula=True)]]})

    snapshot = load_workbook_snapshot(path, formula_engine="native")

    assert snapshot.formula_source is None
    assert not snapshot.formulas_available
    assert "xlsbkernel" in (snapshot.formula_detail or "")


@pytest.mark.skipif(
    not native_formula_available(),
    reason="native/xlsbkernel/ not built in this environment (optional accelerator)",
)
def test_formula_cache_keys_differ_by_resolved_engine(tmp_path: Path) -> None:
    """Switching `formula_engine` must never reuse a cached entry produced by
    a different engine -- a stale cross-engine hit would silently swap which
    adapter's text a run actually uses. ``native`` and ``libreoffice`` both
    run for real on this dev box, so both genuinely populate the cache.
    """
    from qc_tool.io.formula_cache import FormulaExtractionCache

    path = tmp_path / "cache-key.xlsb"
    write_xlsb(path, {"Data": [[StyledCell(1.0), StyledCell(2.0, is_formula=True)]]})
    cache = FormulaExtractionCache(tmp_path / "formula-cache")

    load_workbook_snapshot(path, formula_engine="native", formula_cache=cache)
    assert cache.status()["entry_count"] == 1
    load_workbook_snapshot(path, formula_engine="libreoffice", formula_cache=cache)
    assert cache.status()["entry_count"] == 2  # a distinct entry, not a stale hit

    # Re-running either engine again is a cache hit, not a third entry.
    load_workbook_snapshot(path, formula_engine="native", formula_cache=cache)
    load_workbook_snapshot(path, formula_engine="libreoffice", formula_cache=cache)
    assert cache.status()["entry_count"] == 2


def test_xlsb_native_compat_mode_restricts_to_legacy_engine_coordinates(
    tmp_path: Path,
) -> None:
    """Private, oracle-only switch (plan Criterion 13(a)): combined with
    ``formula_engine="native"`` it must not raise regardless of whether the
    kernel or the legacy engine actually decoded any text for this fixture.
    """
    path = tmp_path / "compat-mode.xlsb"
    write_xlsb(path, {"Data": [[StyledCell(1.0), StyledCell(2.0, is_formula=True)]]})

    snapshot = load_workbook_snapshot(
        path, formula_engine="native", _native_formula_compat_mode=True
    )

    assert snapshot.formula_presence_available


def test_xlsb_typed_formula_backed_date_beside_an_untouched_percent_row(
    tmp_path: Path,
) -> None:
    """Client-cited scenario: a formula-cached numeric serial with a custom
    date-formatted XF must become a real date, while an adjacent
    percentage-formatted row stays a plain float."""
    path = tmp_path / "dates.xlsb"
    write_xlsb(
        path,
        {
            "Data": [
                [StyledCell(0.4557, xf_index=1)],  # A1: percent row
                [StyledCell(46023.0, xf_index=2, is_formula=True)],  # A2: date row
            ]
        },
        cell_xfs=[0, 9, 165],  # 0=General, 9=builtin 0%, 165=custom
        custom_formats={165: "mmm-yy"},
    )

    snap = load_workbook_snapshot(path)

    assert snap.number_formats_available is True
    sheet = snap.sheet("Data")
    percent_cell = sheet.cell("A1")
    assert percent_cell is not None
    assert percent_cell.value == 0.4557  # untouched: not a date format
    assert percent_cell.number_format == "0%"

    date_cell = sheet.cell("A2")
    assert date_cell is not None
    assert date_cell.number_format == "mmm-yy"
    assert date_cell.is_formula is True
    assert date_cell.value == dt.datetime(2026, 1, 1)  # excel serial 46023, 1900 epoch

    if native_kernel_available():
        # Same fixture (percent float + formula-cached date via a custom XF)
        # through the native values engine must agree exactly -- date
        # conversion is unchanged shared code either way, but this proves it
        # end to end rather than by inspection alone.
        native_snap = load_workbook_snapshot(path, _xlsb_values_engine="native")
        native_sheet = native_snap.sheet("Data")
        native_percent = native_sheet.cell("A1")
        assert native_percent is not None
        assert native_percent.value == percent_cell.value
        assert native_percent.number_format == percent_cell.number_format
        native_date = native_sheet.cell("A2")
        assert native_date is not None
        assert native_date.value == date_cell.value
        assert native_date.number_format == date_cell.number_format
        assert native_date.is_formula == date_cell.is_formula


def test_xlsb_1904_epoch_shifts_the_typed_date(tmp_path: Path) -> None:
    path = tmp_path / "dates-1904.xlsb"
    write_xlsb(
        path,
        {"Data": [[StyledCell(46023.0, xf_index=1)]]},
        date1904=True,
        cell_xfs=[0, 14],  # builtin mm-dd-yy
    )

    snap = load_workbook_snapshot(path)

    cell = snap.sheet("Data").cell("A1")
    assert cell is not None
    assert isinstance(cell.value, dt.datetime)
    assert cell.value != dt.datetime(2026, 1, 1)  # would be the 1900-epoch answer

    if native_kernel_available():
        native_cell = (
            load_workbook_snapshot(path, _xlsb_values_engine="native").sheet("Data").cell("A1")
        )
        assert native_cell is not None
        assert native_cell.value == cell.value


def test_xlsb_malformed_styles_degrades_coverage_without_changing_values(
    tmp_path: Path,
) -> None:
    path = tmp_path / "malformed-styles.xlsb"
    write_xlsb(
        path,
        {"Data": [[StyledCell(46023.0, xf_index=1)]]},
        cell_xfs=[0, 165],
        custom_formats={165: "mmm-yy"},
    )
    with zipfile.ZipFile(path) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    parts["xl/styles.bin"] = parts["xl/styles.bin"][:-1]  # truncate
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, content in parts.items():
            archive.writestr(name, content)

    snap = load_workbook_snapshot(path)

    assert snap.number_formats_available is False
    assert "Number-format scan failed" in snap.number_format_detail
    cell = snap.sheet("Data").cell("A1")
    assert cell is not None
    assert cell.value == 46023.0  # cached value is never guessed or dropped
    assert cell.number_format is None


def test_xlsb_scan_failure_preserves_cached_values_and_marks_metrics_unavailable(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    path = tmp_path / "degraded.xlsb"
    write_xlsb(path, {"Data": [["Label", 42.0]]})

    def reject_scan(_data: bytes) -> XlsbFormulaScan:
        raise XlsbFormulaScanError("synthetic structural scan failure")

    monkeypatch.setattr(loader_module, "scan_xlsb_formulas", reject_scan)

    snapshot = load_workbook_snapshot(path)
    cell = snapshot.sheet("Data").cell("B1")

    assert cell is not None and cell.value == 42.0
    assert not snapshot.formula_presence_available
    assert not snapshot.workload.metrics_available
    assert snapshot.workload.detail == "XLSB workload metrics unavailable"


def test_xlsb_workload_refuses_before_materialization_and_override_is_explicit(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    path = tmp_path / "large.xlsb"
    write_xlsb(path, {"Data": [[1.0]]})
    scan = XlsbFormulaScan(
        formula_cells={"Data": frozenset()},
        worksheet_metrics={
            "Data": XlsbWorksheetMetrics(
                cell_count=5_000_000,
                max_row=500_000,
                max_column=10,
                binary_bytes=200 * 1024 * 1024,
            )
        },
        sheet_count=1,
    )
    monkeypatch.setattr(loader_module, "scan_xlsb_formulas", lambda _data: scan)
    real_open_xlsb = loader_module.open_xlsb
    materialized = False

    def tracked_open_xlsb(*args: Any, **kwargs: Any) -> Any:
        nonlocal materialized
        materialized = True
        return real_open_xlsb(*args, **kwargs)

    monkeypatch.setattr(loader_module, "open_xlsb", tracked_open_xlsb)

    with pytest.raises(XLSBWorkloadError, match="retained cells"):
        load_workbook_snapshot(path)

    assert not materialized
    snapshot = load_workbook_snapshot(path, allow_large_workbook=True)

    assert materialized
    assert snapshot.workload.override_used
    assert snapshot.workload.cell_count == 5_000_000
    assert "override accepted" in snapshot.workload.detail


def test_xlsb_enrichment_preserves_original_cached_value(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    path = tmp_path / "formula.xlsb"
    write_xlsb(path, {"Data": [[1.0]]})
    with zipfile.ZipFile(path) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    sheet = bytearray()
    sheet += _record(biff12.WORKSHEET)
    sheet += _record(biff12.DIMENSION, struct.pack("<IIII", 0, 0, 0, 0))
    sheet += _record(biff12.SHEETDATA)
    sheet += _record(biff12.ROW, (0).to_bytes(4, "little"))
    formula_payload = (
        (0).to_bytes(4, "little")
        + (0).to_bytes(4, "little")
        + struct.pack("<d", 42.0)
        + b"\x00\x00"
        + (0).to_bytes(4, "little")
    )
    sheet += _record(biff12.FORMULA_FLOAT, formula_payload)
    sheet += _record(biff12.SHEETDATA_END)
    sheet += _record(biff12.WORKSHEET_END)
    parts["xl/worksheets/sheet1.bin"] = bytes(sheet)
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, data in parts.items():
            archive.writestr(name, data)

    def fake_extract(data: bytes, scan: object, **_kwargs: object) -> FormulaExtraction:
        return FormulaExtraction(
            formulas={"Data": {(1, 1): "=40+2"}},
            engine="libreoffice:test",
            detail="test formula adapter",
        )

    monkeypatch.setattr(loader_module, "_extract_xlsb_formulas", fake_extract)

    snapshot = load_workbook_snapshot(path)

    cell = snapshot.sheet("Data").cells[(1, 1)]
    assert cell.value == 42.0
    assert cell.formula == "=40+2" and cell.has_formula
    assert snapshot.formulas_available
    assert snapshot.formula_source == "libreoffice:test"


def _two_formula_cell_sheet() -> bytes:
    sheet = bytearray()
    sheet += _record(biff12.WORKSHEET)
    sheet += _record(biff12.DIMENSION, struct.pack("<IIII", 0, 1, 0, 0))
    sheet += _record(biff12.SHEETDATA)
    for row in (0, 1):
        sheet += _record(biff12.ROW, row.to_bytes(4, "little"))
        formula_payload = (
            (0).to_bytes(4, "little")
            + (0).to_bytes(4, "little")
            + struct.pack("<d", 42.0 + row)
            + b"\x00\x00"
            + (0).to_bytes(4, "little")
        )
        sheet += _record(biff12.FORMULA_FLOAT, formula_payload)
    sheet += _record(biff12.SHEETDATA_END)
    sheet += _record(biff12.WORKSHEET_END)
    return bytes(sheet)


def test_xlsb_partial_enrichment_merges_available_text_without_raising(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    path = tmp_path / "partial.xlsb"
    write_xlsb(path, {"Data": [[1.0]]})
    with zipfile.ZipFile(path) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    parts["xl/worksheets/sheet1.bin"] = _two_formula_cell_sheet()
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, data in parts.items():
            archive.writestr(name, data)

    def fake_extract(data: bytes, scan: object, **_kwargs: object) -> FormulaExtraction:
        # Only the first of two scanned formula coordinates is returned.
        return FormulaExtraction(
            formulas={"Data": {(1, 1): "=40+2"}},
            engine="libreoffice:test",
            detail="test formula adapter",
        )

    monkeypatch.setattr(loader_module, "_extract_xlsb_formulas", fake_extract)

    snapshot = load_workbook_snapshot(path)

    merged = snapshot.sheet("Data").cells[(1, 1)]
    missing = snapshot.sheet("Data").cells[(2, 1)]
    assert merged.formula == "=40+2" and merged.value == 42.0
    assert missing.formula is None and missing.is_formula and missing.value == 43.0
    assert not snapshot.formulas_available  # partial, not complete
    coverage = snapshot.formula_text_coverage
    assert coverage.state == "partial"
    assert coverage.expected_count == 2
    assert coverage.merged_count == 1
    assert coverage.missing_count == 1


def test_xlsb_reachability_proven_inactive_for_unreferenced_passive_link(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    path = tmp_path / "passive-link.xlsb"
    write_xlsb(path, {"Data": [[1.0]]})
    with zipfile.ZipFile(path) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    sheet = bytearray()
    sheet += _record(biff12.WORKSHEET)
    sheet += _record(biff12.DIMENSION, struct.pack("<IIII", 0, 0, 0, 0))
    sheet += _record(biff12.SHEETDATA)
    sheet += _record(biff12.ROW, (0).to_bytes(4, "little"))
    formula_payload = (
        (0).to_bytes(4, "little")
        + (0).to_bytes(4, "little")
        + struct.pack("<d", 42.0)
        + b"\x00\x00"
        + (0).to_bytes(4, "little")
    )
    sheet += _record(biff12.FORMULA_FLOAT, formula_payload)
    sheet += _record(biff12.SHEETDATA_END)
    sheet += _record(biff12.WORKSHEET_END)
    parts["xl/worksheets/sheet1.bin"] = bytes(sheet)
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, data in parts.items():
            archive.writestr(name, data)
    # Appended after the rewrite so it is never part of the replaced parts dict.
    with zipfile.ZipFile(path, "a") as archive:
        archive.writestr("xl/externalLinks/externalLink1.xml", b"<externalLink/>")

    def fake_extract(data: bytes, scan: object, **_kwargs: object) -> FormulaExtraction:
        return FormulaExtraction(
            formulas={"Data": {(1, 1): "=40+2"}},
            engine="libreoffice:test",
            detail="test formula adapter",
            defined_names=(
                ExtractedDefinedName(name="UnrelatedInternal", target="Data!$A$1"),
            ),
            defined_names_complete=True,
        )

    monkeypatch.setattr(loader_module, "_extract_xlsb_formulas", fake_extract)

    snapshot = load_workbook_snapshot(path)

    assert snapshot.formulas_available
    reachability = snapshot.external_link_reachability
    assert reachability is not None
    assert reachability.proven is True
    assert reachability.live is False


def test_encrypted_workbook(fixture_dir: Path, manifest: FixtureManifest) -> None:
    encrypted = fixture_dir / "current_encrypted.xlsx"
    assert is_encrypted(encrypted)
    assert not is_encrypted(fixture_dir / "current.xlsx")

    snap = load_workbook_snapshot(encrypted, password=manifest.password)
    plain = load_workbook_snapshot(fixture_dir / "current.xlsx")
    assert snap.sheet_names == plain.sheet_names
    assert snap.sheet("Long_Monthly").cells == plain.sheet("Long_Monthly").cells

    with pytest.raises(PasswordRequiredError):
        load_workbook_snapshot(encrypted)
    with pytest.raises(InvalidPasswordError):
        load_workbook_snapshot(encrypted, password="not-the-password")


def test_unsupported_format(tmp_path: Path) -> None:
    bogus = tmp_path / "file.xls"
    bogus.write_bytes(b"PK\x03\x04junk")
    with pytest.raises(UnsupportedFormatError):
        load_workbook_snapshot(bogus)
