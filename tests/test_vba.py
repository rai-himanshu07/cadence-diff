"""VBA project reading, module-text comparison, and coverage boundaries."""

from __future__ import annotations

import io
import zipfile
from pathlib import Path

import pytest
from openpyxl import Workbook

from qc_tool.config.profile import DeliverableProfile
from qc_tool.coverage import CoverageState
from qc_tool.excel.diff_vba import diff_workbook_vba, vba_coverage
from qc_tool.excel.preflight import preflight_workbook
from qc_tool.findings import FindingClass
from qc_tool.io.loader import load_workbook_snapshot
from qc_tool.io.model import WorkbookSnapshot
from qc_tool.io.vba import (
    VbaModule,
    VbaProjectScan,
    VbaReadError,
    decompress_ovba,
    scan_vba_project,
)
from tests.fixtures.vba_writer import build_vba_project, compress_ovba

_MODULE = (
    'Attribute VB_Name = "Module1"\r\n'
    "Sub Refresh()\r\n"
    "    Dim total As Double\r\n"
    "    total = 0\r\n"
    "    total = total + 1\r\n"
    "End Sub\r\n"
)


def _macro_workbook(path: Path, modules: dict[str, str], *, protected: bool = False) -> Path:
    """A real xlsm package: openpyxl content plus a compound-file VBA project."""
    plain = path.with_suffix(".plain.xlsx")
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Data"
    for row in range(1, 4):
        sheet.cell(row=row, column=1, value=row)
    workbook.save(plain)

    project = build_vba_project(modules, protected=protected)
    with zipfile.ZipFile(plain) as source:
        entries = [(item, source.read(item.filename)) for item in source.infolist()]
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as dest:
        for item, payload in entries:
            dest.writestr(item, payload)
        dest.writestr("xl/vbaProject.bin", project)
    plain.unlink()
    return path


# --- MS-OVBA decompression ----------------------------------------------


def test_container_round_trips_through_the_production_reader() -> None:
    payload = (_MODULE * 40).encode("cp1252")

    container = compress_ovba(payload)

    assert len(container) < len(payload)  # copy tokens were used
    assert decompress_ovba(container) == payload


def test_container_without_the_signature_byte_is_refused() -> None:
    with pytest.raises(VbaReadError, match="0x01 signature"):
        decompress_ovba(b"\x02\x00\x30")


def test_chunk_with_a_wrong_signature_is_refused() -> None:
    with pytest.raises(VbaReadError, match="signature is not"):
        decompress_ovba(b"\x01\x00\x00")


def test_chunk_running_past_the_stream_is_refused() -> None:
    with pytest.raises(VbaReadError, match="past the end"):
        decompress_ovba(b"\x01\xff\xbf" + b"\x00" * 4)


# --- compound file -------------------------------------------------------


def test_non_compound_project_reports_unavailable_not_absent() -> None:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("xl/workbook.xml", "<workbook/>")
        archive.writestr("xl/vbaProject.bin", b"definitely not a compound file")

    scan = scan_vba_project(buffer.getvalue())

    assert scan.present is True
    assert scan.available is False
    assert scan.modules == ()
    assert "MS-CFB" in scan.detail


def test_package_without_a_project_is_available_and_absent() -> None:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("xl/workbook.xml", "<workbook/>")

    scan = scan_vba_project(buffer.getvalue())

    assert scan.present is False
    assert scan.available is True
    assert scan.modules == ()


def test_reserved_streams_are_not_reported_as_modules() -> None:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("xl/vbaProject.bin", build_vba_project({"Module1": _MODULE}))

    scan = scan_vba_project(buffer.getvalue())

    assert [module.name for module in scan.modules] == ["Module1"]


def test_locked_project_is_reported_and_still_readable() -> None:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr(
            "xl/vbaProject.bin", build_vba_project({"Module1": _MODULE}, protected=True)
        )

    scan = scan_vba_project(buffer.getvalue())

    assert scan.protected is True
    assert scan.available is True
    assert [module.name for module in scan.modules] == ["Module1"]
    assert "locked for viewing" in scan.detail


def test_modules_are_returned_in_a_stable_order() -> None:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr(
            "xl/vbaProject.bin",
            build_vba_project({"Zeta": _MODULE, "Alpha": _MODULE, "Mid": _MODULE}),
        )

    first = scan_vba_project(buffer.getvalue())
    second = scan_vba_project(buffer.getvalue())

    assert [module.name for module in first.modules] == ["Alpha", "Mid", "Zeta"]
    assert first.modules == second.modules


# --- loader --------------------------------------------------------------


def test_macro_workbook_loads_its_modules(tmp_path: Path) -> None:
    path = _macro_workbook(tmp_path / "book.xlsm", {"Module1": _MODULE})

    snapshot = load_workbook_snapshot(path)

    assert snapshot.vba.present is True
    assert snapshot.vba.available is True
    assert [module.name for module in snapshot.vba.modules] == ["Module1"]
    assert snapshot.vba.modules[0].line_count == 6


def test_workbook_without_macros_reports_no_project(tmp_path: Path) -> None:
    path = tmp_path / "book.xlsx"
    workbook = Workbook()
    workbook.save(path)

    snapshot = load_workbook_snapshot(path)

    assert snapshot.vba.present is False
    assert snapshot.vba.available is True
    assert snapshot.vba.modules == ()


# --- comparison ----------------------------------------------------------


def _snapshot(modules: dict[str, str]) -> WorkbookSnapshot:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("xl/vbaProject.bin", build_vba_project(modules))
    return WorkbookSnapshot(
        source_name="book.xlsm",
        file_format="xlsm",
        formulas_available=True,
        styles_available=True,
        vba=scan_vba_project(buffer.getvalue()),
    )


def test_added_and_removed_modules_are_reported() -> None:
    baseline = _snapshot({"Kept": _MODULE, "Gone": _MODULE})
    current = _snapshot({"Kept": _MODULE, "Fresh": _MODULE})

    findings = diff_workbook_vba(baseline, current)

    assert sorted(finding.element or "" for finding in findings) == ["Fresh", "Gone"]
    assert all(
        finding.finding_class is FindingClass.VBA_MODULE_CHANGED for finding in findings
    )


def test_unchanged_module_produces_no_finding() -> None:
    assert diff_workbook_vba(_snapshot({"Kept": _MODULE}), _snapshot({"Kept": _MODULE})) == []


def test_changed_module_reports_counts_and_line_ranges() -> None:
    changed = _MODULE.replace("total = total + 1", "total = total + 2")
    baseline = _snapshot({"Module1": _MODULE})
    current = _snapshot({"Module1": changed})

    findings = diff_workbook_vba(baseline, current)

    assert len(findings) == 1
    message = findings[0].message
    assert "1 lines added, 1 removed" in message
    assert "at 5" in message


def test_finding_never_carries_module_source_text() -> None:
    changed = _MODULE.replace("Dim total As Double", 'Dim secret As String: secret = "p@ss"')
    findings = diff_workbook_vba(_snapshot({"M": _MODULE}), _snapshot({"M": changed}))

    rendered = " ".join(
        str(value)
        for finding in findings
        for value in (finding.message, finding.baseline_value, finding.current_value)
    )
    assert "p@ss" not in rendered
    assert "secret" not in rendered
    assert "sha256" in rendered


def test_comparison_is_skipped_when_either_side_is_unreadable() -> None:
    readable = _snapshot({"Module1": _MODULE})
    blocked = WorkbookSnapshot(
        source_name="blocked.xlsm",
        file_format="xlsm",
        formulas_available=True,
        styles_available=True,
        vba=VbaProjectScan(present=True, available=False, detail="not an MS-CFB file"),
    )

    assert diff_workbook_vba(blocked, readable) == []
    assert diff_workbook_vba(readable, blocked) == []


def test_comparison_ignores_line_ending_differences() -> None:
    baseline = _snapshot({"Module1": _MODULE})
    current = _snapshot({"Module1": _MODULE.replace("\r\n", "\n")})

    assert diff_workbook_vba(baseline, current) == []


# --- coverage ------------------------------------------------------------


def test_coverage_is_checked_when_no_project_exists() -> None:
    plain = WorkbookSnapshot(
        source_name="plain.xlsx",
        file_format="xlsx",
        formulas_available=True,
        styles_available=True,
        vba=VbaProjectScan(present=False, available=True),
    )

    item = vba_coverage(plain)

    assert item.check_id == "excel-vba"
    assert item.state is CoverageState.CHECKED
    assert item.detail == "no VBA project is present"


def test_coverage_is_unavailable_when_the_project_cannot_be_read() -> None:
    blocked = WorkbookSnapshot(
        source_name="blocked.xlsm",
        file_format="xlsm",
        formulas_available=True,
        styles_available=True,
        vba=VbaProjectScan(present=True, available=False, detail="not an MS-CFB file"),
    )

    item = vba_coverage(blocked)

    assert item.state is CoverageState.UNAVAILABLE
    assert "not an MS-CFB file" in item.detail


def test_coverage_is_degraded_when_some_modules_are_undecodable() -> None:
    partial = WorkbookSnapshot(
        source_name="partial.xlsm",
        file_format="xlsm",
        formulas_available=True,
        styles_available=True,
        vba=VbaProjectScan(
            present=True,
            available=True,
            modules=(VbaModule("Module1", 3, "abc"),),
            unreadable_modules=1,
            detail="1 module streams could not be decoded",
        ),
    )

    item = vba_coverage(partial)

    assert item.state is CoverageState.DEGRADED
    assert "could not be decoded" in item.detail


def test_preflight_registers_the_vba_coverage_item(tmp_path: Path) -> None:
    snapshot = load_workbook_snapshot(
        _macro_workbook(tmp_path / "book.xlsm", {"Module1": _MODULE})
    )

    result = preflight_workbook(snapshot, DeliverableProfile(name="vba-tests"))

    item = next(item for item in result.coverage if item.check_id == "excel-vba")
    assert item.state is CoverageState.CHECKED
    assert item.findings == 1
