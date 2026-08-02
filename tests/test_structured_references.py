"""Structured Excel table references and table-schema QC."""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.formula import Tokenizer
from openpyxl.worksheet.table import Table

from qc_tool.config.profile import default_profile
from qc_tool.coverage import CoverageState
from qc_tool.excel.align import align_workbooks
from qc_tool.excel.dependency import build_dependency_graph, dependents_of
from qc_tool.excel.diff_structure import diff_workbook_structure
from qc_tool.excel.preflight import preflight_workbook
from qc_tool.excel.references import ReferenceStatus, resolve_reference
from qc_tool.findings import FindingClass
from qc_tool.io.loader import load_workbook_snapshot
from qc_tool.io.model import (
    CellRecord,
    ChartDescriptor,
    ChartSeries,
    SheetSnapshot,
    TableDescriptor,
    WorkbookSnapshot,
)


def _table_workbook(*, formula: str | None = None) -> WorkbookSnapshot:
    cells = {
        (1, 1): CellRecord(1, 1, "Key"),
        (1, 2): CellRecord(1, 2, "Amount"),
        (1, 3): CellRecord(1, 3, "Cost"),
        (1, 4): CellRecord(1, 4, "Col]Name"),
        (1, 5): CellRecord(1, 5, "#Items"),
        (1, 6): CellRecord(1, 6, "@Plan"),
    }
    for row in range(2, 6):
        for column in range(1, 7):
            cells[(row, column)] = CellRecord(row, column, row * 10 + column)
    if formula is not None:
        cells[(3, 3)] = CellRecord(3, 3, None, formula=formula)
    return WorkbookSnapshot(
        source_name="tables.xlsx",
        file_format="xlsx",
        formulas_available=True,
        styles_available=True,
        formula_presence_available=True,
        tables_available=True,
        sheets=[SheetSnapshot("Data", "visible", 5, 6, cells)],
        tables=[
            TableDescriptor(
                sheet="Data",
                name="A1Table",
                display_name="A1Table",
                cell_range="A1:F5",
                columns=["Key", "Amount", "Cost", "Col]Name", "#Items", "@Plan"],
                header_row_count=1,
                totals_row_count=1,
                source_id=1,
            )
        ],
    )


@pytest.mark.parametrize(
    ("target", "host", "expected"),
    [
        ("A1Table[Amount]", None, (("Data", 2, 2, 4, 2),)),
        ("A1Table[[#Data],[Amount]]", None, (("Data", 2, 2, 4, 2),)),
        ("A1Table[[#Headers],[Amount]]", None, (("Data", 1, 2, 1, 2),)),
        ("A1Table[[#Totals],[Amount]]", None, (("Data", 5, 2, 5, 2),)),
        ("A1Table[[#All],[Amount]]", None, (("Data", 1, 2, 5, 2),)),
        ("A1Table[[Amount]:[Cost]]", None, (("Data", 2, 2, 4, 3),)),
        ("A1Table[Col]]Name]", None, (("Data", 2, 4, 4, 4),)),
        ("A1Table['#Items]", None, (("Data", 2, 5, 4, 5),)),
        ("A1Table['@Plan]", None, (("Data", 2, 6, 4, 6),)),
        ("[@Amount]", (3, 3), (("Data", 3, 2, 3, 2),)),
        ("A1Table[[#This Row],[Amount]]", (3, 3), (("Data", 3, 2, 3, 2),)),
    ],
)
def test_structured_reference_matrix(
    target: str,
    host: tuple[int, int] | None,
    expected: tuple[tuple[str, int, int, int, int], ...],
) -> None:
    result = resolve_reference(
        _table_workbook(),
        target,
        host_sheet="Data",
        host_cell=host,
    )

    assert result.status is ReferenceStatus.RESOLVED
    assert tuple(
        (item.sheet, item.min_row, item.min_col, item.max_row, item.max_col)
        for item in result.ranges
    ) == expected


@pytest.mark.parametrize(
    "target",
    [
        "A1Table[Amount]",
        "[@Amount]",
        "A1Table[[#Headers],[Amount]]",
        "A1Table[[#Totals],[Amount]]",
        "A1Table[[#Data],[Amount]]",
        "A1Table[[Amount]:[Cost]]",
        "A1Table[Col]]Name]",
        "A1Table['#Items]",
        "A1Table['@Plan]",
    ],
)
def test_openpyxl_tokenizer_preserves_structured_range_operand(target: str) -> None:
    operands = [
        token
        for token in Tokenizer(f"=SUM({target})").items
        if token.type == "OPERAND" and token.subtype == "RANGE"
    ]

    assert [token.value for token in operands] == [target]


@pytest.mark.parametrize(
    ("target", "status"),
    [
        ("A1Table[Missing]", ReferenceStatus.INVALID),
        ("A1Table[[#Data],[Amount]", ReferenceStatus.INVALID),
        (
            "A1Table[[#Data],[Amount],[Cost]]",
            ReferenceStatus.UNSUPPORTED,
        ),
        ("[Book.xlsx]Data!$A$1", ReferenceStatus.UNSUPPORTED),
    ],
)
def test_structured_reference_failures_are_classified(
    target: str, status: ReferenceStatus
) -> None:
    result = resolve_reference(_table_workbook(), target, host_sheet="Data")

    assert result.status is status
    assert result.ranges == ()


def test_structured_chart_references_are_validated_against_table_schema() -> None:
    workbook = _table_workbook()
    workbook.charts.append(
        ChartDescriptor(
            sheet="Data",
            title="Table chart",
            chart_type="LineChart",
            series=[ChartSeries(0, "A1Table[Amount]", "A1Table[Key]")],
        )
    )

    result = preflight_workbook(workbook, default_profile())

    assert not any(
        finding.finding_class
        in {FindingClass.CHART_REFERENCE_INVALID, FindingClass.CHART_LENGTH_MISMATCH}
        for finding in result.findings
    )


def test_ooxml_structured_references_round_trip(tmp_path: Path) -> None:
    path = tmp_path / "structured.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Data"
    sheet.append(["Key", "Amount", "Calculated"])
    sheet.append(["A", 10, "=[@Amount]"])
    sheet.append(["B", 20, "=[@Amount]"])
    sheet.append(["C", 30, "=[@Amount]"])
    sheet.add_table(Table(displayName="MetricsTable", ref="A1:C4"))
    chart = LineChart()
    chart.add_data(
        Reference(sheet, min_col=2, min_row=1, max_row=4),
        titles_from_data=True,
    )
    chart.set_categories(Reference(sheet, min_col=1, min_row=2, max_row=4))
    chart.series[0].val.numRef.f = "MetricsTable[Amount]"
    assert chart.series[0].cat.numRef is not None
    chart.series[0].cat.numRef.f = "MetricsTable[Key]"
    chart.anchor = "E2"
    sheet.add_chart(chart)
    workbook.save(path)

    snapshot = load_workbook_snapshot(path)
    graph = build_dependency_graph(snapshot)
    preflight = preflight_workbook(snapshot, default_profile())

    assert snapshot.tables_available
    assert snapshot.tables[0].columns == ["Key", "Amount", "Calculated"]
    assert snapshot.charts[0].series[0].values_ref == "MetricsTable[Amount]"
    assert snapshot.charts[0].series[0].categories_ref == "MetricsTable[Key]"
    assert dependents_of(graph, "Data", "B2") == ["Data!C2"]
    assert not any(
        finding.finding_class is FindingClass.CHART_REFERENCE_INVALID
        for finding in preflight.findings
    )


def test_unsupported_chart_reference_degrades_structure_coverage() -> None:
    workbook = _table_workbook()
    workbook.charts.append(
        ChartDescriptor(
            sheet="Data",
            title="Union chart",
            chart_type="LineChart",
            series=[
                ChartSeries(0, "A1Table[[#Data],[Amount],[Cost]]", "A1Table[Key]")
            ],
        )
    )

    result = preflight_workbook(workbook, default_profile())
    coverage = next(
        item for item in result.coverage if item.check_id == "excel-intrinsic-structure"
    )

    assert coverage.state is CoverageState.DEGRADED
    assert "unsupported" in coverage.detail.lower()
    assert not any(
        finding.finding_class is FindingClass.CHART_REFERENCE_INVALID
        for finding in result.findings
    )


def test_structured_formula_dependencies_resolve_current_row_and_table_column() -> None:
    workbook = _table_workbook(formula="=[@Amount]+SUM(A1Table[Amount])")

    graph = build_dependency_graph(workbook)

    assert dependents_of(graph, "Data", "B2") == ["Data!C3"]
    assert dependents_of(graph, "Data", "B3") == ["Data!C3"]
    assert graph.coverage_state is CoverageState.CHECKED


def test_unsupported_structured_dependency_degrades_coverage() -> None:
    workbook = _table_workbook(
        formula="=SUM(A1Table[[#Data],[Amount],[Cost]])"
    )

    graph = build_dependency_graph(workbook)
    preflight = preflight_workbook(workbook, default_profile())
    coverage = next(
        item for item in preflight.coverage if item.check_id == "excel-dependencies"
    )

    assert graph.coverage_state is CoverageState.DEGRADED
    assert coverage.state is CoverageState.DEGRADED
    assert graph.unsupported_references == ["structured_reference"]
    assert graph.unsupported_reason_counts == {"structured_reference": 1}


def test_table_structure_diff_detects_rename_range_and_schema_changes() -> None:
    baseline = _table_workbook()
    current = _table_workbook()
    current.tables[0].name = "RenamedTable"
    current.tables[0].display_name = "RenamedTable"
    current.tables[0].cell_range = "A1:E6"
    current.tables[0].columns = ["Key", "Cost", "Amount", "Col]Name", "Margin"]
    current.sheets[0].max_row = 6
    current.sheets[0].max_column = 6

    findings = diff_workbook_structure(
        baseline,
        current,
        align_workbooks(baseline, current),
    )
    table_findings = [
        finding
        for finding in findings
        if finding.finding_class is FindingClass.TABLE_STRUCTURE_CHANGED
    ]

    assert any("renamed" in finding.message for finding in table_findings)
    assert any("range" in finding.message for finding in table_findings)
    assert any("columns" in finding.message for finding in table_findings)


def test_table_structure_diff_detects_addition_and_removal() -> None:
    baseline = _table_workbook()
    current = _table_workbook()
    current.tables[0] = TableDescriptor(
        sheet="Data",
        name="OtherTable",
        display_name="OtherTable",
        cell_range="F1:F5",
        columns=["Other"],
        source_id=2,
    )

    findings = diff_workbook_structure(
        baseline,
        current,
        align_workbooks(baseline, current),
    )
    messages = [
        finding.message
        for finding in findings
        if finding.finding_class is FindingClass.TABLE_STRUCTURE_CHANGED
    ]

    assert any("removed" in message for message in messages)
    assert any("added" in message for message in messages)


def test_table_diff_skips_unavailable_table_metadata() -> None:
    baseline = _table_workbook()
    baseline.tables_available = False
    baseline.tables = []
    current = _table_workbook()

    findings = diff_workbook_structure(
        baseline,
        current,
        align_workbooks(baseline, current),
    )

    assert not any(
        finding.finding_class is FindingClass.TABLE_STRUCTURE_CHANGED
        for finding in findings
    )


def test_unavailable_table_metadata_degrades_preflight_coverage() -> None:
    workbook = _table_workbook()
    workbook.tables_available = False
    workbook.tables = []

    result = preflight_workbook(workbook, default_profile())
    coverage = next(
        item for item in result.coverage if item.check_id == "excel-intrinsic-structure"
    )

    assert coverage.state is CoverageState.DEGRADED
    assert "table metadata" in coverage.detail.lower()


def test_table_structure_diff_detects_header_and_totals_changes() -> None:
    baseline = _table_workbook()
    current = _table_workbook()
    current.tables[0].header_row_count = 0
    current.tables[0].totals_row_count = 0

    findings = diff_workbook_structure(
        baseline,
        current,
        align_workbooks(baseline, current),
    )

    assert any(
        finding.finding_class is FindingClass.TABLE_STRUCTURE_CHANGED
        and "header/totals" in finding.message
        for finding in findings
    )


def test_duplicate_chart_titles_do_not_collapse() -> None:
    baseline = _table_workbook()
    current = _table_workbook()
    chart = ChartDescriptor(
        sheet="Data",
        title="Duplicate",
        chart_type="LineChart",
        series=[ChartSeries(0, "Data!$B$2:$B$4", "Data!$A$2:$A$4")],
    )
    baseline.charts = [chart, chart]
    current.charts = [chart]

    findings = diff_workbook_structure(
        baseline,
        current,
        align_workbooks(baseline, current),
    )

    chart_findings = [
        finding
        for finding in findings
        if finding.finding_class is FindingClass.CHART_STRUCTURE_CHANGED
    ]
    assert len(chart_findings) == 1
    assert chart_findings[0].element == "Duplicate"
    assert "removed" in chart_findings[0].message
