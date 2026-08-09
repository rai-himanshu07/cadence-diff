"""Measured evidence matrix for the large-workbook override policy.

All inputs are synthetic and all output belongs under ignored ``artifacts/``.
The runner exercises current production thresholds without changing them.
"""

from __future__ import annotations

import argparse
import dataclasses
import hashlib
import importlib
import json
import shutil
import subprocess
import sys
import time
import zipfile
from pathlib import Path

from openpyxl import Workbook


def _memory_bytes() -> tuple[int, int]:
    current = peak = 0
    for line in Path("/proc/self/status").read_text(encoding="ascii").splitlines():
        if line.startswith("VmRSS:"):
            current = int(line.split()[1]) * 1024
        elif line.startswith("VmHWM:"):
            peak = int(line.split()[1]) * 1024
    return current, peak


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1 << 20), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _dense(path: Path, *, rows: int, columns: int) -> None:
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet("Data")
    for row in range(1, rows + 1):
        sheet.append([row * columns + column for column in range(columns)])
    workbook.save(path)


def _sparse_full_grid(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    if sheet is None:
        raise RuntimeError("openpyxl did not create a worksheet")
    sheet["A1"] = 1
    sheet["XFD1048576"] = 2
    workbook.save(path)


def _complexity_refusal(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    if sheet is None:
        raise RuntimeError("openpyxl did not create a worksheet")
    sheet["A1"] = "=SUM(B1:B300000001)"
    sheet["B1"] = 1
    workbook.save(path)


def _metadata_only_copy(source: Path, destination: Path) -> None:
    shutil.copyfile(source, destination)
    with zipfile.ZipFile(destination) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    core = parts.get("docProps/core.xml", b"")
    if core:
        core = core.replace(b"</cp:coreProperties>", b" \n</cp:coreProperties>")
        parts["docProps/core.xml"] = core
    with zipfile.ZipFile(destination, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, content in parts.items():
            archive.writestr(name, content)


def _xlsb(path: Path, *, rows: int, columns: int) -> None:
    from tests.fixtures.xlsb_writer import write_xlsb

    values = [
        [f"row-{row}", *(row * columns + column for column in range(1, columns))]
        for row in range(1, rows + 1)
    ]
    write_xlsb(path, {"Data": values})


def _worker(arguments: argparse.Namespace) -> int:
    resource = importlib.import_module("resource")
    limit = int(arguments.limit_gb * 1024**3)
    resource.setrlimit(resource.RLIMIT_AS, (limit, limit))
    started = time.perf_counter()
    outcome = "passed"
    detail: dict[str, object] = {}
    try:
        if arguments.action == "load":
            from qc_tool.io.loader import load_workbook_snapshot

            snapshot = load_workbook_snapshot(
                Path(arguments.input),
                allow_large_workbook=arguments.override,
            )
            detail = {
                "format": snapshot.file_format,
                "cells": sum(len(sheet.cells) for sheet in snapshot.sheets),
                "workload": dataclasses.asdict(snapshot.workload),
            }
        elif arguments.action == "complexity":
            from qc_tool.excel.complexity import assess_workbook_complexity
            from qc_tool.io.loader import load_workbook_snapshot

            snapshot = load_workbook_snapshot(
                Path(arguments.input),
                allow_large_workbook=True,
            )
            complexity = assess_workbook_complexity(
                snapshot,
                allow_complex_workbook=arguments.override,
            )
            detail = {"complexity": dataclasses.asdict(complexity)}
        elif arguments.action == "projection":
            from qc_tool.projection import project_cycle_volume

            projection = project_cycle_volume(
                Path(arguments.input),
                Path(arguments.second_input),
            )
            detail = {
                "projection": (
                    dataclasses.asdict(projection) if projection is not None else None
                )
            }
        else:
            raise ValueError(f"unknown action {arguments.action!r}")
    except Exception as exc:
        outcome = "refused"
        detail = {"error_type": type(exc).__name__, "error": str(exc)}
    current, peak = _memory_bytes()
    print(
        json.dumps(
            {
                "outcome": outcome,
                "override": arguments.override,
                "elapsed_seconds": round(time.perf_counter() - started, 2),
                "current_rss_bytes": current,
                "peak_rss_bytes": peak,
                **detail,
            }
        )
    )
    return 0


def _sample(
    action: str,
    path: Path,
    *,
    override: bool,
    limit_gb: float,
    second: Path | None = None,
) -> dict[str, object]:
    command = [
        sys.executable,
        "-m",
        "tests.workload_policy_probe",
        "worker",
        "--action",
        action,
        "--input",
        str(path),
        "--limit-gb",
        str(limit_gb),
    ]
    if override:
        command.append("--override")
    if second is not None:
        command.extend(("--second-input", str(second)))
    completed = subprocess.run(command, capture_output=True, text=True)
    if completed.returncode != 0:
        raise RuntimeError(completed.stderr.strip())
    return json.loads(completed.stdout.strip().splitlines()[-1])


def _run(arguments: argparse.Namespace) -> int:
    output = arguments.output.resolve()
    inputs = output / "inputs"
    inputs.mkdir(parents=True, exist_ok=True)
    below = inputs / "below-warning.xlsx"
    warning = inputs / "warning-scale.xlsx"
    refusal = inputs / "refusal-scale.xlsx"
    sparse = inputs / "sparse-full-grid.xlsx"
    complexity = inputs / "dependency-refusal.xlsx"
    near_identical = inputs / "warning-scale-metadata-only.xlsx"
    xlsb = inputs / "large-values.xlsb"

    builders = (
        (below, lambda path: _dense(path, rows=20_000, columns=10)),
        (warning, lambda path: _dense(path, rows=100_000, columns=10)),
        (refusal, lambda path: _dense(path, rows=500_000, columns=10)),
        (sparse, _sparse_full_grid),
        (complexity, _complexity_refusal),
        (xlsb, lambda path: _xlsb(path, rows=50_000, columns=8)),
    )
    for path, builder in builders:
        if not path.exists():
            builder(path)
    if not near_identical.exists():
        _metadata_only_copy(warning, near_identical)

    before = {path.name: _sha256(path) for path, _builder in builders}
    before[near_identical.name] = _sha256(near_identical)
    records = {
        "below_default": _sample(
            "load", below, override=False, limit_gb=arguments.limit_gb
        ),
        "warning_default": _sample(
            "load", warning, override=False, limit_gb=arguments.limit_gb
        ),
        "refusal_default": _sample(
            "load", refusal, override=False, limit_gb=arguments.limit_gb
        ),
        "refusal_override": _sample(
            "load", refusal, override=True, limit_gb=arguments.limit_gb
        ),
        "sparse_default": _sample(
            "load", sparse, override=False, limit_gb=arguments.limit_gb
        ),
        "sparse_override": _sample(
            "load", sparse, override=True, limit_gb=arguments.limit_gb
        ),
        "complexity_default": _sample(
            "complexity", complexity, override=False, limit_gb=arguments.limit_gb
        ),
        "complexity_override": _sample(
            "complexity", complexity, override=True, limit_gb=arguments.limit_gb
        ),
        "near_identical_projection": _sample(
            "projection",
            warning,
            second=near_identical,
            override=False,
            limit_gb=arguments.limit_gb,
        ),
        "xlsb_load": _sample(
            "load", xlsb, override=False, limit_gb=arguments.limit_gb
        ),
    }
    existing_diff = Path("artifacts/monster-probe/probe-r40000-l4.json")
    after = {path.name: _sha256(path) for path, _builder in builders}
    after[near_identical.name] = _sha256(near_identical)
    result = {
        "limit_gb": arguments.limit_gb,
        "inputs": {
            path.name: {"bytes": path.stat().st_size, "sha256": before[path.name]}
            for path in (*[item[0] for item in builders], near_identical)
        },
        "records": records,
        "existing_moderate_huge_diff_probe": (
            json.loads(existing_diff.read_text(encoding="utf-8"))
            if existing_diff.exists()
            else None
        ),
        "source_hashes_unchanged": before == after,
        "xlsb_policy_boundary": (
            "Generated XLSB demonstrates retained-cell loading, but no "
            "representative refusal-scale private XLSB is available."
        ),
    }
    target = output / "result.json"
    target.write_text(json.dumps(result, indent=2), encoding="utf-8")
    print(json.dumps(result, indent=2))
    return 0


def _parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description=__doc__)
    subparsers = parser.add_subparsers(dest="command", required=True)
    worker = subparsers.add_parser("worker")
    worker.add_argument("--action", choices=("load", "complexity", "projection"))
    worker.add_argument("--input", required=True)
    worker.add_argument("--second-input", default="")
    worker.add_argument("--override", action="store_true")
    worker.add_argument("--limit-gb", type=float, default=6.0)
    run = subparsers.add_parser("run")
    run.add_argument("--output", type=Path, required=True)
    run.add_argument("--limit-gb", type=float, default=6.0)
    return parser


def main(argv: list[str] | None = None) -> int:
    arguments = _parser().parse_args(argv)
    if arguments.command == "worker":
        return _worker(arguments)
    return _run(arguments)


if __name__ == "__main__":
    raise SystemExit(main())
