"""Run history tests (criterion 13)."""

import datetime as dt
import json
import os
import sqlite3
import zipfile
from pathlib import Path
from typing import Literal

import pytest
from openpyxl.utils.cell import coordinate_to_tuple

from qc_tool.config.profile import DeliverableProfile
from qc_tool.engine import QCRunResult
from qc_tool.findings import (
    Finding,
    FindingClass,
    FindingSubtype,
    FindingTemporalContext,
    Materiality,
    NumericCounterfactualBasis,
    SeriesAnchorV1,
    SeriesAnchorV2,
    Severity,
)
from qc_tool.findings_store import FindingSequence, decode_block, encode_block
from qc_tool.history.carry_forward import apply_carry_forward
from qc_tool.history.longitudinal import (
    DecisionOrigin,
    DossierStatus,
    contract_scope_for_profile,
)
from qc_tool.history.review_state import RunSignoff
from qc_tool.history.store import RunHistory, export_runs_archive, sha256_file
from qc_tool.package import PackageManifest
from qc_tool.review_series import (
    ReviewSlice,
    anchor_matches_finding,
    anchor_segment,
    cluster_confirmation_updates,
)
from qc_tool.run_service import perform_run
from tests.conftest import fixture_profile


def _longitudinal_finding(
    finding_id: str,
    *,
    location: str = "A1",
    value: str = "1",
) -> Finding:
    return Finding(
        finding_id=finding_id,
        artifact="excel",
        finding_class=FindingClass.VALUE_CHANGED,
        severity=Severity.CRITICAL,
        sheet="Data",
        location=location,
        baseline_value="0",
        current_value=value,
        message="value changed",
    )


def _record_longitudinal_run(
    history: RunHistory,
    finding_id: str,
    *,
    profile_name: str = "fixture",
    location: str = "A1",
    value: str = "1",
    rerun_of: int | None = None,
    profile_snapshot: DeliverableProfile | None = None,
) -> int:
    return history.record_run(
        QCRunResult(
            profile_name=profile_name,
            findings=[
                _longitudinal_finding(
                    finding_id,
                    location=location,
                    value=value,
                )
            ],
        ),
        file_hashes={},
        report_paths={},
        rerun_of=rerun_of,
        profile_snapshot=profile_snapshot,
    )


def _finalize(history: RunHistory, run_id: int) -> None:
    history.record_signoff(
        RunSignoff(
            run_id=run_id,
            finalized_at="2026-08-07T00:00:00+00:00",
            review_state_digest="review-state",
            profile_sha256="profile-sha",
            attestation_path="run.qca",
            attestation_sha256="attestation-sha",
        )
    )


def test_record_and_list_runs(
    qc_result: QCRunResult, fixture_dir: Path, tmp_path: Path
) -> None:
    history = RunHistory(tmp_path / "history.sqlite3")
    hashes = {
        "baseline_excel": sha256_file(fixture_dir / "baseline.xlsx"),
        "current_excel": sha256_file(fixture_dir / "current.xlsx"),
    }
    first = history.record_run(
        qc_result, file_hashes=hashes, report_paths={"html": "runs/1/report.html"}
    )
    second = history.record_run(qc_result, file_hashes=hashes, report_paths={})
    assert second > first

    runs = history.list_runs()
    assert [run.run_id for run in runs] == [second, first]
    newest = runs[0]
    assert newest.profile == "fixture"
    assert newest.mode == qc_result.mode
    assert newest.coverage == qc_result.coverage
    assert newest.files["current_excel"] == "current.xlsx"
    assert newest.file_hashes == hashes
    assert newest.counts["critical"] > 0
    assert newest.review_counts["critical"] > 0
    assert newest.review_counts["critical"] <= newest.counts["critical"]
    assert newest.started_at.tzinfo is not None  # timezone-aware UTC
    assert newest.findings == []  # listing stays lightweight


def test_storage_bytes_recorded_backfilled_and_summed(
    qc_result: QCRunResult, fixture_dir: Path, tmp_path: Path
) -> None:
    history = RunHistory(tmp_path / "history.sqlite3")
    hashes = {"current_excel": sha256_file(fixture_dir / "current.xlsx")}
    run_dir = tmp_path / "runs" / "run-1"
    run_dir.mkdir(parents=True)
    report = run_dir / "qc_report.html"
    report.write_bytes(b"x" * 4096)

    run_id = history.record_run(
        qc_result, file_hashes=hashes, report_paths={"html": str(report)}
    )
    record = history.get_run(run_id)
    assert record.storage_bytes is not None
    # At least the findings JSON plus the 4 KiB report file.
    assert record.storage_bytes > 4096

    # A legacy row (recorded before the column existed) backfills on demand.
    import sqlite3

    with sqlite3.connect(tmp_path / "history.sqlite3") as conn:
        conn.execute("UPDATE runs SET storage_bytes = NULL WHERE id = ?", (run_id,))
    assert history.get_run(run_id).storage_bytes is None
    assert history.backfill_storage_bytes() == 1
    refreshed = history.get_run(run_id).storage_bytes
    assert refreshed is not None and refreshed > 4096

    total, archived_bytes, unmeasured = history.storage_summary()
    assert total == refreshed
    assert archived_bytes == 0
    assert unmeasured == 0
    history.set_archived([run_id], True)
    _, archived_bytes, _ = history.storage_summary()
    assert archived_bytes == total


def test_archive_hides_runs_without_losing_them(
    qc_result: QCRunResult, tmp_path: Path
) -> None:
    history = RunHistory(tmp_path / "history.sqlite3")
    keep = history.record_run(qc_result, file_hashes={}, report_paths={})
    retire = history.record_run(qc_result, file_hashes={}, report_paths={})

    assert history.set_archived([retire], True) == 1

    active = history.list_runs(include_archived=False)
    assert [run.run_id for run in active] == [keep]
    assert {run.run_id: run.archived for run in history.list_runs()} == {
        keep: False,
        retire: True,
    }
    # The evidence itself survives archiving.
    assert history.get_run(retire).findings

    assert history.set_archived([retire], False) == 1
    assert [run.run_id for run in history.list_runs(include_archived=False)] == [
        retire,
        keep,
    ]


def test_delete_removes_records_annotations_and_managed_reports(
    qc_result: QCRunResult, tmp_path: Path
) -> None:
    work_dir = tmp_path / "work"
    runs_root = work_dir / "runs"
    report = runs_root / "abc" / "report.html"
    report.parent.mkdir(parents=True)
    report.write_text("<html></html>", encoding="utf-8")
    history = RunHistory(work_dir / "history.sqlite3")
    run_id = history.record_run(
        qc_result, file_hashes={}, report_paths={"html": str(report)}
    )
    finding_id = history.get_run(run_id).findings[0].finding_id
    history.set_annotation(run_id, finding_id, severity="info", comment="checked")

    assert history.delete_runs([run_id], managed_root=runs_root) == 1

    assert history.list_runs() == []
    assert history.get_annotations(run_id) == {}
    assert not report.exists()
    assert not report.parent.exists()  # the emptied run directory goes too


def test_deleting_every_run_does_not_reuse_a_run_id(
    qc_result: QCRunResult,
    tmp_path: Path,
) -> None:
    work_dir = tmp_path / "work"
    history = RunHistory(work_dir / "history.sqlite3")
    first = history.record_run(qc_result, file_hashes={}, report_paths={})

    assert history.delete_runs([first], managed_root=work_dir / "runs") == 1
    second = history.record_run(qc_result, file_hashes={}, report_paths={})

    assert second > first
    assert [record.run_id for record in history.list_runs()] == [second]


def test_delete_never_touches_files_outside_the_managed_root(
    qc_result: QCRunResult, tmp_path: Path
) -> None:
    work_dir = tmp_path / "work"
    runs_root = work_dir / "runs"
    runs_root.mkdir(parents=True)
    outsider = tmp_path / "not-ours.html"
    outsider.write_text("keep me", encoding="utf-8")
    history = RunHistory(work_dir / "history.sqlite3")
    run_id = history.record_run(
        qc_result, file_hashes={}, report_paths={"html": str(outsider)}
    )

    assert history.delete_runs([run_id], managed_root=runs_root) == 1

    assert history.list_runs() == []
    assert outsider.read_text(encoding="utf-8") == "keep me"


def test_export_archive_bundles_reports_and_a_path_free_manifest(
    qc_result: QCRunResult, tmp_path: Path
) -> None:
    work_dir = tmp_path / "work"
    runs_root = work_dir / "runs"
    report = runs_root / "abc" / "report.html"
    report.parent.mkdir(parents=True)
    report.write_text("<html>report</html>", encoding="utf-8")
    outsider = tmp_path / "escape.html"
    outsider.write_text("secret", encoding="utf-8")
    history = RunHistory(work_dir / "history.sqlite3")
    inside = history.record_run(
        qc_result, file_hashes={}, report_paths={"html": str(report)}
    )
    outside = history.record_run(
        qc_result, file_hashes={}, report_paths={"html": str(outsider)}
    )

    bundle = export_runs_archive(
        history.list_runs(), work_dir / "exports" / "runs.zip", managed_root=runs_root
    )

    with zipfile.ZipFile(bundle) as archive:
        names = set(archive.namelist())
        manifest = json.loads(archive.read("manifest.json"))
    assert f"run-{inside}/html.html" in names
    assert not any(name.startswith(f"run-{outside}/") for name in names)
    assert {entry["run_id"] for entry in manifest["runs"]} == {inside, outside}
    # A manifest is evidence, not a path leak.
    assert str(work_dir) not in json.dumps(manifest)
    if os.name == "posix":
        assert bundle.stat().st_mode & 0o077 == 0


def test_get_run_roundtrips_findings(qc_result: QCRunResult, tmp_path: Path) -> None:
    history = RunHistory(tmp_path / "history.sqlite3")
    run_id = history.record_run(
        qc_result, file_hashes={}, report_paths={"excel": "runs/1/report.xlsx"}
    )
    record = history.get_run(run_id)
    assert len(record.findings) == len(qc_result.findings)
    assert record.findings[0] == qc_result.findings[0]  # full pydantic round-trip
    assert record.report_paths == {"excel": "runs/1/report.xlsx"}
    assert record.mode == qc_result.mode
    assert record.coverage == qc_result.coverage


def test_get_missing_run_raises(tmp_path: Path) -> None:
    history = RunHistory(tmp_path / "history.sqlite3")
    with pytest.raises(KeyError):
        history.get_run(999)


def test_history_survives_reopen(qc_result: QCRunResult, tmp_path: Path) -> None:
    db = tmp_path / "history.sqlite3"
    RunHistory(db).record_run(qc_result, file_hashes={}, report_paths={})
    reopened = RunHistory(db)
    runs = reopened.list_runs()
    assert len(runs) == 1
    assert runs[0].started_at <= dt.datetime.now(dt.UTC)


def test_review_sessions_are_explicit_nonoverlapping_and_capped(
    qc_result: QCRunResult,
    tmp_path: Path,
) -> None:
    history = RunHistory(tmp_path / "history.sqlite3")
    first = history.record_run(qc_result, file_hashes={}, report_paths={})
    second = history.record_run(qc_result, file_hashes={}, report_paths={})
    started = dt.datetime(2026, 8, 6, 9, 0, tzinfo=dt.UTC)

    history.start_review_session(first, now=started)
    history.start_review_session(second, now=started + dt.timedelta(minutes=30))

    assert history.review_seconds(first, now=started + dt.timedelta(hours=1)) == 1800
    assert history.active_review_run() == second
    # A missing pause cannot accrue beyond four hours.
    assert history.review_seconds(second, now=started + dt.timedelta(days=1)) == 14400
    assert history.pause_review_sessions(now=started + dt.timedelta(days=1)) == 1
    assert history.active_review_run() is None
    assert history.review_seconds(second) == 14400


def test_pause_review_sessions_can_target_one_run(
    qc_result: QCRunResult,
    tmp_path: Path,
) -> None:
    history = RunHistory(tmp_path / "history.sqlite3")
    first = history.record_run(qc_result, file_hashes={}, report_paths={})
    second = history.record_run(qc_result, file_hashes={}, report_paths={})
    started = dt.datetime(2026, 8, 8, 9, 0, tzinfo=dt.UTC)
    history.start_review_session(second, now=started)

    # pausing an idle run leaves the active session of the other run running
    assert history.pause_review_sessions(run_id=first, now=started) == 0
    assert history.active_review_run() == second
    assert (
        history.pause_review_sessions(
            run_id=second, now=started + dt.timedelta(minutes=5)
        )
        == 1
    )
    assert history.active_review_run() is None
    assert history.review_seconds(second) == 300


def test_legacy_run_rehydrates_after_optional_contract_expansion(tmp_path: Path) -> None:
    database = tmp_path / "legacy.sqlite3"
    with sqlite3.connect(database) as connection:
        connection.executescript(
            """
            CREATE TABLE runs (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                started_at TEXT NOT NULL,
                profile TEXT NOT NULL,
                files TEXT NOT NULL,
                file_hashes TEXT NOT NULL,
                counts TEXT NOT NULL,
                disclosures TEXT NOT NULL,
                verified_crosschecks INTEGER NOT NULL,
                findings TEXT NOT NULL,
                report_paths TEXT NOT NULL
            );
            """
        )
        connection.execute(
            """
            INSERT INTO runs (
                started_at, profile, files, file_hashes, counts,
                disclosures, verified_crosschecks, findings, report_paths
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                "2026-07-26T10:00:00+00:00",
                "legacy",
                json.dumps({"current_excel": "old.xlsx"}),
                "{}",
                json.dumps({"warning": 1}),
                "[]",
                0,
                json.dumps(
                    [
                        {
                            "artifact": "excel",
                            "finding_class": "value_changed",
                            "expected_growth": True,
                            "materiality": "recent_restatement",
                            "message": "legacy finding",
                        }
                    ]
                ),
                "{}",
            ),
        )

    record = RunHistory(database).get_run(1)

    assert record.mode.value == "cycle_comparison"
    assert record.coverage == []
    assert record.review_counts == {}
    assert len(record.findings) == 1
    finding = record.findings[0]
    assert finding.message == "legacy finding"
    assert finding.expected_growth is True
    assert finding.expected_reason is None
    assert finding.materiality is not None
    assert finding.materiality.value == "recent_restatement"
    assert finding.temporal_context is None
    assert finding.evidence_tags == set()
    assert finding.impacts == []
    assert finding.current_excerpt is None


def test_annotations_persist_and_apply(qc_result: QCRunResult, tmp_path: Path) -> None:
    history = RunHistory(tmp_path / "history.sqlite3")
    run_id = history.record_run(qc_result, file_hashes={}, report_paths={})
    target = qc_result.findings[0].finding_id  # a critical finding

    history.set_annotation(
        run_id, target, severity="info", comment="expected per client email 07-25"
    )
    record = history.get_run(run_id)
    annotated = next(f for f in record.findings if f.finding_id == target)
    assert annotated.severity is not None and annotated.severity.value == "info"
    assert annotated.severity_overridden
    assert annotated.analyst_comment == "expected per client email 07-25"
    # Counts recomputed from findings reflect the override.
    from qc_tool.engine import QCRunResult as Result

    rebuilt = Result(profile_name=record.profile, findings=record.findings)
    assert rebuilt.counts != qc_result.counts

    # Upsert: comment-only update keeps engine severity for other findings.
    other = qc_result.findings[1].finding_id
    history.set_annotation(run_id, other, severity=None, comment="checked, fine")
    record = history.get_run(run_id)
    noted = next(f for f in record.findings if f.finding_id == other)
    assert not noted.severity_overridden
    assert noted.analyst_comment == "checked, fine"


def test_bulk_annotations_commit_as_one_group_decision(
    qc_result: QCRunResult, tmp_path: Path
) -> None:
    history = RunHistory(tmp_path / "history.sqlite3")
    run_id = history.record_run(qc_result, file_hashes={}, report_paths={})
    targets = [finding.finding_id for finding in qc_result.findings[:3]]

    history.set_annotations_bulk(
        run_id,
        [(target, "expected", "reviewed as one range") for target in targets],
    )

    record = history.get_run(run_id)
    reviewed = [finding for finding in record.findings if finding.finding_id in targets]
    assert len(reviewed) == 3
    assert all(finding.severity is not None for finding in reviewed)
    assert all(finding.severity.value == "expected" for finding in reviewed if finding.severity)
    assert all(finding.severity_overridden for finding in reviewed)
    assert all(finding.analyst_comment == "reviewed as one range" for finding in reviewed)


def test_annotations_flow_into_reports(qc_result: QCRunResult, tmp_path: Path) -> None:
    from openpyxl import load_workbook

    from qc_tool.engine import QCRunResult as Result
    from qc_tool.report.excel_report import write_excel_report
    from qc_tool.report.html_report import render_html_report

    history = RunHistory(tmp_path / "history.sqlite3")
    run_id = history.record_run(qc_result, file_hashes={}, report_paths={})
    target = qc_result.findings[0].finding_id
    history.set_annotation(run_id, target, severity="info", comment="reviewed & accepted")
    record = history.get_run(run_id)
    result = Result(profile_name=record.profile, findings=record.findings)

    path = tmp_path / "annotated.xlsx"
    write_excel_report(result, path)
    sheet = load_workbook(path)["Findings"]
    assert sheet["B2"].value == "info *"  # override marker
    assert sheet["R2"].value == "reviewed & accepted"

    html = render_html_report(result)
    assert r'"comment": "reviewed \u0026 accepted"' in html
    assert '"overridden": true' in html
    assert "member.severity + (member.overridden ? ' *' : '')" in html


def test_contract_scope_prefers_portable_id_and_labels_legacy_names() -> None:
    contract_id = "a" * 32

    assert contract_scope_for_profile("Old name", contract_id) == (
        f"contract:{contract_id}"
    )
    assert contract_scope_for_profile(" Monthly Pack ", "") == (
        "legacy-profile:monthly pack"
    )


def test_portable_contract_id_joins_renamed_profiles_in_dossier(tmp_path: Path) -> None:
    history = RunHistory(tmp_path / "history.sqlite3")
    contract_id = "b" * 32
    first = _record_longitudinal_run(
        history,
        "F1",
        profile_name="Old name",
        profile_snapshot=DeliverableProfile(
            name="Old name",
            contract_id=contract_id,
        ),
    )
    current = _record_longitudinal_run(
        history,
        "F2",
        profile_name="New name",
        rerun_of=first,
        profile_snapshot=DeliverableProfile(
            name="New name",
            contract_id=contract_id,
        ),
    )

    dossier = history.get_dossier(current, "F2")

    assert dossier.contract_scope == f"contract:{contract_id}"
    assert [entry.status for entry in dossier.entries] == [
        DossierStatus.EXACT,
        DossierStatus.EXACT,
    ]


def test_legacy_profile_names_do_not_join_longitudinal_evidence(tmp_path: Path) -> None:
    history = RunHistory(tmp_path / "history.sqlite3")
    first = _record_longitudinal_run(history, "F1", profile_name="Pack A")
    current = _record_longitudinal_run(
        history,
        "F2",
        profile_name="Pack B",
        rerun_of=first,
    )

    dossier = history.get_dossier(current, "F2")

    assert [entry.status for entry in dossier.entries] == [
        DossierStatus.NO_STORED_OBSERVATION,
        DossierStatus.EXACT,
    ]


def test_record_run_indexes_every_longitudinal_occurrence(tmp_path: Path) -> None:
    database = tmp_path / "history.sqlite3"
    history = RunHistory(database)
    findings = [
        _longitudinal_finding("F1", location="A1"),
        _longitudinal_finding("F2", location="A2"),
    ]
    run_id = history.record_run(
        QCRunResult(profile_name="fixture", findings=findings),
        file_hashes={},
        report_paths={},
    )

    with sqlite3.connect(database) as connection:
        count = connection.execute(
            "SELECT COUNT(*) FROM decision_occurrences WHERE run_id = ?",
            (run_id,),
        ).fetchone()

    assert count is not None and count[0] == len(findings)


def test_record_run_persists_resolved_formula_engines(tmp_path: Path) -> None:
    """Criterion 5: the resolved formula-engine/adapter-fingerprint per excel
    role must survive a history round trip, so Re-QC/carry-forward can later
    disclose a cross-run engine change.
    """
    history = RunHistory(tmp_path / "history.sqlite3")
    run_id = history.record_run(
        QCRunResult(
            profile_name="fixture",
            formula_engines={
                "baseline_excel": "native-biff12:1.2.3",
                "current_excel": "native-biff12:1.2.3",
            },
        ),
        file_hashes={},
        report_paths={},
    )

    record = history.get_run(run_id)

    assert record.formula_engines == {
        "baseline_excel": "native-biff12:1.2.3",
        "current_excel": "native-biff12:1.2.3",
    }


def test_legacy_run_without_formula_engines_defaults_to_an_empty_dict(
    tmp_path: Path,
) -> None:
    """A row recorded before this disclosure existed migrates to `{}`, never
    a crash or a guessed engine."""
    database = tmp_path / "history.sqlite3"
    history = RunHistory(database)
    run_id = history.record_run(
        QCRunResult(profile_name="fixture"), file_hashes={}, report_paths={}
    )
    with sqlite3.connect(database) as connection:
        connection.execute(
            "UPDATE runs SET formula_engines = '{}' WHERE id = ?", (run_id,)
        )

    assert history.get_run(run_id).formula_engines == {}


def test_record_run_reports_fixed_code_subphase_timings(tmp_path: Path) -> None:
    """Step 7 instrumentation: a diagnostic hook only -- optional, additive,
    and never required for correctness (the omitted-callback path above
    proves that)."""
    history = RunHistory(tmp_path / "history.sqlite3")
    findings = [
        _longitudinal_finding("F1", location="A1"),
        _longitudinal_finding("F2", location="A2"),
    ]
    observed: dict[str, float] = {}

    def on_subphase(name: str, elapsed_seconds: float) -> None:
        observed[name] = elapsed_seconds

    history.record_run(
        QCRunResult(profile_name="fixture", findings=findings),
        file_hashes={},
        report_paths={},
        on_subphase=on_subphase,
    )

    expected_names = {
        "main_pass",
        "story_classify_and_replay",
        "sqlite_write",
        "storage_measurement",
        "total",
    }
    assert set(observed) == expected_names
    assert all(elapsed >= 0.0 for elapsed in observed.values())
    # The subphases are strict subsets of the wall clock, not double-counted.
    assert observed["total"] >= observed["main_pass"]
    assert observed["total"] >= observed["story_classify_and_replay"]
    assert observed["total"] >= observed["sqlite_write"]
    assert observed["total"] >= observed["storage_measurement"]


def test_indexed_dossier_reads_positions_without_iterating_runs(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    history = RunHistory(tmp_path / "history.sqlite3")
    run_id = _record_longitudinal_run(history, "F0001")

    def reject_iteration(_self: FindingSequence):
        raise AssertionError("indexed dossier iterated the full run")

    monkeypatch.setattr(FindingSequence, "__iter__", reject_iteration)

    dossier = history.get_dossier(run_id, "F0001")

    assert [entry.finding_id for entry in dossier.entries] == ["F0001"]


def test_longitudinal_backfill_is_bounded_idempotent_and_labels_legacy_manual(
    tmp_path: Path,
) -> None:
    database = tmp_path / "history.sqlite3"
    history = RunHistory(database)
    run_id = _record_longitudinal_run(history, "F1")
    history.set_annotation(run_id, "F1", severity="info", comment="checked")
    with sqlite3.connect(database) as connection:
        connection.execute(
            "DELETE FROM decision_occurrences WHERE run_id = ?",
            (run_id,),
        )
        connection.execute(
            "DELETE FROM decision_origins WHERE run_id = ?",
            (run_id,),
        )

    assert history.backfill_longitudinal() == 2
    assert history.backfill_longitudinal() == 0
    dossier = history.get_dossier(run_id, "F1")
    assert dossier.entries[0].origin is DecisionOrigin.LEGACY_MANUAL
    with pytest.raises(ValueError, match="between 1 and 500"):
        history.backfill_longitudinal(limit=0)
    with pytest.raises(ValueError, match="between 1 and 500"):
        history.backfill_longitudinal(limit=501)


def test_dossier_distinguishes_changed_missing_and_ambiguous_observations(
    tmp_path: Path,
) -> None:
    history = RunHistory(tmp_path / "history.sqlite3")
    first = _record_longitudinal_run(history, "F1", value="1")
    second = _record_longitudinal_run(
        history,
        "F2",
        value="2",
        rerun_of=first,
    )
    missing = _record_longitudinal_run(
        history,
        "F3",
        location="A3",
        value="3",
        rerun_of=second,
    )

    changed_dossier = history.get_dossier(second, "F2")
    missing_dossier = history.get_dossier(missing, "F3")

    assert [entry.status for entry in changed_dossier.entries] == [
        DossierStatus.CHANGED,
        DossierStatus.EXACT,
    ]
    assert [entry.status for entry in missing_dossier.entries] == [
        DossierStatus.NO_STORED_OBSERVATION,
        DossierStatus.NO_STORED_OBSERVATION,
        DossierStatus.EXACT,
    ]

    ambiguous = history.record_run(
        QCRunResult(
            profile_name="fixture",
            findings=[
                _longitudinal_finding("A1", location="B1", value="10"),
                _longitudinal_finding("A2", location="B1", value="10"),
            ],
        ),
        file_hashes={},
        report_paths={},
        rerun_of=missing,
    )
    current = _record_longitudinal_run(
        history,
        "F4",
        location="B1",
        value="10",
        rerun_of=ambiguous,
    )
    ambiguous_dossier = history.get_dossier(current, "F4")
    assert ambiguous_dossier.entries[-2].status is DossierStatus.AMBIGUOUS


def test_recurrence_requires_three_finalized_exact_runs_and_two_manual_decisions(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    history = RunHistory(tmp_path / "history.sqlite3")
    run_ids: list[int] = []
    previous: int | None = None
    for index in range(3):
        run_id = _record_longitudinal_run(
            history,
            f"F{index}",
            rerun_of=previous,
        )
        history.set_annotation(
            run_id,
            f"F{index}",
            severity="info",
            comment=f"manual {index}",
        )
        _finalize(history, run_id)
        run_ids.append(run_id)
        previous = run_id

    eligible = history.recurrence_eligibility(run_ids[-1], "F2")

    assert eligible is not None
    assert eligible.run_ids == tuple(run_ids)
    assert eligible.manual_count == 3
    assert eligible.analyst_severity == "info"

    dossier = history.get_dossier(run_ids[-1], "F2")

    def reject_recomputation(*_args: object, **_kwargs: object):
        raise AssertionError("recurrence recomputed the dossier")

    monkeypatch.setattr(history, "get_dossier", reject_recomputation)
    reused = history.recurrence_eligibility_from_dossier(dossier)

    assert reused == eligible


def test_carried_decisions_do_not_inflate_recurrence_manual_count(
    tmp_path: Path,
) -> None:
    history = RunHistory(tmp_path / "history.sqlite3")
    first = _record_longitudinal_run(history, "F1")
    history.set_annotation(first, "F1", severity="info", comment="manual")
    _finalize(history, first)

    second = _record_longitudinal_run(history, "F2", rerun_of=first)
    assert apply_carry_forward(history, second, {"F2"}) == 1
    _finalize(history, second)

    third = _record_longitudinal_run(history, "F3", rerun_of=second)
    assert apply_carry_forward(history, third, {"F3"}) == 1
    _finalize(history, third)

    assert history.recurrence_eligibility(third, "F3") is None
    dossier = history.get_dossier(third, "F3")
    assert [entry.origin for entry in dossier.entries] == [
        DecisionOrigin.MANUAL,
        DecisionOrigin.CARRIED,
        DecisionOrigin.CARRIED,
    ]


def test_recurrence_rejects_contradictory_or_unfinalized_lineage(
    tmp_path: Path,
) -> None:
    history = RunHistory(tmp_path / "history.sqlite3")
    previous: int | None = None
    run_ids: list[int] = []
    for index, severity in enumerate(("info", "warning", "info")):
        run_id = _record_longitudinal_run(
            history,
            f"F{index}",
            rerun_of=previous,
        )
        history.set_annotation(
            run_id,
            f"F{index}",
            severity=severity,
            comment="manual",
        )
        run_ids.append(run_id)
        previous = run_id
    for run_id in run_ids[:2]:
        _finalize(history, run_id)

    assert history.recurrence_eligibility(run_ids[-1], "F2") is None
    _finalize(history, run_ids[-1])
    assert history.recurrence_eligibility(run_ids[-1], "F2") is None


def test_counterfactual_sidecar_round_trip_and_tamper_detection(
    tmp_path: Path,
) -> None:
    database = tmp_path / "history.sqlite3"
    history = RunHistory(database)
    finding = _longitudinal_finding("F1")
    basis = NumericCounterfactualBasis(
        baseline=1.0,
        current=2.0,
        number_format="0.0",
        sheet="Data",
        location="A1",
    )
    finding.counterfactual_basis = basis
    run_id = history.record_run(
        QCRunResult(profile_name="fixture", findings=[finding]),
        file_hashes={},
        report_paths={},
    )

    assert history.get_counterfactual_bases(run_id) == {"F1": basis}
    assert history.get_raw_run(run_id).findings[0].counterfactual_basis is None
    with sqlite3.connect(database) as connection:
        connection.execute(
            "UPDATE counterfactual_bases SET digest = 'bad' WHERE run_id = ?",
            (run_id,),
        )
    with pytest.raises(ValueError, match="counterfactual digest mismatch"):
        history.get_counterfactual_bases(run_id)


def test_package_archive_preserves_topology_without_paths(tmp_path: Path) -> None:
    work_dir = tmp_path / "work"
    runs_root = work_dir / "runs"
    report = runs_root / "run" / "report.html"
    report.parent.mkdir(parents=True)
    report.write_text("report", encoding="utf-8")
    files = {
        "current_excel:core": tmp_path / "client-core.xlsx",
        "current_excel:ops": tmp_path / "client-ops.xlsx",
    }
    for path in files.values():
        path.write_bytes(b"source")
    package_manifest = PackageManifest.from_role_files(files)
    result = QCRunResult(
        profile_name="package",
        files={role: path.name for role, path in files.items()},
        package_manifest=package_manifest,
    )
    history = RunHistory(work_dir / "history.sqlite3")
    history.record_run(
        result,
        file_hashes={role: sha256_file(path) for role, path in files.items()},
        file_paths={role: str(path) for role, path in files.items()},
        report_paths={"html": str(report)},
    )

    archive_path = export_runs_archive(
        history.list_runs(),
        work_dir / "exports" / "runs.zip",
        managed_root=runs_root,
    )
    with zipfile.ZipFile(archive_path) as archive:
        payload = json.loads(archive.read("manifest.json"))

    archived_run = payload["runs"][0]
    assert archived_run["package_manifest"] == package_manifest.model_dump(
        mode="json"
    )
    serialized = json.dumps(payload)
    assert str(tmp_path) not in serialized
    assert "file_paths" not in serialized


def test_delete_clears_auxiliary_evidence_but_never_source_files(
    tmp_path: Path,
) -> None:
    work_dir = tmp_path / "work"
    runs_root = work_dir / "runs"
    report = runs_root / "run" / "report.html"
    report.parent.mkdir(parents=True)
    report.write_text("report", encoding="utf-8")
    source = tmp_path / "source.xlsx"
    source.write_bytes(b"original source")
    finding = _longitudinal_finding("F1")
    finding.counterfactual_basis = NumericCounterfactualBasis(
        baseline=1.0,
        current=2.0,
        sheet="Data",
        location="A1",
    )
    finding.subtype = FindingSubtype.VALUE_REPLACEMENT
    finding.materiality = Materiality.MATERIAL
    finding.temporal_context = FindingTemporalContext.HISTORICAL
    finding.series_anchor = SeriesAnchorV1(
        sheet="Data",
        current_region_id="Data!A1:D9",
        period_axis="rows",
        series_index=1,
        period_index=1,
    )
    history = RunHistory(work_dir / "history.sqlite3")
    run_id = history.record_run(
        QCRunResult(profile_name="delete", findings=[finding]),
        file_hashes={"current_excel": sha256_file(source)},
        file_paths={"current_excel": str(source)},
        report_paths={"html": str(report)},
    )
    history.set_annotation(run_id, "F1", severity="info", comment="manual")
    history.record_promotion(run_id, "F1", "waiver", "profile-sha")
    history.start_review_session(run_id)
    history.pause_review_sessions()
    _finalize(history, run_id)

    assert history.delete_runs([run_id], managed_root=runs_root) == 1

    assert source.read_bytes() == b"original source"
    assert not report.exists()
    with sqlite3.connect(work_dir / "history.sqlite3") as connection:
        for table in (
            "runs",
            "annotations",
            "run_signoffs",
            "counterfactual_bases",
            "series_anchor_sidecars",
            "decision_occurrences",
            "decision_origins",
            "contract_promotions",
            "review_sessions",
        ):
            count = connection.execute(f"SELECT COUNT(*) FROM {table}").fetchone()
            assert count is not None and count[0] == 0


def test_six_cycle_lineage_oracle_is_deterministic_and_carry_safe(
    tmp_path: Path,
) -> None:
    history = RunHistory(tmp_path / "history.sqlite3")
    run_ids: list[int] = []
    previous: int | None = None
    carried_cycles = {2, 4}
    for cycle in range(1, 7):
        finding_id = f"F{cycle}"
        run_id = _record_longitudinal_run(
            history,
            finding_id,
            rerun_of=previous,
        )
        if cycle in carried_cycles:
            assert apply_carry_forward(history, run_id, {finding_id}) == 1
        else:
            history.set_annotation(
                run_id,
                finding_id,
                severity="info",
                comment=f"manual cycle {cycle}",
            )
        _finalize(history, run_id)
        run_ids.append(run_id)
        previous = run_id

        if cycle == 3:
            eligibility = history.recurrence_eligibility(run_id, finding_id)
            assert eligibility is not None
            assert eligibility.run_ids == tuple(run_ids[:3])
            assert eligibility.manual_count == 2
            history.record_promotion(run_id, finding_id, "waiver", "profile-sha")

    dossier = history.get_dossier(run_ids[-1], "F6")
    assert len(dossier.entries) == 6
    assert all(entry.status is DossierStatus.EXACT for entry in dossier.entries)
    assert [entry.origin for entry in dossier.entries] == [
        DecisionOrigin.MANUAL,
        DecisionOrigin.CARRIED,
        DecisionOrigin.MANUAL,
        DecisionOrigin.CARRIED,
        DecisionOrigin.MANUAL,
        DecisionOrigin.MANUAL,
    ]
    assert dossier.entries[1].carry_source_run_id == run_ids[0]
    assert dossier.entries[3].carry_source_run_id == run_ids[2]
    assert history.recurrence_eligibility(run_ids[3], "F4") is None
    assert history.recurrence_eligibility(run_ids[4], "F5") is None
    final = history.recurrence_eligibility(run_ids[5], "F6")
    assert final is not None
    assert final.run_ids == tuple(run_ids[3:])
    assert final.manual_count == 2


# --- private logical-series anchor sidecar ------------------------------------


def _anchored_finding(
    finding_id: str,
    location: str,
    *,
    member: str = "primary",
    sheet: str = "Data",
    region: str = "Data!A1:D9",
    segment: Literal[
        "restatement", "new_period", "cleared_period"
    ] = "restatement",
    legacy: bool = False,
) -> Finding:
    row, col = coordinate_to_tuple(location)
    finding = Finding(
        finding_id=finding_id,
        artifact="excel",
        artifact_member=member,
        finding_class=FindingClass.VALUE_CHANGED,
        severity=Severity.CRITICAL,
        subtype=(
            FindingSubtype.VALUE_REPLACEMENT
            if segment == "restatement"
            else (
                FindingSubtype.VALUE_ADDED_POPULATION
                if segment == "new_period"
                else FindingSubtype.VALUE_CLEARED_POPULATION
            )
        ),
        materiality=Materiality.MATERIAL if segment == "restatement" else None,
        temporal_context=(
            FindingTemporalContext.HISTORICAL if segment == "restatement" else None
        ),
        sheet=sheet,
        location=location,
        baseline_location=location,
        baseline_value="0" if segment != "new_period" else None,
        current_value=None if segment == "cleared_period" else "1",
        message="value changed",
    )
    if legacy:
        finding.series_anchor = SeriesAnchorV1(
            sheet=sheet,
            current_region_id=region,
            period_axis="rows",
            series_index=col,
            period_index=row,
        )
    else:
        finding.series_anchor = SeriesAnchorV2(
            sheet=sheet,
            current_region_id=region,
            period_axis="rows",
            series_index=col,
            period_index=row,
            segment=segment,
        )
    return finding


def _record_anchored_run(history: RunHistory, findings: list[Finding]) -> int:
    return history.record_run(
        QCRunResult(profile_name="series", findings=findings),
        file_hashes={},
        report_paths={},
    )


def test_series_anchor_sidecar_round_trips_and_stays_out_of_public_findings(
    tmp_path: Path,
) -> None:
    history = RunHistory(tmp_path / "history.sqlite3")
    findings = [_anchored_finding("F1", "B2"), _anchored_finding("F2", "B3")]
    unanchored = _longitudinal_finding("F3")
    run_id = _record_anchored_run(history, [*findings, unanchored])

    anchors = history.get_series_anchors(run_id)
    assert set(anchors) == {"F1", "F2"}
    assert anchors["F1"].period_index == 2
    assert anchors["F2"].period_index == 3
    assert all(anchor.series_index == 2 for anchor in anchors.values())

    record = history.get_run(run_id)
    assert len(record.series_anchor_digest) == 64
    assert all(finding.series_anchor is None for finding in record.findings)
    with sqlite3.connect(tmp_path / "history.sqlite3") as connection:
        payload = connection.execute(
            "SELECT findings FROM runs WHERE id = ?", (run_id,)
        ).fetchone()
    assert payload is not None and "series_anchor" not in payload[0]


def test_run_without_eligible_anchors_binds_an_empty_population(
    tmp_path: Path,
) -> None:
    history = RunHistory(tmp_path / "history.sqlite3")
    run_id = _record_anchored_run(history, [_longitudinal_finding("F1")])

    record = history.get_run(run_id)
    assert record.series_anchor_digest != ""
    assert history.get_series_anchors(run_id) == {}


def test_legacy_run_without_a_digest_reports_no_anchors(
    tmp_path: Path,
    caplog: pytest.LogCaptureFixture,
) -> None:
    history = RunHistory(tmp_path / "history.sqlite3")
    run_id = _record_anchored_run(history, [_anchored_finding("F1", "B2")])
    with sqlite3.connect(tmp_path / "history.sqlite3") as connection:
        connection.execute(
            "UPDATE runs SET series_anchor_digest = '' WHERE id = ?", (run_id,)
        )

    with caplog.at_level("WARNING"):
        assert history.get_series_anchors(run_id) == {}
    assert "legacy_run_with_rows" in caplog.text
    assert "Data!A1:D9" not in caplog.text


@pytest.mark.parametrize(
    ("statement", "code"),
    [
        (
            "UPDATE series_anchor_sidecars SET payload = "
            "'{\"version\":2,\"sheet\":\"Data\",\"current_region_id\":\"Data!A1:D9\","
            "\"period_axis\":\"rows\",\"series_index\":2,\"period_index\":9,"
            "\"segment\":\"restatement\"}'",
            "entry_digest_mismatch",
        ),
        ("UPDATE series_anchor_sidecars SET payload = 'not json'", "invalid_payload"),
        ("UPDATE series_anchor_sidecars SET version = 1", "version_mismatch"),
        ("UPDATE series_anchor_sidecars SET finding_id = 'F404'", "unknown_finding"),
        ("DELETE FROM series_anchor_sidecars", "missing_rows"),
    ],
)
def test_tampered_sidecar_fails_closed_with_a_fixed_code(
    tmp_path: Path,
    caplog: pytest.LogCaptureFixture,
    statement: str,
    code: str,
) -> None:
    history = RunHistory(tmp_path / "history.sqlite3")
    run_id = _record_anchored_run(history, [_anchored_finding("F1", "B2")])
    with sqlite3.connect(tmp_path / "history.sqlite3") as connection:
        connection.execute(statement)

    with caplog.at_level("WARNING"):
        assert history.get_series_anchors(run_id) == {}

    assert code in caplog.text
    assert "Data!A1:D9" not in caplog.text


def test_moved_finding_locator_invalidates_its_anchor(
    tmp_path: Path, caplog: pytest.LogCaptureFixture
) -> None:
    history = RunHistory(tmp_path / "history.sqlite3")
    finding = _anchored_finding("F1", "B2")
    run_id = _record_anchored_run(history, [finding])
    with sqlite3.connect(tmp_path / "history.sqlite3") as connection:
        stored = connection.execute(
            "SELECT blob FROM run_finding_blocks"
            " WHERE run_id = ? AND block_ordinal = 0",
            (run_id,),
        ).fetchone()
        assert stored is not None
        payload = decode_block(bytes(stored[0]))
        row = payload[0]
        assert isinstance(row, dict)
        row["location"] = "B7"
        connection.execute(
            "UPDATE run_finding_blocks SET blob = ?"
            " WHERE run_id = ? AND block_ordinal = 0",
            (encode_block(payload), run_id),
        )

    with caplog.at_level("WARNING"):
        assert history.get_series_anchors(run_id) == {}
    assert "entry_digest_mismatch" in caplog.text


def test_package_members_keep_isolated_anchor_bindings(tmp_path: Path) -> None:
    history = RunHistory(tmp_path / "history.sqlite3")
    run_id = _record_anchored_run(
        history,
        [
            _anchored_finding("F1", "B2", member="core"),
            _anchored_finding("F2", "B2", member="ops"),
        ],
    )

    anchors = history.get_series_anchors(run_id)
    assert set(anchors) == {"F1", "F2"}

    with sqlite3.connect(tmp_path / "history.sqlite3") as connection:
        digests = dict(
            connection.execute(
                "SELECT finding_id, digest FROM series_anchor_sidecars WHERE run_id = ?",
                (run_id,),
            ).fetchall()
        )
    assert digests["F1"] != digests["F2"]


def test_stored_v1_sidecar_still_returns_every_anchor(tmp_path: Path) -> None:
    history = RunHistory(tmp_path / "history.sqlite3")
    findings = [
        _anchored_finding("F1", "B2", legacy=True),
        _anchored_finding("F2", "B3", legacy=True),
    ]
    run_id = _record_anchored_run(history, findings)

    with sqlite3.connect(tmp_path / "history.sqlite3") as connection:
        versions = [
            row[0]
            for row in connection.execute(
                "SELECT version FROM series_anchor_sidecars WHERE run_id = ?",
                (run_id,),
            )
        ]
    assert versions == [1, 1]

    anchors = history.get_series_anchors(run_id)
    assert set(anchors) == {"F1", "F2"}
    assert all(anchor.version == 1 for anchor in anchors.values())
    assert all(anchor_segment(anchor) == "restatement" for anchor in anchors.values())


def test_new_period_anchor_round_trips_as_v2(tmp_path: Path) -> None:
    history = RunHistory(tmp_path / "history.sqlite3")
    run_id = _record_anchored_run(
        history,
        [
            _anchored_finding("F1", "B2"),
            _anchored_finding("F2", "B4", segment="new_period"),
        ],
    )

    anchors = history.get_series_anchors(run_id)
    assert anchor_segment(anchors["F1"]) == "restatement"
    assert anchor_segment(anchors["F2"]) == "new_period"
    assert anchors["F2"].period_index == 4


def test_cleared_period_anchor_round_trips_as_v2(tmp_path: Path) -> None:
    history = RunHistory(tmp_path / "history.sqlite3")
    run_id = _record_anchored_run(
        history,
        [
            _anchored_finding("F1", "B2"),
            _anchored_finding("F2", "B5", segment="cleared_period"),
        ],
    )

    anchors = history.get_series_anchors(run_id)

    assert anchor_segment(anchors["F1"]) == "restatement"
    assert anchor_segment(anchors["F2"]) == "cleared_period"
    assert anchors["F2"].version == 2
    assert anchors["F2"].period_index == 5


def test_segment_disagreeing_with_the_subtype_fails_closed(
    tmp_path: Path, caplog: pytest.LogCaptureFixture
) -> None:
    history = RunHistory(tmp_path / "history.sqlite3")
    run_id = _record_anchored_run(history, [_anchored_finding("F1", "B2")])
    with sqlite3.connect(tmp_path / "history.sqlite3") as connection:
        row = connection.execute(
            "SELECT payload FROM series_anchor_sidecars WHERE run_id = ?", (run_id,)
        ).fetchone()
        payload = json.loads(row[0])
        payload["segment"] = "new_period"
        connection.execute(
            "UPDATE series_anchor_sidecars SET payload = ? WHERE run_id = ?",
            (json.dumps(payload, sort_keys=True, separators=(",", ":")), run_id),
        )

    with caplog.at_level("WARNING"):
        assert history.get_series_anchors(run_id) == {}
    assert "entry_digest_mismatch" in caplog.text


def test_unknown_anchor_version_fails_closed(
    tmp_path: Path, caplog: pytest.LogCaptureFixture
) -> None:
    history = RunHistory(tmp_path / "history.sqlite3")
    run_id = _record_anchored_run(history, [_anchored_finding("F1", "B2")])
    with sqlite3.connect(tmp_path / "history.sqlite3") as connection:
        connection.execute(
            "UPDATE series_anchor_sidecars SET payload = ?, version = 9",
            ('{"version": 9, "sheet": "Data"}',),
        )

    with caplog.at_level("WARNING"):
        assert history.get_series_anchors(run_id) == {}
    assert "invalid_payload" in caplog.text


def test_history_archive_never_carries_the_series_sidecar(tmp_path: Path) -> None:
    work_dir = tmp_path / "work"
    runs_root = work_dir / "runs"
    report = runs_root / "run" / "report.html"
    report.parent.mkdir(parents=True)
    report.write_text("report", encoding="utf-8")
    history = RunHistory(work_dir / "history.sqlite3")
    history.record_run(
        QCRunResult(profile_name="series", findings=[_anchored_finding("F1", "B2")]),
        file_hashes={},
        report_paths={"html": str(report)},
    )

    archive_path = export_runs_archive(
        history.list_runs(),
        work_dir / "exports" / "runs.zip",
        managed_root=runs_root,
    )
    with zipfile.ZipFile(archive_path) as archive:
        blob = "".join(
            archive.read(name).decode("utf-8", "replace") for name in archive.namelist()
        )

    assert "series_anchor" not in blob
    assert "current_region_id" not in blob


def test_comment_free_same_severity_confirmation_reloads_as_reviewed(
    tmp_path: Path,
) -> None:
    history = RunHistory(tmp_path / "history.sqlite3")
    findings = [_anchored_finding("F1", "B2"), _anchored_finding("F2", "B3")]
    run_id = _record_anchored_run(history, findings)

    updates = cluster_confirmation_updates(
        (
            ReviewSlice(
                slice_id="G1~s",
                group_id="G1",
                members=tuple(findings),
                whole_group=True,
            ),
        ),
        "",
    )
    history.set_annotations_bulk(
        run_id,
        [(update.finding_id, update.severity, update.comment) for update in updates],
    )

    record = history.get_run(run_id)
    assert [update.severity for update in updates] == ["critical", "critical"]
    assert all(finding.severity_overridden for finding in record.findings)
    assert all(finding.severity is Severity.CRITICAL for finding in record.findings)
    assert all(finding.analyst_comment == "" for finding in record.findings)


def test_real_run_persists_producer_anchors_end_to_end(
    fixture_dir: Path, tmp_path: Path
) -> None:
    artifacts = perform_run(
        tmp_path / "work",
        {
            "baseline_excel": fixture_dir / "baseline.xlsx",
            "current_excel": fixture_dir / "current.xlsx",
        },
        {},
        fixture_profile(),
    )
    history = RunHistory(tmp_path / "work" / "history.sqlite3")

    anchors = history.get_series_anchors(artifacts.run_id)
    assert anchors
    by_id = {finding.finding_id: finding for finding in artifacts.result.findings}
    for finding_id, anchor in anchors.items():
        assert anchor_matches_finding(by_id[finding_id], anchor)
    # the same measure column spans at least two canonical decisions
    columns = {
        (anchor.sheet, anchor.current_region_id, anchor.series_index)
        for anchor in anchors.values()
    }
    assert len(columns) < len(anchors)


def test_new_runs_store_findings_in_compressed_blocks(tmp_path: Path) -> None:
    """Step 5: block rows replace the giant JSON column for new runs."""
    history = RunHistory(tmp_path / "history.sqlite3")
    findings = [_longitudinal_finding(f"F{index}") for index in range(1, 4)]
    run_id = history.record_run(
        QCRunResult(profile_name="blocks", findings=findings),
        file_hashes={},
        report_paths={},
    )

    with sqlite3.connect(tmp_path / "history.sqlite3") as connection:
        stored = connection.execute(
            "SELECT findings FROM runs WHERE id = ?", (run_id,)
        ).fetchone()
        blocks = connection.execute(
            "SELECT block_ordinal, first_ordinal, count FROM run_finding_blocks"
            " WHERE run_id = ? ORDER BY block_ordinal",
            (run_id,),
        ).fetchall()
    assert stored is not None and stored[0] == "[]"
    assert blocks == [(0, 0, 3)]

    record = history.get_run(run_id)
    assert [finding.finding_id for finding in record.findings] == ["F1", "F2", "F3"]
    assert list(record.findings) == findings


def test_delete_runs_removes_finding_block_rows(tmp_path: Path) -> None:
    history = RunHistory(tmp_path / "history.sqlite3")
    run_id = history.record_run(
        QCRunResult(
            profile_name="blocks",
            findings=[_longitudinal_finding("F1")],
        ),
        file_hashes={},
        report_paths={},
    )

    assert history.delete_runs([run_id], managed_root=tmp_path / "runs") == 1

    with sqlite3.connect(tmp_path / "history.sqlite3") as connection:
        remaining = connection.execute(
            "SELECT COUNT(*) FROM run_finding_blocks WHERE run_id = ?", (run_id,)
        ).fetchone()
    assert remaining is not None and remaining[0] == 0


def test_legacy_json_findings_rows_still_load(tmp_path: Path) -> None:
    """Rows recorded before block storage keep loading from runs.findings."""
    history = RunHistory(tmp_path / "history.sqlite3")
    findings = [_longitudinal_finding("F1"), _longitudinal_finding("F2")]
    run_id = history.record_run(
        QCRunResult(profile_name="legacy", findings=findings),
        file_hashes={},
        report_paths={},
    )
    with sqlite3.connect(tmp_path / "history.sqlite3") as connection:
        connection.execute(
            "DELETE FROM run_finding_blocks WHERE run_id = ?", (run_id,)
        )
        connection.execute(
            "UPDATE runs SET findings = ? WHERE id = ?",
            (
                json.dumps([finding.model_dump(mode="json") for finding in findings]),
                run_id,
            ),
        )

    record = history.get_run(run_id)
    assert isinstance(record.findings, list)
    assert list(record.findings) == findings


def test_view_summaries_round_trip_with_the_run(tmp_path: Path) -> None:
    """Step 2 (16 GB plan): pattern summaries + stories persist at record."""
    from qc_tool.review_stream import (
        stream_stories,
        summarize_pattern_groups_with_priority,
    )

    findings = [
        _longitudinal_finding(f"F{index:04d}", location=f"B{index}")
        for index in range(1, 8)
    ]
    result = QCRunResult(profile_name="summaries", findings=findings)
    history = RunHistory(tmp_path / "history.sqlite3")
    run_id = history.record_run(
        result, file_hashes={"current_excel": "c" * 64}, report_paths={}
    )
    stored = history.get_view_summaries(run_id)
    assert stored is not None
    summaries, aggregates, stories = stored
    expected_summaries, expected_aggregates = (
        summarize_pattern_groups_with_priority(result.findings)
    )
    assert summaries == expected_summaries
    assert aggregates == expected_aggregates
    assert stories == stream_stories(
        [iter(result.findings), iter(result.findings), iter(result.findings)]
    )
    # legacy rows (recorded before the table) yield None
    with sqlite3.connect(tmp_path / "history.sqlite3") as conn:
        conn.execute("DELETE FROM run_view_summaries WHERE run_id = ?", (run_id,))
    assert history.get_view_summaries(run_id) is None
    # unreadable payloads degrade to None, never raise
    with sqlite3.connect(tmp_path / "history.sqlite3") as conn:
        conn.execute(
            "INSERT INTO run_view_summaries (run_id, blob) VALUES (?, ?)",
            (run_id, b"garbage"),
        )
    assert history.get_view_summaries(run_id) is None


def test_reports_defer_past_the_threshold(
    fixture_dir: Path, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """On-demand exports: big runs record without report files."""
    import qc_tool.run_service as run_service

    monkeypatch.setattr(run_service, "REPORT_DEFER_FINDINGS", 0)
    work_dir = tmp_path / "work"
    artifacts = perform_run(
        work_dir,
        {
            "baseline_excel": fixture_dir / "baseline.xlsx",
            "current_excel": fixture_dir / "current.xlsx",
        },
        {},
        fixture_profile(),
    )
    assert artifacts.report_paths == {}
    assert not list((work_dir / "runs").glob("*/qc_report.*"))
    record = RunHistory(work_dir / "history.sqlite3").get_run(artifacts.run_id)
    assert record.report_paths == {}
    assert record.findings

    # the CLI-style override still writes eagerly at the same threshold
    eager = perform_run(
        tmp_path / "eager",
        {
            "baseline_excel": fixture_dir / "baseline.xlsx",
            "current_excel": fixture_dir / "current.xlsx",
        },
        {},
        fixture_profile(),
        write_reports=True,
    )
    assert eager.report_paths["html"].exists()
    assert eager.report_paths["excel"].exists()


def test_set_report_paths_attaches_generated_files(tmp_path: Path) -> None:
    history = RunHistory(tmp_path / "history.sqlite3")
    run_id = history.record_run(
        QCRunResult(
            profile_name="deferred",
            findings=[_longitudinal_finding("F0001")],
        ),
        file_hashes={},
        report_paths={},
    )
    report = tmp_path / "qc_report.html"
    report.write_text("<html></html>", encoding="utf-8")
    history.set_report_paths(run_id, {"html": str(report)})
    record = history.get_run(run_id)
    assert record.report_paths == {"html": str(report)}
    assert (record.storage_bytes or 0) > 0
    with pytest.raises(KeyError):
        history.set_report_paths(9999, {})


def test_volume_projection_flags_only_changed_sheets(
    fixture_dir: Path, tmp_path: Path
) -> None:
    """Identical parts project zero; a modified sheet projects its cells."""
    import shutil

    from openpyxl import load_workbook as opx_load

    from qc_tool.projection import project_cycle_volume

    baseline = fixture_dir / "baseline.xlsx"
    identical = tmp_path / "identical.xlsx"
    shutil.copy(baseline, identical)
    projection = project_cycle_volume(baseline, identical)
    assert projection is not None
    assert projection.projected_max_findings == 0
    assert not projection.changed_sheets
    assert projection.identical_sheets

    modified = tmp_path / "modified.xlsx"
    workbook = opx_load(baseline)
    sheet = workbook.worksheets[0]
    sheet["A1"] = "projection-probe"
    workbook.save(modified)
    projection = project_cycle_volume(baseline, modified)
    assert projection is not None
    assert projection.projected_max_findings > 0
    # the zip rewrite makes CRC comparisons conservative, but the probe
    # sheet must always be flagged
    assert sheet.title in projection.changed_sheets

    assert project_cycle_volume(baseline, tmp_path / "missing.xlsx") is None


def test_legacy_annotation_lineage_migrates_to_versioned_many_to_one(
    tmp_path: Path,
) -> None:
    """A pre-A4 strictly-1:1 `annotation_lineage` row survives the table
    rebuild as one `relation=identity`, `outcome=inherited` row under the
    new many-to-one shape -- a faithful reinterpretation, not data loss.
    """
    db_path = tmp_path / "history.sqlite3"
    with sqlite3.connect(db_path) as conn:
        conn.execute(
            """
            CREATE TABLE annotation_lineage (
                run_id INTEGER NOT NULL,
                finding_id TEXT NOT NULL,
                source_run_id INTEGER NOT NULL,
                source_finding_id TEXT NOT NULL,
                evidence_version INTEGER NOT NULL,
                evidence_digest TEXT NOT NULL,
                applied_at TEXT NOT NULL,
                PRIMARY KEY (run_id, finding_id)
            )
            """
        )
        conn.execute(
            "INSERT INTO annotation_lineage VALUES (?, ?, ?, ?, ?, ?, ?)",
            (2, "N1", 1, "F1", 1, "a" * 64, "2026-01-01T00:00:00+00:00"),
        )
        conn.commit()

    history = RunHistory(db_path)

    lineage = history.get_annotation_lineage(2)["N1"]
    assert len(lineage) == 1
    row = lineage[0]
    assert row.source_run_id == 1
    assert row.source_finding_id == "F1"
    assert row.relation.value == "identity"
    assert row.outcome.value == "inherited"
    assert row.evidence_version == 1
    assert row.source_digest == "a" * 64
    assert row.applied_at == "2026-01-01T00:00:00+00:00"

    # Reopening an already-migrated database is a no-op, not a data loss.
    reopened = RunHistory(db_path)
    assert reopened.get_annotation_lineage(2)["N1"] == lineage
