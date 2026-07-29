"""Fixture-generator verification: determinism, seeded defects, ground truth.

Every seeded defect recorded in the manifest must be physically present in
the generated artifacts — these tests are the contract the QC engines will
be tested against in later plan steps.
"""

import io
from pathlib import Path

from msoffcrypto.format.ooxml import OOXMLFile
from openpyxl import load_workbook
from pptx import Presentation
from pyxlsb import open_workbook

from tests.fixtures import domain
from tests.fixtures.generate import generate
from tests.fixtures.manifest_schema import FixtureManifest

PLAINTEXT_FILES = [
    "baseline.xlsx",
    "current.xlsx",
    "baseline.xlsb",
    "current.xlsb",
    "baseline.pptx",
    "current.pptx",
    "manifest.json",
]


def _slide_by_title(prs, title: str):
    for slide in prs.slides:
        shape = slide.shapes.title
        if shape is not None and shape.text == title:
            return slide
    raise AssertionError(f"no slide titled {title!r}")


def _slide_texts(slide) -> list[str]:
    texts: list[str] = []
    for shape in slide.shapes:
        if shape.has_text_frame:
            for paragraph in shape.text_frame.paragraphs:
                texts.append("".join(run.text for run in paragraph.runs))
    return texts


def _diff_zip_parts(path_a: Path, path_b: Path) -> str:
    """Explain the first differing part of two zip files (for flake diagnosis)."""
    import zipfile

    with zipfile.ZipFile(path_a) as za, zipfile.ZipFile(path_b) as zb:
        names_a, names_b = za.namelist(), zb.namelist()
        if names_a != names_b:
            return f"part lists differ: {set(names_a) ^ set(names_b)}"
        for name in names_a:
            data_a, data_b = za.read(name), zb.read(name)
            if data_a != data_b:
                offset = next(
                    i for i, (x, y) in enumerate(zip(data_a, data_b, strict=False)) if x != y
                )
                context_a = data_a[max(0, offset - 40) : offset + 40]
                context_b = data_b[max(0, offset - 40) : offset + 40]
                return (
                    f"part {name!r} differs at byte {offset}: "
                    f"{context_a!r} != {context_b!r}"
                )
    return "zip containers identical; raw bytes differ (zip metadata?)"


def test_deterministic_output(tmp_path: Path) -> None:
    dir_a, dir_b = tmp_path / "a", tmp_path / "b"
    generate(dir_a)
    generate(dir_b)
    for name in PLAINTEXT_FILES:
        bytes_a, bytes_b = (dir_a / name).read_bytes(), (dir_b / name).read_bytes()
        if bytes_a != bytes_b and name != "manifest.json":
            raise AssertionError(f"{name}: {_diff_zip_parts(dir_a / name, dir_b / name)}")
        assert bytes_a == bytes_b, name


def test_manifest_schema_roundtrip(fixture_dir: Path, manifest: FixtureManifest) -> None:
    assert manifest.schema_version == 1
    for name in manifest.files.values():
        assert (fixture_dir / name).exists(), name
    ids = [d.defect_id for d in manifest.defects]
    assert len(ids) == len(set(ids)), "defect ids must be unique"
    assert manifest.defect("E01").sheet == "Long_Monthly"


def test_workbook_growth_and_value_seeds(fixture_dir: Path, manifest: FixtureManifest) -> None:
    base = load_workbook(fixture_dir / "baseline.xlsx")
    curr = load_workbook(fixture_dir / "current.xlsx")

    assert base["Long_Monthly"].max_row == 21
    assert curr["Long_Monthly"].max_row == 25

    e01 = manifest.defect("E01")
    assert base["Long_Monthly"]["C7"].value == float(e01.baseline or "")
    assert curr["Long_Monthly"]["C7"].value == float(e01.current or "")

    e06 = manifest.defect("E06")
    assert e06.cell is not None and e06.baseline_cell is not None
    assert base["Wide_Weekly"][e06.baseline_cell].value == float(e06.baseline or "")
    assert curr["Wide_Weekly"][e06.cell].value == float(e06.current or "")

    base_weeks = [c.value for c in base["Wide_Weekly"][1][1:]]
    curr_weeks = [c.value for c in curr["Wide_Weekly"][1][1:]]
    assert "W03" in base_weeks and "W03" not in curr_weeks  # E07
    assert "W21" not in base_weeks and "W21" in curr_weeks  # EX02


def test_workbook_formula_seeds(fixture_dir: Path, manifest: FixtureManifest) -> None:
    base = load_workbook(fixture_dir / "baseline.xlsx")
    curr = load_workbook(fixture_dir / "current.xlsx")

    assert base["Long_Monthly"]["E10"].value == "=C10-D10"
    e02 = curr["Long_Monthly"]["E10"]
    assert isinstance(e02.value, float)  # hardcoded constant

    assert curr["Long_Monthly"]["E14"].value == manifest.defect("E03").current
    assert curr["Long_Monthly"]["E22"].value == "=C22-D22"
    assert curr["Long_Monthly"]["E25"].value is None  # E04

    assert curr["Summary"]["C5"].value == "#REF!"  # E05
    assert curr["Summary"]["C5"].data_type == "e"
    assert curr["Summary"]["D5"].value == "=SUM(#REF!)"  # E13

    e08 = manifest.defect("E08")
    assert e08.cell is not None
    assert curr["Wide_Weekly"][e08.cell].value == e08.current

    # EX06: expected range extension, not a logic change.
    assert base["Summary"]["B2"].value == "=SUM(Long_Monthly!C2:C21)"
    assert curr["Summary"]["B2"].value == "=SUM(Long_Monthly!C2:C25)"


def test_workbook_structure_seeds(fixture_dir: Path, manifest: FixtureManifest) -> None:
    base = load_workbook(fixture_dir / "baseline.xlsx")
    curr = load_workbook(fixture_dir / "current.xlsx")

    assert "Old_Sheet" in base.sheetnames and "Old_Sheet" not in curr.sheetnames  # E14
    assert "New_Analysis" not in base.sheetnames and "New_Analysis" in curr.sheetnames  # E15
    assert base["Params"].sheet_state == "visible"
    assert curr["Params"].sheet_state == "hidden"  # E11

    assert base.defined_names["KPI_Margin"].attr_text == "Dashboard!$B$4"
    assert curr.defined_names["KPI_Margin"].attr_text == "Dashboard!$B$5"  # E12
    assert base.defined_names["RevenueData"].attr_text == "Long_Monthly!$A$1:$E$21"
    assert curr.defined_names["RevenueData"].attr_text == "Long_Monthly!$A$1:$E$25"  # EX04

    assert base["Dashboard"]["B4"].number_format == "0.0%"
    assert curr["Dashboard"]["B4"].number_format == "0.00"  # E09
    assert base["Dashboard"]["A2"].fill.fgColor.rgb == "FFFFFF00"
    assert curr["Dashboard"]["A2"].fill.fgColor.rgb == "FFFF0000"  # E18

    def series_ref(wb) -> str:
        chart = wb["Dashboard"]._charts[0]
        return chart.series[0].val.numRef.f

    assert series_ref(base) == manifest.defect("E16").baseline
    assert series_ref(curr) == manifest.defect("E16").current


def test_pivot_parts_injected(fixture_dir: Path, manifest: FixtureManifest) -> None:
    import zipfile

    e17 = manifest.defect("E17")
    for name, expected_ref in (("baseline.xlsx", e17.baseline), ("current.xlsx", e17.current)):
        with zipfile.ZipFile(fixture_dir / name) as zf:
            names = zf.namelist()
            assert "xl/pivotTables/pivotTable1.xml" in names
            cache = zf.read("xl/pivotCache/pivotCacheDefinition1.xml").decode("utf-8")
            assert f'ref="{expected_ref}"' in cache
            assert 'sheet="Long_Monthly"' in cache


def test_encrypted_workbook_roundtrip(fixture_dir: Path, manifest: FixtureManifest) -> None:
    decrypted = io.BytesIO()
    with (fixture_dir / "current_encrypted.xlsx").open("rb") as fin:
        office_file = OOXMLFile(fin)
        assert office_file.is_encrypted()
        office_file.load_key(password=manifest.password)
        office_file.decrypt(decrypted)
    assert decrypted.getvalue() == (fixture_dir / "current.xlsx").read_bytes()
    decrypted.seek(0)
    wb = load_workbook(decrypted)
    assert wb["Long_Monthly"].max_row == 25


def test_xlsb_pair_readback(fixture_dir: Path, manifest: FixtureManifest) -> None:
    def read_rows(name: str) -> list[list[object]]:
        with open_workbook(str(fixture_dir / name)) as wb:
            assert wb.sheets == ["Long_Monthly"]
            with wb.get_sheet(1) as sheet:
                return [[c.v for c in row] for row in sheet.rows()]

    base_rows = read_rows("baseline.xlsb")
    curr_rows = read_rows("current.xlsb")
    assert base_rows[0][:5] == ["Period", "Region", "Revenue", "Cost", "Margin"]
    assert len(base_rows) == 21 and len(curr_rows) == 25

    xb01 = manifest.defect("XB01")
    assert base_rows[6][2] == float(xb01.baseline or "")  # C7 -> row 6, col 2
    assert curr_rows[6][2] == float(xb01.current or "")


def test_deck_slide_seeds(fixture_dir: Path, manifest: FixtureManifest) -> None:
    base = Presentation(str(fixture_dir / "baseline.pptx"))
    curr = Presentation(str(fixture_dir / "current.pptx"))

    def titles(prs) -> list[str]:
        result = []
        for slide in prs.slides:
            shape = slide.shapes.title
            result.append(shape.text if shape is not None else "")
        return result

    base_titles, curr_titles = titles(base), titles(curr)
    assert "Deep Dive Archive" in base_titles and "Deep Dive Archive" not in curr_titles  # P02
    assert "New Initiatives" not in base_titles and "New Initiatives" in curr_titles  # P01
    assert base_titles.index("Notes & Definitions") != curr_titles.index(
        "Notes & Definitions"
    )  # PX01

    p03 = manifest.defect("P03")
    assert p03.baseline in _slide_texts(_slide_by_title(base, "Executive Summary"))
    assert p03.current in _slide_texts(_slide_by_title(curr, "Executive Summary"))


def test_deck_table_and_chart_seeds(fixture_dir: Path, manifest: FixtureManifest) -> None:
    base = Presentation(str(fixture_dir / "baseline.pptx"))
    curr = Presentation(str(fixture_dir / "current.pptx"))

    def table_cell(prs, row: int, col: int) -> str:
        slide = _slide_by_title(prs, "Revenue by Region")
        for shape in slide.shapes:
            if shape.has_table:
                return shape.table.cell(row, col).text
        raise AssertionError("no table found")

    p04 = manifest.defect("P04")
    assert table_cell(base, 1, 4) == p04.baseline  # North / Apr-26
    assert table_cell(curr, 1, 4) == p04.current
    assert table_cell(curr, 0, 6) == "Jun-26"  # PX04

    def chart_values(prs, title: str) -> tuple[list[str], list[float]]:
        slide = _slide_by_title(prs, title)
        for shape in slide.shapes:
            if shape.has_chart:
                plot = shape.chart.plots[0]
                cats = [str(c) for c in plot.categories]
                vals = [float(v) for v in plot.series[0].values]
                return cats, vals
        raise AssertionError(f"no chart on slide {title!r}")

    base_cats, base_vals = chart_values(base, "Revenue Trend")
    curr_cats, curr_vals = chart_values(curr, "Revenue Trend")
    p05 = manifest.defect("P05")
    mar = base_cats.index("Mar-26")
    assert base_vals[mar] == float(p05.baseline or "")
    assert curr_vals[mar] == float(p05.current or "")
    assert curr_cats[-1] == "Jun-26"  # PX03

    base_weeks, _ = chart_values(base, "Weekly Ops")
    curr_weeks, _ = chart_values(curr, "Weekly Ops")
    assert base_weeks == ["W17", "W18", "W19", "W20"]  # PX02
    assert curr_weeks == ["W18", "W19", "W20", "W21"]


def test_crosscheck_ground_truth(fixture_dir: Path, manifest: FixtureManifest) -> None:
    curr = load_workbook(fixture_dir / "current.xlsx", data_only=False)

    revenue_entry = next(e for e in manifest.crosscheck_map if e.figure_label == "Total revenue")
    assert revenue_entry.matches
    workbook_value = curr["Dashboard"]["B2"].value
    assert isinstance(workbook_value, int | float)
    assert domain.fmt_millions(float(workbook_value)) == revenue_entry.figure_text

    margin_entry = next(e for e in manifest.crosscheck_map if e.figure_label == "Margin")
    assert not margin_entry.matches and margin_entry.defect_id == "X03"
    margin_value = curr["Dashboard"]["B4"].value
    assert isinstance(margin_value, int | float)
    assert domain.fmt_pct(float(margin_value)) != margin_entry.figure_text

    cell_entry = next(e for e in manifest.crosscheck_map if e.figure_label == "North / Jan-26")
    assert cell_entry.matches
    lm_value = curr["Long_Monthly"]["C2"].value
    assert isinstance(lm_value, int | float)
    assert f"{lm_value:,.0f}" == cell_entry.figure_text
