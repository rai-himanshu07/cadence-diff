"""End-to-end NiceGUI tests for the full-page mode-aware configuration
workspace (plan-20260913, Step 7): the ``/configure`` route, auto-start
scan, region rendering, and bulk confirmation, exercised through the real
``qc_tool.ui.app.create_pages()`` page registration and a real (tiny)
subprocess-based setup scan -- not a mock.
"""

from __future__ import annotations

import asyncio
import json
import threading
from pathlib import Path

import pytest
from nicegui import events, ui
from nicegui.testing import User
from openpyxl import Workbook

from qc_tool.config.resolved_input import (
    ResolvedColumn,
    ResolvedInputConfigurationV1,
    ResolvedMember,
    ResolvedRegion,
    ResolvedSelector,
    ResolvedSheet,
)
from qc_tool.coverage import QCRunMode
from qc_tool.excel.regions import TableRegion
from qc_tool.history.config_session import ConfigSessionStore, session_key_for
from qc_tool.history.run_state import RunStateStore
from qc_tool.history.store import sha256_file
from qc_tool.projection import VolumeProjection
from qc_tool.setup import coordinator as setup_coordinator_module
from qc_tool.setup.coordinator import SetupStatus, get_setup_coordinator
from qc_tool.setup.models import DetectedRegion, MemberSetupProfile, SheetSetupProfile
from qc_tool.setup.preview_worker import PreviewWindowOutcome
from qc_tool.ui import app as app_module
from qc_tool.ui import config_workspace as config_workspace_module
from qc_tool.ui.app import create_pages
from qc_tool.ui.config_review import member_review_from_scan
from qc_tool.ui.config_workspace import (
    _apply_pending_row_matching,
    _credentials_for_unchanged_sources,
    build_session_choices,
    configuration_export_bytes,
    session_choice_overrides_from_resolved,
)
from tests.fixtures.generate import encrypt_file

pytest_plugins = ["nicegui.testing.user_plugin"]

#: A real subprocess spawn is far slower than should_see's default ~0.3s
#: retry budget; 50 retries (~5s) matches this project's own established
#: convention for waiting on real subprocess-backed UI updates.
_SUBPROCESS_RETRIES = 50


def _emit(element: ui.element, event_type: str, args: object = None) -> None:
    for listener in element._event_listeners.values():
        if listener.type == event_type and listener.handler is not None:
            with element.parent_slot or element.client.layout.default_slot:
                listener.handler(
                    events.GenericEventArguments(
                        sender=element,
                        client=element.client,
                        args={} if args is None else args,
                    )
                )
            return
    raise AssertionError(f"no {event_type!r} listener on {element}")


def _write_simple_workbook(path: Path, *, final_value: int = 30) -> None:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Data"
    sheet.append(["ID", "Value"])
    sheet.append(["A0", 10])
    sheet.append(["A1", 20])
    sheet.append(["A2", final_value])
    workbook.save(path)


def _member_profile_for_pending() -> MemberSetupProfile:
    regions = (
        DetectedRegion(
            region=TableRegion("Data", 1, 1, 2, 2, "block", 1, 1, "none")
        ),
        DetectedRegion(
            region=TableRegion("Data", 4, 1, 8, 2, "block", 4, 1, "none")
        ),
    )
    return MemberSetupProfile(
        member_id="primary",
        baseline_hash="a" * 64,
        current_hash="b" * 64,
        baseline_sheets=(SheetSetupProfile(sheet_name="Data", regions=regions),),
        current_sheets=(SheetSetupProfile(sheet_name="Data", regions=regions),),
    )


def test_file_replacement_retains_only_unchanged_source_credentials() -> None:
    credentials = {
        "baseline_excel": "baseline-secret",
        "current_excel": "current-secret",
    }
    previous_hashes = {
        "baseline_excel": "a" * 64,
        "current_excel": "b" * 64,
    }
    current_hashes = {
        "baseline_excel": "a" * 64,
        "current_excel": "c" * 64,
    }

    assert _credentials_for_unchanged_sources(
        credentials, previous_hashes, current_hashes
    ) == {"baseline_excel": "baseline-secret"}


def test_resolved_configuration_recovers_run_only_regions_and_selectors() -> None:
    resolved = ResolvedInputConfigurationV1(
        mode=QCRunMode.CYCLE_COMPARISON,
        members=(
            ResolvedMember(
                member_id="primary",
                sheets=(
                    ResolvedSheet(
                        sheet_id="data",
                        baseline_sheet_name="Data old",
                        current_sheet_name="Data",
                        regions=(
                            ResolvedRegion(
                                region_id="data_a1_b4",
                                mode="keyed",
                                current_outer_range="A1:B4",
                                baseline_outer_range="C1:D4",
                                columns=(
                                    ResolvedColumn(
                                        column_id="id",
                                        baseline_letter="C",
                                        current_letter="A",
                                        alignment_role="identity",
                                    ),
                                ),
                                coverage="confirmed",
                            ),
                        ),
                        selectors=(
                            ResolvedSelector(
                                selector_id="scenario",
                                baseline_cell="B9",
                                current_cell="B5",
                            ),
                        ),
                    ),
                ),
            ),
        ),
        warnings_acknowledged=("existing-warning",),
    )

    choices = session_choice_overrides_from_resolved(resolved)

    region = choices["region_decisions"]["data_a1_b4"]  # type: ignore[index]
    assert region["mode"] == "keyed"
    assert region["identity_columns"] == ["A"]
    assert region["column_baseline_letters"] == [["A", "C"]]
    selectors = choices["selectors"]["primary"]  # type: ignore[index]
    assert selectors == [
        {
            "sheet_name": "Data",
            "selector_id": "scenario",
            "label": "scenario",
            "cell": "B5",
            "baseline_cell": "B9",
        }
    ]
    assert choices["sheet_renames"] == {"primary": {"Data": "Data old"}}
    assert choices["warnings_acknowledged"] == ["existing-warning"]


def test_pending_row_matching_updates_only_the_target_region() -> None:
    member = member_review_from_scan("primary", _member_profile_for_pending())
    first, second = member.current_sheets[0].regions

    updated = _apply_pending_row_matching(
        member,
        [
            {
                "member_id": "primary",
                "sheet": "Data",
                "anchor_cell": second.anchor_cell,
                "current_range": second.current_range,
                "header_row": 4,
                "identity_columns": ["A"],
                "ordinal_columns": ["B"],
                "duplicate_policy": "occurrence",
            }
        ],
    )

    unchanged, proposed = updated.current_sheets[0].regions
    assert unchanged == first
    assert proposed.mode == "keyed"
    assert proposed.identity_columns == ("A",)
    assert proposed.ordinal_columns == ("B",)
    assert proposed.duplicate_key_policy == "occurrence"
    assert proposed.first_data_row == 5
    assert proposed.confirmed
    assert not proposed.ranked_candidate_pending


def _stage_session(work_dir: Path) -> str:
    files = {
        "baseline_excel": work_dir / "uploads" / "baseline_excel" / "baseline.xlsx",
        "current_excel": work_dir / "uploads" / "current_excel" / "current.xlsx",
    }
    for role, path in files.items():
        path.parent.mkdir(parents=True, exist_ok=True)
        _write_simple_workbook(
            path, final_value=31 if role == "current_excel" else 30
        )
    file_hashes = {role: sha256_file(path) for role, path in files.items()}
    choices = build_session_choices(
        mode=QCRunMode.CYCLE_COMPARISON,
        profile_name="default",
        files={role: str(path) for role, path in files.items()},
        file_hashes=file_hashes,
        output_mode="decision",
        allow_large_workbooks=False,
        allow_dependency_indexing=False,
        acceptance_absolute=0.0,
        acceptance_percent=0.0,
        rerun_of=None,
    )
    session_key = session_key_for(file_hashes)
    ConfigSessionStore(work_dir / "history.sqlite3").save_choices(
        session_key, profile_name="default", choices=choices
    )
    return session_key


def _stage_multi_region_session(work_dir: Path) -> str:
    files = {
        "baseline_excel": work_dir / "uploads" / "baseline_excel" / "baseline.xlsx",
        "current_excel": work_dir / "uploads" / "current_excel" / "current.xlsx",
    }
    for role, path in files.items():
        path.parent.mkdir(parents=True, exist_ok=True)
        workbook = Workbook()
        data = workbook.active
        assert data is not None
        data.title = "Data"
        data.append(["ID", "Value"])
        data.append(["A0", 10])
        data.append(["A1", 20])
        data.append([None, None])
        data.append(["Code", "Amount"])
        data.append(["B0", 30])
        data.append(["B1", 41 if role == "current_excel" else 40])
        summary = workbook.create_sheet("Summary")
        summary.append(["Metric", "Value"])
        summary.append(["Total", 71 if role == "current_excel" else 70])
        workbook.save(path)
    file_hashes = {role: sha256_file(path) for role, path in files.items()}
    choices = build_session_choices(
        mode=QCRunMode.CYCLE_COMPARISON,
        profile_name="default",
        files={role: str(path) for role, path in files.items()},
        file_hashes=file_hashes,
        output_mode="decision",
        allow_large_workbooks=False,
        allow_dependency_indexing=False,
        acceptance_absolute=0.0,
        acceptance_percent=0.0,
        rerun_of=None,
    )
    session_key = session_key_for(file_hashes)
    ConfigSessionStore(work_dir / "history.sqlite3").save_choices(
        session_key, profile_name="default", choices=choices
    )
    return session_key


def _write_workbook_with_ids(path: Path, ids: list[str]) -> None:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Data"
    sheet.append(["ID", "Value"])
    for index, identity in enumerate(ids):
        sheet.append([identity, index * 10])
    workbook.save(path)


def _stage_disjoint_key_session(work_dir: Path) -> str:
    """A baseline/current pair whose 'ID' column shares NO values at all --
    a confirmed identity on column A would measure a 0% overlap ratio
    (Step 12 Fix 5's low-key-overlap acknowledgement scenario).
    """
    files = {
        "baseline_excel": work_dir / "uploads" / "baseline_excel" / "baseline.xlsx",
        "current_excel": work_dir / "uploads" / "current_excel" / "current.xlsx",
    }
    files["baseline_excel"].parent.mkdir(parents=True, exist_ok=True)
    files["current_excel"].parent.mkdir(parents=True, exist_ok=True)
    _write_workbook_with_ids(files["baseline_excel"], ["A0", "A1", "A2"])
    _write_workbook_with_ids(files["current_excel"], ["B0", "B1", "B2"])
    file_hashes = {role: sha256_file(path) for role, path in files.items()}
    choices = build_session_choices(
        mode=QCRunMode.CYCLE_COMPARISON,
        profile_name="default",
        files={role: str(path) for role, path in files.items()},
        file_hashes=file_hashes,
        output_mode="decision",
        allow_large_workbooks=False,
        allow_dependency_indexing=False,
        acceptance_absolute=0.0,
        acceptance_percent=0.0,
        rerun_of=None,
    )
    session_key = session_key_for(file_hashes)
    ConfigSessionStore(work_dir / "history.sqlite3").save_choices(
        session_key, profile_name="default", choices=choices
    )
    return session_key


def _stage_moved_identity_key_session(work_dir: Path) -> str:
    """A baseline/current pair where the identity values current column A
    actually shares are sitting in baseline column C, not baseline column
    A -- proves a baseline-letter override on an identity column actually
    RECOMPUTES the key-overlap ratio (not just the initial computation
    from setting identity_columns) rather than leaving a stale ratio from
    before the override was entered.
    """
    files = {
        "baseline_excel": work_dir / "uploads" / "baseline_excel" / "baseline.xlsx",
        "current_excel": work_dir / "uploads" / "current_excel" / "current.xlsx",
    }
    files["baseline_excel"].parent.mkdir(parents=True, exist_ok=True)
    files["current_excel"].parent.mkdir(parents=True, exist_ok=True)

    baseline_wb = Workbook()
    baseline_sheet = baseline_wb.active
    assert baseline_sheet is not None
    baseline_sheet.title = "Data"
    baseline_sheet.append(["ID", "Value", "ID"])
    for index, identity in enumerate(("A0", "A1", "A2")):
        baseline_sheet.append([identity, index * 10, f"B{index}"])
    baseline_wb.save(files["baseline_excel"])

    current_wb = Workbook()
    current_sheet = current_wb.active
    assert current_sheet is not None
    current_sheet.title = "Data"
    current_sheet.append(["ID", "Value"])
    for index, identity in enumerate(("B0", "B1", "B2")):
        current_sheet.append([identity, index * 10])
    current_wb.save(files["current_excel"])

    file_hashes = {role: sha256_file(path) for role, path in files.items()}
    choices = build_session_choices(
        mode=QCRunMode.CYCLE_COMPARISON,
        profile_name="default",
        files={role: str(path) for role, path in files.items()},
        file_hashes=file_hashes,
        output_mode="decision",
        allow_large_workbooks=False,
        allow_dependency_indexing=False,
        acceptance_absolute=0.0,
        acceptance_percent=0.0,
        rerun_of=None,
    )
    session_key = session_key_for(file_hashes)
    ConfigSessionStore(work_dir / "history.sqlite3").save_choices(
        session_key, profile_name="default", choices=choices
    )
    return session_key


def _write_simple_deck(path: Path, titles: list[str]) -> None:
    from pptx import Presentation

    deck = Presentation()
    layout = deck.slide_layouts[0]
    for title in titles:
        slide = deck.slides.add_slide(layout)
        title_shape = slide.shapes.title
        assert title_shape is not None
        title_shape.text = title
    deck.save(str(path))


def _stage_session_with_ppt(work_dir: Path) -> str:
    files = {
        "baseline_excel": work_dir / "uploads" / "baseline_excel" / "baseline.xlsx",
        "current_excel": work_dir / "uploads" / "current_excel" / "current.xlsx",
        "baseline_ppt": work_dir / "uploads" / "baseline_ppt" / "baseline.pptx",
        "current_ppt": work_dir / "uploads" / "current_ppt" / "current.pptx",
    }
    for role, path in files.items():
        path.parent.mkdir(parents=True, exist_ok=True)
        if role.endswith("_ppt"):
            continue
        _write_simple_workbook(path)
    _write_simple_deck(files["baseline_ppt"], ["Cover", "Old Summary"])
    _write_simple_deck(files["current_ppt"], ["Cover", "New Summary"])
    file_hashes = {role: sha256_file(path) for role, path in files.items()}
    choices = build_session_choices(
        mode=QCRunMode.CYCLE_COMPARISON,
        profile_name="default",
        files={role: str(path) for role, path in files.items()},
        file_hashes=file_hashes,
        output_mode="decision",
        allow_large_workbooks=False,
        allow_dependency_indexing=False,
        acceptance_absolute=0.0,
        acceptance_percent=0.0,
        rerun_of=None,
    )
    session_key = session_key_for(file_hashes)
    ConfigSessionStore(work_dir / "history.sqlite3").save_choices(
        session_key, profile_name="default", choices=choices
    )
    return session_key


def _stage_ppt_only_preflight_session(work_dir: Path) -> str:
    current_ppt = work_dir / "uploads" / "current_ppt" / "current.pptx"
    current_ppt.parent.mkdir(parents=True, exist_ok=True)
    _write_simple_deck(current_ppt, ["Cover", "Summary"])
    files = {"current_ppt": current_ppt}
    file_hashes = {role: sha256_file(path) for role, path in files.items()}
    choices = build_session_choices(
        mode=QCRunMode.CURRENT_FILE_PREFLIGHT,
        profile_name="default",
        files={role: str(path) for role, path in files.items()},
        file_hashes=file_hashes,
        output_mode="decision",
        allow_large_workbooks=False,
        allow_dependency_indexing=False,
        acceptance_absolute=0.0,
        acceptance_percent=0.0,
        rerun_of=None,
    )
    session_key = session_key_for(file_hashes)
    ConfigSessionStore(work_dir / "history.sqlite3").save_choices(
        session_key, profile_name="default", choices=choices
    )
    return session_key


def _stage_encrypted_session(work_dir: Path) -> str:
    baseline = work_dir / "uploads" / "baseline_excel" / "baseline.xlsx"
    current = work_dir / "uploads" / "current_excel" / "current.xlsx"
    current_plain = work_dir / "current-plain.xlsx"
    baseline.parent.mkdir(parents=True, exist_ok=True)
    current.parent.mkdir(parents=True, exist_ok=True)
    _write_simple_workbook(baseline)
    _write_simple_workbook(current_plain, final_value=31)
    encrypt_file(current_plain, current, "hunter2")
    current_plain.unlink()
    files = {"baseline_excel": baseline, "current_excel": current}
    file_hashes = {role: sha256_file(path) for role, path in files.items()}
    choices = build_session_choices(
        mode=QCRunMode.CYCLE_COMPARISON,
        profile_name="default",
        files={role: str(path) for role, path in files.items()},
        file_hashes=file_hashes,
        output_mode="decision",
        allow_large_workbooks=False,
        allow_dependency_indexing=False,
        acceptance_absolute=0.0,
        acceptance_percent=0.0,
        rerun_of=None,
    )
    session_key = session_key_for(file_hashes)
    ConfigSessionStore(work_dir / "history.sqlite3").save_choices(
        session_key, profile_name="default", choices=choices
    )
    return session_key


@pytest.mark.asyncio
async def test_configure_page_reports_an_unknown_session(
    user: User, tmp_path: Path
) -> None:
    work_dir = tmp_path / "work"
    create_pages(work_dir)

    await user.open("/configure?session=does-not-exist")

    await user.should_see("This configuration session was not found")


@pytest.mark.asyncio
async def test_stale_file_role_stays_in_workspace_with_replacement_controls(
    user: User, tmp_path: Path
) -> None:
    work_dir = tmp_path / "work"
    session_key = _stage_session(work_dir)
    record = ConfigSessionStore(work_dir / "history.sqlite3").get(session_key)
    assert record is not None
    files = record.choices["files"]
    assert isinstance(files, dict)
    Path(str(files["current_excel"])).unlink()
    create_pages(work_dir)

    await user.open(f"/configure?session={session_key}")

    await user.should_see("Configure & run")
    await user.should_see("missing or changed")
    assert user.find(kind=ui.upload).elements
    assert user.find(kind=ui.button, content="Back to files").elements
    await user.should_not_see("Back to Compare to select files again")


@pytest.mark.asyncio
async def test_workspace_can_add_a_new_excel_member_role(
    user: User, tmp_path: Path
) -> None:
    work_dir = tmp_path / "work"
    session_key = _stage_session(work_dir)
    create_pages(work_dir)
    await user.open(f"/configure?session={session_key}")
    await user.should_see("New Excel member ID")
    member_input = next(
        element
        for element in user.find(kind=ui.input).elements
        if element.props.get("label") == "New Excel member ID"
    )
    member_input.value = "ops"

    user.find(kind=ui.button, content="Add workbook member").click()

    await user.should_see("Baseline workbook · ops")
    await user.should_see("Current workbook · ops")
    record = ConfigSessionStore(work_dir / "history.sqlite3").get(session_key)
    assert record is not None
    assert record.choices["pending_roles"] == [
        "baseline_excel:ops",
        "current_excel:ops",
    ]


@pytest.mark.asyncio
async def test_back_to_files_preserves_inputs_and_reuses_the_session(
    user: User, tmp_path: Path
) -> None:
    work_dir = tmp_path / "work"
    legacy_key = _stage_session(work_dir)
    store = ConfigSessionStore(work_dir / "history.sqlite3")
    original = store.get(legacy_key)
    assert original is not None
    create_pages(work_dir)

    await user.open(f"/configure?session={legacy_key}")
    await user.should_see("Back to files")
    user.find(kind=ui.button, content="Back to files").click()

    await user.should_see("Compare deliverables")
    await user.should_see("restored")
    user.find(marker="run-qc-button").click()
    await user.should_see("Configure & run")

    resumed = store.get(original.session_id)
    assert resumed is not None
    assert resumed.revision == original.revision + 1
    with store._connect() as connection:
        assert connection.execute("SELECT COUNT(*) FROM config_sessions").fetchone()[0] == 1


@pytest.mark.asyncio
async def test_completed_setup_rehydrates_reviews_after_coordinator_restart(
    user: User, tmp_path: Path
) -> None:
    work_dir = tmp_path / "work"
    session_key = _stage_session(work_dir)
    store = ConfigSessionStore(work_dir / "history.sqlite3")
    create_pages(work_dir)

    await user.open(f"/configure?session={session_key}")
    await user.should_see("Analysis complete", retries=_SUBPROCESS_RETRIES)
    original = store.get(session_key)
    assert original is not None
    original_job = get_setup_coordinator(work_dir).find(
        original.session_id, original.input_generation
    )
    assert original_job is not None
    assert original_job.result_payloads_snapshot()["primary"]["current_sheets"]

    # Simulate browser history restoration after the in-memory coordinator
    # was evicted or the server restarted. Lifecycle state and the full setup
    # payload remain in their separate SQLite stores.
    with setup_coordinator_module._COORDINATORS_LOCK:
        setup_coordinator_module._COORDINATORS.pop(work_dir.resolve(), None)

    await user.open(f"/configure?session={original.session_id}")

    await user.should_see("Analysis complete", retries=_SUBPROCESS_RETRIES)
    await user.should_see(marker="active-region-editor", retries=_SUBPROCESS_RETRIES)
    await user.should_not_see("missing required Excel member review: primary")
    restored_job = get_setup_coordinator(work_dir).find(
        original.session_id, original.input_generation
    )
    assert restored_job is not None
    assert restored_job.result_payloads_snapshot()["primary"]["current_sheets"]


@pytest.mark.asyncio
async def test_ppt_only_preflight_reaches_review_without_an_excel_job(
    user: User, tmp_path: Path
) -> None:
    work_dir = tmp_path / "work"
    session_key = _stage_ppt_only_preflight_session(work_dir)
    create_pages(work_dir)

    await user.open(f"/configure?session={session_key}")

    await user.should_see("PowerPoint slides", retries=_SUBPROCESS_RETRIES)
    await user.should_see("Summary", retries=_SUBPROCESS_RETRIES)
    run_once = user.find(kind=ui.button, content="Run once").elements.pop()
    assert isinstance(run_once, ui.button)
    assert run_once.enabled


@pytest.mark.asyncio
async def test_configure_page_lists_input_roles_and_completes_the_scan(
    user: User, tmp_path: Path
) -> None:
    work_dir = tmp_path / "work"
    session_key = _stage_session(work_dir)
    create_pages(work_dir)

    await user.open(f"/configure?session={session_key}")

    await user.should_see("Baseline workbook")
    await user.should_see("Current workbook")
    await user.should_see("Analysis complete", retries=_SUBPROCESS_RETRIES)

    record = ConfigSessionStore(work_dir / "history.sqlite3").get(session_key)
    assert record is not None
    job = get_setup_coordinator(work_dir).find(
        record.session_id, record.input_generation
    )
    assert job is not None
    assert job.snapshot().status is SetupStatus.COMPLETE
    payload = job.result_payloads_snapshot()["primary"]
    current_sheets = payload["current_sheets"]
    assert isinstance(current_sheets, list)
    data_sheet = next(sheet for sheet in current_sheets if sheet["sheet_name"] == "Data")
    assert len(data_sheet["regions"]) == 1


@pytest.mark.asyncio
async def test_configure_page_mounts_one_active_region_and_integrated_preview(
    user: User, tmp_path: Path
) -> None:
    work_dir = tmp_path / "work"
    session_key = _stage_multi_region_session(work_dir)
    create_pages(work_dir)

    await user.open(f"/configure?session={session_key}")
    await user.should_see("Analysis complete", retries=_SUBPROCESS_RETRIES)
    await user.should_see(marker="active-preview-grid", retries=_SUBPROCESS_RETRIES)

    assert len(user.find(marker="active-region-editor").elements) == 1
    assert len(user.find(marker="active-preview-panel").elements) == 1
    sheet_select = next(iter(user.find(marker="active-sheet-select").elements))
    assert isinstance(sheet_select, ui.select)
    assert len(sheet_select.options) == 2
    region_select = next(iter(user.find(marker="active-region-select").elements))
    assert isinstance(region_select, ui.select)
    assert isinstance(region_select.options, dict)
    assert len(region_select.options) == 2

    second_region = list(region_select.options)[1]
    region_select.value = second_region

    assert len(user.find(marker="active-region-editor").elements) == 1
    updated_region_select = next(
        iter(user.find(marker="active-region-select").elements)
    )
    assert isinstance(updated_region_select, ui.select)
    assert updated_region_select.value == second_region


@pytest.mark.asyncio
async def test_tabbed_region_editor_preserves_tab_and_toggles_exclusion(
    user: User, tmp_path: Path
) -> None:
    work_dir = tmp_path / "work"
    session_key = _stage_session(work_dir)
    create_pages(work_dir)

    await user.open(f"/configure?session={session_key}")
    await user.should_see(marker="active-region-editor", retries=_SUBPROCESS_RETRIES)
    await user.should_see("Rows")
    await user.should_see("Data bounds")
    await user.should_see("Scenario checks")
    await user.should_see("Advanced")

    tabs = next(
        element
        for element in user.find(kind=ui.tabs).elements
        if "config-region-tabs" in element.classes
    )
    tabs.value = "bounds"
    user.find(kind=ui.button, content="Remove from this run").click()

    updated_tabs = next(
        element
        for element in user.find(kind=ui.tabs).elements
        if "config-region-tabs" in element.classes
    )
    assert updated_tabs.value == "bounds"
    await user.should_see("Restore to this run")

    record = ConfigSessionStore(work_dir / "history.sqlite3").get(session_key)
    assert record is not None
    decisions = record.choices["region_decisions"]
    assert isinstance(decisions, dict)
    assert next(iter(decisions.values()))["mode"] == "excluded"

    user.find(kind=ui.button, content="Restore to this run").click()
    restored = ConfigSessionStore(work_dir / "history.sqlite3").get(session_key)
    assert restored is not None
    restored_decisions = restored.choices["region_decisions"]
    assert isinstance(restored_decisions, dict)
    assert next(iter(restored_decisions.values()))["mode"] == "automatic"


@pytest.mark.asyncio
async def test_data_start_commits_on_enter_and_preview_formula_text_reloads(
    user: User, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    work_dir = tmp_path / "work"
    session_key = _stage_session(work_dir)
    preview_requests = []

    def _preview(request):
        preview_requests.append(request)
        formula_text = {"1,2": "=A2*2"} if request.reveal_formulas else {}
        return PreviewWindowOutcome(
            rows=[["cached", "cached"]],
            formula_cells=[[False, True]],
            formula_text=formula_text,
            resolved_min_row=1,
            resolved_min_col=1,
            resolved_max_row=1,
            resolved_max_col=2,
        )

    monkeypatch.setattr(config_workspace_module, "run_preview_window_worker", _preview)
    create_pages(work_dir)

    await user.open(f"/configure?session={session_key}")
    await user.should_see(marker="active-preview-grid", retries=_SUBPROCESS_RETRIES)

    tabs = next(
        element
        for element in user.find(kind=ui.tabs).elements
        if "config-region-tabs" in element.classes
    )
    tabs.value = "bounds"
    data_start = next(
        element
        for element in user.find(kind=ui.input).elements
        if element.props.get("label") == "Data starts on row"
    )
    data_start.value = "2"
    _emit(data_start, "keydown.enter")

    record = ConfigSessionStore(work_dir / "history.sqlite3").get(session_key)
    assert record is not None
    decisions = record.choices["region_decisions"]
    assert isinstance(decisions, dict)
    decision = next(iter(decisions.values()))
    assert decision["header_intent"] == "first_data_row"
    assert decision["first_data_row"] == 2
    assert decision["preamble_rows"] == 0

    reveal = next(
        element
        for element in user.find(kind=ui.checkbox).elements
        if element.text == "Reveal formula text"
    )
    before = len(preview_requests)
    reveal.value = True
    await user.should_see("=A2*2", retries=_SUBPROCESS_RETRIES)
    assert len(preview_requests) > before
    assert preview_requests[-1].reveal_formulas is True


@pytest.mark.asyncio
async def test_preview_latest_request_wins_after_an_inflight_side_change(
    user: User, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    work_dir = tmp_path / "work"
    session_key = _stage_session(work_dir)
    started = threading.Event()
    release = threading.Event()
    requested_sides: list[str] = []

    def _preview(request):
        requested_sides.append(request.side)
        if len(requested_sides) == 1:
            started.set()
            assert release.wait(timeout=5)
        text = "CURRENT-STALE" if request.side == "current" else "BASELINE-LATEST"
        return PreviewWindowOutcome(
            rows=[[text]],
            formula_cells=[[False]],
            resolved_min_row=1,
            resolved_min_col=1,
            resolved_max_row=1,
            resolved_max_col=1,
        )

    monkeypatch.setattr(config_workspace_module, "run_preview_window_worker", _preview)
    create_pages(work_dir)

    await user.open(f"/configure?session={session_key}")
    assert await asyncio.to_thread(started.wait, 10)
    side_toggle = next(
        element
        for element in user.find(kind=ui.toggle).elements
        if isinstance(element.options, dict)
        and {"current", "baseline"}.issubset(element.options)
    )
    side_toggle.value = "baseline"
    release.set()

    await user.should_see("BASELINE-LATEST", retries=_SUBPROCESS_RETRIES)
    await user.should_not_see("CURRENT-STALE")
    assert requested_sides == ["current", "baseline"]


@pytest.mark.asyncio
async def test_profile_policy_editor_opens_in_place_only_once(
    user: User, tmp_path: Path
) -> None:
    work_dir = tmp_path / "work"
    session_key = _stage_session(work_dir)
    create_pages(work_dir)

    await user.open(f"/configure?session={session_key}")
    await user.should_see("Analysis complete", retries=_SUBPROCESS_RETRIES)
    edit_button = next(
        iter(user.find(kind=ui.button, content="Edit profile policy").elements)
    )
    _emit(edit_button, "click")
    _emit(edit_button, "click")

    await user.should_see("Apply saved profile to this setup")
    await user.should_see("New profile name")
    await user.should_see("Create")
    assert (
        len(
            user.find(
                kind=ui.button,
                content="Apply saved profile to this setup",
            ).elements
        )
        == 1
    )
    new_name = next(
        element
        for element in user.find(kind=ui.input).elements
        if element.props.get("label") == "New profile name"
    )
    new_name.value = "embedded-policy"
    user.find(kind=ui.button, content="Create").click()
    await user.should_see("Profile 'embedded-policy' created")
    user.find(
        kind=ui.button,
        content="Apply saved profile to this setup",
    ).click()

    record = ConfigSessionStore(work_dir / "history.sqlite3").get(session_key)
    assert record is not None
    assert record.profile_name == "embedded-policy"


@pytest.mark.asyncio
async def test_retry_analysis_binds_and_starts_the_replacement_job(
    user: User, tmp_path: Path
) -> None:
    work_dir = tmp_path / "work"
    legacy_key = _stage_session(work_dir)
    record = ConfigSessionStore(work_dir / "history.sqlite3").get(legacy_key)
    assert record is not None
    coordinator = get_setup_coordinator(work_dir)
    cancelled_job = coordinator.get_or_create(
        record.session_id, record.input_generation
    )
    coordinator.cancel(record.session_id, record.input_generation)
    create_pages(work_dir)

    await user.open(f"/configure?session={record.session_id}")
    await user.should_see("Setup analysis was cancelled")
    user.find(kind=ui.button, content="Resume analysis").click()

    await user.should_see("Analysis complete", retries=_SUBPROCESS_RETRIES)
    replacement_job = coordinator.find(record.session_id, record.input_generation)
    assert replacement_job is not None
    assert replacement_job is not cancelled_job
    assert replacement_job.snapshot().status is SetupStatus.COMPLETE


@pytest.mark.asyncio
async def test_confirm_all_regions_persists_confirmed_true_for_a_valid_region(
    user: User, tmp_path: Path
) -> None:
    work_dir = tmp_path / "work"
    session_key = _stage_session(work_dir)
    create_pages(work_dir)

    await user.open(f"/configure?session={session_key}")
    await user.should_see("Analysis complete", retries=_SUBPROCESS_RETRIES)
    await user.should_see("Confirm all detected regions", retries=_SUBPROCESS_RETRIES)

    user.find("Confirm all detected regions").click()

    store = ConfigSessionStore(work_dir / "history.sqlite3")
    record = store.get(session_key)
    assert record is not None
    region_decisions = record.choices.get("region_decisions")
    assert isinstance(region_decisions, dict)
    assert region_decisions
    assert all(decision["confirmed"] is True for decision in region_decisions.values())


@pytest.mark.asyncio
async def test_confirming_disjoint_identity_columns_surfaces_a_low_overlap_warning(
    user: User, tmp_path: Path
) -> None:
    """Step 12 Fix 5: setting a keyed region's identity column to one that
    shares NO values between baseline and current triggers a real
    (subprocess-backed) key-overlap query and surfaces a dedicated,
    separately-acknowledged BLOCKING warning: 'Run once' stays disabled
    until the analyst explicitly checks the acknowledgement, then becomes
    enabled again (mirrors ``allow_large_workbooks``'s "acknowledge and
    continue" precedent -- disclosed, not silently ignorable).
    """
    work_dir = tmp_path / "work"
    session_key = _stage_disjoint_key_session(work_dir)
    create_pages(work_dir)

    await user.open(f"/configure?session={session_key}")
    await user.should_see("Analysis complete", retries=_SUBPROCESS_RETRIES)

    mode_toggle = next(iter(user.find(marker="region-mode-toggle").elements))
    assert isinstance(mode_toggle, ui.toggle)
    mode_toggle.value = "keyed"

    identity_select = next(
        element
        for element in user.find(kind=ui.select).elements
        if element.props.get("label") == "Key columns"
    )
    identity_select.value = ["A"]

    # The header row's own "ID" label is one extra shared key on both
    # sides (this diagnostic query includes the whole resolved range --
    # see ``region_key_overlap_query_bounds``'s own disclosed, bounded
    # simplification), so a 3-row fixture with fully disjoint data keys
    # measures ~14% overlap, not exactly 0% -- still well below the 90%
    # threshold, so the warning still correctly fires.
    await user.should_see("only overlap 14%", retries=_SUBPROCESS_RETRIES)

    run_once_button = user.find(kind=ui.button, content="Run once").elements.pop()
    assert isinstance(run_once_button, ui.button)
    assert not run_once_button.enabled

    checkbox = next(
        element
        for element in user.find(kind=ui.checkbox).elements
        if "only overlap 14%" in (element.text or "")
    )
    checkbox.value = True

    run_once_button = user.find(kind=ui.button, content="Run once").elements.pop()
    assert isinstance(run_once_button, ui.button)
    assert run_once_button.enabled

    store = ConfigSessionStore(work_dir / "history.sqlite3")
    record = store.get(session_key)
    assert record is not None
    region_decisions = record.choices.get("region_decisions")
    assert isinstance(region_decisions, dict)
    [decision] = region_decisions.values()
    assert decision["mode"] == "keyed"
    assert decision["identity_columns"] == ["A"]


@pytest.mark.asyncio
async def test_changing_an_identity_columns_baseline_letter_recomputes_the_overlap(
    user: User, tmp_path: Path
) -> None:
    """Regression test for a real bug found via independent review (Step
    12, this session): the key-overlap ratio is computed from ``region.
    baseline_letter_for(letter)`` for each identity column, but the
    baseline-letter-override change handler never re-triggered the query
    -- entering an override after the initial (stale, low) computation
    left the low-overlap warning showing even once the override made the
    real overlap high. Proves the fix: after setting the baseline
    override, the warning disappears and 'Run once' becomes enabled
    without ever checking an acknowledgement box.
    """
    work_dir = tmp_path / "work"
    session_key = _stage_moved_identity_key_session(work_dir)
    create_pages(work_dir)

    await user.open(f"/configure?session={session_key}")
    await user.should_see("Analysis complete", retries=_SUBPROCESS_RETRIES)

    mode_toggle = next(iter(user.find(marker="region-mode-toggle").elements))
    assert isinstance(mode_toggle, ui.toggle)
    mode_toggle.value = "keyed"

    identity_select = next(
        element
        for element in user.find(kind=ui.select).elements
        if element.props.get("label") == "Key columns"
    )
    identity_select.value = ["A"]

    # Baseline column A vs current column A are disjoint -- the initial,
    # stale computation surfaces the same low-overlap warning as the
    # disjoint-key test above.
    await user.should_see("only overlap 14%", retries=_SUBPROCESS_RETRIES)
    run_once_button = user.find(kind=ui.button, content="Run once").elements.pop()
    assert isinstance(run_once_button, ui.button)
    assert not run_once_button.enabled

    user.find("Column letter differs from baseline?").click()
    await user.should_see("A in baseline", retries=_SUBPROCESS_RETRIES)
    baseline_letter_input = next(
        element
        for element in user.find(kind=ui.input).elements
        if element.props.get("label") == "A in baseline"
    )
    baseline_letter_input.value = "C"

    # The real overlap (baseline column C vs current column A) is a full
    # match -- the warning must disappear and the run must become
    # available WITHOUT ever checking an acknowledgement box, proving a
    # fresh query ran against the corrected baseline letter rather than
    # replaying the stale ratio from the pre-override computation.
    await user.should_not_see("only overlap", retries=_SUBPROCESS_RETRIES)
    run_once_button = user.find(kind=ui.button, content="Run once").elements.pop()
    assert isinstance(run_once_button, ui.button)
    assert run_once_button.enabled


@pytest.mark.asyncio
async def test_selector_prerequisite_is_added_and_survives_a_session_reload(
    user: User, tmp_path: Path
) -> None:
    """Step 8: an analyst-declared selector prerequisite must round-trip
    through the session store, not just live in the running page's
    in-memory state -- otherwise resuming a saved configuration session
    (this store's whole purpose) would silently drop it.
    """
    work_dir = tmp_path / "work"
    session_key = _stage_session(work_dir)
    create_pages(work_dir)

    await user.open(f"/configure?session={session_key}")
    await user.should_see("Analysis complete", retries=_SUBPROCESS_RETRIES)
    await user.should_see(
        "Add scenario or parameter check", retries=_SUBPROCESS_RETRIES
    )

    cell_input = next(
        element
        for element in user.find(kind=ui.input).elements
        if element.props.get("label") == "Cell (A1)"
    )
    label_input = next(
        element
        for element in user.find(kind=ui.input).elements
        if element.props.get("label") == "Label"
    )
    cell_input.value = "B5"
    label_input.value = "Scenario"
    user.find(kind=ui.button, content="Add check").click()

    await user.should_see("Scenario (B5)")

    store = ConfigSessionStore(work_dir / "history.sqlite3")
    record = store.get(session_key)
    assert record is not None
    selectors = record.choices.get("selectors")
    assert isinstance(selectors, dict)
    saved_entries = selectors.get("primary")
    assert isinstance(saved_entries, list) and len(saved_entries) == 1
    assert saved_entries[0]["sheet_name"] == "Data"
    assert saved_entries[0]["label"] == "Scenario"
    assert saved_entries[0]["cell"] == "B5"

    # Simulate resuming the session later (a fresh page load re-runs the
    # scan from scratch): the persisted selector must be restored, not
    # silently dropped.
    await user.open(f"/configure?session={session_key}")
    await user.should_see("Analysis complete", retries=_SUBPROCESS_RETRIES)
    await user.should_see("Scenario (B5)")


@pytest.mark.asyncio
async def test_selector_baseline_cell_override_survives_a_session_reload(
    user: User, tmp_path: Path
) -> None:
    """plan-20260913 Step 12 fix: a selector's baseline-side cell can
    differ from its current-side cell, and the override round-trips
    through the session store like every other region/selector choice.
    """
    work_dir = tmp_path / "work"
    session_key = _stage_session(work_dir)
    create_pages(work_dir)

    await user.open(f"/configure?session={session_key}")
    await user.should_see("Analysis complete", retries=_SUBPROCESS_RETRIES)
    await user.should_see(
        "Add scenario or parameter check", retries=_SUBPROCESS_RETRIES
    )

    cell_input = next(
        element
        for element in user.find(kind=ui.input).elements
        if element.props.get("label") == "Cell (A1)"
    )
    label_input = next(
        element
        for element in user.find(kind=ui.input).elements
        if element.props.get("label") == "Label"
    )
    cell_input.value = "B5"
    label_input.value = "Scenario"
    user.find(kind=ui.button, content="Add check").click()
    await user.should_see("Scenario (B5)")

    baseline_cell_input = next(
        element
        for element in user.find(kind=ui.input).elements
        if element.props.get("label") == "Cell in baseline (if moved)"
    )
    baseline_cell_input.value = "B9"

    store = ConfigSessionStore(work_dir / "history.sqlite3")
    record = store.get(session_key)
    assert record is not None
    selectors = record.choices.get("selectors")
    assert isinstance(selectors, dict)
    saved_entries = selectors.get("primary")
    assert isinstance(saved_entries, list) and len(saved_entries) == 1
    assert saved_entries[0]["baseline_cell"] == "B9"

    await user.open(f"/configure?session={session_key}")
    await user.should_see("Analysis complete", retries=_SUBPROCESS_RETRIES)
    await user.should_see("Scenario (B5)")
    restored_baseline_cell_input = next(
        element
        for element in user.find(kind=ui.input).elements
        if element.props.get("label") == "Cell in baseline (if moved)"
    )
    assert restored_baseline_cell_input.value == "B9"


@pytest.mark.asyncio
async def test_ppt_slide_review_shows_inventory_and_added_removed_warning(
    user: User, tmp_path: Path
) -> None:
    """Step 9: cycle mode's PowerPoint slide review shows the deck's slide
    titles and flags an unpaired added/removed slide as an unresolved
    (blocking) warning, exactly mirroring the Excel sheet-pairing
    criterion. Slides are peeked directly (no subprocess), so this
    exercises the section immediately after the Excel scan completes.
    """
    work_dir = tmp_path / "work"
    session_key = _stage_session_with_ppt(work_dir)
    create_pages(work_dir)

    await user.open(f"/configure?session={session_key}")
    await user.should_see("Analysis complete", retries=_SUBPROCESS_RETRIES)
    await user.should_see("PowerPoint slides")
    await user.should_see("Cover")
    await user.should_see("New Summary")
    # An unpaired added/removed slide is an unresolved blocker until the
    # analyst either declares a rename or acknowledges it.
    await user.should_see("Old Summary")


@pytest.mark.asyncio
async def test_ppt_slide_inclusion_toggle_persists_across_a_session_reload(
    user: User, tmp_path: Path
) -> None:
    work_dir = tmp_path / "work"
    session_key = _stage_session_with_ppt(work_dir)
    create_pages(work_dir)

    await user.open(f"/configure?session={session_key}")
    await user.should_see("Analysis complete", retries=_SUBPROCESS_RETRIES)
    await user.should_see("Cover")

    user.find(marker="slide-include-1").click()

    store = ConfigSessionStore(work_dir / "history.sqlite3")
    record = store.get(session_key)
    assert record is not None
    assert record.choices.get("excluded_slides") == [1]

    # Reopening the same session must restore the exclusion, not silently
    # re-include the slide.
    await user.open(f"/configure?session={session_key}")
    await user.should_see("Analysis complete", retries=_SUBPROCESS_RETRIES)
    await user.should_see("Cover", retries=_SUBPROCESS_RETRIES)
    excluded_checkbox = next(iter(user.find(marker="slide-include-1").elements))
    assert isinstance(excluded_checkbox, ui.checkbox)
    assert excluded_checkbox.value is False


@pytest.mark.asyncio
async def test_final_package_flags_a_stale_saved_slide_anchor_as_a_warning(
    user: User, tmp_path: Path
) -> None:
    """Step 9: final-package mode resolves a saved crosscheck mapping's
    slide title against the current deck. A title with zero matches is
    disclosed as a stale, warning-only mismatch -- never a blocker, and
    never guessed at (no other slide is substituted for it).
    """
    from qc_tool.config.profile import (
        CrosscheckMapping,
        CrosscheckProfile,
        DeliverableProfile,
        save_profile,
    )
    from qc_tool.config.profile import profile_path as _profile_path

    work_dir = tmp_path / "work"
    profiles_dir = work_dir / "profiles"
    stale_profile = DeliverableProfile(
        name="stale-mapping",
        crosscheck=CrosscheckProfile(
            mappings=[
                CrosscheckMapping(
                    slide="Old Summary",
                    line_skeleton="Revenue was #",
                    source_sheet="Data",
                    source_cell="B2",
                )
            ]
        ),
    )
    save_profile(stale_profile, _profile_path(profiles_dir, "stale-mapping"))

    files = {
        "current_excel": work_dir / "uploads" / "current_excel" / "current.xlsx",
        "current_ppt": work_dir / "uploads" / "current_ppt" / "current.pptx",
    }
    for path in files.values():
        path.parent.mkdir(parents=True, exist_ok=True)
    _write_simple_workbook(files["current_excel"])
    _write_simple_deck(files["current_ppt"], ["Cover", "New Summary"])
    file_hashes = {role: sha256_file(path) for role, path in files.items()}
    choices = build_session_choices(
        mode=QCRunMode.FINAL_PACKAGE,
        profile_name="stale-mapping",
        files={role: str(path) for role, path in files.items()},
        file_hashes=file_hashes,
        output_mode="decision",
        allow_large_workbooks=False,
        allow_dependency_indexing=False,
        acceptance_absolute=0.0,
        acceptance_percent=0.0,
        rerun_of=None,
    )
    session_key = session_key_for(file_hashes)
    ConfigSessionStore(work_dir / "history.sqlite3").save_choices(
        session_key, profile_name="stale-mapping", choices=choices
    )
    create_pages(work_dir)

    await user.open(f"/configure?session={session_key}")
    await user.should_see("Analysis complete", retries=_SUBPROCESS_RETRIES)
    await user.should_see("New Summary", retries=_SUBPROCESS_RETRIES)
    await user.should_see("Warnings")
    await user.should_see("Old Summary")
    await user.should_see("stale")


@pytest.mark.asyncio
async def test_save_profile_rejects_a_concurrent_edit_from_another_session(
    user: User, tmp_path: Path
) -> None:
    """plan-20260913 Step 11: "optimistic conflict protection" -- saving a
    profile this workspace session opened must refuse, not silently
    clobber, when the file on disk changed since it was opened (e.g. a
    second browser tab saved first).
    """
    from qc_tool.config.profile import DeliverableProfile, NumericTolerance, save_profile
    from qc_tool.config.profile import profile_path as _profile_path

    work_dir = tmp_path / "work"
    profiles_dir = work_dir / "profiles"
    profiles_dir.mkdir(parents=True)
    save_profile(
        DeliverableProfile(name="acme"), _profile_path(profiles_dir, "acme")
    )

    files = {
        "baseline_excel": work_dir / "uploads" / "baseline_excel" / "baseline.xlsx",
        "current_excel": work_dir / "uploads" / "current_excel" / "current.xlsx",
    }
    for path in files.values():
        path.parent.mkdir(parents=True, exist_ok=True)
        _write_simple_workbook(path)
    file_hashes = {role: sha256_file(path) for role, path in files.items()}
    choices = build_session_choices(
        mode=QCRunMode.CYCLE_COMPARISON,
        profile_name="acme",
        files={role: str(path) for role, path in files.items()},
        file_hashes=file_hashes,
        output_mode="decision",
        allow_large_workbooks=False,
        allow_dependency_indexing=False,
        acceptance_absolute=0.0,
        acceptance_percent=0.0,
        rerun_of=None,
    )
    session_key = session_key_for(file_hashes)
    ConfigSessionStore(work_dir / "history.sqlite3").save_choices(
        session_key, profile_name="acme", choices=choices
    )
    create_pages(work_dir)

    await user.open(f"/configure?session={session_key}")
    await user.should_see("Analysis complete", retries=_SUBPROCESS_RETRIES)
    await user.should_see("Confirm all detected regions", retries=_SUBPROCESS_RETRIES)
    user.find("Confirm all detected regions").click()

    # A "concurrent" edit: another session/tab saves the same profile name
    # with different content while this workspace stays open.
    save_profile(
        DeliverableProfile(
            name="acme", tolerance=NumericTolerance(absolute=99.0)
        ),
        _profile_path(profiles_dir, "acme"),
    )

    user.find(kind=ui.button, content="Save profile").click()

    await user.should_see("changed elsewhere while this workspace was open")
    # The concurrent edit's tolerance must survive untouched -- the click
    # above must not have overwritten it.
    from qc_tool.config.profile import load_profile_by_name

    reloaded = load_profile_by_name(profiles_dir, "acme")
    assert reloaded.tolerance.absolute == 99.0


@pytest.mark.asyncio
async def test_save_profile_succeeds_when_nothing_changed_underneath(
    user: User, tmp_path: Path
) -> None:
    """The inverse control: a save with no concurrent edit must still
    succeed normally -- conflict protection must not become a false
    positive on the ordinary, unchanged-file path.
    """
    from qc_tool.config.profile import DeliverableProfile, save_profile
    from qc_tool.config.profile import profile_path as _profile_path

    work_dir = tmp_path / "work"
    profiles_dir = work_dir / "profiles"
    profiles_dir.mkdir(parents=True)
    save_profile(
        DeliverableProfile(name="acme"), _profile_path(profiles_dir, "acme")
    )

    files = {
        "baseline_excel": work_dir / "uploads" / "baseline_excel" / "baseline.xlsx",
        "current_excel": work_dir / "uploads" / "current_excel" / "current.xlsx",
    }
    for path in files.values():
        path.parent.mkdir(parents=True, exist_ok=True)
        _write_simple_workbook(path)
    file_hashes = {role: sha256_file(path) for role, path in files.items()}
    choices = build_session_choices(
        mode=QCRunMode.CYCLE_COMPARISON,
        profile_name="acme",
        files={role: str(path) for role, path in files.items()},
        file_hashes=file_hashes,
        output_mode="decision",
        allow_large_workbooks=False,
        allow_dependency_indexing=False,
        acceptance_absolute=0.0,
        acceptance_percent=0.0,
        rerun_of=None,
    )
    session_key = session_key_for(file_hashes)
    ConfigSessionStore(work_dir / "history.sqlite3").save_choices(
        session_key, profile_name="acme", choices=choices
    )
    create_pages(work_dir)

    await user.open(f"/configure?session={session_key}")
    await user.should_see("Analysis complete", retries=_SUBPROCESS_RETRIES)
    await user.should_see("Confirm all detected regions", retries=_SUBPROCESS_RETRIES)
    user.find("Confirm all detected regions").click()

    user.find(kind=ui.button, content="Save profile").click()

    await user.should_see("Profile 'acme' saved")


@pytest.mark.asyncio
async def test_save_profile_from_default_requires_typing_a_new_name_first(
    user: User, tmp_path: Path
) -> None:
    """Regression test for a real bug found via live browser verification
    (plan-20260913 Step 12, this session): starting from the immutable
    'default' profile, 'Save profile'/'Save profile and run' had NO UI
    mechanism to type a new profile name at all -- both handlers hardcoded
    ``as_new_name=None``, so ``_finalize()``'s ``target_name`` always
    resolved to an empty string and every save attempt failed with "Choose
    a profile name before saving". Every PRE-EXISTING save-profile test
    (see the test immediately above) pre-seeded the session with an
    already-named non-default profile via a direct ``ConfigSessionStore``
    call, so this exact bug was never exercised by any automated test.
    Proves the fix: a "Save as profile name" input lets an analyst type a
    new name and successfully save from the real starting state.
    """
    work_dir = tmp_path / "work"
    session_key = _stage_session(work_dir)  # profile_name="default"
    create_pages(work_dir)

    await user.open(f"/configure?session={session_key}")
    await user.should_see("Analysis complete", retries=_SUBPROCESS_RETRIES)
    await user.should_see("Confirm all detected regions", retries=_SUBPROCESS_RETRIES)
    user.find("Confirm all detected regions").click()

    # Without typing a name, saving from "default" must fail exactly as it
    # did before this fix -- proves the bug is real, not already unreachable.
    user.find(kind=ui.button, content="Save profile").click()
    await user.should_see("Choose a profile name before saving")

    name_input = next(
        element
        for element in user.find(kind=ui.input).elements
        if element.props.get("label") == "Save as profile name"
    )
    name_input.value = "verify-save-as"

    user.find(kind=ui.button, content="Save profile").click()
    await user.should_see("Profile 'verify-save-as' saved")

    assert (work_dir / "profiles" / "verify-save-as.yaml").exists()


class _RecordingQueueManager:
    """UI-test queue boundary that records requests without starting a
    worker -- mirrors ``tests/test_ui.py``'s own stub of the same name.
    """

    def __init__(self, store: RunStateStore) -> None:
        self.store = store
        self.shutdown_hook_installed = False
        self.submitted: list[tuple[object, dict[str, str]]] = []

    def submit(self, request: object, credentials: dict[str, str] | None = None) -> None:
        self.submitted.append((request, dict(credentials or {})))

    def shutdown(self) -> None:
        pass


@pytest.mark.asyncio
async def test_encrypted_setup_prompts_then_submits_with_page_memory_credential(
    user: User, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    work_dir = tmp_path / "work"
    session_key = _stage_encrypted_session(work_dir)
    manager = _RecordingQueueManager(RunStateStore(work_dir / "history.sqlite3"))
    monkeypatch.setattr(app_module, "get_manager", lambda _work_dir: manager)
    create_pages(work_dir)

    await user.open(f"/configure?session={session_key}")
    await user.should_see(
        "A password is required before setup analysis can continue.",
        retries=_SUBPROCESS_RETRIES,
    )
    password_input = next(
        element
        for element in user.find(kind=ui.input).elements
        if element.props.get("label") == "Password for current excel"
    )
    password_input.value = "hunter2"
    user.find(kind=ui.button, content="Resume analysis").click()

    await user.should_see("Analysis complete", retries=_SUBPROCESS_RETRIES)
    await user.open(f"/configure?session={session_key}")
    await user.should_see("Re-enter the required password before running.")
    password_input = next(
        element
        for element in user.find(kind=ui.input).elements
        if element.props.get("label") == "Password for current excel"
    )
    password_input.value = "hunter2"
    user.find(kind=ui.button, content="Use passwords").click()
    user.find("Confirm all detected regions").click()
    user.find(kind=ui.button, content="Run once").click()
    await user.should_see("Run started")

    assert len(manager.submitted) == 1
    _request, credentials = manager.submitted[0]
    assert credentials == {"current_excel": "hunter2"}
    record = ConfigSessionStore(work_dir / "history.sqlite3").get(session_key)
    assert record is not None
    assert "hunter2" not in str(record.choices)


@pytest.mark.asyncio
async def test_run_once_submits_the_run_and_navigates_home(
    user: User, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """The ordinary case: a small comparison with no projected-volume
    warning submits immediately through the workspace's own 'Run once'
    action (plan-20260913 Step 12's wizard-first fix -- this exact
    submission path had no end-to-end test before this).
    """
    work_dir = tmp_path / "work"
    session_key = _stage_session(work_dir)
    manager = _RecordingQueueManager(RunStateStore(work_dir / "history.sqlite3"))
    monkeypatch.setattr(app_module, "get_manager", lambda _work_dir: manager)
    create_pages(work_dir)

    await user.open(f"/configure?session={session_key}")
    await user.should_see("Analysis complete", retries=_SUBPROCESS_RETRIES)
    await user.should_see("Confirm all detected regions", retries=_SUBPROCESS_RETRIES)
    user.find("Confirm all detected regions").click()

    user.find(kind=ui.button, content="Run once").click()

    await user.should_see("Run started")
    assert len(manager.submitted) == 1
    request, credentials = manager.submitted[0]
    assert credentials == {}
    assert request.resolved_input_configuration is not None  # type: ignore[attr-defined]


def test_configuration_export_is_built_as_in_memory_bytes() -> None:
    from qc_tool.config.profile import DeliverableProfile
    from qc_tool.config.resolved_input import ResolvedInputConfigurationV1

    profile = DeliverableProfile(name="export-test")
    resolved = ResolvedInputConfigurationV1()

    exported = configuration_export_bytes(profile, resolved)

    assert isinstance(exported, bytes)
    payload = json.loads(exported)
    assert payload["profile"]["name"] == "export-test"
    assert payload["resolved_input_digest"] == resolved.canonical_sha256()


@pytest.mark.asyncio
async def test_run_once_warns_before_a_very_large_comparison(
    user: User, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Porting the main page's own pre-wizard volume-projection courtesy
    warning into the workspace (Step 12 fix) -- without this, the warning
    would have silently disappeared the moment 'Run QC' always routes
    through here first.
    """
    work_dir = tmp_path / "work"
    session_key = _stage_session(work_dir)
    manager = _RecordingQueueManager(RunStateStore(work_dir / "history.sqlite3"))
    monkeypatch.setattr(app_module, "get_manager", lambda _work_dir: manager)
    huge_projection = VolumeProjection(
        changed_sheets=("Data",),
        added_sheets=(),
        removed_sheets=(),
        identical_sheets=(),
        projected_max_findings=999_999,
    )
    monkeypatch.setattr(
        config_workspace_module, "project_cycle_volume", lambda *_args: huge_projection
    )
    create_pages(work_dir)

    await user.open(f"/configure?session={session_key}")
    await user.should_see("Analysis complete", retries=_SUBPROCESS_RETRIES)
    await user.should_see("Confirm all detected regions", retries=_SUBPROCESS_RETRIES)
    user.find("Confirm all detected regions").click()

    user.find(kind=ui.button, content="Run once").click()

    await user.should_see("This looks like a very large comparison")
    assert manager.submitted == []

    user.find(kind=ui.button, content="Run anyway").click()

    await user.should_see("Run started")
    assert len(manager.submitted) == 1
