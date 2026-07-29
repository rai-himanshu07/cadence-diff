"""Excel data-validation and conditional-format integrity contracts."""

import json
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, FormulaRule
from openpyxl.styles import Alignment, Color, Font, PatternFill
from openpyxl.styles.differential import DifferentialStyle
from openpyxl.worksheet.datavalidation import DataValidation

from qc_tool.coverage import CoverageState
from qc_tool.engine import QCRunResult, run_qc
from qc_tool.excel.interaction import (
    conditional_style_coverage,
    diff_interaction_rules,
    interaction_rule_coverage,
)
from qc_tool.findings import Finding, FindingClass
from qc_tool.fingerprint import fingerprint_file
from qc_tool.history.store import RunHistory
from qc_tool.io.loader import load_workbook_snapshot
from qc_tool.io.model import (
    ConditionalFormatDescriptor,
    DataValidationDescriptor,
    SheetSnapshot,
    WorkbookSnapshot,
)
from qc_tool.io.ooxml_interaction import _differential_style
from qc_tool.report.excel_report import write_excel_report
from qc_tool.report.html_report import render_html_report
from qc_tool.triage.rules import triage


def _build_interaction_workbook(
    path: Path,
    *,
    theme_style: bool = False,
    complex_rule: bool = False,
) -> None:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Data"
    validation = DataValidation(
        type="list",
        formula1='"Yes,No"',
        allow_blank=True,
        showErrorMessage=True,
        showInputMessage=True,
        promptTitle="Choose",
        prompt="Pick one",
    )
    validation.add("A2:A5")
    validation.add("C2:C5")
    sheet.add_data_validation(validation)
    sheet.conditional_formatting.add(
        "B2:B5",
        CellIsRule(
            operator="greaterThan",
            formula=["10"],
            fill=PatternFill("solid", fgColor="FFFF0000"),
        ),
    )
    if theme_style:
        sheet.conditional_formatting.add(
            "D2:D5",
            FormulaRule(
                formula=["D2>0"],
                font=Font(color=Color(theme=1)),
                stopIfTrue=True,
            ),
        )
    if complex_rule:
        sheet.conditional_formatting.add(
            "E2:E5",
            ColorScaleRule(
                start_type="min",
                start_color="FFFF0000",
                end_type="max",
                end_color="FF00FF00",
            ),
        )
    workbook.save(path)


def test_ooxml_snapshot_captures_validation_and_conditional_rules(
    tmp_path: Path,
) -> None:
    path = tmp_path / "interactions.xlsx"
    _build_interaction_workbook(path)

    snapshot = load_workbook_snapshot(path)

    assert snapshot.interaction_rules_available
    assert snapshot.interaction_rules_supported
    assert snapshot.conditional_format_styles_supported
    assert len(snapshot.data_validations) == 1
    validation = snapshot.data_validations[0]
    assert validation.target_ranges == ("A2:A5", "C2:C5")
    assert validation.validation_type == "list"
    assert validation.formula1 == '"Yes,No"'
    assert validation.allow_blank
    assert validation.prompt_title == "Choose"
    assert validation.prompt == "Pick one"

    assert len(snapshot.conditional_formats) == 1
    rule = snapshot.conditional_formats[0]
    assert rule.target_ranges == ("B2:B5",)
    assert rule.rule_type == "cellIs"
    assert rule.operator == "greaterThan"
    assert rule.formulas == ("10",)
    assert rule.priority == 1
    assert rule.style_supported
    assert rule.style_key is not None and "FF0000" in rule.style_key


def test_unresolved_theme_style_degrades_only_style_coverage(tmp_path: Path) -> None:
    path = tmp_path / "theme.xlsx"
    _build_interaction_workbook(path, theme_style=True)

    snapshot = load_workbook_snapshot(path)
    rule_state, _ = interaction_rule_coverage(snapshot)
    style_state, style_detail = conditional_style_coverage(snapshot)

    assert snapshot.interaction_rules_supported
    assert not snapshot.conditional_format_styles_supported
    assert rule_state is CoverageState.CHECKED
    assert style_state is CoverageState.DEGRADED
    assert "theme" in style_detail.lower()
    theme_rule = next(rule for rule in snapshot.conditional_formats if not rule.style_supported)
    assert theme_rule.style_key is None


def test_complex_conditional_rule_degrades_semantic_coverage(tmp_path: Path) -> None:
    path = tmp_path / "complex.xlsx"
    _build_interaction_workbook(path, complex_rule=True)

    snapshot = load_workbook_snapshot(path)
    state, detail = interaction_rule_coverage(snapshot)

    assert not snapshot.interaction_rules_supported
    assert state is CoverageState.DEGRADED
    assert "colorScale" in detail


def _workbook(
    *,
    validations: list[DataValidationDescriptor] | None = None,
    formats: list[ConditionalFormatDescriptor] | None = None,
) -> WorkbookSnapshot:
    return WorkbookSnapshot(
        "rules.xlsx",
        "xlsx",
        True,
        True,
        sheets=[SheetSnapshot("Data", "visible", 20, 20)],
        data_validations=validations or [],
        conditional_formats=formats or [],
    )


def _validation(
    targets: tuple[str, ...],
    *,
    source_index: int = 0,
    formula1: str = '"Yes,No"',
    formula2: str | None = None,
    validation_type: str = "list",
    prompt: str | None = "Pick one",
) -> DataValidationDescriptor:
    return DataValidationDescriptor(
        sheet="Data",
        source_index=source_index,
        source_id=f"Data!validation[{source_index}]",
        target_ranges=targets,
        validation_type=validation_type,
        formula1=formula1,
        formula2=formula2,
        allow_blank=True,
        prompt=prompt,
    )


def _format(
    targets: tuple[str, ...],
    *,
    source_index: int = 0,
    formula: str = "B2>10",
    operator: str | None = None,
    priority: int = 1,
    stop_if_true: bool | None = False,
    style_key: str | None = "rgb-red",
) -> ConditionalFormatDescriptor:
    return ConditionalFormatDescriptor(
        sheet="Data",
        source_index=source_index,
        source_id=f"Data!conditional-format[{source_index}]",
        target_ranges=targets,
        rule_type="expression",
        operator=operator,
        formulas=(formula,),
        priority=priority,
        stop_if_true=stop_if_true,
        style_key=style_key,
    )


def test_data_validation_diff_detects_targets_criteria_display_and_inventory() -> None:
    baseline = _workbook(
        validations=[
            _validation(("A2:A5",), prompt="Pick one"),
            _validation(
                ("D2:D5",),
                source_index=1,
                validation_type="whole",
                formula1="1",
                formula2="100",
            ),
        ]
    )
    current = _workbook(
        validations=[
            _validation(("A2:A6",), prompt="Choose one"),
            _validation(
                ("D2:D5",),
                source_index=1,
                validation_type="whole",
                formula1="1",
                formula2="90",
            ),
            _validation(("F2:F5",), source_index=2, formula1='"A,B"'),
        ]
    )

    findings = diff_interaction_rules(baseline, current)
    validation_findings = [
        finding
        for finding in findings
        if finding.finding_class is FindingClass.DATA_VALIDATION_CHANGED
    ]

    assert any("target" in finding.message for finding in validation_findings)
    assert any("criteria" in finding.message for finding in validation_findings)
    assert any("display text" in finding.message for finding in validation_findings)
    assert any("added" in finding.message for finding in validation_findings)


def test_conditional_format_diff_detects_targets_rule_priority_stop_style_inventory() -> None:
    baseline = _workbook(
        formats=[
            _format(("B2:B5",)),
            _format(("D2:D5",), source_index=1, formula="D2<0"),
        ]
    )
    current = _workbook(
        formats=[
            _format(
                ("B2:B6",),
                formula="B2>20",
                priority=2,
                stop_if_true=True,
                style_key="rgb-green",
            ),
            _format(("F2:F5",), source_index=2, formula="F2=0"),
        ]
    )

    findings = diff_interaction_rules(baseline, current)
    conditional_findings = [
        finding
        for finding in findings
        if finding.finding_class is FindingClass.CONDITIONAL_FORMAT_CHANGED
    ]

    assert any("target" in finding.message for finding in conditional_findings)
    assert any("condition" in finding.message for finding in conditional_findings)
    assert any("priority" in finding.message for finding in conditional_findings)
    assert any("stop-if-true" in finding.message for finding in conditional_findings)
    assert any("style" in finding.message for finding in conditional_findings)
    assert any("removed" in finding.message for finding in conditional_findings)
    assert any("added" in finding.message for finding in conditional_findings)


def test_interaction_findings_roundtrip_history_and_keep_formulas_inert(
    tmp_path: Path,
) -> None:
    result = QCRunResult(profile_name="interaction")
    result.findings = triage(
        [
            Finding(
                artifact="excel",
                finding_class=FindingClass.CONDITIONAL_FORMAT_CHANGED,
                sheet="Data",
                baseline_value="=A1>0",
                current_value="=A1>10",
                message="conditional-format condition changed",
            )
        ]
    )
    history = RunHistory(tmp_path / "history.sqlite3")
    run_id = history.record_run(result, file_hashes={}, report_paths={})
    record = history.get_run(run_id)
    assert record.findings == result.findings

    report_path = tmp_path / "interaction-report.xlsx"
    write_excel_report(result, report_path)
    findings_sheet = load_workbook(report_path, data_only=False)["Findings"]
    assert findings_sheet["H2"].value == "=A1>0"
    assert findings_sheet["H2"].data_type == "s"
    assert findings_sheet["I2"].value == "=A1>10"
    assert findings_sheet["I2"].data_type == "s"
    assert "conditional_format_changed" in render_html_report(result)


def test_cycle_coverage_separates_rule_semantics_from_unresolved_styles(
    tmp_path: Path,
) -> None:
    baseline = tmp_path / "baseline.xlsx"
    current = tmp_path / "current.xlsx"
    _build_interaction_workbook(baseline, theme_style=True)
    _build_interaction_workbook(current, theme_style=True)

    result = run_qc(baseline_excel=baseline, current_excel=current)
    rule_coverage = next(
        item for item in result.coverage if item.check_id == "excel-interaction-rules"
    )
    style_coverage = next(
        item
        for item in result.coverage
        if item.check_id == "excel-conditional-format-styles"
    )

    assert rule_coverage.state is CoverageState.CHECKED
    assert style_coverage.state is CoverageState.DEGRADED
    assert "theme" in style_coverage.detail.lower()


def test_interaction_fingerprint_is_structural_only(tmp_path: Path) -> None:
    path = tmp_path / "interactions.xlsx"
    _build_interaction_workbook(path)

    payload = fingerprint_file(path)
    serialized = json.dumps(payload, sort_keys=True)

    assert payload["data_validation_count"] == 1
    assert payload["conditional_format_count"] == 1
    assert payload["capabilities"]["interaction_semantics"] is True
    assert payload["capabilities"]["conditional_format_styles"] is True
    for forbidden in ("Pick one", "Choose", '"Yes,No"', "A2:A5", "B2:B5"):
        assert forbidden not in serialized


def test_inserted_rule_does_not_steal_simultaneously_changed_rule() -> None:
    baseline = _workbook(
        validations=[
            _validation(
                ("A2:A5",),
                validation_type="whole",
                formula1="1",
                formula2="100",
            )
        ]
    )
    current = _workbook(
        validations=[
            _validation(("Z2:Z5",), formula1='"New,Rule"'),
            _validation(
                ("A2:A6",),
                source_index=1,
                validation_type="whole",
                formula1="1",
                formula2="90",
            ),
        ]
    )

    findings = diff_interaction_rules(baseline, current)

    assert any("target coverage changed" in finding.message for finding in findings)
    assert any("criteria changed" in finding.message for finding in findings)
    added = [finding for finding in findings if "added" in finding.message]
    assert len(added) == 1
    assert added[0].current_value == "Z2:Z5"
    assert not any("removed" in finding.message for finding in findings)


def test_asymmetric_unavailable_pair_reports_unavailable_coverage(
    fixture_dir: Path,
) -> None:
    result = run_qc(
        baseline_excel=fixture_dir / "baseline.xlsb",
        current_excel=fixture_dir / "current.xlsx",
    )
    rule_coverage = next(
        item for item in result.coverage if item.check_id == "excel-interaction-rules"
    )
    style_coverage = next(
        item
        for item in result.coverage
        if item.check_id == "excel-conditional-format-styles"
    )

    assert rule_coverage.state is CoverageState.UNAVAILABLE
    assert style_coverage.state is CoverageState.UNAVAILABLE
    assert not any(
        finding.finding_class
        in {
            FindingClass.DATA_VALIDATION_CHANGED,
            FindingClass.CONDITIONAL_FORMAT_CHANGED,
        }
        for finding in result.findings
    )


def test_differential_style_normalizes_rgb_alpha_and_empty_components() -> None:
    transparent_key, transparent_supported, _ = _differential_style(
        DifferentialStyle(font=Font(color=Color(rgb="00000000")))
    )
    opaque_key, opaque_supported, _ = _differential_style(
        DifferentialStyle(font=Font(color=Color(rgb="FF000000")))
    )
    empty_key, empty_supported, _ = _differential_style(
        DifferentialStyle(
            font=Font(),
            fill=PatternFill(),
            alignment=Alignment(),
        )
    )

    assert transparent_supported and opaque_supported and empty_supported
    assert transparent_key == opaque_key
    assert empty_key is None


def test_duplicate_conditions_match_by_distinct_targets() -> None:
    baseline = _workbook(
        validations=[
            _validation(("A2:A5",), source_index=0),
            _validation(("C2:C5",), source_index=1),
        ]
    )
    current = _workbook(
        validations=[
            _validation(("C2:C6",), source_index=0),
            _validation(("A2:A6",), source_index=1),
        ]
    )

    findings = diff_interaction_rules(baseline, current)

    assert len(findings) == 2
    assert all("target coverage changed" in finding.message for finding in findings)
    assert not any(
        "added" in finding.message or "removed" in finding.message
        for finding in findings
    )


def test_unsupported_matched_rule_only_reports_target_changes() -> None:
    baseline_rule = _format(("B2:B5",), formula="")
    baseline_rule.rule_type = "colorScale"
    baseline_rule.semantic_supported = False
    current_rule = _format(
        ("B2:B6",),
        formula="different",
        priority=2,
        stop_if_true=True,
    )
    current_rule.rule_type = "colorScale"
    current_rule.semantic_supported = False

    findings = diff_interaction_rules(
        _workbook(formats=[baseline_rule]),
        _workbook(formats=[current_rule]),
    )

    assert len(findings) == 1
    assert {finding.finding_class for finding in findings} == {
        FindingClass.CONDITIONAL_FORMAT_CHANGED
    }
    assert any("target coverage changed" in finding.message for finding in findings)
    assert not any(
        phrase in finding.message
        for finding in findings
        for phrase in (
            "condition changed",
            "priority changed",
            "stop-if-true changed",
            "differential style changed",
        )
    )
