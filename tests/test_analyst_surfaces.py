"""Step 9: analyst-facing prioritization across every surface."""

from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook, load_workbook

from qc_tool.coverage import CoverageItem, CoverageState, capability_limited
from qc_tool.engine import QCRunResult, run_qc
from qc_tool.findings import (
    Finding,
    FindingClass,
    FindingEvidenceTag,
    FindingExpectedReason,
    FindingProvenance,
    FindingSubtype,
    FindingTemporalContext,
    Materiality,
    Severity,
)
from qc_tool.history.store import RunHistory
from qc_tool.report.excel_report import write_excel_report
from qc_tool.report.html_report import _member_payload, render_html_report
from qc_tool.report.json_report import result_payload
from qc_tool.review import apply_group_review, build_pattern_groups, count_pattern_groups
from qc_tool.triage.rules import triage


def _cycle_workbook(path: Path, *, offset: int) -> None:
    workbook = Workbook()
    sheet = workbook.active
    if sheet is None:  # pragma: no cover - openpyxl always creates one sheet
        raise RuntimeError("openpyxl did not create a default worksheet")
    sheet.title = "Data"
    sheet.append(["Key", "Amount", "Derived"])
    for index in range(1, 9):
        sheet.append([f"k{index}", index + offset, f"=B{index + 1}*2"])
    workbook.save(path)


def _result() -> QCRunResult:
    findings = triage(
        [
            Finding(
                artifact="excel",
                finding_class=FindingClass.FORMULA_ERROR,
                provenance=FindingProvenance.INHERITED,
                sheet="Data",
                location="C2",
                current_value="#REF!",
                message="Data!C2: error value #REF! (inherited from the baseline)",
            ),
            Finding(
                artifact="excel",
                finding_class=FindingClass.VALUE_CHANGED,
                subtype=FindingSubtype.VALUE_ADDED_POPULATION,
                sheet="Data",
                location="B4",
                baseline_value="",
                current_value="7",
                message="Data!B4: value added to a previously blank cell",
            ),
        ]
    )
    result = QCRunResult(profile_name="surfaces", findings=findings)
    result.coverage = [
        CoverageItem(
            check_id="excel-values",
            label="Excel values",
            artifact="excel",
            state=CoverageState.CHECKED,
        ),
        CoverageItem(
            check_id="excel-charts",
            label="Excel charts",
            artifact="excel",
            state=CoverageState.UNAVAILABLE,
            detail="chart metadata unavailable",
        ),
    ]
    return result


def test_step9_checkpoint_reconciles_with_the_delta_ledger() -> None:
    oracles = Path(__file__).parent / "oracles"
    step8 = json.loads((oracles / "real_workload_step8.json").read_text(encoding="utf-8"))
    step9 = json.loads((oracles / "real_workload_step9.json").read_text(encoding="utf-8"))
    ledger = json.loads(
        (oracles / "step_delta_ledger.json").read_text(encoding="utf-8")
    )["steps"]["9"]

    assert step9["atomic_findings"] - step8["atomic_findings"] == (
        ledger["expected_atomic_delta"]
    )
    assert sum(step9["pattern_review_counts"].values()) - sum(
        step8["pattern_review_counts"].values()
    ) == ledger["expected_pattern_review_delta"]
    assert sum(step9["review_counts"].values()) - sum(
        step8["review_counts"].values()
    ) == ledger["expected_review_count_delta"]
    assert step9["severity"] == step8["severity"]
    assert step9["coverage_states"] == step8["coverage_states"]
    assert step9["mixed_pattern_groups"] == 0
    assert step9["source_hashes_unchanged"] is True


def test_step8_readme_documents_primary_review_and_runtime_controls() -> None:
    readme = (Path(__file__).parents[1] / "README.md").read_text(encoding="utf-8")
    normalized = " ".join(readme.replace("**", "").split())

    assert "semantic pattern groups" in normalized
    assert "Spatial groups remain a separate" in normalized
    assert "--accept-absolute 1 --accept-percent 0.1" in readme
    assert '--sheets "Dashboard,Data" --slides 1,3-5' in readme
    assert "Implicit numeric block refreshes are Warning" in readme


def test_capability_limited_is_reported_whenever_a_check_is_unavailable() -> None:
    result = _result()

    assert capability_limited(result.coverage)
    assert "capability-limited" in render_html_report(result)
    payload = result_payload(result, include_review_summary=True)
    assert payload["review_summary"]["capability_limited"] is True


def test_machine_json_stays_atomic_and_the_review_summary_is_opt_in() -> None:
    result = _result()

    default_payload = result_payload(result)
    summary_payload = result_payload(result, include_review_summary=True)

    assert "review_summary" not in default_payload
    assert default_payload["schema_version"] == 1
    assert len(default_payload["findings"]) == len(result.findings)
    assert summary_payload["findings"] == default_payload["findings"]
    summary = summary_payload["review_summary"]
    assert summary["summary_version"] == 3
    assert sum(summary["pattern_review_counts"].values()) == len(summary["groups"])
    assert sum(summary["atomic_findings_by_severity"].values()) == len(result.findings)
    assert {
        finding_id for group in summary["groups"] for finding_id in group["finding_ids"]
    } == {finding.finding_id for finding in result.findings}
    assert any(group["provenance"] == "inherited" for group in summary["groups"])
    assert any(group["subtype"] == "added_population" for group in summary["groups"])
    assert all("expected_reason" in group for group in summary["groups"])
    assert all("temporal_context" in group for group in summary["groups"])
    assert all("evidence_tags" in group for group in summary["groups"])
    stories = summary["stories"]
    assert stories, "v3 summary must carry change stories"
    assert {
        finding_id for story in stories for finding_id in story["finding_ids"]
    } == {finding.finding_id for finding in result.findings}


def test_every_surface_reproduces_the_same_pattern_counts(tmp_path: Path) -> None:
    result = _result()
    expected = count_pattern_groups(build_pattern_groups(result.findings))
    report = tmp_path / "report.xlsx"

    write_excel_report(result, report)
    summary = load_workbook(report)["Summary"]
    summary_text = "\n".join(
        str(cell.value)
        for row in summary.iter_rows()
        for cell in row
        if cell.value is not None
    )
    html = render_html_report(result)
    payload = result_payload(result, include_review_summary=True)

    for severity, count in expected.review_items.items():
        if not count:
            continue
        assert f"{severity.value.title()} pattern review items\n{count}" in summary_text
        assert payload["review_summary"]["pattern_review_counts"][severity.value] == count
    assert "pattern review items" in html
    assert "capability-limited" in summary_text


def test_provenance_and_subtype_reach_the_excel_evidence_columns(tmp_path: Path) -> None:
    report = tmp_path / "report.xlsx"

    write_excel_report(_result(), report)
    findings = load_workbook(report)["Findings"]
    headers = [cell.value for cell in next(findings.iter_rows(max_row=1))]
    rows = {
        row[0]: dict(zip(headers, row, strict=True))
        for row in findings.iter_rows(min_row=2, values_only=True)
    }

    assert "Provenance" in headers
    assert "Subtype" in headers
    assert "Materiality" in headers
    assert "Temporal context" in headers
    assert "Expected reason" in headers
    assert "Evidence tags" in headers
    assert {str(row["Provenance"] or "") for row in rows.values()} == {"inherited", ""}
    assert {str(row["Subtype"] or "") for row in rows.values()} == {
        "added_population",
        "",
    }


def test_all_evidence_axes_reach_html_atomic_members() -> None:
    finding = Finding(
        finding_id="F0001",
        artifact="excel",
        finding_class=FindingClass.VALUE_CHANGED,
        severity=Severity.EXPECTED,
        expected_reason=FindingExpectedReason.PROFILE_REFRESH,
        provenance=FindingProvenance.INHERITED,
        subtype=FindingSubtype.VALUE_REPLACEMENT,
        materiality=Materiality.MATERIAL,
        temporal_context=FindingTemporalContext.RECENT_WINDOW,
        evidence_tags={
            FindingEvidenceTag.FORMULA_PRESENCE,
            FindingEvidenceTag.CONCENTRATED_POPULATION,
        },
        message="evidence",
    )
    result = QCRunResult(profile_name="evidence", findings=[finding])

    member = next(iter(_member_payload(result).values()))[0]

    assert member["provenance"] == "inherited"
    assert member["subtype"] == "replacement"
    assert member["materiality"] == "material"
    assert member["temporal_context"] == "recent_window"
    assert member["expected_reason"] == "profile_refresh"
    assert member["evidence_tags"] == "concentrated_population; formula_presence"


def test_bulk_disposition_never_crosses_a_semantic_boundary(tmp_path: Path) -> None:
    baseline = tmp_path / "baseline.xlsx"
    current = tmp_path / "current.xlsx"
    _cycle_workbook(baseline, offset=0)
    _cycle_workbook(current, offset=1)
    result = run_qc(baseline_excel=baseline, current_excel=current)
    history = RunHistory(tmp_path / "history.sqlite3")
    run_id = history.record_run(result, file_hashes={}, report_paths={})
    groups = build_pattern_groups(result.findings)
    target = max(groups, key=lambda group: group.member_count)

    annotations = apply_group_review(
        target, severity=Severity.INFO, comment="bulk", replace_existing=True
    )
    for annotation in annotations:
        history.set_annotation(
            run_id,
            annotation.finding_id,
            severity=annotation.severity,
            comment=annotation.comment,
        )
    stored = history.get_run(run_id)
    touched = {
        finding.finding_id
        for finding in stored.findings
        if finding.analyst_comment == "bulk"
    }

    assert touched == {member.finding_id for member in target.members}
    assert all(
        finding.analyst_comment == ""
        for finding in stored.findings
        if finding.finding_id not in touched
    )


def test_legacy_history_regeneration_still_renders_every_surface(tmp_path: Path) -> None:
    result = _result()
    history = RunHistory(tmp_path / "history.sqlite3")
    run_id = history.record_run(result, file_hashes={}, report_paths={})
    legacy = [
        {
            key: value
            for key, value in finding.model_dump(mode="json").items()
            if key not in {"provenance", "subtype", "event_key"}
        }
        for finding in result.findings
    ]
    import sqlite3

    with sqlite3.connect(tmp_path / "history.sqlite3") as connection:
        connection.execute(
            "UPDATE runs SET findings = ? WHERE id = ?",
            (json.dumps(legacy), run_id),
        )

    record = history.get_run(run_id)
    regenerated = QCRunResult(profile_name="legacy", findings=record.findings)
    report = tmp_path / "legacy.xlsx"
    write_excel_report(regenerated, report)

    assert report.exists()
    assert "pattern review items" in render_html_report(regenerated)
    assert result_payload(regenerated, include_review_summary=True)["review_summary"]
