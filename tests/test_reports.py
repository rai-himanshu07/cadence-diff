"""Report generation tests (criterion 11)."""

import json
from pathlib import Path

from openpyxl import load_workbook

from qc_tool.coverage import MappingCoverage
from qc_tool.engine import QCRunResult, run_qc
from qc_tool.findings import Finding, FindingClass, SeriesAnchorV1, Severity
from qc_tool.history.review_state import finding_evidence_digest
from qc_tool.package import (
    PackageArtifact,
    PackageManifest,
    PackageMember,
    PackageSide,
)
from qc_tool.report.excel_report import write_excel_report
from qc_tool.report.html_report import render_html_report, write_html_report
from qc_tool.report.json_report import result_payload
from qc_tool.review import build_pattern_groups


def test_report_cells_never_carry_xml_illegal_characters(tmp_path: Path) -> None:
    from openpyxl import Workbook

    from qc_tool.report.excel_report import _dynamic_cell, _safe_report_text

    assert _safe_report_text("line one\vline two") == "line one\nline two"
    assert _safe_report_text("page\fbreak") == "page\nbreak"
    assert _safe_report_text("bell\x07null\x00escape\x1b") == "bellnullescape"
    assert _safe_report_text("kept\ttab\nnewline\rreturn") == "kept\ttab\nnewline\rreturn"

    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    cell = _dynamic_cell(sheet, row=1, column=1, value="soft\vbreak\x08")
    assert cell.value == "soft\nbreak"
    path = tmp_path / "sanitized.xlsx"
    workbook.save(path)  # raised IllegalCharacterError before fix


def test_excel_report_structure(qc_result: QCRunResult, tmp_path: Path) -> None:
    path = tmp_path / "report.xlsx"
    write_excel_report(qc_result, path)

    workbook = load_workbook(path)
    assert workbook.sheetnames == [
        "Summary",
        "Stories",
        "Coverage",
        "Alignment Trust",
        "Review Groups",
        "Findings",
    ]

    findings_sheet = workbook["Findings"]
    assert findings_sheet.max_row == len(qc_result.findings) + 1
    assert findings_sheet["A1"].value == "ID"
    assert findings_sheet["A2"].value == "F0001"
    assert findings_sheet["B2"].value == "critical"
    assert findings_sheet.auto_filter.ref is not None

    summary = workbook["Summary"]
    summary_text = " ".join(
        str(cell.value) for row in summary.iter_rows() for cell in row if cell.value
    )
    assert "fixture" in summary_text  # profile name
    assert "cycle_comparison" in summary_text
    assert "ooxml-streaming:" in summary_text
    assert "current.xlsx" in summary_text
    assert "Critical pattern review items" in summary_text
    assert "Critical finding records" in summary_text
    assert "Critical represented changes" in summary_text

    coverage = workbook["Coverage"]
    assert coverage["A1"].value == "Artifact"
    assert coverage.max_row == len(qc_result.coverage) + 1

    review_groups = workbook["Review Groups"]
    assert review_groups["A1"].value == "Group"
    assert review_groups.max_row <= len(qc_result.findings) + 1
    assert review_groups["J2"].hyperlink is not None
    assert review_groups["J2"].hyperlink.location.startswith("'Findings'!A")
    assert "Clear any Findings sheet filter" in review_groups["J2"].hyperlink.tooltip
    assert summary["D3"].hyperlink is not None
    assert summary["D3"].hyperlink.location == "'Stories'!A1"
    assert summary["D4"].hyperlink is not None
    assert summary["D4"].hyperlink.location == "'Review Groups'!A1"
    assert summary["D7"].hyperlink is not None
    assert summary["D7"].hyperlink.location == "'Alignment Trust'!A1"

    alignment_trust = workbook["Alignment Trust"]
    assert alignment_trust["A1"].value == "Member"
    assert alignment_trust.max_row > 1

    stories_sheet = workbook["Stories"]
    assert stories_sheet["A1"].value == "Story"
    assert stories_sheet.max_row >= 2  # at least one story for the fixture pair


def test_excel_report_discloses_xlsb_degradation(fixture_dir: Path, tmp_path: Path) -> None:
    result = run_qc(
        baseline_excel=fixture_dir / "baseline.xlsb",
        current_excel=fixture_dir / "current.xlsb",
    )
    path = tmp_path / "xlsb_report.xlsx"
    write_excel_report(result, path)
    summary = load_workbook(path)["Summary"]
    summary_text = " ".join(
        str(cell.value) for row in summary.iter_rows() for cell in row if cell.value
    )
    assert "NOTE:" in summary_text and "degraded" in summary_text


def test_excel_report_keeps_formula_text_inert(tmp_path: Path) -> None:
    result = QCRunResult(profile_name="test")
    result.findings = [
        Finding(
            artifact="excel",
            finding_class=FindingClass.FORMULA_LOGIC_CHANGED,
            baseline_value='=HYPERLINK("https://example.invalid","open")',
            current_value="=1+1",
            message="formula changed",
            severity=Severity.WARNING,
        )
    ]
    path = tmp_path / "formula-text.xlsx"

    write_excel_report(result, path)

    workbook = load_workbook(path, data_only=False)
    findings = workbook["Findings"]
    assert findings["N2"].value.startswith("=HYPERLINK")
    assert findings["N2"].data_type == "s"
    assert findings["O2"].value == "=1+1"
    assert findings["O2"].data_type == "s"


def test_excel_report_population_column(tmp_path: Path) -> None:
    from qc_tool.findings import MembershipCodec, PopulationEvidence

    result = QCRunResult(profile_name="test")
    result.findings = [
        Finding(
            artifact="excel",
            finding_class=FindingClass.FORMULA_LOGIC_CHANGED,
            severity=Severity.WARNING,
            sheet="Data",
            location="B2:B20",
            baseline_location="B2:B20",
            element="population",
            message="19 cells share one population",
            population=PopulationEvidence(
                member_count=19,
                membership=MembershipCodec(
                    current_rectangles=("B2:B20",),
                    baseline_mode="shift",
                    shift=(0, 0),
                    member_count=19,
                ),
                first="B2",
                last="B20",
                shape_before_digest="a" * 64,
                shape_after_digest="b" * 64,
            ),
        ),
        Finding(
            artifact="excel",
            finding_class=FindingClass.VALUE_CHANGED,
            severity=Severity.CRITICAL,
            sheet="Data",
            location="C1",
            baseline_value="1",
            current_value="2",
            message="value changed",
        ),
    ]
    path = tmp_path / "population.xlsx"

    write_excel_report(result, path)

    workbook = load_workbook(path)
    findings = workbook["Findings"]
    assert findings["A1"].value == "ID"
    header_row = [cell.value for cell in findings[1]]
    assert header_row[-1] == "Population"
    assert findings.cell(row=2, column=len(header_row)).value == (
        "19 cells; B2:B20; shift (0, 0)"
    )
    assert findings.cell(row=3, column=len(header_row)).value in (None, "")
    summary_values = {
        row[0].value: row[1].value
        for row in workbook["Summary"].iter_rows(min_col=1, max_col=2)
    }
    assert summary_values["Warning finding records"] == "1"
    assert summary_values["Warning represented changes"] == "19"


def test_excel_report_renders_a_renamed_sheet_alias_beside_the_physical_name(
    tmp_path: Path,
) -> None:
    """plan-20260913 Step 10: "aliases plus physical coordinates render in
    human reports" -- a finding whose logical address resolves to a sheet
    renamed since baseline shows both names, sourced only from the run's
    own resolved configuration.
    """
    from qc_tool.config.input_contract import INPUT_CONTRACT_VERSION
    from qc_tool.config.resolved_input import (
        ResolvedInputConfigurationV1,
        ResolvedMember,
        ResolvedSheet,
    )
    from qc_tool.findings import LogicalFindingAddress

    result = QCRunResult(profile_name="test")
    result.resolved_input_configuration = ResolvedInputConfigurationV1(
        inspection_contract_version=INPUT_CONTRACT_VERSION,
        members=(
            ResolvedMember(
                member_id="primary",
                sheets=(
                    ResolvedSheet(
                        sheet_id="sheet-revenue",
                        baseline_sheet_name="Revenue FY25",
                        current_sheet_name="Revenue FY26",
                    ),
                ),
            ),
        ),
    )
    result.findings = [
        Finding(
            artifact="excel",
            finding_class=FindingClass.VALUE_CHANGED,
            severity=Severity.CRITICAL,
            sheet="Revenue FY26",
            location="C1",
            baseline_value="1",
            current_value="2",
            message="value changed",
            logical_address=LogicalFindingAddress(
                member_id="primary", sheet_id="sheet-revenue"
            ),
        )
    ]
    path = tmp_path / "renamed-sheet.xlsx"

    write_excel_report(result, path)

    workbook = load_workbook(path)
    findings = workbook["Findings"]
    header_row = [cell.value for cell in findings[1]]
    sheet_column = header_row.index("Sheet / Slide") + 1
    assert findings.cell(row=2, column=sheet_column).value == (
        "Revenue FY26 (renamed from Revenue FY25)"
    )


def test_html_report_renders_a_renamed_sheet_alias_beside_the_physical_name() -> None:
    from qc_tool.config.input_contract import INPUT_CONTRACT_VERSION
    from qc_tool.config.resolved_input import (
        ResolvedInputConfigurationV1,
        ResolvedMember,
        ResolvedSheet,
    )
    from qc_tool.findings import LogicalFindingAddress

    result = QCRunResult(profile_name="test")
    result.resolved_input_configuration = ResolvedInputConfigurationV1(
        inspection_contract_version=INPUT_CONTRACT_VERSION,
        members=(
            ResolvedMember(
                member_id="primary",
                sheets=(
                    ResolvedSheet(
                        sheet_id="sheet-revenue",
                        baseline_sheet_name="Revenue FY25",
                        current_sheet_name="Revenue FY26",
                    ),
                ),
            ),
        ),
    )
    result.findings = [
        Finding(
            artifact="excel",
            finding_class=FindingClass.VALUE_CHANGED,
            severity=Severity.CRITICAL,
            sheet="Revenue FY26",
            location="C1",
            baseline_value="1",
            current_value="2",
            message="value changed",
            logical_address=LogicalFindingAddress(
                member_id="primary", sheet_id="sheet-revenue"
            ),
        )
    ]

    html = render_html_report(result)

    assert "Revenue FY26 (renamed from Revenue FY25)" in html


def test_html_report_population_column(tmp_path: Path) -> None:
    from qc_tool.findings import MembershipCodec, PopulationEvidence

    result = QCRunResult(profile_name="test")
    result.findings = [
        Finding(
            artifact="excel",
            finding_class=FindingClass.FORMULA_LOGIC_CHANGED,
            severity=Severity.WARNING,
            sheet="Data",
            location="B2:B20",
            baseline_location="B2:B20",
            element="population",
            message="19 cells share one population",
            population=PopulationEvidence(
                member_count=19,
                membership=MembershipCodec(
                    current_rectangles=("B2:B20",),
                    baseline_mode="shift",
                    shift=(0, 0),
                    member_count=19,
                ),
                first="B2",
                last="B20",
                shape_before_digest="a" * 64,
                shape_after_digest="b" * 64,
            ),
        )
    ]

    html = render_html_report(result)

    assert "<th>Population</th>" in html
    assert "19 cells; B2:B20; shift (0, 0)" in html
    assert "1 finding record" in html
    assert "19 represented changes" in html


def test_mapping_reports_include_unavailable_and_total_surfaces(
    tmp_path: Path,
) -> None:
    result = QCRunResult(
        profile_name="mapping",
        mapping_coverage=MappingCoverage(
            eligible=26,
            unavailable=1,
            mapped=1,
            verified=1,
            unmapped=25,
        ),
        disclosures=["Mapping disclosure remains visible"],
    )

    html = render_html_report(result)
    assert "26 readable" in html
    assert "1 unavailable" in html
    assert "27 total surfaces" in html

    path = tmp_path / "mapping.xlsx"
    write_excel_report(result, path)
    workbook = load_workbook(path)
    summary = workbook["Summary"]
    values = {
        row[0].value: row[1].value
        for row in summary.iter_rows(min_col=1, max_col=2)
        if row[0].value
    }
    assert values["Readable PPT figures"] == 26
    assert values["Unavailable PPT surfaces"] == 1
    assert values["Total PPT surfaces"] == 27
    summary_text = " ".join(
        str(cell.value) for row in summary.iter_rows() for cell in row if cell.value
    )
    assert "Mapping disclosure remains visible" in summary_text


def test_html_report_contents(qc_result: QCRunResult, tmp_path: Path) -> None:
    path = tmp_path / "report.html"
    write_html_report(qc_result, path)
    html = path.read_text(encoding="utf-8")

    assert html.startswith("<!DOCTYPE html>")
    assert "F0001" in html
    assert "cycle_comparison" in html
    assert "Check coverage" in html
    for severity in Severity:
        assert f'class="card {severity.value}"' in html
    # Self-contained: no external asset references.
    assert "http://" not in html and "https://" not in html
    assert "pattern review items" in html
    assert "finding records" in html
    assert "represented changes" in html
    # Every finding record remains available in safely escaped inline JSON.
    assert all(finding.finding_id in html for finding in qc_result.findings)
    assert "data-member-body" in html
    assert "const pageSize = 50" in html


def test_html_report_escapes_client_content() -> None:
    hostile = QCRunResult(profile_name="test")
    finding = Finding(
        artifact="excel",
        finding_class=FindingClass.VALUE_CHANGED,
        severity=Severity.CRITICAL,
        finding_id="F0001",
        sheet="Sheet1",
        location="A1",
        baseline_value="<script>alert('x')</script>",
        current_value="<img src=x onerror=alert(1)>",
        message="cell containing <b>markup</b> changed",
    )
    hostile.findings = [finding]
    html = render_html_report(hostile)
    assert "<script>alert" not in html
    assert "<img src=x" not in html
    assert r"\u003cscript\u003e" in html
    assert r"\u003cimg src=x onerror=alert(1)\u003e" in html
    assert "&lt;b&gt;markup&lt;/b&gt;" in html


def test_json_context_is_private_and_opt_in(qc_result: QCRunResult) -> None:
    private_payload = result_payload(qc_result)
    assert private_payload["schema_version"] == 1
    assert private_payload["context_included"] is False
    assert private_payload["mapping_suggestions"] == []
    assert all(
        "baseline_excerpt" not in finding and "current_excerpt" not in finding
        for finding in private_payload["findings"]
    )
    assert all(
        {"expected_reason", "temporal_context", "evidence_tags"} <= finding.keys()
        for finding in private_payload["findings"]
    )

    diagnostic_payload = result_payload(qc_result, include_context=True)
    assert diagnostic_payload["context_included"] is True
    assert any(
        finding.get("baseline_excerpt") or finding.get("current_excerpt")
        for finding in diagnostic_payload["findings"]
    )


def _package_result() -> QCRunResult:
    manifest = PackageManifest(
        members=(
            PackageMember(
                member_id="core",
                side=PackageSide.CURRENT,
                artifact=PackageArtifact.EXCEL,
                display_name="core.xlsx",
            ),
            PackageMember(
                member_id="ops",
                side=PackageSide.CURRENT,
                artifact=PackageArtifact.EXCEL,
                display_name="ops.xlsx",
            ),
            PackageMember(
                member_id="primary",
                side=PackageSide.CURRENT,
                artifact=PackageArtifact.PPT,
                display_name="deck.pptx",
            ),
        )
    )
    findings = [
        Finding(
            finding_id="F0001",
            artifact="excel",
            artifact_member="core",
            finding_class=FindingClass.VALUE_CHANGED,
            severity=Severity.CRITICAL,
            sheet="Data",
            location="B2",
            baseline_value="1",
            current_value="2",
            message="core changed",
        ),
        Finding(
            finding_id="F0002",
            artifact="excel",
            artifact_member="ops",
            finding_class=FindingClass.FORMULA_LOGIC_CHANGED,
            severity=Severity.WARNING,
            sheet="Data",
            baseline_location="C2",
            location="C2",
            baseline_value="=A1+B1",
            current_value="=A1+C1",
            message="ops formula changed",
        ),
    ]
    return QCRunResult(
        profile_name="package",
        files={member.role_key: member.display_name for member in manifest.members},
        findings=findings,
        package_manifest=manifest,
    )


def test_true_package_json_uses_v2_and_matches_shipped_schema_keys() -> None:
    payload = result_payload(_package_result(), include_review_summary=True)
    schema_path = Path(__file__).parents[1] / "qc_tool/report/findings-v2.schema.json"
    schema = json.loads(schema_path.read_text(encoding="utf-8"))

    assert payload["schema_version"] == 2
    assert payload["schema"].endswith("findings-v2.json")
    assert payload["package_manifest"]["version"] == 1
    assert payload["review_summary"]["summary_version"] == 4
    assert set(payload) <= set(schema["properties"])
    assert set(schema["required"]) <= set(payload)


def test_scalar_json_v1_matches_shipped_schema_keys(qc_result: QCRunResult) -> None:
    payload = result_payload(qc_result)
    schema_path = Path(__file__).parents[1] / "qc_tool/report/findings.schema.json"
    schema = json.loads(schema_path.read_text(encoding="utf-8"))

    assert payload["schema_version"] == 1
    assert "package_manifest" not in payload
    assert "review_summary" not in payload
    assert set(payload) <= set(schema["properties"])
    assert set(schema["required"]) <= set(payload)


def test_population_finding_json_uses_v3_and_matches_shipped_schema_keys() -> None:
    """Criterion 7: schema v3 is used wherever a population finding appears
    -- in a plain scalar run, not only a true-package run -- and every
    payload key stays a subset of the shipped schema's declared properties.
    """
    from qc_tool.findings import MembershipCodec, PopulationEvidence

    result = QCRunResult(
        profile_name="fixture",
        findings=[
            Finding(
                artifact="excel",
                finding_class=FindingClass.FORMULA_LOGIC_CHANGED,
                severity=Severity.WARNING,
                sheet="Data",
                location="B2:B20",
                element="population",
                message="19 cells share one population",
                population=PopulationEvidence(
                    member_count=19,
                    membership=MembershipCodec(
                        current_rectangles=("B2:B20",),
                        baseline_mode="shift",
                        shift=(0, 0),
                        member_count=19,
                    ),
                    first="B2",
                    last="B20",
                    shape_before_digest="a" * 64,
                    shape_after_digest="b" * 64,
                ),
            )
        ],
    )

    payload = result_payload(result)
    schema_path = Path(__file__).parents[1] / "qc_tool/report/findings-v3.schema.json"
    schema = json.loads(schema_path.read_text(encoding="utf-8"))

    assert payload["schema_version"] == 3
    assert payload["schema"].endswith("findings-v3.json")
    assert set(payload) <= set(schema["properties"])
    assert set(schema["required"]) <= set(payload)


def test_json_payload_carries_output_mode_and_matches_schema_for_every_version() -> None:
    """plan-20260910: `requested_output_mode`/`resolved_output_policy` are
    always present and stay within each shipped schema's declared
    properties -- for a null resolved policy (v1) and a real, non-null one
    (v2 via a package manifest) alike.
    """
    from qc_tool.config.profile import ReviewPolicy, resolve_output_policy
    from qc_tool.coverage import FindingOutputMode

    v1_result = QCRunResult(profile_name="fixture")
    v1_payload = result_payload(v1_result)
    v1_schema = json.loads(
        (Path(__file__).parents[1] / "qc_tool/report/findings.schema.json").read_text(
            encoding="utf-8"
        )
    )
    assert v1_payload["requested_output_mode"] == "profile"
    assert v1_payload["resolved_output_policy"] is None
    assert set(v1_payload) <= set(v1_schema["properties"])

    decision_result = _package_result()
    decision_result.requested_output_mode = FindingOutputMode.DECISION
    decision_result.formula_engines = {"current_excel": "native-biff12:1.2.3"}
    decision_result.values_engines = {"current_excel": "native-biff12:1.2.3"}
    resolved = resolve_output_policy(
        FindingOutputMode.DECISION,
        ReviewPolicy(),
    )
    decision_result.resolved_output_policy = resolved
    decision_payload = result_payload(decision_result, include_review_summary=True)
    v2_schema = json.loads(
        (Path(__file__).parents[1] / "qc_tool/report/findings-v2.schema.json").read_text(
            encoding="utf-8"
        )
    )
    assert decision_payload["requested_output_mode"] == "decision"
    # Compare against the production model serialization rather than a
    # duplicated hand-written policy payload.
    assert decision_payload["resolved_output_policy"] == resolved.model_dump(
        mode="json"
    )
    assert decision_payload["resolved_output_policy"]["populations"]["enabled"] is True
    assert decision_payload["formula_engines"] == decision_result.formula_engines
    assert decision_payload["values_engines"] == decision_result.values_engines
    assert set(decision_payload) <= set(v2_schema["properties"])
    assert set(v2_schema["required"]) <= set(decision_payload)


def test_json_payload_carries_resolved_input_configuration_and_matches_schema() -> None:
    """plan-20260913 Step 10: the exact per-run resolved logical
    configuration plus its canonical digest are always present in the JSON
    payload (``None``/"" with no saved input_contract) and stay within
    every shipped schema's declared properties.
    """
    from qc_tool.config.input_contract import INPUT_CONTRACT_VERSION
    from qc_tool.config.resolved_input import ResolvedInputConfigurationV1, ResolvedMember

    v1_result = QCRunResult(profile_name="fixture")
    v1_payload = result_payload(v1_result)
    v1_schema = json.loads(
        (Path(__file__).parents[1] / "qc_tool/report/findings.schema.json").read_text(
            encoding="utf-8"
        )
    )
    assert v1_payload["resolved_input_configuration"] is None
    assert v1_payload["resolved_input_digest"] == ""
    assert set(v1_payload) <= set(v1_schema["properties"])

    resolved = ResolvedInputConfigurationV1(
        inspection_contract_version=INPUT_CONTRACT_VERSION,
        members=(
            ResolvedMember(
                member_id="primary",
                baseline_source_sha256="a" * 64,
                current_source_sha256="b" * 64,
            ),
        ),
    )
    bound_result = QCRunResult(profile_name="fixture")
    bound_result.resolved_input_configuration = resolved
    bound_result.resolved_input_digest = resolved.canonical_sha256()
    bound_payload = result_payload(bound_result)
    assert bound_payload["resolved_input_configuration"] == resolved.model_dump(
        mode="json"
    )
    assert bound_payload["resolved_input_digest"] == resolved.canonical_sha256()
    assert set(bound_payload) <= set(v1_schema["properties"])


def test_multi_member_excel_report_has_package_and_member_columns(
    tmp_path: Path,
) -> None:
    result = _package_result()
    path = tmp_path / "package.xlsx"

    write_excel_report(result, path)
    workbook = load_workbook(path)

    assert "Package" in workbook.sheetnames
    assert workbook["Findings"]["D1"].value == "Member"
    assert workbook["Findings"]["D2"].value == "core"
    assert workbook["Findings"]["D3"].value == "ops"
    assert workbook["Review Groups"]["C1"].value == "Member"
    assert workbook["Review Groups"]["K2"].hyperlink is not None
    package = workbook["Package"]
    assert [package.cell(row=row, column=3).value for row in range(2, 5)] == [
        "core",
        "ops",
        "primary",
    ]
    assert workbook["Summary"]["D8"].hyperlink is not None
    assert workbook["Summary"]["D8"].hyperlink.location == "'Package'!A1"


def test_multi_member_html_is_autoescaped_and_renders_formula_diff() -> None:
    result = _package_result()
    assert result.package_manifest is not None
    hostile = result.package_manifest.members[0].model_copy(
        update={"display_name": "<script>alert(1)</script>.xlsx"}
    )
    result.package_manifest = PackageManifest(
        members=(hostile, *result.package_manifest.members[1:])
    )

    html = render_html_report(result)

    assert "Package members" in html
    assert "<script>alert(1)</script>.xlsx" not in html
    assert "&lt;script&gt;alert(1)&lt;/script&gt;.xlsx" in html
    assert "artifact_member" in html
    assert "Formula token diff" in html
    assert "fdiff-removed" in html
    assert "fdiff-added" in html


def test_series_anchors_leave_every_public_surface_byte_identical(
    qc_result: QCRunResult, tmp_path: Path
) -> None:
    before_json = json.dumps(result_payload(qc_result), sort_keys=True)
    before_html = render_html_report(qc_result)
    before_groups = [group.group_id for group in build_pattern_groups(qc_result.findings)]
    before_digests = [
        finding_evidence_digest(finding) for finding in qc_result.findings
    ]

    anchored = 0
    for finding in qc_result.findings:
        if finding.sheet is None or finding.location is None:
            continue
        anchored += 1
        finding.series_anchor = SeriesAnchorV1(
            sheet=finding.sheet,
            current_region_id=f"{finding.sheet}!A1:Z99",
            period_axis="rows",
            series_index=1,
            period_index=1,
        )
    assert anchored > 0

    assert json.dumps(result_payload(qc_result), sort_keys=True) == before_json
    assert render_html_report(qc_result) == before_html
    assert [
        group.group_id for group in build_pattern_groups(qc_result.findings)
    ] == before_groups
    assert [
        finding_evidence_digest(finding) for finding in qc_result.findings
    ] == before_digests

    path = tmp_path / "anchored.xlsx"
    write_excel_report(qc_result, path)
    workbook = load_workbook(path)
    text = " ".join(
        str(cell.value)
        for sheet in workbook.worksheets
        for row in sheet.iter_rows()
        for cell in row
        if cell.value is not None
    )
    assert "series_anchor" not in text
    assert "current_region_id" not in text


def test_findings_continue_onto_follow_on_sheets_past_the_row_cap(
    tmp_path: Path, monkeypatch
) -> None:
    """Step 6: the Excel report never truncates; it continues onto new sheets."""
    import qc_tool.report.excel_report as excel_report

    monkeypatch.setattr(excel_report, "MAX_FINDINGS_DATA_ROWS", 2)
    findings = [
        Finding(
            finding_id=f"F{index:04d}",
            artifact="excel",
            finding_class=FindingClass.VALUE_CHANGED,
            severity=Severity.CRITICAL,
            sheet="Data",
            location=f"B{index}",
            baseline_value="1",
            current_value="2",
            message=f"value changed {index}",
        )
        for index in range(1, 6)
    ]
    result = QCRunResult(profile_name="cap", findings=findings)
    path = tmp_path / "capped.xlsx"

    write_excel_report(result, path)

    workbook = load_workbook(path)
    assert [name for name in workbook.sheetnames if name.startswith("Findings")] == [
        "Findings",
        "Findings (2)",
        "Findings (3)",
    ]
    assert workbook["Findings"]["A2"].value == "F0001"
    assert workbook["Findings"]["A3"].value == "F0002"
    assert workbook["Findings (2)"]["A2"].value == "F0003"
    assert workbook["Findings (3)"]["A2"].value == "F0005"
    assert workbook["Findings (3)"].max_row == 2
    # Each continuation sheet keeps headers and its own filter range.
    assert workbook["Findings (2)"]["A1"].value == "ID"
    assert workbook["Findings (2)"].auto_filter.ref == "A1:U3"
    # Review-group links resolve into the continuation sheet that holds the row.
    review = workbook["Review Groups"]
    locations: set[str] = set()
    for row in range(2, review.max_row + 1):
        hyperlink = review.cell(row=row, column=10).hyperlink
        if hyperlink is not None and hyperlink.location is not None:
            locations.add(hyperlink.location)
    assert any(location.startswith("'Findings") for location in locations)


def test_streamed_json_report_equals_materialized_payload(
    qc_result: QCRunResult, tmp_path: Path, monkeypatch
) -> None:
    import qc_tool.report.json_report as json_report

    class _FrozenDatetime(json_report.dt.datetime):
        @classmethod
        def now(cls, tz=None):
            return cls(2026, 1, 1, tzinfo=json_report.dt.UTC)

    monkeypatch.setattr(json_report.dt, "datetime", _FrozenDatetime)
    from qc_tool.report.json_report import write_json_report

    path = tmp_path / "streamed.json"
    write_json_report(qc_result, path, include_review_summary=True)
    streamed = json.loads(path.read_text(encoding="utf-8"))
    materialized = result_payload(qc_result, include_review_summary=True)
    assert streamed == materialized
    # Byte-exact framing too, not just structural equality.
    assert path.read_text(encoding="utf-8") == json.dumps(materialized, indent=2)


def test_html_report_past_threshold_keeps_decisions_and_discloses(
    monkeypatch,
) -> None:
    import qc_tool.report.html_report as html_report

    monkeypatch.setattr(html_report, "ATOMICS_INLINE_THRESHOLD", 1)
    findings = [
        Finding(
            finding_id=f"F{index:04d}",
            artifact="excel",
            finding_class=FindingClass.VALUE_CHANGED,
            severity=Severity.CRITICAL,
            sheet="Data",
            location=f"B{index}",
            baseline_value="1",
            current_value="2",
            message=f"value changed {index}",
        )
        for index in range(1, 4)
    ]
    result = QCRunResult(profile_name="cap", findings=findings)

    html = render_html_report(result)

    assert "above the 1-row inline limit" in html
    assert "Excel and JSON exports" in html
    assert "data-member-body" not in html  # no embedded atomic payloads
    assert "value changed 1" not in html  # member text is not inlined
    assert 'class="review-group"' in html or "review-item" in html  # decisions stay


def test_streamed_html_write_matches_monolithic_render(
    qc_result: QCRunResult, tmp_path: Path
) -> None:
    """Step 8: the chunked template write is byte-identical to render()."""
    import re

    path = tmp_path / "streamed.html"
    write_html_report(qc_result, path)
    streamed = path.read_text(encoding="utf-8")
    rendered = render_html_report(qc_result)

    strip_stamp = lambda text: re.sub(  # noqa: E731
        r"\d{4}-\d{2}-\d{2}T\d{2}:\d{2}:\d{2}\+00:00", "STAMP", text
    )
    assert strip_stamp(streamed) == strip_stamp(rendered)
