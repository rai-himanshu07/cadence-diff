"""PPT extraction, fuzzy matching, and diff tests (criteria 7, 8)."""

import json
import struct
import xml.etree.ElementTree as ET
import zipfile
import zlib
from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import load_workbook
from pptx import Presentation
from pptx.util import Inches

from qc_tool.attestation import create_attestation, verify_attestation
from qc_tool.config.profile import DeliverableProfile, PptProfile
from qc_tool.coverage import CoverageState, QCRunMode
from qc_tool.engine import run_qc
from qc_tool.findings import Finding, FindingClass, Severity
from qc_tool.history.store import RunHistory
from qc_tool.ppt.diff import diff_decks
from qc_tool.ppt.extract import DeckSnapshot, SlideContent, load_deck_snapshot
from qc_tool.ppt.match import SlideMatching, match_slides
from qc_tool.ppt.preflight import preflight_deck
from qc_tool.report.excel_report import write_excel_report
from qc_tool.report.html_report import write_html_report
from qc_tool.report.json_report import result_payload
from tests.fixtures.manifest_schema import FixtureManifest
from tests.fixtures.ppt_builder import build_grouped_text_deck


@pytest.fixture(scope="module")
def base_deck(fixture_dir: Path) -> DeckSnapshot:
    return load_deck_snapshot(fixture_dir / "baseline.pptx")


@pytest.fixture(scope="module")
def curr_deck(fixture_dir: Path) -> DeckSnapshot:
    return load_deck_snapshot(fixture_dir / "current.pptx")


@pytest.fixture(scope="module")
def matching(base_deck: DeckSnapshot, curr_deck: DeckSnapshot) -> SlideMatching:
    return match_slides(base_deck, curr_deck)


@pytest.fixture(scope="module")
def findings(matching: SlideMatching) -> list[Finding]:
    return diff_decks(matching)


def _by_class(findings: list[Finding], cls: FindingClass) -> list[Finding]:
    return [f for f in findings if f.finding_class is cls]


def _png(red: int, green: int, blue: int) -> bytes:
    def chunk(kind: bytes, payload: bytes) -> bytes:
        return (
            struct.pack(">I", len(payload))
            + kind
            + payload
            + struct.pack(">I", zlib.crc32(kind + payload) & 0xFFFFFFFF)
        )

    return (
        b"\x89PNG\r\n\x1a\n"
        + chunk(b"IHDR", struct.pack(">IIBBBBB", 1, 1, 8, 2, 0, 0, 0))
        + chunk(b"IDAT", zlib.compress(bytes((0, red, green, blue))))
        + chunk(b"IEND", b"")
    )


def _picture_deck(path: Path, image: bytes) -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    slide.shapes.add_picture(
        BytesIO(image),
        Inches(1),
        Inches(1),
        width=Inches(2),
        height=Inches(1),
    )
    presentation.save(str(path))


def _duplicate_picture_deck(path: Path) -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    for image in (_png(255, 0, 0), _png(0, 0, 255)):
        shape = slide.shapes.add_picture(
            BytesIO(image),
            Inches(1),
            Inches(1),
            width=Inches(2),
            height=Inches(1),
        )
        shape.name = "Duplicate picture"
    presentation.save(str(path))


def _rewrite_picture_relationship(
    source: Path,
    destination: Path,
    *,
    external: bool,
) -> None:
    drawing_ns = "http://schemas.openxmlformats.org/drawingml/2006/main"
    office_rel_ns = (
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    )
    package_rel_ns = "http://schemas.openxmlformats.org/package/2006/relationships"
    slide_member = "ppt/slides/slide1.xml"
    rels_member = "ppt/slides/_rels/slide1.xml.rels"
    relationship_id = ""
    with zipfile.ZipFile(source) as archive, zipfile.ZipFile(
        destination, "w", zipfile.ZIP_DEFLATED
    ) as output:
        for info in archive.infolist():
            data = archive.read(info.filename)
            if info.filename == slide_member:
                root = ET.fromstring(data)
                blip = root.find(f".//{{{drawing_ns}}}blip")
                if blip is None:
                    raise ValueError("picture fixture has no drawing blip")
                embed = f"{{{office_rel_ns}}}embed"
                relationship_id = str(blip.attrib[embed])
                if external:
                    del blip.attrib[embed]
                    blip.set(f"{{{office_rel_ns}}}link", relationship_id)
                else:
                    blip.set(embed, "rId999")
                data = ET.tostring(root, encoding="utf-8", xml_declaration=True)
            elif info.filename == rels_member and external:
                root = ET.fromstring(data)
                relationship = next(
                    item
                    for item in root.findall(f"{{{package_rel_ns}}}Relationship")
                    if item.attrib.get("Id") == relationship_id
                )
                relationship.set("Target", "https://private.invalid/client-image.png")
                relationship.set("TargetMode", "External")
                data = ET.tostring(root, encoding="utf-8", xml_declaration=True)
            output.writestr(info, data)


def test_embedded_picture_bytes_are_hashed_without_decoding(tmp_path: Path) -> None:
    first_path = tmp_path / "first.pptx"
    same_path = tmp_path / "same.pptx"
    changed_path = tmp_path / "changed.pptx"
    _picture_deck(first_path, _png(255, 0, 0))
    _picture_deck(same_path, _png(255, 0, 0))
    _picture_deck(changed_path, _png(0, 0, 255))

    first = load_deck_snapshot(first_path)
    same = load_deck_snapshot(same_path)
    changed = load_deck_snapshot(changed_path)

    first_shape = first.slides[0].shapes[0]
    assert first_shape.media_kind == "image/png"
    assert len(first_shape.media_digest or "") == 64
    assert first_shape.media_digest == same.slides[0].shapes[0].media_digest
    assert first_shape.media_digest != changed.slides[0].shapes[0].media_digest
    assert first.media_available and changed.media_available


def _soft_break_deck(path: Path) -> None:
    """Deck whose text carries PowerPoint soft line breaks (<a:br/> = \\v)."""
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    title = slide.shapes.title
    assert title is not None
    title.text_frame.text = "Weekly Revenue\vRegional Outlook"
    box = slide.shapes.add_textbox(Inches(1), Inches(2), Inches(4), Inches(1))
    box.text_frame.text = "First line\vTBD second line"
    table = slide.shapes.add_table(
        1, 1, Inches(1), Inches(4), Inches(2), Inches(1)
    ).table
    table.cell(0, 0).text = "cell\vbreak"
    notes_frame = slide.notes_slide.notes_text_frame
    assert notes_frame is not None
    notes_frame.text = "note\vbreak"
    presentation.save(str(path))


def test_soft_line_breaks_are_normalized_at_extraction(tmp_path: Path) -> None:
    path = tmp_path / "soft-breaks.pptx"
    _soft_break_deck(path)

    deck = load_deck_snapshot(path)

    slide = deck.slides[0]
    assert slide.title == "Weekly Revenue\nRegional Outlook"
    assert slide.display_name == "Weekly Revenue\nRegional Outlook"
    assert "First line\nTBD second line" in slide.texts
    assert slide.tables[0].rows[0][0] == "cell\nbreak"
    assert slide.notes == ["note\nbreak"]
    everything = [
        slide.title or "",
        *slide.texts,
        *slide.notes,
        *(cell for table in slide.tables for row in table.rows for cell in row),
        *(text for shape in slide.shapes for text in shape.texts),
    ]
    assert not any("\v" in text for text in everything)


def test_nested_group_text_extracts_once_in_depth_first_shape_order(
    tmp_path: Path,
) -> None:
    path = build_grouped_text_deck(
        tmp_path / "grouped.pptx",
        top_before=("Before 1",),
        grouped_lines=("Outer 2",),
        nested_lines=("Nested 3",),
        top_after=("After 4",),
    )

    slide = load_deck_snapshot(path).slides[0]

    assert slide.texts == [
        "Grouped KPIs",
        "Before 1",
        "Outer 2",
        "Nested 3",
        "After 4",
    ]
    assert slide.texts.count("Nested 3") == 1


def test_grouped_text_changes_flow_through_slide_matching_and_diff(
    tmp_path: Path,
) -> None:
    baseline = load_deck_snapshot(
        build_grouped_text_deck(
            tmp_path / "baseline-grouped.pptx",
            grouped_lines=("Revenue $100M for Jan-26",),
            nested_lines=("Margin 20% for Jan-26",),
        )
    )
    current = load_deck_snapshot(
        build_grouped_text_deck(
            tmp_path / "current-grouped.pptx",
            grouped_lines=("Revenue $110M for Jan-26",),
            nested_lines=("Margin 20% for Jan-26",),
        )
    )

    findings = diff_decks(match_slides(baseline, current))
    changed = _by_class(findings, FindingClass.SLIDE_TEXT_CHANGED)

    assert len(changed) == 1
    assert changed[0].baseline_value == "Revenue $100M for Jan-26"
    assert changed[0].current_value == "Revenue $110M for Jan-26"
    assert changed[0].expected_reason is not None

    unchanged = diff_decks(match_slides(current, current))
    assert _by_class(unchanged, FindingClass.SLIDE_TEXT_CHANGED) == []


def test_whitespace_only_ppt_text_alignment_does_not_create_a_finding() -> None:
    baseline = SlideContent(
        index=0,
        title="KPIs",
        texts=["New infections\t1.3 million\t[1.0-1.7 million]"],
        shape_count=1,
    )
    current = SlideContent(
        index=0,
        title="KPIs",
        texts=["New infections\t        1.3 million\t[1.0-1.7 million]"],
        shape_count=1,
    )

    findings = diff_decks(SlideMatching(pairs=[(baseline, current)]))

    assert _by_class(findings, FindingClass.SLIDE_TEXT_CHANGED) == []


def test_preflight_report_survives_soft_line_break_titles(tmp_path: Path) -> None:
    path = tmp_path / "current.pptx"
    _soft_break_deck(path)

    result = run_qc(current_ppt=path, mode=QCRunMode.CURRENT_FILE_PREFLIGHT)

    draft = [
        finding
        for finding in result.findings
        if finding.finding_class is FindingClass.PPT_DRAFT_TOKEN
    ]
    assert draft, "fixture must produce a finding on the soft-break slide"
    assert all("\v" not in (finding.slide or "") for finding in draft)

    report = tmp_path / "report.xlsx"
    write_excel_report(result, report)  # raised IllegalCharacterError before fix

    findings_sheet = load_workbook(report)["Findings"]
    labels = [
        row[4] for row in findings_sheet.iter_rows(min_row=2, values_only=True)
    ]
    assert any(
        label and "Weekly Revenue" in str(label) for label in labels
    )


def test_linked_and_malformed_picture_relationships_degrade_without_leakage(
    tmp_path: Path,
) -> None:
    embedded_path = tmp_path / "embedded.pptx"
    linked_path = tmp_path / "linked.pptx"
    malformed_path = tmp_path / "malformed.pptx"
    _picture_deck(embedded_path, _png(1, 2, 3))
    _rewrite_picture_relationship(embedded_path, linked_path, external=True)
    _rewrite_picture_relationship(embedded_path, malformed_path, external=False)

    linked = load_deck_snapshot(linked_path)
    malformed = load_deck_snapshot(malformed_path)

    linked_shape = linked.slides[0].shapes[0]
    malformed_shape = malformed.slides[0].shapes[0]
    assert linked_shape.media_kind == "linked-image"
    assert linked_shape.media_digest is None
    assert not linked.media_available
    assert "linked-image" in linked.media_detail
    assert "private.invalid" not in linked.media_detail
    assert malformed_shape.media_kind == "unavailable-image"
    assert malformed_shape.media_digest is None
    assert not malformed.media_available
    assert "unreadable-image-relationship" in malformed.media_detail
    linked_coverage = next(
        item
        for item in preflight_deck(linked, PptProfile()).coverage
        if item.check_id == "ppt-media-structural"
    )
    assert linked_coverage.state is CoverageState.DEGRADED
    assert "private.invalid" not in linked_coverage.detail


def test_media_change_roundtrips_without_digest_or_content_leakage(
    tmp_path: Path,
) -> None:
    baseline_path = tmp_path / "baseline.pptx"
    current_path = tmp_path / "current.pptx"
    _picture_deck(baseline_path, _png(255, 0, 0))
    _picture_deck(current_path, _png(0, 0, 255))
    baseline_digest = load_deck_snapshot(baseline_path).slides[0].shapes[0].media_digest
    current_digest = load_deck_snapshot(current_path).slides[0].shapes[0].media_digest
    assert baseline_digest is not None and current_digest is not None

    result = run_qc(baseline_ppt=baseline_path, current_ppt=current_path)

    media = [
        finding
        for finding in result.findings
        if finding.finding_class is FindingClass.PPT_MEDIA_CHANGED
    ]
    assert len(media) == 1
    finding = media[0]
    assert finding.severity is Severity.WARNING
    assert finding.slide_index == 1 and finding.baseline_slide_index == 1
    assert finding.focus_shape_id and finding.baseline_focus_shape_id
    structural = next(
        item
        for item in result.coverage
        if item.check_id == "ppt-media-structural"
    )
    visual = next(
        item for item in result.coverage if item.check_id == "ppt-media-visual"
    )
    assert structural.state is CoverageState.CHECKED and structural.findings == 1
    assert visual.state is CoverageState.UNAVAILABLE

    serialized = json.dumps(result_payload(result, include_context=True))
    assert FindingClass.PPT_MEDIA_CHANGED.value in serialized
    assert baseline_digest not in serialized
    assert current_digest not in serialized
    assert "255, 0, 0" not in serialized and "0, 0, 255" not in serialized

    html_path = tmp_path / "report.html"
    excel_path = tmp_path / "report.xlsx"
    write_html_report(result, html_path)
    write_excel_report(result, excel_path)
    history = RunHistory(tmp_path / "history.sqlite3")
    run_id = history.record_run(result, file_hashes={}, report_paths={})
    restored = history.get_run(run_id)
    assert any(
        item.finding_class is FindingClass.PPT_MEDIA_CHANGED
        for item in restored.findings
    )

    key = b"m" * 32
    attestation = create_attestation(
        tmp_path / "media.qca",
        result=result,
        profile=DeliverableProfile(name="default"),
        input_files={"baseline_ppt": baseline_path, "current_ppt": current_path},
        report_paths={"html": html_path, "excel": excel_path},
        key=key,
    )
    assert verify_attestation(attestation, key=key).valid


def test_duplicate_media_keys_degrade_cycle_coverage_without_guessing(
    tmp_path: Path,
) -> None:
    baseline_path = tmp_path / "baseline-duplicates.pptx"
    current_path = tmp_path / "current-duplicates.pptx"
    _duplicate_picture_deck(baseline_path)
    _duplicate_picture_deck(current_path)

    result = run_qc(baseline_ppt=baseline_path, current_ppt=current_path)

    assert not [
        finding
        for finding in result.findings
        if finding.finding_class is FindingClass.PPT_MEDIA_CHANGED
    ]
    structural = next(
        item
        for item in result.coverage
        if item.check_id == "ppt-media-structural"
    )
    assert structural.state is CoverageState.DEGRADED
    assert "2 media shape(s)" in structural.detail


def test_extraction_shapes(base_deck: DeckSnapshot, curr_deck: DeckSnapshot) -> None:
    assert [s.title for s in base_deck.slides] == [
        "Executive Summary",
        "Revenue by Region",
        "Revenue Trend",
        "Weekly Ops",
        "Notes & Definitions",
        "Deep Dive Archive",
    ]
    trend = next(s for s in curr_deck.slides if s.title == "Revenue Trend")
    assert trend.charts[0].categories[-1] == "Jun-26"
    region = next(s for s in curr_deck.slides if s.title == "Revenue by Region")
    assert region.tables[0].rows[0][0] == "Region"


def test_fuzzy_matching_survives_reorder_add_delete(matching: SlideMatching) -> None:
    pair_titles = {(b.title, c.title) for b, c in matching.pairs}
    assert pair_titles == {
        ("Executive Summary", "Executive Summary"),
        ("Revenue by Region", "Revenue by Region"),
        ("Revenue Trend", "Revenue Trend"),
        ("Weekly Ops", "Weekly Ops"),
        ("Notes & Definitions", "Notes & Definitions"),
    }
    assert [s.title for s in matching.added] == ["New Initiatives"]  # P01
    assert [s.title for s in matching.removed] == ["Deep Dive Archive"]  # P02
    # PX01: only the deliberately moved slide counts as reordered — pure
    # index shifts caused by the insertion do not.
    assert [(b.title, c.index) for b, c in matching.reordered] == [
        ("Notes & Definitions", 1)
    ]


def test_slide_pins_override(base_deck: DeckSnapshot, curr_deck: DeckSnapshot) -> None:
    profile = PptProfile(
        slide_pins={"Deep Dive Archive": "New Initiatives"}, match_threshold=55.0
    )
    pinned = match_slides(base_deck, curr_deck, profile)
    assert ("Deep Dive Archive", "New Initiatives") in {
        (b.title, c.title) for b, c in pinned.pairs
    }
    assert pinned.added == [] and pinned.removed == []


def test_slide_add_remove_reorder_findings(findings: list[Finding]) -> None:
    assert [f.slide for f in _by_class(findings, FindingClass.SLIDE_ADDED)] == [
        "New Initiatives"
    ]
    assert [f.slide for f in _by_class(findings, FindingClass.SLIDE_REMOVED)] == [
        "Deep Dive Archive"
    ]
    reordered = _by_class(findings, FindingClass.SLIDE_REORDERED)
    assert [(f.slide, f.expected_growth) for f in reordered] == [
        ("Notes & Definitions", True)
    ]


def test_reorder_message_distinguishes_relative_displacement() -> None:
    """LIS flags relative displacement; equal positions must not read 'N to N'."""

    def slide(index: int, title: str) -> SlideContent:
        return SlideContent(index=index, title=title, texts=[], shape_count=1)

    matching = SlideMatching(
        reordered=[
            (slide(6, "Held"), slide(6, "Held")),
            (slide(7, "Moved"), slide(5, "Moved")),
        ]
    )

    messages = {
        finding.slide: finding.message
        for finding in _by_class(diff_decks(matching), FindingClass.SLIDE_REORDERED)
    }
    assert messages["Held"] == (
        "slide 'Held' kept position 7 while surrounding slides moved"
    )
    assert messages["Moved"] == "slide 'Moved' moved from position 8 to 6"


def test_text_changes_split_wording_from_figures(
    findings: list[Finding], manifest: FixtureManifest
) -> None:
    texts = _by_class(findings, FindingClass.SLIDE_TEXT_CHANGED)
    unexpected = [f for f in texts if not f.expected_growth]
    p03 = manifest.defect("P03")
    assert [(f.slide, f.baseline_value, f.current_value) for f in unexpected] == [
        ("Executive Summary", p03.baseline, p03.current)
    ]
    # Figure lines (revenue/margin bullets) refresh each cycle: expected.
    expected = [f for f in texts if f.expected_growth]
    assert len(expected) == 2
    assert all(f.slide == "Executive Summary" for f in expected)


def test_table_diff(findings: list[Finding], manifest: FixtureManifest) -> None:
    tables = _by_class(findings, FindingClass.TABLE_VALUE_CHANGED)
    assert tables and all(
        finding.focus_shape_id and finding.baseline_focus_shape_id
        for finding in tables
    )
    p04 = manifest.defect("P04")
    unexpected = [f for f in tables if not f.expected_growth]
    assert [(f.slide, f.element, f.baseline_value, f.current_value) for f in unexpected] == [
        ("Revenue by Region", p04.element, p04.baseline, p04.current)
    ]
    expected = [f for f in tables if f.expected_growth]
    assert [(f.slide, f.element) for f in expected] == [
        ("Revenue by Region", "Jun-26")
    ]  # PX04


def test_chart_diff_full_history(findings: list[Finding], manifest: FixtureManifest) -> None:
    charts = [
        f
        for f in _by_class(findings, FindingClass.CHART_VALUE_CHANGED)
        if f.slide == "Revenue Trend"
    ]
    assert charts and all(
        finding.focus_shape_id and finding.baseline_focus_shape_id
        for finding in charts
    )
    p05 = manifest.defect("P05")
    unexpected = [f for f in charts if not f.expected_growth]
    assert len(unexpected) == 1
    finding = unexpected[0]
    assert finding.element == p05.element  # Mar-26
    assert float(finding.baseline_value or "") == float(p05.baseline or "")
    assert float(finding.current_value or "") == float(p05.current or "")
    expected = [f for f in charts if f.expected_growth]
    assert [f.element for f in expected] == ["Jun-26"]  # PX03


def test_chart_diff_rolling_window(findings: list[Finding]) -> None:
    weekly = [
        f
        for f in _by_class(findings, FindingClass.CHART_VALUE_CHANGED)
        if f.slide == "Weekly Ops"
    ]
    # PX02: the whole rolling advance is expected — no unexpected findings.
    assert weekly and all(f.expected_growth for f in weekly)
    window = next(f for f in weekly if f.element == "window")
    assert window.baseline_value == "W17..W20"
    assert window.current_value == "W18..W21"


def test_chart_window_profile_override(matching: SlideMatching) -> None:
    profile = PptProfile(chart_windows={"Weekly Ops": "full"})
    findings = diff_decks(matching, profile)
    weekly = [
        f
        for f in findings
        if f.finding_class is FindingClass.CHART_VALUE_CHANGED and f.slide == "Weekly Ops"
    ]
    removed = [f for f in weekly if not f.expected_growth and "removed" in f.message]
    assert [f.element for f in removed] == ["W17"]  # forced full-history: drop-off flags
