"""Claim identity, population reconciliation, and mapping durability.

The six cycles below are the ones `docs/COMPETITIVE_STRATEGY.md` names as the
durability proof: slide reorder, wording changes, sheet/range movement, period
roll, restatement, and duplicate equal values. Each is generated here, so the
measurement is executable and carries no private content.
"""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches

from qc_tool.config.profile import CrosscheckMapping, CrosscheckProfile
from qc_tool.coverage import MappingCoverage
from qc_tool.crosscheck.claims import claim_identity, claim_population
from qc_tool.crosscheck.package import reconcile_package
from qc_tool.crosscheck.trace import extract_deck_figures
from qc_tool.io.loader import load_workbook_snapshot
from qc_tool.ppt.extract import load_deck_snapshot
from qc_tool.ppt.model import DeckSnapshot, ShapeContent, SlideContent


def _deck(path: Path, slides: list[tuple[str, list[str]]]) -> Path:
    presentation = Presentation()
    title_only = presentation.slide_layouts[5]
    for title, lines in slides:
        slide = presentation.slides.add_slide(title_only)
        placeholder = slide.shapes.title
        assert placeholder is not None
        placeholder.text = title
        body = slide.shapes.add_textbox(Inches(0.5), Inches(1.8), Inches(9), Inches(4))
        frame = body.text_frame
        frame.text = lines[0]
        for line in lines[1:]:
            frame.add_paragraph().text = line
    presentation.save(str(path))
    return path


def _workbook(path: Path, cells: dict[str, dict[str, float | str]]) -> Path:
    workbook = Workbook()
    default = workbook.active
    assert default is not None
    workbook.remove(default)
    for sheet_name, values in cells.items():
        sheet = workbook.create_sheet(sheet_name)
        for ref, value in values.items():
            sheet[ref] = value
    workbook.save(path)
    return path


_BASE_SLIDES = [
    ("Executive Summary", ["Total revenue $120M", "Margin 41.5%"]),
    ("Detail", ["Headcount 1,250", "Attrition 8.0%"]),
]
_BASE_CELLS = {
    "Dashboard": {"A1": "Total revenue", "B1": 120.0, "A2": "Margin", "B2": 0.415},
    "People": {"A1": "Headcount", "B1": 1250.0, "A2": "Attrition", "B2": 0.08},
}


def _mappings() -> list[CrosscheckMapping]:
    return [
        CrosscheckMapping(
            slide="Executive Summary",
            line_skeleton="Total revenue $#M",
            figure_index=0,
            label="Total revenue",
            source_sheet="Dashboard",
            source_cell="B1",
        ),
        CrosscheckMapping(
            slide="Detail",
            line_skeleton="Headcount #",
            figure_index=0,
            label="Headcount",
            source_sheet="People",
            source_cell="B1",
        ),
    ]


def _profile() -> CrosscheckProfile:
    return CrosscheckProfile(mappings=_mappings())


def _survivors(deck_path: Path) -> set[str]:
    """Which mapped claim identities still exist in this cycle's deck."""
    declared = {
        f"{mapping.slide}|{mapping.line_skeleton}|{mapping.figure_index}"
        for mapping in _mappings()
    }
    present = {
        claim_identity(occurrence).key
        for occurrence in extract_deck_figures(load_deck_snapshot(deck_path))
    }
    return declared & present


# --- identity ------------------------------------------------------------


def test_claim_identity_distinguishes_duplicate_equal_values(tmp_path: Path) -> None:
    path = _deck(
        tmp_path / "dup.pptx",
        [("Executive Summary", ["Revenue 100 and target 100", "Margin 41.5%"])],
    )

    identities = [
        claim_identity(occurrence)
        for occurrence in extract_deck_figures(load_deck_snapshot(path))
    ]
    keys = [identity.key for identity in identities]

    assert len(keys) == len(set(keys)), "equal values on one line collapsed"


def test_claim_identity_key_is_stable_and_readable() -> None:
    from qc_tool.crosscheck.claims import ClaimIdentity

    identity = ClaimIdentity("Executive Summary", "Total revenue $#M", 0)

    assert identity.key == "Executive Summary|Total revenue $#M|0"


# --- population ----------------------------------------------------------


def test_population_reconciles_with_the_mapped_split(tmp_path: Path) -> None:
    deck_path = _deck(tmp_path / "deck.pptx", _BASE_SLIDES)
    workbook = load_workbook_snapshot(_workbook(tmp_path / "book.xlsx", _BASE_CELLS))
    deck = load_deck_snapshot(deck_path)

    result = reconcile_package(workbook, deck, _profile())

    coverage = result.mapping_coverage
    assert coverage.reconciles
    assert coverage.mapped + coverage.unmapped == coverage.eligible


def test_rasterized_surface_is_counted_as_unavailable_not_ignored() -> None:
    deck = DeckSnapshot(
        source_name="deck.pptx",
        slides=[
            SlideContent(
                index=0,
                title="Executive Summary",
                texts=["Total revenue $120M"],
                shapes=[
                    ShapeContent(
                        source_id="s0",
                        source_index=0,
                        shape_id=2,
                        shape_type="PICTURE",
                        name="Screenshot",
                        left=0,
                        top=0,
                        width=10,
                        height=10,
                        z_order=0,
                    )
                ],
            )
        ],
    )

    population = claim_population(deck, extract_deck_figures(deck))

    assert population.unreadable_shapes == 1
    assert population.unreadable_slides == (1,)
    assert population.complete is False


def test_readable_deck_reports_a_complete_population(tmp_path: Path) -> None:
    deck = load_deck_snapshot(_deck(tmp_path / "deck.pptx", _BASE_SLIDES))

    population = claim_population(deck, extract_deck_figures(deck))

    assert population.complete is True
    assert population.unreadable_shapes == 0


def test_unavailable_surfaces_degrade_crosscheck_coverage(tmp_path: Path) -> None:
    workbook = load_workbook_snapshot(_workbook(tmp_path / "book.xlsx", _BASE_CELLS))
    deck = load_deck_snapshot(_deck(tmp_path / "deck.pptx", _BASE_SLIDES))
    deck.slides[0].shapes.append(
        ShapeContent(
            source_id="s9",
            source_index=9,
            shape_id=99,
            shape_type="PICTURE",
            name="Chart image",
            left=0,
            top=0,
            width=10,
            height=10,
            z_order=9,
        )
    )

    result = reconcile_package(workbook, deck, _profile())

    item = next(
        item for item in result.coverage if item.check_id == "excel-ppt-crosscheck"
    )
    assert result.mapping_coverage.unavailable == 1
    assert item.state.value == "degraded"
    assert "cannot be read" in item.detail


def test_legacy_mapping_coverage_defaults_to_no_unavailable_claims() -> None:
    legacy = MappingCoverage.model_validate({"eligible": 4, "mapped": 2, "unmapped": 2})

    assert legacy.unavailable == 0
    assert legacy.reconciles is True


# --- durability across six generated cycles ------------------------------

_CYCLES: dict[str, list[tuple[str, list[str]]]] = {
    "slide_reorder": list(reversed(_BASE_SLIDES)),
    "wording_change": [
        # The mapped line is the one reworded, so this cycle actually exercises
        # a mapping that cannot survive rather than one that trivially does.
        ("Executive Summary", ["Total revenues $120M", "Margin 41.5%"]),
        ("Detail", ["Headcount 1,250", "Attrition 8.0%"]),
    ],
    "period_roll": [
        ("Executive Summary", ["Total revenue $131M", "Margin 42.9%"]),
        ("Detail", ["Headcount 1,310", "Attrition 7.4%"]),
    ],
    "restatement": [
        ("Executive Summary", ["Total revenue $118M", "Margin 40.8%"]),
        ("Detail", ["Headcount 1,244", "Attrition 8.2%"]),
    ],
    "duplicate_equal_values": [
        ("Executive Summary", ["Total revenue $120M", "Margin 41.5%"]),
        ("Detail", ["Headcount 1,250", "Attrition 8.0%", "Target 1,250 and plan 1,250"]),
    ],
    "sheet_range_movement": _BASE_SLIDES,
}

_MOVED_CELLS = {
    "Dashboard 2026": {"A7": "Total revenue", "B7": 120.0, "A8": "Margin", "B8": 0.415},
    "People": {"A5": "Headcount", "B5": 1250.0, "A6": "Attrition", "B6": 0.08},
}


@pytest.mark.parametrize("cycle", sorted(_CYCLES))
def test_each_cycle_reconciles_exactly(cycle: str, tmp_path: Path) -> None:
    deck = load_deck_snapshot(_deck(tmp_path / f"{cycle}.pptx", _CYCLES[cycle]))
    cells = _MOVED_CELLS if cycle == "sheet_range_movement" else _BASE_CELLS
    workbook = load_workbook_snapshot(_workbook(tmp_path / f"{cycle}.xlsx", cells))

    coverage = reconcile_package(workbook, deck, _profile()).mapping_coverage

    assert coverage.reconciles, f"{cycle} population does not reconcile"
    assert coverage.eligible > 0


@pytest.mark.parametrize(
    ("cycle", "expected_survivors"),
    [
        ("slide_reorder", 2),
        ("period_roll", 2),
        ("restatement", 2),
        ("duplicate_equal_values", 2),
        ("sheet_range_movement", 2),
        # Rewording a mapped line changes its skeleton, so the mapping does not
        # survive. That is recorded, not hidden: the claim reappears as unmapped.
        ("wording_change", 1),
    ],
)
def test_mapping_durability_is_measured_per_cycle(
    cycle: str, expected_survivors: int, tmp_path: Path
) -> None:
    path = _deck(tmp_path / f"{cycle}.pptx", _CYCLES[cycle])

    assert len(_survivors(path)) == expected_survivors


def test_a_lost_mapping_becomes_unmapped_rather_than_disappearing(
    tmp_path: Path,
) -> None:
    """Narrowing must be visible: a claim that lost its mapping is still counted."""
    deck = load_deck_snapshot(_deck(tmp_path / "reworded.pptx", _CYCLES["wording_change"]))
    workbook = load_workbook_snapshot(_workbook(tmp_path / "book.xlsx", _BASE_CELLS))

    coverage = reconcile_package(workbook, deck, _profile()).mapping_coverage

    assert coverage.reconciles
    assert coverage.mapped == 1
    assert coverage.unmapped == coverage.eligible - 1
