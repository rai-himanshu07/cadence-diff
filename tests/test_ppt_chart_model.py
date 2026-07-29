"""Complete PowerPoint chart and shape extraction contracts."""

from __future__ import annotations

import copy
import zipfile
from pathlib import Path
from typing import Any, cast
from xml.etree import ElementTree

import pytest
from pptx import Presentation
from pptx.chart.data import ChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.util import Inches

import qc_tool.ppt.extract as ppt_extract
from qc_tool.config.profile import PptProfile
from qc_tool.coverage import CoverageState
from qc_tool.ppt.chart_xml import PptChartParseError, parse_ppt_chart
from qc_tool.ppt.extract import load_deck_snapshot
from qc_tool.ppt.preflight import preflight_deck

_CHART_NS = "http://schemas.openxmlformats.org/drawingml/2006/chart"


def _tag(name: str) -> str:
    return f"{{{_CHART_NS}}}{name}"


def _rewrite_chart_as_combo(path: Path) -> None:
    with zipfile.ZipFile(path) as source:
        members = [(item, source.read(item)) for item in source.infolist()]
    with zipfile.ZipFile(path, "w") as output:
        for item, content in members:
            if item.filename == "ppt/charts/chart1.xml":
                root = ElementTree.fromstring(content)
                plot_area = root.find(f".//{_tag('plotArea')}")
                assert plot_area is not None
                primary = plot_area.find(_tag("lineChart"))
                assert primary is not None
                secondary = copy.deepcopy(primary)
                secondary.tag = _tag("barChart")
                first = next(iter(secondary))
                secondary.insert(0, ElementTree.Element(_tag("barDir"), {"val": "col"}))
                if first.tag == _tag("grouping"):
                    first.set("val", "clustered")
                series = secondary.find(_tag("ser"))
                assert series is not None
                index = series.find(_tag("idx"))
                order = series.find(_tag("order"))
                assert index is not None and order is not None
                index.set("val", "1")
                order.set("val", "1")
                secondary_axis_ids = secondary.findall(_tag("axId"))
                assert len(secondary_axis_ids) == 2
                secondary_axis_ids[0].set("val", "300")
                secondary_axis_ids[1].set("val", "400")
                name_value = series.find(f".//{_tag('tx')}//{_tag('v')}")
                assert name_value is not None
                name_value.text = "Margin"
                for point in series.findall(f".//{_tag('val')}//{_tag('pt')}"):
                    value = point.find(_tag("v"))
                    assert value is not None and value.text is not None
                    value.text = str(float(value.text) / 100.0)
                primary_index = list(plot_area).index(primary)
                plot_area.insert(primary_index + 1, secondary)
                category_axis = plot_area.find(_tag("catAx"))
                value_axis = plot_area.find(_tag("valAx"))
                assert category_axis is not None and value_axis is not None
                secondary_category_axis = copy.deepcopy(category_axis)
                secondary_value_axis = copy.deepcopy(value_axis)
                category_axis_id = secondary_category_axis.find(_tag("axId"))
                category_cross_axis = secondary_category_axis.find(_tag("crossAx"))
                value_axis_id = secondary_value_axis.find(_tag("axId"))
                value_cross_axis = secondary_value_axis.find(_tag("crossAx"))
                value_position = secondary_value_axis.find(_tag("axPos"))
                assert all(
                    item is not None
                    for item in (
                        category_axis_id,
                        category_cross_axis,
                        value_axis_id,
                        value_cross_axis,
                        value_position,
                    )
                )
                assert category_axis_id is not None
                assert category_cross_axis is not None
                assert value_axis_id is not None
                assert value_cross_axis is not None
                assert value_position is not None
                category_axis_id.set("val", "300")
                category_cross_axis.set("val", "400")
                value_axis_id.set("val", "400")
                value_cross_axis.set("val", "300")
                value_position.set("val", "r")
                plot_area.append(secondary_category_axis)
                plot_area.append(secondary_value_axis)
                content = ElementTree.tostring(
                    root,
                    encoding="utf-8",
                    xml_declaration=True,
                )
            output.writestr(item, content)


def _build_combo_deck(path: Path) -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    title = slide.shapes.title
    assert title is not None
    title.text = "Operations"
    text = slide.shapes.add_textbox(Inches(0.5), Inches(1.1), Inches(3), Inches(0.5))
    text.text = "Visible annotation"
    table_shape = slide.shapes.add_table(
        2,
        2,
        Inches(0.5),
        Inches(2),
        Inches(3),
        Inches(1.2),
    )
    table_shape.table.cell(0, 0).text = "Metric"
    table_shape.table.cell(0, 1).text = "Jan"
    table_shape.table.cell(1, 0).text = "Revenue"
    table_shape.table.cell(1, 1).text = "10"

    data = ChartData()
    data.categories = ["Jan", "Feb", "Mar"]
    data.add_series("Revenue", [10, 20, 30])
    chart_shape = cast(
        Any,
        slide.shapes.add_chart(
            XL_CHART_TYPE.LINE,
            Inches(4),
            Inches(1.2),
            Inches(5),
            Inches(4),
            data,
        ),
    )
    chart = chart_shape.chart
    chart.has_legend = True
    assert chart.legend is not None
    chart.legend.position = XL_LEGEND_POSITION.BOTTOM
    plot = chart.plots[0]
    plot.has_data_labels = True
    plot.data_labels.show_value = True
    chart.value_axis.maximum_scale = 40
    chart.value_axis.major_unit = 10

    notes_frame = slide.notes_slide.notes_text_frame
    assert notes_frame is not None
    notes_frame.text = "Internal note only"
    presentation.save(str(path))
    _rewrite_chart_as_combo(path)


def test_ppt_extracts_complete_combo_shapes_notes_and_geometry(tmp_path: Path) -> None:
    path = tmp_path / "combo.pptx"
    _build_combo_deck(path)

    deck = load_deck_snapshot(path)

    assert deck.charts_available
    assert deck.notes_available
    assert len(deck.slides) == 1
    slide = deck.slides[0]
    assert slide.notes == ["Internal note only"]
    assert "Internal note only" not in slide.texts
    assert "Visible annotation" in slide.texts
    assert [shape.z_order for shape in slide.shapes] == list(range(slide.shape_count))
    assert len({shape.source_id for shape in slide.shapes}) == slide.shape_count

    assert len(slide.tables) == 1
    table = slide.tables[0]
    assert table.rows == [["Metric", "Jan"], ["Revenue", "10"]]
    assert table.source_id.startswith("slide[0]/shape[")
    assert table.width > 0 and table.height > 0

    assert len(slide.charts) == 1
    chart = slide.charts[0]
    assert chart.source_id.startswith("slide[0]/shape[")
    assert chart.source_index == 0
    assert chart.z_order > table.z_order
    assert chart.width > 0 and chart.height > 0
    assert [plot.chart_type for plot in chart.plots] == ["lineChart", "barChart"]
    assert [len(plot.series) for plot in chart.plots] == [1, 1]
    assert [series.name for series in chart.all_series] == ["Revenue", "Margin"]
    assert all(series.categories == ["Jan", "Feb", "Mar"] for series in chart.all_series)
    assert chart.all_series[0].values == [10.0, 20.0, 30.0]
    assert chart.all_series[1].values == [0.1, 0.2, 0.3]
    assert [plot.axis_group for plot in chart.plots] == ["primary", "secondary-1"]
    assert chart.plots[0].axis_ids != chart.plots[1].axis_ids
    assert len(chart.axes) == 4
    value_axis = next(
        axis
        for axis in chart.axes
        if axis.axis_type == "valAx" and axis.position == "l"
    )
    assert value_axis.maximum == 40.0
    assert value_axis.major_unit == 10.0
    assert chart.legend is not None and chart.legend.position == "b"
    assert chart.plots[0].data_labels is not None
    assert chart.plots[0].data_labels.show_value
    assert chart.plots[0].visible_label_values == ["10.0", "20.0", "30.0"]


def test_duplicate_named_chart_panels_keep_distinct_source_identities(
    tmp_path: Path,
) -> None:
    path = tmp_path / "panels.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    data = ChartData()
    data.categories = ["Jan", "Feb"]
    data.add_series("Revenue", [10, 20])
    for left in (0.5, 5.0):
        shape = cast(
            Any,
            slide.shapes.add_chart(
                XL_CHART_TYPE.LINE,
                Inches(left),
                Inches(1.5),
                Inches(4),
                Inches(3),
                data,
            ),
        )
        shape.name = "Duplicate panel"
    presentation.save(str(path))

    deck = load_deck_snapshot(path)
    charts = deck.slides[0].charts

    assert len(charts) == 2
    assert [chart.source_index for chart in charts] == [0, 1]
    assert len({chart.source_id for chart in charts}) == 2
    assert len({(chart.left, chart.top) for chart in charts}) == 2


def test_chart_parse_failure_degrades_without_losing_slide_content(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    path = tmp_path / "degraded.pptx"
    _build_combo_deck(path)

    def fail_chart(*args: object, **kwargs: object) -> object:
        raise PptChartParseError("unsupported chart construct")

    monkeypatch.setattr(ppt_extract, "parse_ppt_chart", fail_chart)

    deck = load_deck_snapshot(path)
    preflight = preflight_deck(deck, PptProfile())
    chart_coverage = next(
        item for item in preflight.coverage if item.check_id == "ppt-tables-charts"
    )

    assert not deck.charts_available
    assert "unsupported chart construct" in deck.chart_detail
    assert deck.slides[0].charts == []
    assert "Visible annotation" in deck.slides[0].texts
    assert deck.slides[0].tables[0].rows[1] == ["Revenue", "10"]
    assert chart_coverage.state is CoverageState.DEGRADED


def test_notes_failure_degrades_without_mixing_notes_into_visible_text(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    path = tmp_path / "notes-degraded.pptx"
    _build_combo_deck(path)

    def fail_notes(slide: object) -> list[str]:
        raise ValueError("notes XML unavailable")

    monkeypatch.setattr(ppt_extract, "_notes", fail_notes)

    deck = load_deck_snapshot(path)
    preflight = preflight_deck(deck, PptProfile())
    notes_coverage = next(
        item for item in preflight.coverage if item.check_id == "ppt-notes"
    )

    assert not deck.notes_available
    assert "notes XML unavailable" in deck.notes_detail
    assert deck.slides[0].notes == []
    assert "Visible annotation" in deck.slides[0].texts
    assert notes_coverage.state is CoverageState.DEGRADED


def test_sparse_chart_caches_and_default_label_separator() -> None:
        xml = f"""
        <c:chartSpace xmlns:c="{_CHART_NS}">
            <c:chart><c:plotArea><c:lineChart>
                <c:ser><c:idx val="0"/><c:order val="0"/>
                    <c:tx><c:v>Revenue</c:v></c:tx>
                    <c:cat><c:strLit><c:ptCount val="5"/>
                        <c:pt idx="0"><c:v>Jan</c:v></c:pt>
                        <c:pt idx="4"><c:v>May</c:v></c:pt>
                    </c:strLit></c:cat>
                    <c:val><c:numLit><c:ptCount val="5"/>
                        <c:pt idx="0"><c:v>10</c:v></c:pt>
                        <c:pt idx="4"><c:v>50</c:v></c:pt>
                    </c:numLit></c:val>
                </c:ser>
                <c:dLbls><c:showSerName val="1"/><c:showCatName val="1"/>
                    <c:showVal val="1"/></c:dLbls>
            </c:lineChart></c:plotArea></c:chart>
        </c:chartSpace>
        """.encode()

        chart = parse_ppt_chart(
                xml,
                source_id="slide[0]/shape[0]",
                source_index=0,
                shape_id=1,
                name="Chart",
                left=0,
                top=0,
                width=100,
                height=100,
                z_order=0,
                source_part="/ppt/charts/chart1.xml",
        )

        series = chart.all_series[0]
        assert series.categories == ["Jan", "", "", "", "May"]
        assert series.values == [10.0, None, None, None, 50.0]
        assert chart.plots[0].visible_label_values == [
                "Revenue\nJan\n10.0",
                "Revenue\nMay\n50.0",
        ]


def test_chart_without_supported_plot_fails_closed() -> None:
        xml = f"""
        <c:chartSpace xmlns:c="{_CHART_NS}">
            <c:chart><c:plotArea><c:extLst/></c:plotArea></c:chart>
        </c:chartSpace>
        """.encode()

        with pytest.raises(PptChartParseError, match="no supported plot"):
                parse_ppt_chart(
                        xml,
                        source_id="slide[0]/shape[0]",
                        source_index=0,
                        shape_id=1,
                        name="Chart",
                        left=0,
                        top=0,
                        width=100,
                        height=100,
                        z_order=0,
                        source_part="/ppt/charts/chart1.xml",
                )


def test_chart_with_uncaptured_axis_reference_fails_closed() -> None:
        xml = f"""
        <c:chartSpace xmlns:c="{_CHART_NS}">
            <c:chart><c:plotArea><c:lineChart>
                <c:ser><c:idx val="0"/><c:order val="0"/></c:ser>
                <c:axId val="10"/><c:axId val="20"/>
            </c:lineChart></c:plotArea></c:chart>
        </c:chartSpace>
        """.encode()

        with pytest.raises(PptChartParseError, match="uncaptured axis IDs"):
                parse_ppt_chart(
                        xml,
                        source_id="slide[0]/shape[0]",
                        source_index=0,
                        shape_id=1,
                        name="Chart",
                        left=0,
                        top=0,
                        width=100,
                        height=100,
                        z_order=0,
                        source_part="/ppt/charts/chart1.xml",
                )
