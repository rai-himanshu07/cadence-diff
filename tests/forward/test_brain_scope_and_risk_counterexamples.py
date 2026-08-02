"""Forward-failing scope and intrinsic-risk contracts for Brain Step 1."""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook
from pptx import Presentation

from qc_tool import engine as engine_module
from qc_tool.coverage import QCRunMode
from qc_tool.engine import run_qc
from qc_tool.findings import FindingClass, Severity, limit_findings
from qc_tool.io.model import (
    CellRecord,
    SheetSnapshot,
    WorkbookRisk,
    WorkbookRiskKind,
    WorkbookSnapshot,
)


def _write_scope_workbooks(tmp_path: Path) -> tuple[Path, Path]:
    def build(path: Path, *, current: bool) -> None:
        workbook = Workbook()
        flood = workbook.active
        assert flood is not None
        flood.title = "Flood"
        flood.append(["Value"])
        for row in range(2, 10):
            flood.cell(row, 1, "#N/A" if current else float(row))
        selected = workbook.create_sheet("Selected")
        selected.append(["Value"])
        selected.cell(2, 1, "#REF!" if current else 1.0)
        workbook.save(path)

    baseline = tmp_path / "baseline.xlsx"
    current = tmp_path / "current.xlsx"
    build(baseline, current=False)
    build(current, current=True)
    return baseline, current


def _write_blank_deck(path: Path) -> None:
    presentation = Presentation()
    presentation.slides.add_slide(presentation.slide_layouts[6])
    presentation.save(str(path))


@pytest.mark.parametrize(
    "mode",
    [
        QCRunMode.CURRENT_FILE_PREFLIGHT,
        QCRunMode.CYCLE_COMPARISON,
        QCRunMode.FINAL_PACKAGE,
    ],
)
def test_scope_filters_before_global_budget_in_every_mode(
    mode: QCRunMode,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    baseline, current = _write_scope_workbooks(tmp_path)
    monkeypatch.setattr(
        engine_module,
        "limit_findings",
        lambda findings: limit_findings(
            findings,
            max_per_class_scope=500,
            max_total=2,
        ),
    )
    kwargs: dict[str, object] = {
        "current_excel": current,
        "compare_sheets": ["Selected"],
        "mode": mode,
    }
    if mode is QCRunMode.CYCLE_COMPARISON:
        kwargs["baseline_excel"] = baseline
    elif mode is QCRunMode.FINAL_PACKAGE:
        deck = tmp_path / "current.pptx"
        _write_blank_deck(deck)
        kwargs["current_ppt"] = deck

    result = run_qc(**kwargs)  # type: ignore[arg-type]

    assert any(
        finding.finding_class is FindingClass.FORMULA_ERROR
        and finding.sheet == "Selected"
        for finding in result.findings
    )


def test_unknown_sheet_scope_fails_actionably(tmp_path: Path) -> None:
    _, current = _write_scope_workbooks(tmp_path)

    with pytest.raises(ValueError, match=r"(?i)unknown.*Missing"):
        run_qc(
            current_excel=current,
            compare_sheets=["Missing"],
            mode=QCRunMode.CURRENT_FILE_PREFLIGHT,
        )


def test_unknown_slide_scope_fails_actionably(
    fixture_dir: Path,
) -> None:
    with pytest.raises(ValueError, match=r"(?i)unknown.*999"):
        run_qc(
            baseline_ppt=fixture_dir / "baseline.pptx",
            current_ppt=fixture_dir / "current.pptx",
            compare_slides=[999],
            mode=QCRunMode.CYCLE_COMPARISON,
        )


def _xlsb_snapshot(source_name: str, *, external: bool) -> WorkbookSnapshot:
    return WorkbookSnapshot(
        source_name=source_name,
        file_format="xlsb",
        formulas_available=False,
        styles_available=False,
        formula_presence_available=True,
        tables_available=False,
        charts_available=False,
        interaction_rules_available=False,
        sheets=[
            SheetSnapshot(
                "Data",
                "visible",
                1,
                1,
                {(1, 1): CellRecord(1, 1, 1.0)},
            )
        ],
        intrinsic_risks=(
            [WorkbookRisk(WorkbookRiskKind.EXTERNAL_WORKBOOK_LINK)]
            if external
            else []
        ),
    )


def test_scanner_proven_xlsb_external_is_critical_in_cycle_mode(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    baseline = _xlsb_snapshot("baseline.xlsb", external=False)
    current = _xlsb_snapshot("current.xlsb", external=True)

    def fake_load(path: Path, **_: object) -> WorkbookSnapshot:
        return baseline if path.name == "baseline.xlsb" else current

    monkeypatch.setattr(engine_module, "_load_excel_file", fake_load)
    result = run_qc(
        baseline_excel=Path("baseline.xlsb"),
        current_excel=Path("current.xlsb"),
        mode=QCRunMode.CYCLE_COMPARISON,
    )
    external = [
        finding
        for finding in result.findings
        if finding.finding_class is FindingClass.EXTERNAL_LINK
    ]

    assert external
    assert {finding.severity for finding in external} == {Severity.CRITICAL}
