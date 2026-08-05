"""Cell comments, Power Query definitions, and workbook connections."""

from __future__ import annotations

import base64
import io
import zipfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment

from qc_tool.config.profile import DeliverableProfile
from qc_tool.coverage import CoverageState
from qc_tool.excel.diff_metadata import (
    comment_coverage,
    connection_coverage,
    diff_workbook_metadata,
    external_connection_findings,
    power_query_coverage,
)
from qc_tool.excel.preflight import preflight_workbook
from qc_tool.findings import FindingClass, Severity
from qc_tool.io.loader import load_workbook_snapshot
from qc_tool.io.model import WorkbookSnapshot
from qc_tool.io.ooxml_metadata import WorkbookMetadataScan, scan_workbook_metadata
from qc_tool.triage.rules import DEFAULT_SEVERITIES

_SECRET_TARGET = "Provider=SQLOLEDB;Server=finance-prod-01;Uid=svc_report;Pwd=hunter2"
_SECRET_URL = "https://internal.example.invalid/secret-feed.csv"


def _workbook_with_comments(path: Path, comments: dict[str, tuple[str, str]]) -> Path:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Data"
    for row in range(1, 4):
        sheet.cell(row=row, column=1, value=row)
    for ref, (text, author) in comments.items():
        sheet[ref].comment = Comment(text, author)
    workbook.save(path)
    return path


def _package(parts: dict[str, bytes | str]) -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        for name, payload in parts.items():
            archive.writestr(name, payload)
    return buffer.getvalue()


def _connections_package(body: str) -> bytes:
    return _package(
        {
            "xl/workbook.xml": "<workbook><sheets/></workbook>",
            "xl/connections.xml": body,
        }
    )


def _mashup(section: str) -> bytes:
    inner = io.BytesIO()
    with zipfile.ZipFile(inner, "w") as package:
        package.writestr("Formulas/Section1.m", section)
    parts = inner.getvalue()
    blob = (1).to_bytes(4, "little") + len(parts).to_bytes(4, "little") + parts
    encoded = base64.b64encode(blob).decode()
    return _package(
        {
            "xl/workbook.xml": "<workbook><sheets/></workbook>",
            "customXml/item1.xml": (
                '<DataMashup xmlns="http://schemas.microsoft.com/DataMashup">'
                f"{encoded}</DataMashup>"
            ),
        }
    )


def _snapshot(scan: WorkbookMetadataScan) -> WorkbookSnapshot:
    return WorkbookSnapshot(
        source_name="book.xlsx",
        file_format="xlsx",
        formulas_available=True,
        styles_available=True,
        metadata=scan,
    )


# --- comments ------------------------------------------------------------


def test_comments_are_read_with_sheet_and_author(tmp_path: Path) -> None:
    path = _workbook_with_comments(
        tmp_path / "book.xlsx", {"A2": ("check this number", "Reviewer")}
    )

    snapshot = load_workbook_snapshot(path)

    assert snapshot.metadata.comments_available is True
    comment = snapshot.metadata.comments[0]
    assert comment.sheet == "Data"
    assert comment.ref == "A2"
    assert "check this number" in comment.text
    assert comment.author == "Reviewer"


def test_comment_added_removed_and_changed_are_reported(tmp_path: Path) -> None:
    baseline = load_workbook_snapshot(
        _workbook_with_comments(
            tmp_path / "base.xlsx",
            {"A1": ("stays", "R"), "A2": ("original", "R"), "A3": ("goes", "R")},
        )
    )
    current = load_workbook_snapshot(
        _workbook_with_comments(
            tmp_path / "curr.xlsx",
            {"A1": ("stays", "R"), "A2": ("rewritten", "R"), "B1": ("brand new", "R")},
        )
    )

    findings = [
        finding
        for finding in diff_workbook_metadata(baseline, current)
        if finding.finding_class is FindingClass.CELL_COMMENT_CHANGED
    ]

    assert sorted(finding.element or "" for finding in findings) == [
        "Data!A2",
        "Data!A3",
        "Data!B1",
    ]
    messages = {finding.element: finding.message for finding in findings}
    assert "added" in messages["Data!B1"]
    assert "removed" in messages["Data!A3"]
    assert "changed" in messages["Data!A2"]


def test_comment_finding_carries_a_focusable_cell(tmp_path: Path) -> None:
    baseline = load_workbook_snapshot(_workbook_with_comments(tmp_path / "b.xlsx", {}))
    current = load_workbook_snapshot(
        _workbook_with_comments(tmp_path / "c.xlsx", {"C7": ("note", "R")})
    )

    finding = diff_workbook_metadata(baseline, current)[0]

    assert finding.sheet == "Data"
    assert finding.location == "C7"


def test_unchanged_comment_produces_no_finding(tmp_path: Path) -> None:
    baseline = load_workbook_snapshot(
        _workbook_with_comments(tmp_path / "b.xlsx", {"A1": ("same", "R")})
    )
    current = load_workbook_snapshot(
        _workbook_with_comments(tmp_path / "c.xlsx", {"A1": ("same", "R")})
    )

    assert diff_workbook_metadata(baseline, current) == []


def test_comment_changes_are_informational() -> None:
    assert DEFAULT_SEVERITIES[FindingClass.CELL_COMMENT_CHANGED] is Severity.INFO


# --- power query ---------------------------------------------------------


def test_named_queries_are_extracted_from_the_mashup() -> None:
    section = (
        "section Section1;\r\n\r\n"
        'shared Sales = let Source = Excel.CurrentWorkbook() in Source;\r\n\r\n'
        'shared "Margin Detail" = let Source = Sales in Source;\r\n'
    )

    scan = scan_workbook_metadata(_mashup(section))

    assert scan.queries_available is True
    assert [query.name for query in scan.queries] == ["Margin Detail", "Sales"]
    assert scan.queries_detail == ""


def test_changed_query_definition_is_reported() -> None:
    before = 'section Section1;\r\nshared Sales = let x = 1 in x;\r\n'
    after = 'section Section1;\r\nshared Sales = let x = 2 in x;\r\n'

    findings = diff_workbook_metadata(
        _snapshot(scan_workbook_metadata(_mashup(before))),
        _snapshot(scan_workbook_metadata(_mashup(after))),
    )

    assert [finding.finding_class for finding in findings] == [
        FindingClass.POWER_QUERY_CHANGED
    ]
    assert findings[0].element == "Sales"
    assert "definition changed" in findings[0].message


def test_identical_query_definition_produces_no_finding() -> None:
    section = 'section Section1;\r\nshared Sales = let x = 1 in x;\r\n'
    package = _mashup(section)

    assert (
        diff_workbook_metadata(
            _snapshot(scan_workbook_metadata(package)),
            _snapshot(scan_workbook_metadata(package)),
        )
        == []
    )


def test_undecodable_mashup_degrades_to_a_definition_digest() -> None:
    package = _package(
        {
            "xl/workbook.xml": "<workbook><sheets/></workbook>",
            "customXml/item1.xml": (
                '<DataMashup xmlns="http://schemas.microsoft.com/DataMashup">'
                f"{base64.b64encode(b'not-a-mashup-envelope').decode()}</DataMashup>"
            ),
        }
    )

    scan = scan_workbook_metadata(package)

    assert scan.queries_available is True
    assert [query.name for query in scan.queries] == ["DataMashup1"]
    assert "falls back to a definition digest" in scan.queries_detail


def test_workbook_without_queries_reports_none() -> None:
    scan = scan_workbook_metadata(_package({"xl/workbook.xml": "<workbook/>"}))

    assert scan.queries == ()
    assert scan.queries_available is True


# --- connections ---------------------------------------------------------


def test_connection_target_is_never_retained() -> None:
    scan = scan_workbook_metadata(
        _connections_package(
            "<connections>"
            f'<connection id="1" name="Sales" type="5">'
            f'<dbPr connection="{_SECRET_TARGET}" command="SELECT * FROM Payroll"/>'
            "</connection></connections>"
        )
    )

    connection = scan.connections[0]
    rendered = repr(scan)
    assert connection.name == "Sales"
    assert connection.kind == "oledb"
    assert "hunter2" not in rendered
    assert "finance-prod-01" not in rendered
    assert "Payroll" not in rendered
    assert len(connection.target_digest) == 16


def test_web_connection_is_classified_as_external() -> None:
    scan = scan_workbook_metadata(
        _connections_package(
            "<connections>"
            f'<connection id="1" name="Feed" type="4"><webPr url="{_SECRET_URL}"/>'
            "</connection></connections>"
        )
    )

    assert scan.connections[0].kind == "web"
    assert scan.connections[0].external is True


def test_power_query_connection_is_classified() -> None:
    scan = scan_workbook_metadata(
        _connections_package(
            "<connections>"
            '<connection id="1" name="Query - Sales" type="5">'
            '<dbPr connection="Provider=Microsoft.Mashup.OleDb.1;Location=Sales"/>'
            "</connection></connections>"
        )
    )

    assert scan.connections[0].kind == "power_query"


def test_retargeted_connection_is_reported_without_the_target() -> None:
    before = _connections_package(
        "<connections>"
        '<connection id="1" name="Sales" type="5">'
        '<dbPr connection="Provider=X;Server=old-server"/>'
        "</connection></connections>"
    )
    after = _connections_package(
        "<connections>"
        '<connection id="1" name="Sales" type="5">'
        f'<dbPr connection="{_SECRET_TARGET}"/>'
        "</connection></connections>"
    )

    findings = diff_workbook_metadata(
        _snapshot(scan_workbook_metadata(before)),
        _snapshot(scan_workbook_metadata(after)),
    )

    assert len(findings) == 1
    assert findings[0].finding_class is FindingClass.CONNECTION_CHANGED
    assert "points at a different target" in findings[0].message
    rendered = " ".join(
        str(value)
        for value in (
            findings[0].message,
            findings[0].baseline_value,
            findings[0].current_value,
        )
    )
    assert "finance-prod-01" not in rendered
    assert "hunter2" not in rendered
    assert "old-server" not in rendered


def test_external_connection_presence_is_separate_from_comparison() -> None:
    snapshot = _snapshot(
        scan_workbook_metadata(
            _connections_package(
                "<connections>"
                f'<connection id="1" name="Feed" type="4"><webPr url="{_SECRET_URL}"/>'
                "</connection>"
                '<connection id="2" name="Local" type="6"><textPr/></connection>'
                "</connections>"
            )
        )
    )

    findings = external_connection_findings(snapshot)

    assert [finding.element for finding in findings] == ["Feed"]
    assert findings[0].finding_class is FindingClass.EXTERNAL_CONNECTION
    assert "secret-feed" not in findings[0].message


def test_unreadable_connection_metadata_is_unavailable_not_absent() -> None:
    scan = scan_workbook_metadata(_connections_package("<connections><broken>"))

    assert scan.connections_available is False
    assert scan.connections == ()
    assert "not readable XML" in scan.connections_detail


# --- coverage ------------------------------------------------------------


def test_coverage_items_report_checked_for_a_plain_workbook(tmp_path: Path) -> None:
    snapshot = load_workbook_snapshot(_workbook_with_comments(tmp_path / "b.xlsx", {}))

    for item in (
        comment_coverage(snapshot),
        power_query_coverage(snapshot),
        connection_coverage(snapshot),
    ):
        assert item.state is CoverageState.CHECKED


def test_connection_coverage_is_unavailable_when_metadata_is_broken() -> None:
    snapshot = _snapshot(scan_workbook_metadata(_connections_package("<connections><x>")))

    item = connection_coverage(snapshot)

    assert item.check_id == "excel-connections"
    assert item.state is CoverageState.UNAVAILABLE


def test_power_query_coverage_is_degraded_when_text_cannot_be_opened() -> None:
    snapshot = _snapshot(
        scan_workbook_metadata(
            _package(
                {
                    "xl/workbook.xml": "<workbook><sheets/></workbook>",
                    "customXml/item1.xml": (
                        '<DataMashup xmlns="http://schemas.microsoft.com/DataMashup">'
                        f"{base64.b64encode(b'opaque').decode()}</DataMashup>"
                    ),
                }
            )
        )
    )

    assert power_query_coverage(snapshot).state is CoverageState.DEGRADED


def test_preflight_registers_all_three_coverage_items(tmp_path: Path) -> None:
    snapshot = load_workbook_snapshot(
        _workbook_with_comments(tmp_path / "b.xlsx", {"A1": ("note", "R")})
    )

    result = preflight_workbook(snapshot, DeliverableProfile(name="metadata-tests"))

    registered = {item.check_id for item in result.coverage}
    assert {"excel-comments", "excel-power-query", "excel-connections"} <= registered


# --- compatibility -------------------------------------------------------


def test_legacy_snapshot_without_metadata_still_compares(tmp_path: Path) -> None:
    """A snapshot built before this slice must not break the new comparison."""
    legacy = WorkbookSnapshot(
        source_name="legacy.xlsx",
        file_format="xlsx",
        formulas_available=True,
        styles_available=True,
    )

    assert diff_workbook_metadata(legacy, legacy) == []
    assert external_connection_findings(legacy) == []
    assert comment_coverage(legacy).state is CoverageState.UNAVAILABLE
