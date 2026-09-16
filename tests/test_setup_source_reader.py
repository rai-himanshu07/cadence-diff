"""Synthetic contracts for setup-only workbook readers."""

from __future__ import annotations

import ast
import datetime as dt
from pathlib import Path

import pytest
from openpyxl import Workbook

from qc_tool.io.decrypt import InvalidPasswordError, PasswordRequiredError
from qc_tool.setup.source_reader import open_setup_source
from tests.fixtures.generate import encrypt_file
from tests.fixtures.xlsb_writer import StyledCell, write_xlsb


def _write_ooxml(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Data"
    sheet.sheet_state = "hidden"
    sheet.append(["ID", "When", "Formula"])
    sheet.append(["A", dt.date(2026, 9, 15), "=1+1"])
    workbook.create_sheet("Visible")
    workbook.active = 1
    workbook.save(path)


def test_ooxml_reader_preserves_typed_cached_values_formula_presence_and_visibility(
    tmp_path: Path,
) -> None:
    path = tmp_path / "source.xlsx"
    _write_ooxml(path)

    with open_setup_source(path) as source:
        sheets = list(source.iter_sheets())

    assert source.file_format == "xlsx"
    assert source.xlsb_risk is None
    assert [sheet.name for sheet in sheets] == ["Data", "Visible"]
    data = sheets[0]
    assert data.visibility == "hidden"
    assert data.cells[(2, 1)].value == "A"
    assert data.cells[(2, 2)].value == dt.datetime(2026, 9, 15)
    assert data.cells[(2, 3)].has_formula
    assert data.cells[(2, 3)].formula == "=1+1"


def test_ooxml_reader_requires_and_accepts_an_ephemeral_password(tmp_path: Path) -> None:
    plain = tmp_path / "plain.xlsx"
    encrypted = tmp_path / "encrypted.xlsx"
    _write_ooxml(plain)
    encrypt_file(plain, encrypted, "hunter2")

    with pytest.raises(PasswordRequiredError), open_setup_source(encrypted):
        pass
    with pytest.raises(InvalidPasswordError), open_setup_source(
        encrypted, password="wrong"
    ):
        pass
    with open_setup_source(encrypted, password="hunter2") as source:
        assert len(list(source.iter_sheets())) == 2


def test_xlsb_reader_uses_cached_values_and_formula_presence_without_formula_text(
    tmp_path: Path,
) -> None:
    path = tmp_path / "source.xlsb"
    write_xlsb(
        path,
        {
            "Data": [
                ["ID", "Value"],
                ["A", StyledCell(2, is_formula=True)],
            ]
        },
    )

    with open_setup_source(path) as source:
        [sheet] = list(source.iter_sheets())

    assert source.file_format == "xlsb"
    assert source.xlsb_risk is not None
    assert sheet.cells[(2, 1)].value == "A"
    assert sheet.cells[(2, 2)].value == 2
    assert sheet.cells[(2, 2)].has_formula
    assert sheet.cells[(2, 2)].formula is None


def test_xlsb_reader_preserves_sheet_visibility(tmp_path: Path) -> None:
    path = tmp_path / "source.xlsb"
    write_xlsb(
        path,
        {"Visible": [[1]], "Hidden": [[2]], "VeryHidden": [[3]]},
        sheet_visibility={
            "Hidden": "hidden",
            "VeryHidden": "veryHidden",
        },
    )

    with open_setup_source(path) as source:
        sheets = list(source.iter_sheets())

    assert {sheet.name: sheet.visibility for sheet in sheets} == {
        "Visible": "visible",
        "Hidden": "hidden",
        "VeryHidden": "veryHidden",
    }


def test_xlsb_reader_prefers_native_cached_values_when_available(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    import qc_tool.io.native_kernel as native_kernel_module

    path = tmp_path / "source.xlsb"
    write_xlsb(path, {"Data": [["stored-value"]]})
    monkeypatch.setattr(native_kernel_module, "native_kernel_available", lambda: True)
    monkeypatch.setattr(
        native_kernel_module,
        "raw_values_report",
        lambda _data: [("Data", [(0, 0, None, None, "native-value")])],
    )

    with open_setup_source(path) as source:
        [sheet] = list(source.iter_sheets())

    assert sheet.cells[(1, 1)].value == "native-value"


def test_xlsb_reader_falls_back_when_native_sheet_inventory_differs(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    import qc_tool.io.native_kernel as native_kernel_module

    path = tmp_path / "source.xlsb"
    write_xlsb(path, {"Data": [["fallback-value"]]})
    monkeypatch.setattr(native_kernel_module, "native_kernel_available", lambda: True)
    monkeypatch.setattr(
        native_kernel_module,
        "raw_values_report",
        lambda _data: [("Other", [(0, 0, None, None, "wrong-sheet")])],
    )

    with open_setup_source(path) as source:
        [sheet] = list(source.iter_sheets())

    assert sheet.cells[(1, 1)].value == "fallback-value"


def test_xlsb_reader_falls_back_when_native_cached_values_fail(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    import qc_tool.io.native_kernel as native_kernel_module

    path = tmp_path / "source.xlsb"
    write_xlsb(path, {"Data": [["fallback-value"]]})
    monkeypatch.setattr(native_kernel_module, "native_kernel_available", lambda: True)

    def fail_native(_data: bytes) -> list[object]:
        raise RuntimeError("synthetic native failure")

    monkeypatch.setattr(native_kernel_module, "raw_values_report", fail_native)

    with open_setup_source(path) as source:
        [sheet] = list(source.iter_sheets())

    assert sheet.cells[(1, 1)].value == "fallback-value"


def test_setup_modules_do_not_reference_full_loader_or_office_adapters() -> None:
    root = Path(__file__).parents[1] / "qc_tool" / "setup"
    forbidden_names = {
        "load_workbook_snapshot",
        "extract_formulas_with_excel",
        "extract_formulas_with_libreoffice",
    }

    for path in root.glob("*.py"):
        tree = ast.parse(path.read_text(encoding="utf-8"))
        referenced = {
            node.id for node in ast.walk(tree) if isinstance(node, ast.Name)
        } | {
            node.attr for node in ast.walk(tree) if isinstance(node, ast.Attribute)
        }
        assert not forbidden_names & referenced, path.name
