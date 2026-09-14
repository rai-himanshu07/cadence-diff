"""Alignment engine tests against the fixture pair's known ground truth."""

from pathlib import Path

import pytest
from openpyxl import Workbook

from qc_tool.config.execution import ExecutionBindings
from qc_tool.config.profile import DeliverableProfile
from qc_tool.config.resolved_input import (
    ResolvedColumn,
    ResolvedInputConfigurationV1,
    ResolvedMember,
    ResolvedRegion,
    ResolvedSheet,
)
from qc_tool.coverage import CoverageState
from qc_tool.engine import run_qc
from qc_tool.excel.align import (
    AlignmentTrustManifest,
    AxisAlignment,
    AxisEntry,
    RegionAlignment,
    WorkbookAlignment,
    _align_axis,
    _pair_regions,
    align_workbooks,
    build_alignment_trust_manifest,
)
from qc_tool.excel.periods import parse_period
from qc_tool.excel.regions import TableRegion, detect_regions
from qc_tool.findings import FindingClass
from qc_tool.io.loader import load_workbook_snapshot
from qc_tool.io.model import CellRecord, SheetSnapshot, WorkbookSnapshot


@pytest.fixture(scope="module")
def baseline(fixture_dir: Path) -> WorkbookSnapshot:
    return load_workbook_snapshot(fixture_dir / "baseline.xlsx")


@pytest.fixture(scope="module")
def current(fixture_dir: Path) -> WorkbookSnapshot:
    return load_workbook_snapshot(fixture_dir / "current.xlsx")


@pytest.fixture(scope="module")
def alignment(baseline: WorkbookSnapshot, current: WorkbookSnapshot) -> WorkbookAlignment:
    return align_workbooks(baseline, current)


def _single_region(alignment: WorkbookAlignment, sheet: str) -> RegionAlignment:
    regions = alignment.regions[sheet]
    assert len(regions) == 1, f"expected one region pair on {sheet}"
    return regions[0]


def test_sheet_pairing(alignment: WorkbookAlignment) -> None:
    assert alignment.added_sheets == ["New_Analysis"]  # E15
    assert alignment.removed_sheets == ["Old_Sheet"]  # E14
    assert set(alignment.common_sheets) == {
        "Long_Monthly",
        "Wide_Weekly",
        "Dashboard",
        "Summary",
        "Params",
    }


def test_long_monthly_growth_isolated(alignment: WorkbookAlignment) -> None:
    region = _single_region(alignment, "Long_Monthly")
    assert region.rows.method == "keys"
    # 20 historical rows pair exactly onto themselves; zero false movement.
    assert region.rows.pairs == [(row, row) for row in range(2, 22)]
    assert region.rows.growth == [22, 23, 24, 25]  # EX01: Jun-26 block
    assert region.rows.inserted == []
    assert region.rows.deleted == []
    # Header columns are stable.
    assert region.columns.pairs == [(col, col) for col in range(1, 6)]
    assert region.columns.growth == []


def test_wide_weekly_deletion_and_growth(alignment: WorkbookAlignment) -> None:
    region = _single_region(alignment, "Wide_Weekly")
    columns = region.columns
    assert columns.method == "keys"
    assert columns.deleted == [4]  # E07: W03 lived in column D
    assert columns.growth == [21]  # EX02: W21 appended in column U
    assert columns.inserted == []
    # W04 onward shifted one column left; key alignment must absorb the shift.
    pair_map = dict(columns.pairs)
    assert pair_map[1] == 1  # label column
    assert pair_map[2] == 2 and pair_map[3] == 3  # W01, W02
    assert pair_map[5] == 4  # W04: E -> D
    assert pair_map[21] == 20  # W20: U -> T
    assert region.rows.pairs == [(row, row) for row in range(1, 5)]


def test_yearless_week_rollover_is_growth() -> None:
    baseline = [
        AxisEntry(index=1, key=("W51", 0), periods=(parse_period("W51"),)),
        AxisEntry(index=2, key=("W52", 0), periods=(parse_period("W52"),)),
    ]
    current = [
        *baseline,
        AxisEntry(index=3, key=("W01", 0), periods=(parse_period("W01"),)),
    ]

    aligned = _align_axis(baseline, current)

    assert aligned.growth == [3]
    assert aligned.inserted == []


def test_key_alignment_marks_low_confidence_positional_fallback() -> None:
    baseline = [
        AxisEntry(index=index, key=(label,), periods=(None,))
        for index, label in enumerate(("A", "B", "C", "D"), start=1)
    ]
    current = [
        AxisEntry(index=index, key=(label,), periods=(None,))
        for index, label in enumerate(("W", "X", "Y", "Z"), start=1)
    ]

    aligned = _align_axis(baseline, current)

    assert aligned.method == "positional"
    assert aligned.low_confidence_fallback


def test_disjoint_block_labels_align_positionally_and_remain_visible(
    tmp_path: Path,
) -> None:
    baseline = tmp_path / "baseline.xlsx"
    current = tmp_path / "current.xlsx"

    def write(path: Path, labels: list[str], offset: int) -> None:
        workbook = Workbook()
        sheet = workbook.active
        assert sheet is not None
        sheet.title = "Data"
        sheet.append(["Key", "Value"])
        for index, label in enumerate(labels, start=1):
            sheet.append([label, index + offset])
        workbook.save(path)

    write(baseline, ["A", "B", "C", "D"], 0)
    write(current, ["W", "X", "Y", "Z"], 100)

    result = run_qc(baseline_excel=baseline, current_excel=current)

    value_coverage = next(
        item for item in result.coverage if item.check_id == "excel-values"
    )
    formula_coverage = next(
        item for item in result.coverage if item.check_id == "excel-formulas"
    )
    assert value_coverage.state is CoverageState.CHECKED
    assert formula_coverage.state is CoverageState.CHECKED
    assert not value_coverage.detail
    assert not any(
        finding.finding_class is FindingClass.ALIGNMENT_LOW_CONFIDENCE
        for finding in result.findings
    )
    assert any(
        finding.finding_class is FindingClass.VALUE_CHANGED
        for finding in result.findings
    )


def test_dashboard_blocks_align(alignment: WorkbookAlignment) -> None:
    regions = alignment.regions["Dashboard"]
    assert len(regions) == 2
    kpi, headcount = regions

    kpi_rows = dict(kpi.rows.pairs)
    assert kpi_rows == {row: row for row in range(1, 6)}
    assert kpi.rows.deleted == [] and kpi.rows.inserted == []

    assert headcount.columns.growth == [11]  # EX05: Jun-26 headcount column (K)
    assert headcount.columns.deleted == []
    assert dict(headcount.rows.pairs) == {row: row for row in range(2, 5)}


def test_region_pairing_survives_inserted_block() -> None:
    def sheet(name: str, blocks: list[tuple[int, tuple[str, str]]]) -> SheetSnapshot:
        cells: dict[tuple[int, int], CellRecord] = {}
        for start_row, block_labels in blocks:
            for offset, label in enumerate(block_labels):
                row = start_row + offset
                cells[(row, 1)] = CellRecord(row=row, column=1, value=label)
                cells[(row, 2)] = CellRecord(row=row, column=2, value=100 + row)
        return SheetSnapshot(
            name=name,
            visibility="visible",
            max_row=max(row for row, _ in cells),
            max_column=2,
            cells=cells,
        )

    baseline = sheet("Dashboard", [(1, ("Revenue", "Margin")), (5, ("Headcount", "Attrition"))])
    current = sheet(
        "Dashboard",
        [
            (1, ("New metric", "Other")),
            (5, ("Revenue", "Margin")),
            (9, ("Headcount", "Attrition")),
        ],
    )

    pairs, unpaired_base, unpaired_current = _pair_regions(
        baseline,
        current,
        detect_regions(baseline),
        detect_regions(current),
    )

    assert [(base.cell_range, curr.cell_range) for base, curr in pairs] == [
        ("A1:B2", "A5:B6"),
        ("A5:B6", "A9:B10"),
    ]
    assert unpaired_base == []
    assert [region.cell_range for region in unpaired_current] == ["A1:B2"]


def test_summary_block_alignment(alignment: WorkbookAlignment) -> None:
    region = _single_region(alignment, "Summary")
    assert dict(region.rows.pairs) == {row: row for row in range(2, 7)}
    # C/D columns exist only in current (seeded error cells E05/E13): they are
    # unexpected insertions, not cadence growth.
    assert region.columns.inserted == [3, 4]
    assert region.columns.growth == []


def test_profile_sheet_ignore(baseline: WorkbookSnapshot, current: WorkbookSnapshot) -> None:
    profile = DeliverableProfile.model_validate(
        {"name": "test", "excel": {"ignore_sheets": ["Params", "Old_Sheet", "New_Analysis"]}}
    )
    aligned = align_workbooks(baseline, current, profile)
    assert aligned.added_sheets == []
    assert aligned.removed_sheets == []
    assert "Params" not in aligned.common_sheets


def test_cell_pairs_identity_on_unchanged_history(alignment: WorkbookAlignment) -> None:
    region = _single_region(alignment, "Long_Monthly")
    pairs = list(region.cell_pairs())
    assert ((7, 3), (7, 3)) in pairs  # E01's cell maps onto itself
    assert len(pairs) == 20 * 5
    assert all(base == curr for base, curr in pairs)


def test_alignment_trust_manifest_reconciles_every_paired_region(
    alignment: WorkbookAlignment,
) -> None:
    manifest = build_alignment_trust_manifest(alignment)
    trust_by_region = {
        (region.sheet, region.region_id): region
        for region in manifest.regions
    }

    for sheet, regions in alignment.regions.items():
        for region in regions:
            trust = trust_by_region[(sheet, region.current.region_id)]
            assert trust.baseline_range == region.baseline.cell_range
            assert trust.current_range == region.current.cell_range
            assert trust.row.paired == len(region.rows.pairs)
            assert trust.row.deleted == len(region.rows.deleted)
            assert trust.row.inserted == len(region.rows.inserted)
            assert trust.row.growth == len(region.rows.growth)
            assert trust.column.paired == len(region.columns.pairs)
            assert trust.column.deleted == len(region.columns.deleted)
            assert trust.column.inserted == len(region.columns.inserted)
            assert trust.column.growth == len(region.columns.growth)
            assert trust.comparable_cell_pairs == (
                len(region.rows.pairs) * len(region.columns.pairs)
            )
            assert trust.skipped_low_confidence_cells == (
                trust.comparable_cell_pairs if region.low_confidence else 0
            )
            assert "confidence_score" not in trust.model_dump(mode="json")


def test_alignment_trust_counts_all_low_confidence_pairs_as_skipped() -> None:
    baseline = TableRegion("Data", 1, 1, 4, 2, "block", None, 1, "none")
    current = TableRegion("Data", 1, 1, 4, 2, "block", None, 1, "none")
    rows = AxisAlignment(
        pairs=[(1, 1), (2, 2), (3, 3), (4, 4)],
        method="positional",
        low_confidence_fallback=True,
    )
    columns = AxisAlignment(pairs=[(1, 1), (2, 2)])
    alignment = WorkbookAlignment(
        common_sheets=["Data"],
        regions={
            "Data": [
                RegionAlignment(
                    baseline=baseline,
                    current=current,
                    rows=rows,
                    columns=columns,
                )
            ]
        },
        low_confidence_regions=["Data!A1:B4"],
    )

    trust = build_alignment_trust_manifest(alignment).regions[0]

    assert trust.low_confidence is True
    assert trust.row.method == "positional"
    assert trust.row.low_confidence_fallback is True
    assert trust.comparable_cell_pairs == 8
    assert trust.skipped_low_confidence_cells == 8


def test_alignment_trust_preserves_unpaired_regions_and_member_identity() -> None:
    baseline = TableRegion("Old", 2, 1, 5, 3, "long", 2, 1, "rows")
    current = TableRegion("New", 1, 2, 4, 6, "wide", 1, 2, "columns")
    alignment = WorkbookAlignment(
        unpaired_baseline_regions=[baseline],
        unpaired_current_regions=[current],
    )

    manifest = build_alignment_trust_manifest(alignment, artifact_member="ops")

    assert [entry.side for entry in manifest.unpaired] == ["baseline", "current"]
    assert [entry.artifact_member for entry in manifest.unpaired] == ["ops", "ops"]
    assert manifest.unpaired[0].cell_range == "A2:C5"
    assert manifest.unpaired[1].cell_range == "B1:F4"
    assert manifest.unpaired[0].orientation == "long"
    assert manifest.unpaired[1].orientation == "wide"


def test_alignment_trust_manifest_json_round_trip(
    alignment: WorkbookAlignment,
) -> None:
    manifest = build_alignment_trust_manifest(alignment, artifact_member="ops")

    restored = AlignmentTrustManifest.model_validate(
        manifest.model_dump(mode="json")
    )

    assert restored == manifest
    assert all(region.artifact_member == "ops" for region in restored.regions)


def test_full_cycle_result_carries_alignment_trust(qc_result) -> None:
    trust = qc_result.alignment_trust

    assert trust is not None
    assert trust.version == 1
    assert trust.regions
    assert sum(region.comparable_cell_pairs for region in trust.regions) > 0


# --- plan-20260913 Step 3: execution bindings drive alignment directly -----


def _write_renamed_pair(tmp_path: Path) -> tuple[Path, Path]:
    """A tiny keyed table on a sheet named ``Sheet2025`` in baseline and
    ``Sheet2026`` in current -- the same rows, shuffled, one value changed.
    """
    baseline_path = tmp_path / "renamed_baseline.xlsx"
    current_path = tmp_path / "renamed_current.xlsx"

    base_wb = Workbook()
    base_ws = base_wb.active
    assert base_ws is not None
    base_ws.title = "Sheet2025"
    base_ws.append(["ID", "Amount"])
    base_ws.append(["A1", 10])
    base_ws.append(["A2", 20])
    base_ws.append(["A3", 30])
    base_wb.save(baseline_path)

    curr_wb = Workbook()
    curr_ws = curr_wb.active
    assert curr_ws is not None
    curr_ws.title = "Sheet2026"
    curr_ws.append(["ID", "Amount"])
    curr_ws.append(["A3", 30])
    curr_ws.append(["A1", 10])
    curr_ws.append(["A2", 25])  # genuine value change
    curr_wb.save(current_path)
    return baseline_path, current_path


def _renamed_execution_bindings() -> ExecutionBindings:
    resolved = ResolvedInputConfigurationV1(
        members=(
            ResolvedMember(
                member_id="primary",
                sheets=(
                    ResolvedSheet(
                        sheet_id="data",
                        baseline_sheet_name="Sheet2025",
                        current_sheet_name="Sheet2026",
                        regions=(
                            ResolvedRegion(
                                region_id="r1",
                                mode="keyed",
                                header_intent="first_data_row",
                                current_data_range="A2:B4",
                                current_first_data_row=2,
                                columns=(
                                    ResolvedColumn(
                                        column_id="id",
                                        current_letter="A",
                                        alignment_role="identity",
                                    ),
                                ),
                            ),
                        ),
                    ),
                ),
            ),
        )
    )
    return ExecutionBindings(resolved)


def test_confirmed_rename_pairs_sheets_with_no_add_remove(tmp_path: Path) -> None:
    baseline_path, current_path = _write_renamed_pair(tmp_path)
    base_snapshot = load_workbook_snapshot(baseline_path)
    curr_snapshot = load_workbook_snapshot(current_path)

    alignment = align_workbooks(
        base_snapshot,
        curr_snapshot,
        execution_bindings=_renamed_execution_bindings(),
    )

    assert alignment.added_sheets == []
    assert alignment.removed_sheets == []
    assert alignment.renamed_sheets == {"Sheet2026": "Sheet2025"}
    assert "Sheet2026" in alignment.common_sheets


def test_execution_bound_keyed_region_drives_row_alignment(tmp_path: Path) -> None:
    baseline_path, current_path = _write_renamed_pair(tmp_path)
    base_snapshot = load_workbook_snapshot(baseline_path)
    curr_snapshot = load_workbook_snapshot(current_path)

    alignment = align_workbooks(
        base_snapshot,
        curr_snapshot,
        execution_bindings=_renamed_execution_bindings(),
    )

    region = _single_region(alignment, "Sheet2026")
    assert region.rows.method == "keys"
    assert region.rows.identity_columns == ("A",)
    # Row 1 (the header) pairs positionally, remaining visible outside keyed
    # matching. Shuffled data rows pair by identity: baseline row 4 (A3)
    # pairs with current row 2, baseline row 2 (A1) pairs with current row
    # 3, baseline row 3 (A2) pairs with current row 4 -- pure reorder, zero
    # deletions/insertions.
    assert sorted(region.rows.pairs) == [(1, 1), (2, 3), (3, 4), (4, 2)]
    assert region.rows.deleted == []
    assert region.rows.inserted == []


def test_without_execution_bindings_a_renamed_sheet_is_add_plus_remove(
    tmp_path: Path,
) -> None:
    """Legacy behavior guard: with no execution bindings, a physical rename
    is still ordinary sheet_removed + sheet_added -- byte-identical to
    before this module existed.
    """
    baseline_path, current_path = _write_renamed_pair(tmp_path)
    base_snapshot = load_workbook_snapshot(baseline_path)
    curr_snapshot = load_workbook_snapshot(current_path)

    alignment = align_workbooks(base_snapshot, curr_snapshot)

    assert alignment.added_sheets == ["Sheet2026"]
    assert alignment.removed_sheets == ["Sheet2025"]
    assert alignment.renamed_sheets == {}
    assert alignment.regions == {}


# --- plan-20260913 Step 12 Fix 4: confirmed column mappings drive the ------
# --- alignment engine's column axis, not just resolved-config reporting ---


def _write_column_moved_pair(tmp_path: Path) -> tuple[Path, Path]:
    """Baseline has Amount in column B and Notes in column C; current has
    them swapped (Notes in B, Amount in C). One genuine value change
    (R2's amount) rides along at the mapped location.
    """
    baseline_path = tmp_path / "colmoved_baseline.xlsx"
    current_path = tmp_path / "colmoved_current.xlsx"

    base_wb = Workbook()
    base_ws = base_wb.active
    assert base_ws is not None
    base_ws.title = "Data"
    base_ws.append(["ID", "Amount", "Notes"])
    base_ws.append(["R1", 100, "x"])
    base_ws.append(["R2", 200, "y"])
    base_wb.save(baseline_path)

    curr_wb = Workbook()
    curr_ws = curr_wb.active
    assert curr_ws is not None
    curr_ws.title = "Data"
    curr_ws.append(["ID", "Notes", "Amount"])
    curr_ws.append(["R1", "x", 100])
    curr_ws.append(["R2", "y", 999])  # genuine change: R2's amount 200 -> 999
    curr_wb.save(current_path)
    return baseline_path, current_path


def _column_moved_execution_bindings() -> ExecutionBindings:
    resolved = ResolvedInputConfigurationV1(
        members=(
            ResolvedMember(
                member_id="primary",
                sheets=(
                    ResolvedSheet(
                        sheet_id="data",
                        baseline_sheet_name="Data",
                        current_sheet_name="Data",
                        regions=(
                            ResolvedRegion(
                                region_id="r1",
                                current_data_range="A1:C3",
                                columns=(
                                    ResolvedColumn(
                                        column_id="amount",
                                        baseline_letter="B",
                                        current_letter="C",
                                    ),
                                    ResolvedColumn(
                                        column_id="notes",
                                        baseline_letter="C",
                                        current_letter="B",
                                    ),
                                ),
                            ),
                        ),
                    ),
                ),
            ),
        )
    )
    return ExecutionBindings(resolved)


def test_confirmed_column_mapping_pairs_columns_by_letter_not_position(
    tmp_path: Path,
) -> None:
    baseline_path, current_path = _write_column_moved_pair(tmp_path)
    base_snapshot = load_workbook_snapshot(baseline_path)
    curr_snapshot = load_workbook_snapshot(current_path)

    alignment = align_workbooks(
        base_snapshot,
        curr_snapshot,
        execution_bindings=_column_moved_execution_bindings(),
    )

    region = _single_region(alignment, "Data")
    columns = region.columns
    assert columns.method == "keys"
    pair_map = dict(columns.pairs)
    # Column A (ID) never moved: positional pairing among the unmapped
    # remainder still applies.
    assert pair_map[1] == 1
    # Baseline B (Amount) <-> current C; baseline C (Notes) <-> current B --
    # the OPPOSITE of what plain positional pairing would produce.
    assert pair_map[2] == 3
    assert pair_map[3] == 2
    assert sorted(columns.moved_pairs) == [(2, 3), (3, 2)]
    assert columns.deleted == []
    assert columns.inserted == []


def test_without_confirmed_column_mapping_a_moved_column_is_positional(
    tmp_path: Path,
) -> None:
    """Legacy behavior guard: with no execution bindings, a moved column is
    still ordinary positional pairing -- byte-identical to before this hook
    existed.
    """
    baseline_path, current_path = _write_column_moved_pair(tmp_path)
    base_snapshot = load_workbook_snapshot(baseline_path)
    curr_snapshot = load_workbook_snapshot(current_path)

    alignment = align_workbooks(base_snapshot, curr_snapshot)

    region = _single_region(alignment, "Data")
    assert dict(region.columns.pairs) == {1: 1, 2: 2, 3: 3}
    assert region.columns.moved_pairs == ()


def test_confirmed_column_mapping_changes_what_run_qc_actually_compares(
    tmp_path: Path,
) -> None:
    """The end-to-end proof: the SAME two files produce entirely different
    VALUE_CHANGED evidence depending only on whether a confirmed column
    mapping is threaded through -- proving the override changes what is
    actually COMPARED at run time, not merely what is stored/reported.
    """
    baseline_path, current_path = _write_column_moved_pair(tmp_path)

    unmapped = run_qc(baseline_excel=baseline_path, current_excel=current_path)
    unmapped_value_changes = {
        (finding.location, finding.baseline_value, finding.current_value)
        for finding in unmapped.findings
        if finding.finding_class is FindingClass.VALUE_CHANGED
    }
    # Without the mapping, every swapped cell looks like a spurious change:
    # baseline Amount(header/100/200) vs current Notes(header/"x"/"y") in
    # column B, and baseline Notes(header/"x"/"y") vs current
    # Amount(header/100/999) in column C -- 3 rows x 2 columns = 6.
    assert len(unmapped_value_changes) == 6

    mapped = run_qc(
        baseline_excel=baseline_path,
        current_excel=current_path,
        resolved_input_configuration=ResolvedInputConfigurationV1(
            members=(
                ResolvedMember(
                    member_id="primary",
                    sheets=(
                        ResolvedSheet(
                            sheet_id="data",
                            baseline_sheet_name="Data",
                            current_sheet_name="Data",
                            regions=(
                                ResolvedRegion(
                                    region_id="r1",
                                    current_data_range="A1:C3",
                                    columns=(
                                        ResolvedColumn(
                                            column_id="amount",
                                            baseline_letter="B",
                                            current_letter="C",
                                        ),
                                        ResolvedColumn(
                                            column_id="notes",
                                            baseline_letter="C",
                                            current_letter="B",
                                        ),
                                    ),
                                ),
                            ),
                        ),
                    ),
                ),
            )
        ),
    )
    mapped_value_changes = [
        finding
        for finding in mapped.findings
        if finding.finding_class is FindingClass.VALUE_CHANGED
    ]
    # With the mapping, only the ONE genuine content change survives: R2's
    # amount, correctly compared at its mapped physical location (baseline
    # B3=200 vs current C3=999).
    assert len(mapped_value_changes) == 1
    [only] = mapped_value_changes
    assert only.location == "C3"
    assert only.baseline_value == "200"
    assert only.current_value == "999"

    # The two moved columns each surface exactly one structural disclosure
    # instead of contributing to the cell-level VALUE_CHANGED noise above.
    moved = [
        finding
        for finding in mapped.findings
        if finding.finding_class is FindingClass.COLUMN_MOVED
    ]
    assert {(f.baseline_location, f.location) for f in moved} == {
        ("column B", "column C"),
        ("column C", "column B"),
    }
