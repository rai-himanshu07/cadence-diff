"""End-to-end artifact, history, and read-only guarantees for new QC modes."""

import hashlib
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from qc_tool.config.profile import DeliverableProfile
from qc_tool.coverage import CoverageState, QCRunMode
from qc_tool.engine import run_qc
from qc_tool.findings import FindingClass
from qc_tool.findings_store import FindingSequence
from qc_tool.history.store import RunHistory
from qc_tool.package import PackageManifest
from qc_tool.ui.app import perform_run


def _write_member_workbook(path: Path, values: tuple[int, ...]) -> None:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Data"
    sheet.append(["Key", "Amount", "Derived"])
    for index, value in enumerate(values, start=1):
        sheet.append([f"k{index}", value, f"=B{index + 1}*2"])
    workbook.save(path)


def _multi_cycle_files(tmp_path: Path) -> dict[str, Path]:
    files = {
        "baseline_excel:core": tmp_path / "core-baseline.xlsx",
        "current_excel:core": tmp_path / "core-current.xlsx",
        "baseline_excel:ops": tmp_path / "ops-baseline.xlsx",
        "current_excel:ops": tmp_path / "ops-current.xlsx",
    }
    _write_member_workbook(files["baseline_excel:core"], (10, 20, 30, 40))
    _write_member_workbook(files["current_excel:core"], (11, 20, 30, 40))
    _write_member_workbook(files["baseline_excel:ops"], (100, 200, 300, 400))
    _write_member_workbook(files["current_excel:ops"], (100, 250, 300, 400))
    return files


@pytest.mark.parametrize(
    ("mode", "roles"),
    [
        (QCRunMode.CURRENT_FILE_PREFLIGHT, ("current_excel",)),
        (QCRunMode.CURRENT_FILE_PREFLIGHT, ("current_ppt",)),
        (QCRunMode.FINAL_PACKAGE, ("current_excel", "current_ppt")),
    ],
)
def test_new_modes_write_reports_history_and_preserve_sources(
    mode: QCRunMode,
    roles: tuple[str, ...],
    fixture_dir: Path,
    tmp_path: Path,
) -> None:
    source_by_role = {
        "current_excel": fixture_dir / "current.xlsx",
        "current_ppt": fixture_dir / "current.pptx",
    }
    files = {role: source_by_role[role] for role in roles}
    before = {
        role: hashlib.sha256(path.read_bytes()).hexdigest()
        for role, path in files.items()
    }

    artifacts = perform_run(
        tmp_path / "work",
        files,
        {},
        DeliverableProfile(name="mode-e2e"),
        mode=mode,
        write_reports=True,
    )

    assert artifacts.result.mode is mode
    assert artifacts.result.coverage
    assert all(
        path.exists() and path.stat().st_size > 0
        for path in artifacts.report_paths.values()
    )
    workbook = load_workbook(artifacts.report_paths["excel"], read_only=True)
    assert "Coverage" in workbook.sheetnames
    html = artifacts.report_paths["html"].read_text(encoding="utf-8")
    assert mode.value in html and "Check coverage" in html
    record = RunHistory(tmp_path / "work" / "history.sqlite3").get_run(
        artifacts.run_id
    )
    assert record.mode is mode
    assert record.coverage == artifacts.result.coverage
    after = {
        role: hashlib.sha256(path.read_bytes()).hexdigest()
        for role, path in files.items()
    }
    assert after == before


def test_two_member_cycle_keeps_duplicate_sheet_evidence_isolated(
    tmp_path: Path,
) -> None:
    files = _multi_cycle_files(tmp_path)
    manifest = PackageManifest.from_role_files(files)
    before = {
        role: hashlib.sha256(path.read_bytes()).hexdigest()
        for role, path in files.items()
    }

    result = run_qc(
        profile=DeliverableProfile(name="multi-cycle"),
        package_manifest=manifest,
        package_files=files,
        mode=QCRunMode.CYCLE_COMPARISON,
    )

    member_value_findings = {
        finding.artifact_member: finding
        for finding in result.findings
        if finding.finding_class is FindingClass.VALUE_CHANGED
        and finding.sheet == "Data"
    }
    assert set(member_value_findings) == {"core", "ops"}
    assert member_value_findings["core"].location == "B2"
    assert member_value_findings["ops"].location == "B3"
    assert len({finding.finding_id for finding in result.findings}) == len(
        result.findings
    )
    assert result.package_manifest == manifest
    assert result.alignment_trust is not None
    assert {
        region.artifact_member for region in result.alignment_trust.regions
    } == {"core", "ops"}
    member_coverage = [
        item for item in result.coverage if item.artifact_member != "primary"
    ]
    assert member_coverage
    assert all(item.check_id.startswith("member:") for item in member_coverage)
    after = {
        role: hashlib.sha256(path.read_bytes()).hexdigest()
        for role, path in files.items()
    }
    assert after == before


def test_multi_member_cycle_reports_added_and_removed_members(
    tmp_path: Path,
) -> None:
    files = _multi_cycle_files(tmp_path)
    old_baseline = tmp_path / "old-baseline.xlsx"
    new_current = tmp_path / "new-current.xlsx"
    _write_member_workbook(old_baseline, (1, 2, 3, 4))
    _write_member_workbook(new_current, (5, 6, 7, 8))
    del files["baseline_excel:ops"]
    del files["current_excel:ops"]
    files["baseline_excel:old"] = old_baseline
    files["current_excel:new"] = new_current

    result = run_qc(
        profile=DeliverableProfile(name="member-events"),
        package_manifest=PackageManifest.from_role_files(files),
        package_files=files,
        mode=QCRunMode.CYCLE_COMPARISON,
    )

    events = {
        (finding.finding_class, finding.artifact_member)
        for finding in result.findings
        if finding.finding_class
        in {FindingClass.WORKBOOK_ADDED, FindingClass.WORKBOOK_REMOVED}
    }
    assert events == {
        (FindingClass.WORKBOOK_ADDED, "new"),
        (FindingClass.WORKBOOK_REMOVED, "old"),
    }
    member_coverage = next(
        item for item in result.coverage if item.check_id == "excel-package-members"
    )
    assert member_coverage.findings == 2


def test_multi_member_manifest_round_trips_through_history(tmp_path: Path) -> None:
    files = _multi_cycle_files(tmp_path)
    work_dir = tmp_path / "work"

    artifacts = perform_run(
        work_dir,
        files,
        {},
        DeliverableProfile(name="multi-history"),
        mode=QCRunMode.CYCLE_COMPARISON,
        package_manifest=PackageManifest.from_role_files(files),
    )
    record = RunHistory(work_dir / "history.sqlite3").get_run(artifacts.run_id)

    assert record.package_manifest == artifacts.result.package_manifest
    assert set(record.files) == set(files)
    assert set(record.file_hashes) == set(files)
    assert set(record.file_paths) == set(files)
    assert {finding.artifact_member for finding in record.findings} >= {
        "core",
        "ops",
    }


def test_eight_member_preflight_executes_every_workbook_at_the_cap(
    tmp_path: Path,
) -> None:
    files: dict[str, Path] = {}
    for index in range(8):
        role = f"current_excel:wb{index}"
        path = tmp_path / f"wb{index}.xlsx"
        _write_member_workbook(
            path,
            tuple(index * 100 + value for value in (1, 2, 3, 4)),
        )
        files[role] = path
    before = {role: hashlib.sha256(path.read_bytes()).hexdigest() for role, path in files.items()}

    result = run_qc(
        profile=DeliverableProfile(name="eight-member"),
        package_manifest=PackageManifest.from_role_files(files),
        package_files=files,
        mode=QCRunMode.CURRENT_FILE_PREFLIGHT,
    )

    assert result.package_manifest is not None
    assert len(result.package_manifest.members) == 8
    assert isinstance(result.findings, FindingSequence)
    assert {
        item.artifact_member
        for item in result.coverage
        if item.artifact == "excel"
    } >= {f"wb{index}" for index in range(8)}
    package_coverage = next(
        item for item in result.coverage if item.check_id == "excel-package-members"
    )
    assert "8 current Excel members" in package_coverage.detail
    assert {
        role: hashlib.sha256(path.read_bytes()).hexdigest()
        for role, path in files.items()
    } == before


def test_current_preflight_integrates_circular_findings_and_coverage(
    tmp_path: Path,
) -> None:
    source = tmp_path / "circular.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Data"
    sheet["A1"] = "=B1+1"
    sheet["B1"] = "=A1+1"
    workbook.save(source)

    result = run_qc(
        current_excel=source,
        mode=QCRunMode.CURRENT_FILE_PREFLIGHT,
    )

    circular = [
        finding
        for finding in result.findings
        if finding.finding_class is FindingClass.CIRCULAR_REFERENCE
    ]
    coverage = next(
        item for item in result.coverage
        if item.check_id == "excel-circular-references"
    )
    assert len(circular) == 1
    assert {"Data!A1", "Data!B1"} <= set(
        (circular[0].element or "").split("; ")
    )
    assert coverage.state is CoverageState.CHECKED
    assert coverage.findings == 1


def test_powerpoint_preflight_integrates_repetition_coverage(
    fixture_dir: Path,
) -> None:
    result = run_qc(
        current_ppt=fixture_dir / "current.pptx",
        mode=QCRunMode.CURRENT_FILE_PREFLIGHT,
    )

    coverage = next(
        item for item in result.coverage
        if item.check_id == "ppt-internal-repetition"
    )
    assert coverage.state in {CoverageState.CHECKED, CoverageState.DEGRADED}
    assert "eligible_groups=" in coverage.detail
    assert coverage.findings == sum(
        finding.finding_class is FindingClass.PPT_REPEATED_CLAIM_MISMATCH
        for finding in result.findings
    )
