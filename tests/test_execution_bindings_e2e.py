"""End-to-end proof that engine-native logical bindings drive a real
``run_qc()`` cycle comparison (plan-20260913, Step 3): a confirmed sheet
rename plus an execution-confirmed keyed region.

Deliberately exercises the full pipeline (load -> align -> structure ->
formulas -> values) rather than calling individual producers directly, so
this is real proof the ``resolved_input_configuration`` parameter is wired
correctly end to end, not just at the ``align_workbooks()`` layer.
"""

from __future__ import annotations

from pathlib import Path
from typing import Literal

import pytest
from openpyxl import Workbook

import qc_tool.engine as engine_module
from qc_tool.config.resolved_input import (
    ResolvedColumn,
    ResolvedInputConfigurationV1,
    ResolvedMember,
    ResolvedRegion,
    ResolvedSheet,
)
from qc_tool.coverage import QCRunMode
from qc_tool.engine import run_qc
from qc_tool.findings import FindingClass
from qc_tool.run_action import RunActionReason, RunBlockedError


def _write_pair(tmp_path: Path) -> tuple[Path, Path]:
    baseline_path = tmp_path / "baseline.xlsx"
    current_path = tmp_path / "current.xlsx"

    base_wb = Workbook()
    base_ws = base_wb.active
    assert base_ws is not None
    base_ws.title = "Ledger2025"
    base_ws.append(["ID", "Amount"])
    base_ws.append(["A1", 10])
    base_ws.append(["A2", 20])
    base_ws.append(["A3", 30])
    base_wb.save(baseline_path)

    curr_wb = Workbook()
    curr_ws = curr_wb.active
    assert curr_ws is not None
    curr_ws.title = "Ledger2026"
    curr_ws.append(["ID", "Amount"])
    curr_ws.append(["A3", 30])
    curr_ws.append(["A1", 10])
    curr_ws.append(["A2", 25])  # genuine value change
    curr_wb.save(current_path)
    return baseline_path, current_path


def _resolved_configuration() -> ResolvedInputConfigurationV1:
    return ResolvedInputConfigurationV1(
        mode=QCRunMode.CYCLE_COMPARISON,
        members=(
            ResolvedMember(
                member_id="primary",
                sheets=(
                    ResolvedSheet(
                        sheet_id="ledger",
                        baseline_sheet_name="Ledger2025",
                        current_sheet_name="Ledger2026",
                        regions=(
                            ResolvedRegion(
                                region_id="r1",
                                mode="keyed",
                                header_intent="first_data_row",
                                current_data_range="A2:B4",
                                current_first_data_row=2,
                                columns=(
                                    ResolvedColumn(
                                        column_id="id",
                                        current_letter="A",
                                        alignment_role="identity",
                                    ),
                                ),
                            ),
                        ),
                    ),
                ),
            ),
        ),
    )


def test_run_qc_with_resolved_configuration_reports_rename_and_real_change(
    tmp_path: Path,
) -> None:
    baseline_path, current_path = _write_pair(tmp_path)

    result = run_qc(
        baseline_excel=baseline_path,
        current_excel=current_path,
        mode=QCRunMode.CYCLE_COMPARISON,
        resolved_input_configuration=_resolved_configuration(),
    )

    classes = [f.finding_class for f in result.findings]
    assert classes.count(FindingClass.SHEET_RENAMED) == 1
    assert classes.count(FindingClass.SHEET_ADDED) == 0
    assert classes.count(FindingClass.SHEET_REMOVED) == 0
    # Pure reorder produces no row insertion/deletion noise.
    assert classes.count(FindingClass.ROW_INSERTED) == 0
    assert classes.count(FindingClass.ROW_DELETED) == 0
    # The genuine value edit (A2: 20 -> 25) still reports.
    value_changes = [f for f in result.findings if f.finding_class is FindingClass.VALUE_CHANGED]
    assert len(value_changes) == 1
    assert value_changes[0].sheet == "Ledger2026"


def test_resolved_configuration_identity_matching_is_case_sensitive(
    tmp_path: Path,
) -> None:
    """Step 8: the mode-aware configuration workspace's execution-bindings
    adapter always requests exact typed equality (never the legacy engine's
    own strip+casefold identity normalization) -- an ID differing only in
    case is a genuine identity mismatch, not a match.
    """
    baseline_path = tmp_path / "baseline.xlsx"
    current_path = tmp_path / "current.xlsx"
    base_wb = Workbook()
    base_ws = base_wb.active
    assert base_ws is not None
    base_ws.title = "Ledger2025"
    base_ws.append(["ID", "Amount"])
    base_ws.append(["A1", 10])
    base_wb.save(baseline_path)
    curr_wb = Workbook()
    curr_ws = curr_wb.active
    assert curr_ws is not None
    curr_ws.title = "Ledger2025"
    curr_ws.append(["ID", "Amount"])
    curr_ws.append(["a1", 10])  # same value, different case
    curr_wb.save(current_path)

    resolved = ResolvedInputConfigurationV1(
        mode=QCRunMode.CYCLE_COMPARISON,
        members=(
            ResolvedMember(
                member_id="primary",
                sheets=(
                    ResolvedSheet(
                        sheet_id="ledger",
                        baseline_sheet_name="Ledger2025",
                        current_sheet_name="Ledger2025",
                        regions=(
                            ResolvedRegion(
                                region_id="r1",
                                mode="keyed",
                                header_intent="first_data_row",
                                current_data_range="A2:B2",
                                current_first_data_row=2,
                                columns=(
                                    ResolvedColumn(
                                        column_id="id",
                                        current_letter="A",
                                        alignment_role="identity",
                                    ),
                                ),
                            ),
                        ),
                    ),
                ),
            ),
        ),
    )

    result = run_qc(
        baseline_excel=baseline_path,
        current_excel=current_path,
        mode=QCRunMode.CYCLE_COMPARISON,
        resolved_input_configuration=resolved,
    )

    classes = [f.finding_class for f in result.findings]
    # "A1" and "a1" are different identities under exact typed equality --
    # the position-1 row's key value genuinely changed, not a silent match.
    assert classes.count(FindingClass.ROW_KEY_CHANGED) == 1


def test_structural_and_formula_findings_carry_a_bounded_logical_address(
    tmp_path: Path,
) -> None:
    baseline_path, current_path = _write_pair(tmp_path)
    result = run_qc(
        baseline_excel=baseline_path,
        current_excel=current_path,
        mode=QCRunMode.CYCLE_COMPARISON,
        resolved_input_configuration=_resolved_configuration(),
    )

    renamed = next(
        f for f in result.findings if f.finding_class is FindingClass.SHEET_RENAMED
    )
    assert renamed.logical_address is not None
    assert renamed.logical_address.member_id == "primary"
    assert renamed.logical_address.sheet_id == "ledger"
    # Never a sheet name, range, or value -- content-free ids only.
    assert "Ledger2025" not in renamed.logical_address.model_dump_json()
    assert "Ledger2026" not in renamed.logical_address.model_dump_json()
    # The private sidecar never reaches the public payload.
    assert "logical_address" not in renamed.model_dump(mode="json")


def test_run_qc_without_resolved_configuration_stays_legacy(tmp_path: Path) -> None:
    """Guard: omitting ``resolved_input_configuration`` (the default) keeps
    a physical rename as ordinary sheet_removed + sheet_added.
    """
    baseline_path, current_path = _write_pair(tmp_path)

    result = run_qc(
        baseline_excel=baseline_path,
        current_excel=current_path,
        mode=QCRunMode.CYCLE_COMPARISON,
    )

    classes = [f.finding_class for f in result.findings]
    assert FindingClass.SHEET_RENAMED not in classes
    assert classes.count(FindingClass.SHEET_ADDED) == 1
    assert classes.count(FindingClass.SHEET_REMOVED) == 1


def _resolved_region_mode(mode: Literal["positional", "excluded"]):
    return ResolvedInputConfigurationV1(
        mode=QCRunMode.CYCLE_COMPARISON,
        members=(
            ResolvedMember(
                member_id="primary",
                sheets=(
                    ResolvedSheet(
                        sheet_id="data",
                        baseline_sheet_name="Data",
                        current_sheet_name="Data",
                        regions=(
                            ResolvedRegion(
                                region_id="data_a1_c5",
                                mode=mode,
                                baseline_outer_range="A1:C5",
                                current_outer_range="A1:C5",
                                baseline_data_range="A1:C5",
                                current_data_range="A1:C5",
                                coverage=(
                                    "positional"
                                    if mode == "positional"
                                    else "excluded"
                                ),
                            ),
                        ),
                    ),
                ),
            ),
        ),
    )


def _write_region_mode_pair(tmp_path: Path) -> tuple[Path, Path]:
    baseline_path = tmp_path / "baseline.xlsx"
    current_path = tmp_path / "current.xlsx"
    baseline = Workbook()
    baseline_sheet = baseline.active
    assert baseline_sheet is not None
    baseline_sheet.title = "Data"
    baseline_sheet.append(["ID", "Value", "Formula"])
    for index in range(4):
        baseline_sheet.append([f"K{index}", index, f"=B{index + 2}*2"])
    baseline.save(baseline_path)
    current = Workbook()
    current_sheet = current.active
    assert current_sheet is not None
    current_sheet.title = "Data"
    current_sheet.append(["ID", "Value", "Formula"])
    for index in reversed(range(4)):
        current_sheet.append(
            [
                f"K{index}",
                index + (10 if index == 2 else 0),
                "#REF!" if index == 1 else f"=B{5 - index}*3",
            ]
        )
    current.save(current_path)
    return baseline_path, current_path


def test_explicit_positional_region_skips_ranked_screening(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    baseline_path, current_path = _write_region_mode_pair(tmp_path)

    def ranked_screen_must_not_run(*_args, **_kwargs):
        raise AssertionError("explicit positional regions must not be screened")

    monkeypatch.setattr(
        engine_module,
        "detect_ranked_table_candidate",
        ranked_screen_must_not_run,
    )

    result = run_qc(
        baseline_excel=baseline_path,
        current_excel=current_path,
        mode=QCRunMode.CYCLE_COMPARISON,
        resolved_input_configuration=_resolved_region_mode("positional"),
    )

    assert result.alignment_trust is not None
    assert any(
        region.row.method == "positional"
        for region in result.alignment_trust.regions
    )


def test_excluded_region_produces_no_region_findings_and_discloses_scope(
    tmp_path: Path,
) -> None:
    baseline_path, current_path = _write_region_mode_pair(tmp_path)

    result = run_qc(
        baseline_excel=baseline_path,
        current_excel=current_path,
        mode=QCRunMode.CYCLE_COMPARISON,
        resolved_input_configuration=_resolved_region_mode("excluded"),
    )

    assert not any(
        finding.sheet == "Data" and finding.location
        for finding in result.findings
    )
    exclusion = next(
        item
        for item in result.coverage
        if item.check_id == "excel-configured-region-exclusions"
    )
    assert exclusion.state.value == "not_included"
    assert exclusion.findings == 1


def test_resolved_ignore_and_expected_refresh_columns_reach_the_value_diff(
    tmp_path: Path,
) -> None:
    """Step 8: ignore/expected-refresh column decisions declared through the
    mode-aware configuration workspace fold into the SAME ignore/refresh
    range-set mechanism a legacy profile's ``ignore_ranges``/
    ``refresh_ranges`` already use -- not a parallel, inert code path.
    """
    from qc_tool.findings import FindingExpectedReason, Severity

    baseline_path = tmp_path / "baseline.xlsx"
    current_path = tmp_path / "current.xlsx"
    base_wb = Workbook()
    base_ws = base_wb.active
    assert base_ws is not None
    base_ws.title = "Data"
    base_ws.append(["ID", "Scratch", "Cycle"])
    base_ws.append(["A1", "old-scratch", 100])
    base_wb.save(baseline_path)
    curr_wb = Workbook()
    curr_ws = curr_wb.active
    assert curr_ws is not None
    curr_ws.title = "Data"
    curr_ws.append(["ID", "Scratch", "Cycle"])
    curr_ws.append(["A1", "new-scratch", 200])  # both columns changed
    curr_wb.save(current_path)

    resolved = ResolvedInputConfigurationV1(
        mode=QCRunMode.CYCLE_COMPARISON,
        members=(
            ResolvedMember(
                member_id="primary",
                sheets=(
                    ResolvedSheet(
                        sheet_id="data",
                        baseline_sheet_name="Data",
                        current_sheet_name="Data",
                        regions=(
                            ResolvedRegion(
                                region_id="r1",
                                mode="automatic",
                                current_data_range="A2:C2",
                                columns=(
                                    ResolvedColumn(
                                        column_id="scratch",
                                        current_letter="B",
                                        comparison_policy="ignore",
                                    ),
                                    ResolvedColumn(
                                        column_id="cycle",
                                        current_letter="C",
                                        comparison_policy="expected_refresh",
                                    ),
                                ),
                            ),
                        ),
                    ),
                ),
            ),
        ),
    )

    result = run_qc(
        baseline_excel=baseline_path,
        current_excel=current_path,
        mode=QCRunMode.CYCLE_COMPARISON,
        resolved_input_configuration=resolved,
    )

    value_findings = {
        f.location: f for f in result.findings if f.finding_class is FindingClass.VALUE_CHANGED
    }
    # Column B (ignore) never produces a finding at all.
    assert "B2" not in value_findings
    # Column C (expected_refresh) still reports, but as Expected.
    assert "C2" in value_findings
    assert value_findings["C2"].expected_reason is FindingExpectedReason.PROFILE_REFRESH
    from qc_tool.triage.rules import assign_severity

    assert assign_severity(value_findings["C2"]) is Severity.EXPECTED


def _write_pair_with_a_blank_identity_key(tmp_path: Path) -> tuple[Path, Path]:
    baseline_path = tmp_path / "baseline.xlsx"
    current_path = tmp_path / "current.xlsx"
    base_wb = Workbook()
    base_ws = base_wb.active
    assert base_ws is not None
    base_ws.title = "Data"
    base_ws.append(["ID", "Value"])
    base_ws.append(["A1", 10])
    base_ws.append(["", 20])  # blank identity key
    base_wb.save(baseline_path)
    curr_wb = Workbook()
    curr_ws = curr_wb.active
    assert curr_ws is not None
    curr_ws.title = "Data"
    curr_ws.append(["ID", "Value"])
    curr_ws.append(["A1", 11])
    curr_ws.append(["", 21])  # blank identity key
    curr_wb.save(current_path)
    return baseline_path, current_path


def _resolved_configuration_with_keyed_region(
    *, blank_key_policy: Literal["system_default", "tolerate", "block"]
) -> ResolvedInputConfigurationV1:
    return ResolvedInputConfigurationV1(
        mode=QCRunMode.CYCLE_COMPARISON,
        members=(
            ResolvedMember(
                member_id="primary",
                sheets=(
                    ResolvedSheet(
                        sheet_id="data",
                        baseline_sheet_name="Data",
                        current_sheet_name="Data",
                        regions=(
                            ResolvedRegion(
                                region_id="r1",
                                mode="keyed",
                                current_outer_range="A1:B3",
                                current_data_range="A2:B3",
                                blank_key_policy=blank_key_policy,
                                columns=(
                                    ResolvedColumn(
                                        column_id="id",
                                        current_letter="A",
                                        alignment_role="identity",
                                    ),
                                ),
                            ),
                        ),
                    ),
                ),
            ),
        ),
    )


def test_blank_identity_key_blocks_the_run_when_the_region_policy_is_block(
    tmp_path: Path,
) -> None:
    """Step 8: a region's ``blank_key_policy == "block"`` refuses the run
    before any diff, rather than silently letting the blank-keyed row
    surface as an ordinary insert/delete (today's uniform "tolerate"
    engine behavior regardless of the declared policy).
    """
    baseline_path, current_path = _write_pair_with_a_blank_identity_key(tmp_path)
    resolved = _resolved_configuration_with_keyed_region(blank_key_policy="block")

    with pytest.raises(RunBlockedError) as excinfo:
        run_qc(
            baseline_excel=baseline_path,
            current_excel=current_path,
            mode=QCRunMode.CYCLE_COMPARISON,
            resolved_input_configuration=resolved,
        )

    action = excinfo.value.action_required
    assert action.reason is RunActionReason.BLANK_IDENTITY_KEY_BLOCKED
    assert action.items
    assert action.items[0].sheet == "Data"
    assert action.items[0].label == "r1"
    # Only a bounded row count is ever disclosed, never a key value.
    serialized = action.model_dump_json()
    assert "10" not in serialized
    assert "11" not in serialized


def test_blank_identity_key_does_not_block_under_the_default_policy(
    tmp_path: Path,
) -> None:
    """The SAME blank-keyed fixture, with the region left at its default
    ``"system_default"`` policy, must run to completion exactly like
    today -- a blank key is excluded from matching (an ordinary insert/
    delete), never a refusal, preserving existing behavior byte-for-byte
    for every profile that never opts into ``"block"``.
    """
    baseline_path, current_path = _write_pair_with_a_blank_identity_key(tmp_path)
    resolved = _resolved_configuration_with_keyed_region(blank_key_policy="system_default")

    result = run_qc(
        baseline_excel=baseline_path,
        current_excel=current_path,
        mode=QCRunMode.CYCLE_COMPARISON,
        resolved_input_configuration=resolved,
    )

    assert result is not None
