"""Compact formula-cycle graph and circular-reference evidence contracts."""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook

from qc_tool.coverage import CoverageState
from qc_tool.excel.dependency import (
    DependencyGraph,
    FormulaCycleGraph,
    Node,
    build_dependency_graph,
    detect_circular_references,
)
from qc_tool.excel.references import ReferenceStatus, ResolvedRange
from qc_tool.findings import FindingClass
from qc_tool.io.loader import load_workbook_snapshot
from qc_tool.progress import CancellationToken, RunCancelled


def _cell(node: Node) -> ResolvedRange:
    sheet, row, column = node
    return ResolvedRange(sheet, row, column, row, column)


def test_cycle_graph_preserves_literal_self_reference() -> None:
    node: Node = ("Sheet1", 1, 1)
    graph = FormulaCycleGraph([node])

    graph.add_range(_cell(node), node)

    assert graph.edge_count == 1
    assert graph.strongly_connected_components() == ((node,),)


def test_cycle_graph_reports_mutual_and_disjoint_components_deterministically() -> None:
    first_a: Node = ("S", 1, 1)
    first_b: Node = ("S", 1, 2)
    second_a: Node = ("S", 2, 1)
    second_b: Node = ("S", 2, 2)
    graph = FormulaCycleGraph([second_b, first_b, second_a, first_a])
    for precedent, dependent in (
        (first_b, first_a),
        (first_a, first_b),
        (second_b, second_a),
        (second_a, second_b),
    ):
        graph.add_range(_cell(precedent), dependent)

    components = graph.strongly_connected_components()

    assert components == (
        (first_a, first_b),
        (second_a, second_b),
    )


def test_cycle_graph_range_edges_include_formula_nodes_only() -> None:
    dependent: Node = ("S", 1, 1)
    range_formula: Node = ("S", 2, 3)
    graph = FormulaCycleGraph([dependent, range_formula])

    graph.add_range(ResolvedRange("S", 1, 3, 3, 3), dependent)
    graph.add_range(_cell(dependent), range_formula)

    assert graph.edge_count == 2
    assert graph.strongly_connected_components() == (
        (("S", 1, 1), ("S", 2, 3)),
    )


def test_impact_graph_still_drops_self_edges() -> None:
    node: Node = ("S", 1, 1)
    graph = DependencyGraph([node])

    graph.add_direct(node, node)
    graph.formula_cycles.add_range(_cell(node), node)

    assert graph.direct_dependents(node) == set()
    assert graph.formula_cycles.strongly_connected_components() == ((node,),)


def test_detect_circular_references_emits_one_bounded_finding_per_component() -> None:
    nodes = [("S", row, 1) for row in range(1, 13)]
    graph = DependencyGraph(nodes)
    for index, dependent in enumerate(nodes):
        graph.formula_cycles.add_range(
            _cell(nodes[(index + 1) % len(nodes)]),
            dependent,
        )

    result = detect_circular_references(graph)

    assert len(result.findings) == 1
    finding = result.findings[0]
    assert finding.finding_class is FindingClass.CIRCULAR_REFERENCE
    assert finding.location == "A1"
    assert finding.current_value == "12 formula cells"
    assert finding.event_key.startswith("excel-cycle:")
    assert finding.event_key == finding.root_cause_key
    assert "+4 more cells" in (finding.element or "")
    assert result.coverage.state is CoverageState.CHECKED
    assert result.coverage.findings == 1


def test_circular_event_key_is_deterministic() -> None:
    first: Node = ("S", 1, 1)
    second: Node = ("S", 1, 2)

    def event_key() -> str:
        graph = DependencyGraph([first, second])
        graph.formula_cycles.add_range(_cell(second), first)
        graph.formula_cycles.add_range(_cell(first), second)
        return detect_circular_references(graph).findings[0].event_key

    assert event_key() == event_key()


def test_clean_cycle_graph_has_checked_zero_component_coverage() -> None:
    first: Node = ("S", 1, 1)
    second: Node = ("S", 1, 2)
    graph = DependencyGraph([first, second])
    graph.formula_cycles.add_range(_cell(first), second)

    result = detect_circular_references(graph)

    assert result.findings == ()
    assert result.coverage.state is CoverageState.CHECKED
    assert result.coverage.check_id == "excel-circular-references"
    assert "0 circular components" in result.coverage.detail


def test_unsupported_invalid_parse_and_edge_budget_gaps_degrade_coverage() -> None:
    node: Node = ("S", 1, 1)
    other: Node = ("S", 1, 2)
    graph = DependencyGraph([node, other])
    graph.record_reference(ReferenceStatus.UNSUPPORTED, "dynamic_reference")
    graph.record_reference(ReferenceStatus.INVALID, "missing_sheet")
    graph.record_parse_error()
    graph.formula_cycles = FormulaCycleGraph([node, other], max_edges=1)
    graph.formula_cycles.add_range(ResolvedRange("S", 1, 1, 1, 2), node)

    result = detect_circular_references(graph)

    assert result.coverage.state is CoverageState.DEGRADED
    assert "1 unsupported references" in result.coverage.detail
    assert "1 invalid references" in result.coverage.detail
    assert "1 unparseable formulas" in result.coverage.detail
    assert "1 formula ranges exceeded the cycle-edge budget" in (
        result.coverage.detail
    )


def test_build_dependency_graph_detects_real_self_reference(tmp_path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Data"
    sheet["A1"] = "=A1+1"
    sheet["B1"] = "=A1*2"
    source = tmp_path / "self-cycle.xlsx"
    workbook.save(source)

    graph = build_dependency_graph(load_workbook_snapshot(source))
    result = detect_circular_references(graph)

    assert len(result.findings) == 1
    assert result.findings[0].location == "A1"
    assert "Data!A1" in (result.findings[0].element or "")


def test_cycle_graph_cancellation_is_cooperative() -> None:
    token = CancellationToken()
    token.cancel()
    graph = FormulaCycleGraph(
        [("S", row, 1) for row in range(1, 10_002)]
    )

    with pytest.raises(RunCancelled):
        graph.strongly_connected_components(token)


def test_cycle_graph_rejects_non_formula_dependent() -> None:
    graph = FormulaCycleGraph([("S", 1, 1)])

    with pytest.raises(ValueError, match="not a formula node"):
        graph.add_range(
            ResolvedRange("S", 1, 1, 1, 1),
            ("S", 2, 1),
        )
