"""Step 6: confirmed composite row identity (Acceptance Criteria 8-9).

Real, tiny xlsx fixtures built directly via openpyxl and run through the
actual engine (mirrors the Step 5 `tests/test_comparison_prerequisites.py`
convention) rather than hand-built snapshots, so the full load -> align ->
diff pipeline is exercised end to end.

Every workbook is a "block" region (no period axis) with a plain numeric
Rank column in the leftmost position -- exactly like a real ranked
leaderboard. That column is numeric (not text), so the pre-existing
stable-label block heuristic (`_block_labels_are_stable`, which only ever
counts string values) never key-aligns it, and it is a plain arithmetic
sequence, so the detector's own sequence-like screen would reject it as an
identity candidate too. Both properties mean the region falls back to plain
positional row alignment until a `RowIdentityRule` is confirmed -- exactly
the "unconfigured positional block region" precondition this feature acts on.

Formula checks (`qc_tool.excel.formulas.diff_workbook_formulas`) take no
`ignore`/`value_only_ignore` parameter at all -- they operate purely on the
(correctly identity-aligned) `WorkbookAlignment` and the workbook snapshots,
so they are structurally unaffected by ordinal-column suppression and are
not re-tested here; the style-change test below covers the analogous
"presentation checks stay active" guarantee for the same mechanism.
"""

from __future__ import annotations

import random
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font

from qc_tool.config.profile import (
    DeliverableProfile,
    ExcelProfile,
    RowIdentityRule,
    SheetProfile,
    default_profile,
)
from qc_tool.coverage import QCRunMode
from qc_tool.engine import _aggregate_ranked_package_actions, run_qc
from qc_tool.excel.align import AlignmentTrustManifest, AlignmentTrustManifestV2
from qc_tool.findings import FindingClass
from qc_tool.io.model import CellValue
from qc_tool.package import (
    PackageArtifact,
    PackageManifest,
    PackageMember,
    PackageSide,
)
from qc_tool.run_action import (
    MAX_RUN_ACTION_ITEMS,
    RunActionItem,
    RunActionReason,
    RunActionRequired,
    RunBlockedError,
)
from qc_tool.ui.ranked_table_dialog import apply_view_model, view_model_from_action

#: Large enough that a fully displaced permutation clears the detector's
#: MIN_PROJECTED_MISMATCHES (10,000) floor -- see tests/test_ranked_identity.py
#: for the unit-level threshold proofs; this module only needs one scale big
#: enough to exercise the detector through the real engine.
_LARGE_N = 6000


def _write_panel(
    path: Path,
    rows: list[list[CellValue]],
    *,
    headers: tuple[str, ...] = ("Rank", "ID", "Value"),
) -> None:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Panel"
    for col, header in enumerate(headers, start=1):
        sheet.cell(row=1, column=col, value=header)
    for r, row in enumerate(rows, start=2):
        for c, value in enumerate(row, start=1):
            sheet.cell(row=r, column=c, value=value)
    workbook.save(path)


def _base_rows(n: int) -> list[list[CellValue]]:
    return [[i + 1, f"ID{i}", 100.0 + i] for i in range(n)]


def _detector_rows(n: int) -> list[list[CellValue]]:
    """Like ``_base_rows`` but with two extra non-rank value columns.

    A lone Rank+Value layout's Rank column is positionally self-consistent
    by definition (row i's rank is always i), which cancels out Value's own
    key-vs-position improvement signal once it is the sole other column --
    see the analogous fix in tests/test_ranked_identity.py. Extra value
    columns dilute that so a full shuffle clears MIN_PROJECTED_MISMATCHES.
    """
    return [
        [i + 1, f"ID{i}", 100.0 + i, 200.0 + i * 2, 300.0 - i] for i in range(n)
    ]


def _shuffled(rows: list[list[CellValue]], seed: int) -> list[list[CellValue]]:
    """A reordering of ``rows`` with column 0 (Rank) renumbered to reflect
    the new display order -- exactly how a real rank column behaves."""
    shuffled = [list(row) for row in rows]
    random.Random(seed).shuffle(shuffled)
    for position, row in enumerate(shuffled):
        row[0] = position + 1
    return shuffled


def _profile_with_rule(rule: RowIdentityRule, sheet: str = "Panel") -> DeliverableProfile:
    return DeliverableProfile(
        name="ranked",
        excel=ExcelProfile(sheets={sheet: SheetProfile(row_identity_rules=[rule])}),
    )


def _classes(result, *classes: FindingClass) -> list:
    return [f for f in result.findings if f.finding_class in classes]


# --- confirmed identity realigns correctly --------------------------------


def test_confirmed_identity_pure_permutation_yields_zero_value_findings(
    tmp_path: Path,
) -> None:
    n = 30
    base_rows = _base_rows(n)
    curr_rows = _shuffled(base_rows, seed=1234)
    baseline, current = tmp_path / "baseline.xlsx", tmp_path / "current.xlsx"
    _write_panel(baseline, base_rows)
    _write_panel(current, curr_rows)
    profile = _profile_with_rule(
        RowIdentityRule(anchor_cell="B2", identity_columns=["B"], ordinal_columns=["A"])
    )

    result = run_qc(
        baseline_excel=baseline,
        current_excel=current,
        profile=profile,
        mode=QCRunMode.CYCLE_COMPARISON,
    )

    assert _classes(result, FindingClass.VALUE_CHANGED) == []
    assert _classes(result, FindingClass.ROW_INSERTED, FindingClass.ROW_DELETED) == []


def test_confirmed_identity_still_reports_a_genuine_value_change(
    tmp_path: Path,
) -> None:
    n = 30
    base_rows = _base_rows(n)
    curr_rows = _shuffled(base_rows, seed=1234)
    for row in curr_rows:
        if row[1] == "ID9":
            row[2] = 999.0
            break
    baseline, current = tmp_path / "baseline.xlsx", tmp_path / "current.xlsx"
    _write_panel(baseline, base_rows)
    _write_panel(current, curr_rows)
    profile = _profile_with_rule(
        RowIdentityRule(anchor_cell="B2", identity_columns=["B"], ordinal_columns=["A"])
    )

    result = run_qc(
        baseline_excel=baseline,
        current_excel=current,
        profile=profile,
        mode=QCRunMode.CYCLE_COMPARISON,
    )

    value_findings = _classes(result, FindingClass.VALUE_CHANGED)
    assert len(value_findings) == 1
    assert "999" in (value_findings[0].current_value or "")


def test_composite_identity_resolves_single_column_ambiguity(tmp_path: Path) -> None:
    n = 30
    half = n // 2
    base_rows = [
        [i + 1, f"Team {'A' if i < half else 'B'}", f"Member{i % half}", 100.0 + i]
        for i in range(n)
    ]
    curr_rows = _shuffled(base_rows, seed=7)
    baseline, current = tmp_path / "baseline.xlsx", tmp_path / "current.xlsx"
    _write_panel(baseline, base_rows, headers=("Rank", "Team", "Member", "Value"))
    _write_panel(current, curr_rows, headers=("Rank", "Team", "Member", "Value"))
    profile = _profile_with_rule(
        RowIdentityRule(
            anchor_cell="B2", identity_columns=["B", "C"], ordinal_columns=["A"]
        )
    )

    result = run_qc(
        baseline_excel=baseline,
        current_excel=current,
        profile=profile,
        mode=QCRunMode.CYCLE_COMPARISON,
    )

    # Neither Team nor Member alone is unique -- only the pair is. If the
    # alignment fell back to positional (or key'd on one column), the
    # shuffle would produce spurious value-changed findings.
    assert _classes(result, FindingClass.VALUE_CHANGED) == []


def test_ambiguous_duplicate_group_is_skipped_by_default(tmp_path: Path) -> None:
    n = 10
    base_rows = [[i + 1, f"ID{i}", 100.0 + i] for i in range(n)]
    base_rows[3][1] = "DUP"
    base_rows[7][1] = "DUP"
    curr_rows = _shuffled(base_rows, seed=3)
    baseline, current = tmp_path / "baseline.xlsx", tmp_path / "current.xlsx"
    _write_panel(baseline, base_rows)
    _write_panel(current, curr_rows)
    profile = _profile_with_rule(
        RowIdentityRule(
            anchor_cell="B2",
            identity_columns=["B"],
            ordinal_columns=["A"],
            duplicate_policy="skip",
        )
    )

    result = run_qc(
        baseline_excel=baseline,
        current_excel=current,
        profile=profile,
        mode=QCRunMode.CYCLE_COMPARISON,
    )

    assert isinstance(result.alignment_trust, AlignmentTrustManifestV2)
    region = result.alignment_trust.regions[0]
    assert region.duplicate_policy == "skip"
    assert region.skipped_duplicate_groups == 1
    assert region.skipped_duplicate_rows == 4  # 2 "DUP" rows on each side
    # Every uniquely identified row still aligns cleanly; nothing about the
    # skipped, ambiguous "DUP" rows produces a spurious value comparison.
    assert _classes(result, FindingClass.VALUE_CHANGED) == []


def test_blank_identity_component_is_never_guessed(tmp_path: Path) -> None:
    n = 10
    base_rows = [[i + 1, f"ID{i}", 100.0 + i] for i in range(n)]
    base_rows[5][1] = None  # baseline data row (ID5's slot) has no identity
    curr_rows = _shuffled(base_rows, seed=9)
    # The shuffled copy of that same blank row is given a brand-new, valid
    # identity absent from the baseline entirely -- proves the baseline's
    # blank-key row is never silently paired with anything, valid or not.
    for row in curr_rows:
        if row[1] is None:
            row[1] = "ID_RESTORED"
            row[2] = 999.0
            break
    baseline, current = tmp_path / "baseline.xlsx", tmp_path / "current.xlsx"
    _write_panel(baseline, base_rows)
    _write_panel(current, curr_rows)
    profile = _profile_with_rule(
        RowIdentityRule(anchor_cell="B2", identity_columns=["B"], ordinal_columns=["A"])
    )

    result = run_qc(
        baseline_excel=baseline,
        current_excel=current,
        profile=profile,
        mode=QCRunMode.CYCLE_COMPARISON,
    )

    assert len(_classes(result, FindingClass.ROW_DELETED)) == 1
    assert len(_classes(result, FindingClass.ROW_INSERTED)) == 1
    # No spurious value comparison ever pairs the blank-key baseline row
    # with the differently-identified restored row.
    assert _classes(result, FindingClass.VALUE_CHANGED) == []


def test_insertion_and_deletion_detected_via_identity(tmp_path: Path) -> None:
    n = 10
    base_rows = [[i + 1, f"ID{i}", 100.0 + i] for i in range(n)]
    curr_rows = [list(row) for row in base_rows if row[1] != "ID3"]
    curr_rows.append([len(curr_rows) + 1, "ID_NEW", 500.0])
    baseline, current = tmp_path / "baseline.xlsx", tmp_path / "current.xlsx"
    _write_panel(baseline, base_rows)
    _write_panel(current, curr_rows)
    profile = _profile_with_rule(
        RowIdentityRule(anchor_cell="B2", identity_columns=["B"], ordinal_columns=["A"])
    )

    result = run_qc(
        baseline_excel=baseline,
        current_excel=current,
        profile=profile,
        mode=QCRunMode.CYCLE_COMPARISON,
    )

    assert len(_classes(result, FindingClass.ROW_DELETED)) == 1
    assert len(_classes(result, FindingClass.ROW_INSERTED)) == 1
    assert _classes(result, FindingClass.VALUE_CHANGED) == []


def test_ordinal_column_suppresses_only_value_changed_not_style(
    tmp_path: Path,
) -> None:
    n = 30
    base_rows = _base_rows(n)
    curr_rows = _shuffled(base_rows, seed=1234)
    baseline, current = tmp_path / "baseline.xlsx", tmp_path / "current.xlsx"
    _write_panel(baseline, base_rows)
    _write_panel(current, curr_rows)

    workbook = load_workbook(current)
    sheet = workbook["Panel"]
    for row in sheet.iter_rows(min_row=2):
        if row[1].value == "ID0":
            row[0].font = Font(bold=True)
            break
    workbook.save(current)

    profile = _profile_with_rule(
        RowIdentityRule(anchor_cell="B2", identity_columns=["B"], ordinal_columns=["A"])
    )
    result = run_qc(
        baseline_excel=baseline,
        current_excel=current,
        profile=profile,
        mode=QCRunMode.CYCLE_COMPARISON,
    )

    # The ordinal column's cached-value comparison is suppressed...
    assert _classes(result, FindingClass.VALUE_CHANGED) == []
    # ...but its style is a presentation check and stays fully active.
    assert len(_classes(result, FindingClass.STYLE_CHANGED)) == 1


# --- detector wiring through the full engine ------------------------------


def test_unconfirmed_ranked_table_blocks_for_confirmation(tmp_path: Path) -> None:
    n = _LARGE_N
    base_rows = _detector_rows(n)
    curr_rows = _shuffled(base_rows, seed=99)
    baseline, current = tmp_path / "baseline.xlsx", tmp_path / "current.xlsx"
    headers = ("Rank", "ID", "Value", "Value2", "Value3")
    _write_panel(baseline, base_rows, headers=headers)
    _write_panel(current, curr_rows, headers=headers)

    with pytest.raises(RunBlockedError) as excinfo:
        run_qc(
            baseline_excel=baseline,
            current_excel=current,
            profile=None,
            mode=QCRunMode.CYCLE_COMPARISON,
        )

    action = excinfo.value.action_required
    assert action.version == 2
    assert action.reason is RunActionReason.ROW_IDENTITY_CONFIRMATION_REQUIRED
    assert action.items
    assert action.items[0].sheet == "Panel"
    assert action.items[0].cell == "A1"
    evidence = action.items[0].ranked_table_evidence
    assert evidence is not None
    assert evidence.version == 2
    assert evidence.member_id == "primary"
    assert evidence.sheet == "Panel"
    assert evidence.current_range == f"A1:E{n + 1}"
    assert evidence.data_row_count == n + 1
    assert evidence.available_columns == ("A", "B", "C", "D", "E")
    assert evidence.suggested_identity_columns == ("B",)
    assert evidence.suggested_ordinal_columns == ("A",)
    assert 0.0 <= evidence.non_blank_coverage <= 1.0
    assert 0.0 <= evidence.mismatch_reduction <= 1.0
    assert evidence.projected_positional_mismatches > 0
    assert evidence.projected_avoided_mismatches > 0
    # Legacy flat fields are no longer populated by this producer; a
    # consumer that only understands v1 items must fall back gracefully
    # instead of crashing (checked directly by the consumer-side tests).
    assert action.items[0].suggested_identity_columns == ()
    assert action.items[0].suggested_ordinal_columns == ()
    assert action.items[0].detail == ""
    # Bounded, aggregate evidence (ratios/counts) is expected and fine; the
    # actual identity text must never leak.
    serialized = action.model_dump_json()
    assert "ID0" not in serialized
    assert "ID1" not in serialized


def test_confirmed_identity_never_re_triggers_the_detector(tmp_path: Path) -> None:
    """The same data that blocks unconfirmed must run cleanly once the
    analyst confirms the exact rule the block suggested."""
    n = _LARGE_N
    base_rows = _detector_rows(n)
    curr_rows = _shuffled(base_rows, seed=99)
    baseline, current = tmp_path / "baseline.xlsx", tmp_path / "current.xlsx"
    headers = ("Rank", "ID", "Value", "Value2", "Value3")
    _write_panel(baseline, base_rows, headers=headers)
    _write_panel(current, curr_rows, headers=headers)
    profile = _profile_with_rule(
        RowIdentityRule(anchor_cell="B2", identity_columns=["B"], ordinal_columns=["A"])
    )

    result = run_qc(
        baseline_excel=baseline,
        current_excel=current,
        profile=profile,
        mode=QCRunMode.CYCLE_COMPARISON,
    )

    assert _classes(result, FindingClass.VALUE_CHANGED) == []


def test_perform_run_writes_no_history_or_reports_when_blocked_by_detector(
    tmp_path: Path,
) -> None:
    from qc_tool.history.store import RunHistory
    from qc_tool.run_service import perform_run

    n = _LARGE_N
    base_rows = _detector_rows(n)
    curr_rows = _shuffled(base_rows, seed=99)
    baseline, current = tmp_path / "baseline.xlsx", tmp_path / "current.xlsx"
    headers = ("Rank", "ID", "Value", "Value2", "Value3")
    _write_panel(baseline, base_rows, headers=headers)
    _write_panel(current, curr_rows, headers=headers)
    work_dir = tmp_path / "work"
    work_dir.mkdir()

    with pytest.raises(RunBlockedError):
        perform_run(
            work_dir,
            {"baseline_excel": baseline, "current_excel": current},
            {},
            default_profile(),
            mode=QCRunMode.CYCLE_COMPARISON,
        )

    history_db = work_dir / "history.sqlite3"
    if history_db.exists():
        assert RunHistory(history_db).list_runs() == []
    runs_dir = work_dir / "runs"
    assert not runs_dir.exists() or not any(runs_dir.iterdir())


def test_ranked_package_action_cap_represents_every_blocked_member() -> None:
    actions = [
        (
            f"member{member_index}",
            RunActionRequired(
                version=2,
                reason=RunActionReason.ROW_IDENTITY_CONFIRMATION_REQUIRED,
                items=[
                    RunActionItem(sheet="Panel", cell=f"A{item_index + 1}")
                    for item_index in range(20)
                ],
            ),
        )
        for member_index in range(8)
    ]

    action = _aggregate_ranked_package_actions(actions)

    assert len(action.items) == MAX_RUN_ACTION_ITEMS
    assert {item.member_id for item in action.items} == {
        f"member{member_index}" for member_index in range(8)
    }
    assert action.omitted_items == 144


def test_multi_package_reports_and_resolves_all_ranked_members_in_one_cycle(
    tmp_path: Path,
) -> None:
    members: list[PackageMember] = []
    files: dict[str, Path] = {}
    headers = ("Rank", "ID", "Value", "Value2", "Value3")
    for member_index, member_id in enumerate(("ops", "primary")):
        base_rows = _detector_rows(_LARGE_N)
        for row in base_rows:
            row[1] = f"M{member_index}-{row[1]}"
        current_rows = _shuffled(base_rows, seed=99 + member_index)
        baseline = tmp_path / f"baseline-{member_id}.xlsx"
        current = tmp_path / f"current-{member_id}.xlsx"
        _write_panel(baseline, base_rows, headers=headers)
        _write_panel(current, current_rows, headers=headers)
        for side, path in (
            (PackageSide.BASELINE, baseline),
            (PackageSide.CURRENT, current),
        ):
            member = PackageMember(
                member_id=member_id,
                side=side,
                artifact=PackageArtifact.EXCEL,
                display_name=path.name,
            )
            members.append(member)
            files[member.role_key] = path
    manifest = PackageManifest(members=tuple(members))

    with pytest.raises(RunBlockedError) as excinfo:
        run_qc(
            package_manifest=manifest,
            package_files=files,
            mode=QCRunMode.CYCLE_COMPARISON,
        )

    action = excinfo.value.action_required
    assert action.reason is RunActionReason.ROW_IDENTITY_CONFIRMATION_REQUIRED
    assert {item.member_id for item in action.items} == {"ops", "primary"}
    assert {
        item.ranked_table_evidence.member_id
        for item in action.items
        if item.ranked_table_evidence is not None
    } == {"ops", "primary"}
    view_model = view_model_from_action(
        action.model_dump(mode="json"),
        source_profile="default",
    )
    assert view_model is not None
    view_model = view_model.with_profile_name("ranked-package")
    assert view_model.is_valid
    profile = apply_view_model(
        default_profile(),
        view_model,
        workbook_count=2,
    ).model_copy(update={"name": "ranked-package"})

    result = run_qc(
        package_manifest=manifest,
        package_files=files,
        profile=profile,
        mode=QCRunMode.CYCLE_COMPARISON,
    )

    assert _classes(result, FindingClass.VALUE_CHANGED) == []


# --- V1/V2 alignment-trust manifest ----------------------------------------


def test_ordinary_run_keeps_writing_v1_alignment_trust(tmp_path: Path) -> None:
    from qc_tool.history.store import RunHistory
    from qc_tool.run_service import perform_run

    baseline, current = tmp_path / "baseline.xlsx", tmp_path / "current.xlsx"
    _write_panel(baseline, [["Header"], ["A"], ["B"]], headers=("Col",))
    _write_panel(current, [["Header"], ["A"], ["B2"]], headers=("Col",))
    work_dir = tmp_path / "work"
    work_dir.mkdir()

    artifacts = perform_run(
        work_dir,
        {"baseline_excel": baseline, "current_excel": current},
        {},
        default_profile(),
        mode=QCRunMode.CYCLE_COMPARISON,
    )

    assert isinstance(artifacts.result.alignment_trust, AlignmentTrustManifest)
    assert not isinstance(artifacts.result.alignment_trust, AlignmentTrustManifestV2)
    record = RunHistory(work_dir / "history.sqlite3").get_run(artifacts.run_id)
    assert isinstance(record.alignment_trust, AlignmentTrustManifest)
    assert not isinstance(record.alignment_trust, AlignmentTrustManifestV2)


def test_confirmed_identity_run_writes_v2_alignment_trust_round_trip(
    tmp_path: Path,
) -> None:
    from qc_tool.history.store import RunHistory
    from qc_tool.run_service import perform_run

    n = 30
    base_rows = _base_rows(n)
    curr_rows = _shuffled(base_rows, seed=1234)
    baseline, current = tmp_path / "baseline.xlsx", tmp_path / "current.xlsx"
    _write_panel(baseline, base_rows)
    _write_panel(current, curr_rows)
    profile = _profile_with_rule(
        RowIdentityRule(anchor_cell="B2", identity_columns=["B"], ordinal_columns=["A"])
    )
    work_dir = tmp_path / "work"
    work_dir.mkdir()

    artifacts = perform_run(
        work_dir,
        {"baseline_excel": baseline, "current_excel": current},
        {},
        profile,
        mode=QCRunMode.CYCLE_COMPARISON,
    )

    assert isinstance(artifacts.result.alignment_trust, AlignmentTrustManifestV2)
    record = RunHistory(work_dir / "history.sqlite3").get_run(artifacts.run_id)
    assert isinstance(record.alignment_trust, AlignmentTrustManifestV2)
    region = record.alignment_trust.regions[0]
    assert region.identity_columns == ("B",)
    assert region.ordinal_columns == ("A",)
    assert region.duplicate_policy == "skip"
