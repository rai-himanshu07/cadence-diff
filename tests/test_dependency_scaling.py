"""Plan 2026-08-13 step 1: whole-column twin contracts and oracle parity.

These contracts freeze the observable behavior of the dependency index on the
whole-column-aggregate shape BEFORE the scaling rework, so steps 2-3 can only
change cost, never answers. Reported occurrence counts (coverage inputs) are
pinned as invariants.
"""

from __future__ import annotations

import time

import pytest
from openpyxl.utils import get_column_letter

from qc_tool.excel.dependency import (
    DependencyGraph,
    FormulaCycleGraph,
    build_dependency_graph,
    dependent_nodes_of,
    dependents_of,
    detect_circular_references,
)
from qc_tool.excel.references import ResolvedRange
from tests.dependency_oracle import build_oracle_graph
from tests.fixtures.whole_column import (
    CALC_SHEET,
    DATA_SHEET,
    equality_twin,
    pathology_twin,
)

_SYMBOLIC_FORCE = 100  # equality twin columns hold 120 cells -> symbolic


def _oracle_direct(oracle, node) -> set[tuple[str, int, int]]:
    direct = set(oracle.successors(node)) if oracle.has_node(node) else set()
    for resolved, dependent in oracle.graph["symbolic"]:
        if (
            node[0] == resolved.sheet
            and resolved.min_row <= node[1] <= resolved.max_row
            and resolved.min_col <= node[2] <= resolved.max_col
        ):
            direct.add(dependent)
    return direct


def test_direct_dependents_match_the_oracle_on_the_equality_twin() -> None:
    twin = equality_twin()
    graph = build_dependency_graph(twin, max_range_cells=_SYMBOLIC_FORCE)
    oracle = build_oracle_graph(twin, max_range_cells=_SYMBOLIC_FORCE)

    probes = [
        (DATA_SHEET, 1, 1),  # $A$1: direct edge from every formula
        (DATA_SHEET, 17, 1),  # inside whole column A only
        (DATA_SHEET, 17, 2),  # inside whole column B and the literal range
        (DATA_SHEET, 90, 2),  # inside whole column B, past the literal range
        (DATA_SHEET, 17, 3),  # inside whole column C only
        (CALC_SHEET, 1, 1),  # formula cell with no dependents
    ]
    for node in probes:
        assert graph.direct_dependents(node) == _oracle_direct(oracle, node), node


def test_transitive_dependents_match_the_oracle_display_strings() -> None:
    twin = equality_twin()
    graph = build_dependency_graph(twin, max_range_cells=_SYMBOLIC_FORCE)
    oracle = build_oracle_graph(twin, max_range_cells=_SYMBOLIC_FORCE)

    for row, column in ((1, 1), (17, 2), (90, 3)):
        ref = f"{get_column_letter(column)}{row}"
        expected = sorted(
            f"{sheet}!{get_column_letter(col)}{r}"
            for sheet, r, col in _oracle_direct(oracle, (DATA_SHEET, row, column))
        )
        # The twin has no formula-to-formula chains, so transitive == direct.
        assert dependents_of(graph, DATA_SHEET, ref) == expected
        assert dependent_nodes_of(graph, (DATA_SHEET, row, column)) == set(
            _oracle_direct(oracle, (DATA_SHEET, row, column))
        )


def test_pathology_twin_pins_reported_population_counts() -> None:
    graph = build_dependency_graph(pathology_twin())

    # 2000 formulas x (Calc!$A$1 criteria + Data!$A$1) minus Calc!A1's self-ref.
    assert graph.direct_edge_count == 3999
    assert graph.range_descriptor_count == 2000
    assert len(graph.symbolic_references) == 2000
    assert graph.bounded_reference_count == 2000
    assert graph.formula_cycles.node_count == 2000
    # Calc!$A$1 is itself a formula node, so every criteria ref is a cycle edge.
    assert graph.formula_cycles.edge_count == 2000
    assert graph.formula_cycles.omitted_range_references == 0


def test_pathology_twin_query_answers_are_stable() -> None:
    graph = build_dependency_graph(pathology_twin())

    # A data cell inside every whole column's bounded extent and the literal
    # range: its dependents are exactly the formulas referencing that column
    # or the literal range, never the whole calc population.
    inside_b = graph.direct_dependents((DATA_SHEET, 25_000, 2))
    assert len(inside_b) == 2000  # 400 SUMIF($B:$B) + all 2000 literal SUM($B$…)
    inside_e = graph.direct_dependents((DATA_SHEET, 25_000, 5))
    assert len(inside_e) == 400  # only the SUMIF($E:$E) rotation
    beyond_extent = graph.direct_dependents((DATA_SHEET, 60_000, 5))
    assert beyond_extent == set()


def test_pathology_twin_build_and_query_stay_inside_ci_budget() -> None:
    started = time.perf_counter()
    graph = build_dependency_graph(pathology_twin())
    build_seconds = time.perf_counter() - started

    started = time.perf_counter()
    for row in range(1_000, 11_000, 10):
        graph.direct_dependents((DATA_SHEET, row, 3))
    query_seconds = time.perf_counter() - started

    # Generous ceilings: catch order-of-magnitude regressions, not jitter.
    assert build_seconds < 30.0
    assert query_seconds < 20.0


def test_cycle_graph_memoizes_repeated_rectangles() -> None:
    graph = build_dependency_graph(pathology_twin())

    # 2000 formulas x 4 references each collapse to 8 unique rectangles.
    assert graph.formula_cycles.candidate_walks == 8
    assert graph.formula_cycles.candidate_hits == 7_992


def test_cycle_candidate_cache_bounds_empty_entries() -> None:
    dependent = ("Formula", 1, 1)
    graph = FormulaCycleGraph(
        [dependent],
        max_candidate_cache_entries=3,
        max_candidate_cache_memberships=0,
    )

    for row in range(1, 6):
        graph.add_range(ResolvedRange("Empty", row, 1, row, 1), dependent)

    assert len(graph._candidate_cache) == 3
    assert graph._candidate_cache_memberships == 0
    assert graph.candidate_evictions == 2
    assert graph.candidate_admission_skips == 0


def test_cycle_candidate_cache_zero_entry_limit_disables_admission() -> None:
    nodes = [("S", row, 1) for row in range(1, 3)]
    graph = FormulaCycleGraph(
        nodes,
        max_candidate_cache_entries=0,
        max_candidate_cache_memberships=100,
    )
    resolved = ResolvedRange("S", 1, 1, 2, 1)

    graph.add_range(resolved, nodes[0])
    graph.add_range(resolved, nodes[0])

    assert graph._candidate_cache == {}
    assert graph._candidate_cache_memberships == 0
    assert graph.candidate_admission_skips == 2
    assert graph.candidate_walks == 2
    assert graph.candidate_hits == 0
    assert graph.edge_count == 2


@pytest.mark.parametrize(
    ("entry_limit", "membership_limit"),
    [(-1, 0), (0, -1)],
)
def test_cycle_candidate_cache_rejects_negative_limits(
    entry_limit: int,
    membership_limit: int,
) -> None:
    with pytest.raises(ValueError, match="cache limits must be non-negative"):
        FormulaCycleGraph(
            [("S", 1, 1)],
            max_candidate_cache_entries=entry_limit,
            max_candidate_cache_memberships=membership_limit,
        )


def test_cycle_candidate_cache_bounds_retained_memberships() -> None:
    nodes = [("S", row, 1) for row in range(1, 7)]
    graph = FormulaCycleGraph(
        nodes,
        max_candidate_cache_entries=10,
        max_candidate_cache_memberships=12,
    )

    for max_row in range(6, 11):
        graph.add_range(ResolvedRange("S", 1, 1, max_row, 1), nodes[0])

    assert len(graph._candidate_cache) == 2
    assert graph._candidate_cache_memberships == 12
    assert graph.candidate_evictions == 3
    assert graph.candidate_admission_skips == 0
    assert graph.edge_count == 6


def test_cycle_candidate_cache_does_not_admit_one_oversized_entry() -> None:
    nodes = [("S", row, 1) for row in range(1, 5)]
    graph = FormulaCycleGraph(
        nodes,
        max_candidate_cache_entries=10,
        max_candidate_cache_memberships=3,
    )
    resolved = ResolvedRange("S", 1, 1, 4, 1)

    graph.add_range(resolved, nodes[0])
    graph.add_range(resolved, nodes[0])

    assert graph._candidate_cache == {}
    assert graph._candidate_cache_memberships == 0
    assert graph.candidate_admission_skips == 2
    assert graph.candidate_walks == 2
    assert graph.candidate_hits == 0
    assert graph.edge_count == 4
    assert graph.strongly_connected_components() == ((nodes[0],),)


def test_cycle_candidate_cache_eviction_preserves_graph_and_coverage() -> None:
    nodes = [("S", row, 1) for row in range(1, 5)]
    unbounded = FormulaCycleGraph(
        nodes,
        max_edges=5,
        max_candidate_cache_entries=100,
        max_candidate_cache_memberships=100,
    )
    constrained = FormulaCycleGraph(
        nodes,
        max_edges=5,
        max_candidate_cache_entries=1,
        max_candidate_cache_memberships=4,
    )
    operations = (
        (ResolvedRange("S", 1, 1, 2, 1), nodes[2]),
        (ResolvedRange("S", 2, 1, 4, 1), nodes[0]),
        (ResolvedRange("S", 1, 1, 4, 1), nodes[1]),
        (ResolvedRange("S", 1, 1, 2, 1), nodes[3]),
        (ResolvedRange("S", 1, 1, 4, 1), nodes[1]),
    )
    for resolved, dependent in operations:
        unbounded.add_range(resolved, dependent)
        constrained.add_range(resolved, dependent)

    assert constrained._adjacency == unbounded._adjacency
    assert constrained.edge_count == unbounded.edge_count
    assert constrained.omitted_range_references == unbounded.omitted_range_references
    assert (
        constrained.strongly_connected_components()
        == unbounded.strongly_connected_components()
    )
    unbounded_graph = DependencyGraph()
    unbounded_graph.formula_cycles = unbounded
    constrained_graph = DependencyGraph()
    constrained_graph.formula_cycles = constrained
    assert detect_circular_references(constrained_graph) == detect_circular_references(
        unbounded_graph
    )


def test_repeated_rectangles_are_stored_once_with_grouped_dependents() -> None:
    graph = build_dependency_graph(pathology_twin())

    # 5 unique whole-column rectangles serve all 2000 symbolic occurrences,
    # and the one literal rectangle groups its 2000 dependents while the
    # logical descriptor count keeps its per-reference meaning.
    assert len(graph._symbolic_rect_bounds) == 5
    assert len(graph._rect_bounds) == 1
    assert len(graph._rect_dependents[0]) == 2000
    assert graph.range_descriptor_count == 2000


def test_impact_accumulator_matches_annotate_then_limit_exactly() -> None:
    from qc_tool.excel.dependency import (
        ImpactAccumulator,
        annotate_impacts,
        limit_impacts,
    )
    from qc_tool.findings import Finding, FindingClass

    twin = equality_twin()

    def fresh_findings() -> list:
        rows = [
            # >25 dependents through whole columns: cap + sentinel path.
            Finding(
                artifact="excel",
                finding_class=FindingClass.VALUE_CHANGED,
                sheet=DATA_SHEET,
                location="A17",
                message="probe",
            ),
            # No dependents at all: plain extras path.
            Finding(
                artifact="excel",
                finding_class=FindingClass.VALUE_CHANGED,
                sheet=CALC_SHEET,
                location="H120",
                message="probe",
            ),
            # Non-annotated class keeps extras-only handling.
            Finding(
                artifact="excel",
                finding_class=FindingClass.SHEET_ADDED,
                sheet=DATA_SHEET,
                location="B2",
                message="probe",
            ),
        ]
        # Chart-style extras, one duplicating a real dependent display.
        rows[0].impacts = ["Excel chart 'X' series 'S'", f"{CALC_SHEET}!A1"]
        rows[1].impacts = [f"chart {i}" for i in range(30)]
        rows[2].impacts = ["zeta", "alpha", "alpha"]
        return rows

    graph = build_dependency_graph(twin, max_range_cells=_SYMBOLIC_FORCE)
    classic = fresh_findings()
    annotate_impacts(classic, graph)
    limit_impacts(classic)

    graph2 = build_dependency_graph(twin, max_range_cells=_SYMBOLIC_FORCE)
    accumulated = fresh_findings()
    accumulator = ImpactAccumulator(graph2)
    accumulator.annotate(accumulated)
    accumulator.finalize(accumulated)

    for old, new in zip(classic, accumulated, strict=True):
        assert new.impacts == old.impacts, (old.sheet, old.location)
