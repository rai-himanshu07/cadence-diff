"""Sheet-scoped defined-name extraction, comparison, coverage, and redaction."""

from __future__ import annotations

import io
import zipfile
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.workbook.defined_name import DefinedName

from qc_tool.config.profile import DeliverableProfile
from qc_tool.coverage import CoverageState
from qc_tool.excel.align import align_workbooks
from qc_tool.excel.diff_structure import diff_workbook_structure
from qc_tool.excel.preflight import defined_name_scope_coverage, preflight_workbook
from qc_tool.excel.references import build_reference_index
from qc_tool.findings import FindingClass
from qc_tool.io.loader import load_workbook_snapshot
from qc_tool.io.model import NamedRange, WorkbookSnapshot
from qc_tool.io.ooxml_names import scan_defined_names
from qc_tool.sanitize import sanitize_workbook


def _build(
    path: Path,
    *,
    global_names: dict[str, str] | None = None,
    local_names: dict[str, dict[str, str]] | None = None,
) -> Path:
    workbook = Workbook()
    data = workbook.active
    assert data is not None
    data.title = "Data"
    summary = workbook.create_sheet("Summary")
    for row in range(1, 6):
        data.cell(row=row, column=1, value=f"key-{row}")
        data.cell(row=row, column=2, value=row * 10)
        summary.cell(row=row, column=3, value=row * 2)
        summary.cell(row=row, column=4, value=f"label-{row}")
    for name, target in (global_names or {}).items():
        workbook.defined_names[name] = DefinedName(name, attr_text=target)
    for sheet_name, names in (local_names or {}).items():
        sheet = workbook[sheet_name]
        for name, target in names.items():
            sheet.defined_names[name] = DefinedName(name, attr_text=target)
    workbook.save(path)
    return path


def _scoped_pair(path: Path) -> Path:
    return _build(
        path,
        global_names={"SharedName": "Data!$A$1:$A$5"},
        local_names={
            "Data": {"SharedName": "Data!$B$1:$B$5"},
            "Summary": {"SharedName": "Summary!$C$1:$C$5", "SummaryOnly": "Summary!$D$1"},
        },
    )


def _profile() -> DeliverableProfile:
    return DeliverableProfile(name="defined-name-tests")


def _workbook_part(data: bytes) -> str:
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        return archive.read("xl/workbook.xml").decode()


def _rewrite_workbook_part(data: bytes, old: str, new: str) -> bytes:
    """Replace text inside ``xl/workbook.xml``; the package is a zip, not raw XML."""
    buffer = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as source:
        assert old in source.read("xl/workbook.xml").decode()
        with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as dest:
            for item in source.infolist():
                payload = source.read(item.filename)
                if item.filename == "xl/workbook.xml":
                    payload = payload.decode().replace(old, new).encode()
                dest.writestr(item, payload)
    return buffer.getvalue()


# --- raw package scan ----------------------------------------------------


def test_scan_separates_workbook_and_sheet_scope(tmp_path: Path) -> None:
    scan = scan_defined_names(_scoped_pair(tmp_path / "book.xlsx").read_bytes())

    assert scan.available is True
    assert [(item.name, item.target) for item in scan.workbook_scoped] == [
        ("SharedName", "Data!$A$1:$A$5")
    ]
    assert [(item.sheet, item.name, item.target) for item in scan.sheet_scoped] == [
        ("Data", "SharedName", "Data!$B$1:$B$5"),
        ("Summary", "SharedName", "Summary!$C$1:$C$5"),
        ("Summary", "SummaryOnly", "Summary!$D$1"),
    ]


def test_scan_workbook_scope_agrees_with_openpyxl(tmp_path: Path) -> None:
    """The raw scan must be equivalent to today's producer for workbook scope."""
    path = _scoped_pair(tmp_path / "book.xlsx")
    scan = scan_defined_names(path.read_bytes())
    reference = load_workbook(path, data_only=False)

    assert {(item.name, item.target) for item in scan.workbook_scoped} == {
        (name, str(defined.attr_text))
        for name, defined in reference.defined_names.items()
        if not name.startswith("_xlnm")
    }


def test_scan_skips_builtin_print_names(tmp_path: Path) -> None:
    path = tmp_path / "book.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.print_area = "A1:B2"
    workbook.save(path)

    scan = scan_defined_names(path.read_bytes())

    assert scan.builtin_skipped >= 1
    assert scan.sheet_scoped == ()
    assert scan.workbook_scoped == ()


def test_scan_reports_unavailable_for_non_ooxml_bytes() -> None:
    scan = scan_defined_names(b"not-a-package")

    assert scan.available is False
    assert scan.detail
    assert scan.sheet_scoped == ()


def test_scan_refuses_out_of_range_scope_instead_of_guessing(tmp_path: Path) -> None:
    path = _build(tmp_path / "book.xlsx", local_names={"Summary": {"Local": "Summary!$C$1"}})
    broken = _rewrite_workbook_part(
        path.read_bytes(), 'localSheetId="1"', 'localSheetId="97"'
    )

    scan = scan_defined_names(broken)

    assert scan.available is True
    assert scan.sheet_scoped == ()
    assert scan.unresolved_scopes == 1
    assert "unreadable scope" in scan.detail


def test_scan_omits_macro_bound_names(tmp_path: Path) -> None:
    path = _build(tmp_path / "book.xlsx", local_names={"Summary": {"Local": "Summary!$C$1"}})
    broken = _rewrite_workbook_part(
        path.read_bytes(), 'name="Local"', 'name="Local" function="1"'
    )

    scan = scan_defined_names(broken)

    assert scan.sheet_scoped == ()
    assert scan.macro_skipped == 1
    assert "macro-bound" in scan.detail


# --- identity ------------------------------------------------------------


def test_named_range_identity_stays_distinct_across_scopes() -> None:
    workbook_scoped = NamedRange("SharedName", "Data!$A$1")
    data_scoped = NamedRange("SharedName", "Data!$B$1", sheet="Data")
    summary_scoped = NamedRange("SharedName", "Summary!$C$1", sheet="Summary")

    identities = {
        item.qualified_name for item in (workbook_scoped, data_scoped, summary_scoped)
    }

    assert identities == {"SharedName", "Data!SharedName", "Summary!SharedName"}
    assert workbook_scoped.scope_label == "workbook"
    assert summary_scoped.scope_label == "Summary"


def test_legacy_positional_named_range_construction_still_works() -> None:
    legacy = NamedRange("Revenue", "Data!A1")

    assert legacy.sheet is None
    assert legacy.hidden is False
    assert legacy.qualified_name == "Revenue"


# --- loader --------------------------------------------------------------


@pytest.mark.parametrize("loader", ["streaming", "oracle"])
def test_loader_captures_all_three_scopes(tmp_path: Path, loader: str) -> None:
    path = _scoped_pair(tmp_path / "book.xlsx")

    snapshot = load_workbook_snapshot(path, _ooxml_loader=loader)  # type: ignore[arg-type]

    assert snapshot.defined_name_scope_available is True
    assert {item.qualified_name for item in snapshot.named_ranges} == {
        "SharedName",
        "Data!SharedName",
        "Summary!SharedName",
        "Summary!SummaryOnly",
    }


def test_workbook_without_local_names_is_unchanged(tmp_path: Path) -> None:
    path = _build(tmp_path / "book.xlsx", global_names={"Only": "Data!$A$1:$A$5"})

    snapshot = load_workbook_snapshot(path)

    assert [(item.name, item.target, item.sheet) for item in snapshot.named_ranges] == [
        ("Only", "Data!$A$1:$A$5", None)
    ]
    assert snapshot.defined_name_scope_available is True
    assert snapshot.defined_name_scope_detail == ""


def test_reference_index_ignores_sheet_local_names(tmp_path: Path) -> None:
    """Resolution must not move: only workbook-scoped names answer a lookup."""
    snapshot = load_workbook_snapshot(_scoped_pair(tmp_path / "book.xlsx"))

    index = build_reference_index(snapshot)

    resolved = index.named_ranges["sharedname"]
    assert len(resolved) == 1
    assert resolved[0].sheet is None
    assert resolved[0].target == "Data!$A$1:$A$5"
    assert "summaryonly" not in index.named_ranges


# --- comparison ----------------------------------------------------------


def _structure_findings(
    baseline: WorkbookSnapshot, current: WorkbookSnapshot
) -> list[str]:
    profile = _profile()
    alignment = align_workbooks(baseline, current, profile)
    return [
        finding.element or ""
        for finding in diff_workbook_structure(baseline, current, alignment, profile)
        if finding.finding_class is FindingClass.NAMED_RANGE_CHANGED
    ]


def test_sheet_scoped_repoint_is_reported_against_its_own_scope(tmp_path: Path) -> None:
    baseline = load_workbook_snapshot(_scoped_pair(tmp_path / "base.xlsx"))
    current = load_workbook_snapshot(
        _build(
            tmp_path / "curr.xlsx",
            global_names={"SharedName": "Data!$A$1:$A$5"},
            local_names={
                "Data": {"SharedName": "Data!$B$1:$B$5"},
                "Summary": {
                    "SharedName": "Summary!$D$1:$D$5",
                    "SummaryOnly": "Summary!$D$1",
                },
            },
        )
    )

    assert _structure_findings(baseline, current) == ["Summary!SharedName"]


def test_identical_name_in_two_scopes_does_not_cross_contaminate(tmp_path: Path) -> None:
    baseline = load_workbook_snapshot(
        _build(
            tmp_path / "base.xlsx",
            global_names={"SharedName": "Data!$A$1"},
            local_names={"Summary": {"SharedName": "Summary!$C$1"}},
        )
    )
    current = load_workbook_snapshot(
        _build(
            tmp_path / "curr.xlsx",
            global_names={"SharedName": "Summary!$C$1"},
            local_names={"Summary": {"SharedName": "Data!$A$1"}},
        )
    )

    assert sorted(_structure_findings(baseline, current)) == [
        "SharedName",
        "Summary!SharedName",
    ]


def test_workbook_scoped_finding_message_and_sheet_are_unchanged(tmp_path: Path) -> None:
    baseline = load_workbook_snapshot(
        _build(tmp_path / "base.xlsx", global_names={"Only": "Data!$A$1"})
    )
    current = load_workbook_snapshot(
        _build(tmp_path / "curr.xlsx", global_names={"Only": "Data!$A$2"})
    )
    profile = _profile()
    alignment = align_workbooks(baseline, current, profile)

    findings = [
        finding
        for finding in diff_workbook_structure(baseline, current, alignment, profile)
        if finding.finding_class is FindingClass.NAMED_RANGE_CHANGED
    ]

    assert len(findings) == 1
    assert findings[0].sheet is None
    assert findings[0].element == "Only"
    assert findings[0].message == "named range 'Only' repointed"


def test_invalid_sheet_scoped_target_is_reported_with_its_scope(tmp_path: Path) -> None:
    snapshot = load_workbook_snapshot(
        _build(tmp_path / "book.xlsx", local_names={"Summary": {"Broken": "Ghost!$A$1"}})
    )

    findings = preflight_workbook(snapshot, _profile()).findings
    invalid = [
        finding
        for finding in findings
        if finding.finding_class is FindingClass.NAMED_RANGE_INVALID
    ]

    assert [finding.element for finding in invalid] == ["Summary!Broken"]
    assert invalid[0].sheet == "Summary"


# --- coverage ------------------------------------------------------------


def test_coverage_is_checked_when_scope_is_readable(tmp_path: Path) -> None:
    snapshot = load_workbook_snapshot(_scoped_pair(tmp_path / "book.xlsx"))

    item = defined_name_scope_coverage(snapshot)

    assert item.check_id == "excel-defined-name-scope"
    assert item.state is CoverageState.CHECKED
    assert item.findings == 3


def test_coverage_is_unavailable_when_scope_cannot_be_read() -> None:
    blocked = WorkbookSnapshot(
        source_name="legacy.xlsb",
        file_format="xlsb",
        formulas_available=False,
        styles_available=False,
        defined_name_scope_available=False,
        defined_name_scope_detail="binary workbook part",
    )

    item = defined_name_scope_coverage(blocked)

    assert item.state is CoverageState.UNAVAILABLE
    assert "binary workbook part" in item.detail


def test_coverage_is_degraded_when_some_scopes_are_unreadable() -> None:
    partial = WorkbookSnapshot(
        source_name="partial.xlsx",
        file_format="xlsx",
        formulas_available=True,
        styles_available=True,
        defined_name_scope_available=True,
        defined_name_scope_detail="1 defined names declare an unreadable scope",
    )

    item = defined_name_scope_coverage(partial)

    assert item.state is CoverageState.DEGRADED
    assert "unreadable scope" in item.detail


def test_preflight_registers_the_scope_coverage_item(tmp_path: Path) -> None:
    snapshot = load_workbook_snapshot(_scoped_pair(tmp_path / "book.xlsx"))

    result = preflight_workbook(snapshot, _profile())

    assert any(item.check_id == "excel-defined-name-scope" for item in result.coverage)


# --- privacy -------------------------------------------------------------


def test_strict_redaction_removes_sheet_scoped_names(tmp_path: Path) -> None:
    source = _scoped_pair(tmp_path / "book.xlsx")
    dest = tmp_path / "redacted.xlsx"

    sanitize_workbook(source, dest, redact_text=True)

    redacted = load_workbook(io.BytesIO(dest.read_bytes()), data_only=False)
    assert dict(redacted.defined_names) == {}
    assert all(not sheet.defined_names for sheet in redacted.worksheets)
    assert "SummaryOnly" not in _workbook_part(dest.read_bytes())
