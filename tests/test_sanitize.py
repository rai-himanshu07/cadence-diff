"""Sanitizer tests: determinism, structure preservation, safety."""

import hashlib
from pathlib import Path
from typing import Any, cast

import pytest
from openpyxl import load_workbook
from pptx import Presentation

from qc_tool.privacy import verify_sanitized
from qc_tool.sanitize import SanitizeError, sanitize_file


def test_workbook_sanitize_scrambles_numbers_keeps_structure(
    fixture_dir: Path, tmp_path: Path
) -> None:
    source = fixture_dir / "current.xlsx"
    before = hashlib.sha256(source.read_bytes()).hexdigest()
    output, stats = sanitize_file(source, tmp_path / "out.xlsx", seed=7)

    assert hashlib.sha256(source.read_bytes()).hexdigest() == before  # read-only
    assert stats.numbers_scrambled > 0

    original = load_workbook(source)
    scrambled = load_workbook(output)
    orig_sheet = original["Long_Monthly"]
    new_sheet = scrambled["Long_Monthly"]

    assert new_sheet["C7"].value != orig_sheet["C7"].value  # constant scrambled
    orig_c7 = float(orig_sheet["C7"].value)
    assert abs(float(new_sheet["C7"].value)) == pytest.approx(
        abs(orig_c7), rel=0.16
    )  # magnitude preserved
    assert new_sheet["E14"].value == orig_sheet["E14"].value  # formula untouched
    assert new_sheet["A2"].value == "Jan-26"  # period labels kept
    assert new_sheet["B2"].value == orig_sheet["B2"].value  # text labels kept
    assert scrambled.sheetnames == original.sheetnames


def test_workbook_sanitize_is_deterministic(fixture_dir: Path, tmp_path: Path) -> None:
    source = fixture_dir / "current.xlsx"
    out_a, _ = sanitize_file(source, tmp_path / "a.xlsx", seed=42)
    out_b, _ = sanitize_file(source, tmp_path / "b.xlsx", seed=42)
    out_c, _ = sanitize_file(source, tmp_path / "c.xlsx", seed=43)

    value = lambda p: load_workbook(p)["Long_Monthly"]["C7"].value  # noqa: E731
    assert value(out_a) == value(out_b)
    assert value(out_a) != value(out_c)


def test_workbook_redact_text_keeps_periods(fixture_dir: Path, tmp_path: Path) -> None:
    output, stats = sanitize_file(
        fixture_dir / "current.xlsx", tmp_path / "r.xlsx", seed=1, redact_text=True
    )
    sheet = load_workbook(output)["Sheet_001"]
    assert sheet["A2"].value == "Jan-26"  # period kept
    assert str(sheet["B2"].value).startswith("TXT_")  # label redacted
    assert sheet.title.startswith("Sheet_")
    assert stats.texts_redacted > 0
    assert verify_sanitized(output).safe


def test_deck_sanitize(fixture_dir: Path, tmp_path: Path) -> None:
    source = fixture_dir / "current.pptx"
    output, stats = sanitize_file(source, tmp_path / "out.pptx", seed=5)
    assert stats.numbers_scrambled > 0
    assert stats.charts_rebuilt > 0

    def chart_values(path: Path) -> tuple[list[str], list[float]]:
        for slide in Presentation(str(path)).slides:
            title = slide.shapes.title
            if title is not None and title.text == "Revenue Trend":
                for shape in slide.shapes:
                    if shape.has_chart:
                        plot = cast(Any, shape).chart.plots[0]
                        return (
                            [str(c) for c in plot.categories],
                            [float(v) for v in plot.series[0].values],
                        )
        raise AssertionError("trend chart not found")

    orig_cats, orig_vals = chart_values(source)
    new_cats, new_vals = chart_values(output)
    assert new_cats == orig_cats  # categories (periods) kept
    assert new_vals != orig_vals  # values scrambled

    titles = [
        s.shapes.title.text
        for s in Presentation(str(output)).slides
        if s.shapes.title is not None
    ]
    assert "Executive Summary" in titles  # titles never touched


def test_unsupported_formats_rejected(fixture_dir: Path, tmp_path: Path) -> None:
    with pytest.raises(SanitizeError, match="re-save"):
        sanitize_file(fixture_dir / "current.xlsb", tmp_path / "x.xlsb")
    with pytest.raises(SanitizeError, match="differ from the source"):
        sanitize_file(fixture_dir / "current.xlsx", fixture_dir / "current.xlsx")
    with pytest.raises(SanitizeError, match="must match source extension"):
        sanitize_file(fixture_dir / "current.xlsx", tmp_path / "wrong.pptx")


def test_redaction_strips_client_metadata_and_mixed_titles(tmp_path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.comments import Comment

    workbook_path = tmp_path / "client.xlsx"
    workbook_output = tmp_path / "client.sanitized.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Client_Alpha_Revenue"
    sheet["A1"] = "Client Alpha"
    sheet["A2"] = "Jun-26"
    sheet["A3"] = 1
    sheet["A3"].comment = Comment("Call Jane at Client Alpha", "John Reviewer")
    sheet["A4"] = "https://client-alpha.example/data"
    sheet["A4"].hyperlink = sheet["A4"].value
    cast(Any, sheet.oddHeader).center.text = "Client Alpha confidential"
    workbook.properties.creator = "Jane Analyst"
    workbook.properties.subject = "Client Alpha confidential"
    workbook.save(workbook_path)

    sanitize_file(workbook_path, workbook_output, seed=2, redact_text=True)
    sanitized_workbook = load_workbook(workbook_output)
    sanitized_sheet = sanitized_workbook[sanitized_workbook.sheetnames[0]]
    assert sanitized_sheet.title == "Sheet_001"
    assert str(sanitized_sheet["A1"].value).startswith("TXT_")
    assert sanitized_sheet["A2"].value == "Jun-26"
    assert sanitized_sheet["A3"].value != 1
    assert sanitized_sheet["A3"].comment is None
    assert sanitized_sheet["A4"].hyperlink is None
    assert not cast(Any, sanitized_sheet.oddHeader).center.text
    assert not sanitized_workbook.properties.creator
    assert not sanitized_workbook.properties.subject

    deck_path = tmp_path / "client.pptx"
    deck_output = tmp_path / "client.sanitized.pptx"
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[0])
    title = slide.shapes.title
    assert title is not None
    title.text = "Client Alpha Q1 2026 — $45M"
    subtitle = cast(Any, slide.placeholders[1])
    subtitle.text = "Prepared for Client Alpha"
    deck.core_properties.author = "Jane Analyst"
    deck.core_properties.subject = "Client Alpha confidential"
    deck.save(str(deck_path))

    sanitize_file(deck_path, deck_output, seed=2, redact_text=True)
    sanitized_deck = Presentation(str(deck_output))
    sanitized_title = sanitized_deck.slides[0].shapes.title
    assert sanitized_title is not None
    assert sanitized_title.text.startswith("TXT_")
    assert "Q1 2026" in sanitized_title.text
    assert "Client Alpha" not in sanitized_title.text
    sanitized_subtitle = cast(Any, sanitized_deck.slides[0].placeholders[1])
    assert sanitized_subtitle.text.startswith("TXT_")
    assert not sanitized_deck.core_properties.author
    assert not sanitized_deck.core_properties.subject
    assert verify_sanitized(
        deck_output, forbidden_tokens=["Client Alpha", "Jane Analyst"]
    ).safe


def test_macro_enabled_workbook_is_refused(tmp_path: Path) -> None:
    source = tmp_path / "client.xlsm"
    source.write_bytes(b"not parsed because policy rejects first")
    with pytest.raises(SanitizeError, match="macro-enabled"):
        sanitize_file(source, tmp_path / "out.xlsm", redact_text=True)


def test_privacy_verifier_rejects_unsanitized_fixture(fixture_dir: Path) -> None:
    report = verify_sanitized(
        fixture_dir / "current.xlsx", forbidden_tokens=["Long_Monthly"]
    )
    assert not report.safe
    assert any(issue.code in {"sheet-name", "forbidden-token"} for issue in report.issues)
