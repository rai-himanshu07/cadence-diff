"""Complete dependency coverage, symbolic aggregates, and chart impacts."""

from pathlib import Path
from typing import Any, cast

import pytest
from openpyxl import Workbook
from pptx import Presentation
from pptx.chart.data import ChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches

from qc_tool.config.profile import (
    CrosscheckMapping,
    CrosscheckProfile,
    DeliverableProfile,
    default_profile,
)
from qc_tool.coverage import CoverageState
from qc_tool.crosscheck.trace import annotate_ppt_chart_impacts
from qc_tool.engine import run_qc
from qc_tool.excel.charts import (
    annotate_chart_impacts,
    chart_reference_coverage,
    diff_charts,
)
from qc_tool.excel.dependency import (
    annotate_impacts,
    build_dependency_graph,
    dependents_of,
    limit_impacts,
)
from qc_tool.excel.preflight import preflight_workbook
from qc_tool.excel.references import ReferenceStatus, resolve_reference
from qc_tool.findings import Finding, FindingClass
from qc_tool.io.model import (
    CellRecord,
    ChartDescriptor,
    ChartPlot,
    ChartSeries,
    FormulaRangeDescriptor,
    NamedRange,
    SheetSnapshot,
    WorkbookSnapshot,
)
from qc_tool.ppt.model import (
    ChartContent,
    DeckSnapshot,
    PptChartLabel,
    PptChartPlot,
    PptChartSeries,
    SlideContent,
)


def _sheet(
    *,
    max_row: int = 100,
    max_column: int = 10,
    cells: dict[tuple[int, int], CellRecord] | None = None,
) -> WorkbookSnapshot:
    return WorkbookSnapshot(
        "dependency.xlsx",
        "xlsx",
        True,
        True,
        sheets=[
            SheetSnapshot(
                "Data",
                "visible",
                max_row,
                max_column,
                cells or {},
            )
        ],
    )


def test_whole_column_and_row_references_are_bounded_to_sheet() -> None:
    workbook = _sheet(max_row=100, max_column=10)

    whole_column = resolve_reference(
        workbook,
        "Data!A:C",
        host_sheet="Data",
        require_within_sheet=False,
    )
    whole_row = resolve_reference(
        workbook,
        "Data!2:10",
        host_sheet="Data",
        require_within_sheet=False,
    )

    assert whole_column.status is ReferenceStatus.RESOLVED
    assert whole_column.ranges[0].min_row == 1
    assert whole_column.ranges[0].max_row == 100
    assert whole_column.ranges[0].min_col == 1
    assert whole_column.ranges[0].max_col == 3
    assert "bounded" in whole_column.detail
    assert whole_row.status is ReferenceStatus.RESOLVED
    assert whole_row.ranges[0].min_row == 2
    assert whole_row.ranges[0].max_row == 10
    assert whole_row.ranges[0].min_col == 1
    assert whole_row.ranges[0].max_col == 10


def test_bounded_whole_column_dependency_is_queryable() -> None:
    workbook = _sheet(
        cells={(1, 3): CellRecord(1, 3, None, formula="=SUM(A:A)")},
    )

    graph = build_dependency_graph(workbook)

    assert dependents_of(graph, "Data", "A75") == ["Data!C1"]
    assert graph.coverage_state is CoverageState.CHECKED
    assert "bounded whole-row/column" in graph.coverage_detail


def test_oversized_ranges_are_symbolic_and_transitive() -> None:
    workbook = _sheet(
        max_row=100_000,
        cells={
            (1, 2): CellRecord(1, 2, None, formula="=SUM(A:A)"),
            (1, 3): CellRecord(1, 3, None, formula="=SUM(B:B)"),
        },
    )

    graph = build_dependency_graph(workbook, max_range_cells=100)

    assert len(graph.symbolic_references) == 2
    assert dependents_of(graph, "Data", "A50000") == ["Data!B1", "Data!C1"]
    assert graph.coverage_state is CoverageState.CHECKED
    assert "2 symbolic aggregate" in graph.coverage_detail


def test_standalone_dependency_coverage_discloses_symbolic_aggregates() -> None:
    workbook = _sheet(
        max_row=100_000,
        cells={(1, 2): CellRecord(1, 2, None, formula="=SUM(A:A)")},
    )

    result = preflight_workbook(workbook, default_profile())
    coverage = next(
        item for item in result.coverage if item.check_id == "excel-dependencies"
    )

    assert coverage.state is CoverageState.CHECKED
    assert "symbolic aggregate" in coverage.detail


def test_unsupported_and_invalid_references_degrade_dependency_coverage() -> None:
    workbook = _sheet(
        cells={
            (1, 2): CellRecord(
                1,
                2,
                None,
                formula="=SUM([Book.xlsx]Data!A1)",
            ),
            (1, 3): CellRecord(1, 3, None, formula="=Missing!A1"),
        }
    )

    graph = build_dependency_graph(workbook)

    assert graph.coverage_state is CoverageState.DEGRADED
    assert len(graph.unsupported_references) == 1
    assert len(graph.invalid_references) == 1
    assert "unsupported" in graph.coverage_detail
    assert "invalid" in graph.coverage_detail


def test_invalid_reference_telemetry_is_occurrence_accurate_and_bounded() -> None:
    occurrence_count = 250
    workbook = _sheet(
        max_row=occurrence_count,
        cells={
            (row, 2): CellRecord(row, 2, None, formula="=MissingReference")
            for row in range(1, occurrence_count + 1)
        },
    )

    graph = build_dependency_graph(workbook)

    assert graph.reference_status_counts[ReferenceStatus.INVALID] == occurrence_count
    assert graph.invalid_reason_counts == {"invalid_a1_reference": occurrence_count}
    assert 0 < len(graph.invalid_references) <= 8
    assert set(graph.invalid_references) == {"invalid_a1_reference"}
    assert f"{occurrence_count} invalid references" in graph.coverage_detail


def test_repeated_formula_template_is_extracted_once_but_projected_per_host(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    import qc_tool.excel.dependency as dependency

    workbook = _sheet(
        max_row=3,
        cells={
            (1, 2): CellRecord(1, 2, None, formula="=@A1:A3"),
            (3, 2): CellRecord(3, 2, None, formula="=@a1:a3"),
        },
    )
    calls = 0
    original = dependency.extract_formula_precedents

    def counted(formula: str):
        nonlocal calls
        calls += 1
        return original(formula)

    monkeypatch.setattr(dependency, "extract_formula_precedents", counted)

    graph = build_dependency_graph(workbook)

    assert calls == 1
    assert dependents_of(graph, "Data", "A1") == ["Data!B1"]
    assert dependents_of(graph, "Data", "A3") == ["Data!B3"]


def test_dependency_metadata_index_resolves_sheet_names_case_insensitively() -> None:
    workbook = _sheet(
        cells={
            (1, 1): CellRecord(1, 1, 10),
            (1, 2): CellRecord(1, 2, None, formula="=data!a1"),
        },
    )

    graph = build_dependency_graph(workbook)

    assert dependents_of(graph, "Data", "A1") == ["Data!B1"]
    assert graph.coverage_state is CoverageState.CHECKED


def test_lexical_symbol_shadows_same_named_workbook_range() -> None:
    workbook = _sheet(
        cells={
            (1, 1): CellRecord(1, 1, 10),
            (1, 2): CellRecord(1, 2, 20),
            (1, 3): CellRecord(1, 3, 30),
            (1, 4): CellRecord(1, 4, None, formula="=LET(rate,B1,rate+C1)"),
        },
    )
    workbook.named_ranges.append(NamedRange("rate", "Data!A1"))

    graph = build_dependency_graph(workbook)

    assert dependents_of(graph, "Data", "A1") == []
    assert dependents_of(graph, "Data", "B1") == ["Data!D1"]
    assert dependents_of(graph, "Data", "C1") == ["Data!D1"]
    assert graph.coverage_state is CoverageState.CHECKED


def test_let_value_can_resolve_same_named_workbook_range_before_binding() -> None:
    workbook = _sheet(
        cells={
            (1, 1): CellRecord(1, 1, 10),
            (1, 2): CellRecord(1, 2, 20),
            (1, 3): CellRecord(1, 3, 30),
            (1, 4): CellRecord(
                1,
                4,
                None,
                formula="=LET(rate,rate+B1,rate+C1)",
            ),
        },
    )
    workbook.named_ranges.append(NamedRange("rate", "Data!A1"))

    graph = build_dependency_graph(workbook)

    assert dependents_of(graph, "Data", "A1") == ["Data!D1"]
    assert dependents_of(graph, "Data", "B1") == ["Data!D1"]
    assert dependents_of(graph, "Data", "C1") == ["Data!D1"]
    assert graph.coverage_state is CoverageState.CHECKED


def test_malformed_lexical_formula_degrades_without_false_invalid_reference() -> None:
    workbook = _sheet(
        cells={(1, 2): CellRecord(1, 2, None, formula="=LET(item,A1)")},
    )

    graph = build_dependency_graph(workbook)

    assert graph.invalid_reason_counts == {}
    assert graph.unsupported_reason_counts == {"malformed_let": 1}
    assert graph.coverage_state is CoverageState.DEGRADED


def test_lexical_formula_dependency_checkpoint_after_step3() -> None:
    workbook = _sheet(
        cells={
            (1, 4): CellRecord(
                1,
                4,
                None,
                formula="=LET(rate,A1,scaled,rate*B1,LAMBDA(item,item+scaled)(C1))",
            ),
        }
    )

    graph = build_dependency_graph(workbook)

    assert dependents_of(graph, "Data", "A1") == ["Data!D1"]
    assert dependents_of(graph, "Data", "B1") == ["Data!D1"]
    assert dependents_of(graph, "Data", "C1") == ["Data!D1"]
    assert graph.invalid_references == []
    assert graph.invalid_reason_counts == {}
    assert graph.reference_status_counts[ReferenceStatus.INVALID] == 0
    assert graph.local_symbol_count == 6
    assert graph.coverage_state is CoverageState.CHECKED


def test_final_impact_limit_has_one_accurate_trailing_marker() -> None:
    cells = {
        (1, 1): CellRecord(1, 1, 10),
        **{
            (row, 2): CellRecord(row, 2, None, formula="=A1")
            for row in range(1, 31)
        },
    }
    workbook = _sheet(max_row=30, cells=cells)
    finding = Finding(
        artifact="excel",
        finding_class=FindingClass.VALUE_CHANGED,
        sheet="Data",
        location="A1",
        impacts=["Excel chart 'Trend' series 'Revenue'"],
        message="changed",
    )

    annotate_impacts([finding], build_dependency_graph(workbook))
    limit_impacts([finding])

    assert len(finding.impacts) == 26
    assert finding.impacts[-1] == "... 6 more"
    assert not finding.impacts[0].startswith("...")


def test_invalid_bubble_size_reference_is_a_chart_finding() -> None:
    workbook = _chart_workbook()
    series = workbook.charts[0].series[0]
    series.bubble_size_ref = "Missing!A1:A2"

    findings = diff_charts(workbook, workbook)

    assert any(
        finding.finding_class is FindingClass.CHART_REFERENCE_INVALID
        and "bubble size" in finding.message
        for finding in findings
    )


def test_empty_sheet_whole_column_is_invalid() -> None:
    workbook = _sheet(max_row=0, max_column=0)

    resolution = resolve_reference(
        workbook,
        "Data!A:A",
        host_sheet="Data",
        require_within_sheet=False,
    )

    assert resolution.status is ReferenceStatus.INVALID
    assert "empty sheet" in resolution.detail


def test_parse_error_degrades_dependency_coverage() -> None:
    workbook = _sheet(
        cells={(1, 2): CellRecord(1, 2, None, formula="=###")},
    )

    graph = build_dependency_graph(workbook)

    assert graph.coverage_state is CoverageState.DEGRADED
    assert graph.parse_errors
    assert "unparseable" in graph.coverage_detail


def test_spill_and_anchorarray_resolve_only_with_declared_extent() -> None:
    workbook = _sheet(max_row=10, max_column=5)
    workbook.formula_ranges.append(
        FormulaRangeDescriptor(
            sheet="Data",
            anchor_row=2,
            anchor_column=2,
            cell_range="B2:B5",
            formula_type="array",
            always_calculate=True,
        )
    )

    spill = resolve_reference(workbook, "B2#", host_sheet="Data")
    wrapper = resolve_reference(
        workbook,
        "_xlfn.ANCHORARRAY(B2)",
        host_sheet="Data",
    )

    assert spill.status is wrapper.status is ReferenceStatus.RESOLVED
    assert spill.ranges == wrapper.ranges
    assert spill.ranges[0].min_row == 2 and spill.ranges[0].max_row == 5

    workbook.formula_ranges.clear()
    unsupported = resolve_reference(workbook, "B2#", host_sheet="Data")
    assert unsupported.status is ReferenceStatus.UNSUPPORTED
    assert "spill extent" in unsupported.detail


def test_implicit_intersection_is_bounded_by_host_context() -> None:
    workbook = _sheet(max_row=10, max_column=5)

    direct = resolve_reference(workbook, "@A1", host_sheet="Data", host_cell=(5, 5))
    vertical = resolve_reference(
        workbook,
        "@A1:A5",
        host_sheet="Data",
        host_cell=(3, 4),
    )
    ambiguous = resolve_reference(
        workbook,
        "@A1:B5",
        host_sheet="Data",
        host_cell=(3, 1),
    )

    assert direct.status is ReferenceStatus.RESOLVED and direct.size == 1
    assert vertical.status is ReferenceStatus.RESOLVED
    assert (vertical.ranges[0].min_row, vertical.ranges[0].min_col) == (3, 1)
    assert ambiguous.status is ReferenceStatus.UNSUPPORTED


def test_spill_dependency_expands_declared_range_without_parse_degradation() -> None:
    workbook = _sheet(
        max_row=10,
        max_column=5,
        cells={(1, 4): CellRecord(1, 4, None, formula="=SUM(B2#)")},
    )
    workbook.formula_ranges.append(
        FormulaRangeDescriptor(
            sheet="Data",
            anchor_row=2,
            anchor_column=2,
            cell_range="B2:B5",
            formula_type="array",
        )
    )

    graph = build_dependency_graph(workbook)

    assert dependents_of(graph, "Data", "B4") == ["Data!D1"]
    assert graph.coverage_state is CoverageState.CHECKED
    assert graph.parse_errors == []


def test_unproven_spill_dependency_is_unsupported_not_unparseable() -> None:
    workbook = _sheet(
        cells={(1, 4): CellRecord(1, 4, None, formula="=SUM(B2#)")},
    )

    graph = build_dependency_graph(workbook)

    assert graph.coverage_state is CoverageState.DEGRADED
    assert len(graph.unsupported_references) == 1
    assert graph.parse_errors == []

    result = preflight_workbook(workbook, default_profile())
    coverage = next(
        item for item in result.coverage if item.check_id == "excel-dependencies"
    )
    assert coverage.state is CoverageState.DEGRADED
    assert "unsupported" in coverage.detail


def test_chart_spill_source_uses_the_shared_declared_extent() -> None:
    workbook = _chart_workbook()
    workbook.charts[0].series[0].values_ref = "Data!$B$2#"
    workbook.sheets[0].max_row = 10
    workbook.formula_ranges.append(
        FormulaRangeDescriptor(
            sheet="Data",
            anchor_row=2,
            anchor_column=2,
            cell_range="B2:B5",
            formula_type="array",
        )
    )

    state, detail = chart_reference_coverage(workbook)

    assert state is CoverageState.CHECKED
    assert "validated" in detail


def _chart_workbook() -> WorkbookSnapshot:
    series = ChartSeries(
        index=0,
        values_ref="Data!$B$1",
        categories_ref=None,
        name_text="Revenue",
    )
    plot = ChartPlot(index=0, chart_type="lineChart", series=[series])
    return WorkbookSnapshot(
        "charts.xlsx",
        "xlsx",
        True,
        True,
        sheets=[
            SheetSnapshot(
                "Data",
                "visible",
                10,
                5,
                {
                    (1, 1): CellRecord(1, 1, 10),
                    (1, 2): CellRecord(1, 2, None, formula="=A1"),
                },
            )
        ],
        charts=[
            ChartDescriptor(
                sheet="Data",
                title="Trend",
                chart_type="lineChart",
                series=[series],
                plots=[plot],
            )
        ],
    )


def test_precedent_findings_include_transitive_excel_chart_impact() -> None:
    workbook = _chart_workbook()
    graph = build_dependency_graph(workbook)
    finding = Finding(
        artifact="excel",
        finding_class=FindingClass.VALUE_CHANGED,
        sheet="Data",
        location="A1",
        message="changed",
    )

    annotate_impacts([finding], graph)
    annotate_chart_impacts([finding], workbook, graph)

    assert "Data!B1" in finding.impacts
    assert "Excel chart 'Trend' series 'Revenue'" in finding.impacts


def test_formula_findings_include_downstream_excel_chart_impact() -> None:
    workbook = _chart_workbook()
    sheet = workbook.sheet("Data")
    sheet.cells[(1, 3)] = CellRecord(1, 3, None, formula="=B1")
    workbook.charts[0].series[0].values_ref = "Data!$C$1"
    graph = build_dependency_graph(workbook)
    finding = Finding(
        artifact="excel",
        finding_class=FindingClass.FORMULA_LOGIC_CHANGED,
        sheet="Data",
        location="B1",
        message="changed",
    )

    annotate_chart_impacts([finding], workbook, graph)

    assert "Excel chart 'Trend' series 'Revenue'" in finding.impacts


def test_confirmed_chart_label_mapping_adds_transitive_ppt_chart_impact() -> None:
    workbook = _chart_workbook()
    graph = build_dependency_graph(workbook)
    ppt_series = PptChartSeries(
        index=0,
        order=0,
        name="Revenue",
        categories=["Jan-26"],
        values=[10.0],
        plot_index=0,
        source_index=0,
        source_id="series-0",
    )
    ppt_plot = PptChartPlot(
        index=0,
        chart_type="lineChart",
        series=[ppt_series],
        visible_labels=[PptChartLabel("10", 0, 0, "Jan-26")],
        visible_label_values=["10"],
    )
    deck = DeckSnapshot(
        "deck.pptx",
        slides=[
            SlideContent(
                index=0,
                title="Dashboard",
                texts=[],
                charts=[
                    ChartContent(
                        chart_type="lineChart",
                        categories=["Jan-26"],
                        series=[("Revenue", [10.0])],
                        title="Trend",
                        plots=[ppt_plot],
                    )
                ],
            )
        ],
    )
    profile = CrosscheckProfile(
        mappings=[
            CrosscheckMapping(
                slide="Dashboard",
                line_skeleton="chart:Trend/Revenue/Jan-26",
                figure_index=0,
                label="Trend revenue",
                source_sheet="Data",
                source_cell="B1",
            )
        ]
    )
    finding = Finding(
        artifact="excel",
        finding_class=FindingClass.VALUE_CHANGED,
        sheet="Data",
        location="A1",
        message="changed",
    )

    annotate_ppt_chart_impacts([finding], deck, profile, graph)

    assert finding.impacts == [
        "PowerPoint chart label 'Trend/Revenue/Jan-26' on slide 'Dashboard'"
    ]


def test_missing_chart_mapping_occurrence_adds_no_spurious_impact() -> None:
    workbook = _chart_workbook()
    finding = Finding(
        artifact="excel",
        finding_class=FindingClass.VALUE_CHANGED,
        sheet="Data",
        location="A1",
        message="changed",
    )
    profile = CrosscheckProfile(
        mappings=[
            CrosscheckMapping(
                slide="Missing",
                line_skeleton="chart:Ghost/Series/Point",
                figure_index=0,
                source_sheet="Data",
                source_cell="B1",
            )
        ]
    )

    annotate_ppt_chart_impacts(
        [finding],
        DeckSnapshot("empty.pptx"),
        profile,
        build_dependency_graph(workbook),
    )

    assert finding.impacts == []


def _save_dependency_workbook(path: Path, value: int) -> None:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Data"
    sheet["A1"] = value
    sheet["B1"] = "=A1"
    workbook.save(path)


def _save_label_deck(path: Path) -> None:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    title = slide.shapes.title
    assert title is not None
    title.text = "Dashboard"
    data = ChartData()
    data.categories = ["Jan-26"]
    data.add_series("Revenue", [10])
    shape = cast(
        Any,
        slide.shapes.add_chart(
            XL_CHART_TYPE.LINE,
            Inches(1),
            Inches(1.5),
            Inches(6),
            Inches(4),
            data,
        ),
    )
    plot = shape.chart.plots[0]
    plot.has_data_labels = True
    plot.data_labels.show_value = True
    presentation.save(str(path))


def test_full_run_attaches_transitive_ppt_chart_label_impact(tmp_path: Path) -> None:
    baseline_excel = tmp_path / "baseline.xlsx"
    current_excel = tmp_path / "current.xlsx"
    baseline_ppt = tmp_path / "baseline.pptx"
    current_ppt = tmp_path / "current.pptx"
    _save_dependency_workbook(baseline_excel, 9)
    _save_dependency_workbook(current_excel, 10)
    _save_label_deck(baseline_ppt)
    _save_label_deck(current_ppt)
    profile = DeliverableProfile.model_validate(
        {
            "name": "transitive-chart-impact",
            "crosscheck": {
                "mappings": [
                    {
                        "slide": "Dashboard",
                        "line_skeleton": "chart:chart[0]/Revenue/Jan-26",
                        "figure_index": 0,
                        "label": "Trend revenue",
                        "source_sheet": "Data",
                        "source_cell": "B1",
                    }
                ]
            },
        }
    )

    result = run_qc(
        baseline_excel=baseline_excel,
        current_excel=current_excel,
        baseline_ppt=baseline_ppt,
        current_ppt=current_ppt,
        profile=profile,
    )
    finding = next(
        item
        for item in result.findings
        if item.finding_class is FindingClass.VALUE_CHANGED
        and item.sheet == "Data"
        and item.location == "A1"
    )

    assert (
        "PowerPoint chart label 'chart[0]/Revenue/Jan-26' on slide 'Dashboard'"
        in finding.impacts
    )
