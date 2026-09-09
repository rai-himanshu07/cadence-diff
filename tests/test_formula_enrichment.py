"""Formula enrichment parity, merging, and sandbox-adapter behavior."""

import json
import os
import subprocess
import sys
import types
from pathlib import Path
from typing import Any

import pytest
from openpyxl import Workbook
from openpyxl.workbook.defined_name import DefinedName

import qc_tool.io.libreoffice_formula as libreoffice_formula_module
from qc_tool.io.excel_formula import (
    _load_worker_result,
    _parse_defined_names,
    _terminate_owned_excel,
    extract_formulas_with_excel,
)
from qc_tool.io.excel_formula_worker import (
    _collect_defined_names,
    _grid,
    _read_formula_grid,
    _rectangles,
    _set_manual_calculation,
)
from qc_tool.io.formula_enrichment import (
    ExtractedDefinedName,
    FormulaEnrichmentError,
    FormulaExtraction,
    bounded_defined_names,
    classify_external_reachability,
    compute_formula_text_coverage,
    merge_formula_extraction,
)
from qc_tool.io.libreoffice_formula import (
    _convert_with_libreoffice,
    _extract_formulas_with_converter,
)
from qc_tool.io.model import CellRecord, FormulaTextCoverage, SheetSnapshot, WorkbookSnapshot
from qc_tool.io.xlsb_formula import XlsbFormulaScan


def _snapshot() -> WorkbookSnapshot:
    return WorkbookSnapshot(
        source_name="source.xlsb",
        file_format="xlsb",
        formulas_available=False,
        styles_available=False,
        formula_presence_available=True,
        sheets=[
            SheetSnapshot(
                name="Data",
                visibility="visible",
                max_row=1,
                max_column=1,
                cells={(1, 1): CellRecord(1, 1, 42.0, is_formula=True)},
            )
        ],
    )


def _scan(*, risky: tuple[str, ...] = ()) -> XlsbFormulaScan:
    return XlsbFormulaScan(
        formula_cells={"Data": frozenset({(1, 1)})},
        risky_features=risky,
        blocking_features=risky,
    )


def test_merge_preserves_original_cached_value() -> None:
    snapshot = _snapshot()
    extraction = FormulaExtraction(
        formulas={"Data": {(1, 1): "=20+22"}},
        engine="test-engine",
        detail="test extraction",
    )

    merge_formula_extraction(snapshot, _scan(), extraction)

    cell = snapshot.sheet("Data").cells[(1, 1)]
    assert cell.value == 42.0
    assert cell.formula == "=20+22" and cell.has_formula
    assert snapshot.formulas_available
    assert snapshot.formula_source == "test-engine"
    assert snapshot.formula_text_coverage.state == "complete"
    assert snapshot.formula_text_coverage.missing_count == 0


def test_merge_sets_formula_r1c1_when_the_adapter_supplies_it() -> None:
    snapshot = _snapshot()
    extraction = FormulaExtraction(
        formulas={"Data": {(1, 1): "=20+22"}},
        engine="test-engine",
        detail="test extraction",
        formulas_r1c1={"Data": {(1, 1): "=20+22"}},
    )

    merge_formula_extraction(snapshot, _scan(), extraction)

    assert snapshot.sheet("Data").cells[(1, 1)].formula_r1c1 == "=20+22"


def test_merge_leaves_formula_r1c1_none_when_the_adapter_never_supplies_it() -> None:
    snapshot = _snapshot()
    extraction = FormulaExtraction(
        formulas={"Data": {(1, 1): "=20+22"}},
        engine="test-engine",
        detail="test extraction",
    )

    merge_formula_extraction(snapshot, _scan(), extraction)

    assert snapshot.sheet("Data").cells[(1, 1)].formula_r1c1 is None


def test_coordinate_mismatch_fails_before_mutation() -> None:
    snapshot = _snapshot()
    extraction = FormulaExtraction(
        formulas={"Data": {(1, 2): "=20+22"}},
        engine="test-engine",
        detail="test extraction",
    )

    with pytest.raises(FormulaEnrichmentError, match="coordinate mismatch"):
        merge_formula_extraction(snapshot, _scan(), extraction)

    assert snapshot.sheet("Data").cells[(1, 1)].formula is None
    assert not snapshot.formulas_available
    # Fatal rejection happens before any mutation, including coverage state.
    assert snapshot.formula_text_coverage.state == "none"


# --- Step 4b: partial formula-text coverage -------------------------------


def _two_cell_snapshot() -> WorkbookSnapshot:
    return WorkbookSnapshot(
        source_name="source.xlsb",
        file_format="xlsb",
        formulas_available=False,
        styles_available=False,
        formula_presence_available=True,
        sheets=[
            SheetSnapshot(
                name="Data",
                visibility="visible",
                max_row=2,
                max_column=1,
                cells={
                    (1, 1): CellRecord(1, 1, 42.0, is_formula=True),
                    (2, 1): CellRecord(2, 1, 7.0, is_formula=True),
                },
            )
        ],
    )


def _two_cell_scan(*, passive: tuple[str, ...] = ()) -> XlsbFormulaScan:
    return XlsbFormulaScan(
        formula_cells={"Data": frozenset({(1, 1), (2, 1)})},
        passive_features=passive,
    )


def test_missing_coordinate_is_tolerated_as_partial_coverage() -> None:
    snapshot = _two_cell_snapshot()
    extraction = FormulaExtraction(
        formulas={"Data": {(1, 1): "=20+22"}},
        engine="test-engine",
        detail="test extraction",
    )

    merge_formula_extraction(snapshot, _two_cell_scan(), extraction)

    merged = snapshot.sheet("Data").cells[(1, 1)]
    missing = snapshot.sheet("Data").cells[(2, 1)]
    assert merged.formula == "=20+22" and merged.value == 42.0
    assert missing.formula is None and missing.is_formula and missing.value == 7.0
    assert not snapshot.formulas_available  # historical meaning: complete only
    coverage = snapshot.formula_text_coverage
    assert coverage.state == "partial"
    assert coverage.expected_count == 2
    assert coverage.merged_count == 1
    assert coverage.missing_count == 1


def test_formula_text_coverage_transitions_none_complete_partial() -> None:
    assert FormulaTextCoverage().state == "none"

    scan = _two_cell_scan()
    complete = compute_formula_text_coverage(
        scan,
        FormulaExtraction(
            formulas={"Data": {(1, 1): "=1", (2, 1): "=2"}}, engine="e", detail="d"
        ),
    )
    assert complete.state == "complete"
    assert complete.missing_count == 0

    partial = compute_formula_text_coverage(
        scan,
        FormulaExtraction(formulas={"Data": {(1, 1): "=1"}}, engine="e", detail="d"),
    )
    assert partial.state == "partial"
    assert partial.missing_count == 1


def test_bounded_defined_names_drops_duplicates_and_oversized_targets() -> None:
    names = [
        ExtractedDefinedName(name="A", target="Sheet1!$A$1"),
        ExtractedDefinedName(name="A", target="Sheet1!$B$1"),  # duplicate identity
        ExtractedDefinedName(name="B", target="x" * 5000),  # oversized target
        ExtractedDefinedName(name="C", target="Sheet1!$C$1"),
    ]

    bounded, complete = bounded_defined_names(names)

    assert [item.name for item in bounded] == ["A", "C"]
    assert complete is False


def test_bounded_defined_names_reports_complete_when_nothing_dropped() -> None:
    names = [ExtractedDefinedName(name="A", target="Sheet1!$A$1")]

    bounded, complete = bounded_defined_names(names)

    assert bounded == tuple(names)
    assert complete is True


def test_reachability_unproven_when_coverage_is_partial() -> None:
    coverage = FormulaTextCoverage(
        state="partial", expected_count=2, merged_count=1, missing_count=1
    )
    extraction = FormulaExtraction(
        formulas={"Data": {(1, 1): "=1"}},
        engine="e",
        detail="d",
        defined_names=(ExtractedDefinedName(name="External1", target="[1]Sheet1!A1"),),
        defined_names_complete=True,
    )

    verdict = classify_external_reachability(extraction, coverage)

    assert verdict.proven is False
    assert verdict.live is False


def test_reachability_unproven_when_defined_names_incomplete() -> None:
    coverage = FormulaTextCoverage(
        state="complete", expected_count=1, merged_count=1, missing_count=0
    )
    extraction = FormulaExtraction(
        formulas={"Data": {(1, 1): "=1"}},
        engine="e",
        detail="d",
        defined_names=(),
        defined_names_complete=False,
    )

    verdict = classify_external_reachability(extraction, coverage)

    assert verdict.proven is False


def test_reachability_proven_inactive_stale_unused_name() -> None:
    coverage = FormulaTextCoverage(
        state="complete", expected_count=2, merged_count=2, missing_count=0
    )
    extraction = FormulaExtraction(
        formulas={"Data": {(1, 1): "=SUM(A2:A10)", (2, 1): "=A1*2"}},
        engine="e",
        detail="d",
        defined_names=(ExtractedDefinedName(name="StaleExternal", target="[1]Sheet1!A1"),),
        defined_names_complete=True,
    )

    verdict = classify_external_reachability(extraction, coverage)

    assert verdict.proven is True
    assert verdict.live is False
    assert verdict.direct_reference_count == 0
    assert verdict.transitive_reference_count == 0


def test_reachability_proven_live_direct_external_reference() -> None:
    coverage = FormulaTextCoverage(
        state="complete", expected_count=1, merged_count=1, missing_count=0
    )
    extraction = FormulaExtraction(
        formulas={"Data": {(1, 1): "='[1]Sheet1'!A1"}},
        engine="e",
        detail="d",
        defined_names_complete=True,
    )

    verdict = classify_external_reachability(extraction, coverage)

    assert verdict.proven is True
    assert verdict.live is True
    assert verdict.direct_reference_count == 1
    assert verdict.transitive_reference_count == 0


@pytest.mark.parametrize(
    "formula",
    [
        "='[Other.xlsx]Sheet1'!A1",
        "='file:///C:/reports/other.xlsx'#$Sheet1.A1",
    ],
)
def test_reachability_recognizes_named_file_and_url_external_syntax(
    formula: str,
) -> None:
    coverage = FormulaTextCoverage(
        state="complete", expected_count=1, merged_count=1, missing_count=0
    )
    extraction = FormulaExtraction(
        formulas={"Data": {(1, 1): formula}},
        engine="e",
        detail="d",
        defined_names_complete=True,
    )

    verdict = classify_external_reachability(extraction, coverage)

    assert verdict.proven is True
    assert verdict.live is True
    assert verdict.direct_reference_count == 1


def test_reachability_proven_live_transitive_named_reference() -> None:
    coverage = FormulaTextCoverage(
        state="complete", expected_count=1, merged_count=1, missing_count=0
    )
    extraction = FormulaExtraction(
        formulas={"Data": {(1, 1): "=LiveExternalName+1"}},
        engine="e",
        detail="d",
        defined_names=(
            ExtractedDefinedName(name="LiveExternalName", target="[1]Sheet1!A1"),
        ),
        defined_names_complete=True,
    )

    verdict = classify_external_reachability(extraction, coverage)

    assert verdict.proven is True
    assert verdict.live is True
    assert verdict.direct_reference_count == 0
    assert verdict.transitive_reference_count == 1


def test_reachability_closes_over_defined_name_alias_chain() -> None:
    coverage = FormulaTextCoverage(
        state="complete", expected_count=1, merged_count=1, missing_count=0
    )
    extraction = FormulaExtraction(
        formulas={"Data": {(1, 1): "=LocalAlias+1"}},
        engine="e",
        detail="d",
        defined_names=(
            ExtractedDefinedName(name="ExternalSeed", target="[1]Sheet1!A1"),
            ExtractedDefinedName(name="WorkbookAlias", target="=ExternalSeed"),
            ExtractedDefinedName(name="LocalAlias", target="=WorkbookAlias"),
        ),
        defined_names_complete=True,
    )

    verdict = classify_external_reachability(extraction, coverage)

    assert verdict.proven is True
    assert verdict.live is True
    assert verdict.direct_reference_count == 0
    assert verdict.transitive_reference_count == 1


def test_reachability_is_unproven_above_name_closure_bound() -> None:
    from qc_tool.io.formula_enrichment import MAX_REACHABILITY_DEFINED_NAMES

    coverage = FormulaTextCoverage(
        state="complete", expected_count=1, merged_count=1, missing_count=0
    )
    extraction = FormulaExtraction(
        formulas={"Data": {(1, 1): "=1"}},
        engine="e",
        detail="d",
        defined_names=tuple(
            ExtractedDefinedName(name=f"Name{index}", target="=1")
            for index in range(MAX_REACHABILITY_DEFINED_NAMES + 1)
        ),
        defined_names_complete=True,
    )

    verdict = classify_external_reachability(extraction, coverage)

    assert verdict.proven is False


def test_name_matcher_matches_per_name_search_exactly() -> None:
    """The single-alternation matcher must agree with the per-name regex everywhere."""
    from qc_tool.io.formula_enrichment import _compile_name_matcher, _references_name

    names = {"Rate", "Rate.FX", "rate_2", "A+B", "x[1]", "Ext(1)", "SUM", "a.b.c", "Z"}
    texts = [
        "=Rate*2",
        "=rate.fx+1",
        "=RATE_2",
        "=Rate2",
        "=SUM(A1:A3)",
        "=MySUM(1)",
        "=SUMX",
        "='A+B'!C1",
        "=A+B",
        "=x[1]",
        "=Ext(1)",
        "=a.b.c-a.b",
        "=Z",
        "=ZZ",
        "=1.Z",
        "=Q!Z1",
        "",
        "=",
    ]
    matcher = _compile_name_matcher(names)
    assert matcher is not None
    for text in texts:
        expected = any(_references_name(text, name) for name in names)
        assert (matcher.search(text) is not None) is expected, text
    assert _compile_name_matcher(set()) is None


def test_merge_computes_reachability_only_when_scan_has_passive_features() -> None:
    snapshot = _two_cell_snapshot()
    extraction = FormulaExtraction(
        formulas={"Data": {(1, 1): "=20+22", (2, 1): "=1"}},
        engine="test-engine",
        detail="test extraction",
        defined_names_complete=True,
    )

    merge_formula_extraction(
        snapshot, _two_cell_scan(passive=("external workbook links",)), extraction
    )

    assert snapshot.external_link_reachability is not None
    assert snapshot.external_link_reachability.proven is True
    assert snapshot.external_link_reachability.live is False

    snapshot_no_passive = _two_cell_snapshot()
    merge_formula_extraction(snapshot_no_passive, _two_cell_scan(), extraction)

    assert snapshot_no_passive.external_link_reachability is None



def test_libreoffice_adapter_reads_formula_only_and_checks_parity() -> None:
    def fake_converter(work_dir: Path, timeout: float) -> tuple[str, str]:
        assert timeout == 12.0
        if os.name == "posix":
            assert (work_dir / "input.xlsb").stat().st_mode & 0o777 == 0o600
        workbook = Workbook()
        sheet = workbook.active
        assert sheet is not None
        sheet.title = "Data"
        sheet["A1"] = "=20+22"
        workbook.save(work_dir / "output" / "input.xlsx")
        return "fake-libreoffice", "sandbox test"

    extraction = _extract_formulas_with_converter(
        b"decrypted xlsb bytes",
        _scan(),
        timeout=12.0,
        converter=fake_converter,
    )

    assert extraction.formulas == {"Data": {(1, 1): "=20+22"}}
    assert extraction.engine == "fake-libreoffice"


def test_libreoffice_adapter_rejects_risky_content_before_conversion() -> None:
    called = False

    def fake_converter(work_dir: Path, timeout: float) -> tuple[str, str]:
        nonlocal called
        called = True
        return "unused", "unused"

    with pytest.raises(FormulaEnrichmentError, match="VBA project"):
        _extract_formulas_with_converter(
            b"xlsb",
            _scan(risky=("VBA project",)),
            timeout=300.0,
            converter=fake_converter,
        )

    assert not called


def test_libreoffice_adapter_rejects_formula_coordinate_mismatch() -> None:
    def fake_converter(work_dir: Path, timeout: float) -> tuple[str, str]:
        workbook = Workbook()
        sheet = workbook.active
        assert sheet is not None
        sheet.title = "Data"
        sheet["B1"] = "=20+22"
        workbook.save(work_dir / "output" / "input.xlsx")
        return "fake-libreoffice", "mismatch test"

    with pytest.raises(FormulaEnrichmentError, match="coordinate mismatch"):
        _extract_formulas_with_converter(
            b"xlsb",
            _scan(),
            timeout=300.0,
            converter=fake_converter,
        )


@pytest.mark.skipif(os.name != "posix", reason="POSIX process-group cleanup")
def test_libreoffice_timeout_race_still_returns_enrichment_error(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    class FinishedDuringTimeout:
        pid = 12345
        returncode = None

        def communicate(self, timeout: float | None = None) -> tuple[str, str]:
            if timeout is not None:
                raise subprocess.TimeoutExpired("libreoffice", timeout)
            self.returncode = 0
            return "", ""

    monkeypatch.setattr(
        libreoffice_formula_module,
        "_sandbox_command",
        lambda work_dir: (["libreoffice"], "/usr/bin/libreoffice"),
    )
    monkeypatch.setattr(
        libreoffice_formula_module.subprocess,
        "run",
        lambda *args, **kwargs: subprocess.CompletedProcess(
            args[0], 0, "LibreOffice test\n", ""
        ),
    )
    monkeypatch.setattr(
        libreoffice_formula_module.subprocess,
        "Popen",
        lambda *args, **kwargs: FinishedDuringTimeout(),
    )
    monkeypatch.setattr(
        libreoffice_formula_module.os,
        "killpg",
        lambda *args: (_ for _ in ()).throw(ProcessLookupError()),
    )

    with pytest.raises(FormulaEnrichmentError, match="exceeded"):
        _convert_with_libreoffice(tmp_path, 1.0)


def test_excel_adapter_protocol_is_testable_without_windows() -> None:
    def fake_worker(
        work_dir: Path, request: dict[str, object], timeout: float
    ) -> FormulaExtraction:
        assert timeout == 7.0
        assert request["formula_count"] == 1
        assert request["formula_cells"] == {"Data": [[1, 1]]}
        assert (work_dir / "input.xlsb").read_bytes() == b"xlsb"
        return FormulaExtraction(
            formulas={"Data": {(1, 1): "=20+22"}},
            engine="excel:test:formula2",
            detail="test Excel worker",
        )

    extraction = extract_formulas_with_excel(
        b"xlsb", _scan(), timeout=7.0, worker_runner=fake_worker
    )

    assert extraction.formulas == {"Data": {(1, 1): "=20+22"}}


def test_excel_cleanup_contains_process_termination_denial(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    class FakeWin32Error(Exception):
        pass

    status_path = tmp_path / "excel-status.json"
    status_path.write_text(json.dumps({"pid": 1234, "created": 10.0}), encoding="utf-8")
    handle = object()
    terminated = False
    closed = False

    def terminate_process(process_handle: object, exit_code: int) -> None:
        nonlocal terminated
        assert process_handle is handle and exit_code == 1
        terminated = True
        raise FakeWin32Error(5, "Access is denied")

    def close_handle(process_handle: object) -> None:
        nonlocal closed
        assert process_handle is handle
        closed = True

    monkeypatch.setitem(
        sys.modules,
        "pywintypes",
        types.SimpleNamespace(error=FakeWin32Error),
    )
    monkeypatch.setitem(
        sys.modules,
        "win32api",
        types.SimpleNamespace(
            OpenProcess=lambda access, inherit, pid: handle,
            TerminateProcess=terminate_process,
            CloseHandle=close_handle,
        ),
    )
    monkeypatch.setitem(
        sys.modules,
        "win32con",
        types.SimpleNamespace(
            PROCESS_QUERY_INFORMATION=1,
            PROCESS_TERMINATE=2,
            SYNCHRONIZE=8,
            WAIT_TIMEOUT=258,
        ),
    )
    monkeypatch.setitem(
        sys.modules,
        "win32event",
        types.SimpleNamespace(WaitForSingleObject=lambda process_handle, timeout: 258),
    )
    monkeypatch.setitem(
        sys.modules,
        "win32process",
        types.SimpleNamespace(
            GetProcessTimes=lambda process_handle: {
                "CreationTime": types.SimpleNamespace(timestamp=lambda: 10.0)
            },
            GetModuleFileNameEx=lambda *args: (_ for _ in ()).throw(
                AssertionError("process image must not be queried")
            ),
        ),
    )

    _terminate_owned_excel(status_path)

    assert terminated
    assert closed


def test_excel_formula_rectangles_and_shapes() -> None:
    assert _rectangles({(1, 1), (1, 2), (2, 1), (2, 2), (4, 3)}) == [
        (1, 1, 2, 2),
        (4, 3, 4, 3),
    ]
    assert _grid("=A1", 1, 1) == [["=A1"]]
    assert _grid(("=A1", "=B1"), 1, 2) == [["=A1", "=B1"]]
    assert _grid((("=A1",), ("=A2",)), 2, 1) == [["=A1"], ["=A2"]]

    class FakeComError(Exception):
        pass

    class FormulaRange:
        Formula2: object = (("=A1", "=B1"),)

        @property
        def HasFormula(self) -> object:
            raise AssertionError("HasFormula must not gate Formula2 extraction")

    assert _read_formula_grid(FormulaRange(), 1, 2, FakeComError) == (
        [["=A1", "=B1"]],
        0,
    )

    FormulaRange.Formula2 = (("=A1", 42),)
    grid, invalid_count = _read_formula_grid(FormulaRange(), 1, 2, FakeComError)
    assert grid == [["=A1", None]]
    assert invalid_count == 1


def test_excel_manual_calculation_retries_with_a_guard_workbook() -> None:
    class FakeComError(Exception):
        pass

    class Guard:
        closed = False

        def Close(self, *, SaveChanges: bool) -> None:
            assert SaveChanges is False
            self.closed = True

    class Workbooks:
        def __init__(self, app: Any) -> None:
            self.app = app
            self.guard = Guard()

        def Add(self) -> Guard:
            self.app.workbook_exists = True
            return self.guard

    class App:
        def __init__(self) -> None:
            self.workbook_exists = False
            self.calculation = 0
            self.Workbooks = Workbooks(self)

        @property
        def Calculation(self) -> int:
            return self.calculation

        @Calculation.setter
        def Calculation(self, value: int) -> None:
            if not self.workbook_exists:
                raise FakeComError
            self.calculation = value

    app = App()

    guard = _set_manual_calculation(app, FakeComError)

    assert guard is app.Workbooks.guard
    assert app.Calculation == -4135
    assert not app.Workbooks.guard.closed


def test_excel_manual_calculation_retry_stays_fail_closed() -> None:
    class FakeComError(Exception):
        pass

    class Guard:
        closed = False

        def Close(self, *, SaveChanges: bool) -> None:
            assert SaveChanges is False
            self.closed = True

    class Workbooks:
        def __init__(self) -> None:
            self.guard = Guard()

        def Add(self) -> Guard:
            return self.guard

    class App:
        def __init__(self) -> None:
            self.Workbooks = Workbooks()

        @property
        def Calculation(self) -> int:
            return 0

        @Calculation.setter
        def Calculation(self, value: int) -> None:
            raise FakeComError

    app = App()

    with pytest.raises(FakeComError):
        _set_manual_calculation(app, FakeComError)

    assert app.Workbooks.guard.closed


@pytest.mark.parametrize(
    "payload",
    [
        {"schema_version": 2, "ok": True},
        {"schema_version": 1, "ok": False, "error": "Excel unavailable"},
        {"schema_version": 1, "ok": True, "formulas": {}, "engine": 1, "detail": "x"},
    ],
)
def test_excel_worker_result_validation_fails_closed(
    payload: dict[str, object], tmp_path: Path
) -> None:
    path = tmp_path / "result.json"
    path.write_text(json.dumps(payload), encoding="utf-8")

    with pytest.raises(FormulaEnrichmentError):
        _load_worker_result(path)


# --- Step 4b: defined-name collection (Windows worker + LibreOffice reuse) --


def test_collect_defined_names_splits_scope_and_bounds_duplicates() -> None:
    class Name:
        def __init__(self, name: str, refers_to: str, visible: bool = True) -> None:
            self.Name = name
            self.RefersTo = refers_to
            self.Visible = visible

    names = [
        Name("Workbook1", "=Sheet1!$A$1"),
        Name("Sheet1!Local1", "=Sheet1!$B$1", visible=False),
        Name("Workbook1", "=Sheet1!$C$1"),  # duplicate identity
    ]

    class FakeWorkbook:
        Names = names

    collected, complete = _collect_defined_names(FakeWorkbook())

    assert collected == [
        {"name": "Workbook1", "target": "=Sheet1!$A$1", "sheet": None, "hidden": False},
        {"name": "Local1", "target": "=Sheet1!$B$1", "sheet": "Sheet1", "hidden": True},
    ]
    assert complete is False  # the duplicate identity was dropped


def test_collect_defined_names_degrades_on_com_failure_without_raising() -> None:
    class FakeWorkbook:
        @property
        def Names(self) -> object:
            raise RuntimeError("COM call failed")

    collected, complete = _collect_defined_names(FakeWorkbook())

    assert collected == []
    assert complete is False


def test_parse_defined_names_success_path() -> None:
    payload: dict[str, object] = {
        "defined_names": [
            {"name": "A", "target": "=Sheet1!$A$1", "sheet": None, "hidden": False},
        ],
        "defined_names_complete": True,
    }

    names, complete = _parse_defined_names(payload)

    assert names == (ExtractedDefinedName(name="A", target="=Sheet1!$A$1"),)
    assert complete is True


def test_parse_defined_names_degrades_on_malformed_entry() -> None:
    payload: dict[str, object] = {
        "defined_names": [{"name": "A", "target": 42, "sheet": None, "hidden": False}],
        "defined_names_complete": True,
    }

    names, complete = _parse_defined_names(payload)

    assert names == ()
    assert complete is False


def test_parse_defined_names_absent_key_degrades_quietly() -> None:
    names, complete = _parse_defined_names({})

    assert names == ()
    assert complete is False


def test_libreoffice_extracted_defined_names_reuses_ooxml_scan(tmp_path: Path) -> None:
    workbook = Workbook()
    workbook.defined_names["Global1"] = DefinedName("Global1", attr_text="Sheet!$A$1")
    path = tmp_path / "converted.xlsx"
    workbook.save(path)

    names, complete = libreoffice_formula_module._extracted_defined_names(
        path.read_bytes()
    )

    assert complete is True
    assert names == (ExtractedDefinedName(name="Global1", target="Sheet!$A$1"),)
