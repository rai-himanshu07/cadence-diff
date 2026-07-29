"""Complete Excel chart extraction from raw OOXML package parts."""

import zipfile
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.label import DataLabelList

from qc_tool.config.profile import default_profile
from qc_tool.coverage import CoverageState
from qc_tool.excel.align import align_workbooks
from qc_tool.excel.diff_structure import diff_workbook_structure
from qc_tool.excel.preflight import preflight_workbook
from qc_tool.findings import FindingClass
from qc_tool.io.loader import load_workbook_snapshot
from qc_tool.io.model import (
    ChartDescriptor,
    ChartPlot,
    ChartSeries,
    SheetSnapshot,
    WorkbookSnapshot,
)


def _combo_workbook(path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Data"
    sheet.append(["Period", "Revenue", "Margin"])
    sheet.append(["Jan", 10, 0.1])
    sheet.append(["Feb", 20, 0.2])
    sheet.append(["Mar", 30, 0.3])

    columns = BarChart()
    columns.title = "Revenue and margin"
    columns.add_data(
        Reference(sheet, min_col=2, min_row=1, max_row=4),
        titles_from_data=True,
    )
    columns.set_categories(Reference(sheet, min_col=1, min_row=2, max_row=4))
    columns.dLbls = DataLabelList(showVal=True)
    columns.y_axis.scaling.min = 0
    columns.y_axis.scaling.max = 40
    columns.y_axis.majorUnit = 10
    assert columns.legend is not None
    columns.legend.position = "b"

    line = LineChart()
    line.add_data(
        Reference(sheet, min_col=3, min_row=1, max_row=4),
        titles_from_data=True,
    )
    line.set_categories(Reference(sheet, min_col=1, min_row=2, max_row=4))
    line.y_axis.axId = 200
    line.y_axis.crosses = "max"
    line.y_axis.scaling.min = 0
    line.y_axis.scaling.max = 1
    line.y_axis.majorUnit = 0.2

    columns += line
    columns.anchor = "E2"
    sheet.add_chart(columns)
    workbook.save(path)


def test_combo_chart_extracts_every_plot_series_axis_and_geometry(
    tmp_path: Path,
) -> None:
    path = tmp_path / "combo.xlsx"
    _combo_workbook(path)

    snapshot = load_workbook_snapshot(path)

    assert snapshot.charts_available
    assert len(snapshot.charts) == 1
    chart = snapshot.charts[0]
    assert chart.title == "Revenue and margin"
    assert chart.chart_type == "barChart"
    assert chart.source_index == 0
    assert chart.source_id.startswith("Data!chart[0]@")
    assert chart.anchor is not None
    assert chart.anchor.anchor_type == "oneCellAnchor"
    assert (chart.anchor.from_col, chart.anchor.from_row) == (4, 1)
    assert chart.anchor.width is not None and chart.anchor.width > 0
    assert chart.anchor.height is not None and chart.anchor.height > 0

    assert [plot.chart_type for plot in chart.plots] == ["barChart", "lineChart"]
    assert [plot.axis_ids for plot in chart.plots] == [
        ("10", "100"),
        ("10", "200"),
    ]
    assert [plot.axis_group for plot in chart.plots] == ["primary", "secondary-1"]
    assert chart.plots[0].data_labels is not None
    assert chart.plots[0].data_labels.show_value
    assert len(chart.series) == 2
    assert [series.plot_index for series in chart.series] == [0, 1]
    assert [series.source_index for series in chart.series] == [0, 0]
    assert chart.series[0].source_id.endswith("plot[0]/series[0]")
    assert chart.series[1].source_id.endswith("plot[1]/series[0]")
    assert [series.name_ref for series in chart.series] == [
        "'Data'!B1",
        "'Data'!C1",
    ]
    assert [series.values_ref for series in chart.series] == [
        "'Data'!$B$2:$B$4",
        "'Data'!$C$2:$C$4",
    ]
    assert [series.categories_ref for series in chart.series] == [
        "'Data'!$A$2:$A$4",
        "'Data'!$A$2:$A$4",
    ]

    axes = {axis.axis_id: axis for axis in chart.axes}
    assert set(axes) == {"10", "100", "200"}
    assert axes["100"].axis_type == "valAx"
    assert (axes["100"].minimum, axes["100"].maximum) == (0.0, 40.0)
    assert axes["100"].major_unit == 10.0
    assert axes["200"].crosses == "max"
    assert axes["200"].major_unit == 0.2
    assert chart.legend is not None
    assert chart.legend.position == "b"


def test_duplicate_untitled_charts_have_distinct_source_identities(
    tmp_path: Path,
) -> None:
    path = tmp_path / "duplicates.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Data"
    sheet.append(["Period", "Value"])
    sheet.append(["Jan", 10])
    sheet.append(["Feb", 20])
    for anchor in ("D2", "D18"):
        chart = LineChart()
        chart.add_data(Reference(sheet, min_col=2, min_row=1, max_row=3))
        chart.set_categories(Reference(sheet, min_col=1, min_row=2, max_row=3))
        chart.anchor = anchor
        sheet.add_chart(chart)
    workbook.save(path)

    snapshot = load_workbook_snapshot(path)

    assert len(snapshot.charts) == 2
    assert [chart.source_index for chart in snapshot.charts] == [0, 1]
    assert len({chart.source_id for chart in snapshot.charts}) == 2
    assert len({chart.anchor.signature for chart in snapshot.charts if chart.anchor}) == 2


def test_broken_linked_chart_degrades_without_losing_workbook_cells(
    tmp_path: Path,
) -> None:
    source = tmp_path / "source.xlsx"
    broken = tmp_path / "broken.xlsx"
    _combo_workbook(source)
    with zipfile.ZipFile(source) as source_archive, zipfile.ZipFile(broken, "w") as out:
        for member in source_archive.infolist():
            if member.filename == "xl/charts/chart1.xml":
                continue
            out.writestr(member, source_archive.read(member))

    snapshot = load_workbook_snapshot(broken)
    preflight = preflight_workbook(snapshot, default_profile())
    structure = next(
        item for item in preflight.coverage if item.check_id == "excel-intrinsic-structure"
    )

    assert snapshot.sheet("Data").cell("B2") is not None
    assert not snapshot.charts_available
    assert snapshot.charts == []
    assert "missing" in snapshot.chart_detail.lower()
    assert structure.state is CoverageState.DEGRADED
    assert "chart metadata" in structure.detail.lower()


def test_added_primary_series_does_not_cross_pair_secondary_series() -> None:
    sheet = SheetSnapshot("Data", "visible", 10, 10)
    baseline_chart = ChartDescriptor(
        sheet="Data",
        title="Combo",
        chart_type="barChart",
        plots=[
            ChartPlot(
                index=0,
                chart_type="barChart",
                series=[ChartSeries(0, "Data!B2:B4", "Data!A2:A4")],
            ),
            ChartPlot(
                index=1,
                chart_type="lineChart",
                series=[ChartSeries(1, "Data!C2:C4", "Data!A2:A4")],
            ),
        ],
    )
    baseline_chart.series = [
        series for plot in baseline_chart.plots for series in plot.series
    ]
    current_chart = ChartDescriptor(
        sheet="Data",
        title="Combo",
        chart_type="barChart",
        plots=[
            ChartPlot(
                index=0,
                chart_type="barChart",
                series=[
                    ChartSeries(0, "Data!B2:B4", "Data!A2:A4"),
                    ChartSeries(2, "Data!D2:D4", "Data!A2:A4"),
                ],
            ),
            ChartPlot(
                index=1,
                chart_type="lineChart",
                series=[ChartSeries(1, "Data!C2:C4", "Data!A2:A4")],
            ),
        ],
    )
    current_chart.series = [
        series for plot in current_chart.plots for series in plot.series
    ]
    baseline = WorkbookSnapshot(
        "baseline.xlsx",
        "xlsx",
        True,
        True,
        sheets=[sheet],
        charts=[baseline_chart],
    )
    current = WorkbookSnapshot(
        "current.xlsx",
        "xlsx",
        True,
        True,
        sheets=[sheet],
        charts=[current_chart],
    )

    findings = diff_workbook_structure(
        baseline,
        current,
        align_workbooks(baseline, current),
    )
    chart_findings = [
        finding
        for finding in findings
        if finding.finding_class is FindingClass.CHART_SERIES_CHANGED
    ]

    assert len(chart_findings) == 1
    assert "series" in chart_findings[0].message
    assert "added" in chart_findings[0].message
    assert not any("Data!C2:C4" in str(finding.model_dump()) for finding in chart_findings)
