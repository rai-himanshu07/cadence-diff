"""Deterministic analyst review groups over atomic findings."""

from __future__ import annotations

import random

from openpyxl.utils import get_column_letter

from qc_tool.findings import Finding, FindingClass, Severity
from qc_tool.review import apply_group_review, build_review_groups, review_counts


def _finding(
    row: int,
    column: int,
    *,
    finding_class: FindingClass = FindingClass.VALUE_CHANGED,
    severity: Severity = Severity.CRITICAL,
    baseline_row: int | None = None,
    element: str = "",
) -> Finding:
    location = f"{get_column_letter(column)}{row}"
    baseline_location = (
        f"{get_column_letter(column)}{baseline_row}"
        if baseline_row is not None
        else None
    )
    return Finding(
        finding_id=f"F{row:04d}{column:03d}",
        artifact="excel",
        finding_class=finding_class,
        severity=severity,
        sheet="Data",
        location=location,
        baseline_location=baseline_location,
        element=element,
        baseline_value=str(row + column),
        current_value=str(row + column + 1),
        message=f"Data!{location}: value changed",
    )


def _signature(groups: list) -> list[tuple[object, ...]]:
    return [
        (
            group.group_id,
            group.finding_class,
            group.severity,
            group.ranges,
            group.baseline_ranges,
            tuple(member.finding_id for member in group.members),
        )
        for group in groups
    ]


def test_two_thousand_contiguous_cells_form_one_review_group() -> None:
    findings = [
        _finding(row, column, baseline_row=row - 1)
        for row in range(2, 42)
        for column in range(1, 51)
    ]

    groups = build_review_groups(findings)

    assert len(groups) == 1
    group = groups[0]
    assert group.ranges == ("A2:AX41",)
    assert group.baseline_ranges == ("A1:AX40",)
    assert group.bounding_range == "A2:AX41"
    assert group.member_count == 2_000
    assert {member.finding_id for member in group.members} == {
        finding.finding_id for finding in findings
    }
    counts = review_counts(groups)
    assert counts.review_items[Severity.CRITICAL] == 1
    assert counts.atomic_findings[Severity.CRITICAL] == 2_000


def test_grouping_is_deterministic_for_shuffled_input() -> None:
    findings = [
        _finding(row, column, baseline_row=row - 2)
        for row in range(3, 8)
        for column in range(2, 6)
    ]
    shuffled = findings.copy()
    random.Random(7).shuffle(shuffled)

    assert _signature(build_review_groups(shuffled)) == _signature(
        build_review_groups(findings)
    )


def test_diagonal_gapped_and_mixed_findings_do_not_merge() -> None:
    findings = [
        _finding(1, 1),
        _finding(2, 2),  # diagonal only
        _finding(4, 1),  # row gap
        _finding(4, 2, finding_class=FindingClass.FORMULA_LOGIC_CHANGED),
        _finding(5, 1, severity=Severity.WARNING),
        _finding(6, 1, element="Control A"),
        _finding(7, 1, element="Control B"),
    ]

    groups = build_review_groups(findings)

    assert len(groups) == len(findings)
    assert {group.member_count for group in groups} == {1}


def test_irregular_component_uses_exact_disjoint_rectangles() -> None:
    findings = [
        _finding(1, 1),
        _finding(1, 2),
        _finding(2, 1),
        _finding(3, 1),
        _finding(3, 2),
    ]

    group = build_review_groups(findings)[0]

    assert group.bounding_range == "A1:B3"
    assert group.ranges == ("A1:B1", "A2", "A3:B3")
    assert group.member_count == 5


def test_baseline_translation_splits_adjacent_cells() -> None:
    groups = build_review_groups(
        [
            _finding(10, 1, baseline_row=9),
            _finding(11, 1, baseline_row=9),
        ]
    )

    assert len(groups) == 2


def test_non_cell_and_non_excel_findings_remain_singletons() -> None:
    findings = [
        Finding(
            finding_id="F0001",
            artifact="excel",
            finding_class=FindingClass.COLUMN_DELETED,
            severity=Severity.CRITICAL,
            sheet="Data",
            location="column B",
            message="column deleted",
        ),
        Finding(
            finding_id="F0002",
            artifact="ppt",
            finding_class=FindingClass.SLIDE_TEXT_CHANGED,
            severity=Severity.WARNING,
            slide="Summary",
            message="text changed",
        ),
    ]

    groups = build_review_groups(findings)

    assert len(groups) == 2
    assert all(group.member_count == 1 for group in groups)
    assert {member.finding_id for group in groups for member in group.members} == {
        "F0001",
        "F0002",
    }


def test_group_review_preserves_or_replaces_existing_decisions_explicitly() -> None:
    first = _finding(1, 1)
    first.severity = Severity.CRITICAL
    first.severity_overridden = True
    first.analyst_comment = "individual decision"
    second = _finding(2, 1)
    group = build_review_groups([first, second])[0]

    blank_only = apply_group_review(
        group,
        severity=Severity.WARNING,
        comment="group decision",
        replace_existing=False,
    )

    assert [update.finding_id for update in blank_only] == [second.finding_id]
    assert first.severity is Severity.CRITICAL
    assert first.analyst_comment == "individual decision"
    assert second.severity is Severity.WARNING
    assert second.analyst_comment == "group decision"

    replaced = apply_group_review(
        group,
        severity=Severity.EXPECTED,
        comment="approved group",
        replace_existing=True,
    )

    assert len(replaced) == 2
    assert all(member.severity is Severity.EXPECTED for member in group.members)
    assert all(member.analyst_comment == "approved group" for member in group.members)


def test_cap_degraded_badge_is_scoped_to_the_cap_summary_group() -> None:
    from qc_tool.ui.app import _review_group_rows

    normal = _finding(1, 1)
    capped = Finding(
        finding_id="F9999",
        artifact="excel",
        finding_class=FindingClass.FINDINGS_CAPPED,
        severity=Severity.WARNING,
        sheet="Data",
        message="retained details only",
    )

    rows = _review_group_rows(build_review_groups([normal, capped]))
    by_class = {row["class"]: row for row in rows}

    assert by_class[FindingClass.FINDINGS_CAPPED.value]["cap_degraded"] is True
    assert by_class[FindingClass.VALUE_CHANGED.value]["cap_degraded"] is False


def test_duplicate_non_cell_identities_get_unique_deterministic_group_ids() -> None:
    findings = [
        Finding(
            finding_id="F0002",
            artifact="ppt",
            finding_class=FindingClass.SLIDE_TEXT_CHANGED,
            severity=Severity.WARNING,
            slide="Summary",
            message="second text change",
        ),
        Finding(
            finding_id="F0001",
            artifact="ppt",
            finding_class=FindingClass.SLIDE_TEXT_CHANGED,
            severity=Severity.WARNING,
            slide="Summary",
            message="first text change",
        ),
    ]

    forward = build_review_groups(findings)
    reversed_groups = build_review_groups(list(reversed(findings)))

    assert len({group.group_id for group in forward}) == 2
    assert _signature(forward) == _signature(reversed_groups)
    assert {
        member.finding_id for group in forward for member in group.members
    } == {"F0001", "F0002"}
