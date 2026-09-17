"""UI smoke tests (criterion 12): pages render; perform_run produces artifacts."""

import asyncio
import datetime as dt
import hashlib
import inspect
import logging
import os
import re
from pathlib import Path
from types import SimpleNamespace
from typing import Literal, cast

import pytest
import yaml
from nicegui import app, events, ui
from nicegui.helpers import warnings as nicegui_warnings
from nicegui.testing import User
from openpyxl.utils import get_column_letter

import qc_tool.run_service as run_service
import qc_tool.ui.app as app_module
from qc_tool.config.profile import (
    DeliverableProfile,
    PopulationPolicy,
    ReviewPolicy,
    resolve_output_policy,
    save_profile,
)
from qc_tool.coverage import (
    CoverageItem,
    CoverageState,
    FindingOutputMode,
    MappingCoverage,
    QCRunMode,
)
from qc_tool.crosscheck.trace import MappingSuggestion, SuggestedSource
from qc_tool.engine import QCRunResult
from qc_tool.findings import (
    Finding,
    FindingClass,
    FindingEvidenceTag,
    FindingProvenance,
    FindingSubtype,
    FindingTemporalContext,
    GridExcerpt,
    Materiality,
    Severity,
)
from qc_tool.history.config_session import ConfigSessionStore, session_key_for
from qc_tool.history.run_state import RunStateRecord, RunStateStore, RunStatus
from qc_tool.history.store import RunHistory, RunRecord, sha256_file
from qc_tool.progress import CancellationToken, ProgressEvent, RunCancelled, RunPhase
from qc_tool.review import (
    ReviewGroup,
    apply_group_review,
    build_pattern_groups,
    build_review_groups,
)
from qc_tool.review_series import SeriesReviewLens, build_series_review_lens
from qc_tool.runqueue import RunRequest, get_exclusive_slot
from qc_tool.security import secure_managed_tree
from qc_tool.server_config import NetworkMode, ServerConfig
from qc_tool.signoff import finalize_run, required_acknowledgements
from qc_tool.ui.app import (
    LensEntry,
    _acceptance_summary,
    _capability_summary,
    _class_filter_label,
    _cluster_context_html,
    _context_grid_html,
    _evidence_axes,
    _files_for_mode,
    _filter_review_rows,
    _flatten_lens_rows,
    _format_bytes,
    _history_row,
    _history_trend_row,
    _initial_mode,
    _input_cautions,
    _lens_entries,
    _managed_request_paths,
    _mapping_stats,
    _outcome_summary,
    _persist_expired_lan_config,
    _population_members,
    _population_source_roles,
    _profile_path,
    _queue_status_line,
    _ranked_action_needs_refresh,
    _relative_time,
    _rerun_delta,
    _rerun_profile_choice,
    _review_group_rows,
    _role_requirement,
    _run_blockers,
    _safe_upload_name,
    _scope_summary,
    _set_desktop_focus_preference,
    _storage_prompt_due,
    _storage_secret,
    _temporary_profile_name,
    _temporary_row_matching_base,
    _terminal_request_summary,
    build_cluster_context,
    create_pages,
    list_profiles,
    load_profile_by_name,
    perform_run,
    persist_confirmed_mapping,
)
from qc_tool.ui.guide import (
    COMMON_TASKS,
    GUIDE_SCRIPT,
    GUIDE_SECTIONS,
    PROFILE_CONTROLS_EXAMPLE,
    render_guide,
)
from qc_tool.ui.theme import CSS, REVIEW_GROUPS_BODY_SLOT, page_frame
from tests.conftest import fixture_profile
from tests.fixtures.ranked_table_action_v2 import ranked_table_evidence_payload_v2
from tests.test_review_series import series_oracle

pytest_plugins = ["nicegui.testing.user_plugin"]


class _RecordingQueueManager:
    """UI-test queue boundary that records requests without starting a worker."""

    def __init__(self, store: RunStateStore) -> None:
        self.store = store
        self.shutdown_hook_installed = False
        self.submitted: list[tuple[RunRequest, dict[str, str]]] = []

    def submit(
        self,
        request: RunRequest,
        credentials: dict[str, str] | None = None,
    ) -> RunStateRecord:
        self.submitted.append((request, dict(credentials or {})))
        return RunStateRecord(
            request_id=request.request_id,
            created_at=dt.datetime.now(dt.UTC),
            status=RunStatus.QUEUED,
            mode=request.mode,
            profile=request.profile_name,
            files=dict(request.display_files),
            profile_snapshot=dict(request.profile),
            requested_output_mode=request.requested_output_mode,
        )

    def shutdown(self) -> None:
        pass


def _write_managed_retry_pair(work_dir: Path) -> dict[str, Path]:
    files = {
        "baseline_excel": work_dir / "uploads" / "baseline_excel" / "baseline.xlsx",
        "current_excel": work_dir / "uploads" / "current_excel" / "current.xlsx",
    }
    for role, path in files.items():
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_bytes(role.encode("ascii"))
    return files


def _write_valid_managed_retry_pair(work_dir: Path) -> dict[str, Path]:
    from openpyxl import Workbook

    files = {
        "baseline_excel": work_dir / "uploads" / "baseline_excel" / "baseline.xlsx",
        "current_excel": work_dir / "uploads" / "current_excel" / "current.xlsx",
    }
    for role, path in files.items():
        path.parent.mkdir(parents=True, exist_ok=True)
        workbook = Workbook()
        sheet = workbook.active
        assert sheet is not None
        sheet.title = "Data"
        sheet.append(["ID", "Value"])
        sheet.append(["A0", 10])
        sheet.append(["A1", 20])
        sheet.append(["A2", 31 if role == "current_excel" else 30])
        workbook.save(path)
    return files


def test_mode_toggle_pins_content_color_against_quasar() -> None:
    # Quasar paints the pressed button bg-primary + text-white, and --q-primary
    # is light in dark mode, so every toggle needs its own pinned colours.
    assert ".q-btn-toggle .q-btn .q-btn__content" in CSS
    assert '.q-btn-toggle .q-btn[aria-pressed="true"] .q-btn__content' in CSS
    assert '.q-btn-toggle .q-btn[aria-pressed="true"] { background: var(--btn-bg)' in CSS
    for field in (
        "provenance",
        "subtype",
        "materiality",
        "temporal_context",
        "expected_reason",
        "evidence_tags",
    ):
        assert field in app_module.FINDINGS_BODY_SLOT


def test_initial_mode_defaults_to_preflight_then_restores_valid_choices() -> None:
    assert _initial_mode(None) is QCRunMode.CURRENT_FILE_PREFLIGHT
    assert _initial_mode("invalid") is QCRunMode.CURRENT_FILE_PREFLIGHT
    assert _initial_mode("final_package") is QCRunMode.FINAL_PACKAGE
    assert _initial_mode(
        "final_package",
        rerun_mode=QCRunMode.CYCLE_COMPARISON,
    ) is QCRunMode.CYCLE_COMPARISON


def test_focus_preference_revokes_before_persisting_off(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    observed: list[str] = []

    class Service:
        def set_enabled(self, enabled: bool) -> None:
            observed.append(f"service:{enabled}")

    def save(_root: Path, _config: object) -> Path:
        observed.append("save")
        return tmp_path / "server-config.json"

    monkeypatch.setattr(app_module, "save_server_config", save)

    _set_desktop_focus_preference(tmp_path, Service(), False)  # type: ignore[arg-type]
    assert observed == ["service:False", "save"]
    observed.clear()
    _set_desktop_focus_preference(tmp_path, Service(), True)  # type: ignore[arg-type]
    assert observed == ["save", "service:True"]


@pytest.mark.parametrize("remembered", [False, True])
def test_lan_expiry_preserves_latest_focus_preference(
    tmp_path: Path,
    remembered: bool,
) -> None:
    app_module.save_server_config(
        tmp_path,
        ServerConfig(network=NetworkMode.LAN, desktop_focus=remembered),
    )

    _persist_expired_lan_config(tmp_path)

    config = app_module.load_server_config(tmp_path)
    assert config.network is NetworkMode.LOCAL
    assert config.desktop_focus is remembered


def test_evidence_axes_cover_every_typed_axis() -> None:
    finding = Finding(
        finding_id="F1",
        finding_class=FindingClass.VALUE_CHANGED,
        artifact="excel",
        message="value changed",
        severity=Severity.CRITICAL,
        sheet="Data",
        location="C7",
        baseline_value="1",
        current_value="2",
        element="cell",
        impacts=["Data!C8"],
        root_cause_key="rc-1",
        provenance=FindingProvenance.CHANGED,
        subtype=FindingSubtype.VALUE_REPLACEMENT,
        materiality=Materiality.MATERIAL,
        temporal_context=FindingTemporalContext.HISTORICAL,
        evidence_tags={FindingEvidenceTag.FORMULA_TEXT},
        analyst_comment="checked",
    )

    axes = dict(_evidence_axes(finding))

    for key in (
        "class",
        "artifact",
        "location",
        "baseline",
        "current",
        "element",
        "impacts",
        "root cause",
        "provenance",
        "subtype",
        "materiality",
        "temporal context",
        "evidence",
        "analyst comment",
    ):
        assert key in axes
    # Empty axes are dropped rather than rendered as blank rows.
    assert "waiver" not in axes
    assert "expected reason" not in axes


def test_population_members_decodes_shift_mode_across_rectangles() -> None:
    from qc_tool.findings import MembershipCodec, PopulationEvidence

    finding = Finding(
        artifact="excel",
        finding_class=FindingClass.FORMULA_LOGIC_CHANGED,
        sheet="Data",
        location="B2:B4",
        element="population",
        message="3 cells share one population",
        population=PopulationEvidence(
            member_count=3,
            membership=MembershipCodec(
                current_rectangles=("B2:B4",),
                baseline_mode="shift",
                shift=(-1, 0),
                member_count=3,
            ),
            first="B2",
            last="B4",
            shape_before_digest="a" * 64,
            shape_after_digest="b" * 64,
        ),
    )

    assert _population_members(finding) == [
        ("B2", "B1"),
        ("B3", "B2"),
        ("B4", "B3"),
    ]


def test_population_members_returns_explicit_pairs_verbatim() -> None:
    from qc_tool.findings import MembershipCodec, PopulationEvidence

    finding = Finding(
        artifact="excel",
        finding_class=FindingClass.FORMULA_LOGIC_CHANGED,
        sheet="Data",
        location="B2:D2",
        element="population",
        message="2 cells share one population",
        population=PopulationEvidence(
            member_count=2,
            membership=MembershipCodec(
                current_rectangles=("B2:D2",),
                baseline_mode="pairs",
                pairs=(("B2", "A1"), ("D2", "C4")),
                member_count=2,
            ),
            first="B2",
            last="D2",
            shape_before_digest="a" * 64,
            shape_after_digest="b" * 64,
        ),
    )

    assert _population_members(finding) == [("B2", "A1"), ("D2", "C4")]


def test_population_members_returns_empty_for_atomic_findings() -> None:
    finding = Finding(
        artifact="excel",
        finding_class=FindingClass.VALUE_CHANGED,
        sheet="Data",
        location="B2",
        message="value changed",
    )

    assert _population_members(finding) == []


def test_population_members_page_matches_full_decode_across_rectangle_boundaries() -> None:
    """Criterion 9: `population_members_page` must decode the exact same
    members `population_members` would at every position -- including a
    page that spans two rectangles -- without ever materializing the rest.
    """
    from qc_tool.findings import MembershipCodec, PopulationEvidence
    from qc_tool.review import population_members, population_members_page

    finding = Finding(
        artifact="excel",
        finding_class=FindingClass.FORMULA_LOGIC_CHANGED,
        sheet="Data",
        location="B2:B4,D2:D4",
        element="population",
        message="6 cells share one population",
        population=PopulationEvidence(
            member_count=6,
            membership=MembershipCodec(
                current_rectangles=("B2:B4", "D2:D4"),
                baseline_mode="shift",
                shift=(-1, 0),
                member_count=6,
            ),
            first="B2",
            last="D4",
            shape_before_digest="a" * 64,
            shape_after_digest="b" * 64,
        ),
    )
    full = population_members(finding)
    assert len(full) == 6

    for start in range(0, 7):
        for count in (1, 2, 3, 10):
            assert population_members_page(finding, start, count) == full[start : start + count]

    assert population_members_page(finding, 0, 0) == ()
    assert population_members_page(finding, -1, 3) == ()


def test_population_source_roles_for_primary_and_package_member() -> None:
    primary = Finding(
        artifact="excel",
        finding_class=FindingClass.VALUE_CHANGED,
        sheet="Data",
        location="B2",
        message="value changed",
    )
    member = Finding(
        artifact="excel",
        artifact_member="unit1",
        finding_class=FindingClass.VALUE_CHANGED,
        sheet="Data",
        location="B2",
        message="value changed",
    )

    assert _population_source_roles(primary) == ("baseline_excel", "current_excel")
    assert _population_source_roles(member) == (
        "baseline_excel:unit1",
        "current_excel:unit1",
    )


def _population_finding_and_record(
    work_dir: Path,
    baseline: Path,
    current: Path,
) -> tuple[Finding, RunRecord]:
    from qc_tool.config.profile import PopulationPolicy, ReviewPolicy

    profile = DeliverableProfile(
        name="population-excerpts-pure",
        review_policy=ReviewPolicy(
            populations=PopulationPolicy(enabled=True, threshold=10)
        ),
    )
    artifacts = perform_run(
        work_dir,
        {"baseline_excel": baseline, "current_excel": current},
        {},
        profile,
    )
    history = RunHistory(work_dir / "history.sqlite3")
    record = history.get_run(artifacts.run_id)
    assert record is not None
    finding = next(f for f in record.findings if f.population is not None)
    return finding, record


def test_load_population_sample_excerpts_builds_baseline_and_current_grids(
    tmp_path: Path,
) -> None:
    from qc_tool.ui.app import _load_population_sample_excerpts

    baseline, current = _build_population_fixture(tmp_path)
    finding, record = _population_finding_and_record(
        tmp_path / "work", baseline, current
    )

    loaded = _load_population_sample_excerpts(finding, record)

    assert loaded.disclosure == ""
    assert loaded.excerpts
    baseline_excerpt, current_excerpt = next(iter(loaded.excerpts.values()))
    assert baseline_excerpt is not None
    assert current_excerpt is not None


def test_load_population_sample_excerpts_discloses_missing_path(
    tmp_path: Path,
) -> None:
    from qc_tool.ui.app import _load_population_sample_excerpts

    baseline, current = _build_population_fixture(tmp_path)
    finding, record = _population_finding_and_record(
        tmp_path / "work", baseline, current
    )
    current.unlink()

    loaded = _load_population_sample_excerpts(finding, record)

    assert loaded.excerpts == {}
    assert "no longer at its recorded location" in loaded.disclosure


def test_load_population_sample_excerpts_discloses_hash_mismatch(
    tmp_path: Path,
) -> None:
    from qc_tool.ui.app import _load_population_sample_excerpts

    baseline, current = _build_population_fixture(tmp_path)
    finding, record = _population_finding_and_record(
        tmp_path / "work", baseline, current
    )
    # Overwrite the recorded source in place so the path still exists but the
    # content -- and therefore the hash -- no longer matches the run record.
    current.write_bytes(baseline.read_bytes())

    loaded = _load_population_sample_excerpts(finding, record)

    assert loaded.excerpts == {}
    assert "changed since this run recorded it" in loaded.disclosure


def test_load_population_sample_excerpts_discloses_no_samples(
    tmp_path: Path,
) -> None:
    from qc_tool.ui.app import _load_population_sample_excerpts

    # An atomic finding (population is None) never has samples to load.
    finding = Finding(
        artifact="excel",
        finding_class=FindingClass.VALUE_CHANGED,
        sheet="Data",
        location="B2",
        message="value changed",
    )
    record = RunRecord(
        run_id=1,
        started_at=dt.datetime(2026, 1, 1, tzinfo=dt.UTC),
        profile="default",
        mode=QCRunMode.CYCLE_COMPARISON,
        files={},
        file_hashes={},
        counts={},
        review_counts={},
        disclosures=[],
        verified_crosschecks=0,
        report_paths={},
    )

    loaded = _load_population_sample_excerpts(finding, record)

    assert loaded.excerpts == {}
    assert loaded.disclosure == "no samples recorded"


def test_profile_controls_guide_example_is_valid_yaml() -> None:
    payload = yaml.safe_load(PROFILE_CONTROLS_EXAMPLE)
    tie_outs = payload["excel"]["controls"]["tie_outs"]

    assert tie_outs[0]["components"] == ["Data!C2:C5"]
    assert [term["operation"] for term in tie_outs[1]["terms"]] == [
        "add",
        "subtract",
    ]


def test_guide_search_matches_all_tokens_and_bounded_aliases() -> None:
    assert "tokens.every" in GUIDE_SCRIPT
    assert "aliases[s.id]" in GUIDE_SCRIPT
    assert "mapping unavailable opaque" in GUIDE_SCRIPT


def test_guide_is_task_first_and_documents_desktop_launch() -> None:
    assert GUIDE_SECTIONS[0] == ("launch", "Launch and first run")
    assert ("focus", "Desktop Office focus") in GUIDE_SECTIONS
    assert ("reference", "Advanced capability reference") in GUIDE_SECTIONS
    assert ("Launch QC Tool", "launch") in COMMON_TASKS
    assert ("Use Desktop Office focus", "focus") in COMMON_TASKS
    source = inspect.getsource(render_guide)
    for text in (
        "python -m qc_tool",
        "shortcut install",
        "--desktop-focus",
        "Run IDs are permanent",
        "Bind",
        "Confirm binding",
    ):
        assert text in source


def test_guide_documents_data_directory_defaults_and_restart_boundary() -> None:
    source = inspect.getsource(render_guide)

    for text in (
        "%LOCALAPPDATA%\\\\qc-tool",
        "$XDG_DATA_HOME/qc-tool",
        "~/.local/share/qc-tool",
        "&lt;repository&gt;/data",
        "python -m qc_tool --data-dir",
        "cannot be switched from Local app settings",
        "does not migrate or delete the old directory",
        "Changing the path does not migrate history",
    ):
        assert text in source


def test_review_group_rows_do_not_embed_atomic_member_payloads(qc_result) -> None:
    groups = build_review_groups(qc_result.findings)

    rows = _review_group_rows(groups)

    assert rows
    assert sum(group.member_count for group in groups) == len(qc_result.findings)
    assert all("member_ids" not in row and "excerpts" not in row for row in rows)


def _review_group(
    *members: Finding,
    finding_class: FindingClass = FindingClass.VALUE_CHANGED,
) -> ReviewGroup:
    return ReviewGroup(
        group_id="G001",
        finding_class=finding_class,
        severity=Severity.CRITICAL,
        artifact="excel",
        sheet="Data",
        slide=None,
        element="cell",
        expected_growth=False,
        ranges=("A1",),
        bounding_range="A1",
        baseline_ranges=(),
        baseline_bounding_range="",
        baseline_mixed=False,
        members=tuple(members),
        spatial=False,
    )


def _filter_row(
    row_id: str,
    *,
    severity: str = "critical",
    finding_class: str = "value_changed",
    review_state: str = "needs_review",
) -> dict[str, object]:
    return {
        "id": row_id,
        "severity": severity,
        "class": finding_class,
        "review_state": review_state,
        "where": "Data",
        "location": "A1",
        "message": "changed",
    }


def test_review_filters_compose_class_state_text_and_story_with_and_semantics() -> None:
    rows = [
        _filter_row("G1", review_state="reviewed"),
        _filter_row("G2"),
        _filter_row("G3", finding_class="formula_error"),
        {
            **_filter_row("G4"),
            "where": "Other",
            "message": "different",
        },
    ]

    filtered = _filter_review_rows(
        rows,
        severities={"critical"},
        finding_classes={"value_changed"},
        review_state="needs_review",
        needle="A1 changed",
        story_scope={"G2", "G4"},
    )

    assert [row["id"] for row in filtered] == ["G2"]


def test_empty_class_filter_hides_all_review_rows() -> None:
    assert _filter_review_rows(
        [_filter_row("G1")],
        severities={"critical"},
        finding_classes=set(),
        review_state="all",
    ) == []


def test_cap_only_group_is_visible_only_in_all_review_state() -> None:
    cap = Finding(
        finding_id="F1",
        artifact="run",
        finding_class=FindingClass.FINDINGS_CAPPED,
        message="additional findings omitted",
    )
    row = _review_group_rows(
        [_review_group(cap, finding_class=FindingClass.FINDINGS_CAPPED)]
    )[0]

    assert row["reviewable_members"] == 0
    assert row["reviewed"] == 0
    assert row["review_state"] == "all"
    assert _filter_review_rows(
        [row],
        severities={"critical"},
        finding_classes={FindingClass.FINDINGS_CAPPED.value},
        review_state="all",
    ) == [row]
    assert _filter_review_rows(
        [row],
        severities={"critical"},
        finding_classes={FindingClass.FINDINGS_CAPPED.value},
        review_state="needs_review",
    ) == []
    assert _filter_review_rows(
        [row],
        severities={"critical"},
        finding_classes={FindingClass.FINDINGS_CAPPED.value},
        review_state="reviewed",
    ) == []


def test_partial_group_remains_needs_review_and_cap_members_do_not_count() -> None:
    reviewed = Finding(
        finding_id="F1",
        artifact="excel",
        finding_class=FindingClass.VALUE_CHANGED,
        severity=Severity.CRITICAL,
        severity_overridden=True,
        message="reviewed",
    )
    pending = Finding(
        finding_id="F2",
        artifact="excel",
        finding_class=FindingClass.VALUE_CHANGED,
        severity=Severity.CRITICAL,
        message="pending",
    )
    cap = Finding(
        finding_id="F3",
        artifact="run",
        finding_class=FindingClass.FINDINGS_CAPPED,
        message="omitted",
    )

    partial = _review_group_rows([_review_group(reviewed, pending, cap)])[0]
    complete = _review_group_rows([_review_group(reviewed, cap)])[0]

    assert partial["reviewable_members"] == 2
    assert partial["reviewed"] == 1
    assert partial["review_state"] == "needs_review"
    assert complete["reviewable_members"] == 1
    assert complete["reviewed"] == 1
    assert complete["review_state"] == "reviewed"


def test_unreviewed_only_group_action_preserves_existing_decisions() -> None:
    reviewed = Finding(
        finding_id="F1",
        artifact="excel",
        finding_class=FindingClass.VALUE_CHANGED,
        severity=Severity.INFO,
        severity_overridden=True,
        analyst_comment="already checked",
        message="reviewed",
    )
    pending = Finding(
        finding_id="F2",
        artifact="excel",
        finding_class=FindingClass.VALUE_CHANGED,
        severity=Severity.CRITICAL,
        message="pending",
    )

    updates = apply_group_review(
        _review_group(reviewed, pending),
        severity=Severity.WARNING,
        comment="",
        replace_existing=False,
    )

    assert [update.finding_id for update in updates] == ["F2"]
    assert reviewed.severity is Severity.INFO
    assert reviewed.analyst_comment == "already checked"
    assert pending.severity is Severity.WARNING
    assert pending.severity_overridden is True


def test_managed_names_cannot_escape_storage(tmp_path: Path) -> None:
    assert _safe_upload_name("../../quarterly.xlsx") == "quarterly.xlsx"
    assert _safe_upload_name(r"..\..\quarterly.xlsx") == "quarterly.xlsx"
    assert _profile_path(tmp_path, "monthly-pack") == tmp_path / "monthly-pack.yaml"
    with pytest.raises(ValueError, match="profile name"):
        _profile_path(tmp_path, "../../outside")


def test_terminal_request_files_restore_only_from_managed_uploads(
    tmp_path: Path,
) -> None:
    uploads = tmp_path / "uploads"
    baseline = uploads / "baseline_excel" / "baseline.xlsx"
    member = uploads / "current_excel-ops" / "current.xlsx"
    baseline.parent.mkdir(parents=True)
    member.parent.mkdir(parents=True)
    baseline.write_bytes(b"baseline")
    member.write_bytes(b"current")

    restored, unavailable = _managed_request_paths(
        uploads,
        {
            "baseline_excel": "baseline.xlsx",
            "current_excel:ops": "current.xlsx",
            "current_ppt": "missing.pptx",
            "baseline_ppt": "../outside.pptx",
            "current_excel:Bad ID": "invalid.xlsx",
        },
    )

    assert restored == {
        "baseline_excel": baseline.resolve(),
        "current_excel:ops": member.resolve(),
    }
    assert unavailable == (
        "baseline_ppt",
        "current_excel:Bad ID",
        "current_ppt",
    )


def test_storage_secret_is_private_and_stable(tmp_path: Path) -> None:
    first = _storage_secret(tmp_path)
    second = _storage_secret(tmp_path)

    assert first == second
    assert len(first) >= 32
    if os.name == "posix":
        assert (tmp_path / ".nicegui-storage-secret").stat().st_mode & 0o777 == 0o600


def test_existing_managed_tree_is_migrated_private(tmp_path: Path) -> None:
    nested = tmp_path / "runs" / "old"
    nested.mkdir(parents=True)
    artifact = nested / "report.html"
    artifact.write_text("private", encoding="utf-8")
    if os.name == "posix":
        nested.chmod(0o775)
        artifact.chmod(0o644)

    secure_managed_tree(tmp_path)

    if os.name == "posix":
        assert nested.stat().st_mode & 0o777 == 0o700
        assert artifact.stat().st_mode & 0o777 == 0o600


def test_profile_listing_and_loading(tmp_path: Path) -> None:
    profiles_dir = tmp_path / "profiles"
    assert list_profiles(profiles_dir) == ["default"]
    save_profile(DeliverableProfile(name="monthly-pack"), profiles_dir / "monthly-pack.yaml")
    assert list_profiles(profiles_dir) == ["default", "monthly-pack"]
    loaded = load_profile_by_name(profiles_dir, "monthly-pack")
    assert loaded.name == "monthly-pack"
    assert load_profile_by_name(profiles_dir, "default").name == "default"


def test_mode_file_requirements(tmp_path: Path) -> None:
    files = {
        "baseline_excel": tmp_path / "base.xlsx",
        "current_excel": tmp_path / "current.xlsx",
        "current_ppt": tmp_path / "current.pptx",
    }
    assert set(_files_for_mode(QCRunMode.CYCLE_COMPARISON, files)) == {
        "baseline_excel",
        "current_excel",
    }
    assert set(_files_for_mode(QCRunMode.CURRENT_FILE_PREFLIGHT, files)) == {
        "current_excel",
        "current_ppt",
    }
    assert set(_files_for_mode(QCRunMode.FINAL_PACKAGE, files)) == {
        "current_excel",
        "current_ppt",
    }
    with pytest.raises(ValueError, match="both current"):
        _files_for_mode(
            QCRunMode.FINAL_PACKAGE, {"current_excel": tmp_path / "current.xlsx"}
        )


def test_confirmed_mapping_persists_to_named_profile(tmp_path: Path) -> None:
    profiles_dir = tmp_path / "profiles"
    profiles_dir.mkdir()
    save_profile(DeliverableProfile(name="monthly"), profiles_dir / "monthly.yaml")
    suggestion = MappingSuggestion(
        slide="Summary",
        line="Revenue $1.2M",
        line_skeleton="Revenue $#M",
        figure_index=0,
        figure_raw="$1.2M",
    )
    candidate = SuggestedSource(
        sheet="Dashboard",
        cell="B2",
        value=1_200_000,
        display_match=True,
        label_score=100,
        rel_diff=0,
        row_label="Revenue",
        column_label="Jun-26",
    )

    mapping = persist_confirmed_mapping(
        profiles_dir, "monthly", suggestion, candidate
    )

    assert mapping.source_sheet == "Dashboard" and mapping.source_cell == "B2"
    stored = load_profile_by_name(profiles_dir, "monthly")
    assert stored.crosscheck.mappings == [mapping]
    with pytest.raises(ValueError, match="named profile"):
        persist_confirmed_mapping(profiles_dir, "default", suggestion, candidate)


def test_confirmed_mapping_persists_source_member(tmp_path: Path) -> None:
    profiles_dir = tmp_path / "profiles"
    profiles_dir.mkdir()
    save_profile(DeliverableProfile(name="monthly"), profiles_dir / "monthly.yaml")
    suggestion = MappingSuggestion(
        slide="Summary",
        line="Revenue $1.2M",
        line_skeleton="Revenue $#M",
        figure_index=0,
        figure_raw="$1.2M",
    )
    candidate = SuggestedSource(
        sheet="Dashboard",
        cell="B2",
        value=1_200_000,
        display_match=True,
        label_score=100,
        rel_diff=0,
        source_member="ops",
    )

    mapping = persist_confirmed_mapping(
        profiles_dir,
        "monthly",
        suggestion,
        candidate,
    )

    assert mapping.source_member == "ops"
    assert load_profile_by_name(
        profiles_dir,
        "monthly",
    ).crosscheck.mappings[0].source_member == "ops"


def test_perform_run_produces_artifacts(fixture_dir: Path, tmp_path: Path) -> None:
    work_dir = tmp_path / "work"
    files = {
        "baseline_excel": fixture_dir / "baseline.xlsx",
        "current_excel": fixture_dir / "current.xlsx",
        "baseline_ppt": fixture_dir / "baseline.pptx",
        "current_ppt": fixture_dir / "current.pptx",
    }
    artifacts = perform_run(work_dir, files, {}, fixture_profile())

    assert artifacts.result.counts[Severity.CRITICAL] > 0
    for path in artifacts.report_paths.values():
        assert path.exists() and path.stat().st_size > 0

    history = RunHistory(work_dir / "history.sqlite3")
    record = history.get_run(artifacts.run_id)
    assert record.profile == "fixture"
    assert set(record.file_hashes) == set(files)


def test_perform_run_reports_ordered_progress(
    fixture_dir: Path,
    tmp_path: Path,
) -> None:
    events: list[ProgressEvent] = []

    perform_run(
        tmp_path / "work",
        {"current_excel": fixture_dir / "current.xlsx"},
        {},
        DeliverableProfile(name="progress"),
        mode=QCRunMode.CURRENT_FILE_PREFLIGHT,
        on_progress=events.append,
    )

    completed_phases = [
        event.phase
        for event in events
        if event.total and event.processed == event.total
    ]
    # Reports are on-demand by default, so the phase completes instantly
    # but still appears in order for progress consumers.
    assert completed_phases == [
        RunPhase.PREPARING,
        RunPhase.LOADING_CURRENT_EXCEL,
        RunPhase.ANALYZING_EXCEL,
        RunPhase.WRITING_REPORTS,
        RunPhase.RECORDING_HISTORY,
        RunPhase.COMPLETE,
    ]


def test_cancelled_report_cleans_owned_directory_and_records_no_history(
    fixture_dir: Path,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    work_dir = tmp_path / "work"
    token = CancellationToken()

    def cancel_after_partial_report(result, path: Path) -> None:
        path.write_text("partial", encoding="utf-8")
        token.cancel()

    monkeypatch.setattr(run_service, "write_excel_report", cancel_after_partial_report)

    with pytest.raises(RunCancelled):
        perform_run(
            work_dir,
            {"current_excel": fixture_dir / "current.xlsx"},
            {},
            DeliverableProfile(name="cancelled"),
            mode=QCRunMode.CURRENT_FILE_PREFLIGHT,
            cancellation_token=token,
            write_reports=True,
        )

    runs_dir = work_dir / "runs"
    assert runs_dir.is_dir()
    assert list(runs_dir.iterdir()) == []
    assert not (work_dir / "history.sqlite3").exists()


def test_pre_cancelled_run_creates_no_artifacts(
    fixture_dir: Path,
    tmp_path: Path,
) -> None:
    token = CancellationToken()
    token.cancel()
    work_dir = tmp_path / "work"

    with pytest.raises(RunCancelled):
        perform_run(
            work_dir,
            {"current_excel": fixture_dir / "current.xlsx"},
            {},
            DeliverableProfile(name="cancelled"),
            mode=QCRunMode.CURRENT_FILE_PREFLIGHT,
            cancellation_token=token,
        )

    assert not (work_dir / "runs").exists()
    assert not (work_dir / "history.sqlite3").exists()


def test_immediate_runs_get_distinct_private_report_paths(
    fixture_dir: Path, tmp_path: Path
) -> None:
    work_dir = tmp_path / "work"
    files = {"current_excel": fixture_dir / "current.xlsx"}
    first = perform_run(
        work_dir,
        files,
        {},
        DeliverableProfile(name="first"),
        mode=QCRunMode.CURRENT_FILE_PREFLIGHT,
        write_reports=True,
    )
    second = perform_run(
        work_dir,
        files,
        {},
        DeliverableProfile(name="second"),
        mode=QCRunMode.CURRENT_FILE_PREFLIGHT,
        write_reports=True,
    )

    assert first.report_paths != second.report_paths
    assert first.report_paths["html"].parent != second.report_paths["html"].parent
    if os.name == "posix":
        assert first.report_paths["html"].stat().st_mode & 0o777 == 0o600
        assert first.report_paths["html"].parent.stat().st_mode & 0o777 == 0o700


def test_perform_run_with_encrypted_current(
    fixture_dir: Path, tmp_path: Path, manifest
) -> None:
    files = {
        "baseline_excel": fixture_dir / "baseline.xlsx",
        "current_excel": fixture_dir / "current_encrypted.xlsx",
    }
    artifacts = perform_run(
        tmp_path / "work",
        files,
        {"current_excel": manifest.password},
        DeliverableProfile(name="enc-test"),
    )
    assert artifacts.result.counts[Severity.CRITICAL] > 0


@pytest.mark.asyncio
async def test_main_page_renders(user: User, tmp_path: Path) -> None:
    create_pages(tmp_path / "work")
    await user.open("/")
    await user.should_see("QC Tool")
    await user.should_see("Run QC")
    await user.should_see("Deliverable profile")
    await user.should_see("Override workbook workload refusals")
    await user.should_see("Force full dependency indexing")
    await user.should_see("Current-file preflight")
    await user.should_see("Cycle comparison")
    await user.should_see("Final-package QC")
    await user.should_see("Guide")
    await user.should_see("How to choose a mode")


def test_queue_status_line_shows_state_without_source_paths() -> None:
    created = dt.datetime.now(dt.UTC) - dt.timedelta(seconds=12)
    record = RunStateRecord(
        request_id="0123456789abcdef",
        created_at=created,
        status=RunStatus.RUNNING,
        mode=QCRunMode.CYCLE_COMPARISON.value,
        profile="fixture",
        files={"current_excel": "current.xlsx"},
        phase=RunPhase.COMPARING_FORMULAS.value,
        processed=3,
        total=7,
    )

    line = _queue_status_line(record)

    assert "#01234567" in line
    assert "Running" in line
    assert "Comparing formulas (3/7)" in line
    assert "s elapsed" in line
    assert "current.xlsx" in line
    assert "\\" not in line and "/current.xlsx" not in line


def test_queue_status_line_reports_position_for_queued_requests() -> None:
    record = RunStateRecord(
        request_id="fedcba9876543210",
        created_at=dt.datetime.now(dt.UTC),
        status=RunStatus.QUEUED,
        mode=QCRunMode.CYCLE_COMPARISON.value,
        profile="fixture",
        queue_position=2,
    )

    assert "Queued" in _queue_status_line(record)
    assert "position 2" in _queue_status_line(record)


def test_terminal_request_summary_explains_blocked_attempt_without_a_run() -> None:
    record = RunStateRecord(
        request_id="0123456789abcdef",
        created_at=dt.datetime.now(dt.UTC),
        status=RunStatus.BLOCKED,
        mode=QCRunMode.CYCLE_COMPARISON.value,
        profile="fixture",
        action_required={"reason": "row_identity_confirmation_required"},
    )

    summary = _terminal_request_summary(record)

    assert "paused for row matching" in summary
    assert "No completed run was recorded" in summary
    assert "return to setup" in summary
    assert "row-matching dialog" not in summary


def test_terminal_request_summary_explains_complexity_failure_without_details() -> None:
    record = RunStateRecord(
        request_id="fedcba9876543210",
        created_at=dt.datetime.now(dt.UTC),
        status=RunStatus.FAILED,
        mode=QCRunMode.CYCLE_COMPARISON.value,
        profile="fixture",
        phase=RunPhase.COMPARING_FORMULAS.value,
        error="WorkbookComplexityError: private detail must not render",
    )

    summary = _terminal_request_summary(record)

    assert "Comparing formulas" in summary
    assert "Override workbook workload refusals" in summary
    assert "No completed run was recorded" in summary
    assert "private detail" not in summary


@pytest.mark.asyncio
async def test_active_request_hides_terminal_actions_and_stale_discard_is_refused(
    user: User, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    work_dir = tmp_path / "work"
    store = RunStateStore(work_dir / "history.sqlite3")
    old = store.enqueue(
        "old-terminal-attempt",
        mode=QCRunMode.CYCLE_COMPARISON.value,
        profile="default",
        files={},
        queue_position=0,
    )
    store.finish(old.request_id, RunStatus.FAILED, error="safe failure")
    manager = _RecordingQueueManager(store)
    monkeypatch.setattr(app_module, "get_manager", lambda _work_dir: manager)
    create_pages(work_dir)

    await user.open("/")
    await user.should_see("Discard attempt")
    stale_discard = next(
        iter(user.find(kind=ui.button, content="Discard attempt").elements)
    )
    stale_listener = next(
        listener
        for listener in stale_discard._event_listeners.values()
        if listener.type == "click" and listener.handler is not None
    )
    assert stale_listener.handler is not None
    stale_handler = stale_listener.handler
    stale_slot = stale_discard.parent_slot or stale_discard.client.layout.default_slot

    active = store.enqueue(
        "new-active-attempt",
        mode=QCRunMode.CYCLE_COMPARISON.value,
        profile="default",
        files={},
        queue_position=1,
    )
    await asyncio.sleep(0.6)
    await user.should_not_see("Discard attempt")

    with stale_slot:
        stale_handler(
            events.GenericEventArguments(
                sender=stale_discard,
                client=stale_discard.client,
                args={},
            )
        )
    assert store.get(old.request_id) is not None

    store.finish(active.request_id, RunStatus.CANCELLED)
    await user.should_see("Discard attempt", retries=20)
    await user.should_see("Last attempt #new-acti was cancelled", retries=20)


@pytest.mark.asyncio
async def test_active_request_becoming_blocked_opens_setup_prompt(
    user: User, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    work_dir = tmp_path / "work"
    store = RunStateStore(work_dir / "history.sqlite3")
    active = store.enqueue(
        "active-then-blocked",
        mode=QCRunMode.CYCLE_COMPARISON.value,
        profile="default",
        files={},
        queue_position=0,
    )
    manager = _RecordingQueueManager(store)
    monkeypatch.setattr(app_module, "get_manager", lambda _work_dir: manager)
    create_pages(work_dir)

    await user.open("/")
    await user.should_not_see("QC needs more setup")
    store.finalize_blocked(
        active.request_id,
        {
            "version": 2,
            "reason": "row_identity_confirmation_required",
            "message": "Review row matching.",
            "items": [
                {
                    "member_id": "primary",
                    "sheet": "Data",
                    "cell": "A1",
                    "ranked_table_evidence": ranked_table_evidence_payload_v2(),
                }
            ],
        },
    )

    await user.should_see("QC needs more setup", retries=20)
    await user.should_see("Return to setup", retries=20)


def test_old_typed_ranked_action_requires_a_detector_refresh() -> None:
    evidence = ranked_table_evidence_payload_v2()
    current_action: dict[str, object] = {
        "items": [{"ranked_table_evidence": evidence}]
    }
    old_evidence = dict(evidence)
    old_evidence.pop("manual_review")
    old_action: dict[str, object] = {
        "items": [{"ranked_table_evidence": old_evidence}]
    }

    assert not _ranked_action_needs_refresh(current_action)
    assert _ranked_action_needs_refresh(old_action)
    assert not _ranked_action_needs_refresh(
        {"items": [{"suggested_identity_columns": ["B"]}]}
    )


def test_rerun_profile_choice_restores_an_unsaved_snapshot() -> None:
    snapshot = DeliverableProfile(name="default (temporary)")
    record = RunRecord(
        run_id=1,
        started_at=dt.datetime.now(dt.UTC),
        profile=snapshot.name,
        mode=QCRunMode.CYCLE_COMPARISON,
        files={},
        file_hashes={},
        counts={},
        review_counts={},
        disclosures=[],
        verified_crosschecks=0,
        report_paths={},
        profile_snapshot=snapshot,
    )

    options, selected, override = _rerun_profile_choice(record, ["default"])

    assert options == ["default", "default (temporary)"]
    assert selected == "default (temporary)"
    assert override == snapshot
    assert override is not snapshot


def test_rerun_profile_choice_prefers_a_current_saved_profile() -> None:
    snapshot = DeliverableProfile(name="saved")
    record = RunRecord(
        run_id=1,
        started_at=dt.datetime.now(dt.UTC),
        profile="saved",
        mode=QCRunMode.CYCLE_COMPARISON,
        files={},
        file_hashes={},
        counts={},
        review_counts={},
        disclosures=[],
        verified_crosschecks=0,
        report_paths={},
        profile_snapshot=snapshot,
    )

    options, selected, override = _rerun_profile_choice(
        record, ["default", "saved"]
    )

    assert options == ["default", "saved"]
    assert selected == "saved"
    assert override is None


def test_migrated_default_temporary_profile_recovers_without_yaml(
    tmp_path: Path,
) -> None:
    profile = _temporary_row_matching_base(
        tmp_path,
        "default (temporary)",
        None,
    )

    assert profile.name == "default"
    assert _temporary_profile_name(profile.name) == "default - temporary"
    assert _temporary_profile_name("default (temporary)") == "default - temporary"
    assert _temporary_profile_name("default - temporary") == "default - temporary"
    assert _profile_path(tmp_path, _temporary_profile_name(profile.name)).name == (
        "default - temporary.yaml"
    )


def test_missing_non_default_temporary_profile_has_targeted_error(
    tmp_path: Path,
) -> None:
    with pytest.raises(ValueError, match="no longer available"):
        _temporary_row_matching_base(tmp_path, "missing (temporary)", None)


def test_temporary_profile_snapshot_takes_precedence_over_disk(
    tmp_path: Path,
) -> None:
    snapshot = DeliverableProfile(name="unsaved (temporary)")

    recovered = _temporary_row_matching_base(
        tmp_path,
        snapshot.name,
        snapshot,
    )

    assert recovered == snapshot
    assert recovered is not snapshot


@pytest.mark.asyncio
async def test_run_page_reconnects_to_queue_state_from_another_tab(
    user: User, tmp_path: Path
) -> None:
    work_dir = tmp_path / "work"
    create_pages(work_dir)
    # A request submitted elsewhere: this page must show it, not resubmit it.
    RunStateStore(work_dir / "history.sqlite3").enqueue(
        "abcdef0123456789",
        mode=QCRunMode.CYCLE_COMPARISON.value,
        profile="fixture",
        files={"current_excel": "current.xlsx"},
        queue_position=1,
    )

    await user.open("/")

    await user.should_see("#abcdef01")
    await user.should_see("Cancel #abcdef01")


@pytest.mark.asyncio
async def test_main_page_periodically_evicts_unsubscribed_terminal_setup_jobs(
    user: User, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    class MaintenanceCoordinator:
        shutdown_hook_installed = False

        def __init__(self) -> None:
            self.evictions = 0

        def shutdown(self) -> None:
            pass

        def evict_terminal(self, *, max_age_seconds: float) -> int:
            assert max_age_seconds == 60.0
            self.evictions += 1
            return 0

    coordinator = MaintenanceCoordinator()
    monkeypatch.setattr(
        app_module, "get_setup_coordinator", lambda _work_dir: coordinator
    )
    create_pages(tmp_path / "work")

    await user.open("/")
    await asyncio.sleep(0.6)

    assert coordinator.evictions >= 1


@pytest.mark.asyncio
async def test_complexity_failure_offers_an_explicit_override_confirmation(
    user: User, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    work_dir = tmp_path / "work"
    files = _write_managed_retry_pair(work_dir)
    store = RunStateStore(work_dir / "history.sqlite3")
    profile = DeliverableProfile(name="temporary-rule")
    record = store.enqueue(
        "failed-complexity",
        mode=QCRunMode.CYCLE_COMPARISON.value,
        profile=profile.name,
        files={role: path.name for role, path in files.items()},
        queue_position=0,
        profile_snapshot=profile.model_dump(mode="json"),
        requested_output_mode=FindingOutputMode.DECISION.value,
    )
    store.finish(
        record.request_id,
        RunStatus.FAILED,
        error="WorkbookComplexityError: private detail",
    )
    manager = _RecordingQueueManager(store)
    monkeypatch.setattr(app_module, "get_manager", lambda _work_dir: manager)
    monkeypatch.setattr(app_module, "project_cycle_volume", lambda *_args: None)
    create_pages(work_dir)

    await user.open("/")
    await user.should_see("workbook complexity exceeded the safety limit")
    user.find("Run with override").click()

    await user.should_see("Workbook safety limit reached")
    await user.should_see("substantially more memory and time")
    await user.should_see("source files remain read-only")
    confirm = next(
        button
        for button in user.find(kind=ui.button).elements
        if button.text == "Run with override" and "runbtn" in button.classes
    )
    _emit(confirm, "click", {})
    await user.should_see("Configure & run")

    assert manager.submitted == []
    file_hashes = {role: sha256_file(path) for role, path in files.items()}
    config = ConfigSessionStore(work_dir / "history.sqlite3").latest_for_source_set(
        file_hashes
    )
    assert config is not None
    assert config.choices["allow_large_workbooks"] is True
    assert config.choices["profile_snapshot"] == profile.model_dump(mode="json")


@pytest.mark.asyncio
async def test_stale_row_suggestions_refresh_with_restored_request_context(
    user: User, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    work_dir = tmp_path / "work"
    files = _write_managed_retry_pair(work_dir)
    store = RunStateStore(work_dir / "history.sqlite3")
    profile = DeliverableProfile(name="temporary-rule")
    record = store.enqueue(
        "stale-row-suggestions",
        mode=QCRunMode.CYCLE_COMPARISON.value,
        profile=profile.name,
        files={role: path.name for role, path in files.items()},
        queue_position=0,
        profile_snapshot=profile.model_dump(mode="json"),
        requested_output_mode=FindingOutputMode.ATOMIC.value,
    )
    old_evidence = dict(ranked_table_evidence_payload_v2())
    old_evidence.pop("manual_review")
    store.finalize_blocked(
        record.request_id,
        {
            "version": 2,
            "reason": "row_identity_confirmation_required",
            "items": [
                {
                    "member_id": "primary",
                    "sheet": "Panel",
                    "cell": "A1",
                    "ranked_table_evidence": old_evidence,
                }
            ],
        },
    )
    manager = _RecordingQueueManager(store)
    monkeypatch.setattr(app_module, "get_manager", lambda _work_dir: manager)
    monkeypatch.setattr(app_module, "project_cycle_volume", lambda *_args: None)
    create_pages(work_dir)

    await user.open("/")
    await user.should_see("QC needs more setup")
    user.find(kind=ui.button, content="Return to setup").click()
    await user.should_see("Configure & run")

    assert manager.submitted == []
    file_hashes = {role: sha256_file(path) for role, path in files.items()}
    config = ConfigSessionStore(work_dir / "history.sqlite3").latest_for_source_set(
        file_hashes
    )
    assert config is not None
    assert config.choices["output_mode"] == FindingOutputMode.ATOMIC.value
    assert config.choices["profile_snapshot"] == profile.model_dump(mode="json")


@pytest.mark.asyncio
async def test_run_qc_button_is_blocked_with_no_files_selected(
    user: User, tmp_path: Path
) -> None:
    work_dir = tmp_path / "work"
    create_pages(work_dir)

    await user.open("/")
    user.find(marker="run-qc-button").click()

    await user.should_see("Upload a current Excel workbook")
    # No configuration session was created for a blocked click.
    store = ConfigSessionStore(work_dir / "history.sqlite3")
    assert store.get(session_key_for({})) is None


@pytest.mark.asyncio
async def test_run_qc_button_creates_a_configuration_session(
    user: User, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    work_dir = tmp_path / "work"
    files = _write_managed_retry_pair(work_dir)
    store = RunStateStore(work_dir / "history.sqlite3")
    profile = DeliverableProfile(name="temporary-rule")
    record = store.enqueue(
        "stale-row-suggestions",
        mode=QCRunMode.CYCLE_COMPARISON.value,
        profile=profile.name,
        files={role: path.name for role, path in files.items()},
        queue_position=0,
        profile_snapshot=profile.model_dump(mode="json"),
        requested_output_mode=FindingOutputMode.ATOMIC.value,
    )
    old_evidence = dict(ranked_table_evidence_payload_v2())
    old_evidence.pop("manual_review")
    store.finalize_blocked(
        record.request_id,
        {
            "version": 2,
            "reason": "row_identity_confirmation_required",
            "items": [
                {
                    "member_id": "primary",
                    "sheet": "Panel",
                    "cell": "A1",
                    "ranked_table_evidence": old_evidence,
                }
            ],
        },
    )
    manager = _RecordingQueueManager(store)
    monkeypatch.setattr(app_module, "get_manager", lambda _work_dir: manager)
    monkeypatch.setattr(app_module, "project_cycle_volume", lambda *_args: None)
    create_pages(work_dir)

    await user.open("/")
    await user.should_see("QC needs more setup")
    # Populate state.files/state.file_hashes the same established way this
    # file's own row-suggestions-refresh test does, without depending on a
    # real browser upload widget (none exists in this test suite).
    user.find(kind=ui.button, content="Return to setup").click()
    await user.should_see("Configure & run")

    file_hashes = {role: sha256_file(path) for role, path in files.items()}
    config_store = ConfigSessionStore(work_dir / "history.sqlite3")
    record = config_store.latest_for_source_set(file_hashes)
    assert record is not None
    assert record.choices["mode"] == QCRunMode.CYCLE_COMPARISON.value
    assert record.choices["files"] == {
        role: str(path.resolve()) for role, path in files.items()
    }


@pytest.mark.asyncio
async def test_back_to_files_can_discard_setup_and_keep_files_and_scope(
    user: User, tmp_path: Path
) -> None:
    from qc_tool.ui.config_workspace import build_session_choices

    work_dir = tmp_path / "work"
    files = _write_valid_managed_retry_pair(work_dir)
    file_hashes = {role: sha256_file(path) for role, path in files.items()}
    choices = build_session_choices(
        mode=QCRunMode.CYCLE_COMPARISON,
        profile_name="default",
        files={role: str(path) for role, path in files.items()},
        file_hashes=file_hashes,
        output_mode=FindingOutputMode.DECISION.value,
        allow_large_workbooks=False,
        allow_dependency_indexing=False,
        acceptance_absolute=0.0,
        acceptance_percent=0.0,
        rerun_of=None,
        selected_sheets=("Data",),
    )
    config_store = ConfigSessionStore(work_dir / "history.sqlite3")
    original = config_store.create_session(
        file_hashes=file_hashes,
        profile_name="default",
        choices=choices,
    )
    create_pages(work_dir)

    await user.open(f"/?config_session={original.session_id}")
    await user.should_see("Your in-progress setup is restored")
    await user.should_see("Resume setup")
    await user.should_see("Discard setup")
    source = inspect.getsource(app_module.create_pages)
    assert source.index('classes("readybar")') < source.index(
        'mark("resumed-setup-controls")'
    )

    user.find(kind=ui.button, content="Discard setup").click()

    await user.should_see("Setup discarded; selected files remain on this page.")
    assert config_store.get(original.session_id) is None
    user.find(marker="run-qc-button").click()
    await user.should_see("Configure & run")

    replacement = config_store.latest_for_source_set(file_hashes)
    assert replacement is not None
    assert replacement.session_id != original.session_id
    assert replacement.choices["selected_sheets"] == ["Data"]



@pytest.mark.asyncio
async def test_blocked_row_matching_opens_persistent_setup_prompt(
    user: User, tmp_path: Path
) -> None:
    work_dir = tmp_path / "work"
    store = RunStateStore(work_dir / "history.sqlite3")
    profile = DeliverableProfile(name="default")
    record = store.enqueue(
        "blocked-row-matching",
        mode=QCRunMode.CYCLE_COMPARISON.value,
        profile=profile.name,
        files={"baseline_excel": "baseline.xlsx", "current_excel": "current.xlsx"},
        queue_position=0,
        profile_snapshot=profile.model_dump(mode="json"),
        requested_output_mode=FindingOutputMode.DECISION.value,
    )
    evidence_a = ranked_table_evidence_payload_v2(
        sheet="Panel A",
        current_range="A1:E6001",
        header_row=None,
        manual_review=True,
        column_headers=("", "", "", "", ""),
        formula_ratio=1.0,
    )
    evidence_b = ranked_table_evidence_payload_v2(
        sheet="Panel B",
        current_range="A1:E7001",
        header_row=2,
        column_headers=("Order", "Code", "Value", "Value 2", "Value 3"),
    )
    evidence_c = ranked_table_evidence_payload_v2(
        sheet="Panel C",
        current_range="A1:E8001",
        header_row=4,
        column_headers=("Position", "Key", "Value", "Value 2", "Value 3"),
    )
    store.finalize_blocked(
        record.request_id,
        {
            "version": 2,
            "reason": "row_identity_confirmation_required",
            "message": "Review row matching.",
            "items": [
                {
                    "member_id": "primary",
                    "sheet": evidence["sheet"],
                    "cell": "A1",
                    "label": f"{evidence['sheet']}!{evidence['current_range']}",
                    "ranked_table_evidence": evidence,
                }
                for evidence in (evidence_a, evidence_b, evidence_c)
            ],
        },
    )
    create_pages(work_dir)

    await user.open("/")
    await user.should_see("QC needs more setup")
    await user.should_see("Your files and previous setup choices are preserved")
    await user.should_see("Return to setup")
    assert all(
        button.text != "Review row matching"
        for button in user.find(kind=ui.button).elements
    )


@pytest.mark.asyncio
async def test_blocked_row_matching_consumes_attempt_and_preserves_existing_setup(
    user: User, tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    from qc_tool.config.resolved_input import (
        ResolvedInputConfigurationV1,
        ResolvedMember,
        ResolvedRegion,
        ResolvedSelector,
        ResolvedSheet,
    )
    from qc_tool.ui.config_workspace import build_session_choices

    work_dir = tmp_path / "work"
    files = _write_valid_managed_retry_pair(work_dir)
    file_hashes = {role: sha256_file(path) for role, path in files.items()}
    profile = DeliverableProfile(name="default")
    resolved = ResolvedInputConfigurationV1(
        mode=QCRunMode.CYCLE_COMPARISON,
        profile_name="default",
        members=(
            ResolvedMember(
                member_id="primary",
                baseline_source_sha256=file_hashes["baseline_excel"],
                current_source_sha256=file_hashes["current_excel"],
                sheets=(
                    ResolvedSheet(
                        sheet_id="primary_data",
                        baseline_sheet_name="Data",
                        current_sheet_name="Data",
                        regions=(
                            ResolvedRegion(
                                region_id="data_a1_b4",
                                baseline_outer_range="A1:B4",
                                current_outer_range="A1:B4",
                            ),
                        ),
                        selectors=(
                            ResolvedSelector(
                                selector_id="scenario",
                                baseline_cell="B1",
                                current_cell="B1",
                                equal=True,
                            ),
                        ),
                    ),
                ),
            ),
        ),
    )
    choices = build_session_choices(
        mode=QCRunMode.CYCLE_COMPARISON,
        profile_name="default",
        files={role: str(path) for role, path in files.items()},
        file_hashes=file_hashes,
        output_mode=FindingOutputMode.DECISION.value,
        allow_large_workbooks=False,
        allow_dependency_indexing=False,
        acceptance_absolute=0.0,
        acceptance_percent=0.0,
        rerun_of=None,
        profile_snapshot=profile.model_dump(mode="json"),
        resolved_input_configuration=resolved.model_dump(mode="json"),
    )
    choices["selectors"] = {
        "primary": [
            {
                "sheet_name": "Data",
                "selector_id": "scenario",
                "label": "Scenario choice",
                "cell": "B1",
                "baseline_cell": "B1",
            }
        ]
    }
    choices["region_decisions"] = {
        "data_a1_b4": {
            "mode": "keyed",
            "identity_columns": ["B"],
            "confirmed": True,
        }
    }
    config_store = ConfigSessionStore(work_dir / "history.sqlite3")
    config_session = config_store.create_session(
        file_hashes=file_hashes,
        profile_name="default",
        choices=choices,
    )

    run_store = RunStateStore(work_dir / "history.sqlite3")
    record = run_store.enqueue(
        "blocked-temporary-row-matching",
        mode=QCRunMode.CYCLE_COMPARISON.value,
        profile="default",
        files={role: path.name for role, path in files.items()},
        queue_position=0,
        profile_snapshot=profile.model_dump(mode="json"),
        requested_output_mode=FindingOutputMode.DECISION.value,
        resolved_input_configuration=resolved.model_dump(mode="json"),
        resolved_input_digest=resolved.canonical_sha256(),
    )
    evidence = ranked_table_evidence_payload_v2(
        sheet="Data",
        current_range="A1:B4",
        data_row_count=3,
        available_columns=("A", "B"),
        column_headers=("ID", "Value"),
        suggested_identity_columns=("A",),
        suggested_ordinal_columns=(),
    )
    run_store.finalize_blocked(
        record.request_id,
        {
            "version": 2,
            "reason": "row_identity_confirmation_required",
            "message": "Review row matching.",
            "items": [
                {
                    "member_id": "primary",
                    "sheet": "Data",
                    "cell": "A1",
                    "ranked_table_evidence": evidence,
                }
            ],
        },
    )
    manager = _RecordingQueueManager(run_store)
    monkeypatch.setattr(app_module, "get_manager", lambda _work_dir: manager)
    create_pages(work_dir)

    await user.open(f"/?config_session={config_session.session_id}")
    await user.should_see("QC needs more setup")
    await user.should_not_see("Discard attempt")
    await user.should_not_see("Resume setup")
    await user.should_not_see("Discard setup")
    assert all(
        button.text != "Review row matching"
        for button in user.find(kind=ui.button).elements
    )
    user.find(kind=ui.button, content="Return to setup").click()

    await user.should_see("Configure & run", retries=20)
    await user.should_see("Analysis complete", retries=50)
    await user.should_see("Scenario choice (B1)", retries=20)
    await user.should_see("Confirm row setup", retries=20)
    assert manager.submitted == []
    recovered = config_store.latest_for_source_set(file_hashes)
    assert recovered is not None
    assert run_store.get(record.request_id) is None
    assert recovered.choices["continued_request_id"] == record.request_id
    assert recovered.choices["profile_name"] == "default"
    profile_snapshot = recovered.choices["profile_snapshot"]
    assert isinstance(profile_snapshot, dict)
    assert profile_snapshot["name"] == "default"
    assert recovered.choices["selectors"] == choices["selectors"]
    pending = recovered.choices["pending_row_matching"]
    assert isinstance(pending, list)
    [proposal] = pending
    assert isinstance(proposal, dict)
    assert proposal["identity_columns"] == ["A"]

    mode_toggle = next(iter(user.find(marker="region-mode-toggle").elements))
    assert isinstance(mode_toggle, ui.toggle)
    assert mode_toggle.value == "keyed"
    identity_select = next(iter(user.find(marker="identity-columns").elements))
    assert isinstance(identity_select, ui.select)
    assert identity_select.value == ["B"]
    user.find(kind=ui.button, content="Confirm row setup").click()
    revised = config_store.latest_for_source_set(file_hashes)
    assert revised is not None
    assert "pending_row_matching" not in revised.choices
    region_choices = revised.choices["region_decisions"]
    assert isinstance(region_choices, dict)
    assert any(
        isinstance(choice, dict)
        and choice.get("mode") == "keyed"
        and choice.get("identity_columns") == ["B"]
        and choice.get("confirmed") is True
        for choice in region_choices.values()
    )

    await user.open(f"/?config_session={recovered.session_id}")
    await user.should_see("Resume setup")
    await user.should_not_see("QC needs more setup")
    assert all(
        button.text != "Discard attempt"
        for button in user.find(kind=ui.button).elements
    )


@pytest.mark.asyncio
async def test_guide_page_renders_packaged_operator_content(
    user: User, tmp_path: Path
) -> None:
    create_pages(tmp_path / "work")
    await user.open("/guide")
    await user.should_see("QC Tool guide")
    await user.should_see("Installed on Windows")
    await user.should_see(r"%LOCALAPPDATA%\qc-tool")
    await user.should_see("Installed on Linux")
    await user.should_see("$XDG_DATA_HOME/qc-tool")
    await user.should_see("~/.local/share/qc-tool")
    await user.should_see("Changing the path does not migrate history")
    await user.should_see("Choose the right QC mode")
    await user.should_see("they do not block read-only QC")
    await user.should_see("Profiles, controls, and waivers")
    await user.should_see("The Files page then offers")
    await user.should_see("Discard setup")
    await user.should_see("Rank/order columns")
    await user.should_see("B · Account ID")
    await user.should_see("The profile editor is intentionally advanced")
    await user.should_see("This is not another row-matching review")
    await user.should_see("Pattern review-item counts are analyst decisions")
    await user.should_see("Grouping never makes an error safer")
    await user.should_see("Scope narrows only Excel and PowerPoint findings")
    await user.should_see("Workbook workload override: use it only after a refusal")
    await user.should_see("ask a senior reviewer")
    await user.should_see("Every finding is retained")
    await user.should_see("operation: subtract")
    await user.should_see("Availability controls blankness only")
    await user.should_see("Coverage and severity")
    await user.should_see("Safeguards are visible")
    await user.should_see("Excel to PowerPoint mappings")
    await user.should_see("Treat the list as a search aid, not as proof of source")
    await user.should_see("Suggestions are not proof")
    await user.should_see("merged or multi-row headers")
    await user.should_see("if no single saved cell is the defensible source")
    await user.should_see("Privacy, sharing, and attestations")
    await user.should_see("CLI and automation")
    await user.should_see("Troubleshooting")
    await user.should_see("Before sign-off")
    await user.should_see("No built-in authentication or TLS")
    # Search, common tasks, and the glossary are packaged and offline.
    await user.should_see("Search the guide")
    await user.should_see("Common tasks")
    await user.should_see("Glossary")
    await user.should_see("pattern group")
    # A worked example with dummy data, and the history shelf actions.
    await user.should_see("Worked example")
    await user.should_see("formula replaced by a constant")
    await user.should_see("Archive before you delete")


def test_workload_override_copy_names_its_full_run_scope() -> None:
    source = inspect.getsource(app_module.create_pages)

    assert "Override workbook workload refusals" in source
    assert "formula-link safety limits" in source
    assert "for every workbook in this run" in source
    assert "large-workbook refusal overridden" not in source


@pytest.mark.asyncio
async def test_history_page_renders(user: User, tmp_path: Path) -> None:
    create_pages(tmp_path / "work")
    await user.open("/history")
    await user.should_see("Run history")
    await user.should_see("No runs recorded yet.")


@pytest.mark.asyncio
async def test_history_page_offers_search_and_filters(
    user: User, fixture_dir: Path, tmp_path: Path
) -> None:
    work_dir = tmp_path / "work"
    perform_run(
        work_dir,
        {
            "baseline_excel": fixture_dir / "baseline.xlsx",
            "current_excel": fixture_dir / "current.xlsx",
        },
        {},
        fixture_profile(),
    )
    create_pages(work_dir)
    await user.open("/history")
    await user.should_see("search run id, profile, or file")
    for label in ("Mode", "Profile", "Capability", "Date"):
        await user.should_see(label)


def test_run_blockers_report_every_missing_role() -> None:
    assert _run_blockers(QCRunMode.CYCLE_COMPARISON, {}) == [
        "Upload a baseline and current Excel pair and/or PowerPoint pair"
    ]
    assert _run_blockers(QCRunMode.FINAL_PACKAGE, {"current_excel": Path("a.xlsx")}) == [
        "Upload both current Excel and current PowerPoint files"
    ]
    ready = {"baseline_excel": Path("b.xlsx"), "current_excel": Path("c.xlsx")}
    assert _run_blockers(QCRunMode.CYCLE_COMPARISON, ready) == []


def test_run_blockers_wait_for_hashes_and_reject_identical_pairs() -> None:
    files = {
        "baseline_excel": Path("baseline.xlsx"),
        "current_excel": Path("current.xlsx"),
    }

    assert _run_blockers(
        QCRunMode.CYCLE_COMPARISON,
        files,
        file_hashes={"baseline_excel": "a"},
    ) == ["Verifying selected file bytes: Current — Excel workbook"]
    assert _run_blockers(
        QCRunMode.CYCLE_COMPARISON,
        files,
        file_hashes={"baseline_excel": "same", "current_excel": "same"},
    ) == [
        "Baseline and current Excel for member 'primary' are byte-identical; "
        "a comparison would prove nothing"
    ]
    assert _run_blockers(
        QCRunMode.CYCLE_COMPARISON,
        files,
        file_hashes={"baseline_excel": "a", "current_excel": "b"},
    ) == []


def test_run_blockers_validate_dynamic_members_and_duplicate_bytes() -> None:
    files = {
        "baseline_excel:core": Path("core-old.xlsx"),
        "current_excel:core": Path("core-new.xlsx"),
        "current_excel:ops": Path("ops.xlsx"),
    }

    blockers = _run_blockers(
        QCRunMode.CYCLE_COMPARISON,
        files,
        file_hashes={
            "baseline_excel:core": "same-core",
            "current_excel:core": "same-core",
            "current_excel:ops": "same-core",
        },
    )

    assert blockers == [
        "Duplicate bytes for excel on the current side: "
        "Current — Excel workbook · core and Current — Excel workbook · ops",
        "Baseline and current Excel for member 'core' are byte-identical; "
        "a comparison would prove nothing",
    ]


def test_files_for_mode_keeps_all_selected_current_workbook_members() -> None:
    files = {
        "current_excel:core": Path("core.xlsx"),
        "current_excel:ops": Path("ops.xlsx"),
        "current_ppt": Path("deck.pptx"),
        "baseline_excel:core": Path("old.xlsx"),
    }

    assert _files_for_mode(QCRunMode.FINAL_PACKAGE, files) == {
        "current_excel:core": Path("core.xlsx"),
        "current_excel:ops": Path("ops.xlsx"),
        "current_ppt": Path("deck.pptx"),
    }


def test_run_blockers_name_the_files_a_rerun_still_needs() -> None:
    blockers = _run_blockers(
        QCRunMode.CYCLE_COMPARISON,
        {"baseline_excel": Path("b.xlsx"), "current_excel": Path("c.xlsx")},
        rerun_of=7,
        rerun_required=frozenset({"baseline_excel", "current_excel", "current_ppt"}),
    )
    assert blockers == [
        "Re-QC of run #7 is blocked until these files are selected again: "
        "Current — PowerPoint deck"
    ]


def test_role_requirement_changes_with_mode() -> None:
    assert _role_requirement(QCRunMode.FINAL_PACKAGE, "current_excel") == "required"
    assert (
        _role_requirement(QCRunMode.CURRENT_FILE_PREFLIGHT, "current_ppt")
        == "Excel and/or PPT"
    )
    assert _role_requirement(QCRunMode.CYCLE_COMPARISON, "baseline_excel") == "Excel pair"
    assert _role_requirement(QCRunMode.CYCLE_COMPARISON, "baseline_ppt") == "PPT pair"


def test_policy_summaries_describe_what_the_engine_receives() -> None:
    state = app_module.SessionState()
    assert _scope_summary(state) == "scope: everything"
    assert _acceptance_summary(state) == "acceptance: strict"
    state.available_sheets = ["A", "B", "C"]
    state.selected_sheets = {"A"}
    state.acceptance_absolute = 2.5
    state.acceptance_percent = 0.1
    assert _scope_summary(state) == "scope: 1/3 sheets"
    assert _acceptance_summary(state) == "acceptance: ±2.5 or ±0.1%"
    # A full selection is not a narrowed scope, so it must not be disclosed as one.
    state.selected_sheets = {"A", "B", "C"}
    assert _scope_summary(state) == "scope: everything"


def test_outcome_summary_leads_with_the_decision() -> None:
    assert _outcome_summary({Severity.CRITICAL: 3, Severity.WARNING: 4})[:2] == (
        "attention",
        "Review required",
    )
    assert _outcome_summary({Severity.WARNING: 2})[:2] == ("limited", "Review required")
    assert _outcome_summary({Severity.INFO: 9})[:2] == ("ok", "No blocking findings")


def test_mapping_stats_reconcile_readable_and_unavailable_surfaces() -> None:
    stats = dict(
        _mapping_stats(
            MappingCoverage(
                eligible=26,
                unavailable=1,
                mapped=1,
                verified=1,
                unmapped=25,
            )
        )
    )

    assert stats["readable"] == 26
    assert stats["unavailable"] == 1
    assert stats["total surfaces"] == 27


def test_capability_summary_separates_omitted_inputs_from_failed_checks() -> None:
    omitted = CoverageItem(
        check_id="ppt-comparison",
        label="PowerPoint comparison",
        artifact="ppt",
        state=CoverageState.NOT_INCLUDED,
    )

    assert _capability_summary([omitted]) == (
        "ok",
        "Included checks ran",
        "1 not included in this run",
    )

    unavailable = CoverageItem(
        check_id="excel-formulas",
        label="Excel formulas",
        artifact="excel",
        state=CoverageState.UNAVAILABLE,
    )
    tone, title, detail = _capability_summary([omitted, unavailable])
    assert tone == "limited"
    assert title == "Capability limited"
    assert "1 unavailable" in detail
    assert "1 not included in this run" in detail


def test_context_grid_html_escapes_source_values() -> None:
    excerpt = GridExcerpt(
        rows=[1],
        cols=["<A>"],
        cells=[["<script>alert(1)</script>"]],
        hit_row=0,
        hit_col=0,
    )
    markup = _context_grid_html(excerpt, "baseline")
    assert "<script>" not in markup
    assert "&lt;script&gt;alert(1)&lt;/script&gt;" in markup
    assert "&lt;A&gt;" in markup
    assert 'class="hit"' in markup


def test_relative_time_keeps_the_exact_stamp_available() -> None:
    now = dt.datetime.now(dt.UTC)
    assert _relative_time(now) == "just now"
    assert _relative_time(now - dt.timedelta(minutes=20)) == "20 min ago"
    assert _relative_time(now - dt.timedelta(hours=5)) == "5 h ago"
    old = now - dt.timedelta(days=400)
    assert _relative_time(old) == old.date().isoformat()


def test_history_row_exposes_mode_profile_capability_and_decisions(
    fixture_dir: Path, tmp_path: Path
) -> None:
    work_dir = tmp_path / "work"
    artifacts = perform_run(
        work_dir,
        {
            "baseline_excel": fixture_dir / "baseline.xlsx",
            "current_excel": fixture_dir / "current.xlsx",
        },
        {},
        fixture_profile(),
    )
    record = RunHistory(work_dir / "history.sqlite3").get_run(artifacts.run_id)

    row = _history_row(record)

    assert row["id"] == artifacts.run_id
    assert row["mode_key"] == QCRunMode.CYCLE_COMPARISON.value
    assert row["profile"] == "fixture"
    assert row["capability"] in {"limited", "complete"}
    decisions = row["decisions"]
    assert isinstance(decisions, list)
    assert decisions and all(set(entry) == {"k", "n"} for entry in decisions)
    assert row["started"] == record.started_at.isoformat(timespec="seconds")
    assert "baseline.xlsx" in str(row["files"])
    size = row["size"]
    assert isinstance(size, int) and size > 0  # measured at record time
    assert str(row["size_label"]).endswith(("B", "KB", "MB", "GB"))


def test_format_bytes_covers_unknown_and_scales() -> None:
    assert _format_bytes(None) == "\u2014"
    assert _format_bytes(512) == "512 B"
    assert _format_bytes(4096) == "4.0 KB"
    assert _format_bytes(5 * 1024 * 1024) == "5.0 MB"
    assert _format_bytes(3 * 1024**3) == "3.0 GB"


def test_storage_prompt_fires_at_threshold_then_waits_for_growth() -> None:
    threshold = 1_000
    assert not _storage_prompt_due(999, threshold=threshold)
    assert _storage_prompt_due(1_000, threshold=threshold)
    # Dismissal silences the prompt until roughly 10% further growth.
    assert not _storage_prompt_due(
        1_050, threshold=threshold, dismissed_at_bytes=1_000
    )
    assert _storage_prompt_due(1_100, threshold=threshold, dismissed_at_bytes=1_000)
    # Cleanup below the threshold silences it entirely.
    assert not _storage_prompt_due(900, threshold=threshold, dismissed_at_bytes=1_000)


def test_history_trend_row_contains_only_fixed_aggregate_fields(
    fixture_dir: Path,
    tmp_path: Path,
) -> None:
    work_dir = tmp_path / "work"
    artifacts = perform_run(
        work_dir,
        {"current_excel": fixture_dir / "current.xlsx"},
        {},
        DeliverableProfile(name="trend"),
        mode=QCRunMode.CURRENT_FILE_PREFLIGHT,
    )
    history = RunHistory(work_dir / "history.sqlite3")

    row = _history_trend_row(history.get_run(artifacts.run_id), history)

    assert set(row) == {
        "run",
        "profile",
        "atomics",
        "decisions",
        "limited",
        "mapped",
        "verified",
        "exact_reusable",
        "carried",
        "finalized",
        "review_minutes",
    }
    assert row["review_minutes"] == "unknown"
    assert not ({"files", "file_paths", "report_paths", "findings"} & row.keys())


@pytest.mark.asyncio
async def test_colophon_is_scoped_to_the_guide(user: User, tmp_path: Path) -> None:
    create_pages(tmp_path / "work")
    await user.open("/")
    await user.should_see("source files are never modified")
    await user.should_not_see("Himanshu")

    await user.open("/guide")
    await user.should_see("Curated by Himanshu")


def test_input_cautions_flag_a_baseline_current_mix_up() -> None:
    state = app_module.SessionState()
    assert _input_cautions(state) == []
    state.files = {
        "baseline_excel": Path("/tmp/a/pack.xlsx"),
        "current_excel": Path("/tmp/b/pack.xlsx"),
    }
    state.file_sizes = {"baseline_excel": 4096, "current_excel": 4096}
    state.file_hashes = {"baseline_excel": "a", "current_excel": "b"}

    cautions = _input_cautions(state)

    assert len(cautions) == 1
    assert "pack.xlsx" in cautions[0]
    assert "prove nothing" in cautions[0]
    # A genuinely different current file is not a caution.
    state.file_sizes["current_excel"] = 5120
    assert _input_cautions(state) == []


def test_input_groups_separate_baseline_from_current() -> None:
    groups = {group: roles for group, _, _, roles in app_module.ROLE_GROUPS}

    assert groups["baseline"] == ("baseline_excel", "baseline_ppt")
    assert groups["current"] == ("current_excel", "current_ppt")
    assert set(groups["baseline"]) | set(groups["current"]) == set(app_module.ROLES)


@pytest.mark.asyncio
async def test_inputs_render_as_baseline_and_current_panels(
    user: User, tmp_path: Path
) -> None:
    app.storage.general.pop("qc_mode", None)
    create_pages(tmp_path / "work")
    await user.open("/")
    await user.should_see("Current")
    await user.should_see("the cycle you are signing off")
    await user.should_not_see("the previous cycle you compare against")

    mode_toggle = next(iter(user.find(marker="mode-toggle").elements))
    model_update = next(
        listener.type
        for listener in mode_toggle._event_listeners.values()
        if listener.handler is not None and listener.type.startswith("update:")
    )
    _emit(mode_toggle, model_update, 1)

    await user.should_see("the previous cycle you compare against")


@pytest.mark.asyncio
async def test_manage_profiles_uses_complete_typed_editor_and_dirty_close(
    user: User,
    tmp_path: Path,
) -> None:
    create_pages(tmp_path / "work")
    await user.open("/")

    user.find("Manage profiles").click()

    await user.should_see("Core contract")
    await user.should_see("Advanced Excel/PPT")
    await user.should_see("Advanced YAML")
    await user.should_see("The built-in default profile is immutable")
    await user.should_not_see("Waivers (YAML list)")
    save_button = user.find("Save profile").elements.pop()
    assert isinstance(save_button, ui.button)
    assert not save_button.enabled

    new_name = next(
        element
        for element in user.find(kind=ui.input).elements
        if element.props.get("label") == "New profile name"
    )
    new_name.value = "typed-contract"
    user.find("Create").click()

    save_button = user.find("Save profile").elements.pop()
    assert isinstance(save_button, ui.button)
    assert save_button.enabled
    user.find("Add Waivers item").click()
    await user.should_see("Finding Class")
    await user.should_see("Unsaved profile changes")

    user.find("Close").click()
    await user.should_see("Discard profile changes?")
    user.find("Discard and close").click()
    dialogs = user.find(kind=ui.dialog).elements
    assert dialogs and all(not dialog.value for dialog in dialogs)


@pytest.mark.asyncio
async def test_shutdown_control_confirms_before_stopping(
    user: User, tmp_path: Path
) -> None:
    create_pages(tmp_path / "work")
    await user.open("/")

    user.find(marker="quit").click()

    await user.should_see("Stop the QC Tool server?")
    await user.should_see("Stop server")
    await user.should_see("Keep running")


@pytest.mark.asyncio
async def test_app_settings_are_discoverable_from_the_header(
    user: User,
    tmp_path: Path,
) -> None:
    create_pages(tmp_path / "work")
    await user.open("/")

    user.find(marker="app-settings").click()

    await user.should_see("Local app settings")
    await user.should_see("Local storage · active")
    await user.should_see("use the QC Tool data directory on Linux and Windows")
    await user.should_see("with --data-dir")
    if os.name == "nt":
        await user.should_see("Desktop Office focus · Off")
        await user.should_see("Desktop shortcut")
    else:
        await user.should_see(
            "Desktop Office focus and Desktop shortcut controls are Windows-only"
        )


@pytest.mark.asyncio
async def test_app_settings_report_and_clear_the_formula_cache(
    user: User,
    tmp_path: Path,
) -> None:
    from qc_tool.io.formula_cache import (
        FormulaCacheKey,
        FormulaExtractionCache,
        coordinate_digest,
    )
    from qc_tool.io.formula_enrichment import FormulaExtraction
    from qc_tool.io.xlsb_formula import XlsbFormulaScan

    work_dir = tmp_path / "work"
    scan = XlsbFormulaScan(formula_cells={"Data": frozenset({(1, 1)})})
    cache = FormulaExtractionCache(work_dir / "formula-cache")
    cache.store(
        FormulaCacheKey(
            package_sha256="a" * 64,
            coordinate_digest=coordinate_digest(scan),
            coordinate_count=scan.formula_count,
            adapter_family="libreoffice",
            adapter_fingerprint="libreoffice:test",
        ),
        FormulaExtraction(
            formulas={"Data": {(1, 1): "=A2+1"}}, engine="test:1.0", detail="test"
        ),
    )
    assert cache.status()["entry_count"] == 1

    create_pages(work_dir)
    await user.open("/")

    user.find(marker="app-settings").click()
    await user.should_see("Formula-extraction cache · 1 entries")
    await user.should_see("Native formula engine")

    user.find("Clear formula cache").click()
    await user.should_see("Formula-extraction cache · 0 entries, 0 B")
    assert cache.status()["entry_count"] == 0


@pytest.mark.asyncio
async def test_compare_page_remembers_manual_mode_selection(
    user: User,
    tmp_path: Path,
) -> None:
    app.storage.general.pop("qc_mode", None)
    create_pages(tmp_path / "work")
    await user.open("/")
    await user.should_see("Current-file preflight — cannot run yet")

    mode_toggle = next(iter(user.find(marker="mode-toggle").elements))
    model_update = next(
        listener.type
        for listener in mode_toggle._event_listeners.values()
        if listener.handler is not None and listener.type.startswith("update:")
    )
    _emit(mode_toggle, model_update, 2)

    assert app.storage.general.get("qc_mode") == QCRunMode.FINAL_PACKAGE.value


@pytest.mark.asyncio
async def test_shutdown_dialog_discloses_unfinished_work(
    user: User, tmp_path: Path
) -> None:
    work_dir = tmp_path / "work"
    create_pages(work_dir)
    RunStateStore(work_dir / "history.sqlite3").enqueue(
        "abcdef0123456789",
        mode=QCRunMode.CYCLE_COMPARISON.value,
        profile="fixture",
        files={"current_excel": "current.xlsx"},
        queue_position=1,
    )

    await user.open("/")
    user.find(marker="quit").click()

    await user.should_see("recorded as unfinished")
    await user.should_see("#abcdef01")


@pytest.mark.asyncio
async def test_lan_mode_shows_exposure_warning(user: User, tmp_path: Path) -> None:
    create_pages(
        tmp_path / "work",
        network_mode=NetworkMode.LAN,
        expires_at=dt.datetime.now(dt.UTC) + dt.timedelta(minutes=10),
    )
    await user.open("/")
    await user.should_see("NETWORK ACCESS ENABLED")
    await user.should_see("no built-in authentication or TLS")


@pytest.mark.asyncio
async def test_dark_mode_toggle_persists(user: User, tmp_path: Path) -> None:

    app.storage.general.pop("dark_mode", None)
    create_pages(tmp_path / "work")
    await user.open("/")
    user.find(marker="dark-toggle").click()
    await user.should_see(marker="dark-toggle")  # page still alive after toggle
    assert app.storage.general.get("dark_mode") is True
    user.find(marker="dark-toggle").click()
    await user.should_see(marker="dark-toggle")
    assert app.storage.general.get("dark_mode") is False


@pytest.mark.asyncio
async def test_run_detail_page(user: User, fixture_dir: Path, tmp_path: Path) -> None:
    work_dir = tmp_path / "work"
    artifacts = perform_run(
        work_dir,
        {
            "baseline_excel": fixture_dir / "baseline.xlsx",
            "current_excel": fixture_dir / "current.xlsx",
        },
        {},
        fixture_profile(),
    )
    create_pages(work_dir)
    await user.open(f"/runs/{artifacts.run_id}")
    await user.should_see(f"Run #{artifacts.run_id}")
    await user.should_see("profile 'fixture'")
    await user.should_see("cross-checks ok")
    await user.should_see("Review queue")
    await user.should_see("Atomic evidence")
    await user.should_see("finding records")
    await user.should_see("represented changes")
    await user.should_see("review time")
    await user.should_see("0:00")
    await user.should_see("Start")
    # Review queue is the default view; evidence tabs are opt-in.
    panels = user.find(kind=ui.tab_panels).elements.pop()
    assert panels.value == "review"


@pytest.mark.asyncio
async def test_first_run_hides_specialist_controls_in_advanced_sections(
    user: User, tmp_path: Path
) -> None:
    create_pages(tmp_path / "work")

    await user.open("/")

    await user.should_see("Advanced finding output")
    await user.should_see("Advanced comparison and safety options")


@pytest.mark.asyncio
async def test_run_detail_page_not_found(user: User, tmp_path: Path) -> None:
    create_pages(tmp_path / "work")
    await user.open("/runs/999")
    await user.should_see("Run not found")


@pytest.mark.asyncio
async def test_population_finding_detail_shows_membership_summary(
    user: User, tmp_path: Path
) -> None:
    """A population's review row renders the new membership evidence panel."""
    from openpyxl import Workbook

    from qc_tool.config.profile import PopulationPolicy, ReviewPolicy

    def build(path: Path, *, multiplier: int) -> None:
        workbook = Workbook()
        sheet = workbook.active
        assert sheet is not None
        sheet.title = "Data"
        sheet.append(["Input", "Output"])
        for row in range(2, 21):  # 19 uniform formula changes
            sheet.append([100, f"=A{row}*{multiplier}"])
        workbook.save(path)

    work_dir = tmp_path / "work"
    baseline = tmp_path / "base.xlsx"
    current = tmp_path / "curr.xlsx"
    build(baseline, multiplier=2)
    build(current, multiplier=3)
    profile = DeliverableProfile(
        name="population-ui",
        review_policy=ReviewPolicy(
            populations=PopulationPolicy(enabled=True, threshold=10)
        ),
    )
    artifacts = perform_run(
        work_dir,
        {"baseline_excel": baseline, "current_excel": current},
        {},
        profile,
    )
    create_pages(work_dir)
    await user.open(f"/runs/{artifacts.run_id}")
    await user.should_see("Review queue")

    group_table = next(
        element
        for element in user.find(kind=ui.table).elements
        if "review-groups-table" in element.classes
    )
    population_row = next(
        row for row in group_table.rows if row.get("class") == "formula_logic_changed"
    )
    _emit(group_table, "select", {"id": str(population_row["id"])})
    await user.should_see("Population membership")
    await user.should_see("19 cells")


@pytest.mark.asyncio
async def test_population_expand_members_pager_pages_every_member(
    user: User, tmp_path: Path
) -> None:
    """"Expand members" decodes the codec and pages through all 19 cells."""
    from openpyxl import Workbook

    from qc_tool.config.profile import PopulationPolicy, ReviewPolicy

    def build(path: Path, *, multiplier: int) -> None:
        workbook = Workbook()
        sheet = workbook.active
        assert sheet is not None
        sheet.title = "Data"
        sheet.append(["Input", "Output"])
        for row in range(2, 21):
            sheet.append([100, f"=A{row}*{multiplier}"])
        workbook.save(path)

    work_dir = tmp_path / "work"
    baseline = tmp_path / "base.xlsx"
    current = tmp_path / "curr.xlsx"
    build(baseline, multiplier=2)
    build(current, multiplier=3)
    profile = DeliverableProfile(
        name="population-pager",
        review_policy=ReviewPolicy(
            populations=PopulationPolicy(enabled=True, threshold=10)
        ),
    )
    artifacts = perform_run(
        work_dir,
        {"baseline_excel": baseline, "current_excel": current},
        {},
        profile,
    )
    create_pages(work_dir)
    await user.open(f"/runs/{artifacts.run_id}")
    group_table = next(
        element
        for element in user.find(kind=ui.table).elements
        if "review-groups-table" in element.classes
    )
    population_row = next(
        row for row in group_table.rows if row.get("class") == "formula_logic_changed"
    )
    _emit(group_table, "select", {"id": str(population_row["id"])})
    await user.should_see("Population membership")

    user.find("Expand members").click()
    await user.should_see("19 total")
    await user.should_see("Showing 1-19 of 19")
    await user.should_see("B2")
    await user.should_see("B20")


def _build_population_fixture(
    tmp_path: Path, *, multiplier_current: int = 3
) -> tuple[Path, Path]:
    from openpyxl import Workbook

    def build(path: Path, *, multiplier: int) -> None:
        workbook = Workbook()
        sheet = workbook.active
        assert sheet is not None
        sheet.title = "Data"
        sheet.append(["Input", "Output"])
        for row in range(2, 21):  # 19 uniform formula changes
            sheet.append([100, f"=A{row}*{multiplier}"])
        workbook.save(path)

    baseline = tmp_path / "base.xlsx"
    current = tmp_path / "curr.xlsx"
    build(baseline, multiplier=2)
    build(current, multiplier=multiplier_current)
    return baseline, current


@pytest.mark.asyncio
async def test_population_load_sample_excerpts_renders_grids(
    user: User, tmp_path: Path
) -> None:
    """"Load sample excerpts" reopens the recorded sources and renders grids."""
    from qc_tool.config.profile import PopulationPolicy, ReviewPolicy

    work_dir = tmp_path / "work"
    baseline, current = _build_population_fixture(tmp_path)
    profile = DeliverableProfile(
        name="population-excerpts",
        review_policy=ReviewPolicy(
            populations=PopulationPolicy(enabled=True, threshold=10)
        ),
    )
    artifacts = perform_run(
        work_dir,
        {"baseline_excel": baseline, "current_excel": current},
        {},
        profile,
    )
    create_pages(work_dir)
    await user.open(f"/runs/{artifacts.run_id}")
    group_table = next(
        element
        for element in user.find(kind=ui.table).elements
        if "review-groups-table" in element.classes
    )
    population_row = next(
        row for row in group_table.rows if row.get("class") == "formula_logic_changed"
    )
    _emit(group_table, "select", {"id": str(population_row["id"])})
    await user.should_see("Population membership")

    user.find("Load sample excerpts").click()
    # Criterion 9: excerpts now load via a disposable child process (real
    # spawn + reimport + workbook load), noticeably slower than the old
    # in-process call -- `should_see`'s default ~0.3s retry budget is not
    # enough headroom for that, so give it more retries here.
    await user.should_see("baseline", retries=50)
    await user.should_see("current", retries=50)
    await user.should_see("B2", retries=50)


@pytest.mark.asyncio
async def test_population_load_sample_excerpts_declines_while_the_slot_is_busy(
    user: User, tmp_path: Path
) -> None:
    """plan-20260913 Step 5: reopening sources for an excerpt shares the
    shared exclusive slot a QC run (or another such load) holds -- it must
    decline with a friendly notice, never crash, while the slot is busy.
    """
    from qc_tool.config.profile import PopulationPolicy, ReviewPolicy

    work_dir = tmp_path / "work"
    baseline, current = _build_population_fixture(tmp_path)
    profile = DeliverableProfile(
        name="population-excerpts-busy",
        review_policy=ReviewPolicy(
            populations=PopulationPolicy(enabled=True, threshold=10)
        ),
    )
    artifacts = perform_run(
        work_dir,
        {"baseline_excel": baseline, "current_excel": current},
        {},
        profile,
    )
    create_pages(work_dir)
    await user.open(f"/runs/{artifacts.run_id}")
    group_table = next(
        element
        for element in user.find(kind=ui.table).elements
        if "review-groups-table" in element.classes
    )
    population_row = next(
        row for row in group_table.rows if row.get("class") == "formula_logic_changed"
    )
    _emit(group_table, "select", {"id": str(population_row["id"])})
    await user.should_see("Population membership")

    slot = get_exclusive_slot(work_dir)
    assert slot.try_acquire("a-qc-run-in-progress")
    try:
        user.find("Load sample excerpts").click()
        await user.should_see("try again once it finishes")
    finally:
        slot.release("a-qc-run-in-progress")


@pytest.mark.asyncio
async def test_population_load_sample_excerpts_discloses_moved_source(
    user: User, tmp_path: Path
) -> None:
    """A moved/deleted recorded source fails closed with a plain disclosure."""
    from qc_tool.config.profile import PopulationPolicy, ReviewPolicy

    work_dir = tmp_path / "work"
    baseline, current = _build_population_fixture(tmp_path)
    profile = DeliverableProfile(
        name="population-excerpts-missing",
        review_policy=ReviewPolicy(
            populations=PopulationPolicy(enabled=True, threshold=10)
        ),
    )
    artifacts = perform_run(
        work_dir,
        {"baseline_excel": baseline, "current_excel": current},
        {},
        profile,
    )
    current.unlink()  # source moves/disappears after the run recorded it

    create_pages(work_dir)
    await user.open(f"/runs/{artifacts.run_id}")
    group_table = next(
        element
        for element in user.find(kind=ui.table).elements
        if "review-groups-table" in element.classes
    )
    population_row = next(
        row for row in group_table.rows if row.get("class") == "formula_logic_changed"
    )
    _emit(group_table, "select", {"id": str(population_row["id"])})
    await user.should_see("Population membership")

    user.find("Load sample excerpts").click()
    await user.should_see("no longer at its recorded location")


@pytest.mark.asyncio
async def test_signoff_dialog_reports_unreviewed_blockers(
    user: User,
    fixture_dir: Path,
    tmp_path: Path,
) -> None:
    work_dir = tmp_path / "work"
    artifacts = perform_run(
        work_dir,
        {"current_excel": fixture_dir / "current.xlsx"},
        {},
        DeliverableProfile(name="default"),
        mode=QCRunMode.CURRENT_FILE_PREFLIGHT,
    )
    create_pages(work_dir)
    await user.open(f"/runs/{artifacts.run_id}")

    user.find("Finalize review").click()

    await user.should_see("Finalize reviewed run")
    await user.should_see("findings still need an analyst decision")


@pytest.mark.asyncio
async def test_finalized_run_renders_frozen_evidence_state(
    user: User,
    fixture_dir: Path,
    tmp_path: Path,
) -> None:
    work_dir = tmp_path / "work"
    artifacts = perform_run(
        work_dir,
        {"current_excel": fixture_dir / "current.xlsx"},
        {},
        DeliverableProfile(name="default"),
        mode=QCRunMode.CURRENT_FILE_PREFLIGHT,
    )
    history = RunHistory(work_dir / "history.sqlite3")
    record = history.get_run(artifacts.run_id)
    history.set_annotations_bulk(
        artifacts.run_id,
        [
            (
                finding.finding_id,
                finding.severity.value if finding.severity else None,
                "reviewed",
            )
            for finding in record.findings
            if finding.severity in {Severity.CRITICAL, Severity.WARNING}
        ],
    )
    record = history.get_run(artifacts.run_id)
    finalize_run(
        work_dir,
        artifacts.run_id,
        set(required_acknowledgements(record)),
    )
    create_pages(work_dir)
    await user.open(f"/runs/{artifacts.run_id}")

    await user.should_see("Finalized")
    await user.should_see("Download attestation")
    await user.should_not_see("Finalize review")
    assert _history_row(history.get_run(artifacts.run_id))["finalized"] is True


@pytest.mark.asyncio
async def test_reqc_offers_explicit_prior_decision_preview(
    user: User,
    tmp_path: Path,
) -> None:
    work_dir = tmp_path / "work"
    history = RunHistory(work_dir / "history.sqlite3")
    finding = Finding(
        finding_id="F1",
        artifact="excel",
        finding_class=FindingClass.VALUE_CHANGED,
        severity=Severity.CRITICAL,
        sheet="Data",
        location="B2",
        baseline_value="1",
        current_value="2",
        message="changed",
    )
    previous = history.record_run(
        QCRunResult(profile_name="fixture", findings=[finding]),
        file_hashes={},
        report_paths={},
    )
    history.set_annotation(previous, "F1", severity="info", comment="reviewed")
    current = history.record_run(
        QCRunResult(
            profile_name="fixture",
            findings=[finding.model_copy(update={"finding_id": "N1"})],
        ),
        file_hashes={},
        report_paths={},
        rerun_of=previous,
    )
    create_pages(work_dir)
    await user.open(f"/runs/{current}")

    user.find("Review prior decisions").click()

    await user.should_see("1 exact reusable")
    await user.should_see("source run was not finalized")
    await user.should_see("N1 · info")


def test_rerun_delta_discloses_atomic_to_decision_mode_change(tmp_path: Path) -> None:
    """Criterion 5: a naive resolved/new comparison across an output-mode
    change compares disjoint identity spaces (location-keyed atomics vs
    shape-digest-keyed populations) -- it must be suppressed and disclosed,
    not silently shown as noise.
    """
    work_dir = tmp_path / "work"
    history = RunHistory(work_dir / "history.sqlite3")
    atomic_findings = [
        Finding(
            finding_id="A1",
            artifact="excel",
            finding_class=FindingClass.FORMULA_LOGIC_CHANGED,
            severity=Severity.WARNING,
            sheet="Data",
            location="B2",
            baseline_value="=A1",
            current_value="=A1+1",
            message="changed",
        ),
        Finding(
            finding_id="A2",
            artifact="excel",
            finding_class=FindingClass.FORMULA_LOGIC_CHANGED,
            severity=Severity.WARNING,
            sheet="Data",
            location="B3",
            baseline_value="=A2",
            current_value="=A2+1",
            message="changed",
        ),
    ]
    previous_id = history.record_run(
        QCRunResult(
            profile_name="fixture",
            findings=atomic_findings,
            requested_output_mode=FindingOutputMode.ATOMIC,
        ),
        file_hashes={},
        report_paths={},
    )
    current_id = history.record_run(
        QCRunResult(
            profile_name="fixture",
            findings=atomic_findings,
            requested_output_mode=FindingOutputMode.DECISION,
        ),
        file_hashes={},
        report_paths={},
        rerun_of=previous_id,
    )
    record = history.get_run(current_id)

    delta, note = _rerun_delta(history, record)

    assert delta is None
    assert "output mode changed" in note
    assert "Atomic" in note and "Decision" in note


def test_rerun_delta_discloses_decision_to_atomic_mode_change(tmp_path: Path) -> None:
    """Same disclosure, reversed direction (decision -> atomic)."""
    work_dir = tmp_path / "work"
    history = RunHistory(work_dir / "history.sqlite3")
    finding = Finding(
        finding_id="F1",
        artifact="excel",
        finding_class=FindingClass.VALUE_CHANGED,
        severity=Severity.CRITICAL,
        sheet="Data",
        location="B2",
        baseline_value="1",
        current_value="2",
        message="changed",
    )
    previous_id = history.record_run(
        QCRunResult(
            profile_name="fixture",
            findings=[finding],
            requested_output_mode=FindingOutputMode.DECISION,
        ),
        file_hashes={},
        report_paths={},
    )
    current_id = history.record_run(
        QCRunResult(
            profile_name="fixture",
            findings=[finding],
            requested_output_mode=FindingOutputMode.ATOMIC,
        ),
        file_hashes={},
        report_paths={},
        rerun_of=previous_id,
    )
    record = history.get_run(current_id)

    delta, note = _rerun_delta(history, record)

    assert delta is None
    assert "output mode changed" in note
    assert "Decision" in note and "Atomic" in note


def test_rerun_delta_still_compares_normally_within_the_same_output_mode(
    tmp_path: Path,
) -> None:
    """Same-mode reruns keep the real resolved/new/persisting comparison --
    the mode check must not suppress the ordinary, correct case.
    """
    work_dir = tmp_path / "work"
    history = RunHistory(work_dir / "history.sqlite3")
    finding = Finding(
        finding_id="F1",
        artifact="excel",
        finding_class=FindingClass.VALUE_CHANGED,
        severity=Severity.CRITICAL,
        sheet="Data",
        location="B2",
        baseline_value="1",
        current_value="2",
        message="changed",
    )
    previous_id = history.record_run(
        QCRunResult(
            profile_name="fixture",
            findings=[finding],
            requested_output_mode=FindingOutputMode.ATOMIC,
        ),
        file_hashes={},
        report_paths={},
    )
    current_id = history.record_run(
        QCRunResult(
            profile_name="fixture",
            findings=[finding.model_copy(update={"finding_id": "N1"})],
            requested_output_mode=FindingOutputMode.ATOMIC,
        ),
        file_hashes={},
        report_paths={},
        rerun_of=previous_id,
    )
    record = history.get_run(current_id)

    delta, note = _rerun_delta(history, record)

    assert note == ""
    assert delta is not None
    assert delta.persisting == 1
    assert delta.resolved == 0
    assert delta.new == 0


def test_rerun_delta_skips_same_mode_effective_policy_change(tmp_path: Path) -> None:
    work_dir = tmp_path / "work"
    history = RunHistory(work_dir / "history.sqlite3")
    finding = Finding(
        finding_id="F1",
        artifact="excel",
        finding_class=FindingClass.FORMULA_LOGIC_CHANGED,
        severity=Severity.WARNING,
        sheet="Data",
        location="B2",
        baseline_location="B2",
        baseline_value="=A2*2",
        current_value="=A2*3",
        message="changed",
    )
    previous_policy = resolve_output_policy(
        FindingOutputMode.PROFILE,
        ReviewPolicy(populations=PopulationPolicy(enabled=True, threshold=10)),
    )
    current_policy = resolve_output_policy(
        FindingOutputMode.PROFILE,
        ReviewPolicy(populations=PopulationPolicy(enabled=True, threshold=20)),
    )
    previous_id = history.record_run(
        QCRunResult(
            profile_name="fixture",
            findings=[finding],
            requested_output_mode=FindingOutputMode.PROFILE,
            resolved_output_policy=previous_policy,
        ),
        file_hashes={},
        report_paths={},
    )
    current_id = history.record_run(
        QCRunResult(
            profile_name="fixture",
            findings=[finding.model_copy(update={"finding_id": "N1"})],
            requested_output_mode=FindingOutputMode.PROFILE,
            resolved_output_policy=current_policy,
        ),
        file_hashes={},
        report_paths={},
        rerun_of=previous_id,
    )

    delta, note = _rerun_delta(history, history.get_run(current_id))

    assert delta is None
    assert "effective output policy" in note


def test_rerun_delta_skips_same_policy_threshold_crossing(tmp_path: Path) -> None:
    from qc_tool.findings import MembershipCodec, PopulationEvidence

    work_dir = tmp_path / "work"
    history = RunHistory(work_dir / "history.sqlite3")
    policy = resolve_output_policy(
        FindingOutputMode.DECISION,
        ReviewPolicy(populations=PopulationPolicy(enabled=True, threshold=10)),
    )
    atomic = Finding(
        finding_id="A1",
        artifact="excel",
        finding_class=FindingClass.FORMULA_LOGIC_CHANGED,
        severity=Severity.WARNING,
        sheet="Data",
        location="B2",
        baseline_location="B2",
        baseline_value="=A2*2",
        current_value="=A2*3",
        message="changed",
    )
    shape_before = hashlib.sha256(b"=RC[-1]*2").hexdigest()
    shape_after = hashlib.sha256(b"=RC[-1]*3").hexdigest()
    population = Finding(
        finding_id="P1",
        artifact="excel",
        finding_class=FindingClass.FORMULA_LOGIC_CHANGED,
        severity=Severity.WARNING,
        sheet="Data",
        location="B2:B11",
        element="population",
        message="population",
        population=PopulationEvidence(
            member_count=10,
            membership=MembershipCodec(
                current_rectangles=("B2:B11",),
                baseline_mode="shift",
                shift=(0, 0),
                member_count=10,
            ),
            first="B2",
            last="B11",
            shape_before_digest=shape_before,
            shape_after_digest=shape_after,
        ),
    )
    previous_id = history.record_run(
        QCRunResult(
            profile_name="fixture",
            findings=[population],
            requested_output_mode=FindingOutputMode.DECISION,
            resolved_output_policy=policy,
        ),
        file_hashes={},
        report_paths={},
    )
    current_id = history.record_run(
        QCRunResult(
            profile_name="fixture",
            findings=[atomic],
            requested_output_mode=FindingOutputMode.DECISION,
            resolved_output_policy=policy,
        ),
        file_hashes={},
        report_paths={},
        rerun_of=previous_id,
    )

    delta, note = _rerun_delta(history, history.get_run(current_id))

    assert delta is None
    assert "population grouping changed" in note


@pytest.mark.asyncio
async def test_reqc_page_discloses_output_mode_change_instead_of_a_delta_banner(
    user: User, tmp_path: Path
) -> None:
    """End-to-end: the results page shows the disclosure note, not a
    misleading resolved/new banner, when Re-QC changed output mode."""
    work_dir = tmp_path / "work"
    history = RunHistory(work_dir / "history.sqlite3")
    finding = Finding(
        finding_id="F1",
        artifact="excel",
        finding_class=FindingClass.VALUE_CHANGED,
        severity=Severity.CRITICAL,
        sheet="Data",
        location="B2",
        baseline_value="1",
        current_value="2",
        message="changed",
    )
    previous = history.record_run(
        QCRunResult(
            profile_name="fixture",
            findings=[finding],
            requested_output_mode=FindingOutputMode.ATOMIC,
        ),
        file_hashes={},
        report_paths={},
    )
    current = history.record_run(
        QCRunResult(
            profile_name="fixture",
            findings=[finding.model_copy(update={"finding_id": "N1"})],
            requested_output_mode=FindingOutputMode.DECISION,
        ),
        file_hashes={},
        report_paths={},
        rerun_of=previous,
    )
    create_pages(work_dir)
    await user.open(f"/runs/{current}")

    await user.should_see("output mode changed")
    await user.should_not_see("resolved")
    await user.should_not_see("new ·")


@pytest.mark.asyncio
async def test_run_page_renders_a_predecessor_configuration_diff(
    user: User, tmp_path: Path
) -> None:
    """plan-20260913 Step 10: "History surfaces config diff versus
    predecessor using persisted resolved snapshots, not heuristics" -- a
    region whose resolved mode changed between the predecessor and this
    run is disclosed by name, using only the two runs' own persisted
    resolved configurations.
    """
    from qc_tool.config.input_contract import INPUT_CONTRACT_VERSION
    from qc_tool.config.resolved_input import (
        ResolvedInputConfigurationV1,
        ResolvedMember,
        ResolvedRegion,
        ResolvedSheet,
    )

    work_dir = tmp_path / "work"
    history = RunHistory(work_dir / "history.sqlite3")
    finding = Finding(
        finding_id="F1",
        artifact="excel",
        finding_class=FindingClass.VALUE_CHANGED,
        severity=Severity.CRITICAL,
        sheet="Data",
        location="B2",
        baseline_value="1",
        current_value="2",
        message="changed",
    )

    def _resolved(
        mode: Literal["automatic", "keyed", "positional", "excluded"],
    ) -> ResolvedInputConfigurationV1:
        return ResolvedInputConfigurationV1(
            inspection_contract_version=INPUT_CONTRACT_VERSION,
            members=(
                ResolvedMember(
                    member_id="primary",
                    sheets=(
                        ResolvedSheet(
                            sheet_id="sheet-data",
                            baseline_sheet_name="Data",
                            current_sheet_name="Data",
                            regions=(ResolvedRegion(region_id="r1", mode=mode),),
                        ),
                    ),
                ),
            ),
        )

    previous = history.record_run(
        QCRunResult(
            profile_name="fixture",
            findings=[finding],
            resolved_input_configuration=_resolved("automatic"),
        ),
        file_hashes={},
        report_paths={},
    )
    current = history.record_run(
        QCRunResult(
            profile_name="fixture",
            findings=[finding.model_copy(update={"finding_id": "N1"})],
            resolved_input_configuration=_resolved("keyed"),
        ),
        file_hashes={},
        report_paths={},
        rerun_of=previous,
    )
    create_pages(work_dir)
    await user.open(f"/runs/{current}")

    await user.should_see(f"Configuration changes vs run #{previous}")


@pytest.mark.asyncio
async def test_run_page_saves_a_new_profile_from_the_runs_configuration(
    user: User, tmp_path: Path
) -> None:
    """plan-20260913 Step 11: a completed run's own resolved configuration
    can be saved as a named profile's configuration through the run-detail
    page. The run's own recorded evidence is untouched -- this only writes
    a NEW saved profile file.
    """
    from qc_tool.config.input_contract import INPUT_CONTRACT_VERSION
    from qc_tool.config.profile import load_profile_by_name
    from qc_tool.config.resolved_input import (
        ResolvedInputConfigurationV1,
        ResolvedMember,
        ResolvedRegion,
        ResolvedSheet,
    )

    work_dir = tmp_path / "work"
    (work_dir / "profiles").mkdir(parents=True)
    history = RunHistory(work_dir / "history.sqlite3")
    finding = Finding(
        finding_id="F1",
        artifact="excel",
        finding_class=FindingClass.VALUE_CHANGED,
        severity=Severity.CRITICAL,
        sheet="Data",
        location="B2",
        baseline_value="1",
        current_value="2",
        message="changed",
    )
    run_id = history.record_run(
        QCRunResult(
            profile_name="default",
            findings=[finding],
            resolved_input_configuration=ResolvedInputConfigurationV1(
                inspection_contract_version=INPUT_CONTRACT_VERSION,
                members=(
                    ResolvedMember(
                        member_id="primary",
                        sheets=(
                            ResolvedSheet(
                                sheet_id="sheet_data",
                                current_sheet_name="Data",
                                regions=(
                                    ResolvedRegion(
                                        region_id="region_data",
                                        mode="keyed",
                                        current_outer_range="A1:C10",
                                    ),
                                ),
                            ),
                        ),
                    ),
                ),
            ),
        ),
        file_hashes={},
        report_paths={},
    )
    create_pages(work_dir)

    await user.open(f"/runs/{run_id}")
    user.find("Save configuration to profile").click()

    await user.should_see("Save this run's configuration to a profile")
    name_input = next(
        element
        for element in user.find(kind=ui.input).elements
        if element.props.get("label") == "Profile name"
    )
    name_input.value = "from-this-run"
    user.find(marker="save-configuration-to-profile-confirm").click()

    await user.should_see("Profile 'from-this-run' saved from this run's configuration")
    saved = load_profile_by_name(work_dir / "profiles", "from-this-run")
    assert saved.input_contract is not None
    assert len(saved.input_contract.members[0].sheets[0].regions) == 1
    assert saved.input_contract.members[0].sheets[0].regions[0].mode == "keyed"


@pytest.mark.asyncio
async def test_final_package_detail_shows_mapping_review(
    user: User, fixture_dir: Path, tmp_path: Path
) -> None:
    work_dir = tmp_path / "work"
    profiles_dir = work_dir / "profiles"
    profiles_dir.mkdir(parents=True)
    profile = fixture_profile()
    save_profile(profile, profiles_dir / "fixture.yaml")
    artifacts = perform_run(
        work_dir,
        {
            "current_excel": fixture_dir / "current.xlsx",
            "current_ppt": fixture_dir / "current.pptx",
        },
        {},
        profile,
        mode=QCRunMode.FINAL_PACKAGE,
    )

    create_pages(work_dir)
    await user.open(f"/runs/{artifacts.run_id}")
    await user.should_see("Final-package QC")
    await user.should_see("Check coverage")
    await user.should_see("Mapping review")
    await user.should_see("readable")
    await user.should_see("unavailable")
    await user.should_see("total surfaces")


@pytest.mark.asyncio
async def test_mapping_confirmation_survives_the_repaint_it_triggers(
    user: User,
    fixture_dir: Path,
    tmp_path: Path,
    caplog: pytest.LogCaptureFixture,
) -> None:
    work_dir = tmp_path / "work"
    profiles_dir = work_dir / "profiles"
    profiles_dir.mkdir(parents=True)
    profile = DeliverableProfile(name="mapping-review")
    save_profile(profile, profiles_dir / "mapping-review.yaml")
    artifacts = perform_run(
        work_dir,
        {
            "current_excel": fixture_dir / "current.xlsx",
            "current_ppt": fixture_dir / "current.pptx",
        },
        {},
        profile,
        mode=QCRunMode.FINAL_PACKAGE,
    )
    before = RunHistory(work_dir / "history.sqlite3").get_run(artifacts.run_id)
    assert before.mapping_coverage is not None and before.mapping_suggestions

    create_pages(work_dir)
    await user.open(f"/runs/{artifacts.run_id}")
    buttons = user.find("Confirm").elements
    assert buttons

    nicegui_warnings.reset()
    with caplog.at_level(logging.WARNING, logger="nicegui"):
        _emit(next(iter(buttons)), "click", {})

    assert not caplog.records, [record.getMessage() for record in caplog.records]
    after = RunHistory(work_dir / "history.sqlite3").get_run(artifacts.run_id)
    assert after.mapping_coverage is not None
    assert after.mapping_coverage.mapped == 1
    assert len(after.mapping_suggestions) == len(before.mapping_suggestions) - 1


def _emit(element: ui.element, event_type: str, args: object) -> None:
    """Drive a slot-emitted table event the way the browser would."""
    for listener in element._event_listeners.values():
        if listener.type == event_type and listener.handler is not None:
            # Handlers may schedule client JavaScript, which needs a slot.
            with element.parent_slot or element.client.layout.default_slot:
                listener.handler(
                    events.GenericEventArguments(
                        sender=element, client=element.client, args=args
                    )
                )
            return
    raise AssertionError(f"no {event_type!r} listener on {element}")


@pytest.mark.asyncio
async def test_group_review_survives_the_panel_refresh_it_triggers(
    user: User,
    fixture_dir: Path,
    tmp_path: Path,
    caplog: pytest.LogCaptureFixture,
) -> None:
    """Applying a group review refreshes the panel that owns the dialog.

    Reporting or closing after that refresh touches deleted elements, which
    NiceGUI surfaces as a use-after-free warning and a RuntimeError.
    """
    work_dir = tmp_path / "work"
    artifacts = perform_run(
        work_dir,
        {
            "baseline_excel": fixture_dir / "baseline.xlsx",
            "current_excel": fixture_dir / "current.xlsx",
        },
        {},
        fixture_profile(),
    )
    create_pages(work_dir)
    await user.open(f"/runs/{artifacts.run_id}")

    group_table = next(
        element
        for element in user.find(kind=ui.table).elements
        if "review-groups-table" in element.classes
    )
    # Related-series parents are a navigation lens with no group review action.
    reviewable = next(
        row for row in group_table.rows if row.get("kind") != "cluster"
    )
    _emit(group_table, "select", {"id": str(reviewable["id"])})
    await user.should_see("Review group")

    user.find("Review group").click()
    await user.should_see("Apply review")
    severity_select = next(
        element
        for element in user.find(kind=ui.select).elements
        if element.props.get("label") == "Group severity"
    )
    severity_select.value = Severity.INFO.value

    # Closing or reporting after the refresh touches deregistered elements,
    # which NiceGUI reports by logging rather than by raising to the caller.
    nicegui_warnings.reset()  # the use-after-free warning is otherwise once-only
    with caplog.at_level(logging.WARNING, logger="nicegui"):
        user.find("Apply review").click()

    assert not caplog.records, [record.getMessage() for record in caplog.records]


def test_comment_only_review_is_visible_in_the_queue(qc_result: QCRunResult) -> None:
    """Choosing "keep" records a decision without setting severity_overridden.

    Nothing in the queue used to indicate that, so the decision looked lost.
    """
    findings = [finding.model_copy(deep=True) for finding in qc_result.findings]
    findings[0].analyst_comment = "checked against the source pack"

    rows = _review_group_rows(build_pattern_groups(findings))
    reviewed = {str(row["id"]): int(str(row["reviewed"])) for row in rows}
    members = {str(row["id"]): int(str(row["members"])) for row in rows}

    assert any(reviewed.values()), "a comment-only decision must mark its group reviewed"
    assert all(reviewed[key] <= members[key] for key in reviewed)
    assert sum(reviewed.values()) == 1


def test_untouched_findings_are_not_marked_reviewed(qc_result: QCRunResult) -> None:
    findings = [finding.model_copy(deep=True) for finding in qc_result.findings]

    rows = _review_group_rows(build_pattern_groups(findings))

    assert all(row["reviewed"] == 0 for row in rows)


def test_row_slots_show_a_note_without_a_severity_override() -> None:
    for slot in (app_module.REVIEW_MEMBER_ROWS_SLOT, app_module.FINDINGS_BODY_SLOT):
        assert 'v-else-if="props.row.comment"' in slot
    assert 'v-if="props.row.reviewed"' in app_module.REVIEW_GROUPS_BODY_SLOT


def test_review_rows_have_keyboard_and_focus_contracts() -> None:
    assert 'role="button" tabindex="0"' in app_module.REVIEW_GROUPS_BODY_SLOT
    assert "@keydown.enter.prevent" in app_module.REVIEW_GROUPS_BODY_SLOT
    assert "@keydown.space.prevent" in app_module.REVIEW_GROUPS_BODY_SLOT
    assert ':data-review-id="props.row.id"' in app_module.REVIEW_GROUPS_BODY_SLOT
    assert 'role="button" tabindex="0"' in app_module.REVIEW_MEMBER_ROWS_SLOT
    assert ':data-member-id="props.row.id"' in app_module.REVIEW_MEMBER_ROWS_SLOT
    assert "@media (max-width: 1350px)" in CSS
    assert ".reviewclass-inline { display: inline; }" in CSS


def test_shared_wordmark_is_an_accessible_home_link() -> None:
    source = inspect.getsource(page_frame)

    assert 'ui.link("QC Tool", "/").classes("wordmark")' in source
    assert ".wordmark:focus-visible" in CSS


def test_keyboard_triage_script_has_all_focus_and_dialog_guards() -> None:
    source = inspect.getsource(app_module._render_result_view)

    for contract in (
        "if not mutable:",
        "activeElementIsEditable()",
        "a.isContentEditable",
        "dialogIsOpen()",
        "if (!visibleIds().length) return",
        "if (activeElementIsEditable() || dialogIsOpen()) return",
        "persist_group_review(",
        "replace_existing=False",
    ):
        assert contract in source


def test_longitudinal_panel_reuses_the_loaded_dossier_for_recurrence() -> None:
    source = inspect.getsource(app_module._render_result_view)
    body = source.split("def load_longitudinal", 1)[1].split(
        "expansion.on_value_change", 1
    )[0]

    assert "recurrence_eligibility_from_dossier" in body
    assert "active_history.recurrence_eligibility," not in body


def test_atomic_slot_can_confirm_the_current_severity() -> None:
    assert 'label="Confirm severity"' in app_module.FINDINGS_BODY_SLOT
    assert "value: props.row.severity" in app_module.FINDINGS_BODY_SLOT


# --- bounded logical-series context panels ------------------------------------


def _context_finding(
    finding_id: str,
    location: str,
    severity: Severity,
    *,
    excerpt: GridExcerpt | None = None,
) -> Finding:
    return Finding(
        finding_id=finding_id,
        artifact="excel",
        finding_class=FindingClass.VALUE_CHANGED,
        severity=severity,
        sheet="Ops",
        location=location,
        baseline_location=location,
        baseline_value="1",
        current_value="2",
        message=f"Ops!{location}: value changed",
        current_excerpt=excerpt,
    )


def _window(rows: list[int], cols: list[str], hit: tuple[int, int]) -> GridExcerpt:
    return GridExcerpt(
        cols=cols,
        rows=rows,
        cells=[[f"{col}{row}" for col in cols] for row in rows],
        hit_row=hit[0],
        hit_col=hit[1],
    )


def test_overlapping_excerpts_merge_into_one_panel() -> None:
    members = [
        _context_finding(
            "F1",
            "B3",
            Severity.CRITICAL,
            excerpt=_window([2, 3, 4], ["A", "B", "C"], (1, 1)),
        ),
        _context_finding(
            "F2",
            "B5",
            Severity.WARNING,
            excerpt=_window([4, 5, 6], ["A", "B", "C"], (1, 1)),
        ),
    ]

    context = build_cluster_context(members, side="current", selected_finding_id="F2")

    assert context.merged
    assert len(context.panels) == 1
    panel = context.panels[0]
    assert panel.rows == (2, 3, 4, 5, 6)
    assert panel.cols == ("A", "B", "C")
    assert panel.marks == {(1, 1): "critical", (3, 1): "warning"}
    assert panel.selected == (3, 1)
    assert context.shown == 2
    assert context.total == 2
    assert context.complete


def test_disconnected_windows_stay_separate_panels() -> None:
    members = [
        _context_finding(
            "F1",
            "B3",
            Severity.CRITICAL,
            excerpt=_window([2, 3, 4], ["A", "B", "C"], (1, 1)),
        ),
        _context_finding(
            "F2",
            "B40",
            Severity.WARNING,
            excerpt=_window([39, 40, 41], ["A", "B", "C"], (1, 1)),
        ),
    ]

    context = build_cluster_context(members, side="current")

    assert [panel.rows for panel in context.panels] == [(2, 3, 4), (39, 40, 41)]
    assert context.shown == 2


def test_panels_are_bounded_to_twenty_five_rows_by_fifteen_columns() -> None:
    rows = list(range(1, 61))
    cols = [get_column_letter(index) for index in range(1, 21)]
    members = [
        _context_finding(
            "F1", "A1", Severity.CRITICAL, excerpt=_window(rows, cols, (0, 0))
        )
    ]

    context = build_cluster_context(members, side="current")

    assert context.panels
    for panel in context.panels:
        assert len(panel.rows) <= 25
        assert len(panel.cols) <= 15
    assert sum(len(panel.rows) for panel in context.panels) >= 60


def test_conflicting_stored_values_fall_back_to_individual_excerpts() -> None:
    first = _window([2, 3, 4], ["A", "B", "C"], (1, 1))
    second = _window([2, 3, 4], ["A", "B", "C"], (1, 1))
    second.cells[1][1] = "contradicting"
    members = [
        _context_finding("F1", "B3", Severity.CRITICAL, excerpt=first),
        _context_finding("F2", "B4", Severity.WARNING, excerpt=second),
    ]

    context = build_cluster_context(members, side="current")

    assert not context.merged
    assert len(context.panels) == 2


def test_context_coverage_is_reported_when_excerpts_are_missing() -> None:
    members = [
        _context_finding(
            "F1",
            "B3",
            Severity.CRITICAL,
            excerpt=_window([2, 3, 4], ["A", "B", "C"], (1, 1)),
        ),
        _context_finding("F2", "B90", Severity.WARNING),
    ]

    context = build_cluster_context(members, side="current")

    assert not context.complete
    assert context.coverage_note == "1 of 2 related finding cells shown"


def test_panel_html_escapes_values_and_labels_every_severity_cell() -> None:
    excerpt = _window([2, 3, 4], ["A", "B", "C"], (1, 1))
    excerpt.cells[1][1] = "<script>alert(1)</script>"
    members = [_context_finding("F1", "B3", Severity.CRITICAL, excerpt=excerpt)]

    context = build_cluster_context(members, side="current", selected_finding_id="F1")
    rendered = _cluster_context_html(context.panels[0], "current")

    assert "<script>alert(1)</script>" not in rendered
    assert "&lt;script&gt;alert(1)&lt;/script&gt;" in rendered
    assert 'class="sev-critical hit"' in rendered
    assert 'aria-label="critical finding at B3"' in rendered
    assert 'title="critical finding at B3"' in rendered


def test_theme_defines_a_non_color_only_severity_legend() -> None:
    for rule in (
        ".ctxgrid td.sev-critical",
        ".ctxgrid td.sev-warning",
        ".ctxgrid td.sev-info",
        ".ctxgrid td.sev-expected",
        ".ctxlegend",
    ):
        assert rule in CSS


# --- UI craftsmanship pass (2026-08-08) ----------------------------------------


def test_display_cell_text_tames_repr_noise_but_keeps_the_exact_value() -> None:
    short, full = app_module._display_cell_text("3428.8569518319678")
    assert short == "3428.857"
    assert full == "3428.8569518319678"

    date, full_date = app_module._display_cell_text("2024-02-23T00:00:00")
    assert date == "2024-02-23"
    assert full_date == "2024-02-23T00:00:00"

    # short numbers, text, and non-midnight stamps stay verbatim
    assert app_module._display_cell_text("1216") == ("1216", None)
    assert app_module._display_cell_text("=B4*C4") == ("=B4*C4", None)
    assert app_module._display_cell_text("2024-02-23T10:30:00") == (
        "2024-02-23T10:30:00",
        None,
    )


def test_class_filter_label_summarizes_instead_of_chipping() -> None:
    assert _class_filter_label(21, 21) == "All classes"
    assert _class_filter_label(0, 21) == "No classes"
    assert _class_filter_label(5, 21) == "5 of 21 classes"
    assert _class_filter_label(0, 0) == "All classes"


def test_warning_is_amber_not_a_second_red() -> None:
    assert "--warning: #9a6700" in CSS  # light
    assert "--warning: #e2b54b" in CSS  # dark
    # non-color outline cues in context grids
    assert "outline: 2px dashed var(--warning)" in CSS
    assert "outline: 2px dotted var(--info)" in CSS


def test_toolbar_is_sticky_and_the_queue_scrolls_internally() -> None:
    toolbar = CSS.split(".reviewtoolbar {", 1)[1].split("}", 1)[0]
    assert "position: sticky" in toolbar
    assert ".review-groups-table .q-table__middle { max-height:" in CSS
    assert (
        ".review-groups-table thead tr th { position: sticky; top: 0;" in CSS
    )


def test_detail_panel_pins_header_and_scrolls_body() -> None:
    assert ".detailhead" in CSS
    assert ".detailbody" in CSS
    body_rule = CSS.split(".detailbody {", 1)[1].split("}", 1)[0]
    assert "overflow-y: auto" in body_rule
    source = inspect.getsource(app_module._render_result_view)
    # every detail renderer builds head + body, with actions in the head
    assert source.count('classes("detailhead")') >= 4
    assert source.count('classes("detailbody")') >= 4
    assert 'classes("detailactions")' in source


def test_parent_row_renders_mix_chips() -> None:
    template = app_module.REVIEW_GROUPS_BODY_SLOT.split("<q-tr v-else", 1)[0]
    assert "props.row.mix" in template
    assert "mixchip mix-" in template


def test_cluster_legend_lists_only_markers_that_can_appear() -> None:
    # The cluster view opens a whole series, never one finding, so its grids
    # cannot show the amber "open finding" cell; single-finding excerpts show
    # the box without a legend and need none. User-flagged 2026-08-08.
    source = inspect.getsource(app_module._render_result_view)
    assert "open finding" not in source
    assert "selected finding<" not in source
    assert 'swatch selected' not in source
    # the legend is pinned in the cluster detail head, above the scroll
    head = source.split("def _render_cluster_detail", 1)[1].split(
        'classes("detailbody")', 1
    )[0]
    assert "_severity_legend_html()" in head


# --- follow-up review pass (2026-08-08) -----------------------------------------


def test_format_review_time_reads_like_a_clock() -> None:
    assert app_module._format_review_time(None) == "0:00"
    assert app_module._format_review_time(0) == "0:00"
    assert app_module._format_review_time(59.9) == "0:59"
    assert app_module._format_review_time(605) == "10:05"
    assert app_module._format_review_time(3600) == "1:00:00"
    assert app_module._format_review_time(3725) == "1:02:05"


def test_timer_kpi_is_a_stat_box_pinned_top_right() -> None:
    source = inspect.getsource(app_module._render_result_view)
    assert 'classes("stat timerkpi")' in source
    assert "review time" in source
    # a finalized run shows its frozen time and offers no controls
    assert 'ui.label("finalized").classes("a")' in source
    kpi_rule = CSS.split(".timerkpi {", 1)[1].split("}", 1)[0]
    assert "position: absolute" in kpi_rule
    assert "right: 0" in kpi_rule
    header_rule = CSS.split(".runheader {", 1)[1].split("}", 1)[0]
    assert "position: relative" in header_rule


def test_cluster_row_states_each_fact_once() -> None:
    entries = _entries()
    row = entries[0].row

    # LOCATION carries only the axis; the period count lives in the # column,
    # so no number is printed three times per row.
    assert str(row["location"]).startswith(("column ", "row "))
    assert "period" not in str(row["location"])
    message = str(row["message"])
    assert "findings" not in message
    assert "lens" not in message
    assert "decision" in message


def test_queue_pager_sits_under_the_queue_it_turns() -> None:
    # In the toolbar the pager read as detail-panel chrome; it pages the
    # queue, so it lives with the queue, like the members dialog pager.
    source = inspect.getsource(app_module._render_result_view)
    assert '"items-center queuepager"' in source
    table = source.index('add_slot("body", REVIEW_GROUPS_BODY_SLOT)')
    pager = source.index('"items-center queuepager"')
    detail = source.index('classes("detailpanel")')
    assert table < pager < detail
    assert "lenspager" not in source
    # centred under the table, with a persisted rows-per-page choice
    pager_rule = CSS.split(".queuepager {", 1)[1].split("}", 1)[0]
    assert "justify-content: center" in pager_rule
    assert app_module._TOP_PAGE_SIZES == (10, 25, 50, 100)
    assert 'app.storage.general["review_page_size"]' in source
    assert 'start // size' in source  # resizing keeps the first row stable


def test_queue_columns_have_light_boundaries_and_manual_resize() -> None:
    # boundaries use the theme line variable so both modes stay quiet
    assert (
        ".review-groups-table th:not(:last-child),\n"
        ".review-groups-table td:not(:last-child) "
        "{ border-right: 1px solid var(--line-soft); }" in CSS
    )
    from qc_tool.ui.theme import COL_RESIZE_JS

    # drag a header boundary to trade width between neighbours; double-click
    # restores the stylesheet defaults; a grab never toggles the sort
    assert "col-resize" in COL_RESIZE_JS
    assert "dblclick" in COL_RESIZE_JS
    assert "stopPropagation" in COL_RESIZE_JS
    assert "MIN = 48" in COL_RESIZE_JS
    source = inspect.getsource(app_module._render_result_view)
    assert "ui.add_body_html(COL_RESIZE_JS)" in source


def test_sheet_facet_filters_children_and_keeps_matching_parents() -> None:
    # the fixture series live on the Ops sheet: constraining to Ops keeps
    # everything, constraining elsewhere empties the queue
    assert len(_entries(sheets={"Ops"})) == 4
    assert _entries(sheets={"Elsewhere"}) == []
    assert _entries(sheets=set()) == []
    assert len(_entries(sheets=None)) == 4  # unconstrained default


def test_where_facet_helpers_name_what_they_filter() -> None:
    assert app_module._facet_filter_label(3, 9, "sheets") == "3 of 9 sheets"
    assert app_module._facet_filter_label(9, 9, "sheets") == "All sheets"
    assert app_module._facet_filter_label(0, 9, "slides") == "No slides"
    assert _class_filter_label(5, 21) == "5 of 21 classes"  # wrapper unchanged

    def _group(sheet: str | None, slide: str | None) -> ReviewGroup:
        return cast(
            "ReviewGroup", SimpleNamespace(sheet=sheet, slide=slide)
        )

    # slide labels are titles when the deck has them, so the noun must come
    # from the group fields, never from string matching
    assert app_module._where_noun([_group("Ops", None)]) == "sheets"
    assert app_module._where_noun([_group(None, "Executive Summary")]) == "slides"
    assert (
        app_module._where_noun(
            [_group("Ops", None), _group(None, "slide 2")]
        )
        == "sheets/slides"
    )
    assert app_module._where_noun([]) == "sheets"

    ordered = sorted(["slide 10", "slide 2", "Ops"], key=app_module._natural_key)
    assert ordered == ["Ops", "slide 2", "slide 10"]


def test_toolbar_offers_a_sheet_facet_like_the_class_facet() -> None:
    source = inspect.getsource(app_module._render_result_view)
    assert "where_button" in source
    assert "_all_wheres" in source
    assert "_set_where" in source
    assert '"whole file"' in source  # blank locations stay reachable
    assert "sheets=set(selected_wheres)" in source


def test_queue_sort_never_splits_a_series_from_its_children() -> None:
    # Quasar client-side header sort scattered child rows away from their
    # parents (user-found 2026-08-08): the queue's flat row array interleaves
    # families, so ordering must be Python-owned at the entry level.
    assert all("sortable" not in column for column in app_module._REVIEW_GROUP_COLUMNS)

    entries = _entries()
    by_severity = app_module._sort_lens_entries(entries, "severity")
    # families intact: every entry keeps exactly its own children
    assert {id(e) for e in by_severity} == {id(e) for e in entries}
    ranks = [
        [s.value for s in Severity].index(str(e.row["severity"]))
        for e in by_severity
    ]
    assert ranks == sorted(ranks)

    by_findings = app_module._sort_lens_entries(entries, "findings", descending=True)
    counts = [int(str(e.row["members"])) for e in by_findings]
    assert counts == sorted(counts, reverse=True)

    by_location = app_module._sort_lens_entries(entries, "location")
    locations = [str(e.row["location"]) for e in by_location]
    assert locations == sorted(locations, key=app_module._natural_key)

    # priority keeps the evidence order; descending is an honest reversal
    assert app_module._sort_lens_entries(entries, "priority") == entries
    assert app_module._sort_lens_entries(entries, "priority", descending=True) == list(
        reversed(entries)
    )

    source = inspect.getsource(app_module._render_result_view)
    assert "_sort_lens_entries(" in source
    assert "order: priority" in source
    assert "Reverse queue order" in source


def test_review_count_column_renders_comma_grouped_and_stays_wide_and_nowrap() -> None:
    """Client feedback: large `#` counts must never wrap/overflow, but the
    underlying row value stays a plain int so findings-order sort (above) and
    every other int(str(row["members"])) reader keep working unchanged."""
    # exactly two <q-td key="members"> cells: the cluster-row template and the
    # ordinary row/child template; both must format via JS toLocaleString.
    assert REVIEW_GROUPS_BODY_SLOT.count('key="members"') == 2
    assert REVIEW_GROUPS_BODY_SLOT.count(
        "Number(props.row.members).toLocaleString('en-US')"
    ) == 2
    assert "{{ props.row.members }}" not in REVIEW_GROUPS_BODY_SLOT

    fourth_column_rules = [
        line
        for line in CSS.splitlines()
        if ".review-groups-table" in line and "nth-child(4)" in line
    ]
    assert fourth_column_rules, "the # column must have dedicated CSS rules"
    assert any("white-space: nowrap" in line for line in fourth_column_rules)
    widths = [
        int(match.group(1))
        for line in fourth_column_rules
        if (match := re.search(r"width:\s*(\d+)%", line))
    ]
    # 1,234,567 (9 chars, monospace tabular-nums) needs real room at every
    # supported desktop breakpoint, not the old 4%/6%.
    assert widths and all(width >= 7 for width in widths)


def test_run_qc_submission_is_single_flight() -> None:
    """Client feedback: rapid repeated Run QC clicks must produce at most one
    projection, one dialog, and one queued request, released only when the
    owned request's RunStateRecord.is_active goes false (never a terminal
    allowlist) or the analyst explicitly cancels the projection dialog."""
    source = inspect.getsource(app_module.create_pages)
    start_run_source = source.split("async def start_run(", 1)[1].split(
        "def _open_projection_dialog(", 1
    )[0].split(") -> None:", 1)[1]

    # the guard is the very first statement, before any `await`, so a second
    # concurrently scheduled click sees the lock before it can act
    guard_line = start_run_source.strip().splitlines()[0]
    assert guard_line == "if _run_ui_busy():"
    assert "await" not in start_run_source.split(guard_line, 1)[0]

    # phase transitions use is_active/ACTIVE_STATUSES semantics, never an
    # enumerated terminal-status allowlist
    assert '"phase": "idle"' in source
    assert 'run_lock["phase"] = "projecting"' in start_run_source
    assert 'run_lock["phase"] = "dialog"' in start_run_source
    assert "if run_lock[\"request_id\"] == request_id:" in source
    assert "_unlock_run()" in source
    assert "record.is_active" in source
    assert "ACTIVE_STATUSES" not in source  # terminality is never enumerated here

    # the volume-projection dialog can only resolve via its own Cancel/Run
    # buttons, so the lock always releases deterministically
    dialog_source = source.split("def _open_projection_dialog(", 1)[1].split(
        "def submit_run(", 1
    )[0]
    assert 'ui.dialog().props("persistent")' in dialog_source
    assert "def cancel_dialog() -> None:" in dialog_source
    assert "_unlock_run()" in dialog_source.split("def cancel_dialog", 1)[1]

    # submit_run releases the lock on QueueBusyError and locks to "submitted"
    # on success, so it is never left stuck in "projecting"/"dialog"
    submit_run_source = source.split("def submit_run(", 1)[1]
    assert 'run_lock["phase"] = "submitted"' in submit_run_source
    busy_error_branch = submit_run_source.split("except QueueBusyError as exc:", 1)[1].split(
        "run_lock", 1
    )[0]
    assert "_unlock_run()" in busy_error_branch


def test_column_letters_normalizes_and_deduplicates() -> None:
    from qc_tool.ui.app import _column_letters

    assert _column_letters(" b, A;B ") == ["B", "A"]


def test_ranked_block_routes_back_to_the_single_configuration_workspace() -> None:
    source = inspect.getsource(app_module.create_pages)
    action_branch = source.split("def _refresh_terminal_actions(", 1)[1].split(
        "if rerun_banner_actions is not None:", 1
    )[0]

    assert "def _open_setup_attention_prompt(" in source
    assert '"QC needs more setup"' in source
    assert '"Return to setup"' in action_branch
    assert '"Review row matching"' not in action_branch
    assert "await _restore_request_context(current)" in source
    assert "_row_matching_choice_updates_from_action(" in source
    assert "choice_updates=choice_updates" in source
    assert "reuse_source_session=True" in source
    assert "await start_run(" not in action_branch


def test_complexity_failure_branch_prompts_and_retries_only_after_confirmation() -> None:
    source = inspect.getsource(app_module.create_pages)
    failure_branch = source.split(
        "elif record.status is RunStatus.FAILED:", 1
    )[1].split("else:", 1)[0]
    override_dialog = source.split("def _open_complexity_override(", 1)[1].split(
        "def _refresh_terminal_actions(", 1
    )[0]

    assert "_open_complexity_override(record)" in failure_branch
    assert 'ui.dialog().props("persistent")' in override_dialog
    assert '"Run with override"' in override_dialog
    assert "await _restore_request_context(record)" in override_dialog
    assert "state.allow_large_workbooks = True" in override_dialog
    assert "open_configuration_workspace(" in override_dialog
    assert "await start_run(" not in override_dialog


def test_stale_ranked_action_refreshes_through_managed_request_context() -> None:
    source = inspect.getsource(app_module.create_pages)
    action_branch = source.split("def _refresh_terminal_actions(", 1)[1].split(
        "_refresh_terminal_actions(latest_terminal)", 1
    )[0]

    assert "_ranked_action_needs_refresh(action)" in action_branch
    assert '"Refresh setup"' in action_branch
    assert "_continue_blocked_setup(record)" in action_branch
    assert "await start_run(" not in action_branch


def test_guide_describes_manual_prerequisites_and_ranked_setup() -> None:
    from qc_tool.ui import guide

    source = inspect.getsource(guide.render_guide)
    flat = " ".join(source.split())
    assert "These cells are manually pinned" in source
    assert "Scenario checks" in source
    assert "Ranked or sorted tables" in source
    assert "QC needs more " in source
    assert '"setup message. Return to setup' in source
    assert "There is no second" in source
    assert '"row-matching editor and QC never restarts automatically.' in flat
    assert '"is consumed and cannot prompt again on a later reload.' in flat
    assert '"Configuration never starts QC automatically;' in flat
    assert "Discard " in source
    assert "attempt removes only" in source
    assert "formula result is never presented as a header" in source
    assert "Different selections can change reference columns" in source
    assert "Rank/order columns" in source
    assert "automatic dropdown suggestions are OOXML-only" not in source


def test_mix_chips_are_quiet_and_allowed_to_wrap() -> None:
    chip_rule = CSS.split(".mixchip {", 1)[1].split("}", 1)[0]
    assert "border" not in chip_rule  # dot + count, no pill chrome
    assert (
        ".review-groups-table tr.clusterrow > td:first-child { white-space: normal; }"
        in CSS
    )


# --- related-series lens rows -------------------------------------------------


def _lens_fixture() -> tuple[
    SeriesReviewLens, dict[str, ReviewGroup], dict[str, str]
]:
    findings, anchors = series_oracle()
    groups = build_pattern_groups(findings)
    lens = build_series_review_lens(groups, anchors)
    return lens, {group.group_id: group for group in groups}, {}


def test_parent_detail_names_new_and_cleared_period_segments() -> None:
    source = inspect.getsource(app_module._render_result_view)
    body = source.split("def _render_cluster_detail", 1)[1].split(
        "def _render_slice_detail", 1
    )[0]

    assert 'segment == "new_period"' in body
    assert 'band = "new period"' in body
    assert 'segment == "cleared_period"' in body
    assert 'band = "cleared period"' in body
    assert "member.temporal_context.value" in body


def test_context_pager_turns_both_sides_together() -> None:
    # One shared index: independent per-side pagers made analysts compare
    # panel 3 of current against panel 1 of baseline. Live-found 2026-08-08.
    source = inspect.getsource(app_module._render_result_view)
    body = source.split("def _render_cluster_context", 1)[1].split(
        "def _render_cluster_detail", 1
    )[0]

    assert "shared_index" in body
    assert "def draw_sides" in body
    assert "turn: Callable[[int], None] = turn" in body
    assert "on_click=_turn_handler(-1)" in body
    assert "on_click=_turn_handler(1)" in body
    assert "on_click=lambda: turn" not in body
    # exactly one pager for both sides, hosted in the slot BETWEEN the grids
    assert body.count("chevron_left") == 1
    assert body.count("chevron_right") == 1
    assert "pager_slot" in body
    assert body.index("pager_slot = ui.element") < body.index('f"{label} context"')
    assert '" · both sides"' in body


def _entries(**overrides: object) -> list[LensEntry]:
    lens, groups_by_id, rationales = _lens_fixture()
    kwargs: dict[str, object] = {
        "severities": {severity.value for severity in Severity},
        "finding_classes": {FindingClass.VALUE_CHANGED.value},
        "review_state": "all",
    }
    kwargs.update(overrides)
    return _lens_entries(lens, groups_by_id, rationales, **kwargs)  # type: ignore[arg-type]


def test_lens_entries_expose_four_structural_parents() -> None:
    entries = _entries()

    assert len(entries) == 4
    assert [len(entry.children) for entry in entries] == [2, 2, 3, 3] or sorted(
        len(entry.children) for entry in entries
    ) == [2, 2, 3, 3]
    for entry in entries:
        assert entry.row["kind"] == "cluster"
        assert entry.row["class"] == "related_series"
        assert entry.row["parent"] == ""
        assert entry.hidden == 0
        # WHERE carries the sheet once; LOCATION carries the axis position
        assert entry.row["where"] == "Ops"
        assert str(entry.row["location"]).startswith("column ")
        assert entry.row["mix"], "severity mix chips must be populated"
    assert sum(int(str(entry.row["members"])) for entry in entries) == 19


def test_collapsed_parents_render_no_child_rows() -> None:
    entries = _entries()

    rows = _flatten_lens_rows(entries, expanded=set(), child_pages={})

    assert len(rows) == 4
    assert all(row["kind"] == "cluster" for row in rows)
    assert all(row["expanded"] is False for row in rows)


def test_expanding_one_parent_adds_only_its_children() -> None:
    entries = _entries()
    target = next(
        entry for entry in entries if str(entry.row["location"]).endswith("column M")
        or "column M" in str(entry.row["location"])
    )
    parent_id = str(target.row["id"])

    rows = _flatten_lens_rows(entries, expanded={parent_id}, child_pages={})

    children = [row for row in rows if row["kind"] == "child"]
    assert len(rows) == 4 + len(target.children)
    assert {row["parent"] for row in children} == {parent_id}
    assert [int(str(row["members"])) for row in children] == [2, 1, 1]
    assert [row["severity"] for row in children] == ["critical", "warning", "warning"]
    parent_row = next(row for row in rows if row["id"] == parent_id)
    assert parent_row["expanded"] is True
    assert parent_row["sevmix"] == "2 critical · 2 warning"


def test_top_level_pagination_counts_parents_not_children() -> None:
    entries = _entries()

    first = _flatten_lens_rows(
        entries, expanded=set(), child_pages={}, page=0, page_size=2
    )
    second = _flatten_lens_rows(
        entries, expanded=set(), child_pages={}, page=1, page_size=2
    )

    assert len(first) == 2
    assert len(second) == 2
    assert {row["id"] for row in first}.isdisjoint({row["id"] for row in second})


def test_child_rows_are_paged_within_an_expanded_parent(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(app_module, "_CHILD_PAGE_SIZE", 2)
    entries = _entries()
    target = next(entry for entry in entries if len(entry.children) == 3)
    parent_id = str(target.row["id"])

    first = _flatten_lens_rows(entries, expanded={parent_id}, child_pages={})
    second = _flatten_lens_rows(
        entries, expanded={parent_id}, child_pages={parent_id: 1}
    )

    assert len([row for row in first if row["kind"] == "child"]) == 2
    assert len([row for row in second if row["kind"] == "child"]) == 1
    parent_row = next(row for row in first if row["id"] == parent_id)
    assert parent_row["child_page"] == "1-2 of 3"


def test_lens_rows_never_duplicate_or_lose_an_atomic_finding() -> None:
    lens, _groups_by_id, _rationales = _lens_fixture()
    findings, _anchors = series_oracle()

    seen = [
        member.finding_id
        for cluster in lens.clusters
        for child in cluster.slices
        for member in child.members
    ]

    assert sorted(seen) == sorted(finding.finding_id for finding in findings)


def test_parent_row_template_toggles_and_omits_triage_shortcuts() -> None:
    template = app_module.REVIEW_GROUPS_BODY_SLOT
    parent, child = template.split("<q-tr v-else", 1)

    assert "aria-expanded" in parent
    assert "$parent.$emit('toggle'" in parent
    assert "keydown.enter.prevent=\"$parent.$emit('toggle'" in parent
    assert "keydown.space.prevent=\"$parent.$emit('toggle'" in parent
    assert "triage-key-action" not in parent
    assert "$parent.$emit('childpage'" in parent
    assert "triage-key-action" in child
    assert "$parent.$emit('select'" in child


# --- lens filters, story pins and keyboard semantics ---------------------------


def test_severity_filter_hides_children_and_discloses_the_count() -> None:
    entries = _entries(severities={Severity.WARNING.value})

    assert len(entries) == 4
    m_series = next(
        entry for entry in entries if "column M" in str(entry.row["location"])
    )
    assert [int(str(child["members"])) for child in m_series.children] == [1, 1]
    assert m_series.hidden == 1
    assert int(str(m_series.row["hidden_children"])) == 1
    assert int(str(m_series.row["members"])) == 2
    assert m_series.row["sevmix"] == "2 warning"


def test_a_parent_disappears_when_every_child_is_filtered_out() -> None:
    entries = _entries(severities={Severity.EXPECTED.value})

    assert entries == []


def test_text_filter_applies_at_child_level() -> None:
    entries = _entries(needle="M137")

    assert len(entries) == 1
    assert "column M" in str(entries[0].row["location"])
    assert [int(str(child["members"])) for child in entries[0].children] == [1]
    assert entries[0].hidden == 2


def test_story_pin_by_canonical_group_surfaces_every_derived_slice() -> None:
    lens, groups_by_id, _rationales = _lens_fixture()
    historical = next(
        group for group in groups_by_id.values() if group.member_count == 8
    )

    entries = _lens_entries(
        lens,
        groups_by_id,
        {},
        severities={severity.value for severity in Severity},
        finding_classes={FindingClass.VALUE_CHANGED.value},
        review_state="all",
        story_scope={historical.group_id},
    )

    assert len(entries) == 4
    surfaced = [
        child
        for entry in entries
        for child in entry.children
    ]
    assert {child["group_id"] for child in surfaced} == {historical.group_id}
    assert sum(int(str(child["members"])) for child in surfaced) == 8


def test_parent_review_state_aggregates_its_visible_children() -> None:
    findings, anchors = series_oracle()
    for finding in findings:
        if finding.location in {"M132", "M136"}:
            finding.analyst_comment = "checked"
    groups = build_pattern_groups(findings)
    lens = build_series_review_lens(groups, anchors)

    entries = _lens_entries(
        lens,
        {group.group_id: group for group in groups},
        {},
        severities={severity.value for severity in Severity},
        finding_classes={FindingClass.VALUE_CHANGED.value},
        review_state="all",
    )
    m_series = next(
        entry for entry in entries if "column M" in str(entry.row["location"])
    )

    assert m_series.row["review_state"] == "needs_review"
    assert int(str(m_series.row["reviewed"])) == 2
    assert int(str(m_series.row["reviewable_members"])) == 4

    reviewed_only = _lens_entries(
        lens,
        {group.group_id: group for group in groups},
        {},
        severities={severity.value for severity in Severity},
        finding_classes={FindingClass.VALUE_CHANGED.value},
        review_state="reviewed",
    )
    reviewed_m = next(
        entry
        for entry in reviewed_only
        if "column M" in str(entry.row["location"])
    )
    assert reviewed_m.row["review_state"] == "reviewed"
    assert reviewed_m.hidden == 2


def test_keyboard_navigation_focuses_parents_instead_of_toggling_them() -> None:
    source = inspect.getsource(app_module._render_result_view)

    assert "function activate(id)" in source
    assert "el.dataset.reviewKind === 'cluster'" in source
    assert "activate(ids[next])" in source
    assert "activate(ids[previous])" in source


def test_row_level_actions_never_resolve_a_parent_row() -> None:
    source = inspect.getsource(app_module._render_result_view)

    assert "def _reviewable_group(row_id: str)" in source
    assert "if row_id in lens.cluster_by_id:" in source
    assert "dataclasses.replace(group, members=item.members)" in source
    assert "group = _reviewable_group(group_id)" in source


# --- confirm visible ----------------------------------------------------------


def test_confirm_visible_commits_before_touching_memory() -> None:
    source = inspect.getsource(app_module._render_result_view)
    body = source.split("def open_cluster_confirmation", 1)[1].split(
        "def open_group_review", 1
    )[0]

    commit = body.index("history.set_annotations_bulk")
    mutate = body.index("finding.severity_overridden = True")
    rebuild = body.index("refresh_review_groups()")

    assert commit < mutate < rebuild
    assert "cluster_confirmation_updates(visible" in body
    assert "not an override or a replacement" in body
    assert "hidden by the\n" in body or "hidden by the " in body
    assert "Your decisions were saved. Reload this run before " in body
    assert 'ui.run_javascript("window.location.reload()")' in body


@pytest.mark.asyncio
async def test_confirm_visible_preserves_each_severity_and_skips_hidden(
    user: User,
    fixture_dir: Path,
    tmp_path: Path,
) -> None:
    work_dir = tmp_path / "work"
    artifacts = perform_run(
        work_dir,
        {
            "baseline_excel": fixture_dir / "baseline.xlsx",
            "current_excel": fixture_dir / "current.xlsx",
        },
        {},
        fixture_profile(),
    )
    history = RunHistory(work_dir / "history.sqlite3")
    assert history.get_series_anchors(artifacts.run_id)
    severities_before = {
        finding.finding_id: finding.severity
        for finding in history.get_run(artifacts.run_id).findings
    }

    create_pages(work_dir)
    await user.open(f"/runs/{artifacts.run_id}")

    group_table = next(
        element
        for element in user.find(kind=ui.table).elements
        if "review-groups-table" in element.classes
    )
    parent = next(row for row in group_table.rows if row["kind"] == "cluster")
    _emit(group_table, "toggle", {"id": str(parent["id"])})
    await user.should_see("Confirm visible")

    user.find("Confirm visible").click()
    await user.should_see("not an override or a replacement")

    user.find("Confirm listed findings").click()
    await user.should_see("Confirmed")

    annotations = history.get_annotations(artifacts.run_id)
    assert len(annotations) == int(str(parent["members"]))
    for finding_id, (severity, _comment) in annotations.items():
        engine_severity = severities_before[finding_id]
        assert engine_severity is not None
        assert severity == engine_severity.value
    assert {
        finding.severity
        for finding in history.get_run(artifacts.run_id).findings
        if finding.finding_id in annotations
    } == {
        severities_before[finding_id] for finding_id in annotations
    }


def test_guide_documents_the_related_series_lens() -> None:
    from qc_tool.ui.guide import render_guide

    text = inspect.getsource(render_guide)

    for contract in (
        "Related series",
        "navigation lens, not a decision",
        "structural only",
        "Confirm visible",
        "N of M related finding cells shown",
        "keep today's grouping",
        "new period",
        "cleared period",
        "periods affected",
        "wiped period row",
    ):
        assert contract in text


async def _open_cluster_confirmation(
    user: User, work_dir: Path, run_id: int
) -> dict[str, object]:
    create_pages(work_dir)
    await user.open(f"/runs/{run_id}")
    group_table = next(
        element
        for element in user.find(kind=ui.table).elements
        if "review-groups-table" in element.classes
    )
    parent = next(row for row in group_table.rows if row["kind"] == "cluster")
    _emit(group_table, "toggle", {"id": str(parent["id"])})
    await user.should_see("Confirm visible")
    user.find("Confirm visible").click()
    await user.should_see("not an override or a replacement")
    return parent


@pytest.mark.asyncio
async def test_confirm_visible_database_failure_leaves_memory_unchanged(
    user: User,
    fixture_dir: Path,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    work_dir = tmp_path / "work"
    artifacts = perform_run(
        work_dir,
        {
            "baseline_excel": fixture_dir / "baseline.xlsx",
            "current_excel": fixture_dir / "current.xlsx",
        },
        {},
        fixture_profile(),
    )
    history = RunHistory(work_dir / "history.sqlite3")
    await _open_cluster_confirmation(user, work_dir, artifacts.run_id)

    def boom(*args: object, **kwargs: object) -> None:
        raise RuntimeError("annotation store is unavailable")

    monkeypatch.setattr(RunHistory, "set_annotations_bulk", boom)
    user.find("Confirm listed findings").click()
    await user.should_see("annotation store is unavailable")

    assert history.get_annotations(artifacts.run_id) == {}
    assert all(
        not finding.severity_overridden and not finding.analyst_comment
        for finding in history.get_run(artifacts.run_id).findings
    )
    # the dialog stays open, so the selection is untouched
    await user.should_see("not an override or a replacement")


@pytest.mark.asyncio
async def test_confirm_visible_post_commit_failure_reports_saved_and_reloads(
    user: User,
    fixture_dir: Path,
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    caplog: pytest.LogCaptureFixture,
) -> None:
    work_dir = tmp_path / "work"
    artifacts = perform_run(
        work_dir,
        {
            "baseline_excel": fixture_dir / "baseline.xlsx",
            "current_excel": fixture_dir / "current.xlsx",
        },
        {},
        fixture_profile(),
    )
    history = RunHistory(work_dir / "history.sqlite3")
    await _open_cluster_confirmation(user, work_dir, artifacts.run_id)

    def boom(*args: object, **kwargs: object) -> list[ReviewGroup]:
        raise RuntimeError("lens rebuild failed")

    monkeypatch.setattr(app_module, "build_pattern_groups", boom)
    with caplog.at_level(logging.ERROR, logger="qc_tool.ui.app"):
        user.find("Confirm listed findings").click()
        await user.should_see("Your decisions were saved")

    assert "cluster-confirmation-rebuild-failed" in caplog.text
    caplog.clear()
    # the transaction is authoritative; nothing was rolled back
    assert history.get_annotations(artifacts.run_id)


def test_atomic_pager_slices_the_sequence_serverside() -> None:
    """Step 7: past the inline threshold the atomic tab pages server-side."""
    from qc_tool.findings_store import BLOCK_FINDINGS
    from qc_tool.ui.app import (
        _ATOMIC_INLINE_THRESHOLD,
        _ATOMIC_PAGE_SIZE,
        _atomic_page_rows,
        _FindingIndex,
    )

    # The inline threshold must not exceed the sequence page cache (4 blocks),
    # below which lazy reads hand out stable cached objects like a list.
    assert _ATOMIC_INLINE_THRESHOLD <= 4 * BLOCK_FINDINGS

    findings = [
        Finding(
            finding_id=f"F{index:04d}",
            artifact="excel",
            finding_class=FindingClass.VALUE_CHANGED,
            severity=Severity.CRITICAL,
            sheet="Data",
            location=f"B{index}",
            message=f"value changed {index}",
        )
        for index in range(1, 302)
    ]
    result = QCRunResult(profile_name="paged", findings=findings)

    first = _atomic_page_rows(result, 0)
    assert len(first) == _ATOMIC_PAGE_SIZE
    assert first[0]["id"] == "F0001"
    last = _atomic_page_rows(result, 3)
    assert [row["id"] for row in last] == ["F0301"]

    index = _FindingIndex(result.findings)
    found = index.get("F0250")
    assert found is not None and found.location == "B250"
    assert index.get("F9999") is None


def test_workbench_budget_refuses_fat_monster_runs() -> None:
    """OOM regression: the interactive workbench is bounded on 16 GB machines."""
    from qc_tool.ui.app import (
        _WORKBENCH_BUDGET_BYTES,
        _estimated_workbench_bytes,
        _workbench_within_budget,
    )

    def batch(count: int, payload: str) -> list[Finding]:
        return [
            Finding(
                finding_id=f"F{index:04d}",
                artifact="excel",
                finding_class=FindingClass.VALUE_CHANGED,
                severity=Severity.CRITICAL,
                sheet="Data",
                location=f"B{index}",
                message=payload,
            )
            for index in range(1, count + 1)
        ]

    # Below the paging threshold the estimate is zero: always interactive.
    assert _estimated_workbench_bytes(batch(50, "x" * 4000)) == 0
    assert _workbench_within_budget(batch(50, "x" * 4000))

    # A fat monster run (UNAIDS shape: ~1.4 KB raw JSON per finding at 948k
    # findings materialized to ~13.7 GB) must refuse the full workbench.
    fat = batch(30_000, "x" * 1400)
    projected = _estimated_workbench_bytes(fat)
    per_finding = projected // (len(fat) * 8)
    assert per_finding >= 1400
    assert not _workbench_within_budget(
        fat * (1 + _WORKBENCH_BUDGET_BYTES // projected)
    )

    # A skinny run of the same count stays interactive.
    assert _workbench_within_budget(batch(30_000, "value moved"))


def test_summary_review_rows_match_materialized_rows() -> None:
    """Step 2 (16 GB plan): the summary queue rows equal canonical rows."""
    from qc_tool.review import build_pattern_groups, prioritize_review
    from qc_tool.review_stream import (
        prioritize_review_summaries,
        summarize_pattern_groups_with_priority,
    )
    from qc_tool.story import build_stories
    from qc_tool.ui.app import _review_group_rows, _summary_review_row

    findings = []
    for index in range(1, 8):
        findings.append(
            Finding(
                finding_id=f"F{index:04d}",
                artifact="excel",
                finding_class=(
                    FindingClass.VALUE_CHANGED
                    if index % 2
                    else FindingClass.STYLE_CHANGED
                ),
                severity=Severity.CRITICAL if index % 3 else Severity.INFO,
                sheet="Data",
                location=f"B{index}",
                baseline_location=f"B{index}",
                message=f"changed {index}",
            )
        )
    stories = build_stories(findings)
    prioritized = prioritize_review(build_pattern_groups(findings), stories)
    canonical_rows = _review_group_rows(
        [item.group for item in prioritized],
        {item.group.group_id: item.rationale for item in prioritized},
    )
    summaries, aggregates = summarize_pattern_groups_with_priority(iter(findings))
    by_id = {finding.finding_id: finding for finding in findings}
    summary_rows = []
    for item in prioritize_review_summaries(summaries, aggregates, stories):
        message = ""
        if item.summary.member_count == 1:
            message = by_id[item.summary.member_finding_ids[0]].message
        summary_rows.append(
            _summary_review_row(item.summary, item.rationale, frozenset(), message)
        )
    assert summary_rows == canonical_rows


def test_style_key_legend_matches_the_loader_field_count(tmp_path: Path) -> None:
    """The decoded legend must track the loader's pipe-joined style key."""
    from openpyxl import Workbook

    from qc_tool.io.loader import _style_key
    from qc_tool.ui.app import _STYLE_KEY_FIELDS

    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    cell = worksheet["A1"]
    cell.value = "probe"
    key = _style_key(cell)
    assert len(key.split("|")) == len(_STYLE_KEY_FIELDS)
