"""End-to-end validation (step 15): the manifest is the contract.

Two-sided assertion over a full pipeline run:

1. every seeded manifest defect is detected as a non-expected finding, and
2. every non-expected finding is explained by a seeded defect or one of
   its documented consequences — nothing unexplained, nothing missed.

Plus: whole-pipeline read-only guarantee and a large-workbook run.
"""

import hashlib
import time
from dataclasses import dataclass, replace
from pathlib import Path
from typing import Any

import psutil
import pytest
from openpyxl import Workbook

import qc_tool.engine as engine_module
import qc_tool.io.loader as loader_module
from qc_tool.config.profile import DeliverableProfile
from qc_tool.engine import QCRunResult, run_qc
from qc_tool.findings import (
    Finding,
    FindingClass,
    FindingTemporalContext,
    Materiality,
    Severity,
)
from qc_tool.progress import CancellationToken
from qc_tool.ui.app import perform_run
from tests.conftest import fixture_profile
from tests.fixtures.manifest_schema import FixtureManifest


@dataclass(frozen=True, slots=True)
class Expectation:
    """How one seeded defect (or consequence) must surface as a finding."""

    defect_id: str
    finding_class: FindingClass
    where: str | None = None  # sheet or slide title
    token: str | None = None  # must appear in location/baseline_location/element

    def matches(self, finding: Finding) -> bool:
        if finding.finding_class is not self.finding_class:
            return False
        if self.where is not None and self.where not in (finding.sheet, finding.slide):
            return False
        if self.token is not None:
            haystack = " | ".join(
                str(part)
                for part in (
                    finding.location,
                    finding.baseline_location,
                    finding.element,
                )
                if part
            )
            if self.token not in haystack:
                return False
        return True


#: One expectation per manifest defect in the xlsx/pptx/crosscheck run.
EXPECTATIONS = [
    Expectation("E01", FindingClass.VALUE_CHANGED, "Long_Monthly", "C7"),
    Expectation("E19", FindingClass.VALUE_CHANGED, "Long_Monthly", "D4"),
    Expectation("E20", FindingClass.VALUE_CHANGED, "Long_Monthly", "C21"),
    Expectation("E02", FindingClass.FORMULA_HARDCODED, "Long_Monthly", "E10"),
    Expectation("E03", FindingClass.FORMULA_LOGIC_CHANGED, "Long_Monthly", "E14"),
    Expectation("E04", FindingClass.FORMULA_NOT_EXTENDED, "Long_Monthly", "E25"),
    Expectation("E05", FindingClass.FORMULA_ERROR, "Summary", "C5"),
    Expectation("E13", FindingClass.FORMULA_ERROR, "Summary", "D5"),
    Expectation("E06", FindingClass.VALUE_CHANGED, "Wide_Weekly", "E2"),
    Expectation("E07", FindingClass.COLUMN_DELETED, "Wide_Weekly", "column D"),
    Expectation("E08", FindingClass.FORMULA_INCONSISTENT, "Wide_Weekly", "O4"),
    Expectation("E09", FindingClass.NUMBER_FORMAT_CHANGED, "Dashboard", "B4"),
    Expectation("E11", FindingClass.HIDDEN_CHANGED, "Params"),
    Expectation("E12", FindingClass.NAMED_RANGE_CHANGED, None, "KPI_Margin"),
    Expectation("E14", FindingClass.SHEET_REMOVED, "Old_Sheet"),
    Expectation("E15", FindingClass.SHEET_ADDED, "New_Analysis"),
    Expectation("E16", FindingClass.CHART_SERIES_CHANGED, "Dashboard"),
    Expectation("E17", FindingClass.PIVOT_SOURCE_CHANGED, None, "RevenuePivot"),
    Expectation("E18", FindingClass.STYLE_CHANGED, "Dashboard", "A2"),
    Expectation("P01", FindingClass.SLIDE_ADDED, "New Initiatives"),
    Expectation("P02", FindingClass.SLIDE_REMOVED, "Deep Dive Archive"),
    Expectation("P03", FindingClass.SLIDE_TEXT_CHANGED, "Executive Summary"),
    Expectation("P04", FindingClass.TABLE_VALUE_CHANGED, "Revenue by Region", "Apr-26"),
    Expectation("P05", FindingClass.CHART_VALUE_CHANGED, "Revenue Trend", "Mar-26"),
    Expectation("X03", FindingClass.CROSSCHECK_MISMATCH, None, "Margin"),
]

#: Documented consequences of seeded defects (also legitimate findings).
CONSEQUENCES = [
    # E03's altered formula also deviates from its column pattern.
    Expectation("E03*", FindingClass.FORMULA_INCONSISTENT, "Long_Monthly", "E14"),
    # E08's deviant formula also differs from its baseline counterpart.
    Expectation("E08*", FindingClass.FORMULA_LOGIC_CHANGED, "Wide_Weekly", "O4"),
    # E05/E13 arrive in current-only columns of the Summary block.
    Expectation("E05*", FindingClass.COLUMN_INSERTED, "Summary", "column C"),
    Expectation("E13*", FindingClass.COLUMN_INSERTED, "Summary", "column D"),
]

#: Total non-expected findings on the fixture pair — pinned deliberately:
#: a change here must be a conscious engine-behavior decision.
#: 27 -> 29 (2026-08-01): E19 noise-tier and E20 restatement-tier value
#: defects were seeded; both stay visible (INFO / WARNING), never hidden.
EXPECTED_NON_EXPECTED_COUNT = 29


def _non_expected(result: QCRunResult) -> list[Finding]:
    return [f for f in result.findings if f.severity is not Severity.EXPECTED]


def test_every_manifest_defect_detected(
    qc_result: QCRunResult, manifest: FixtureManifest
) -> None:
    findings = _non_expected(qc_result)
    covered_ids = {e.defect_id for e in EXPECTATIONS}
    manifest_ids = {d.defect_id for d in manifest.defects if d.artifact != "xlsb"}
    assert manifest_ids == covered_ids, "expectation table out of sync with manifest"

    missed = [
        expectation.defect_id
        for expectation in EXPECTATIONS
        if not any(expectation.matches(f) for f in findings)
    ]
    assert missed == [], f"seeded defects not detected: {missed}"


def test_no_unexplained_findings(qc_result: QCRunResult) -> None:
    findings = _non_expected(qc_result)
    explained = EXPECTATIONS + CONSEQUENCES
    unexplained = [
        f"{f.finding_id} {f.finding_class.value} {f.sheet or f.slide} "
        f"{f.location or f.element or ''}: {f.message}"
        for f in findings
        if not any(e.matches(f) for e in explained)
    ]
    assert unexplained == [], f"findings without a seeded cause: {unexplained}"
    assert len(findings) == EXPECTED_NON_EXPECTED_COUNT


def test_materiality_tiers_reach_severity_end_to_end(qc_result: QCRunResult) -> None:
    """E01 old-history stays critical; E19 is noise/INFO; E20 restates/WARNING."""
    values = {
        f.location: f
        for f in qc_result.findings
        if f.finding_class is FindingClass.VALUE_CHANGED and f.sheet == "Long_Monthly"
    }

    assert values["C7"].materiality is Materiality.MATERIAL
    assert values["C7"].severity is Severity.CRITICAL
    assert values["D4"].materiality is Materiality.NOISE
    assert values["D4"].severity is Severity.INFO
    assert values["C21"].materiality is Materiality.MATERIAL
    assert values["C21"].temporal_context is FindingTemporalContext.RECENT_WINDOW
    assert values["C21"].severity is Severity.WARNING


def test_streaming_and_oracle_produce_equal_e2e_results(
    fixture_dir: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    baseline_excel = fixture_dir / "baseline.xlsx"
    current_excel = fixture_dir / "current.xlsx"
    baseline_ppt = fixture_dir / "baseline.pptx"
    current_ppt = fixture_dir / "current.pptx"
    streaming = run_qc(
        baseline_excel=baseline_excel,
        current_excel=current_excel,
        baseline_ppt=baseline_ppt,
        current_ppt=current_ppt,
    )

    def load_oracle(
        path: Path,
        *,
        password: str | None = None,
        allow_large_workbook: bool = False,
        cancellation_token: CancellationToken | None = None,
        formula_cache: Any = None,
        formula_engine: Any = "auto",
        _native_formula_compat_mode: bool = False,
        _xlsb_values_engine: Any = "pyxlsb",
    ) -> Any:
        return loader_module.load_workbook_snapshot(
            path,
            password=password,
            allow_large_workbook=allow_large_workbook,
            _ooxml_loader="oracle",
            cancellation_token=cancellation_token,
            formula_cache=formula_cache,
            formula_engine=formula_engine,
            _native_formula_compat_mode=_native_formula_compat_mode,
            _xlsb_values_engine=_xlsb_values_engine,
        )

    monkeypatch.setattr(engine_module, "load_workbook_snapshot", load_oracle)
    oracle = run_qc(
        baseline_excel=baseline_excel,
        current_excel=current_excel,
        baseline_ppt=baseline_ppt,
        current_ppt=current_ppt,
    )

    roles = {"baseline_excel", "current_excel"}
    assert set(streaming.values_engines) == roles
    assert set(oracle.values_engines) == roles
    assert all(
        engine.startswith("ooxml-streaming:")
        for engine in streaming.values_engines.values()
    )
    assert all(
        engine.startswith("openpyxl-oracle:")
        for engine in oracle.values_engines.values()
    )
    assert replace(streaming, values_engines={}) == replace(
        oracle,
        values_engines={},
    )


def test_xlsb_defect_detected(fixture_dir: Path, manifest: FixtureManifest) -> None:
    result = run_qc(
        baseline_excel=fixture_dir / "baseline.xlsb",
        current_excel=fixture_dir / "current.xlsb",
    )
    xb01 = manifest.defect("XB01")
    non_expected = _non_expected(result)
    value_locations = {
        f.location for f in non_expected if f.finding_class is FindingClass.VALUE_CHANGED
    }
    # C7 is the seed; E7 is its stored derived margin (values-only format).
    assert value_locations == {xb01.cell, "E7"}
    assert result.disclosures, "xlsb degradation must be disclosed"


def test_whole_pipeline_is_read_only(fixture_dir: Path, tmp_path: Path) -> None:
    tracked = sorted(p for p in fixture_dir.iterdir() if p.suffix != ".json")

    def hashes() -> dict[str, str]:
        return {p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p in tracked}

    before = hashes()
    perform_run(
        tmp_path / "work",
        {
            "baseline_excel": fixture_dir / "baseline.xlsx",
            "current_excel": fixture_dir / "current.xlsx",
            "baseline_ppt": fixture_dir / "baseline.pptx",
            "current_ppt": fixture_dir / "current.pptx",
        },
        {},
        fixture_profile(),
    )
    perform_run(
        tmp_path / "work",
        {
            "baseline_excel": fixture_dir / "baseline.xlsx",
            "current_excel": fixture_dir / "current_encrypted.xlsx",
        },
        {"current_excel": "qc-test"},
        DeliverableProfile(name="enc"),
    )
    run_qc(
        baseline_excel=fixture_dir / "baseline.xlsb",
        current_excel=fixture_dir / "current.xlsb",
    )
    assert hashes() == before


def _write_large_pair(dest: Path) -> tuple[Path, Path]:
    """A values-heavy long-format pair: ~20k rows x 10 cols per workbook."""
    months = [f"{2025 + m // 12}-{m % 12 + 1:02d}" for m in range(23)]  # ..2026-11
    regions = ["North", "South", "East", "West"]

    def build(path: Path, *, extra_rows: int, edited_row: int | None) -> None:
        workbook = Workbook()
        sheet = workbook.active
        assert sheet is not None
        sheet.title = "Data"
        sheet.append(["Period", "Region", *[f"Metric{i}" for i in range(1, 9)]])
        total = 20_000 + extra_rows
        for index in range(total):
            period = "2026-12" if index >= 20_000 else months[index % len(months)]
            row = [
                period,
                regions[index % len(regions)],
                *[float(index * 8 + metric) for metric in range(8)],
            ]
            if edited_row is not None and index == edited_row:
                row[2] += 5.0
            sheet.append(row)
        workbook.save(path)

    base, curr = dest / "large_base.xlsx", dest / "large_curr.xlsx"
    build(base, extra_rows=0, edited_row=None)
    build(curr, extra_rows=10, edited_row=10_000)
    return base, curr


def test_large_workbook_run(tmp_path: Path) -> None:
    base, curr = _write_large_pair(tmp_path)
    process = psutil.Process()
    rss_before = process.memory_info().rss
    started = time.monotonic()
    result = run_qc(baseline_excel=base, current_excel=curr)
    duration = time.monotonic() - started
    rss_after = process.memory_info().rss

    value_changes = [
        f
        for f in _non_expected(result)
        if f.finding_class is FindingClass.VALUE_CHANGED
    ]
    assert len(value_changes) == 1  # the single seeded edit, no false positives
    growth = [f for f in result.findings if f.finding_class is FindingClass.ROW_GROWTH]
    assert len(growth) == 10  # the appended 2026-12 rows

    assert duration < 90, f"large run took {duration:.1f}s"
    assert (rss_after - rss_before) < 4 * 1024**3, "memory growth exceeded 4GB"
