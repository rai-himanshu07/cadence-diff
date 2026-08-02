"""Step 6: compact dependency index with proven parity against the old graph."""

from __future__ import annotations

import json
from pathlib import Path
from typing import Any

import pytest

from qc_tool.excel.dependency import (
    build_dependency_graph,
    dependent_nodes_of,
    dependents_of,
)
from qc_tool.io.loader import load_workbook_snapshot
from qc_tool.io.model import CellRecord, SheetSnapshot, WorkbookSnapshot
from tests.dependency_oracle import build_oracle_graph, oracle_dependents

_ROOT = Path(__file__).resolve().parents[1]
_ORACLE_IMPORT = "tests.dependency_oracle"
_ALLOWED_ORACLE_IMPORTERS = {"tests/test_dependency_index.py"}


def _oracle(name: str) -> dict[str, Any]:
    return json.loads((_ROOT / "tests" / "oracles" / name).read_text(encoding="utf-8"))


def _workbook(
    cells: dict[tuple[int, int], CellRecord],
    *,
    rows: int = 60,
    columns: int = 10,
    sheet: str = "Data",
) -> WorkbookSnapshot:
    return WorkbookSnapshot(
        "synthetic.xlsx",
        "xlsx",
        True,
        True,
        formula_presence_available=True,
        formula_source="openpyxl",
        sheets=[SheetSnapshot(sheet, "visible", rows, columns, cells)],
    )


def test_step6_checkpoint_reconciles_with_the_delta_ledger() -> None:
    step5 = _oracle("real_workload_step5.json")
    step6 = _oracle("real_workload_step6.json")
    ledger = _oracle("step_delta_ledger.json")["steps"]["6"]

    assert step6["atomic_findings"] - step5["atomic_findings"] == (
        ledger["expected_atomic_delta"]
    )
    assert step6["severity"] == step5["severity"]
    assert step6["review_counts"] == step5["review_counts"]
    assert step6["coverage_states"] == step5["coverage_states"]
    assert step6["event_groups"] == {
        "axis": step5["axis_event_groups"],
        "object": step5["object_event_groups"],
    }
    assert {
        status: step6["reference_statuses"][status] - step5["reference_statuses"][status]
        for status in step5["reference_statuses"]
    } == ledger["expected_reference_status_delta"]
    assert step6["dependency_index"]["materialized_member_edges"] == (
        ledger["expected_materialized_member_edges"]
    )
    assert (
        step6["dependency_index"]["direct_edges"]
        + step6["dependency_index"]["range_descriptors"]
        < step6["projected_concrete_edges"] // 50
    )
    targets = ledger["performance_targets"]
    assert step6["median_elapsed_seconds"] <= targets["max_median_elapsed_seconds"]
    assert step6["median_peak_rss_mib"] <= targets["max_median_peak_rss_mib"]
    assert step6["trials"] >= 3
    assert step6["source_hashes_unchanged"] is True


def test_only_the_parity_tests_import_the_test_only_oracle() -> None:
    offenders = sorted(
        str(path.relative_to(_ROOT))
        for path in _ROOT.rglob("*.py")
        if "__pycache__" not in path.parts
        and _ORACLE_IMPORT in path.read_text(encoding="utf-8")
        and str(path.relative_to(_ROOT)) not in _ALLOWED_ORACLE_IMPORTERS
        and path.name != "dependency_oracle.py"
    )

    assert offenders == []


def test_production_dependency_module_does_not_import_networkx() -> None:
    source = (_ROOT / "qc_tool" / "excel" / "dependency.py").read_text(encoding="utf-8")

    assert "networkx" not in source
    assert not any(
        "networkx" in path.read_text(encoding="utf-8")
        for path in (_ROOT / "qc_tool").rglob("*.py")
    )


def test_ranges_are_stored_as_descriptors_not_per_cell_edges() -> None:
    cells: dict[tuple[int, int], CellRecord] = {
        (row, 1): CellRecord(row, 1, row) for row in range(1, 51)
    }
    cells[(1, 2)] = CellRecord(1, 2, 1275, formula="=SUM(A1:A50)")
    cells[(2, 2)] = CellRecord(2, 2, 1275, formula="=SUM(A1:A50)*2")

    graph = build_dependency_graph(_workbook(cells))

    assert graph.range_descriptor_count == 2
    assert graph.direct_edge_count == 0
    assert dependents_of(graph, "Data", "A25") == ["Data!B1", "Data!B2"]


def test_single_cell_references_remain_direct_edges() -> None:
    cells = {
        (1, 1): CellRecord(1, 1, 5),
        (1, 2): CellRecord(1, 2, 5, formula="=A1"),
        (1, 3): CellRecord(1, 3, 5, formula="=B1+1"),
    }

    graph = build_dependency_graph(_workbook(cells))

    assert graph.direct_edge_count == 2
    assert graph.range_descriptor_count == 0
    assert dependents_of(graph, "Data", "A1") == ["Data!B1", "Data!C1"]


def test_closures_are_memoized_only_for_requested_sources() -> None:
    cells = {
        (1, 1): CellRecord(1, 1, 1),
        (1, 2): CellRecord(1, 2, 1, formula="=A1"),
        (1, 3): CellRecord(1, 3, 1, formula="=B1"),
    }
    graph = build_dependency_graph(_workbook(cells))

    first = dependent_nodes_of(graph, ("Data", 1, 1))
    second = dependent_nodes_of(graph, ("Data", 1, 1))

    assert first == second == {("Data", 1, 2), ("Data", 1, 3)}
    assert set(graph._closure_cache) == {("Data", 1, 1)}


@pytest.mark.parametrize(
    "formulas",
    [
        {(1, 3): "=SUM(A1:B20)", (2, 3): "=C1*2", (3, 3): "=SUM(C1:C2)"},
        {(1, 3): "=A1+A2", (2, 3): "=SUM(A1:A5)", (3, 3): "=C1+C2"},
        {(1, 3): "=SUM(A:A)", (2, 3): "=C1", (3, 3): "=SUM(C1:C2)+A3"},
        {(1, 3): "=SUM(A1:A3)+SUM(B1:B3)", (2, 3): "=C1", (3, 3): "=C2"},
    ],
)
def test_compact_index_matches_the_oracle_on_supported_references(
    formulas: dict[tuple[int, int], str],
) -> None:
    cells: dict[tuple[int, int], CellRecord] = {
        (row, column): CellRecord(row, column, row * column)
        for row in range(1, 21)
        for column in (1, 2)
    }
    for (row, column), formula in formulas.items():
        cells[(row, column)] = CellRecord(row, column, 0, formula=formula)
    workbook = _workbook(cells)

    index = build_dependency_graph(workbook)
    oracle = build_oracle_graph(workbook)

    for row in range(1, 21):
        for column in (1, 2, 3):
            ref = f"{'ABC'[column - 1]}{row}"
            assert dependents_of(index, "Data", ref) == oracle_dependents(
                oracle, "Data", ref
            )


def test_compact_index_matches_the_oracle_on_the_generated_fixture(
    fixture_dir: Path,
) -> None:
    workbook = load_workbook_snapshot(fixture_dir / "current.xlsx")

    index = build_dependency_graph(workbook)
    oracle = build_oracle_graph(workbook)

    for sheet in workbook.sheets:
        for row, column in sheet.cells:
            ref = f"{_column_letter(column)}{row}"
            assert dependents_of(index, sheet.name, ref) == oracle_dependents(
                oracle, sheet.name, ref
            ), f"impact parity broke at {sheet.name}!{ref}"


def _column_letter(column: int) -> str:
    from openpyxl.utils import get_column_letter

    return get_column_letter(column)
