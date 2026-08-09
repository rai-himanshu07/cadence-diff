"""Step 5: analyst comparison scope for sheets and slides (disclosed, lossless load)."""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook

from qc_tool.cli import _parse_slide_list
from qc_tool.config.profile import default_profile
from qc_tool.coverage import CoverageState, QCRunMode
from qc_tool.engine import QCRunResult, run_qc
from qc_tool.findings import Finding, FindingClass
from qc_tool.history.store import RunHistory
from qc_tool.io.loader import load_workbook_snapshot
from qc_tool.io.peek import peek_sheet_names, peek_slide_titles
from qc_tool.package import PackageManifest
from qc_tool.report.json_report import result_payload
from qc_tool.review import build_pattern_groups
from qc_tool.scope import ComparisonScope
from qc_tool.triage.rules import triage


def _write_pair(tmp_path: Path) -> tuple[Path, Path]:
    def build(path: Path, edited: bool) -> None:
        workbook = Workbook()
        one = workbook.active
        assert one is not None
        one.title = "Alpha"
        one.append(["Month", "A"])
        one.append(["Jan-25", 200.0 if edited else 100.0])
        one.append(["Feb-25", 51.0])
        one.append(["Mar-25", 52.0])
        two = workbook.create_sheet("Beta")
        two.append(["Month", "B"])
        two.append(["Jan-25", 999.0 if edited else 10.0])
        two.append(["Feb-25", 11.0])
        two.append(["Mar-25", 12.0])
        workbook.save(path)

    baseline = tmp_path / "baseline.xlsx"
    current = tmp_path / "current.xlsx"
    build(baseline, edited=False)
    build(current, edited=True)
    return baseline, current


def test_sheet_scope_filters_findings_and_discloses(tmp_path: Path) -> None:
    baseline, current = _write_pair(tmp_path)
    full = run_qc(
        baseline_excel=baseline, current_excel=current, profile=default_profile()
    )
    scoped = run_qc(
        baseline_excel=baseline,
        current_excel=current,
        profile=default_profile(),
        compare_sheets=["Alpha"],
    )

    full_sheets = {f.sheet for f in full.findings if f.sheet}
    scoped_sheets = {f.sheet for f in scoped.findings if f.sheet}
    assert "Beta" in full_sheets
    assert scoped_sheets <= {"Alpha"}
    assert any(f.sheet == "Alpha" for f in scoped.findings)
    assert any("scope narrowed by analyst" in d for d in scoped.disclosures)
    assert not any("scope narrowed" in d for d in full.disclosures)


def test_scope_keeps_workbook_level_findings(tmp_path: Path) -> None:
    baseline, current = _write_pair(tmp_path)
    scoped = run_qc(
        baseline_excel=baseline,
        current_excel=current,
        profile=default_profile(),
        compare_sheets=["Alpha"],
    )
    # Nothing artifact-level lost: every retained excel finding is Alpha or unattributed.
    assert all(
        f.sheet in (None, "Alpha") for f in scoped.findings if f.artifact == "excel"
    )


def test_peek_sheet_names(tmp_path: Path) -> None:
    baseline, _ = _write_pair(tmp_path)
    assert peek_sheet_names(baseline) == ["Alpha", "Beta"]
    assert peek_sheet_names(tmp_path / "missing.xlsx") == []


def test_peek_slide_titles_unreadable_is_empty(tmp_path: Path) -> None:
    bogus = tmp_path / "deck.pptx"
    bogus.write_bytes(b"not a deck")
    assert peek_slide_titles(bogus) == []


def test_slide_scope_on_fixture_deck(fixture_dir: Path) -> None:
    baseline = fixture_dir / "baseline.pptx"
    current = fixture_dir / "current.pptx"
    titles = peek_slide_titles(current)
    assert titles, "fixture deck must peek"

    full = run_qc(
        baseline_ppt=baseline, current_ppt=current, profile=default_profile()
    )
    scoped = run_qc(
        baseline_ppt=baseline,
        current_ppt=current,
        profile=default_profile(),
        compare_slides=[1],
    )
    scoped_titles = {f.slide for f in scoped.findings if f.slide}
    assert len(scoped.findings) <= len(full.findings)
    current_titles = {title for _, title in titles}
    # Every slide-attributed finding filters by stable numeric index; only
    # genuinely deck-level findings remain unconditionally visible.
    assert {
        title for title in scoped_titles if title in current_titles
    } <= {titles[0][1]}
    assert all(
        finding.slide_index in (None, 1)
        for finding in scoped.findings
        if finding.artifact == "ppt"
    )
    removed = [
        finding
        for finding in scoped.findings
        if finding.finding_class is FindingClass.SLIDE_REMOVED
    ]
    assert removed
    assert all(finding.slide_index is None for finding in removed)
    assert all(finding.baseline_slide_index is not None for finding in removed)
    assert all(finding.baseline_location for finding in removed)


@pytest.mark.parametrize(
    "kwargs",
    [
        {"compare_sheets": []},
        {"compare_slides": []},
    ],
)
def test_explicit_empty_scope_is_invalid(
    tmp_path: Path, kwargs: dict[str, list[object]]
) -> None:
    _, current = _write_pair(tmp_path)

    with pytest.raises(ValueError, match=r"(?i)scope.*empty"):
        run_qc(
            current_excel=current,
            mode=QCRunMode.CURRENT_FILE_PREFLIGHT,
            **kwargs,  # type: ignore[arg-type]
        )


def test_scope_coverage_uses_exact_selected_total_detail(tmp_path: Path) -> None:
    baseline, current = _write_pair(tmp_path)

    result = run_qc(
        baseline_excel=baseline,
        current_excel=current,
        compare_sheets=["Alpha"],
    )
    coverage = next(
        item for item in result.coverage if item.check_id == "comparison-scope"
    )

    assert coverage.state is CoverageState.CHECKED
    assert coverage.detail == "Excel 1/2; files loaded fully"
    assert result.comparison_scope.excel_sheets == ("Alpha",)
    assert result.comparison_scope.ppt_slide_indices is None


def test_duplicate_slide_titles_filter_and_group_by_numeric_index() -> None:
    findings = triage(
        [
            Finding(
                artifact="ppt",
                finding_class=FindingClass.SLIDE_TEXT_CHANGED,
                slide="Overview",
                slide_index=index,
                message=f"slide {index} changed",
            )
            for index in (1, 2)
        ]
    )
    scope = ComparisonScope(ppt_slide_indices=(2,))

    filtered = scope.filter_findings(findings)

    assert [finding.slide_index for finding in filtered] == [2]
    assert len(build_pattern_groups(findings)) == 2


def test_scope_round_trips_history_and_json(tmp_path: Path) -> None:
    scope = ComparisonScope(excel_sheets=("Alpha",), ppt_slide_indices=(2,))
    result = QCRunResult(profile_name="scope", comparison_scope=scope)
    history = RunHistory(tmp_path / "history.sqlite3")

    run_id = history.record_run(result, file_hashes={}, report_paths={})
    stored = history.get_run(run_id)
    payload = result_payload(result)

    assert stored.comparison_scope == scope
    assert payload["comparison_scope"] == {
        "excel_sheets": ["Alpha"],
        "ppt_slide_indices": [2],
    }


def test_legacy_history_defaults_to_unrestricted_scope(tmp_path: Path) -> None:
    history = RunHistory(tmp_path / "history.sqlite3")
    run_id = history.record_run(
        QCRunResult(profile_name="legacy"), file_hashes={}, report_paths={}
    )
    import sqlite3

    with sqlite3.connect(tmp_path / "history.sqlite3") as connection:
        connection.execute(
            "UPDATE runs SET comparison_scope = '{}' WHERE id = ?", (run_id,)
        )

    assert history.get_run(run_id).comparison_scope == ComparisonScope()


def test_member_scope_validates_and_filters_duplicate_sheet_names(
    tmp_path: Path,
) -> None:
    baseline, current = _write_pair(tmp_path)
    workbooks = {
        "core": load_workbook_snapshot(baseline),
        "ops": load_workbook_snapshot(current),
    }
    scope = ComparisonScope(
        excel_member_sheets={"core": ("Alpha",), "ops": ("Beta",)}
    ).validate_loaded(workbooks_by_member=workbooks)
    findings = [
        Finding(
            artifact="excel",
            artifact_member=member,
            finding_class=FindingClass.VALUE_CHANGED,
            sheet=sheet,
            message=f"{member}/{sheet}",
        )
        for member in ("core", "ops")
        for sheet in ("Alpha", "Beta")
    ]

    filtered = scope.filter_findings(findings)

    assert [(finding.artifact_member, finding.sheet) for finding in filtered] == [
        ("core", "Alpha"),
        ("ops", "Beta"),
    ]
    assert scope.disclosure() == (
        "Comparison scope narrowed by analyst - Excel member core: Alpha; "
        "Excel member ops: Beta; unlisted Excel members use all sheets. "
        "Excel/PPT findings outside the selected scope are not reported; "
        "package analyses remain whole-package; files loaded fully."
    )
    assert scope.coverage_item(workbooks_by_member=workbooks).detail == (
        "Excel 2/4 sheets across 2/2 members; files loaded fully"
    )


def test_member_scope_rejects_unknown_members_sheets_and_legacy_scope(
    tmp_path: Path,
) -> None:
    baseline, current = _write_pair(tmp_path)
    workbooks = {
        "core": load_workbook_snapshot(baseline),
        "ops": load_workbook_snapshot(current),
    }

    with pytest.raises(ValueError, match="unknown workbook members"):
        ComparisonScope(
            excel_member_sheets={"missing": ("Alpha",)}
        ).validate_loaded(workbooks_by_member=workbooks)
    with pytest.raises(ValueError, match="unknown Excel sheet scope for member ops"):
        ComparisonScope(
            excel_member_sheets={"ops": ("Missing",)}
        ).validate_loaded(workbooks_by_member=workbooks)
    with pytest.raises(ValueError, match="legacy Excel sheet scope is ambiguous"):
        ComparisonScope(excel_sheets=("Alpha",)).validate_loaded(
            workbooks_by_member=workbooks
        )


def test_integrated_member_scope_rejects_unknown_member_before_qc(
    tmp_path: Path,
) -> None:
    _baseline, current = _write_pair(tmp_path)
    files = {"current_excel:ops": current}

    with pytest.raises(ValueError, match="unknown workbook members in scope"):
        run_qc(
            profile=default_profile(),
            package_manifest=PackageManifest.from_role_files(files),
            package_files=files,
            mode=QCRunMode.CURRENT_FILE_PREFLIGHT,
            compare_member_sheets={"missing": ("Alpha",)},
        )


def test_partial_member_scope_leaves_unlisted_members_and_crosschecks_unrestricted(
) -> None:
    scope = ComparisonScope(excel_member_sheets={"core": ("Alpha",)})
    findings = [
        Finding(
            artifact="excel",
            artifact_member=member,
            finding_class=FindingClass.VALUE_CHANGED,
            sheet=sheet,
            message=f"{member}/{sheet}",
        )
        for member in ("core", "ops")
        for sheet in ("Alpha", "Beta")
    ]
    findings.append(
        Finding(
            artifact="crosscheck",
            artifact_member="core",
            finding_class=FindingClass.CROSSCHECK_MISMATCH,
            sheet="Beta",
            message="whole-package mapping mismatch",
        )
    )

    filtered = scope.filter_findings(findings)

    assert [
        (finding.artifact, finding.artifact_member, finding.sheet)
        for finding in filtered
    ] == [
        ("excel", "core", "Alpha"),
        ("excel", "ops", "Alpha"),
        ("excel", "ops", "Beta"),
        ("crosscheck", "core", "Beta"),
    ]
    disclosure = scope.disclosure()
    assert disclosure is not None
    assert "unlisted Excel members use all sheets" in disclosure
    assert "package analyses remain whole-package" in disclosure


@pytest.mark.parametrize(
    ("raw", "expected"),
    [
        ("", None),
        ("1", [1]),
        ("1,3-5", [1, 3, 4, 5]),
        (" 2 , 4 ", [2, 4]),
    ],
)
def test_parse_slide_list(raw: str, expected: list[int] | None) -> None:
    assert _parse_slide_list(raw) == expected


@pytest.mark.parametrize("raw", ["0", "3-1", "x", "1-"])
def test_parse_slide_list_rejects_invalid(raw: str) -> None:
    with pytest.raises(ValueError):
        _parse_slide_list(raw)
