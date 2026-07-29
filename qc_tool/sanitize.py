"""Structure-preserving numeric scrambling and verified strict redaction.

Default mode scrambles numeric constants for local testing and is NOT a privacy
boundary. Strict ``redact_text=True`` also pseudonymizes visible identifiers,
removes unsupported content/metadata, and must pass the fail-closed OOXML
privacy verifier before output is returned.

Disclosed limitations (inherent to the openpyxl round-trip):

- cached formula results are cleared (Excel recalculates on open);
- pivot tables may be dropped;
- exotic chart types in decks are left unscrambled (counted as skipped).

Sources are never modified; output must be a different path.
"""

import contextlib
import datetime as dt
import hashlib
import logging
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, cast

from openpyxl import load_workbook
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.shapes import MSO_SHAPE_TYPE

from qc_tool.excel.periods import parse_period
from qc_tool.privacy import verify_sanitized
from qc_tool.security import private_directory, private_file

logger = logging.getLogger(__name__)

_NUMBER_RE = re.compile(r"[-+]?\d[\d,]*(?:\.\d+)?")
_DISPLAY_NUMBER_RE = re.compile(
    r"(?<!\w)\(?\s*[$€£₹]?\s*[-+]?\d[\d,]*(?:\.\d+)?\s*"
    r"(?:%|bn|mn|k|m|b)?\s*\)?(?!\w)",
    re.IGNORECASE,
)
_PERIOD_TOKEN_RE = re.compile(
    r"\b(?:"
    r"(?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)[A-Za-z]*[-_ ]?'?\d{2,4}"
    r"|\d{4}-\d{2}(?:-\d{2})?"
    r"|(?:CW|WK|W)[-_ ]?\d{1,2}(?:[-_ ]?'?\d{2,4})?"
    r"|Q[1-4](?:[-_ ]?(?:FY)?[-_ ]?'?\d{2,4})?"
    r")\b",
    re.IGNORECASE,
)
_FORMULA_STRING_RE = re.compile(r'"(?:[^"]|"")*"')
_FACTOR_LOW, _FACTOR_SPAN = 0.85, 0.30


class SanitizeError(Exception):
    """The file cannot be sanitized (unsupported format or bad paths)."""


@dataclass(slots=True)
class SanitizeStats:
    numbers_scrambled: int = 0
    texts_redacted: int = 0
    charts_rebuilt: int = 0
    charts_skipped: int = 0
    metadata_removed: int = 0
    comments_removed: int = 0
    links_removed: int = 0
    media_removed: int = 0
    notes_redacted: int = 0
    identifiers_redacted: int = 0
    notes: list[str] = field(default_factory=list)


def _factor(seed: int, *parts: object) -> float:
    digest = hashlib.sha256(
        "|".join([str(seed), *map(str, parts)]).encode("utf-8")
    ).digest()
    fraction = int.from_bytes(digest[:8], "big") / 2**64
    return _FACTOR_LOW + fraction * _FACTOR_SPAN


def _scramble_number(value: float | int, seed: int, *parts: object) -> float | int:
    scrambled = value * _factor(seed, *parts)
    if isinstance(value, int):
        result = round(scrambled)
        if value != 0 and result == value:
            result += 1 if value > 0 else -1
        return result
    result = round(scrambled, 6)
    if value != 0 and result == value:
        increment = max(abs(value) * 1e-6, 1e-6)
        result = value + (increment if value > 0 else -increment)
    return result


def _redact_text(value: str, seed: int, *parts: object) -> str:
    token = hashlib.sha256(
        "|".join([str(seed), value, *map(str, parts)]).encode("utf-8")
    ).hexdigest()[:6]
    return f"TXT_{token}"


def _redact_visible_text(value: str, seed: int, *parts: object) -> str:
    """Pseudonymize text while retaining period and transformed figure tokens."""
    candidates = sorted(
        [*_PERIOD_TOKEN_RE.finditer(value), *_DISPLAY_NUMBER_RE.finditer(value)],
        key=lambda match: (match.start(), -(match.end() - match.start())),
    )
    matches: list[re.Match[str]] = []
    for match in candidates:
        if matches and match.start() < matches[-1].end():
            continue
        matches.append(match)
    if not matches:
        return _redact_text(value, seed, *parts)
    output: list[str] = []
    cursor = 0
    for index, match in enumerate(matches):
        prefix = value[cursor : match.start()].strip()
        if prefix:
            output.append(_redact_text(prefix, seed, *parts, "segment", index))
        output.append(match.group(0))
        cursor = match.end()
    suffix = value[cursor:].strip()
    if suffix:
        output.append(_redact_text(suffix, seed, *parts, "segment", len(matches)))
    return " ".join(output)


def _check_paths(source: Path, dest: Path) -> None:
    if source.resolve() == dest.resolve():
        raise SanitizeError("output path must differ from the source (read-only)")
    if source.suffix.lower() != dest.suffix.lower():
        raise SanitizeError(
            f"output extension {dest.suffix!r} must match source extension "
            f"{source.suffix!r}"
        )
    if dest.exists() and source.samefile(dest):
        raise SanitizeError("output path resolves to the same file as the source")


def default_output(source: Path) -> Path:
    return source.with_name(f"{source.stem}.sanitized{source.suffix}")


def _clear_properties(properties: object) -> int:
    removed = 0
    string_fields = (
        "author",
        "creator",
        "last_modified_by",
        "lastModifiedBy",
        "title",
        "subject",
        "description",
        "comments",
        "keywords",
        "category",
        "content_status",
        "contentStatus",
        "identifier",
        "language",
        "version",
    )
    for name in string_fields:
        if not hasattr(properties, name):
            continue
        value = getattr(properties, name)
        if value:
            removed += 1
        try:
            setattr(properties, name, "")
        except (TypeError, ValueError):
            setattr(properties, name, None)
    fixed = dt.datetime(2000, 1, 1)
    for name in ("created", "modified", "last_printed", "lastPrinted"):
        if hasattr(properties, name):
            with contextlib.suppress(TypeError, ValueError):
                setattr(properties, name, fixed)
    if hasattr(properties, "revision"):
        with contextlib.suppress(TypeError, ValueError):
            cast(Any, properties).revision = 1
    return removed


def _rewrite_formula(
    formula: str, replacements: dict[str, str], seed: int, *parts: object
) -> str:
    if "[" in formula and "]" in formula:
        return "=#REF!"
    rewritten = formula
    for old, new in sorted(replacements.items(), key=lambda item: -len(item[0])):
        escaped_old = old.replace("'", "''")
        escaped_new = new.replace("'", "''")
        rewritten = rewritten.replace(f"'{escaped_old}'!", f"'{escaped_new}'!")
        rewritten = rewritten.replace(f"{old}!", f"{new}!")
        rewritten = re.sub(
            rf"(?<![A-Za-z0-9_.]){re.escape(old)}(?![A-Za-z0-9_.])",
            new,
            rewritten,
        )

    def replace_string(match: re.Match[str]) -> str:
        return f'"{_redact_text(match.group(0), seed, *parts)}"'

    return _FORMULA_STRING_RE.sub(replace_string, rewritten)


def _local_name(tag: str) -> str:
    return tag.rsplit("}", 1)[-1]


def _scrub_shape_identifiers(base_shape: object, replacement: str) -> int:
    changed = 0
    for element in base_shape._element.iter():  # type: ignore[attr-defined]
        if _local_name(element.tag) != "cNvPr":
            continue
        if element.get("name") != replacement:
            element.set("name", replacement)
            changed += 1
        for attribute in ("descr", "title"):
            if element.get(attribute):
                element.attrib.pop(attribute, None)
                changed += 1
    return changed


def _rewrite_chart_source(source: object, replacements: dict[str, str]) -> int:
    if source is None:
        return 0
    changed = 0
    for attribute in ("numRef", "strRef"):
        reference = getattr(source, attribute, None)
        if reference is None:
            continue
        formula = getattr(reference, "f", None)
        if not formula:
            continue
        rewritten = str(formula)
        for old, new in replacements.items():
            escaped_old = old.replace("'", "''")
            escaped_new = new.replace("'", "''")
            rewritten = rewritten.replace(f"'{escaped_old}'!", f"'{escaped_new}'!")
            rewritten = rewritten.replace(f"{old}!", f"{new}!")
        if rewritten != formula:
            reference.f = rewritten
            changed += 1
    return changed


def sanitize_workbook(
    source: Path, dest: Path, *, seed: int = 0, redact_text: bool = False
) -> SanitizeStats:
    """Write a scrambled copy of an xlsx workbook."""
    suffix = source.suffix.lower()
    if suffix == ".xlsm":
        raise SanitizeError(
            f"{source.name}: macro-enabled workbooks are refused because VBA can "
            "contain credentials and client identifiers; save a reviewed macro-free "
            ".xlsx copy first"
        )
    if suffix != ".xlsx":
        raise SanitizeError(
            f"{source.name}: only .xlsx can be sanitized (openpyxl cannot "
            "write other formats) — re-save the file as .xlsx first"
        )
    _check_paths(source, dest)
    stats = SanitizeStats()
    workbook = load_workbook(source, keep_links=False)
    dynamic_workbook = cast(Any, workbook)
    stats.metadata_removed += _clear_properties(workbook.properties)
    if len(dynamic_workbook.custom_doc_props):
        stats.metadata_removed += len(dynamic_workbook.custom_doc_props)
        dynamic_workbook.custom_doc_props.props.clear()
    sheet_replacements = (
        {sheet.title: f"Sheet_{index:03d}" for index, sheet in enumerate(workbook, 1)}
        if redact_text
        else {}
    )
    table_replacements: dict[str, str] = {}
    if redact_text:
        table_index = 1
        for sheet in workbook.worksheets:
            for table in sheet.tables.values():
                table_replacements[table.name] = f"Table_{table_index:03d}"
                table_index += 1
    name_replacements = (
        {
            name: f"Name_{index:03d}"
            for index, name in enumerate(workbook.defined_names, start=1)
        }
        if redact_text
        else {}
    )
    replacements = {**sheet_replacements, **table_replacements, **name_replacements}
    for sheet in workbook.worksheets:
        for header_footer_name in (
            "oddHeader",
            "evenHeader",
            "firstHeader",
            "oddFooter",
            "evenFooter",
            "firstFooter",
        ):
            header_footer = getattr(sheet, header_footer_name)
            for section_name in ("left", "center", "right"):
                section = getattr(header_footer, section_name)
                if section.text:
                    section.text = None
                    stats.metadata_removed += 1
        if sheet.data_validations.dataValidation:
            stats.metadata_removed += len(sheet.data_validations.dataValidation)
            sheet.data_validations.dataValidation = []
        conditional_rules = getattr(sheet.conditional_formatting, "_cf_rules", {})
        if conditional_rules:
            stats.metadata_removed += len(conditional_rules)
            conditional_rules.clear()
        for row in sheet.iter_rows():
            for cell in row:
                if cell.comment is not None:
                    cell.comment = None
                    stats.comments_removed += 1
                if cell.hyperlink is not None:
                    cell.hyperlink = None
                    stats.links_removed += 1
                if cell.value is None:
                    continue
                if cell.data_type == "f":
                    if redact_text:
                        cell.value = _rewrite_formula(
                            str(cell.value), replacements, seed, sheet.title, cell.coordinate
                        )
                    continue
                value = cell.value
                if isinstance(value, bool):
                    continue
                if isinstance(value, int | float):
                    cell.value = _scramble_number(
                        value, seed, sheet.title, cell.coordinate
                    )
                    stats.numbers_scrambled += 1
                elif (
                    redact_text
                    and isinstance(value, str)
                    and value.strip()
                    and parse_period(value) is None
                ):
                    scrambled_text, hits = _scramble_text_numbers(
                        value, seed, sheet.title, cell.coordinate
                    )
                    cell.value = _redact_visible_text(
                        scrambled_text, seed, sheet.title, cell.coordinate
                    )
                    stats.numbers_scrambled += hits
                    stats.texts_redacted += 1
        if redact_text:
            for chart_index, chart in enumerate(getattr(sheet, "_charts", []), start=1):
                if chart.title is not None:
                    chart.title = f"Chart_{chart_index:03d}"
                    stats.identifiers_redacted += 1
                for series in chart.series:
                    for attribute in ("val", "cat", "xVal", "yVal", "bubbleSize"):
                        stats.identifiers_redacted += _rewrite_chart_source(
                            getattr(series, attribute, None), sheet_replacements
                        )
            dynamic_sheet = cast(Any, sheet)
            stats.media_removed += len(getattr(dynamic_sheet, "_images", []))
            dynamic_sheet._images = []
            for table in sheet.tables.values():
                old_name = table.name
                if old_name in table_replacements:
                    table.name = table_replacements[old_name]
                    table.displayName = table_replacements[old_name]
                    stats.identifiers_redacted += 1
                for column_index, column in enumerate(table.tableColumns, start=1):
                    column.name = f"Column_{column_index:03d}"
                    stats.identifiers_redacted += 1
    if redact_text:
        stats.identifiers_redacted += len(workbook.defined_names)
        workbook.defined_names.clear()
        for sheet in workbook.worksheets:
            old_title = sheet.title
            sheet.title = sheet_replacements[old_title]
            stats.identifiers_redacted += 1
        for style_index, style in enumerate(dynamic_workbook._named_styles, start=1):
            if style.name != "Normal":
                style.name = f"Style_{style_index:03d}"
                stats.identifiers_redacted += 1
    private_directory(dest.parent)
    workbook.save(dest)
    private_file(dest)
    stats.notes.append("cached formula results cleared (Excel recalculates on open)")
    stats.notes.append("pivot tables may be dropped by the openpyxl round-trip")
    return stats


def _scramble_text_numbers(text: str, seed: int, *parts: object) -> tuple[str, int]:
    count = 0
    protected = [match.span() for match in _PERIOD_TOKEN_RE.finditer(text)]

    def replace(match: re.Match[str]) -> str:
        nonlocal count
        raw = match.group(0)
        if any(match.start() < end and match.end() > start for start, end in protected):
            return raw
        cleaned = raw.replace(",", "").lstrip("+")
        try:
            value = float(cleaned)
        except ValueError:
            return raw
        decimals = len(cleaned.rsplit(".", 1)[1]) if "." in cleaned else 0
        scrambled = value * _factor(seed, *parts, count, raw)
        count += 1
        if "," in raw:
            return f"{scrambled:,.{decimals}f}"
        return f"{scrambled:.{decimals}f}"

    return _NUMBER_RE.sub(replace, text), count


def sanitize_deck(
    source: Path, dest: Path, *, seed: int = 0, redact_text: bool = False
) -> SanitizeStats:
    """Write a scrambled copy of a pptx deck (text, tables, chart values)."""
    if source.suffix.lower() != ".pptx":
        raise SanitizeError(f"{source.name}: only .pptx decks can be sanitized")
    _check_paths(source, dest)
    stats = SanitizeStats()
    presentation = Presentation(str(source))
    stats.metadata_removed += _clear_properties(presentation.core_properties)
    removable_types = {
        getattr(MSO_SHAPE_TYPE, name)
        for name in ("PICTURE", "LINKED_PICTURE", "EMBEDDED_OLE_OBJECT", "MEDIA")
        if hasattr(MSO_SHAPE_TYPE, name)
    }

    def scrub_frame(frame: object, *parts: object) -> None:
        for p_idx, paragraph in enumerate(frame.paragraphs):  # type: ignore[attr-defined]
            for r_idx, run in enumerate(paragraph.runs):
                if not run.text:
                    continue
                new_text, hits = _scramble_text_numbers(
                    run.text, seed, *parts, p_idx, r_idx
                )
                if redact_text and parse_period(new_text.strip()) is None:
                    new_text = _redact_visible_text(
                        new_text, seed, *parts, p_idx, r_idx
                    )
                    stats.texts_redacted += 1
                if hits:
                    stats.numbers_scrambled += hits
                if new_text != run.text:
                    run.text = new_text
    if redact_text:
        for master_index, master in enumerate(presentation.slide_masters, start=1):
            for shape_index, base_shape in enumerate(list(master.shapes), start=1):
                stats.identifiers_redacted += _scrub_shape_identifiers(
                    base_shape, f"Shape_M{master_index}_{shape_index}"
                )
                if base_shape.shape_type in removable_types:
                    base_shape._element.getparent().remove(base_shape._element)
                    stats.media_removed += 1
                    continue
                if base_shape.has_text_frame:
                    scrub_frame(
                        cast(Any, base_shape).text_frame,
                        "master",
                        master_index,
                        shape_index,
                    )
            for layout_index, layout in enumerate(master.slide_layouts, start=1):
                for shape_index, base_shape in enumerate(list(layout.shapes), start=1):
                    stats.identifiers_redacted += _scrub_shape_identifiers(
                        base_shape,
                        f"Shape_M{master_index}_L{layout_index}_{shape_index}",
                    )
                    if base_shape.shape_type in removable_types:
                        base_shape._element.getparent().remove(base_shape._element)
                        stats.media_removed += 1
                        continue
                    if base_shape.has_text_frame:
                        scrub_frame(
                            cast(Any, base_shape).text_frame,
                            "layout",
                            master_index,
                            layout_index,
                            shape_index,
                        )
    for s_idx, slide in enumerate(presentation.slides):
        title_shape = slide.shapes.title
        if redact_text and slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame
            if notes is not None and notes.text.strip():
                notes.clear()
                stats.notes_redacted += 1
        for sh_idx, base_shape in enumerate(list(slide.shapes)):
            # python-pptx has_* guards don't narrow BaseShape for type checkers.
            shape = cast(Any, base_shape)
            if redact_text:
                stats.identifiers_redacted += _scrub_shape_identifiers(
                    base_shape, f"Shape_{s_idx + 1}_{sh_idx + 1}"
                )
            is_title = bool(
                title_shape is not None and base_shape._element is title_shape._element
            )
            if redact_text and base_shape.shape_type in removable_types:
                base_shape._element.getparent().remove(base_shape._element)
                stats.media_removed += 1
                continue
            if base_shape.has_text_frame and (redact_text or not is_title):
                scrub_frame(shape.text_frame, s_idx, sh_idx)
            if base_shape.has_table:
                for row_idx, table_row in enumerate(shape.table.rows):
                    for col_idx, cell in enumerate(table_row.cells):
                        scrub_frame(cell.text_frame, s_idx, sh_idx, row_idx, col_idx)
            if base_shape.has_chart:
                chart = shape.chart
                try:
                    plots = list(chart.plots)
                    if not plots:
                        continue
                    chart_data = CategoryChartData()
                    categories = [str(c) for c in plots[0].categories]
                    if redact_text:
                        categories = [
                            category
                            if parse_period(category) is not None
                            else _redact_text(category, seed, s_idx, sh_idx, "category", idx)
                            for idx, category in enumerate(categories)
                        ]
                    chart_data.categories = categories
                    for series_idx, series in enumerate(plots[0].series):
                        series_name = series.name or f"Series {series_idx + 1}"
                        if redact_text:
                            series_name = _redact_text(
                                series_name, seed, s_idx, sh_idx, "series", series_idx
                            )
                        chart_data.add_series(
                            series_name,
                            [
                                None
                                if value is None
                                else _scramble_number(
                                    float(value), seed, s_idx, sh_idx, series_idx, v_idx
                                )
                                for v_idx, value in enumerate(series.values)
                            ],
                        )
                        stats.numbers_scrambled += sum(
                            1 for value in series.values if value is not None
                        )
                    chart.replace_data(chart_data)
                    if redact_text and chart.has_title:
                        chart.chart_title.text_frame.text = f"Chart_{s_idx + 1}_{sh_idx + 1}"
                        stats.identifiers_redacted += 1
                    stats.charts_rebuilt += 1
                except Exception:  # chart types replace_data cannot handle
                    logger.warning("skipping chart on slide %d", s_idx + 1)
                    stats.charts_skipped += 1
    private_directory(dest.parent)
    presentation.save(str(dest))
    private_file(dest)
    return stats


def sanitize_file(
    source: Path, dest: Path | None = None, *, seed: int = 0, redact_text: bool = False
) -> tuple[Path, SanitizeStats]:
    """Dispatch on file type; returns (output path, stats)."""
    if not source.exists():
        raise SanitizeError(f"{source}: file not found")
    output = dest or default_output(source)
    suffix = source.suffix.lower()
    if suffix in {".xlsx", ".xlsm"}:
        result = sanitize_workbook(
            source, output, seed=seed, redact_text=redact_text
        )
    elif suffix == ".pptx":
        result = sanitize_deck(source, output, seed=seed, redact_text=redact_text)
    else:
        raise SanitizeError(
            f"{source.name}: unsupported format {suffix!r} — sanitize supports "
            ".xlsx/.pptx (re-save .xlsb/.xls/.xlsm as macro-free .xlsx first)"
        )
    if redact_text:
        privacy = verify_sanitized(output)
        if not privacy.safe:
            detail = "; ".join(
                f"{issue.code}@{issue.location}: {issue.message}"
                for issue in privacy.issues[:5]
            )
            output.unlink(missing_ok=True)
            raise SanitizeError(f"strict privacy verification failed: {detail}")
        result.notes.append(
            f"strict privacy verification passed ({len(privacy.checks)} checks)"
        )
    return output, result
