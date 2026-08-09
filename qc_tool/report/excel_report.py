"""Annotated Excel diff report: summary sheet + filterable findings sheet.

Consumes the triaged `QCRunResult` only — no diff logic here. The output
is a brand-new workbook; sources are never touched. The workbook is built
in openpyxl write_only mode: bounded sheets compose through a small grid
buffer, while the Findings sheet streams row by row and continues onto
``Findings (N)`` sheets past the Excel row ceiling, so report memory
follows the largest bounded sheet, never the findings population.
"""

import datetime as dt
import logging
import re
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import Cell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.hyperlink import Hyperlink
from openpyxl.worksheet.worksheet import Worksheet

from qc_tool.coverage import capability_limited
from qc_tool.engine import QCRunResult
from qc_tool.findings import Severity
from qc_tool.findings_store import finding_by_id, finding_ordinal
from qc_tool.review import format_ranges
from qc_tool.review_stream import (
    GroupSummary,
    counts_from_summaries,
    stream_stories,
    summarize_pattern_groups,
)
from qc_tool.security import private_directory, private_file

logger = logging.getLogger(__name__)

_SEVERITY_FILLS = {
    Severity.CRITICAL: PatternFill(fill_type="solid", fgColor="FFC00000"),
    Severity.WARNING: PatternFill(fill_type="solid", fgColor="FFED7D31"),
    Severity.INFO: PatternFill(fill_type="solid", fgColor="FF2E75B6"),
    Severity.EXPECTED: PatternFill(fill_type="solid", fgColor="FF548235"),
}
_WHITE_BOLD = Font(bold=True, color="FFFFFFFF")
_HEADER_FILL = PatternFill(fill_type="solid", fgColor="FF262626")
_TOP = Alignment(vertical="top", wrap_text=False)
_TOP_WRAP = Alignment(vertical="top", wrap_text=True)
_ReportCellValue = str | int | float | bool | dt.date | dt.datetime | None

#: Data rows per findings sheet: Excel's 1,048,576-row grid minus the header.
MAX_FINDINGS_DATA_ROWS = 1_048_575

_COLUMNS = [
    ("ID", 8),
    ("Severity", 12),
    ("Artifact", 10),
    ("Class", 24),
    ("Sheet / Slide", 20),
    ("Location", 14),
    ("Element", 22),
    ("Provenance", 18),
    ("Subtype", 20),
    ("Materiality", 18),
    ("Temporal context", 20),
    ("Expected reason", 22),
    ("Evidence tags", 50),
    ("Baseline", 32),
    ("Current", 32),
    ("Message", 60),
    ("Downstream impact", 40),
    ("Analyst comment", 36),
    ("Root cause", 36),
    ("Waiver", 50),
]


# Control characters openpyxl refuses because XML 1.0 cannot carry them.
_XML_ILLEGAL_CHARS = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def _safe_report_text(value: str) -> str:
    """Keep line-break intent (PPT soft breaks arrive as \v), drop the rest."""
    value = value.replace("\v", "\n").replace("\f", "\n")
    return _XML_ILLEGAL_CHARS.sub("", value)


def _dynamic_cell(
    sheet: Worksheet, *, row: int, column: int, value: _ReportCellValue
) -> Cell:
    """Write dynamic content as data, never as an executable spreadsheet formula."""
    if isinstance(value, str):
        value = _safe_report_text(value)
    cell = sheet.cell(row=row, column=column, value=value)
    if isinstance(value, str):
        cell.data_type = "s"
    return cell


@dataclass(slots=True)
class _CellSpec:
    value: str | int | float | bool | None = None
    font: Font | None = None
    fill: PatternFill | None = None
    alignment: Alignment | None = None
    link_location: str | None = None
    link_label: str | None = None
    link_tooltip: str | None = None


@dataclass(slots=True)
class _SheetGrid:
    """Random-access buffer replicating normal-mode writes for bounded sheets.

    Last write per coordinate wins, exactly like ``Worksheet.cell``, so the
    composed layout (including any deliberate overwrites) matches the old
    normal-mode report byte for byte in content.
    """

    cells: dict[tuple[int, int], _CellSpec] = field(default_factory=dict)
    max_row: int = 0
    max_column: int = 0

    def set(
        self,
        row: int,
        column: int,
        value: str | int | float | bool | None,
        *,
        font: Font | None = None,
        fill: PatternFill | None = None,
        alignment: Alignment | None = None,
    ) -> None:
        self.cells[(row, column)] = _CellSpec(
            value=value, font=font, fill=fill, alignment=alignment
        )
        self.max_row = max(self.max_row, row)
        self.max_column = max(self.max_column, column)

    def link(
        self,
        row: int,
        column: int,
        *,
        location: str,
        label: str,
        tooltip: str | None = None,
    ) -> None:
        """Workbook-local navigation without an executable formula."""
        self.cells[(row, column)] = _CellSpec(
            value=label,
            link_location=location,
            link_label=label,
            link_tooltip=tooltip,
        )
        self.max_row = max(self.max_row, row)
        self.max_column = max(self.max_column, column)

    def flush(self, sheet) -> None:
        for row in range(1, self.max_row + 1):
            emitted: list[Cell | None] = []
            last_filled = 0
            for column in range(1, self.max_column + 1):
                spec = self.cells.get((row, column))
                if spec is None:
                    emitted.append(None)
                    continue
                value = spec.value
                if isinstance(value, str):
                    value = _safe_report_text(value)
                cell = WriteOnlyCell(sheet, value=value)
                if isinstance(value, str):
                    cell.data_type = "s"
                if spec.font is not None:
                    cell.font = spec.font
                if spec.fill is not None:
                    cell.fill = spec.fill
                if spec.alignment is not None:
                    cell.alignment = spec.alignment
                if spec.link_location is not None:
                    cell.hyperlink = Hyperlink(
                        ref=f"{get_column_letter(column)}{row}",
                        location=spec.link_location,
                        display=spec.link_label,
                        tooltip=spec.link_tooltip,
                    )
                    cell.style = "Hyperlink"
                emitted.append(cell)
                last_filled = column
            sheet.append(emitted[:last_filled])


def _findings_sheet_title(sheet_index: int) -> str:
    return "Findings" if sheet_index == 0 else f"Findings ({sheet_index + 1})"


def _finding_cell_address(ordinal: int) -> tuple[str, int]:
    """(sheet title, row) of a triaged finding by its 1-based ordinal."""
    sheet_index = (ordinal - 1) // MAX_FINDINGS_DATA_ROWS
    row = (ordinal - 1) % MAX_FINDINGS_DATA_ROWS + 2
    return _findings_sheet_title(sheet_index), row


def _write_summary_sheet(
    workbook: Workbook,
    result: QCRunResult,
    counts,
    story_count: int,
    multi_member: bool,
) -> None:
    grid = _SheetGrid()
    grid.set(1, 1, "QC Report", font=Font(bold=True, size=16))
    rows: list[tuple[str, str]] = [
        ("Generated (UTC)", dt.datetime.now(dt.UTC).isoformat(timespec="seconds")),
        ("Profile", result.profile_name),
        ("Mode", result.mode.value),
        *((f"File: {role}", name) for role, name in result.files.items()),
        (
            "Result status",
            "capability-limited"
            if capability_limited(result.coverage)
            else "all required checks ran",
        ),
        ("Verified cross-checks", str(result.verified_crosschecks)),
        ("Change stories", str(story_count)),
        *(
            (f"{severity.value.title()} pattern review items", str(count))
            for severity, count in counts.review_items.items()
        ),
        *(
            (f"{severity.value.title()} atomic findings", str(count))
            for severity, count in counts.atomic_findings.items()
        ),
    ]
    for offset, (label, value) in enumerate(rows, start=3):
        grid.set(offset, 1, label, font=Font(bold=True))
        grid.set(offset, 2, value)
    if result.mapping_coverage is not None:
        mapping = result.mapping_coverage
        mapping_rows = (
            ("Eligible PPT figures", mapping.eligible),
            ("Mapped figures", mapping.mapped),
            ("Verified figures", mapping.verified),
            ("Mismatched figures", mapping.mismatched),
            ("Unresolved mappings", mapping.unresolved),
            ("Unmapped figures", mapping.unmapped),
        )
        start = 3 + len(rows)
        for offset, (label, value) in enumerate(mapping_rows, start=start):
            grid.set(offset, 1, label, font=Font(bold=True))
            grid.set(offset, 2, value)
    disclosure_row = 3 + len(rows) + 1
    for offset, disclosure in enumerate(result.disclosures):
        grid.set(
            disclosure_row + offset,
            1,
            f"NOTE: {disclosure}",
            font=Font(italic=True, color="FFC00000"),
        )
    grid.link(3, 4, location="'Stories'!A1", label="Change stories")
    grid.link(4, 4, location="'Review Groups'!A1", label="Review groups")
    grid.link(5, 4, location="'Findings'!A1", label="Atomic findings")
    grid.link(6, 4, location="'Coverage'!A1", label="Coverage")
    if result.alignment_trust is not None:
        grid.link(7, 4, location="'Alignment Trust'!A1", label="Alignment trust")
    if multi_member:
        grid.link(8, 4, location="'Package'!A1", label="Package")

    summary = workbook.create_sheet("Summary")
    summary.column_dimensions["A"].width = 28
    summary.column_dimensions["B"].width = 60
    summary.column_dimensions["D"].width = 20
    grid.flush(summary)


def _write_stories_sheet(workbook: Workbook, stories) -> None:
    grid = _SheetGrid()
    story_headers = (
        "Story",
        "Kind",
        "Title",
        "Narrative and evidence",
        "Findings",
        "Severity mix",
    )
    for col, header in enumerate(story_headers, start=1):
        grid.set(1, col, header, font=_WHITE_BOLD, fill=_HEADER_FILL)
    for row, story in enumerate(stories, start=2):
        narrative = story.description
        if story.evidence:
            narrative += "\n" + "\n".join(story.evidence)
        mix = " | ".join(
            f"{name}: {count}"
            for name, count in story.severity_counts.items()
            if count
        )
        values = (
            story.story_id,
            story.kind.value,
            story.title,
            narrative,
            story.member_count,
            mix,
        )
        for col, value in enumerate(values, start=1):
            grid.set(row, col, value)
    sheet = workbook.create_sheet("Stories")
    sheet.freeze_panes = "A2"
    sheet.column_dimensions["C"].width = 44
    sheet.column_dimensions["D"].width = 80
    grid.flush(sheet)


def _write_coverage_sheet(workbook: Workbook, result: QCRunResult) -> None:
    grid = _SheetGrid()
    coverage_headers = ("Artifact", "Check", "Status", "Findings", "Detail")
    for col, header in enumerate(coverage_headers, start=1):
        grid.set(1, col, header, font=_WHITE_BOLD, fill=_HEADER_FILL)
    for row, item in enumerate(result.coverage, start=2):
        values = (item.artifact, item.label, item.state.value, item.findings, item.detail)
        for col, value in enumerate(values, start=1):
            grid.set(row, col, value)
    sheet = workbook.create_sheet("Coverage")
    sheet.freeze_panes = "A2"
    for col, width in enumerate((14, 34, 14, 12, 60), start=1):
        sheet.column_dimensions[get_column_letter(col)].width = width
    grid.flush(sheet)


def _write_alignment_trust_sheet(workbook: Workbook, result: QCRunResult) -> None:
    if result.alignment_trust is None:
        return
    grid = _SheetGrid()
    headers = (
        "Member",
        "Sheet",
        "Region",
        "Baseline range",
        "Current range",
        "Row method",
        "Row fallback",
        "Col method",
        "Col fallback",
        "Paired rows",
        "Deleted rows",
        "Inserted rows",
        "Growth rows",
        "Paired cols",
        "Deleted cols",
        "Inserted cols",
        "Growth cols",
        "Cell pairs",
        "Skipped cells",
        "Low confidence",
    )
    for col, header in enumerate(headers, start=1):
        grid.set(1, col, header, font=_WHITE_BOLD, fill=_HEADER_FILL)
    row = 2
    for r in result.alignment_trust.regions:
        values = (
            r.artifact_member,
            r.sheet,
            r.region_id,
            r.baseline_range,
            r.current_range,
            r.row.method,
            r.row.low_confidence_fallback,
            r.column.method,
            r.column.low_confidence_fallback,
            r.row.paired,
            r.row.deleted,
            r.row.inserted,
            r.row.growth,
            r.column.paired,
            r.column.deleted,
            r.column.inserted,
            r.column.growth,
            r.comparable_cell_pairs,
            r.skipped_low_confidence_cells,
            r.low_confidence,
        )
        for col, value in enumerate(values, start=1):
            grid.set(row, col, value)
        row += 1
    if result.alignment_trust.unpaired:
        row += 1
        up_headers = ("Side", "Member", "Sheet", "Region", "Range", "Orientation")
        for col, header in enumerate(up_headers, start=1):
            grid.set(row, col, header, font=_WHITE_BOLD, fill=_HEADER_FILL)
        row += 1
        for u in result.alignment_trust.unpaired:
            values = (
                u.side,
                u.artifact_member,
                u.sheet,
                u.region_id,
                u.cell_range,
                u.orientation,
            )
            for col, value in enumerate(values, start=1):
                grid.set(row, col, value)
            row += 1
    sheet = workbook.create_sheet("Alignment Trust")
    sheet.freeze_panes = "A2"
    widths = (
        14, 18, 28, 18, 18, 12, 12, 12, 12, 12,
        12, 12, 12, 12, 12, 12, 12, 14, 14, 14,
    )
    for col, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(col)].width = width
    grid.flush(sheet)


def _write_mapping_suggestions_sheet(workbook: Workbook, result: QCRunResult) -> None:
    if not result.mapping_suggestions:
        return
    grid = _SheetGrid()
    headers = ("Slide", "Figure", "Context", "Candidate", "Value", "Display match")
    for col, header in enumerate(headers, start=1):
        grid.set(1, col, header, font=_WHITE_BOLD, fill=_HEADER_FILL)
    row = 2
    for suggestion in result.mapping_suggestions:
        if not suggestion.candidates:
            values = (
                suggestion.slide,
                suggestion.figure_raw,
                suggestion.line,
                "",
                "",
                "",
            )
            for column, value in enumerate(values, start=1):
                grid.set(row, column, value)
            row += 1
            continue
        for candidate in suggestion.candidates:
            values = (
                suggestion.slide,
                suggestion.figure_raw,
                suggestion.line,
                (
                    f"{candidate.source_member}:"
                    if candidate.source_member != "primary"
                    else ""
                )
                + f"{candidate.sheet}!{candidate.cell}",
                candidate.value,
                candidate.display_match,
            )
            for column, value in enumerate(values, start=1):
                grid.set(row, column, value)
            row += 1
    sheet = workbook.create_sheet("Mapping Suggestions")
    for col, width in enumerate((24, 14, 50, 24, 18, 14), start=1):
        sheet.column_dimensions[get_column_letter(col)].width = width
    grid.flush(sheet)


def _write_review_groups_sheet(
    workbook: Workbook,
    result: QCRunResult,
    summaries: list[GroupSummary],
    multi_member: bool,
) -> None:
    grid = _SheetGrid()
    review_headers = (
        "Group",
        "Severity",
        *(("Member",) if multi_member else ()),
        "Class",
        "Sheet / Slide",
        "Affected range",
        "Findings",
        "Baseline range",
        "Element",
        "Review item",
        "Detail",
    )
    review_widths = (
        18,
        12,
        *((14,) if multi_member else ()),
        25,
        20,
        28,
        12,
        28,
        24,
        58,
        20,
    )
    for column, header in enumerate(review_headers, start=1):
        grid.set(1, column, header, font=_WHITE_BOLD, fill=_HEADER_FILL)
    offset = 1 if multi_member else 0
    for row, group in enumerate(summaries, start=2):
        if group.member_count == 1:
            first = finding_by_id(result.findings, group.member_finding_ids[0])
            review_item = first.message if first is not None else ""
        else:
            review_item = (
                f"{group.member_count:,} "
                f"{group.finding_class.value.replace('_', ' ')} findings"
            )
        values = (
            group.group_id,
            group.severity.value,
            *((group.artifact_member,) if multi_member else ()),
            group.finding_class.value,
            group.sheet or group.slide or "",
            format_ranges(group.ranges, group.bounding_range),
            group.member_count,
            "; ".join(group.baseline_ranges),
            group.element,
            review_item,
            "",
        )
        for column, value in enumerate(values, start=1):
            grid.set(
                row,
                column,
                value,
                alignment=Alignment(
                    vertical="top",
                    wrap_text=column in {5 + offset, 7 + offset, 9 + offset},
                ),
                font=_WHITE_BOLD if column == 2 else None,
                fill=_SEVERITY_FILLS[group.severity] if column == 2 else None,
            )
        target_sheet, target_row = _finding_cell_address(
            max(finding_ordinal(group.member_finding_ids[0]), 1)
        )
        grid.link(
            row,
            10 + offset,
            location=f"'{target_sheet}'!A{target_row}",
            label="View atomic detail",
            tooltip="Clear any Findings sheet filter if the target row is hidden.",
        )
    sheet = workbook.create_sheet("Review Groups")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = (
        f"A1:{get_column_letter(len(review_headers))}{len(summaries) + 1}"
    )
    for column, width in enumerate(review_widths, start=1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    grid.flush(sheet)


def _new_findings_sheet(
    workbook: Workbook,
    finding_columns: list[tuple[str, int]],
    sheet_index: int,
    data_rows: int,
):
    sheet = workbook.create_sheet(_findings_sheet_title(sheet_index))
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = (
        f"A1:{get_column_letter(len(finding_columns))}{data_rows + 1}"
    )
    headers = []
    for col, (header, width) in enumerate(finding_columns, start=1):
        sheet.column_dimensions[get_column_letter(col)].width = width
        cell = WriteOnlyCell(sheet, value=header)
        cell.font = _WHITE_BOLD
        cell.fill = _HEADER_FILL
        headers.append(cell)
    sheet.append(headers)
    return sheet


def _write_findings_sheets(
    workbook: Workbook, result: QCRunResult, multi_member: bool
) -> None:
    """Stream every finding; continue onto follow-on sheets past the row cap."""
    finding_columns = [*_COLUMNS]
    if multi_member:
        finding_columns.insert(3, ("Member", 14))
    total = len(result.findings)
    sheet = None
    sheet_index = 0
    rows_in_sheet = 0
    overridden_seen = False
    for position, finding in enumerate(result.findings):
        if sheet is None or rows_in_sheet == MAX_FINDINGS_DATA_ROWS:
            data_rows = min(total - position, MAX_FINDINGS_DATA_ROWS)
            sheet = _new_findings_sheet(
                workbook, finding_columns, sheet_index, data_rows
            )
            sheet_index += 1
            rows_in_sheet = 0
        severity = finding.severity or Severity.WARNING
        severity_label = severity.value + (" *" if finding.severity_overridden else "")
        overridden_seen = overridden_seen or finding.severity_overridden
        values = [
            finding.finding_id,
            severity_label,
            finding.artifact,
            *(
                [finding.artifact_member]
                if multi_member
                else []
            ),
            finding.finding_class.value,
            finding.sheet or finding.slide or "",
            finding.location or finding.baseline_location or "",
            finding.element or "",
            finding.provenance.value if finding.provenance is not None else "",
            finding.subtype.value if finding.subtype is not None else "",
            finding.materiality.value if finding.materiality is not None else "",
            (
                finding.temporal_context.value
                if finding.temporal_context is not None
                else ""
            ),
            (
                finding.expected_reason.value
                if finding.expected_reason is not None
                else ""
            ),
            "; ".join(sorted(tag.value for tag in finding.evidence_tags)),
            finding.baseline_value or "",
            finding.current_value or "",
            finding.message,
            "; ".join(finding.impacts),
            finding.analyst_comment,
            finding.root_cause_key,
            (
                f"{finding.waiver_reason} (expires {finding.waiver_expires})"
                if finding.waiver_reason
                else ""
            ),
        ]
        cells = []
        for col, value in enumerate(values, start=1):
            if isinstance(value, str):
                value = _safe_report_text(value)
            cell = WriteOnlyCell(sheet, value=value)
            if isinstance(value, str):
                cell.data_type = "s"
            cell.alignment = _TOP_WRAP if col >= 8 else _TOP
            if col == 2:
                cell.fill = _SEVERITY_FILLS[severity]
                cell.font = _WHITE_BOLD
            cells.append(cell)
        sheet.append(cells)
        rows_in_sheet += 1
    if sheet is None:
        sheet = _new_findings_sheet(workbook, finding_columns, 0, 0)
    if overridden_seen:
        sheet.append(())
        note = WriteOnlyCell(sheet, value="* severity manually set by the analyst")
        note.font = Font(italic=True)
        sheet.append([note])


def _write_package_sheet(workbook: Workbook, result: QCRunResult) -> None:
    if result.package_manifest is None:
        return
    grid = _SheetGrid()
    headers = ("Side", "Artifact", "Member", "Display name", "Role key")
    for column, header in enumerate(headers, start=1):
        grid.set(1, column, header, font=_WHITE_BOLD, fill=_HEADER_FILL)
    for row, member in enumerate(result.package_manifest.members, start=2):
        for column, value in enumerate(
            (
                member.side.value,
                member.artifact.value,
                member.member_id,
                member.display_name,
                member.role_key,
            ),
            start=1,
        ):
            grid.set(row, column, value)
    sheet = workbook.create_sheet("Package")
    sheet.freeze_panes = "A2"
    for column, width in enumerate((12, 12, 18, 36, 32), start=1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    grid.flush(sheet)


def write_excel_report(result: QCRunResult, path: Path) -> None:
    summaries = summarize_pattern_groups(result.findings)
    counts = counts_from_summaries(summaries)
    stories = stream_stories(
        [iter(result.findings), iter(result.findings), iter(result.findings)]
    )
    multi_member = bool(
        result.package_manifest is not None
        and not result.package_manifest.is_legacy_projection
    )
    workbook = Workbook(write_only=True)
    _write_summary_sheet(workbook, result, counts, len(stories), multi_member)
    _write_stories_sheet(workbook, stories)
    _write_coverage_sheet(workbook, result)
    _write_alignment_trust_sheet(workbook, result)
    _write_mapping_suggestions_sheet(workbook, result)
    _write_review_groups_sheet(workbook, result, summaries, multi_member)
    _write_findings_sheets(workbook, result, multi_member)
    if multi_member:
        _write_package_sheet(workbook, result)

    private_directory(path.parent)
    workbook.save(path)
    private_file(path)
    logger.info(
        "Excel report written to %s (%d review items, %d findings)",
        path,
        len(summaries),
        len(result.findings),
    )
