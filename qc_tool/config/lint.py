"""Profile lint: validate a deliverable profile, optionally against files.

Static checks catch unparseable ranges, duplicate mapping anchors, and
expired waivers. With ``--against`` files, references are resolved for
real: sheets, cells, control ranges, tie-out references, required slides,
and cross-check mapping anchors.
"""

import datetime as dt
import re
from dataclasses import dataclass
from typing import Literal

from openpyxl.utils.cell import coordinate_to_tuple, range_boundaries

from qc_tool.config.profile import DeliverableProfile
from qc_tool.crosscheck.trace import extract_deck_figures
from qc_tool.excel.periods import parse_period
from qc_tool.io.model import WorkbookSnapshot
from qc_tool.ppt.extract import DeckSnapshot

_SHEET_REF_RE = re.compile(r"^(?:'(?P<quoted>[^']+)'|(?P<plain>[^'!]+))!(?P<ref>.+)$")


@dataclass(frozen=True, slots=True)
class LintIssue:
    level: Literal["error", "warning"]
    where: str
    message: str


def _check_range(where: str, cell_range: str, issues: list[LintIssue]) -> None:
    try:
        range_boundaries(cell_range)
    except ValueError:
        issues.append(
            LintIssue("error", where, f"range {cell_range!r} is not a valid A1 range")
        )


def _range_bounds(cell_range: str) -> tuple[int, int, int, int] | None:
    try:
        min_col, min_row, max_col, max_row = range_boundaries(cell_range)
    except ValueError:
        return None
    if min_col is None or min_row is None or max_col is None or max_row is None:
        return None
    return min_col, min_row, max_col, max_row


def _split_sheet_ref(text: str) -> tuple[str, str] | None:
    match = _SHEET_REF_RE.match(text)
    if match is None:
        return None
    return (match["quoted"] or match["plain"]), match["ref"]


def _check_sheet_cell(
    where: str,
    reference: str,
    workbook: WorkbookSnapshot | None,
    issues: list[LintIssue],
    *,
    numeric: bool = False,
) -> None:
    parts = _split_sheet_ref(reference)
    if parts is None:
        issues.append(
            LintIssue("error", where, f"{reference!r} is not a Sheet!Cell reference")
        )
        return
    sheet_name, ref = parts
    try:
        coordinate_to_tuple(ref)
    except ValueError:
        issues.append(LintIssue("error", where, f"{ref!r} is not a valid cell"))
        return
    if workbook is None:
        return
    try:
        sheet = workbook.sheet(sheet_name)
    except KeyError:
        issues.append(
            LintIssue("error", where, f"sheet {sheet_name!r} not found in workbook")
        )
        return
    cell = sheet.cell(ref)
    if cell is None:
        issues.append(
            LintIssue("error", where, f"{sheet_name}!{ref} is empty in the workbook")
        )
    elif numeric and not isinstance(cell.value, int | float):
        issues.append(
            LintIssue(
                "warning",
                where,
                f"{sheet_name}!{ref} holds no numeric value (formula without a "
                "cached result, or text)",
            )
        )


def _check_sheet_range(
    where: str,
    reference: str,
    workbook: WorkbookSnapshot | None,
    issues: list[LintIssue],
    *,
    numeric: bool = False,
) -> None:
    parts = _split_sheet_ref(reference)
    if parts is None:
        issues.append(
            LintIssue("error", where, f"{reference!r} is not a Sheet!Range reference")
        )
        return
    sheet_name, cell_range = parts
    bounds = _range_bounds(cell_range)
    if bounds is None:
        issues.append(
            LintIssue("error", where, f"{cell_range!r} is not a valid A1 range")
        )
        return
    if workbook is None:
        return
    try:
        sheet = workbook.sheet(sheet_name)
    except KeyError:
        issues.append(
            LintIssue("error", where, f"sheet {sheet_name!r} not found in workbook")
        )
        return
    min_col, min_row, max_col, max_row = bounds
    if max_row > sheet.max_row or max_col > sheet.max_column:
        issues.append(
            LintIssue(
                "error",
                where,
                f"{sheet_name}!{cell_range} extends beyond populated workbook bounds "
                f"({sheet.max_row} rows x {sheet.max_column} columns)",
            )
        )
        return
    if numeric:
        non_numeric = 0
        for row in range(min_row, max_row + 1):
            for col in range(min_col, max_col + 1):
                cell = sheet.cells.get((row, col))
                value = None if cell is None else cell.value
                if isinstance(value, bool) or not isinstance(value, int | float):
                    non_numeric += 1
        if non_numeric:
            issues.append(
                LintIssue(
                    "warning",
                    where,
                    f"{sheet_name}!{cell_range} contains {non_numeric} non-numeric or "
                    "uncached cells",
                )
            )


def lint_profile(
    profile: DeliverableProfile,
    *,
    workbook: WorkbookSnapshot | None = None,
    deck: DeckSnapshot | None = None,
) -> list[LintIssue]:
    issues: list[LintIssue] = []
    sheet_names = set(workbook.sheet_names) if workbook is not None else None

    def check_sheet_exists(where: str, name: str) -> None:
        if sheet_names is not None and name not in sheet_names:
            issues.append(
                LintIssue("error", where, f"sheet {name!r} not found in workbook")
            )

    # --- excel section -----------------------------------------------------
    for name in profile.excel.ignore_sheets:
        check_sheet_exists("excel.ignore_sheets", name)
    for sheet_name, sheet_profile in profile.excel.sheets.items():
        where = f"excel.sheets[{sheet_name}]"
        check_sheet_exists(where, sheet_name)
        for cell_range in sheet_profile.ignore_ranges:
            _check_range(f"{where}.ignore_ranges", cell_range, issues)
        for cell_range in sheet_profile.refresh_ranges:
            _check_range(f"{where}.refresh_ranges", cell_range, issues)
        for region in sheet_profile.regions:
            _check_range(f"{where}.regions", region.cell_range, issues)
        for band in sheet_profile.cadence_bands:
            _check_range(f"{where}.cadence_bands", band.cell_range, issues)
        for index, rule in enumerate(sheet_profile.availability_rules):
            rule_where = f"{where}.availability_rules[{rule.name or index}]"
            _check_range(f"{rule_where}.range", rule.cell_range, issues)
            _check_range(f"{rule_where}.periods", rule.period_range, issues)
            if parse_period(rule.required_through) is None:
                issues.append(
                    LintIssue(
                        "error",
                        rule_where,
                        (
                            f"required_through {rule.required_through!r} is not "
                            "a supported period"
                        ),
                    )
                )
            target_bounds = _range_bounds(rule.cell_range)
            period_bounds = _range_bounds(rule.period_range)
            if target_bounds is None or period_bounds is None:
                continue
            target_min_col, target_min_row, target_max_col, target_max_row = (
                target_bounds
            )
            period_min_col, period_min_row, period_max_col, period_max_row = (
                period_bounds
            )
            target_width = target_max_col - target_min_col + 1
            target_height = target_max_row - target_min_row + 1
            period_width = period_max_col - period_min_col + 1
            period_height = period_max_row - period_min_row + 1
            if period_width > 1 and period_height > 1:
                issues.append(
                    LintIssue(
                        "error",
                        rule_where,
                        "periods must be a single row or single column",
                    )
                )
            elif period_height == 1 and period_width > 1 and period_width != target_width:
                issues.append(
                    LintIssue(
                        "error",
                        rule_where,
                        (
                            f"period range width {period_width} does not match "
                            f"target range width {target_width}"
                        ),
                    )
                )
            elif period_width == 1 and period_height > 1 and period_height != target_height:
                issues.append(
                    LintIssue(
                        "error",
                        rule_where,
                        (
                            f"period range height {period_height} does not match "
                            f"target range height {target_height}"
                        ),
                    )
                )
            if workbook is not None:
                try:
                    availability_sheet = workbook.sheet(sheet_name)
                except KeyError:
                    availability_sheet = None
                required_period = parse_period(rule.required_through)
                if availability_sheet is not None and required_period is not None:
                    observed_kinds: set[str] = set()
                    unparseable = 0
                    for row in range(period_min_row, period_max_row + 1):
                        for column in range(period_min_col, period_max_col + 1):
                            cell = availability_sheet.cells.get((row, column))
                            if cell is None or cell.value in {None, ""}:
                                continue
                            period = parse_period(cell.value)
                            if period is None:
                                unparseable += 1
                            else:
                                observed_kinds.add(period.kind)
                    if unparseable:
                        issues.append(
                            LintIssue(
                                "error",
                                rule_where,
                                (
                                    f"period range contains {unparseable} "
                                    "unparseable nonblank labels"
                                ),
                            )
                        )
                    mismatched_kinds = observed_kinds - {required_period.kind}
                    if mismatched_kinds:
                        issues.append(
                            LintIssue(
                                "error",
                                rule_where,
                                (
                                    f"period range kinds {sorted(observed_kinds)} "
                                    "do not match required_through kind "
                                    f"{required_period.kind!r}"
                                ),
                            )
                        )

    controls = profile.excel.controls
    for group_name, group in (
        ("required_ranges", controls.required_ranges),
        ("unique_ranges", controls.unique_ranges),
        ("numeric_bounds", controls.numeric_bounds),
    ):
        for control in group:
            where = f"excel.controls.{group_name}[{control.name or control.sheet}]"
            check_sheet_exists(where, control.sheet)
            _check_range(where, control.cell_range, issues)
            if workbook is not None:
                _check_sheet_range(
                    where,
                    f"{control.sheet}!{control.cell_range}",
                    workbook,
                    issues,
                    numeric=group_name == "numeric_bounds",
                )
            if group_name == "numeric_bounds":
                minimum = getattr(control, "minimum", None)
                maximum = getattr(control, "maximum", None)
                if minimum is None and maximum is None:
                    issues.append(
                        LintIssue("error", where, "numeric bounds define no minimum or maximum")
                    )
                elif minimum is not None and maximum is not None and minimum > maximum:
                    issues.append(
                        LintIssue(
                            "error",
                            where,
                            f"minimum {minimum} exceeds maximum {maximum}",
                        )
                    )
    for tie_out in controls.tie_outs:
        where = f"excel.controls.tie_outs[{tie_out.name}]"
        _check_sheet_cell(where, tie_out.target, workbook, issues, numeric=True)
        if not tie_out.components:
            issues.append(LintIssue("error", where, "tie-out has no components"))
        for component in tie_out.components:
            _check_sheet_range(where, component, workbook, issues, numeric=True)
        if tie_out.absolute_tolerance < 0 or tie_out.relative_tolerance < 0:
            issues.append(LintIssue("error", where, "tie-out tolerances cannot be negative"))

    # --- crosscheck mappings -------------------------------------------------
    seen_anchors: set[tuple[str, str, int]] = set()
    deck_occurrences = extract_deck_figures(deck) if deck is not None else None
    for mapping in profile.crosscheck.mappings:
        where = f"crosscheck.mappings[{mapping.label or mapping.line_skeleton}]"
        anchor = (mapping.slide, mapping.line_skeleton, mapping.figure_index)
        if anchor in seen_anchors:
            issues.append(LintIssue("error", where, "duplicate mapping anchor"))
        seen_anchors.add(anchor)
        _check_sheet_cell(
            where,
            f"{mapping.source_sheet}!{mapping.source_cell}",
            workbook,
            issues,
            numeric=True,
        )
        if deck_occurrences is not None and not any(
            occ.slide == mapping.slide
            and occ.line_skeleton == mapping.line_skeleton
            and occ.figure_index == mapping.figure_index
            for occ in deck_occurrences
        ):
            issues.append(
                LintIssue(
                    "error",
                    where,
                    "mapping anchor does not resolve in the deck (slide or "
                    "wording changed) — re-confirm it",
                )
            )

    # --- ppt section -----------------------------------------------------------
    for index, rule in enumerate(profile.ppt.availability_rules):
        where = f"ppt.availability_rules[{rule.name or index}]"
        if parse_period(rule.required_through) is None:
            issues.append(
                LintIssue(
                    "error",
                    where,
                    (
                        f"required_through {rule.required_through!r} is not "
                        "a supported period"
                    ),
                )
            )
    if deck is not None:
        titles = {slide.title for slide in deck.slides if slide.title}
        slides_by_display = {
            slide.display_name.casefold(): slide for slide in deck.slides
        }
        for required in profile.ppt.required_slides:
            if required not in titles:
                issues.append(
                    LintIssue(
                        "error",
                        "ppt.required_slides",
                        f"required slide {required!r} not found in deck",
                    )
                )
        for index, rule in enumerate(profile.ppt.availability_rules):
            rule_where = f"ppt.availability_rules[{rule.name or index}]"
            slide = slides_by_display.get(rule.slide.casefold())
            if slide is None:
                issues.append(
                    LintIssue(
                        "error",
                        rule_where,
                        f"slide {rule.slide!r} not found in deck",
                    )
                )
                continue
            element_series: list[tuple[str, set[str]]] = []
            if rule.scope == "table":
                element_series = [
                    (
                        table.name or f"table[{table.source_index}]",
                        {
                            row[0]
                            for row in table.rows[1:]
                            if row and row[0]
                        },
                    )
                    for table in slide.tables
                ]
            elif rule.scope == "chart":
                element_series = [
                    (
                        chart.title
                        or chart.name
                        or f"chart[{chart.source_index}]",
                        {
                            series.name
                            for series in chart.all_series
                            if series.name
                        },
                    )
                    for chart in slide.charts
                ]
            matching_elements = [
                (element, series_names)
                for element, series_names in element_series
                if not rule.element
                or element.casefold() == rule.element.casefold()
            ]
            if rule.element and not matching_elements:
                issues.append(
                    LintIssue(
                        "error",
                        rule_where,
                        (
                            f"{rule.scope} element {rule.element!r} not found "
                            f"on slide {rule.slide!r}"
                        ),
                    )
                )
                continue
            if rule.series and not any(
                rule.series.casefold()
                in {name.casefold() for name in series_names}
                for _, series_names in matching_elements
            ):
                issues.append(
                    LintIssue(
                        "error",
                        rule_where,
                        (
                            f"series/row {rule.series!r} not found in matching "
                            f"{rule.scope} elements"
                        ),
                    )
                )

    # --- waivers ------------------------------------------------------------------
    today = dt.date.today()
    for waiver in profile.waivers:
        where = f"waivers[{waiver.finding_class.value}]"
        if waiver.expires < today:
            issues.append(
                LintIssue(
                    "warning", where, f"waiver expired on {waiver.expires.isoformat()}"
                )
            )
    return issues
