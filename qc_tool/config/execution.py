"""Engine-native execution bindings resolved from a
``ResolvedInputConfigurationV1`` (plan-20260913, Step 3).

Bridges the per-run resolved configuration to the alignment/structural/
formula comparison engine: confirmed sheet renames and execution-derived
keyed-region identity rules. Keyed regions are adapted into the existing,
already-proven ``RowIdentityRule`` shape so they flow through
``qc_tool.excel.align``'s current keyed-alignment machinery unchanged,
rather than duplicating it.

Every consumer of ``ExecutionBindings`` treats it as fully optional
(``None`` = do exactly what happened before this module existed), so a
legacy run with no saved ``input_contract`` -- and therefore no resolved
member/sheet/region bindings -- stays byte-identical.

Step 8 extends the adapter to consume ``ResolvedRegion``'s per-side
preamble/footer/first-data-row facts directly, producing genuinely
independent ``baseline_header_row``/``baseline_footer_row`` overrides
(``qc_tool.excel.align``'s bottom/top-aligned positional comparison) instead
of assuming one shared boundary for both sides.

Step 12 (Fix 4) extends the adapter with ``confirmed_column_mappings``: a
region's per-column ``ResolvedColumn.baseline_letter`` override (Step 12
Fix 3) now drives the alignment engine's column axis directly, not merely
storage/reporting -- mirroring ``confirmed_sheet_renames`` exactly, one
level down.
"""

from __future__ import annotations

from typing import Literal

from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import coordinate_to_tuple, range_boundaries

from qc_tool.config.profile import RowIdentityRule
from qc_tool.config.resolved_input import (
    ResolvedInputConfigurationV1,
    ResolvedMember,
    ResolvedRegion,
)


def _top_left(cell_range: str) -> tuple[int, int] | None:
    """(row, col) of a range or single-cell string's top-left corner, or
    ``None`` if unparseable.
    """
    if not cell_range:
        return None
    try:
        if ":" in cell_range:
            min_col, min_row, _, _ = range_boundaries(cell_range)
            if min_col is None or min_row is None:
                return None
            return (min_row, min_col)
        row, col = coordinate_to_tuple(cell_range)
        return (row, col)
    except ValueError:
        return None


def _bottom_row(cell_range: str) -> int | None:
    """The max row of a range or single-cell string, or ``None`` if
    unparseable.
    """
    if not cell_range:
        return None
    try:
        if ":" in cell_range:
            _, _, _, max_row = range_boundaries(cell_range)
            return max_row
        row, _ = coordinate_to_tuple(cell_range)
        return row
    except ValueError:
        return None


def _side_header_row(
    *, outer_range: str | None, first_data_row: int | None, preamble_rows: int
) -> int | None:
    """The last preamble row for one side, or ``None`` when this side has no
    explicit preamble boundary at all (an untouched automatic region).
    """
    if first_data_row is not None:
        return first_data_row - 1
    if preamble_rows > 0 and outer_range:
        top = _top_left(outer_range)
        if top is not None:
            return top[0] + preamble_rows - 1
    return None


def _side_footer_row(*, outer_range: str | None, footer_rows: int) -> int | None:
    """The first footer row for one side, or ``None`` when this side has no
    explicit footer boundary.
    """
    if footer_rows > 0 and outer_range:
        bottom = _bottom_row(outer_range)
        if bottom is not None:
            return bottom - footer_rows + 1
    return None


def region_as_row_identity_rule(region: ResolvedRegion) -> RowIdentityRule | None:
    """Adapt one keyed, execution-confirmed ``ResolvedRegion`` into a
    ``RowIdentityRule``, or ``None`` when this region carries no
    execution-confirmed row identity (mode is not ``"keyed"``, or no column
    has ``alignment_role == "identity"``).
    """
    if region.mode != "keyed":
        return None
    identity_columns = [
        column.current_letter
        for column in region.columns
        if column.alignment_role == "identity" and column.current_letter
    ]
    if not identity_columns:
        return None
    ordinal_columns = [
        column.current_letter
        for column in region.columns
        if column.alignment_role == "ordinal" and column.current_letter
    ]
    anchor_source = region.current_data_range or region.current_outer_range
    if not anchor_source:
        return None
    anchor = _top_left(anchor_source)
    if anchor is None:
        return None
    anchor_row, anchor_col = anchor
    anchor_cell = f"{get_column_letter(anchor_col)}{anchor_row}"
    current_header_row = _side_header_row(
        outer_range=region.current_outer_range,
        first_data_row=region.current_first_data_row,
        preamble_rows=region.current_preamble_rows,
    )
    current_footer_row = _side_footer_row(
        outer_range=region.current_outer_range, footer_rows=region.current_footer_rows
    )
    baseline_header_row = _side_header_row(
        outer_range=region.baseline_outer_range,
        first_data_row=region.baseline_first_data_row,
        preamble_rows=region.baseline_preamble_rows,
    )
    baseline_footer_row = _side_footer_row(
        outer_range=region.baseline_outer_range, footer_rows=region.baseline_footer_rows
    )
    trim_identity_whitespace = any(
        column.trim_outer_whitespace
        for column in region.columns
        if column.alignment_role == "identity"
    )
    return RowIdentityRule(
        anchor_cell=anchor_cell,
        header_row=current_header_row,
        footer_row=current_footer_row,
        # Only carry an explicit per-side override when this region actually
        # has independent baseline-side facts; otherwise the alignment
        # layer's own "fall back to the shared value" default applies.
        baseline_header_row=baseline_header_row if region.baseline_outer_range else None,
        baseline_footer_row=baseline_footer_row if region.baseline_outer_range else None,
        identity_columns=identity_columns,
        ordinal_columns=ordinal_columns,
        duplicate_policy=region.duplicate_key_policy,
        # This adapter is the ONLY producer of execution-confirmed rules
        # from the mode-aware configuration workspace; every rule saved
        # through the legacy YAML profile editor bypasses this function
        # entirely and keeps its original strip+casefold behavior (Step 8's
        # "exact typed equality plus optional outer-whitespace trim only"
        # criterion applies only to this new pathway, never retroactively).
        exact_typed_equality=True,
        trim_identity_whitespace=trim_identity_whitespace,
    )


class RegionColumnPolicies:
    """One region's non-normal column comparison policies, keyed by its
    stable anchor cell (Step 8) -- consumed by
    ``qc_tool.excel.diff_values.region_range_sets`` the same way a
    confirmed ``RowIdentityRule``'s ordinal columns already are.
    """

    __slots__ = ("anchor_cell", "expected_refresh_columns", "ignore_columns")

    def __init__(
        self,
        anchor_cell: str,
        ignore_columns: tuple[str, ...] = (),
        expected_refresh_columns: tuple[str, ...] = (),
    ) -> None:
        self.anchor_cell = anchor_cell
        self.ignore_columns = ignore_columns
        self.expected_refresh_columns = expected_refresh_columns


class RegionDisposition:
    """One explicit non-automatic region decision keyed by anchor cell."""

    __slots__ = ("anchor_cell", "baseline_anchor_cell", "mode")
    anchor_cell: str
    baseline_anchor_cell: str | None
    mode: Literal["positional", "excluded"]

    def __init__(
        self,
        anchor_cell: str,
        mode: Literal["positional", "excluded"],
        baseline_anchor_cell: str | None = None,
    ) -> None:
        self.anchor_cell = anchor_cell
        self.baseline_anchor_cell = baseline_anchor_cell
        self.mode = mode


def region_disposition(region: ResolvedRegion) -> RegionDisposition | None:
    if region.mode == "positional":
        mode: Literal["positional", "excluded"] = "positional"
    elif region.mode == "excluded":
        mode = "excluded"
    else:
        return None
    anchor_source = region.current_data_range or region.current_outer_range
    if not anchor_source:
        return None
    anchor = _top_left(anchor_source)
    if anchor is None:
        return None
    anchor_row, anchor_col = anchor
    baseline_anchor_source = region.baseline_data_range or region.baseline_outer_range
    baseline_anchor = (
        _top_left(baseline_anchor_source) if baseline_anchor_source else None
    )
    return RegionDisposition(
        anchor_cell=f"{get_column_letter(anchor_col)}{anchor_row}",
        mode=mode,
        baseline_anchor_cell=(
            f"{get_column_letter(baseline_anchor[1])}{baseline_anchor[0]}"
            if baseline_anchor is not None
            else None
        ),
    )


def region_column_policies(region: ResolvedRegion) -> RegionColumnPolicies | None:
    """Adapt one region's ignore/expected-refresh column columns into
    ``RegionColumnPolicies``, or ``None`` when it declares neither.
    """
    ignore_columns = tuple(
        column.current_letter
        for column in region.columns
        if column.comparison_policy == "ignore" and column.current_letter
    )
    expected_refresh_columns = tuple(
        column.current_letter
        for column in region.columns
        if column.comparison_policy == "expected_refresh" and column.current_letter
    )
    if not ignore_columns and not expected_refresh_columns:
        return None
    anchor_source = region.current_data_range or region.current_outer_range
    if not anchor_source:
        return None
    anchor = _top_left(anchor_source)
    if anchor is None:
        return None
    anchor_row, anchor_col = anchor
    return RegionColumnPolicies(
        anchor_cell=f"{get_column_letter(anchor_col)}{anchor_row}",
        ignore_columns=ignore_columns,
        expected_refresh_columns=expected_refresh_columns,
    )


class ConfirmedColumnMapping:
    """One region's confirmed baseline-letter overrides, keyed by its
    stable anchor cell (plan-20260913, Step 12 Fix 4) -- consumed by
    ``qc_tool.excel.align``'s column axis the same way a confirmed
    ``RowIdentityRule`` is matched by anchor for the row axis.
    """

    __slots__ = ("anchor_cell", "mapping")

    def __init__(self, anchor_cell: str, mapping: dict[str, str]) -> None:
        self.anchor_cell = anchor_cell
        #: ``{current_letter: baseline_letter}`` -- every entry is a
        #: genuine move (the two letters differ); a column with no
        #: baseline-letter override never appears here.
        self.mapping = mapping


def region_confirmed_column_mapping(region: ResolvedRegion) -> ConfirmedColumnMapping | None:
    """Adapt one region's per-column baseline-letter overrides into a
    ``ConfirmedColumnMapping``, or ``None`` when every column in this
    region shares the same letter on both sides (today's default).
    """
    mapping = {
        column.current_letter: column.baseline_letter
        for column in region.columns
        if (
            column.current_letter
            and column.baseline_letter
            and column.current_letter != column.baseline_letter
        )
    }
    if not mapping:
        return None
    anchor_source = region.current_data_range or region.current_outer_range
    if not anchor_source:
        return None
    anchor = _top_left(anchor_source)
    if anchor is None:
        return None
    anchor_row, anchor_col = anchor
    return ConfirmedColumnMapping(
        anchor_cell=f"{get_column_letter(anchor_col)}{anchor_row}",
        mapping=mapping,
    )


class ExecutionBindings:
    """Read-only lookup over one run's resolved logical member/sheet/region
    bindings, organized for direct consumption by ``qc_tool.excel.align``,
    ``qc_tool.excel.diff_structure``, and ``qc_tool.excel.formulas``.
    """

    __slots__ = ("_members",)

    def __init__(self, resolved: ResolvedInputConfigurationV1) -> None:
        members: dict[str, ResolvedMember] = {}
        for member in resolved.members:
            members[member.member_id] = member
        self._members = members

    def confirmed_sheet_renames(self, member_id: str) -> dict[str, str]:
        """``{current_sheet_name: baseline_sheet_name}`` for every sheet of
        ``member_id`` whose two physical names both resolved and genuinely
        differ. A sheet sharing one name on both sides is ordinary
        same-name pairing, never a rename.
        """
        member = self._members.get(member_id)
        if member is None:
            return {}
        renames: dict[str, str] = {}
        for sheet in member.sheets:
            if (
                sheet.baseline_sheet_name
                and sheet.current_sheet_name
                and sheet.baseline_sheet_name != sheet.current_sheet_name
            ):
                renames[sheet.current_sheet_name] = sheet.baseline_sheet_name
        return renames

    def column_policies(
        self, member_id: str, current_sheet_name: str
    ) -> tuple[RegionColumnPolicies, ...]:
        """Ignore/expected-refresh column policies for every region on
        ``current_sheet_name`` that declares at least one (Step 8).
        """
        member = self._members.get(member_id)
        if member is None:
            return ()
        policies: list[RegionColumnPolicies] = []
        for sheet in member.sheets:
            if sheet.current_sheet_name != current_sheet_name:
                continue
            for region in sheet.regions:
                policy = region_column_policies(region)
                if policy is not None:
                    policies.append(policy)
        return tuple(policies)

    def region_dispositions(
        self, member_id: str, current_sheet_name: str
    ) -> tuple[RegionDisposition, ...]:
        """Explicit positional/excluded decisions for one physical sheet."""
        member = self._members.get(member_id)
        if member is None:
            return ()
        dispositions: list[RegionDisposition] = []
        for sheet in member.sheets:
            if sheet.current_sheet_name != current_sheet_name:
                continue
            for region in sheet.regions:
                disposition = region_disposition(region)
                if disposition is not None:
                    dispositions.append(disposition)
        return tuple(dispositions)

    def confirmed_column_mappings(
        self, member_id: str, current_sheet_name: str
    ) -> tuple[ConfirmedColumnMapping, ...]:
        """Confirmed baseline-letter overrides for every region on
        ``current_sheet_name`` that declares at least one (Step 12 Fix 4).

        Mirrors ``row_identity_rules``'s lookup shape exactly; consumed by
        ``qc_tool.excel.align``'s column axis so a confirmed column move
        actually changes which physical baseline cell a current cell is
        compared against, not merely what is stored/reported.
        """
        member = self._members.get(member_id)
        if member is None:
            return ()
        mappings: list[ConfirmedColumnMapping] = []
        for sheet in member.sheets:
            if sheet.current_sheet_name != current_sheet_name:
                continue
            for region in sheet.regions:
                mapping = region_confirmed_column_mapping(region)
                if mapping is not None:
                    mappings.append(mapping)
        return tuple(mappings)

    def row_identity_rules(
        self, member_id: str, current_sheet_name: str
    ) -> tuple[RowIdentityRule, ...]:
        """Execution-confirmed keyed regions on ``current_sheet_name``,
        adapted into ``RowIdentityRule`` so they flow through the existing,
        already-proven keyed-alignment machinery unchanged.
        """
        member = self._members.get(member_id)
        if member is None:
            return ()
        rules: list[RowIdentityRule] = []
        for sheet in member.sheets:
            if sheet.current_sheet_name != current_sheet_name:
                continue
            for region in sheet.regions:
                rule = region_as_row_identity_rule(region)
                if rule is not None:
                    rules.append(rule)
        return tuple(rules)

    def logical_sheet_id(
        self, member_id: str, current_sheet_name: str
    ) -> str | None:
        """The saved logical ``sheet_id`` bound to this physical current
        sheet name, or ``None`` when unresolved. Used to build a bounded,
        content-free ``LogicalFindingAddress``.
        """
        member = self._members.get(member_id)
        if member is None:
            return None
        for sheet in member.sheets:
            if sheet.current_sheet_name == current_sheet_name:
                return sheet.sheet_id
        return None


def build_execution_bindings(
    resolved: ResolvedInputConfigurationV1 | None,
) -> ExecutionBindings | None:
    """``None`` in, ``None`` out -- the universal legacy-preserving guard
    every engine call site uses before doing anything binding-aware.
    """
    if resolved is None or resolved.source == "legacy_default":
        return None
    return ExecutionBindings(resolved)
