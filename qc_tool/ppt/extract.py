"""Read-only PowerPoint extraction with complete native chart semantics."""

from __future__ import annotations

import hashlib
import logging
from pathlib import Path
from typing import Any, cast

from pptx import Presentation
from pptx.oxml.ns import qn

from qc_tool.io.decrypt import open_decrypted
from qc_tool.ppt.chart_xml import PptChartParseError, parse_ppt_chart
from qc_tool.ppt.model import (
    ChartContent,
    DeckSnapshot,
    PptChartPlot,
    PptChartSeries,
    ShapeContent,
    SlideContent,
    TableContent,
)
from qc_tool.progress import CancellationToken, check_cancelled

logger = logging.getLogger(__name__)

__all__ = [
    "ChartContent",
    "DeckSnapshot",
    "PptChartPlot",
    "PptChartSeries",
    "ShapeContent",
    "SlideContent",
    "TableContent",
    "load_deck_snapshot",
]


def _geometry(shape: Any) -> tuple[int, int, int, int]:
    return (
        int(shape.left or 0),
        int(shape.top or 0),
        int(shape.width or 0),
        int(shape.height or 0),
    )


def _shape_type(shape: Any) -> str:
    shape_type = shape.shape_type
    return str(getattr(shape_type, "name", shape_type))


def _shape_source_id(
    slide_index: int,
    shape_index: int,
    shape_type: str,
    geometry: tuple[int, int, int, int],
) -> str:
    left, top, width, height = geometry
    return (
        f"slide[{slide_index}]/shape[{shape_index}]:{shape_type}"
        f"@{left}:{top}:{width}:{height}"
    )


def _shape_texts(shape: Any) -> list[str]:
    if not shape.has_text_frame:
        return []
    return [
        paragraph.text
        for paragraph in shape.text_frame.paragraphs
        if paragraph.text.strip()
    ]


def _placeholder_type(shape: Any) -> str | None:
    if not shape.is_placeholder:
        return None
    placeholder_type = shape.placeholder_format.type
    return str(getattr(placeholder_type, "name", placeholder_type))


def _embedded_media(shape: Any) -> tuple[str | None, str | None, str | None]:
    """Return media kind, SHA-256, and a fixed degradation code for one shape."""
    blips = list(shape._element.xpath(".//a:blip"))
    if not blips:
        return None, None, None
    if len(blips) != 1:
        return "unavailable-image", None, "multiple-image-relationships"
    blip = blips[0]
    embed_id = blip.get(qn("r:embed"))
    link_id = blip.get(qn("r:link"))
    if not embed_id:
        return (
            "linked-image" if link_id else "unavailable-image",
            None,
            "linked-image" if link_id else "missing-image-relationship",
        )
    try:
        part = shape.part.related_part(embed_id)
        blob = bytes(part.blob)
        kind = str(part.content_type)
    except (AttributeError, KeyError, TypeError, ValueError):
        return "unavailable-image", None, "unreadable-image-relationship"
    return kind, hashlib.sha256(blob).hexdigest(), None


def _notes(slide: Any) -> list[str]:
    if not slide.has_notes_slide:
        return []
    frame = slide.notes_slide.notes_text_frame
    if frame is None:
        return []
    return [
        paragraph.text
        for paragraph in frame.paragraphs
        if paragraph.text.strip()
    ]


def load_deck_snapshot(
    path: Path,
    *,
    password: str | None = None,
    cancellation_token: CancellationToken | None = None,
) -> DeckSnapshot:
    check_cancelled(cancellation_token)
    stream = open_decrypted(path, password)
    presentation = Presentation(stream)
    snapshot = DeckSnapshot(source_name=path.name)
    chart_errors: list[str] = []
    media_errors: list[str] = []
    for index, slide in enumerate(presentation.slides):
        check_cancelled(cancellation_token)
        title_shape = slide.shapes.title
        title = title_shape.text if title_shape is not None else None
        try:
            notes = _notes(slide)
        except (AttributeError, KeyError, TypeError, ValueError) as exc:
            notes = []
            snapshot.notes_available = False
            snapshot.notes_detail = (
                f"Speaker notes unavailable on slide {index + 1}: {exc}"
            )
            logger.warning(
                "%s: PowerPoint notes extraction failed at slide %d: %s",
                path.name,
                index + 1,
                exc,
            )
        content = SlideContent(
            index=index,
            title=title,
            texts=[],
            shape_count=len(slide.shapes),
            notes=notes,
        )
        chart_index = 0
        table_index = 0
        for shape_index, base_shape in enumerate(slide.shapes):
            shape = cast(Any, base_shape)
            geometry = _geometry(shape)
            shape_type = _shape_type(shape)
            source_id = _shape_source_id(index, shape_index, shape_type, geometry)
            texts = _shape_texts(shape)
            media_kind, media_digest, media_error = _embedded_media(shape)
            if media_error is not None:
                media_errors.append(
                    f"slide {index + 1} shape {shape_index + 1}: {media_error}"
                )
            content.shapes.append(
                ShapeContent(
                    source_id=source_id,
                    source_index=shape_index,
                    shape_id=int(shape.shape_id),
                    shape_type=shape_type,
                    name=str(shape.name),
                    left=geometry[0],
                    top=geometry[1],
                    width=geometry[2],
                    height=geometry[3],
                    z_order=shape_index,
                    texts=texts,
                    is_placeholder=bool(shape.is_placeholder),
                    placeholder_type=_placeholder_type(shape),
                    media_kind=media_kind,
                    media_digest=media_digest,
                )
            )
            if base_shape.has_text_frame and base_shape is not title_shape:
                content.texts.extend(texts)
            if base_shape.has_table:
                content.tables.append(
                    TableContent(
                        rows=[
                            [cell.text for cell in row.cells]
                            for row in shape.table.rows
                        ],
                        source_id=source_id,
                        source_index=table_index,
                        shape_id=int(shape.shape_id),
                        name=str(shape.name),
                        left=geometry[0],
                        top=geometry[1],
                        width=geometry[2],
                        height=geometry[3],
                        z_order=shape_index,
                    )
                )
                table_index += 1
            if base_shape.has_chart:
                check_cancelled(cancellation_token)
                chart_part = cast(Any, shape.chart.part)
                try:
                    content.charts.append(
                        parse_ppt_chart(
                            chart_part.blob,
                            source_id=source_id,
                            source_index=chart_index,
                            shape_id=int(shape.shape_id),
                            name=str(shape.name),
                            left=geometry[0],
                            top=geometry[1],
                            width=geometry[2],
                            height=geometry[3],
                            z_order=shape_index,
                            source_part=str(chart_part.partname),
                        )
                    )
                except PptChartParseError as exc:
                    chart_errors.append(
                        f"slide {index + 1}, shape {shape_index + 1}: {exc}"
                    )
                    logger.warning(
                        "%s: PowerPoint chart extraction failed at slide %d shape %d: %s",
                        path.name,
                        index + 1,
                        shape_index + 1,
                        exc,
                    )
                chart_index += 1
        snapshot.slides.append(content)
    if chart_errors:
        snapshot.charts_available = False
        snapshot.chart_detail = "; ".join(chart_errors)
    if media_errors:
        snapshot.media_available = False
        snapshot.media_detail = "; ".join(media_errors)
    return snapshot
