"""Package-aware strict sanitization for a current Excel/PPT deliverable pair."""

import dataclasses
import datetime as dt
import hashlib
from collections.abc import Mapping
from pathlib import Path
from typing import Any, cast

from pptx import Presentation
from pydantic import BaseModel, Field

from qc_tool.config.profile import CrosscheckMapping, DeliverableProfile
from qc_tool.crosscheck.numbers import format_figure_like, replace_figure_ordinals
from qc_tool.crosscheck.trace import (
    extract_deck_figures,
    mapping_identity,
    occurrence_identity,
    verify_mappings,
)
from qc_tool.findings import Finding
from qc_tool.io.loader import load_workbook_snapshot
from qc_tool.package import (
    MEMBER_ID_PATTERN,
    PackageArtifact,
    PackageManifest,
    PackageMember,
    PackageSide,
    paths_by_member,
    structural_member_aliases,
)
from qc_tool.ppt.extract import load_deck_snapshot
from qc_tool.ppt.shapes import iter_text_leaf_shapes
from qc_tool.privacy import PrivacyReport, verify_sanitized
from qc_tool.sanitize import SanitizeError, sanitize_file
from qc_tool.security import private_directory, private_file


class PackageMappingResult(BaseModel):
    mapping_digest: str
    source_sheet_ordinal: int | None = None
    source_cell: str
    occurrence_ordinal: int | None = None
    status: str
    detail: str = ""
    source_member: str = Field(
        default="primary",
        pattern=MEMBER_ID_PATTERN,
        exclude_if=lambda value: value == "primary",
    )


class PackageSanitizeManifest(BaseModel):
    schema_version: int = 1
    generated_at: str
    seed: int
    privacy_safe: bool
    mappings_verified: int = 0
    mappings_unverifiable: int = 0
    inputs: dict[str, dict[str, str | int]]
    outputs: dict[str, dict[str, str | int]]
    mapping_results: list[PackageMappingResult] = Field(default_factory=list)
    sanitizer_stats: dict[str, dict] = Field(default_factory=dict)
    privacy: dict[str, PrivacyReport]
    package_manifest: PackageManifest | None = Field(
        default=None,
        exclude_if=lambda value: value is None,
    )


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _mapping_digest(
    mapping: CrosscheckMapping,
    *,
    source_member: str | None = None,
) -> str:
    payload: tuple[object, ...] = (
        mapping.slide,
        mapping.line_skeleton,
        mapping.figure_index,
        mapping.source_sheet,
        mapping.source_cell,
    )
    member_id = mapping.source_member if source_member is None else source_member
    if member_id != "primary":
        payload = (*payload, member_id)
    return hashlib.sha256(repr(payload).encode()).hexdigest()[:20]


def _sanitized_output_name(
    member: PackageMember,
    source: Path,
    *,
    legacy_projection: bool,
    member_alias: str,
) -> str:
    prefix = "workbook" if member.artifact is PackageArtifact.EXCEL else "deck"
    member_suffix = "" if legacy_projection else f"-{member_alias}"
    return f"{prefix}{member_suffix}.sanitized{source.suffix.casefold()}"


def _replace_deck_figures(path: Path, replacements: dict[int, str]) -> None:
    presentation = Presentation(str(path))
    figure_index = 0
    for slide in presentation.slides:
        title_shape = slide.shapes.title
        title_element = (
            title_shape._element if title_shape is not None else None
        )
        for base_shape in slide.shapes:
            for text_shape in iter_text_leaf_shapes(cast(Any, base_shape)):
                if text_shape._element is title_element:
                    continue
                for paragraph in text_shape.text_frame.paragraphs:
                    if not paragraph.text.strip():
                        continue
                    updated, figure_index = replace_figure_ordinals(
                        paragraph.text, replacements, start_index=figure_index
                    )
                    if updated != paragraph.text:
                        paragraph.text = updated
        # `extract_deck_figures` enumerates every ordinary text occurrence
        # before table cells, so rewriting must use the same two-pass order.
        for base_shape in slide.shapes:
            shape = cast(Any, base_shape)
            if base_shape.has_table:
                for row in list(shape.table.rows)[1:]:
                    for cell in list(row.cells)[1:]:
                        updated, figure_index = replace_figure_ordinals(
                            cell.text, replacements, start_index=figure_index
                        )
                        if updated != cell.text:
                            cell.text = updated
    presentation.save(str(path))
    private_file(path)


def sanitize_package(
    excel_source: Path,
    ppt_source: Path,
    output_dir: Path,
    profile: DeliverableProfile,
    *,
    seed: int = 0,
    forbidden_tokens: list[str] | None = None,
) -> PackageSanitizeManifest:
    """Strictly sanitize one legacy Excel/PPT pair with v1-compatible evidence."""
    files = {
        "current_excel": excel_source,
        "current_ppt": ppt_source,
    }
    return sanitize_package_files(
        files,
        PackageManifest.from_role_files(files),
        output_dir,
        profile,
        seed=seed,
        forbidden_tokens=forbidden_tokens,
    )


def sanitize_package_files(
    package_files: Mapping[str, Path],
    package_manifest: PackageManifest,
    output_dir: Path,
    profile: DeliverableProfile,
    *,
    seed: int = 0,
    forbidden_tokens: list[str] | None = None,
) -> PackageSanitizeManifest:
    """Strictly sanitize a bounded current package and preserve mapping topology."""
    paths_by_member(package_files, package_manifest)
    if package_manifest.members_for(PackageSide.BASELINE, PackageArtifact.EXCEL):
        raise ValueError("sanitize-package accepts current workbook members only")
    if package_manifest.members_for(PackageSide.BASELINE, PackageArtifact.PPT):
        raise ValueError("sanitize-package accepts a current PowerPoint only")
    workbook_members = package_manifest.members_for(
        PackageSide.CURRENT,
        PackageArtifact.EXCEL,
    )
    deck_members = package_manifest.members_for(
        PackageSide.CURRENT,
        PackageArtifact.PPT,
    )
    if not workbook_members or len(deck_members) != 1:
        raise ValueError(
            "sanitize-package needs one or more current workbooks and one current deck"
        )

    output_dir = private_directory(output_dir)
    manifest_path = output_dir / "redaction-manifest.json"
    legacy_projection = package_manifest.is_legacy_projection
    member_aliases = structural_member_aliases(package_manifest)
    unknown_member_aliases = {
        member_id: f"excel_missing_{index:03d}"
        for index, member_id in enumerate(
            sorted(
                {
                    mapping.source_member
                    for mapping in profile.crosscheck.mappings
                    if mapping.source_member
                    not in {member.member_id for member in workbook_members}
                }
            ),
            start=1,
        )
    }

    def source_alias(member_id: str) -> str:
        alias = member_aliases.get((PackageArtifact.EXCEL, member_id))
        return alias if alias is not None else unknown_member_aliases[member_id]

    output_paths = {
        member.role_key: output_dir
        / _sanitized_output_name(
            member,
            package_files[member.role_key],
            legacy_projection=legacy_projection,
            member_alias=member_aliases[(member.artifact, member.member_id)],
        )
        for member in package_manifest.members
    }
    deck_member = deck_members[0]
    ppt_source = package_files[deck_member.role_key]
    ppt_output = output_paths[deck_member.role_key]

    original_deck = load_deck_snapshot(ppt_source)
    original_occurrences = extract_deck_figures(original_deck)
    occurrence_indices = {
        occurrence_identity(occurrence): index
        for index, occurrence in enumerate(original_occurrences)
    }
    _, ppt_stats = sanitize_file(ppt_source, ppt_output, seed=seed, redact_text=True)

    replacements: dict[int, str] = {}
    mapping_results: list[PackageMappingResult] = []
    pending: list[tuple[CrosscheckMapping, int, int, str]] = []
    sanitizer_stats: dict[str, dict] = {
        deck_member.role_key: dataclasses.asdict(ppt_stats)
    }
    known_workbook_members = {member.member_id for member in workbook_members}
    for mapping in profile.crosscheck.mappings:
        if mapping.source_member in known_workbook_members:
            continue
        mapping_results.append(
            PackageMappingResult(
                mapping_digest=_mapping_digest(
                    mapping,
                    source_member=source_alias(mapping.source_member),
                ),
                source_member=source_alias(mapping.source_member),
                source_cell=mapping.source_cell,
                occurrence_ordinal=occurrence_indices.get(
                    mapping_identity(mapping)
                ),
                status="unverifiable",
                detail="mapping source member is not present in the package",
            )
        )

    for member in workbook_members:
        excel_source = package_files[member.role_key]
        excel_output = output_paths[member.role_key]
        original_workbook = load_workbook_snapshot(excel_source)
        _, excel_stats = sanitize_file(
            excel_source,
            excel_output,
            seed=seed,
            redact_text=True,
        )
        sanitizer_stats[member.role_key] = dataclasses.asdict(excel_stats)
        sanitized_workbook = load_workbook_snapshot(excel_output)
        for mapping in profile.crosscheck.mappings:
            if mapping.source_member != member.member_id:
                continue
            safe_member_id = source_alias(mapping.source_member)
            digest = _mapping_digest(mapping, source_member=safe_member_id)
            ordinal = occurrence_indices.get(mapping_identity(mapping))
            try:
                sheet_ordinal = (
                    original_workbook.sheet_names.index(mapping.source_sheet) + 1
                )
            except ValueError:
                sheet_ordinal = None
            if ordinal is None or sheet_ordinal is None:
                mapping_results.append(
                    PackageMappingResult(
                        mapping_digest=digest,
                        source_member=safe_member_id,
                        source_sheet_ordinal=sheet_ordinal,
                        source_cell=mapping.source_cell,
                        occurrence_ordinal=ordinal,
                        status="unverifiable",
                        detail="mapping anchor or source sheet does not resolve",
                    )
                )
                continue
            source_sheet = f"Sheet_{sheet_ordinal:03d}"
            cell = sanitized_workbook.sheet(source_sheet).cell(mapping.source_cell)
            value = None if cell is None else cell.value
            if not isinstance(value, bool) and isinstance(value, int | float):
                replacements[ordinal] = format_figure_like(
                    original_occurrences[ordinal].figure,
                    float(value),
                )
                pending.append((mapping, ordinal, sheet_ordinal, digest))
                continue
            mapping_results.append(
                PackageMappingResult(
                    mapping_digest=digest,
                    source_member=safe_member_id,
                    source_sheet_ordinal=sheet_ordinal,
                    source_cell=mapping.source_cell,
                    occurrence_ordinal=ordinal,
                    status="unverifiable",
                    detail="sanitized source cell has no numeric cached value",
                )
            )

    if replacements:
        _replace_deck_figures(ppt_output, replacements)

    sanitized_deck = load_deck_snapshot(ppt_output)
    sanitized_occurrences = extract_deck_figures(sanitized_deck)
    sanitized_mappings: list[CrosscheckMapping] = []
    pending_by_digest: dict[str, tuple[int, int]] = {}
    for original_mapping, ordinal, sheet_ordinal, digest in pending:
        if ordinal >= len(sanitized_occurrences):
            mapping_results.append(
                PackageMappingResult(
                    mapping_digest=digest,
                    source_member=source_alias(original_mapping.source_member),
                    source_sheet_ordinal=sheet_ordinal,
                    source_cell=original_mapping.source_cell,
                    occurrence_ordinal=ordinal,
                    status="unverifiable",
                    detail="sanitized deck occurrence ordering changed",
                )
            )
            continue
        occurrence = sanitized_occurrences[ordinal]
        sanitized_mappings.append(
            CrosscheckMapping(
                slide=occurrence.slide,
                line_skeleton=occurrence.line_skeleton,
                figure_index=occurrence.figure_index,
                label=f"mapping_{digest}",
                source_sheet=f"Sheet_{sheet_ordinal:03d}",
                source_cell=original_mapping.source_cell,
                source_member=original_mapping.source_member,
            )
        )
        pending_by_digest[digest] = (ordinal, sheet_ordinal)

    verified_labels: set[str] = set()
    findings_by_label: dict[str, Finding] = {}
    for member in workbook_members:
        member_mappings = [
            mapping
            for mapping in sanitized_mappings
            if mapping.source_member == member.member_id
        ]
        if not member_mappings:
            continue
        sanitized_profile = profile.crosscheck.model_copy(
            update={"mappings": member_mappings}
        )
        sanitized_workbook = load_workbook_snapshot(output_paths[member.role_key])
        verification = verify_mappings(
            sanitized_deck,
            sanitized_workbook,
            sanitized_profile,
        )
        verified_labels.update(mapping.label for mapping in verification.verified)
        findings_by_label.update(
            {finding.element or "": finding for finding in verification.findings}
        )
    for mapping in sanitized_mappings:
        digest = mapping.label.removeprefix("mapping_")
        ordinal, sheet_ordinal = pending_by_digest[digest]
        finding = findings_by_label.get(mapping.label)
        verified = mapping.label in verified_labels
        mapping_results.append(
            PackageMappingResult(
                mapping_digest=digest,
                source_member=source_alias(mapping.source_member),
                source_sheet_ordinal=sheet_ordinal,
                source_cell=mapping.source_cell,
                occurrence_ordinal=ordinal,
                status="verified" if verified else "unverifiable",
                detail="" if verified else (finding.message if finding else "verification failed"),
            )
        )

    tokens = forbidden_tokens or []
    privacy_by_role = {
        member.role_key: verify_sanitized(
            output_paths[member.role_key],
            forbidden_tokens=tokens,
        )
        for member in package_manifest.members
    }
    privacy_safe = all(report.safe for report in privacy_by_role.values())
    verified_count = sum(item.status == "verified" for item in mapping_results)
    unverifiable_count = sum(item.status != "verified" for item in mapping_results)
    def redacted_member(member: PackageMember) -> PackageMember:
        alias = member_aliases[(member.artifact, member.member_id)]
        return PackageMember(
            member_id=alias,
            side=member.side,
            artifact=member.artifact,
            display_name=output_paths[member.role_key].name,
        )

    def evidence_key(member: PackageMember) -> str:
        return (
            member.artifact.value
            if legacy_projection
            else redacted_member(member).role_key
        )

    sanitized_manifest = PackageManifest(
        members=tuple(redacted_member(member) for member in package_manifest.members)
    )
    manifest = PackageSanitizeManifest(
        schema_version=1 if legacy_projection else 2,
        generated_at=dt.datetime.now(dt.UTC).isoformat(timespec="seconds"),
        seed=seed,
        privacy_safe=privacy_safe,
        mappings_verified=verified_count,
        mappings_unverifiable=unverifiable_count,
        inputs={
            evidence_key(member): {
                "sha256": _sha256(package_files[member.role_key]),
                "size": package_files[member.role_key].stat().st_size,
            }
            for member in package_manifest.members
        },
        outputs={
            evidence_key(member): {
                "sha256": _sha256(output_paths[member.role_key]),
                "size": output_paths[member.role_key].stat().st_size,
            }
            for member in package_manifest.members
        },
        mapping_results=sorted(mapping_results, key=lambda item: item.mapping_digest),
        sanitizer_stats={
            evidence_key(member): sanitizer_stats[member.role_key]
            for member in package_manifest.members
        },
        privacy={
            evidence_key(member): privacy_by_role[member.role_key]
            for member in package_manifest.members
        },
        package_manifest=None if legacy_projection else sanitized_manifest,
    )
    manifest_path.write_text(manifest.model_dump_json(indent=2), encoding="utf-8")
    private_file(manifest_path)
    if not privacy_safe:
        raise SanitizeError("package privacy verification failed; see redaction-manifest.json")
    return manifest
