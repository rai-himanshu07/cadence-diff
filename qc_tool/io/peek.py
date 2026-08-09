"""Lightweight name peeks for scope selection — never a full load.

Sheet names come from the workbook manifest only; slide titles from a
python-pptx open. Failures (encryption, corruption) return empty lists so
callers disable scoping instead of guessing.
"""

from __future__ import annotations

import io
import logging
import zipfile
from pathlib import Path
from xml.etree import ElementTree

logger = logging.getLogger(__name__)

_MAIN_NS = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"


def peek_sheet_names(path: Path) -> list[str]:
    """Worksheet names of an xlsx/xlsm/xlsb workbook, or [] when unreadable."""
    suffix = path.suffix.lower()
    try:
        if suffix in (".xlsx", ".xlsm"):
            with zipfile.ZipFile(path) as archive:
                root = ElementTree.fromstring(archive.read("xl/workbook.xml"))
            return [
                name
                for sheet in root.iter(f"{_MAIN_NS}sheet")
                if (name := sheet.get("name"))
            ]
        if suffix == ".xlsb":
            from qc_tool.io.xlsb_formula import _workbook_sheets

            with zipfile.ZipFile(path) as archive:
                data = archive.read("xl/workbook.bin")
            return [name for name, _ in _workbook_sheets(data)]
    except Exception as exc:
        logger.info("sheet peek unavailable for %s: %s", path.name, exc)
    return []


def peek_slide_titles(path: Path) -> list[tuple[int, str]]:
    """(1-based index, title) per slide, or [] when unreadable."""
    try:
        from pptx import Presentation

        deck = Presentation(io.BytesIO(path.read_bytes()))
        titles: list[tuple[int, str]] = []
        for index, slide in enumerate(deck.slides, start=1):
            title = ""
            shape = getattr(slide.shapes, "title", None)
            if shape is not None and getattr(shape, "has_text_frame", False):
                # collapse soft line breaks (\v) and newlines into a one-line label
                title = " ".join(shape.text_frame.text.split())
            titles.append((index, title or f"Slide {index}"))
        return titles
    except Exception as exc:
        logger.info("slide peek unavailable for %s: %s", path.name, exc)
        return []
