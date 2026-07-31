"""Read-only workbook loading into `WorkbookSnapshot` for xlsx/xlsm/xlsb.

xlsx/xlsm: one read-only openpyxl pass supplies formulas, constants, styles,
and names. Raw package scans supply cached formula results, worksheet metadata,
tables, charts, pivots, and interaction rules. Pivot descriptors are parsed
from the raw package parts
(``xl/pivotTables/*.xml`` + ``xl/pivotCache/pivotCacheDefinition*.xml``)
because openpyxl's pivot object model is unreliable for arbitrary files.

xlsb: cached values via pyxlsb plus an independent BIFF12 formula-record
scan. Platform adapters may enrich formula text later; cached values always
remain those read from the original xlsb.
"""

import io
import logging
import posixpath
import re
import sys
import zipfile
from collections.abc import Iterator
from pathlib import Path
from typing import Any, Literal
from xml.etree import ElementTree

from openpyxl import load_workbook
from openpyxl.cell.read_only import ReadOnlyCell
from openpyxl.styles.numbers import is_date_format, is_timedelta_format
from openpyxl.utils.cell import column_index_from_string, coordinate_to_tuple
from openpyxl.utils.datetime import from_excel, from_ISO8601
from pyxlsb import open_workbook as open_xlsb

from qc_tool.io.decrypt import open_decrypted
from qc_tool.io.formula_enrichment import (
    FormulaEnrichmentError,
    FormulaExtraction,
    merge_formula_extraction,
)
from qc_tool.io.model import (
    CellRecord,
    CellValue,
    NamedRange,
    PivotDescriptor,
    SheetSnapshot,
    TableDescriptor,
    WorkbookSnapshot,
    WorkbookWorkload,
    is_cell_value,
)
from qc_tool.io.ooxml_chart import ChartParseError, parse_ooxml_charts
from qc_tool.io.ooxml_interaction import extract_worksheet_interactions
from qc_tool.io.ooxml_worksheet import (
    OOXMLMetadataError,
    WorkbookMetadata,
    WorksheetMetadata,
    parse_ooxml_worksheet_metadata,
)
from qc_tool.io.xlsb_formula import (
    XlsbFormulaScan,
    XlsbFormulaScanError,
    scan_xlsb_formulas,
)
from qc_tool.progress import CancellationToken, check_cancelled

logger = logging.getLogger(__name__)

SUPPORTED_SUFFIXES = {".xlsx", ".xlsm", ".xlsb"}

_PIVOT_TABLE_RE = re.compile(r"^xl/pivotTables/pivotTable\d+\.xml$")
_PIVOT_CACHE_RE = re.compile(r"^xl/pivotCache/pivotCacheDefinition\d+\.xml$")
_MIB = 1024 * 1024
_WORKLOAD_LIMITS = (
    ("cell_count", 1_000_000, 5_000_000, "physical cells"),
    ("worksheet_xml_bytes", 128 * _MIB, 512 * _MIB, "worksheet XML bytes"),
    ("shared_string_bytes", 64 * _MIB, 256 * _MIB, "shared-string bytes"),
    ("styles_bytes", 8 * _MIB, 32 * _MIB, "style XML bytes"),
    ("style_count", 50_000, 250_000, "cell styles"),
    ("largest_sheet_area", 10_000_000, 50_000_000, "largest sheet area"),
)

#: BIFF12 BOOLERR codes (pyxlsb yields them as hex strings) -> error literals.
_XLSB_ERRORS = {
    "0x0": "#NULL!",
    "0x7": "#DIV/0!",
    "0xf": "#VALUE!",
    "0x17": "#REF!",
    "0x1d": "#NAME?",
    "0x24": "#NUM!",
    "0x2a": "#N/A",
}


class UnsupportedFormatError(Exception):
    """The file extension is not a supported workbook format."""


class OOXMLCellStreamError(ValueError):
    """Formula and cached-value streams disagree on physical cell inventory."""


class OOXMLWorkloadError(ValueError):
    """An OOXML package exceeds safe local processing limits."""


def _constant_cell_value(value: object) -> CellValue:
    if is_cell_value(value):
        return value
    return str(value)


def _formula_text(value: object) -> str:
    text = getattr(value, "text", None)
    return text if isinstance(text, str) else str(value)


def _without_chart_drawings(data: bytes) -> bytes:
    """Remove chart anchors from an in-memory OOXML copy before openpyxl reads it."""
    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(data)) as source, zipfile.ZipFile(output, "w") as target:
        for member in source.infolist():
            content = source.read(member)
            if member.filename.startswith("xl/drawings/") and member.filename.endswith(".xml"):
                root = ElementTree.fromstring(content)
                for child in list(root):
                    if any(_local_name(item.tag) == "chart" for item in child.iter()):
                        root.remove(child)
                content = ElementTree.tostring(
                    root,
                    encoding="utf-8",
                    xml_declaration=True,
                )
            target.writestr(member, content)
    return output.getvalue()


def load_workbook_snapshot(
    path: Path,
    *,
    password: str | None = None,
    allow_large_workbook: bool = False,
    _ooxml_loader: Literal["streaming", "oracle"] = "streaming",
    cancellation_token: CancellationToken | None = None,
) -> WorkbookSnapshot:
    """Load any supported workbook into a snapshot without touching the source."""
    check_cancelled(cancellation_token)
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise UnsupportedFormatError(
            f"{path.name}: unsupported format {suffix!r}; expected one of "
            f"{sorted(SUPPORTED_SUFFIXES)}"
        )
    stream = open_decrypted(path, password)
    data = stream.getvalue()
    check_cancelled(cancellation_token)
    if suffix == ".xlsb":
        return _load_xlsb(
            data,
            source_name=path.name,
            cancellation_token=cancellation_token,
        )
    loader = (
        _load_ooxml_streaming
        if _ooxml_loader == "streaming"
        else _load_ooxml_oracle
    )
    return loader(
        data,
        source_name=path.name,
        file_format=suffix.lstrip("."),
        allow_large_workbook=allow_large_workbook,
        cancellation_token=cancellation_token,
    )


# --- xlsx / xlsm ---------------------------------------------------------


def _assess_ooxml_workload(
    metadata: WorkbookMetadata,
    *,
    source_name: str,
    allow_large_workbook: bool,
) -> WorkbookWorkload:
    workload = WorkbookWorkload(
        cell_count=sum(sheet.cell_count for sheet in metadata.sheets),
        worksheet_xml_bytes=sum(sheet.xml_bytes for sheet in metadata.sheets),
        shared_string_bytes=metadata.shared_string_bytes,
        styles_bytes=metadata.styles_bytes,
        style_count=metadata.style_count,
        largest_sheet_area=max(
            (sheet.max_row * sheet.max_column for sheet in metadata.sheets),
            default=0,
        ),
    )
    warnings: list[str] = []
    refusals: list[str] = []
    for field_name, warning_limit, refusal_limit, label in _WORKLOAD_LIMITS:
        value = int(getattr(workload, field_name))
        if value >= refusal_limit:
            refusals.append(f"{label} {value:,} >= refusal limit {refusal_limit:,}")
        elif value >= warning_limit:
            warnings.append(f"{label} {value:,} >= warning limit {warning_limit:,}")
    if refusals and not allow_large_workbook:
        raise OOXMLWorkloadError(
            f"{source_name}: workbook workload refused: {'; '.join(refusals)}. "
            "Review the workbook and rerun with the explicit local "
            "allow_large_workbook override only when sufficient memory is available."
        )
    if refusals:
        warnings.extend(f"override accepted: {reason}" for reason in refusals)
        workload.override_used = True
    workload.warning_reasons = tuple(warnings)
    return workload


def _style_key(cell: Any) -> str:
    fill = cell.fill
    font = cell.font
    border = cell.border
    parts = (
        getattr(fill, "patternType", None),
        getattr(getattr(fill, "fgColor", None), "rgb", None),
        font.b,
        font.i,
        font.sz,
        getattr(font.color, "rgb", None) if font.color is not None else None,
        border.left.style,
        border.right.style,
        border.top.style,
        border.bottom.style,
    )
    return "|".join(str(p) for p in parts)


def _openpyxl_hidden_columns(worksheet: Any) -> frozenset[int]:
    hidden: set[int] = set()
    for letter, dimension in worksheet.column_dimensions.items():
        if not dimension.hidden:
            continue
        fallback = column_index_from_string(letter)
        minimum = int(dimension.min or fallback)
        maximum = int(dimension.max or fallback)
        hidden.update(range(minimum, maximum + 1))
    return frozenset(hidden)


def _load_ooxml_oracle(
    data: bytes,
    *,
    source_name: str,
    file_format: str,
    allow_large_workbook: bool = False,
    cancellation_token: CancellationToken | None = None,
) -> WorkbookSnapshot:
    check_cancelled(cancellation_token)
    metadata = parse_ooxml_worksheet_metadata(
        data,
        cancellation_token=cancellation_token,
    )
    workload = _assess_ooxml_workload(
        metadata,
        source_name=source_name,
        allow_large_workbook=allow_large_workbook,
    )
    workbook_data = _without_chart_drawings(data)
    wb_formulas = load_workbook(io.BytesIO(workbook_data), data_only=False)
    check_cancelled(cancellation_token)
    wb_values = load_workbook(io.BytesIO(workbook_data), data_only=True)
    check_cancelled(cancellation_token)

    snapshot = WorkbookSnapshot(
        source_name=source_name,
        file_format=file_format,
        formulas_available=True,
        styles_available=True,
        formula_presence_available=True,
        tables_available=True,
        charts_available=True,
        interaction_rules_available=True,
        interaction_rules_supported=True,
        conditional_format_styles_supported=True,
        formula_source="openpyxl",
        formula_detail="Formula text read directly from OOXML",
        calculation_mode=wb_formulas.calculation.calcMode,
        full_calc_on_load=wb_formulas.calculation.fullCalcOnLoad,
        external_links=_parse_external_links(data),
        workload=workload,
    )
    snapshot.formula_ranges.extend(
        descriptor
        for sheet_metadata in metadata.sheets
        for descriptor in sheet_metadata.formula_ranges
    )
    interaction_details: list[str] = []
    conditional_style_details: list[str] = []

    for name, defined in wb_formulas.defined_names.items():
        if name.startswith("_xlnm"):
            continue
        snapshot.named_ranges.append(NamedRange(name=name, target=str(defined.attr_text)))

    for sheet_name in wb_formulas.sheetnames:
        check_cancelled(cancellation_token)
        ws_f = wb_formulas[sheet_name]
        ws_v = wb_values[sheet_name]

        cells: dict[tuple[int, int], CellRecord] = {}
        for row in ws_f.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                key = (cell.row, cell.column)
                formula: str | None = None
                value: CellValue
                if cell.data_type == "f":
                    formula = _formula_text(cell.value)
                    cached = ws_v.cell(row=cell.row, column=cell.column).value
                    value = cached if is_cell_value(cached) else None
                else:
                    value = _constant_cell_value(cell.value)
                cells[key] = CellRecord(
                    row=cell.row,
                    column=cell.column,
                    value=value,
                    formula=formula,
                    is_formula=formula is not None,
                    number_format=str(cell.number_format),
                    style_key=_style_key(cell),
                )

        hidden_rows = frozenset(index for index, dim in ws_f.row_dimensions.items() if dim.hidden)
        hidden_columns = _openpyxl_hidden_columns(ws_f)
        snapshot.sheets.append(
            SheetSnapshot(
                name=sheet_name,
                visibility=str(ws_f.sheet_state),
                max_row=int(ws_f.max_row or 0),
                max_column=int(ws_f.max_column or 0),
                cells=cells,
                hidden_rows=hidden_rows,
                hidden_columns=hidden_columns,
            )
        )

        for raw_table in ws_f.tables.values():
            table = raw_table
            snapshot.tables.append(
                TableDescriptor(
                    sheet=sheet_name,
                    name=str(table.name),
                    display_name=str(table.displayName),
                    cell_range=str(table.ref),
                    columns=[str(column.name) for column in table.tableColumns],
                    header_row_count=(
                        int(table.headerRowCount) if table.headerRowCount is not None else 1
                    ),
                    totals_row_count=(
                        int(table.totalsRowCount)
                        if table.totalsRowCount is not None
                        else int(bool(table.totalsRowShown))
                    ),
                    source_id=int(table.id) if table.id is not None else None,
                )
            )

        try:
            interaction = extract_worksheet_interactions(ws_f, sheet_name)
        except (AttributeError, TypeError, ValueError) as exc:
            snapshot.interaction_rules_available = False
            snapshot.interaction_rules_supported = False
            snapshot.conditional_format_styles_supported = False
            interaction_details.append(f"{sheet_name}: interaction extraction failed: {exc}")
            conditional_style_details.append(
                f"{sheet_name}: conditional-format style extraction failed: {exc}"
            )
        else:
            snapshot.data_validations.extend(interaction.data_validations)
            snapshot.conditional_formats.extend(interaction.conditional_formats)
            if not interaction.rules_supported:
                snapshot.interaction_rules_supported = False
                interaction_details.extend(
                    f"{sheet_name}: {detail}" for detail in interaction.rule_details
                )
            if not interaction.styles_supported:
                snapshot.conditional_format_styles_supported = False
                conditional_style_details.extend(
                    f"{sheet_name}: {detail}" for detail in interaction.style_details
                )

    try:
        snapshot.charts = parse_ooxml_charts(
            data,
            cancellation_token=cancellation_token,
        )
        snapshot.chart_detail = "Charts parsed from raw OOXML package parts"
    except ChartParseError as exc:
        logger.warning("%s: complete chart extraction failed: %s", source_name, exc)
        snapshot.charts_available = False
        snapshot.chart_detail = f"Complete chart extraction unavailable: {exc}"
    snapshot.interaction_rule_detail = "; ".join(sorted(set(interaction_details)))
    snapshot.conditional_format_style_detail = "; ".join(sorted(set(conditional_style_details)))
    snapshot.pivots.extend(_parse_pivots(data))
    return snapshot


def _physical_cells(worksheet: Any) -> Iterator[ReadOnlyCell]:
    for row in worksheet.iter_rows():
        for cell in row:
            if isinstance(cell, ReadOnlyCell):
                yield cell


def _cell_coordinate(cell: ReadOnlyCell) -> tuple[int, int]:
    return int(cell.row), int(cell.column)


def _stream_formula_cells(
    formula_worksheet: Any,
    *,
    sheet_name: str,
    style_cache: dict[int, tuple[str, str]],
    cancellation_token: CancellationToken | None = None,
) -> dict[tuple[int, int], CellRecord]:
    cells: dict[tuple[int, int], CellRecord] = {}
    for index, formula_cell in enumerate(_physical_cells(formula_worksheet), start=1):
        if index % 10_000 == 0:
            check_cancelled(cancellation_token)
        if formula_cell.value is None:
            continue
        style_id = getattr(formula_cell, "_style_id", None)
        if not isinstance(style_id, int):
            raise OOXMLCellStreamError(
                f"{sheet_name}!{formula_cell.coordinate}: read-only cell "
                "does not expose a workbook-local style identifier"
            )
        style = style_cache.get(style_id)
        if style is None:
            style = (_style_key(formula_cell), str(formula_cell.number_format))
            style_cache[style_id] = style
        formula = (
            _formula_text(formula_cell.value)
            if formula_cell.data_type == "f"
            else None
        )
        value = (
            None
            if formula is not None
            else _constant_cell_value(formula_cell.value)
        )
        row, column = _cell_coordinate(formula_cell)
        cells[(row, column)] = CellRecord(
            row=row,
            column=column,
            value=value,
            formula=formula,
            is_formula=formula is not None,
            number_format=style[1],
            style_key=style[0],
        )
    return cells


def _numeric_value(value: str) -> int | float:
    return float(value) if "." in value or "E" in value.upper() else int(value)


def _cached_formula_value(
    raw_value: str | None,
    data_type: str,
    number_format: str | None,
    epoch: Any,
    shared_strings: list[str],
) -> CellValue:
    if raw_value is None:
        return None
    if data_type == "n":
        numeric = _numeric_value(raw_value)
        if number_format and is_date_format(number_format):
            converted = from_excel(
                numeric,
                epoch,
                timedelta=is_timedelta_format(number_format),
            )
            return converted if is_cell_value(converted) else None
        return numeric
    if data_type == "b":
        return bool(int(raw_value))
    if data_type in {"e", "str"}:
        return raw_value
    if data_type == "s":
        try:
            return shared_strings[int(raw_value)]
        except (IndexError, ValueError) as exc:
            raise OOXMLCellStreamError(
                f"invalid shared-string cached formula result {raw_value!r}"
            ) from exc
    if data_type == "d":
        converted = from_ISO8601(raw_value)
        return converted if is_cell_value(converted) else None
    raise OOXMLCellStreamError(
        f"unsupported cached formula result type {data_type!r}"
    )


def _shared_strings(
    archive: zipfile.ZipFile,
    cancellation_token: CancellationToken | None,
) -> list[str]:
    part = "xl/sharedStrings.xml"
    if part not in archive.namelist():
        return []
    values: list[str] = []
    retained_depth = 0
    try:
        with archive.open(part) as stream:
            for event, element in ElementTree.iterparse(
                stream,
                events=("start", "end"),
            ):
                local_name = _local_name(element.tag)
                if event == "start":
                    if retained_depth:
                        retained_depth += 1
                    elif local_name == "si":
                        retained_depth = 1
                    continue
                if retained_depth:
                    retained_depth -= 1
                    if retained_depth:
                        continue
                    values.append(
                        "".join(
                            child.text or ""
                            for child in element.iter()
                            if _local_name(child.tag) == "t"
                        )
                    )
                    if len(values) % 10_000 == 0:
                        check_cancelled(cancellation_token)
                    element.clear()
                    continue
                element.clear()
    except ElementTree.ParseError as exc:
        raise OOXMLCellStreamError(
            f"malformed shared-string cached-value stream: {exc}"
        ) from exc
    return values


def _merge_cached_formula_values(
    data: bytes,
    metadata: list[WorksheetMetadata],
    snapshot: WorkbookSnapshot,
    *,
    epoch: Any,
    cancellation_token: CancellationToken | None = None,
) -> None:
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        shared_strings = _shared_strings(archive, cancellation_token)
        for sheet_metadata, sheet in zip(metadata, snapshot.sheets, strict=True):
            check_cancelled(cancellation_token)
            expected = {
                coordinate
                for coordinate, record in sheet.cells.items()
                if record.formula is not None
            }
            seen: set[tuple[int, int]] = set()
            retained_depth = 0
            try:
                with archive.open(sheet_metadata.part) as stream:
                    for event, element in ElementTree.iterparse(
                        stream,
                        events=("start", "end"),
                    ):
                        local_name = _local_name(element.tag)
                        if event == "start":
                            if retained_depth:
                                retained_depth += 1
                            elif local_name == "c":
                                retained_depth = 1
                            continue
                        if retained_depth:
                            retained_depth -= 1
                            if retained_depth:
                                continue
                            reference = element.get("r")
                            formula = next(
                                (
                                    child
                                    for child in element
                                    if _local_name(child.tag) == "f"
                                ),
                                None,
                            )
                            if formula is not None and reference:
                                row, column = coordinate_to_tuple(reference)
                                coordinate = (row, column)
                                record = sheet.cells.get(coordinate)
                                if record is None or record.formula is None:
                                    raise OOXMLCellStreamError(
                                        f"{sheet.name}!{reference}: raw formula cell "
                                        "is absent from the formula stream"
                                    )
                                raw_value = next(
                                    (
                                        child.text
                                        for child in element
                                        if _local_name(child.tag) == "v"
                                    ),
                                    None,
                                )
                                record.value = _cached_formula_value(
                                    raw_value,
                                    element.get("t", "n"),
                                    record.number_format,
                                    epoch,
                                    shared_strings,
                                )
                                seen.add(coordinate)
                                if len(seen) % 10_000 == 0:
                                    check_cancelled(cancellation_token)
                            element.clear()
                            continue
                        element.clear()
            except ElementTree.ParseError as exc:
                raise OOXMLCellStreamError(
                    f"{sheet.name}: malformed worksheet cached-value stream: {exc}"
                ) from exc
            missing = expected.difference(seen)
            if missing:
                row, column = min(missing)
                raise OOXMLCellStreamError(
                    f"{sheet.name}!R{row}C{column}: formula stream cell is absent "
                    "from raw worksheet XML"
                )


def _append_streaming_metadata(
    snapshot: WorkbookSnapshot,
    metadata: WorksheetMetadata,
) -> None:
    snapshot.tables.extend(metadata.tables)
    snapshot.formula_ranges.extend(metadata.formula_ranges)
    snapshot.data_validations.extend(metadata.interactions.data_validations)
    snapshot.conditional_formats.extend(metadata.interactions.conditional_formats)
    if not metadata.interactions.rules_supported:
        snapshot.interaction_rules_supported = False
    if not metadata.interactions.styles_supported:
        snapshot.conditional_format_styles_supported = False


def _load_ooxml_streaming(
    data: bytes,
    *,
    source_name: str,
    file_format: str,
    allow_large_workbook: bool = False,
    cancellation_token: CancellationToken | None = None,
) -> WorkbookSnapshot:
    check_cancelled(cancellation_token)
    metadata = parse_ooxml_worksheet_metadata(
        data,
        cancellation_token=cancellation_token,
    )
    workload = _assess_ooxml_workload(
        metadata,
        source_name=source_name,
        allow_large_workbook=allow_large_workbook,
    )
    wb_formulas = load_workbook(
        io.BytesIO(data),
        data_only=False,
        read_only=True,
        keep_links=False,
    )
    try:
        check_cancelled(cancellation_token)
        expected_sheet_names = [sheet.name for sheet in metadata.sheets]
        if wb_formulas.sheetnames != expected_sheet_names:
            raise OOXMLMetadataError(
                "formula workbook sheet inventory does not match raw package metadata"
            )
        snapshot = WorkbookSnapshot(
            source_name=source_name,
            file_format=file_format,
            formulas_available=True,
            styles_available=True,
            formula_presence_available=True,
            tables_available=True,
            charts_available=True,
            interaction_rules_available=True,
            interaction_rules_supported=True,
            conditional_format_styles_supported=True,
            formula_source="openpyxl",
            formula_detail="Formula text read directly from OOXML",
            calculation_mode=wb_formulas.calculation.calcMode,
            full_calc_on_load=wb_formulas.calculation.fullCalcOnLoad,
            external_links=_parse_external_links(data),
            workload=workload,
        )
        epoch = wb_formulas.epoch
        for name, defined in wb_formulas.defined_names.items():
            if name.startswith("_xlnm"):
                continue
            snapshot.named_ranges.append(NamedRange(name=name, target=str(defined.attr_text)))

        interaction_details: list[str] = []
        conditional_style_details: list[str] = []
        style_cache: dict[int, tuple[str, str]] = {}
        for sheet_metadata in metadata.sheets:
            check_cancelled(cancellation_token)
            formula_worksheet = wb_formulas[sheet_metadata.name]
            formula_worksheet.reset_dimensions()
            cells = _stream_formula_cells(
                formula_worksheet,
                sheet_name=sheet_metadata.name,
                style_cache=style_cache,
                cancellation_token=cancellation_token,
            )
            snapshot.sheets.append(
                SheetSnapshot(
                    name=sheet_metadata.name,
                    visibility=sheet_metadata.visibility,
                    max_row=sheet_metadata.max_row,
                    max_column=sheet_metadata.max_column,
                    cells=cells,
                    hidden_rows=sheet_metadata.hidden_rows,
                    hidden_columns=sheet_metadata.hidden_columns,
                )
            )
            _append_streaming_metadata(snapshot, sheet_metadata)
            interaction_details.extend(
                f"{sheet_metadata.name}: {detail}"
                for detail in sheet_metadata.interactions.rule_details
            )
            conditional_style_details.extend(
                f"{sheet_metadata.name}: {detail}"
                for detail in sheet_metadata.interactions.style_details
            )
    finally:
        wb_formulas.close()

    _merge_cached_formula_values(
        data,
        metadata.sheets,
        snapshot,
        epoch=epoch,
        cancellation_token=cancellation_token,
    )

    try:
        check_cancelled(cancellation_token)
        snapshot.charts = parse_ooxml_charts(
            data,
            cancellation_token=cancellation_token,
        )
        snapshot.chart_detail = "Charts parsed from raw OOXML package parts"
    except ChartParseError as exc:
        logger.warning("%s: complete chart extraction failed: %s", source_name, exc)
        snapshot.charts_available = False
        snapshot.chart_detail = f"Complete chart extraction unavailable: {exc}"
    snapshot.interaction_rule_detail = "; ".join(sorted(set(interaction_details)))
    snapshot.conditional_format_style_detail = "; ".join(
        sorted(set(conditional_style_details))
    )
    snapshot.pivots.extend(_parse_pivots(data))
    return snapshot


def _local_name(tag: str) -> str:
    return tag.rsplit("}", 1)[-1]


def _relationship_targets(
    zf: zipfile.ZipFile, relationship_part: str, source_part: str
) -> dict[str, str]:
    if relationship_part not in zf.namelist():
        return {}
    root = ElementTree.fromstring(zf.read(relationship_part))
    targets: dict[str, str] = {}
    for relationship in root:
        if _local_name(relationship.tag) != "Relationship":
            continue
        relationship_id = relationship.get("Id")
        target = relationship.get("Target")
        if not relationship_id or not target or relationship.get("TargetMode") == "External":
            continue
        if target.startswith("/"):
            resolved = target.lstrip("/")
        else:
            resolved = posixpath.normpath(posixpath.join(posixpath.dirname(source_part), target))
        targets[relationship_id] = resolved
    return targets


def _parse_external_links(data: bytes) -> list[str]:
    """Return external workbook targets retained in the OOXML package."""
    targets: set[str] = set()
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        for part in zf.namelist():
            if not (part.startswith("xl/externalLinks/_rels/") and part.endswith(".rels")):
                continue
            root = ElementTree.fromstring(zf.read(part))
            for relationship in root:
                target = relationship.get("Target")
                if relationship.get("TargetMode") == "External" and target:
                    targets.add(target)
    return sorted(targets)


def _parse_pivots(data: bytes) -> list[PivotDescriptor]:
    """Parse pivot descriptors from raw parts (loader contract from fixtures)."""
    tables: list[tuple[str, str, str | None, str | None]] = []
    caches: dict[str, tuple[str | None, str | None]] = {}
    cache_parts_by_id: dict[str, str] = {}
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        for part in sorted(zf.namelist()):
            if _PIVOT_TABLE_RE.match(part):
                root = ElementTree.fromstring(zf.read(part))
                location_ref = None
                for element in root.iter():
                    if _local_name(element.tag) == "location":
                        location_ref = element.get("ref")
                        break
                tables.append((part, root.get("name") or part, location_ref, root.get("cacheId")))
            elif _PIVOT_CACHE_RE.match(part):
                root = ElementTree.fromstring(zf.read(part))
                source_sheet = source_ref = None
                for element in root.iter():
                    if _local_name(element.tag) == "worksheetSource":
                        source_sheet = element.get("sheet")
                        source_ref = element.get("ref")
                        break
                caches[part] = (source_sheet, source_ref)

        if "xl/workbook.xml" in zf.namelist():
            workbook_relationships = _relationship_targets(
                zf, "xl/_rels/workbook.xml.rels", "xl/workbook.xml"
            )
            workbook = ElementTree.fromstring(zf.read("xl/workbook.xml"))
            for element in workbook.iter():
                if _local_name(element.tag) != "pivotCache":
                    continue
                cache_id = element.get("cacheId")
                relationship_id = next(
                    (
                        value
                        for attribute, value in element.attrib.items()
                        if _local_name(attribute) == "id"
                    ),
                    None,
                )
                if cache_id and relationship_id in workbook_relationships:
                    cache_parts_by_id[cache_id] = workbook_relationships[relationship_id]

        pivots = []
        for table_part, name, location_ref, cache_id in tables:
            cache_part = cache_parts_by_id.get(cache_id or "")
            if cache_part not in caches:
                relationship_part = (
                    f"{posixpath.dirname(table_part)}/_rels/{posixpath.basename(table_part)}.rels"
                )
                related_parts = _relationship_targets(zf, relationship_part, table_part)
                cache_part = next((part for part in related_parts.values() if part in caches), None)
            if cache_part not in caches and len(caches) == 1:
                cache_part = next(iter(caches))
            source_sheet, source_ref = caches.get(cache_part or "", (None, None))
            pivots.append(
                PivotDescriptor(
                    name=name,
                    location_ref=location_ref,
                    source_sheet=source_sheet,
                    source_ref=source_ref,
                )
            )
    return pivots


# --- xlsb ----------------------------------------------------------------


def _xlsb_value(raw: object) -> CellValue:
    if isinstance(raw, str):
        return _XLSB_ERRORS.get(raw.lower(), raw)
    if isinstance(raw, bool | int | float):
        return raw
    if raw is None:
        return None
    return str(raw)


def _scan_xlsb(data: bytes, source_name: str) -> tuple[XlsbFormulaScan | None, str]:
    try:
        return scan_xlsb_formulas(data), ""
    except XlsbFormulaScanError as exc:
        logger.warning("%s: XLSB formula-presence scan failed: %s", source_name, exc)
        return None, f"Formula-presence scan failed: {exc}"


def _extract_xlsb_formulas(data: bytes, formula_scan: XlsbFormulaScan) -> FormulaExtraction:
    if sys.platform == "win32":
        from qc_tool.io.excel_formula import extract_formulas_with_excel

        return extract_formulas_with_excel(data, formula_scan)
    if sys.platform.startswith("linux"):
        from qc_tool.io.libreoffice_formula import extract_formulas_with_libreoffice

        return extract_formulas_with_libreoffice(data, formula_scan)
    raise FormulaEnrichmentError(f"no XLSB formula adapter is configured for {sys.platform}")


def _load_xlsb(
    data: bytes,
    *,
    source_name: str,
    cancellation_token: CancellationToken | None = None,
) -> WorkbookSnapshot:
    check_cancelled(cancellation_token)
    formula_scan, scan_detail = _scan_xlsb(data, source_name)
    snapshot = WorkbookSnapshot(
        source_name=source_name,
        file_format="xlsb",
        formulas_available=False,
        styles_available=False,
        formula_presence_available=formula_scan is not None,
        tables_available=False,
        charts_available=False,
        interaction_rules_available=False,
        interaction_rules_supported=False,
        conditional_format_styles_supported=False,
        formula_detail=scan_detail or "Formula records identified; formula text unavailable",
        chart_detail="XLSB chart metadata is unavailable",
        interaction_rule_detail="XLSB interaction-rule metadata is unavailable",
        conditional_format_style_detail=("XLSB conditional-format style metadata is unavailable"),
    )
    with open_xlsb(io.BytesIO(data)) as wb:
        for sheet_name in wb.sheets:
            check_cancelled(cancellation_token)
            cells: dict[tuple[int, int], CellRecord] = {}
            max_row = max_column = 0
            formula_cells = (
                formula_scan.formula_cells.get(sheet_name, frozenset())
                if formula_scan is not None
                else frozenset()
            )
            with wb.get_sheet(sheet_name) as sheet:
                for row in sheet.rows(sparse=True):
                    for cell in row:
                        if cell.v is None:
                            continue
                        row_1 = cell.r + 1
                        col_1 = cell.c + 1
                        max_row = max(max_row, row_1)
                        max_column = max(max_column, col_1)
                        cells[(row_1, col_1)] = CellRecord(
                            row=row_1,
                            column=col_1,
                            value=_xlsb_value(cell.v),
                            is_formula=(row_1, col_1) in formula_cells,
                        )
            for row_1, col_1 in formula_cells.difference(cells):
                cells[(row_1, col_1)] = CellRecord(
                    row=row_1,
                    column=col_1,
                    value=None,
                    is_formula=True,
                )
                max_row = max(max_row, row_1)
                max_column = max(max_column, col_1)
            snapshot.sheets.append(
                SheetSnapshot(
                    name=sheet_name,
                    visibility="visible",  # pyxlsb does not expose sheet state
                    max_row=max_row,
                    max_column=max_column,
                    cells=cells,
                )
            )
    if formula_scan is not None and formula_scan.formula_count:
        check_cancelled(cancellation_token)
        try:
            extraction = _extract_xlsb_formulas(data, formula_scan)
            merge_formula_extraction(snapshot, formula_scan, extraction)
        except FormulaEnrichmentError as exc:
            logger.warning("%s: XLSB formula enrichment unavailable: %s", source_name, exc)
            snapshot.formula_detail = f"Formula presence checked; text unavailable: {exc}"
    return snapshot
