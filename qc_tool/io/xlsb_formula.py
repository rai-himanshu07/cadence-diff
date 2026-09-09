"""Bounds-checked structural scanning for XLSB formula enrichment.

This module deliberately does not decode BIFF12 formula tokens. It identifies
formula-cell records and package features that make opening a workbook in an
external Office process unsafe. Formula text from Excel or LibreOffice is
accepted only when its coordinates exactly match this independent scan.
"""

from __future__ import annotations

import io
import posixpath
import zipfile
from collections.abc import Iterator
from dataclasses import dataclass, field
from xml.etree import ElementTree

_ROW_RECORD = 0x0000
_RETAINED_CELL_RECORDS = frozenset(
    {
        0x0002,  # BrtCellRk
        0x0003,  # BrtCellError
        0x0004,  # BrtCellBool
        0x0005,  # BrtCellReal
        0x0007,  # BrtCellIsst
        0x0008,  # BrtFmlaString
        0x0009,  # BrtFmlaNum
        0x000A,  # BrtFmlaBool
        0x000B,  # BrtFmlaError
    }
)
_FORMULA_RECORDS = frozenset({0x0008, 0x0009, 0x000A, 0x000B})
_DIMENSION_RECORD = 0x0194
_SHEET_RECORD = 0x019C
#: Record IDs verified empirically against an independent, spec-correct reader
#: (python-calamine, MIT-licensed) in an isolated throwaway environment, not
#: against their officially-documented spec numbers directly: pyxlsb's own
#: `BIFF12Reader.read_id` (which this module's `_read_record_id` mirrors) does
#: NOT mask off each byte's continuation bit before shifting, so the value
#: that must be matched is the naive full-byte-shift decode of the wire bytes,
#: not the record's spec-documented number. `BrtWbProp` (spec 0x0099) naive-
#: decodes to 0x0199; `BrtBeginFmts` (spec 0x0267) to 0x04E7; `BrtBeginCellXfs`
#: (spec 0x0269) to 0x04E9 -- the last exactly matching pyxlsb's own untested
#: `CELLXFS` constant, confirming pyxlsb's whole styles-adjacent constant
#: table already uses this same naive convention. pyxlsb never reads
#: `xl/styles.bin` or the workbook date-system flag at all (no registered
#: handler), so this module owns them independently, the same way it already
#: owns formula coordinate scanning. `BrtFmt`/`BrtXF` need no correction: both
#: are single-byte IDs (< 0x80), which both decode schemes agree on.
_WORKBOOK_PROP_RECORD = 0x0199  # BrtWbProp; bit 0 of byte 0 is fDate1904
_FMTS_BEGIN_RECORD = 0x04E7  # BrtBeginFmts; payload = u32 count of BrtFmt
_FMT_RECORD = 0x002C  # BrtFmt: u16 ifmt + XLWideString stFmtCode
_CELLXFS_BEGIN_RECORD = 0x04E9  # BrtBeginCellXfs; payload = u32 count of BrtXF
_XF_RECORD = 0x002F  # BrtXF: u16 ixfParent + u16 iFmt (only iFmt is read)
_MAX_RECORD_BYTES = 64 * 1024 * 1024
_MAX_STRING_CHARS = 1_000_000
_MAX_ROW = 1_048_575
_MAX_COLUMN = 16_383
#: Bounds a malformed styles.bin count from causing an unbounded allocation.
_MAX_STYLE_RECORDS = 1_000_000


_RISKY_PARTS = {
    "xl/vbaproject.bin": "VBA project",
    "xl/connections.bin": "external data connections",
}
_RISKY_PATH_SEGMENTS = {
    "/activex/": "ActiveX controls",
    "/ctrlprops/": "control properties",
    "/dialogsheets/": "dialog sheets",
    "/embeddings/": "embedded OLE content",
    "/externalconnections/": "external data connections",
    "/externallinks/": "external workbook links",
    "/macrosheets/": "Excel 4.0 macro sheets",
    "/querytables/": "external query tables",
    "/customui/": "custom Office UI",
}
_RISKY_RELATIONSHIP_KINDS = {
    "activexcontrol": "ActiveX controls",
    "attachedtoolbars": "custom Office toolbars",
    "connections": "external data connections",
    "ctrlprop": "control properties",
    "customui": "custom Office UI",
    "dialogsheet": "dialog sheets",
    "externallink": "external workbook links",
    "externallinklongpath": "external workbook links",
    "externallinkpath": "external workbook links",
    "macrosheet": "Excel 4.0 macro sheets",
    "oleobject": "embedded OLE content",
    "querytable": "external query tables",
    "vbaproject": "VBA project",
}
_PASSIVE_EXTERNAL_LINK_KINDS = frozenset(
    {"externallink", "externallinklongpath", "externallinkpath"}
)
#: Relationship kinds allowed to target a package part outside ``xl/``.
#: These are metadata parts opened for risk inspection only; they are never
#: treated as worksheets. Every other non-``xl/`` target fails closed.
_NON_XL_RELATIONSHIP_KINDS = {"customxml": "customXml/"}


class XlsbFormulaScanError(ValueError):
    """The XLSB package cannot be scanned without guessing."""


@dataclass(frozen=True, slots=True)
class XlsbWorksheetMetrics:
    """Bounded structural workload facts for one binary worksheet."""

    cell_count: int = 0
    max_row: int = 0
    max_column: int = 0
    binary_bytes: int = 0


@dataclass(frozen=True, slots=True)
class XlsbFormulaScan:
    """Formula coordinates and risky package features found in an XLSB."""

    formula_cells: dict[str, frozenset[tuple[int, int]]]
    risky_features: tuple[str, ...] = ()
    #: Recognized external-workbook-link metadata; safe for the isolated
    #: adapters. Never implies the link is inactive -- that requires complete
    #: formula/defined-name reachability evidence (Step 4b).
    passive_features: tuple[str, ...] = ()
    #: Active or unrecognized content; always refuses the adapter.
    blocking_features: tuple[str, ...] = ()
    #: A non-hyperlink external relationship whose kind is not the recognized
    #: `externalLinkPath`; conservatively refuses until proven otherwise.
    unknown_external_features: tuple[str, ...] = ()
    worksheet_metrics: dict[str, XlsbWorksheetMetrics] = field(default_factory=dict)
    package_worksheet_binary_bytes: int = 0
    shared_string_bytes: int = 0
    styles_bytes: int = 0
    sheet_count: int = 0

    @property
    def formula_count(self) -> int:
        return sum(len(cells) for cells in self.formula_cells.values())

    @property
    def cell_count(self) -> int:
        return sum(metrics.cell_count for metrics in self.worksheet_metrics.values())

    @property
    def worksheet_binary_bytes(self) -> int:
        return self.package_worksheet_binary_bytes or sum(
            metrics.binary_bytes for metrics in self.worksheet_metrics.values()
        )

    @property
    def largest_sheet_area(self) -> int:
        return max(
            (
                metrics.max_row * metrics.max_column
                for metrics in self.worksheet_metrics.values()
            ),
            default=0,
        )

    @property
    def largest_sheet_dimensions(self) -> tuple[int, int]:
        if not self.worksheet_metrics:
            return 0, 0
        metrics = max(
            self.worksheet_metrics.values(),
            key=lambda item: item.max_row * item.max_column,
        )
        return metrics.max_row, metrics.max_column

    @property
    def workload_metrics_available(self) -> bool:
        return self.sheet_count > 0 and len(self.worksheet_metrics) == self.sheet_count

    @property
    def safe_for_external_engine(self) -> bool:
        """Whether the isolated adapter may be invoked at all.

        Passive-only packages (only recognized external-workbook-link
        metadata) are safe to reach the adapter; this is never described as
        risk-free, and any blocking or unknown-external feature refuses.
        """
        classified = {
            *self.passive_features,
            *self.blocking_features,
            *self.unknown_external_features,
        }
        unclassified_legacy_features = set(self.risky_features) - classified
        return (
            not self.blocking_features
            and not self.unknown_external_features
            and not unclassified_legacy_features
        )


@dataclass(frozen=True, slots=True)
class XlsbStyleTable:
    """Bounded number-format facts from one XLSB's ``xl/styles.bin``.

    Deliberately narrower than full style coverage: only the custom-format
    table and the direct cell-XF format-id array are kept, never fonts,
    fills, borders, or the cell-style (named-style) XFs. ``cell_xf_format_ids``
    is positional: a cell's XF index (``iStyleRef``) is its position in this
    tuple, matching how BrtCell records reference ``BrtBeginCellXfs`` entries.
    """

    custom_formats: dict[int, str] = field(default_factory=dict)
    cell_xf_format_ids: tuple[int, ...] = ()

    def number_format(self, xf_index: int) -> str | None:
        """The format code for one cell-XF index, or ``None`` if unresolvable."""
        if xf_index < 0 or xf_index >= len(self.cell_xf_format_ids):
            return None
        ifmt = self.cell_xf_format_ids[xf_index]
        if ifmt in self.custom_formats:
            return self.custom_formats[ifmt]
        from openpyxl.styles.numbers import builtin_format_code

        return builtin_format_code(ifmt)


@dataclass(frozen=True, slots=True)
class _Relationship:
    target: str
    kind: str


def _local_name(tag: str) -> str:
    return tag.rsplit("}", 1)[-1]


def _read_record_id(data: bytes, position: int) -> tuple[int, int]:
    value = 0
    for index in range(4):
        if position >= len(data):
            raise XlsbFormulaScanError("truncated BIFF12 record id")
        byte = data[position]
        position += 1
        value += byte << (8 * index)
        if byte & 0x80 == 0:
            return value, position
    raise XlsbFormulaScanError("BIFF12 record id exceeds four bytes")


def _read_record_length(data: bytes, position: int) -> tuple[int, int]:
    value = 0
    for index in range(4):
        if position >= len(data):
            raise XlsbFormulaScanError("truncated BIFF12 record length")
        byte = data[position]
        position += 1
        value += (byte & 0x7F) << (7 * index)
        if byte & 0x80 == 0:
            if value > _MAX_RECORD_BYTES:
                raise XlsbFormulaScanError(
                    f"BIFF12 record length {value} exceeds the safety limit"
                )
            return value, position
    raise XlsbFormulaScanError("BIFF12 record length exceeds four bytes")


def _records(data: bytes) -> Iterator[tuple[int, bytes]]:
    position = 0
    while position < len(data):
        record_id, position = _read_record_id(data, position)
        length, position = _read_record_length(data, position)
        end = position + length
        if end > len(data):
            raise XlsbFormulaScanError(
                f"BIFF12 record {record_id:#x} is truncated: needs {length} bytes"
            )
        yield record_id, data[position:end]
        position = end


def _read_wide_string(payload: bytes, position: int) -> tuple[str, int]:
    if position + 4 > len(payload):
        raise XlsbFormulaScanError("truncated BIFF12 string length")
    length = int.from_bytes(payload[position : position + 4], "little")
    if length > _MAX_STRING_CHARS:
        raise XlsbFormulaScanError(f"BIFF12 string has unsafe length {length}")
    position += 4
    end = position + length * 2
    if end > len(payload):
        raise XlsbFormulaScanError("truncated BIFF12 UTF-16 string")
    try:
        value = payload[position:end].decode("utf-16-le")
    except UnicodeDecodeError as exc:
        raise XlsbFormulaScanError("invalid BIFF12 UTF-16 string") from exc
    return value, end


def _resolve_target(source_part: str, target: str, *, kind: str = "") -> str:
    normalized_target = target.replace("\\", "/")
    if normalized_target.startswith("/"):
        resolved = posixpath.normpath(normalized_target.lstrip("/"))
    else:
        resolved = posixpath.normpath(
            posixpath.join(posixpath.dirname(source_part), normalized_target)
        )
    if resolved == ".." or resolved.startswith("../"):
        raise XlsbFormulaScanError(f"unsafe workbook relationship target {target!r}")
    if resolved.startswith("xl/"):
        return resolved
    prefix = _NON_XL_RELATIONSHIP_KINDS.get(kind)
    if prefix is not None and resolved.startswith(prefix):
        return resolved
    raise XlsbFormulaScanError(f"unsafe workbook relationship target {target!r}")


def _workbook_relationships(archive: zipfile.ZipFile) -> dict[str, _Relationship]:
    part = "xl/_rels/workbook.bin.rels"
    try:
        root = ElementTree.fromstring(archive.read(part))
    except KeyError as exc:
        raise XlsbFormulaScanError(f"missing required XLSB part {part}") from exc
    except ElementTree.ParseError as exc:
        raise XlsbFormulaScanError(f"invalid relationship XML in {part}") from exc

    relationships: dict[str, _Relationship] = {}
    for element in root:
        if _local_name(element.tag) != "Relationship":
            continue
        relationship_id = element.get("Id")
        target = element.get("Target")
        relationship_type = element.get("Type", "")
        if not relationship_id or not target:
            raise XlsbFormulaScanError("workbook relationship is missing Id or Target")
        if element.get("TargetMode") == "External":
            continue
        if relationship_id in relationships:
            raise XlsbFormulaScanError(
                f"duplicate workbook relationship id {relationship_id!r}"
            )
        kind = relationship_type.rsplit("/", 1)[-1].lower()
        relationships[relationship_id] = _Relationship(
            target=_resolve_target("xl/workbook.bin", target, kind=kind),
            kind=kind,
        )
    return relationships


def _workbook_sheets(data: bytes) -> list[tuple[str, str]]:
    sheets: list[tuple[str, str]] = []
    names: set[str] = set()
    for record_id, payload in _records(data):
        if record_id != _SHEET_RECORD:
            continue
        if len(payload) < 12:
            raise XlsbFormulaScanError("truncated BrtBundleSh record")
        relationship_id, position = _read_wide_string(payload, 8)
        name, _ = _read_wide_string(payload, position)
        if not name or not relationship_id:
            raise XlsbFormulaScanError("BrtBundleSh has an empty name or relationship id")
        if name in names:
            raise XlsbFormulaScanError(f"duplicate workbook sheet name {name!r}")
        names.add(name)
        sheets.append((name, relationship_id))
    if not sheets:
        raise XlsbFormulaScanError("workbook contains no BrtBundleSh records")
    return sheets


def _worksheet_scan(
    data: bytes, sheet_name: str
) -> tuple[frozenset[tuple[int, int]], XlsbWorksheetMetrics]:
    row: int | None = None
    cells: set[tuple[int, int]] = set()
    cell_count = 0
    max_row = 0
    max_column = 0
    for record_id, payload in _records(data):
        if record_id == _ROW_RECORD:
            if len(payload) < 4:
                raise XlsbFormulaScanError(f"{sheet_name}: truncated BrtRowHdr record")
            row = int.from_bytes(payload[:4], "little")
            if row > _MAX_ROW:
                raise XlsbFormulaScanError(f"{sheet_name}: row {row} is out of bounds")
        elif record_id == _DIMENSION_RECORD:
            if len(payload) < 16:
                raise XlsbFormulaScanError(
                    f"{sheet_name}: truncated BrtWsDim record"
                )
            first_row = int.from_bytes(payload[0:4], "little")
            last_row = int.from_bytes(payload[4:8], "little")
            first_column = int.from_bytes(payload[8:12], "little")
            last_column = int.from_bytes(payload[12:16], "little")
            if first_row > last_row or last_row > _MAX_ROW:
                raise XlsbFormulaScanError(
                    f"{sheet_name}: worksheet row dimensions are out of bounds"
                )
            if first_column > last_column or last_column > _MAX_COLUMN:
                raise XlsbFormulaScanError(
                    f"{sheet_name}: worksheet column dimensions are out of bounds"
                )
            max_row = max(max_row, last_row + 1)
            max_column = max(max_column, last_column + 1)
        elif record_id in _RETAINED_CELL_RECORDS:
            if row is None:
                raise XlsbFormulaScanError(
                    f"{sheet_name}: cell record appears before a row header"
                )
            if len(payload) < 4:
                raise XlsbFormulaScanError(f"{sheet_name}: truncated cell record")
            column = int.from_bytes(payload[:4], "little")
            if column > _MAX_COLUMN:
                raise XlsbFormulaScanError(
                    f"{sheet_name}: column {column} is out of bounds"
                )
            cell_count += 1
            max_row = max(max_row, row + 1)
            max_column = max(max_column, column + 1)
            if record_id not in _FORMULA_RECORDS:
                continue
            coordinate = (row + 1, column + 1)
            if coordinate in cells:
                raise XlsbFormulaScanError(
                    f"{sheet_name}: duplicate formula cell at {coordinate}"
                )
            cells.add(coordinate)
    return frozenset(cells), XlsbWorksheetMetrics(
        cell_count=cell_count,
        max_row=max_row,
        max_column=max_column,
        binary_bytes=len(data),
    )


def parse_xlsb_date_system(data: bytes) -> bool:
    """Return ``True`` for the 1904 epoch; ``False`` (1900) if unreadable."""
    for record_id, payload in _records(data):
        if record_id == _WORKBOOK_PROP_RECORD:
            return bool(payload) and (payload[0] & 0x1) != 0
    return False


def parse_xlsb_styles(data: bytes) -> XlsbStyleTable:
    """Parse ``xl/styles.bin``'s custom formats and direct cell-XF format ids.

    A flat sequential walk: entering ``BrtBeginFmts``/``BrtBeginCellXfs``
    reads a declared count, then that many immediately following records
    must be ``BrtFmt``/``BrtXF`` respectively (matches a verified working
    reader's own consumption pattern). Any other record — including the
    unrelated ``BrtBeginCellStyleXfs`` group's own ``BrtXF`` records — is
    simply skipped because it is never collected while no count is pending.
    """
    custom_formats: dict[int, str] = {}
    cell_xf_format_ids: list[int] = []
    pending_formats = 0
    pending_xfs = 0
    for record_id, payload in _records(data):
        if pending_formats > 0:
            if record_id != _FMT_RECORD:
                raise XlsbFormulaScanError(
                    "styles.bin: BrtBeginFmts count does not match its records"
                )
            if len(payload) < 6:
                raise XlsbFormulaScanError("styles.bin: truncated BrtFmt record")
            ifmt = int.from_bytes(payload[0:2], "little")
            code, _ = _read_wide_string(payload, 2)
            custom_formats[ifmt] = code
            pending_formats -= 1
            continue
        if pending_xfs > 0:
            if record_id != _XF_RECORD:
                raise XlsbFormulaScanError(
                    "styles.bin: BrtBeginCellXfs count does not match its records"
                )
            if len(payload) < 4:
                raise XlsbFormulaScanError("styles.bin: truncated BrtXF record")
            cell_xf_format_ids.append(int.from_bytes(payload[2:4], "little"))
            pending_xfs -= 1
            continue
        if record_id == _FMTS_BEGIN_RECORD:
            if len(payload) < 4:
                raise XlsbFormulaScanError("styles.bin: truncated BrtBeginFmts record")
            count = int.from_bytes(payload[0:4], "little")
            if count > _MAX_STYLE_RECORDS:
                raise XlsbFormulaScanError(
                    f"styles.bin: unsafe custom-format count {count}"
                )
            pending_formats = count
        elif record_id == _CELLXFS_BEGIN_RECORD:
            if len(payload) < 4:
                raise XlsbFormulaScanError(
                    "styles.bin: truncated BrtBeginCellXfs record"
                )
            count = int.from_bytes(payload[0:4], "little")
            if count > _MAX_STYLE_RECORDS:
                raise XlsbFormulaScanError(f"styles.bin: unsafe cell-XF count {count}")
            pending_xfs = count
    if pending_formats or pending_xfs:
        raise XlsbFormulaScanError("styles.bin: truncated before its declared record count")
    return XlsbStyleTable(
        custom_formats=custom_formats,
        cell_xf_format_ids=tuple(cell_xf_format_ids),
    )


def parse_worksheet_cell_styles(
    data: bytes, sheet_name: str
) -> dict[tuple[int, int], int]:
    """One worksheet's ``(row, column) -> cell-XF index`` (1-based, bounded).

    Reads only that one worksheet's bytes and returns only that one
    worksheet's map; the caller discards it before moving to the next sheet
    so no whole-workbook coordinate-to-style dictionary is ever materialized.
    """
    row: int | None = None
    styles: dict[tuple[int, int], int] = {}
    for record_id, payload in _records(data):
        if record_id == _ROW_RECORD:
            if len(payload) < 4:
                raise XlsbFormulaScanError(f"{sheet_name}: truncated BrtRowHdr record")
            row = int.from_bytes(payload[:4], "little")
        elif record_id in _RETAINED_CELL_RECORDS:
            if row is None:
                raise XlsbFormulaScanError(
                    f"{sheet_name}: cell record appears before a row header"
                )
            if len(payload) < 7:
                raise XlsbFormulaScanError(f"{sheet_name}: truncated cell record")
            column = int.from_bytes(payload[0:4], "little")
            if column > _MAX_COLUMN:
                raise XlsbFormulaScanError(f"{sheet_name}: column {column} is out of bounds")
            # iStyleRef is a 24-bit little-endian integer starting at byte 4.
            xf_index = int.from_bytes(payload[4:7], "little")
            styles[(row + 1, column + 1)] = xf_index
    return styles


def resolve_worksheet_targets(data: bytes) -> dict[str, str]:
    """Sheet name -> zip member path, worksheet relationships only.

    Reuses the same relationship-resolution the formula/risk scan already
    performs so the loader can fetch one worksheet's raw bytes at a time for
    per-cell style metadata without re-deriving the workbook/rels mapping.
    """
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            relationships = _workbook_relationships(archive)
            try:
                workbook_data = archive.read("xl/workbook.bin")
            except KeyError as exc:
                raise XlsbFormulaScanError(
                    "missing required XLSB part xl/workbook.bin"
                ) from exc
            targets: dict[str, str] = {}
            for sheet_name, relationship_id in _workbook_sheets(workbook_data):
                relationship = relationships.get(relationship_id)
                if relationship is not None and relationship.kind == "worksheet":
                    targets[sheet_name] = relationship.target
            return targets
    except zipfile.BadZipFile as exc:
        raise XlsbFormulaScanError("invalid XLSB ZIP package") from exc


def _package_part_bytes(archive: zipfile.ZipFile, part: str) -> int:
    normalized = part.casefold()
    return sum(
        info.file_size
        for info in archive.infolist()
        if info.filename.replace("\\", "/").casefold() == normalized
    )


def _package_worksheet_bytes(archive: zipfile.ZipFile) -> int:
    return sum(
        info.file_size
        for info in archive.infolist()
        if (name := info.filename.replace("\\", "/").casefold()).startswith(
            "xl/worksheets/"
        )
        and name.endswith(".bin")
    )


@dataclass(frozen=True, slots=True)
class PackageRiskClassification:
    """Package risks classified at the source, not subtracted from labels.

    ``aggregate`` is the historical union kept for existing findings/messages.
    ``passive``, ``blocking``, and ``unknown_external`` are disjoint: passive
    is recognized external-workbook-link metadata (safe for the isolated
    adapters, never described as risk-free); blocking is active or
    unrecognized content that always refuses; unknown_external is a non-
    hyperlink ``TargetMode="External"`` relationship whose kind is not the
    recognized ``externalLinkPath``, which conservatively refuses until later
    reachability evidence (Step 4b) can prove otherwise.
    """

    aggregate: tuple[str, ...] = ()
    passive: tuple[str, ...] = ()
    blocking: tuple[str, ...] = ()
    unknown_external: tuple[str, ...] = ()


def _risky_features(
    archive: zipfile.ZipFile, relationships: dict[str, _Relationship]
) -> PackageRiskClassification:
    aggregate: set[str] = set()
    passive: set[str] = set()
    blocking: set[str] = set()
    unknown_external: set[str] = set()

    def add_passive(label: str) -> None:
        aggregate.add(label)
        passive.add(label)

    def add_blocking(label: str) -> None:
        aggregate.add(label)
        blocking.add(label)

    for name in archive.namelist():
        lowered = "/" + name.replace("\\", "/").lower().lstrip("/")
        exact = lowered.lstrip("/")
        if exact in _RISKY_PARTS:
            add_blocking(_RISKY_PARTS[exact])
        for segment, label in _RISKY_PATH_SEGMENTS.items():
            if segment not in lowered:
                continue
            if segment == "/externallinks/":
                add_passive(label)
            else:
                add_blocking(label)
    for relationship in relationships.values():
        label = _RISKY_RELATIONSHIP_KINDS.get(relationship.kind)
        if label is None:
            continue
        if relationship.kind in _PASSIVE_EXTERNAL_LINK_KINDS:
            add_passive(label)
        else:
            add_blocking(label)
    for part in archive.namelist():
        if not part.lower().endswith(".rels"):
            continue
        try:
            root = ElementTree.fromstring(archive.read(part))
        except ElementTree.ParseError:
            add_blocking("unreadable relationship metadata")
            continue
        for element in root:
            if _local_name(element.tag) != "Relationship":
                continue
            kind = element.get("Type", "").rsplit("/", 1)[-1].lower()
            label = _RISKY_RELATIONSHIP_KINDS.get(kind)
            if label is not None:
                if kind in _PASSIVE_EXTERNAL_LINK_KINDS:
                    add_passive(label)
                else:
                    add_blocking(label)
            if element.get("TargetMode") == "External" and kind != "hyperlink":
                aggregate.add("external relationships")
                # The standard filesystem targets are externalLinkPath and
                # externalLinkLongPath. A plain externalLink relationship with
                # TargetMode=External is atypical: keep it unknown/fail-closed
                # even though package-internal externalLink metadata is passive.
                if kind in {"externallinklongpath", "externallinkpath"}:
                    passive.add("external relationships")
                else:
                    unknown_external.add("external relationships")
    return PackageRiskClassification(
        aggregate=tuple(sorted(aggregate)),
        passive=tuple(sorted(passive)),
        blocking=tuple(sorted(blocking)),
        unknown_external=tuple(sorted(unknown_external)),
    )


def scan_xlsb_formulas(data: bytes) -> XlsbFormulaScan:
    """Return formula coordinates and risky features without decoding formulas."""
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            relationships = _workbook_relationships(archive)
            try:
                workbook_data = archive.read("xl/workbook.bin")
            except KeyError as exc:
                raise XlsbFormulaScanError("missing required XLSB part xl/workbook.bin") from exc

            sheets = _workbook_sheets(workbook_data)
            formula_cells: dict[str, frozenset[tuple[int, int]]] = {}
            worksheet_metrics: dict[str, XlsbWorksheetMetrics] = {}
            for sheet_name, relationship_id in sheets:
                relationship = relationships.get(relationship_id)
                if relationship is None:
                    raise XlsbFormulaScanError(
                        f"sheet {sheet_name!r} has unresolved relationship {relationship_id!r}"
                    )
                if relationship.kind != "worksheet":
                    formula_cells[sheet_name] = frozenset()
                    worksheet_metrics[sheet_name] = XlsbWorksheetMetrics()
                    continue
                try:
                    sheet_data = archive.read(relationship.target)
                except KeyError as exc:
                    raise XlsbFormulaScanError(
                        f"sheet {sheet_name!r} is missing part {relationship.target!r}"
                    ) from exc
                formulas, metrics = _worksheet_scan(sheet_data, sheet_name)
                formula_cells[sheet_name] = formulas
                worksheet_metrics[sheet_name] = metrics

            risk = _risky_features(archive, relationships)
            return XlsbFormulaScan(
                formula_cells=formula_cells,
                risky_features=risk.aggregate,
                passive_features=risk.passive,
                blocking_features=risk.blocking,
                unknown_external_features=risk.unknown_external,
                worksheet_metrics=worksheet_metrics,
                package_worksheet_binary_bytes=_package_worksheet_bytes(archive),
                shared_string_bytes=_package_part_bytes(
                    archive, "xl/sharedStrings.bin"
                ),
                styles_bytes=_package_part_bytes(archive, "xl/styles.bin"),
                sheet_count=len(sheets),
            )
    except zipfile.BadZipFile as exc:
        raise XlsbFormulaScanError("invalid XLSB ZIP package") from exc
