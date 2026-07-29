"""Sandboxed Linux XLSB-to-OOXML formula extraction via LibreOffice."""

from __future__ import annotations

import os
import shutil
import signal
import subprocess
import tempfile
from collections.abc import Callable
from contextlib import suppress
from pathlib import Path
from typing import TypeAlias

from openpyxl import load_workbook

from qc_tool.io.formula_enrichment import (
    FormulaEnrichmentError,
    FormulaExtraction,
    FormulaMap,
    validate_formula_extraction,
)
from qc_tool.io.xlsb_formula import XlsbFormulaScan
from qc_tool.security import private_directory, private_file

DEFAULT_CONVERSION_TIMEOUT_SECONDS = 300.0
_Converter: TypeAlias = Callable[[Path, float], tuple[str, str]]


def _required_executable(name: str) -> str:
    executable = shutil.which(name)
    if executable is None:
        raise FormulaEnrichmentError(f"required Linux formula adapter {name!r} is unavailable")
    return executable


def _sandbox_command(work_dir: Path) -> tuple[list[str], str]:
    bubblewrap = _required_executable("bwrap")
    libreoffice = _required_executable("libreoffice")
    command = [
        bubblewrap,
        "--die-with-parent",
        "--new-session",
        "--unshare-all",
        "--clearenv",
    ]
    for system_path in ("/usr", "/bin", "/lib", "/lib64", "/etc"):
        if Path(system_path).exists():
            command.extend(("--ro-bind", system_path, system_path))
    command.extend(
        (
            "--proc",
            "/proc",
            "--dev",
            "/dev",
            "--tmpfs",
            "/tmp",
            "--bind",
            str(work_dir),
            "/work",
            "--setenv",
            "HOME",
            "/work/home",
            "--setenv",
            "TMPDIR",
            "/tmp",
            "--setenv",
            "LANG",
            "C.UTF-8",
            "--setenv",
            "LC_ALL",
            "C.UTF-8",
            "--chdir",
            "/work",
            libreoffice,
            "--headless",
            "--nologo",
            "--nodefault",
            "--nolockcheck",
            "--norestore",
            "--nofirststartwizard",
            "-env:UserInstallation=file:///work/profile",
            "--convert-to",
            "xlsx:Calc MS Excel 2007 XML",
            "--outdir",
            "/work/output",
            "/work/input.xlsb",
        )
    )
    return command, libreoffice


def _convert_with_libreoffice(work_dir: Path, timeout: float) -> tuple[str, str]:
    command, libreoffice = _sandbox_command(work_dir)
    try:
        version_result = subprocess.run(
            [libreoffice, "--version"],
            stdin=subprocess.DEVNULL,
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
            timeout=min(timeout, 15.0),
            check=True,
        )
    except (OSError, subprocess.SubprocessError) as exc:
        raise FormulaEnrichmentError("could not identify the LibreOffice version") from exc
    version = version_result.stdout.strip().splitlines()[0]
    if not version:
        raise FormulaEnrichmentError("LibreOffice returned an empty version")
    process = subprocess.Popen(
        command,
        stdin=subprocess.DEVNULL,
        stdout=subprocess.PIPE,
        stderr=subprocess.PIPE,
        text=True,
        encoding="utf-8",
        errors="replace",
        env={"PATH": "/usr/bin:/bin", "LANG": "C.UTF-8", "LC_ALL": "C.UTF-8"},
        start_new_session=True,
    )
    try:
        _, stderr = process.communicate(timeout=timeout)
    except subprocess.TimeoutExpired as exc:
        with suppress(OSError):
            os.killpg(process.pid, signal.SIGKILL)
        process.communicate()
        raise FormulaEnrichmentError(
            f"LibreOffice formula extraction exceeded {timeout:g} seconds"
        ) from exc
    if process.returncode != 0:
        detail = stderr.strip().splitlines()[-1] if stderr.strip() else "no diagnostic"
        raise FormulaEnrichmentError(
            f"LibreOffice formula extraction failed with exit code {process.returncode}: "
            f"{detail[:200]}"
        )
    return f"libreoffice:{version}", f"Formula text extracted by {version} in sandbox"


def _formula_text(raw: object, sheet_name: str, coordinate: str) -> str:
    value = raw if isinstance(raw, str) else getattr(raw, "text", None)
    if not isinstance(value, str) or not value.startswith("="):
        raise FormulaEnrichmentError(
            f"{sheet_name}!{coordinate}: LibreOffice produced unsupported formula data"
        )
    return value


def _read_ooxml_formulas(path: Path) -> FormulaMap:
    try:
        workbook = load_workbook(path, data_only=False, read_only=True, keep_links=False)
    except Exception as exc:
        raise FormulaEnrichmentError("LibreOffice output is not a readable OOXML workbook") from exc
    formulas: FormulaMap = {}
    try:
        for sheet in workbook.worksheets:
            sheet_formulas: dict[tuple[int, int], str] = {}
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.data_type != "f":
                        continue
                    if cell.row is None or cell.column is None:
                        raise FormulaEnrichmentError(
                            f"{sheet.title}: LibreOffice returned a formula without coordinates"
                        )
                    sheet_formulas[(cell.row, cell.column)] = _formula_text(
                        cell.value, sheet.title, cell.coordinate
                    )
            if sheet_formulas:
                formulas[sheet.title] = sheet_formulas
    finally:
        workbook.close()
    return formulas


def extract_formulas_with_libreoffice(
    data: bytes,
    scan: XlsbFormulaScan,
    *,
    timeout: float = DEFAULT_CONVERSION_TIMEOUT_SECONDS,
    converter: _Converter | None = None,
) -> FormulaExtraction:
    """Extract XLSB formula text in an isolated, networkless LibreOffice process."""
    if os.name != "posix":
        raise FormulaEnrichmentError("LibreOffice XLSB enrichment is enabled on POSIX only")
    if not scan.safe_for_external_engine:
        raise FormulaEnrichmentError(
            "XLSB contains active or external content: " + ", ".join(scan.risky_features)
        )

    with tempfile.TemporaryDirectory(prefix="qc-tool-xlsb-") as temporary:
        work_dir = private_directory(Path(temporary))
        private_directory(work_dir / "output")
        private_directory(work_dir / "home")
        input_path = work_dir / "input.xlsb"
        input_path.write_bytes(data)
        private_file(input_path)

        engine, detail = (converter or _convert_with_libreoffice)(work_dir, timeout)
        output_path = work_dir / "output" / "input.xlsx"
        if not output_path.is_file() or output_path.stat().st_size == 0:
            raise FormulaEnrichmentError("LibreOffice did not produce the expected XLSX output")
        private_file(output_path)
        extraction = FormulaExtraction(
            formulas=_read_ooxml_formulas(output_path),
            engine=engine,
            detail=detail,
        )
        validate_formula_extraction(scan, extraction)
        return extraction
