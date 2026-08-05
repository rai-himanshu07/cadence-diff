"""Excel-to-PPT figure tracing: suggest source cells, verify saved mappings.

Workflow: the tool extracts every figure occurrence from the current deck,
suggests ranked candidate source cells from the current workbook (display
match first, then label affinity), the analyst confirms, and confirmed
mappings persist in the deliverable profile. Each cycle the saved mappings
re-verify automatically: figures are relocated by their stable line
skeleton, so refreshed numbers do not break the anchor.
"""

import logging
from collections import Counter
from dataclasses import dataclass, field

from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import coordinate_to_tuple
from pydantic import BaseModel, Field
from rapidfuzz import fuzz

from qc_tool.config.profile import CrosscheckMapping, CrosscheckProfile
from qc_tool.crosscheck.numbers import (
    ParsedFigure,
    display_matches,
    extract_figures,
    numeric_skeleton,
    relative_difference,
)
from qc_tool.excel.dependency import DependencyGraph, Node, dependent_nodes_of
from qc_tool.findings import Finding, FindingClass
from qc_tool.io.model import SheetSnapshot, WorkbookSnapshot, display_cell_value
from qc_tool.ppt.extract import DeckSnapshot

logger = logging.getLogger(__name__)

_NEAR_MISS_LIMIT = 0.05  # near-miss candidates within 5% of the figure


@dataclass(frozen=True, slots=True)
class FigureOccurrence:
    slide: str
    line: str
    line_skeleton: str
    figure_index: int
    figure: ParsedFigure
    #: 1-based slide position in the deck this occurrence was extracted from.
    slide_index: int
    #: Exact enclosing table/chart shape; flattened ordinary text has none.
    shape_id: int | None = None

    @property
    def context(self) -> str:
        return f"{self.slide} {self.line_skeleton}"


@dataclass(frozen=True, slots=True)
class SourceCandidate:
    sheet: str
    cell: str
    value: float
    display_match: bool
    label_score: float
    rel_diff: float
    row_label: str = ""
    column_label: str = ""


class SuggestedSource(BaseModel):
    sheet: str
    cell: str
    value: float
    display_match: bool
    label_score: float
    rel_diff: float
    row_label: str = ""
    column_label: str = ""

    @classmethod
    def from_candidate(cls, candidate: SourceCandidate) -> "SuggestedSource":
        return cls(
            sheet=candidate.sheet,
            cell=candidate.cell,
            value=candidate.value,
            display_match=candidate.display_match,
            label_score=candidate.label_score,
            rel_diff=candidate.rel_diff,
            row_label=candidate.row_label,
            column_label=candidate.column_label,
        )


class MappingSuggestion(BaseModel):
    slide: str
    line: str
    line_skeleton: str
    figure_index: int
    figure_raw: str
    candidates: list[SuggestedSource] = Field(default_factory=list)


def occurrence_identity(occurrence: FigureOccurrence) -> tuple[str, str, int]:
    return occurrence.slide, occurrence.line_skeleton, occurrence.figure_index


def mapping_identity(mapping: CrosscheckMapping) -> tuple[str, str, int]:
    return mapping.slide, mapping.line_skeleton, mapping.figure_index


@dataclass(slots=True)
class CrosscheckResult:
    findings: list[Finding] = field(default_factory=list)
    verified: list[CrosscheckMapping] = field(default_factory=list)


def extract_deck_figures(deck: DeckSnapshot) -> list[FigureOccurrence]:
    """Every visible text, table-cell, and native chart-label figure."""
    occurrences: list[FigureOccurrence] = []
    for slide in deck.slides:
        ordinary_values: Counter[float] = Counter()
        for line in slide.texts:
            for index, figure in enumerate(extract_figures(line)):
                ordinary_values[round(figure.value, 12)] += 1
                occurrences.append(
                    FigureOccurrence(
                        slide=slide.display_name,
                        line=line,
                        line_skeleton=numeric_skeleton(line),
                        figure_index=index,
                        figure=figure,
                        slide_index=slide.index + 1,
                    )
                )
        for table in slide.tables:
            if not table.rows:
                continue
            headers = table.rows[0]
            for row in table.rows[1:]:
                row_label = row[0] if row else ""
                for col_idx, cell_text in enumerate(row[1:], start=1):
                    figures = extract_figures(cell_text)
                    if len(figures) != 1:
                        continue
                    ordinary_values[round(figures[0].value, 12)] += 1
                    header = headers[col_idx] if col_idx < len(headers) else str(col_idx)
                    occurrences.append(
                        FigureOccurrence(
                            slide=slide.display_name,
                            line=cell_text,
                            line_skeleton=f"table:{row_label}/{header}",
                            figure_index=0,
                            figure=figures[0],
                            slide_index=slide.index + 1,
                            shape_id=table.shape_id or None,
                        )
                    )
        for chart in slide.charts:
            chart_anchor = chart.title or f"chart[{chart.source_index}]"
            for plot in chart.plots:
                series_by_index = {
                    series.source_index: series for series in plot.series
                }
                for label in plot.visible_labels:
                    series = series_by_index.get(label.series_source_index)
                    series_anchor = (
                        series.name
                        if series is not None and series.name
                        else f"series[{label.series_source_index}]"
                    )
                    category_anchor = (
                        label.category
                        if label.category
                        else f"point[{label.category_index}]"
                    )
                    for figure_index, figure in enumerate(extract_figures(label.text)):
                        rounded = round(figure.value, 12)
                        if ordinary_values[rounded] > 0:
                            ordinary_values[rounded] -= 1
                            continue
                        occurrences.append(
                            FigureOccurrence(
                                slide=slide.display_name,
                                line=label.text,
                                line_skeleton=(
                                    f"chart:{chart_anchor}/{series_anchor}/"
                                    f"{category_anchor}"
                                ),
                                figure_index=figure_index,
                                figure=figure,
                                slide_index=slide.index + 1,
                                shape_id=chart.shape_id or None,
                            )
                        )
    return occurrences


def _nearest_left_label(sheet: SheetSnapshot, row: int, col: int) -> str:
    for candidate_col in range(col - 1, 0, -1):
        cell = sheet.cells.get((row, candidate_col))
        if cell is not None and isinstance(cell.value, str):
            return cell.value
    return ""


def _nearest_above_label(sheet: SheetSnapshot, row: int, col: int) -> str:
    for candidate_row in range(row - 1, 0, -1):
        cell = sheet.cells.get((candidate_row, col))
        if cell is not None and isinstance(cell.value, str):
            return cell.value
    return ""


def suggest_sources(
    occurrence: FigureOccurrence, workbook: WorkbookSnapshot, *, limit: int = 5
) -> list[SourceCandidate]:
    """Ranked candidate source cells: display matches first, then near misses."""
    candidates: list[SourceCandidate] = []
    for sheet in workbook.sheets:
        for (row, col), cell in sheet.cells.items():
            value = cell.value
            if isinstance(value, bool) or not isinstance(value, int | float):
                continue
            value = float(value)
            matches = display_matches(occurrence.figure, value)
            rel_diff = relative_difference(occurrence.figure, value)
            if not matches and rel_diff > _NEAR_MISS_LIMIT:
                continue
            row_label = _nearest_left_label(sheet, row, col)
            column_label = _nearest_above_label(sheet, row, col)
            label_score = fuzz.token_set_ratio(
                occurrence.context, f"{sheet.name} {row_label} {column_label}"
            )
            candidates.append(
                SourceCandidate(
                    sheet=sheet.name,
                    cell=f"{get_column_letter(col)}{row}",
                    value=value,
                    display_match=matches,
                    label_score=label_score,
                    rel_diff=rel_diff,
                    row_label=row_label,
                    column_label=column_label,
                )
            )
    candidates.sort(
        key=lambda c: (not c.display_match, -c.label_score, c.rel_diff, c.sheet, c.cell)
    )
    return candidates[:limit]


def build_mapping(
    occurrence: FigureOccurrence, candidate: SourceCandidate, *, label: str = ""
) -> CrosscheckMapping:
    """A profile-persistable mapping from a confirmed suggestion."""
    return CrosscheckMapping(
        slide=occurrence.slide,
        line_skeleton=occurrence.line_skeleton,
        figure_index=occurrence.figure_index,
        label=label or f"{candidate.row_label} {candidate.column_label}".strip(),
        source_sheet=candidate.sheet,
        source_cell=candidate.cell,
    )


def _locate_occurrence(
    occurrences: list[FigureOccurrence], mapping: CrosscheckMapping
) -> FigureOccurrence | None:
    for occurrence in occurrences:
        if (
            occurrence.slide == mapping.slide
            and occurrence.line_skeleton == mapping.line_skeleton
            and occurrence.figure_index == mapping.figure_index
        ):
            return occurrence
    return None


def annotate_ppt_chart_impacts(
    findings: list[Finding],
    deck: DeckSnapshot,
    profile: CrosscheckProfile,
    dependency_graph: DependencyGraph,
) -> None:
    """Attach confirmed visible chart-label mappings as downstream impacts."""
    occurrences = extract_deck_figures(deck)
    chart_mappings: list[tuple[Node, str]] = []
    for mapping in profile.mappings:
        occurrence = _locate_occurrence(occurrences, mapping)
        if occurrence is None or not occurrence.line_skeleton.startswith("chart:"):
            continue
        try:
            row, column = coordinate_to_tuple(mapping.source_cell)
        except ValueError:
            continue
        chart_anchor = occurrence.line_skeleton.removeprefix("chart:")
        chart_mappings.append(
            (
                (mapping.source_sheet, row, column),
                (
                    f"PowerPoint chart label {chart_anchor!r} on slide "
                    f"{occurrence.slide!r}"
                ),
            )
        )
    for finding in findings:
        if finding.sheet is None or finding.location is None:
            continue
        try:
            row, column = coordinate_to_tuple(finding.location)
        except ValueError:
            continue
        source: Node = (finding.sheet, row, column)
        affected = {source, *dependent_nodes_of(dependency_graph, source)}
        impacts = {
            impact
            for mapped_source, impact in chart_mappings
            if mapped_source in affected
        }
        finding.impacts = sorted({*finding.impacts, *impacts})


def verify_mappings(
    deck: DeckSnapshot, workbook: WorkbookSnapshot, profile: CrosscheckProfile
) -> CrosscheckResult:
    """Re-verify every saved mapping against the current deck + workbook."""
    result = CrosscheckResult()
    occurrences = extract_deck_figures(deck)
    for mapping in profile.mappings:
        display = mapping.label or f"{mapping.slide}: {mapping.line_skeleton}"
        occurrence = _locate_occurrence(occurrences, mapping)
        if occurrence is None:
            result.findings.append(
                Finding(
                    artifact="crosscheck",
                    finding_class=FindingClass.CROSSCHECK_UNRESOLVED,
                    slide=mapping.slide,
                    element=display,
                    message=(
                        f"{display}: mapped figure not found in the current deck "
                        "(slide or wording changed); re-confirm the mapping"
                    ),
                )
            )
            continue
        try:
            sheet = workbook.sheet(mapping.source_sheet)
        except KeyError:
            sheet = None
        cell = sheet.cell(mapping.source_cell) if sheet is not None else None
        value = cell.value if cell is not None else None
        if isinstance(value, bool) or not isinstance(value, int | float):
            result.findings.append(
                Finding(
                    artifact="crosscheck",
                    finding_class=FindingClass.CROSSCHECK_UNRESOLVED,
                    slide=mapping.slide,
                    slide_index=occurrence.slide_index,
                    focus_shape_id=occurrence.shape_id,
                    element=display,
                    location=f"{mapping.source_sheet}!{mapping.source_cell}",
                    message=(
                        f"{display}: source cell {mapping.source_sheet}!"
                        f"{mapping.source_cell} has no numeric value (formula "
                        "without cached result, or cell moved)"
                    ),
                )
            )
            continue
        if display_matches(occurrence.figure, float(value)):
            result.verified.append(mapping)
            continue
        result.findings.append(
            Finding(
                artifact="crosscheck",
                finding_class=FindingClass.CROSSCHECK_MISMATCH,
                slide=mapping.slide,
                slide_index=occurrence.slide_index,
                focus_shape_id=occurrence.shape_id,
                element=display,
                location=f"{mapping.source_sheet}!{mapping.source_cell}",
                baseline_value=display_cell_value(value),
                current_value=occurrence.figure.raw,
                message=(
                    f"{display}: deck shows {occurrence.figure.raw} but "
                    f"{mapping.source_sheet}!{mapping.source_cell} holds {value}"
                ),
            )
        )
    return result
