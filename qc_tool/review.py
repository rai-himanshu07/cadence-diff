"""Deterministic analyst review groups derived from atomic findings."""

from __future__ import annotations

import hashlib
import json
import re
from collections import defaultdict, deque
from collections.abc import Iterator, Sequence
from dataclasses import dataclass, replace
from typing import TypeAlias

from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import coordinate_to_tuple, range_boundaries

from qc_tool.excel.formula_tokens import tokenize_formula
from qc_tool.findings import (
    Finding,
    FindingClass,
    FindingProvenance,
    FindingSubtype,
    FindingTemporalContext,
    Materiality,
    Severity,
)
from qc_tool.story import ChangeStory, StoryKind

Coordinate: TypeAlias = tuple[int, int]
Rectangle: TypeAlias = tuple[int, int, int, int]
FindingIdentity: TypeAlias = tuple[str, ...]

_A1_RE = re.compile(r"^\$?[A-Z]{1,3}\$?[1-9]\d*$", re.IGNORECASE)
_FINDING_ID_RE = re.compile(r"^F(\d+)$")
_PRESENTATION_CLASSES = frozenset(
    {
        FindingClass.NUMBER_FORMAT_CHANGED,
        FindingClass.STYLE_CHANGED,
    }
)
_GROUPABLE_CLASSES = frozenset(
    {
        FindingClass.VALUE_CHANGED,
        FindingClass.FORMULA_ERROR,
        FindingClass.FORMULA_HARDCODED,
        FindingClass.FORMULA_REMOVED,
        FindingClass.FORMULA_MISSING,
        FindingClass.FORMULA_CACHE_MISSING,
        FindingClass.FORMULA_NOT_EXTENDED,
        FindingClass.FORMULA_LOGIC_CHANGED,
        FindingClass.FORMULA_INCONSISTENT,
        FindingClass.NUMBER_FORMAT_CHANGED,
        FindingClass.STYLE_CHANGED,
        FindingClass.REQUIRED_VALUE_MISSING,
        FindingClass.NUMERIC_BOUND_VIOLATION,
    }
)
_SEVERITY_RANK = {
    Severity.CRITICAL: 0,
    Severity.WARNING: 1,
    Severity.INFO: 2,
    Severity.EXPECTED: 3,
}


@dataclass(frozen=True, slots=True)
class ReviewGroup:
    """One analyst decision backed by one or more atomic findings."""

    group_id: str
    finding_class: FindingClass
    severity: Severity
    artifact: str
    sheet: str | None
    slide: str | None
    element: str
    expected_growth: bool
    ranges: tuple[str, ...]
    bounding_range: str
    baseline_ranges: tuple[str, ...]
    baseline_bounding_range: str
    baseline_mixed: bool
    members: tuple[Finding, ...]
    spatial: bool
    artifact_member: str = "primary"

    @property
    def member_count(self) -> int:
        return len(self.members)


@dataclass(frozen=True, slots=True)
class ReviewCounts:
    """Review decisions, finding records, and represented changes by severity.

    ``atomic_findings`` retains its historical field name for compatibility;
    its values count stored ``Finding`` records. One population record can
    represent many underlying changes.
    """

    review_items: dict[Severity, int]
    atomic_findings: dict[Severity, int]
    represented_changes: dict[Severity, int] | None = None


def represented_change_count(finding: Finding) -> int:
    """Number of underlying changes represented by one finding record."""
    if finding.population is not None:
        return finding.population.member_count
    return 1


@dataclass(frozen=True, slots=True)
class ReviewAnnotation:
    """One atomic annotation update produced by a group review action."""

    finding_id: str
    severity: str | None
    comment: str


def finding_identity_key(finding: Finding) -> FindingIdentity:
    """Stable identity shared with Re-QC and review-group fingerprints."""
    base = (
        finding.artifact,
        finding.finding_class.value,
        finding.sheet or "",
        finding.slide or "",
        finding.location or finding.baseline_location or "",
        finding.element or "",
    )
    if finding.artifact_member != "primary":
        return (*base, finding.artifact_member)
    return base


def requeue_identity_key(finding: Finding) -> FindingIdentity:
    """Cross-run pairing identity for Re-QC delta and carry-forward only.

    Populations use `population_identity_digest`, which excludes geometry
    so member-set churn (a row inserted/deleted/permuted between runs) does
    not break identity across runs (Criterion 5: identity stability).
    Atomic findings keep the existing location-based `finding_identity_key`
    unchanged. Do not use this for within-run grouping -- `review_stream.py`
    and `build_review_groups` need the location-based key even for
    populations, since a run only ever contains one instance of each.
    """
    if finding.population is not None:
        return ("population", population_identity_digest(finding))
    return finding_identity_key(finding)


def _population_pairs(finding: Finding) -> Iterator[tuple[str, str]]:
    """Shared core: yields every (current, baseline) pair one at a time.

    Pairs mode already stores the explicit list. Shift mode enumerates every
    cell in every current-side rectangle and applies the constant offset --
    the same reconstruction the plan's membership codec guarantees exactly.
    Iteration order is fixed (rectangle order, then row-major within each
    rectangle) so a bounded page always names the same members regardless
    of population size.
    """
    population = finding.population
    if population is None:
        return
    membership = population.membership
    if membership.baseline_mode == "pairs":
        yield from membership.pairs or ()
        return
    dr, dc = membership.shift or (0, 0)
    for rectangle in membership.current_rectangles:
        min_col, min_row, max_col, max_row = range_boundaries(rectangle)
        if min_col is None or min_row is None or max_col is None or max_row is None:
            continue
        for row in range(min_row, max_row + 1):
            for col in range(min_col, max_col + 1):
                current = f"{get_column_letter(col)}{row}"
                baseline = f"{get_column_letter(col + dc)}{row + dr}"
                yield (current, baseline)


def population_members(finding: Finding) -> tuple[tuple[str, str], ...]:
    """Decode a population's ENTIRE membership codec into (current, baseline)
    pairs.

    Materializes every member -- safe only when the caller already knows the
    population is small (e.g. tests, or after checking `member_count`).
    Prefer `population_members_iter` (streaming) or `population_members_page`
    (bounded random access) for a population whose size is not already known
    to be small; see Criterion 9.
    """
    return tuple(_population_pairs(finding))


def population_members_iter(finding: Finding) -> Iterator[tuple[str, str]]:
    """Streaming decode: O(1) extra memory regardless of population size.

    For callers (e.g. carry-forward matching) that must visit every member
    exactly once but never need more than the current member resident.
    """
    return _population_pairs(finding)


def population_members_page(
    finding: Finding, start: int, count: int
) -> tuple[tuple[str, str], ...]:
    """Decode only ``[start, start + count)`` members without ever
    materializing the rest -- a UI pager stays bounded regardless of how
    large the population is (Criterion 9).
    """
    if count <= 0 or start < 0:
        return ()
    population = finding.population
    if population is None:
        return ()
    membership = population.membership
    if membership.baseline_mode == "pairs":
        return tuple((membership.pairs or ())[start : start + count])
    dr, dc = membership.shift or (0, 0)
    end = start + count
    pairs: list[tuple[str, str]] = []
    seen = 0
    for rectangle in membership.current_rectangles:
        min_col, min_row, max_col, max_row = range_boundaries(rectangle)
        if min_col is None or min_row is None or max_col is None or max_row is None:
            continue
        width = max_col - min_col + 1
        height = max_row - min_row + 1
        rect_size = width * height
        if seen + rect_size <= start:
            seen += rect_size
            continue
        if seen >= end:
            break
        for offset in range(max(0, start - seen), min(rect_size, end - seen)):
            row = min_row + offset // width
            col = min_col + offset % width
            current = f"{get_column_letter(col)}{row}"
            baseline = f"{get_column_letter(col + dc)}{row + dr}"
            pairs.append((current, baseline))
        seen += rect_size
        if seen >= end:
            break
    return tuple(pairs)


def _coordinate(location: str | None) -> Coordinate | None:
    if location is None or ":" in location or _A1_RE.fullmatch(location) is None:
        return None
    try:
        return coordinate_to_tuple(location.replace("$", ""))
    except ValueError:
        return None


def _baseline_mode(finding: Finding, current: Coordinate) -> tuple[object, ...]:
    if finding.baseline_location is None:
        return ("no-baseline",)
    baseline = _coordinate(finding.baseline_location)
    if baseline is None:
        return ("marker", finding.baseline_location)
    return ("translation", current[0] - baseline[0], current[1] - baseline[1])


def _producer_subtype(finding: Finding) -> str:
    if finding.finding_class is FindingClass.FORMULA_ERROR:
        value = finding.current_value or ""
        return "formula-text" if value.startswith("=") else "saved-value"
    return ""


def _partition_key(finding: Finding, current: Coordinate) -> tuple[object, ...]:
    severity = finding.severity or Severity.WARNING
    base = (
        finding.artifact,
        finding.sheet or "",
        finding.finding_class.value,
        severity.value,
        finding.expected_growth,
        (finding.element or "").strip(),
        finding.waiver_reason,
        finding.waiver_expires,
        _producer_subtype(finding),
        _baseline_mode(finding, current),
    )
    return (*base, finding.artifact_member) if finding.artifact_member != "primary" else base


def _eligible(finding: Finding) -> Coordinate | None:
    if (
        finding.artifact != "excel"
        or finding.sheet is None
        or finding.finding_class not in _GROUPABLE_CLASSES
    ):
        return None
    return _coordinate(finding.location)


def _finding_order(finding: Finding) -> tuple[object, ...]:
    match = _FINDING_ID_RE.fullmatch(finding.finding_id)
    if match is not None:
        return (0, int(match.group(1)))
    return (1, finding_identity_key(finding), finding.finding_id)


def _horizontal_runs(columns: list[int]) -> list[tuple[int, int]]:
    if not columns:
        return []
    runs: list[tuple[int, int]] = []
    start = previous = columns[0]
    for column in columns[1:]:
        if column == previous + 1:
            previous = column
            continue
        runs.append((start, previous))
        start = previous = column
    runs.append((start, previous))
    return runs


def _rectangles(coordinates: set[Coordinate]) -> list[Rectangle]:
    by_row: dict[int, list[int]] = defaultdict(list)
    for row, column in coordinates:
        by_row[row].append(column)

    rectangles: list[Rectangle] = []
    active: dict[tuple[int, int], int] = {}
    previous_row: int | None = None
    for row in sorted(by_row):
        if previous_row is not None and row != previous_row + 1:
            rectangles.extend(
                (start_row, start_col, previous_row, end_col)
                for (start_col, end_col), start_row in active.items()
            )
            active.clear()
        runs = set(_horizontal_runs(sorted(set(by_row[row]))))
        for run in set(active).difference(runs):
            start_col, end_col = run
            rectangles.append((active.pop(run), start_col, row - 1, end_col))
        for run in runs.difference(active):
            active[run] = row
        previous_row = row

    if previous_row is not None:
        rectangles.extend(
            (start_row, start_col, previous_row, end_col)
            for (start_col, end_col), start_row in active.items()
        )
    return sorted(rectangles)


def _range(rectangle: Rectangle) -> str:
    min_row, min_col, max_row, max_col = rectangle
    start = f"{get_column_letter(min_col)}{min_row}"
    end = f"{get_column_letter(max_col)}{max_row}"
    return start if start == end else f"{start}:{end}"


def _bounding_range(coordinates: set[Coordinate]) -> str:
    if not coordinates:
        return ""
    rows = [row for row, _ in coordinates]
    columns = [column for _, column in coordinates]
    return _range((min(rows), min(columns), max(rows), max(columns)))


def _group_id(key: tuple[object, ...], members: tuple[Finding, ...]) -> str:
    return _group_id_from_identities(
        key, tuple(finding_identity_key(member) for member in members)
    )


def _group_id_from_identities(
    key: tuple[object, ...],
    identities: tuple[FindingIdentity, ...],
) -> str:
    payload = json.dumps(
        {"key": key, "members": sorted(identities)},
        separators=(",", ":"),
        sort_keys=True,
    ).encode("utf-8")
    return "G" + hashlib.sha256(payload).hexdigest()[:12].upper()


def _component_groups(
    key: tuple[object, ...],
    entries: list[tuple[Coordinate, Finding]],
) -> list[ReviewGroup]:
    by_coordinate: dict[Coordinate, list[Finding]] = defaultdict(list)
    for coordinate, finding in entries:
        by_coordinate[coordinate].append(finding)
    pending = set(by_coordinate)
    groups: list[ReviewGroup] = []
    while pending:
        first = min(pending)
        pending.remove(first)
        component = {first}
        queue = deque([first])
        while queue:
            row, column = queue.popleft()
            for neighbor in (
                (row - 1, column),
                (row + 1, column),
                (row, column - 1),
                (row, column + 1),
            ):
                if neighbor not in pending:
                    continue
                pending.remove(neighbor)
                component.add(neighbor)
                queue.append(neighbor)

        members = tuple(
            sorted(
                (
                    finding
                    for coordinate in component
                    for finding in by_coordinate[coordinate]
                ),
                key=_finding_order,
            )
        )
        current_ranges = tuple(_range(item) for item in _rectangles(component))
        baseline_coordinates = {
            coordinate
            for member in members
            if (coordinate := _coordinate(member.baseline_location)) is not None
        }
        baseline_mixed = 0 < len(baseline_coordinates) < len(members)
        baseline_ranges = (
            tuple(_range(item) for item in _rectangles(baseline_coordinates))
            if len(baseline_coordinates) == len(members)
            else ()
        )
        representative = members[0]
        severity = representative.severity or Severity.WARNING
        groups.append(
            ReviewGroup(
                group_id=_group_id(key, members),
                finding_class=representative.finding_class,
                severity=severity,
                artifact=representative.artifact,
                sheet=representative.sheet,
                slide=representative.slide,
                element=representative.element or "",
                expected_growth=representative.expected_growth,
                ranges=current_ranges,
                bounding_range=_bounding_range(component),
                baseline_ranges=baseline_ranges,
                baseline_bounding_range=(
                    _bounding_range(baseline_coordinates)
                    if baseline_ranges
                    else ""
                ),
                baseline_mixed=baseline_mixed,
                members=members,
                spatial=True,
                artifact_member=representative.artifact_member,
            )
        )
    return groups


def _singleton(finding: Finding) -> ReviewGroup:
    severity = finding.severity or Severity.WARNING
    current = _coordinate(finding.location)
    baseline = _coordinate(finding.baseline_location)
    key = (
        "singleton",
        finding_identity_key(finding),
        severity.value,
        finding.message,
        finding.baseline_value or "",
        finding.current_value or "",
        tuple(finding.impacts),
    )
    return ReviewGroup(
        group_id=_group_id(key, (finding,)),
        finding_class=finding.finding_class,
        severity=severity,
        artifact=finding.artifact,
        sheet=finding.sheet,
        slide=finding.slide,
        element=finding.element or "",
        expected_growth=finding.expected_growth,
        ranges=(_range((*current, *current)),) if current is not None else (),
        bounding_range=(
            _range((*current, *current))
            if current is not None
            else finding.location or finding.baseline_location or ""
        ),
        baseline_ranges=(
            (_range((*baseline, *baseline)),) if baseline is not None else ()
        ),
        baseline_bounding_range=(
            _range((*baseline, *baseline)) if baseline is not None else ""
        ),
        baseline_mixed=False,
        members=(finding,),
        spatial=False,
        artifact_member=getattr(finding, "artifact_member", "primary"),
    )


def _review_sort_key(group: ReviewGroup) -> tuple[object, ...]:
    coordinate = _coordinate(group.bounding_range.partition(":")[0])
    return (
        _SEVERITY_RANK[group.severity],
        group.artifact,
        group.sheet or group.slide or "",
        coordinate or (10**9, 10**9),
        group.finding_class.value,
        group.element,
        group.group_id,
        tuple(
            (
                member.message,
                member.baseline_value or "",
                member.current_value or "",
                tuple(member.impacts),
                member.finding_id,
            )
            for member in group.members
        ),
    )


def build_review_groups(findings: Sequence[Finding]) -> list[ReviewGroup]:
    """Build a lossless deterministic review layer over atomic findings."""
    partitions: dict[tuple[object, ...], list[tuple[Coordinate, Finding]]] = (
        defaultdict(list)
    )
    groups: list[ReviewGroup] = []
    for finding in findings:
        coordinate = _eligible(finding)
        if coordinate is None:
            groups.append(_singleton(finding))
            continue
        partitions[_partition_key(finding, coordinate)].append((coordinate, finding))
    for key, entries in partitions.items():
        groups.extend(_component_groups(key, entries))
    ordered = sorted(groups, key=_review_sort_key)
    seen: dict[str, int] = defaultdict(int)
    unique: list[ReviewGroup] = []
    for group in ordered:
        seen[group.group_id] += 1
        occurrence = seen[group.group_id]
        unique.append(
            group
            if occurrence == 1
            else replace(group, group_id=f"{group.group_id}-{occurrence}")
        )
    return unique


def review_counts(groups: list[ReviewGroup]) -> ReviewCounts:
    """Count decisions, finding records, and represented changes by severity."""
    review_items = dict.fromkeys(Severity, 0)
    atomic_findings = dict.fromkeys(Severity, 0)
    represented_changes = dict.fromkeys(Severity, 0)
    for group in groups:
        review_items[group.severity] += 1
        atomic_findings[group.severity] += group.member_count
        represented_changes[group.severity] += sum(
            represented_change_count(finding) for finding in group.members
        )
    return ReviewCounts(
        review_items=review_items,
        atomic_findings=atomic_findings,
        represented_changes=represented_changes,
    )


# --- semantic pattern groups -------------------------------------------------


def _formula_signature(formula: str | None) -> tuple[str, ...]:
    """Token shape with concrete references and literals abstracted away."""
    if not formula:
        return ()
    try:
        tokens = tokenize_formula(formula)
    except Exception:  # malformed formulas keep their own stable shape
        return ("<unparsed>",)
    shape: list[str] = []
    for token in tokens:
        if token.type == "OPERAND" and token.subtype == "RANGE":
            shape.append("REF")
        elif token.type == "OPERAND" and token.subtype == "NUMBER":
            shape.append("NUM")
        elif token.type == "OPERAND" and token.subtype == "TEXT":
            shape.append("TEXT")
        else:
            shape.append(token.value.casefold())
    return tuple(shape)


def _transformation_shape(finding: Finding) -> str:
    if finding.finding_class is FindingClass.FORMULA_INCONSISTENT:
        # The dominant pattern stays exact so one group never mixes two patterns.
        payload: tuple[tuple[str, ...], tuple[str, ...]] = (
            (finding.baseline_value or "",),
            _formula_signature(finding.current_value),
        )
    elif finding.finding_class is FindingClass.FORMULA_LOGIC_CHANGED:
        if (
            finding.subtype
            in (FindingSubtype.FORMULA_WRAPPED, FindingSubtype.FORMULA_UNWRAPPED)
            and finding.event_key
        ):
            # One wrapper rollout = one skeleton, regardless of the varied
            # inner formulas it wrapped.
            payload = ((finding.event_key,), ())
        else:
            payload = (
                _formula_signature(finding.baseline_value),
                _formula_signature(finding.current_value),
            )
    elif finding.finding_class in _PRESENTATION_CLASSES:
        payload = ((finding.baseline_value or "",), (finding.current_value or "",))
    else:
        return ""
    encoded = json.dumps(payload, separators=(",", ":"), sort_keys=True).encode("utf-8")
    return hashlib.sha256(encoded).hexdigest()[:12]


def _baseline_translation_mode(finding: Finding) -> str:
    if finding.baseline_location is None:
        return "no-baseline"
    baseline = _coordinate(finding.baseline_location)
    current = _coordinate(finding.location)
    if baseline is None or current is None:
        return "marker"
    return "in-place" if baseline == current else "translated"


def _pattern_key(finding: Finding) -> tuple[object, ...]:
    severity = finding.severity or Severity.WARNING
    population_event = (
        finding.event_key
        if finding.subtype is FindingSubtype.COLUMNAR_ERROR_POPULATION
        else ""
    )
    base = (
        finding.artifact,
        finding.sheet or "",
        finding.slide or "",
        finding.slide_index or 0,
        finding.baseline_slide_index or 0,
        finding.finding_class.value,
        severity.value,
        finding.expected_growth,
        finding.expected_reason.value if finding.expected_reason is not None else "",
        finding.provenance.value if finding.provenance is not None else "",
        finding.subtype.value if finding.subtype is not None else "",
        finding.materiality.value if finding.materiality is not None else "",
        (
            finding.temporal_context.value
            if finding.temporal_context is not None
            else ""
        ),
        tuple(sorted(tag.value for tag in finding.evidence_tags)),
        population_event,
        _transformation_shape(finding),
        _baseline_translation_mode(finding),
        finding.waiver_reason,
        finding.waiver_expires,
    )
    return (*base, finding.artifact_member) if finding.artifact_member != "primary" else base


def population_key(finding: Finding, shape_pair: tuple[str, str]) -> tuple[object, ...]:
    """Pre-story population identity/grouping key (group-first plan, A2).

    ``_pattern_key`` minus ``evidence_tags`` -- story tagging sets additional
    tags at finalize, so a value that changes afterward cannot be part of a
    key decided beforehand -- with the transformation shape replaced by the
    caller's exact digest pair (``shape_pair``). This is stricter than
    ``_transformation_shape``'s token abstraction: two members are grouped
    together only when their exact R1C1/format text is identical, per the
    plan's Criterion 4.
    """
    severity = finding.severity or Severity.WARNING
    population_event = (
        finding.event_key
        if finding.subtype is FindingSubtype.COLUMNAR_ERROR_POPULATION
        else ""
    )
    base = (
        finding.artifact,
        finding.sheet or "",
        finding.slide or "",
        finding.slide_index or 0,
        finding.baseline_slide_index or 0,
        finding.finding_class.value,
        severity.value,
        finding.expected_growth,
        finding.expected_reason.value if finding.expected_reason is not None else "",
        finding.provenance.value if finding.provenance is not None else "",
        finding.subtype.value if finding.subtype is not None else "",
        finding.materiality.value if finding.materiality is not None else "",
        (
            finding.temporal_context.value
            if finding.temporal_context is not None
            else ""
        ),
        population_event,
        shape_pair,
        _baseline_translation_mode(finding),
        finding.waiver_reason,
        finding.waiver_expires,
    )
    return (*base, finding.artifact_member) if finding.artifact_member != "primary" else base


def population_identity_digest(finding: Finding) -> str:
    """Stable digest of a population finding's identity (group-first plan).

    Identity = ``(artifact, finding_class, sheet, shape_before_digest,
    shape_after_digest, artifact_member)`` -- narrower than
    ``population_key`` (no severity/expected/provenance/subtype/materiality/
    temporal/waiver state). Shared by ``Finding.root_cause_key`` and the
    attestation v4 population manifest so both name the same population with
    the same value.
    """
    assert finding.population is not None
    payload = {
        "artifact": finding.artifact,
        "finding_class": finding.finding_class.value,
        "sheet": finding.sheet,
        "shape_before_digest": finding.population.shape_before_digest,
        "shape_after_digest": finding.population.shape_after_digest,
        "artifact_member": finding.artifact_member,
    }
    encoded = json.dumps(payload, sort_keys=True, separators=(",", ":")).encode("utf-8")
    return hashlib.sha256(encoded).hexdigest()


def _pattern_group(key: tuple[object, ...], members: tuple[Finding, ...]) -> ReviewGroup:
    representative = members[0]
    coordinates = {
        coordinate
        for member in members
        if (coordinate := _coordinate(member.location)) is not None
    }
    baseline_coordinates = {
        coordinate
        for member in members
        if (coordinate := _coordinate(member.baseline_location)) is not None
    }
    ranges = tuple(_range(item) for item in _rectangles(coordinates))
    baseline_ranges = (
        tuple(_range(item) for item in _rectangles(baseline_coordinates))
        if len(baseline_coordinates) == len(members)
        else ()
    )
    return ReviewGroup(
        group_id=_group_id(key, members),
        finding_class=representative.finding_class,
        severity=representative.severity or Severity.WARNING,
        artifact=representative.artifact,
        sheet=representative.sheet,
        slide=representative.slide,
        element=representative.element or "",
        expected_growth=representative.expected_growth,
        ranges=ranges,
        bounding_range=(
            _bounding_range(coordinates)
            if coordinates
            else representative.location or representative.baseline_location or ""
        ),
        baseline_ranges=baseline_ranges,
        baseline_bounding_range=(
            _bounding_range(baseline_coordinates) if baseline_ranges else ""
        ),
        baseline_mixed=0 < len(baseline_coordinates) < len(members),
        members=members,
        spatial=False,
        artifact_member=representative.artifact_member,
    )


def build_pattern_groups(findings: Sequence[Finding]) -> list[ReviewGroup]:
    """Semantically pure review groups keyed on meaning, not adjacency.

    One group may span disconnected ranges. Every member shares its finding
    class, severity, expected state, provenance, subtype, transformation shape,
    baseline-translation mode, and waiver state.
    """
    partitions: dict[tuple[object, ...], list[Finding]] = defaultdict(list)
    for finding in findings:
        partitions[_pattern_key(finding)].append(finding)
    groups = [
        _pattern_group(key, tuple(sorted(members, key=_finding_order)))
        for key, members in partitions.items()
    ]
    ordered = sorted(groups, key=_review_sort_key)
    seen: dict[str, int] = defaultdict(int)
    unique: list[ReviewGroup] = []
    for group in ordered:
        seen[group.group_id] += 1
        occurrence = seen[group.group_id]
        unique.append(
            group
            if occurrence == 1
            else replace(group, group_id=f"{group.group_id}-{occurrence}")
        )
    return unique


def count_pattern_groups(groups: list[ReviewGroup]) -> ReviewCounts:
    """Count semantic pattern decisions and their atomic members by severity."""
    return review_counts(groups)


def format_ranges(
    ranges: Sequence[str], bounding_range: str, *, max_spans: int = 3
) -> str:
    """Compact exact group geometry for human-facing tables."""
    if not ranges:
        return bounding_range
    if len(ranges) <= max_spans:
        return "; ".join(ranges)
    return f"{'; '.join(ranges[:max_spans])}; +{len(ranges) - max_spans} spans"


def format_group_ranges(group: ReviewGroup, *, max_spans: int = 3) -> str:
    """Compact exact group geometry for human-facing tables."""
    return format_ranges(group.ranges, group.bounding_range, max_spans=max_spans)


def population_summary_text(finding: Finding, *, max_spans: int = 3) -> str:
    """One-line population summary shared by the UI panel and both reports."""
    population = finding.population
    if population is None:
        return ""
    membership = population.membership
    mapping = (
        f"shift {membership.shift}"
        if membership.baseline_mode == "shift"
        else f"{len(membership.pairs or ())} explicit baseline pairs"
    )
    ranges = format_ranges(
        membership.current_rectangles, finding.location or "", max_spans=max_spans
    )
    return f"{population.member_count:,} cells; {ranges}; {mapping}"


def apply_group_review(
    group: ReviewGroup,
    *,
    severity: Severity | None,
    comment: str,
    replace_existing: bool,
) -> list[ReviewAnnotation]:
    """Apply one explicit group decision and return changed atomic annotations."""
    updates: list[ReviewAnnotation] = []
    normalized_comment = comment.strip()
    for finding in group.members:
        severity_changed = bool(
            severity is not None
            and (replace_existing or not finding.severity_overridden)
        )
        comment_changed = bool(
            normalized_comment
            and (replace_existing or not finding.analyst_comment)
        )
        if severity_changed:
            finding.severity = severity
            finding.severity_overridden = True
        if comment_changed:
            finding.analyst_comment = normalized_comment
        if not severity_changed and not comment_changed:
            continue
        updates.append(
            ReviewAnnotation(
                finding_id=finding.finding_id,
                severity=(
                    finding.severity.value
                    if finding.severity_overridden and finding.severity is not None
                    else None
                ),
                comment=finding.analyst_comment,
            )
        )
    return updates


# --- guided review prioritization --------------------------------------------

#: Every signal is derived from a field a producer already set, so a rationale
#: can always be traced back to observed evidence rather than to a heuristic.
_PRIORITY_WEIGHTS: dict[str, int] = {
    "severity_critical": 1000,
    "severity_warning": 400,
    "severity_info": 50,
    "material_delta": 300,
    "historical_change": 250,
    "new_defect": 200,
    "downstream_impacts": 150,
    "unexplained_residual": 120,
    "wide_population": 80,
    "inherited_only": -200,
    "noise_only": -300,
    "within_tolerance_only": -250,
    "story_explained": -100,
    "already_reviewed": 0,
    "waived": 0,
    "expected_growth": 0,
}

#: Signals that mean the analyst has no new decision to make. These defer a
#: group behind every open one regardless of severity, because a waived
#: critical is settled while an unreviewed style change is not.
_DEFERRAL_TIERS: dict[str, int] = {
    "already_reviewed": 1,
    "waived": 2,
    "expected_growth": 3,
}

_SIGNAL_PHRASES: dict[str, str] = {
    "severity_critical": "critical severity",
    "severity_warning": "warning severity",
    "severity_info": "informational severity",
    "material_delta": "{material} member(s) exceed the materiality threshold",
    "historical_change": "{historical} member(s) change historical periods",
    "new_defect": "{new} member(s) have no baseline counterpart",
    "downstream_impacts": "{impacts} recorded downstream impact(s)",
    "unexplained_residual": "no story explains this change",
    "wide_population": "{members} atomic findings in one decision",
    "inherited_only": "every member was already present in the baseline",
    "noise_only": "every numeric delta is display noise",
    "within_tolerance_only": "every numeric delta is within the accepted band",
    "story_explained": "explained by story {story}",
    "already_reviewed": "an analyst has already recorded a decision",
    "waived": "covered by an active waiver",
    "expected_growth": "expected new-cycle growth",
}

#: A decision holding at least this many atomics is worth surfacing early
#: because one judgement retires a large share of the queue.
_WIDE_POPULATION = 25


@dataclass(frozen=True, slots=True)
class ReviewPriority:
    """One review group placed in guided order, with its evidence cited."""

    group: ReviewGroup
    rank: int
    score: int
    signals: tuple[str, ...]
    rationale: str


def _priority_signals(
    group: ReviewGroup, story_by_finding: dict[str, str]
) -> tuple[list[str], dict[str, object]]:
    members = group.members
    counts: dict[str, object] = {
        "members": len(members),
        "material": sum(1 for item in members if item.materiality is Materiality.MATERIAL),
        "historical": sum(
            1
            for item in members
            if item.temporal_context is FindingTemporalContext.HISTORICAL
        ),
        "new": sum(1 for item in members if item.provenance is FindingProvenance.NEW),
        "impacts": sum(len(item.impacts) for item in members),
    }
    signals: list[str] = []
    if group.severity is Severity.CRITICAL:
        signals.append("severity_critical")
    elif group.severity is Severity.WARNING:
        signals.append("severity_warning")
    elif group.severity is Severity.INFO:
        signals.append("severity_info")

    if counts["material"]:
        signals.append("material_delta")
    if counts["historical"]:
        signals.append("historical_change")
    if counts["new"]:
        signals.append("new_defect")
    if counts["impacts"]:
        signals.append("downstream_impacts")
    if len(members) >= _WIDE_POPULATION:
        signals.append("wide_population")

    tiers = {item.materiality for item in members if item.materiality is not None}
    if tiers == {Materiality.NOISE}:
        signals.append("noise_only")
    elif tiers == {Materiality.WITHIN_TOLERANCE}:
        signals.append("within_tolerance_only")
    if members and all(
        item.provenance is FindingProvenance.INHERITED for item in members
    ):
        signals.append("inherited_only")

    story = next(
        (
            story_by_finding[item.finding_id]
            for item in members
            if item.finding_id in story_by_finding
        ),
        None,
    )
    if story is None:
        signals.append("unexplained_residual")
    else:
        signals.append("story_explained")
        counts["story"] = story

    if any(item.waiver_reason for item in members):
        signals.append("waived")
    if any(item.severity_overridden or item.analyst_comment for item in members):
        signals.append("already_reviewed")
    if group.expected_growth or group.severity is Severity.EXPECTED:
        signals.append("expected_growth")
    return signals, counts


def _rationale(signals: list[str], counts: dict[str, object]) -> str:
    return "; ".join(
        _SIGNAL_PHRASES[signal].format(**counts) for signal in signals
    ) or "no prioritization signal applies"


def prioritize_review(
    groups: list[ReviewGroup],
    stories: list[ChangeStory] | None = None,
) -> list[ReviewPriority]:
    """Order review decisions by evidence, without dropping or hiding any.

    The result is always a permutation of ``groups``: prioritization may order,
    explain, and de-emphasize, but every atomic finding stays reachable.
    """
    story_by_finding: dict[str, str] = {}
    for story in stories or ():
        if story.kind is StoryKind.RESIDUAL:
            continue
        for finding_id in story.finding_ids:
            story_by_finding.setdefault(finding_id, story.story_id)

    scored: list[
        tuple[int, int, tuple[object, ...], ReviewGroup, list[str], dict[str, object]]
    ] = []
    for group in groups:
        signals, counts = _priority_signals(group, story_by_finding)
        score = sum(_PRIORITY_WEIGHTS[signal] for signal in signals)
        tier = max((_DEFERRAL_TIERS.get(signal, 0) for signal in signals), default=0)
        scored.append((tier, score, _review_sort_key(group), group, signals, counts))

    scored.sort(key=lambda entry: (entry[0], -entry[1], entry[2]))
    return [
        ReviewPriority(
            group=group,
            rank=index,
            score=score,
            signals=tuple(signals),
            rationale=_rationale(signals, counts),
        )
        for index, (_tier, score, _order, group, signals, counts) in enumerate(
            scored, start=1
        )
    ]
