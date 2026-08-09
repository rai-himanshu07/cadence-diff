"""Structural-only fingerprints for sharing real deliverable shapes safely."""

import datetime as dt
import hashlib
import json
import re
from collections import Counter
from collections.abc import Mapping
from pathlib import Path

from openpyxl.utils.cell import range_boundaries

from qc_tool.crosscheck.numbers import extract_figures
from qc_tool.excel.formulas import to_r1c1
from qc_tool.excel.periods import parse_period
from qc_tool.excel.regions import detect_regions
from qc_tool.io.loader import load_workbook_snapshot
from qc_tool.io.model import CellRecord, WorkbookSnapshot, display_cell_value
from qc_tool.package import (
    PackageManifest,
    PackageMember,
    paths_by_member,
    structural_member_aliases,
)
from qc_tool.ppt.extract import DeckSnapshot, load_deck_snapshot
from qc_tool.security import private_directory, private_file

_QUOTED_STRING_RE = re.compile(r'"(?:[^"]|"")*"')
_MONTH_NAME_RE = re.compile(
    r"^(?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)[A-Za-z]*[-_ ]?'?\d{2,4}$",
    re.I,
)
_ISO_DATE_RE = re.compile(r"^\d{4}-\d{2}(?:-\d{2})?$")
_WEEK_RE = re.compile(r"^(?:\d{4}[-_ ]?)?(?:CW|WK|W)[-_ ]?\d{1,2}(?:[-_ ]?'?\d{2,4})?$", re.I)
_QUARTER_RE = re.compile(r"^Q[1-4](?:[-_ ]?(?:FY)?[-_ ]?'?\d{2,4})?$", re.I)


def _period_grammar(value: object) -> str | None:
    period = parse_period(value)
    if period is None:
        return None
    text = display_cell_value(value).strip()
    if _MONTH_NAME_RE.match(text):
        return "month-name-year"
    if _ISO_DATE_RE.match(text):
        return "iso-date" if text.count("-") == 2 else "iso-month"
    if _WEEK_RE.match(text):
        return "week-label"
    if _QUARTER_RE.match(text):
        return "quarter-label"
    return period.kind


def _cell_kind(cell: CellRecord) -> str:
    if cell.formula is not None:
        return "formula"
    if cell.is_error:
        return "error"
    if isinstance(cell.value, bool):
        return "boolean"
    if isinstance(cell.value, dt.datetime | dt.date):
        return "temporal"
    if isinstance(cell.value, int | float):
        return "numeric"
    if isinstance(cell.value, str):
        return "text"
    return "other"


def _range_shape(reference: str | None) -> dict[str, int] | None:
    if not reference:
        return None
    _, _, raw_range = reference.rpartition("!")
    try:
        min_col, min_row, max_col, max_row = range_boundaries(raw_range.replace("$", ""))
    except ValueError:
        return None
    if min_col is None or min_row is None or max_col is None or max_row is None:
        return None
    return {"rows": max_row - min_row + 1, "columns": max_col - min_col + 1}


def _formula_digest(
    formula: str,
    row: int,
    col: int,
    workbook: WorkbookSnapshot,
) -> str:
    canonical = to_r1c1(formula, row, col)
    canonical = _QUOTED_STRING_RE.sub('"TEXT"', canonical)
    replacements = [*workbook.sheet_names, *(item.name for item in workbook.named_ranges)]
    for index, value in enumerate(sorted(replacements, key=len, reverse=True), start=1):
        canonical = re.sub(
            rf"(?<![A-Za-z0-9_.])'?{re.escape(value)}'?(?![A-Za-z0-9_.])",
            f"ID{index}",
            canonical,
            flags=re.IGNORECASE,
        )
    return hashlib.sha256(f"cadence-diff/formula/v1|{canonical}".encode()).hexdigest()[:20]


def workbook_fingerprint(path: Path, *, password: str | None = None) -> dict:
    workbook = load_workbook_snapshot(path, password=password)
    sheets: list[dict] = []
    all_formula_patterns: Counter[str] = Counter()
    for sheet_index, sheet in enumerate(workbook.sheets, start=1):
        kinds = Counter(_cell_kind(cell) for cell in sheet.cells.values())
        periods = Counter(
            grammar
            for cell in sheet.cells.values()
            if (grammar := _period_grammar(cell.value)) is not None
        )
        formula_patterns = Counter(
            _formula_digest(cell.formula, row, col, workbook)
            for (row, col), cell in sheet.cells.items()
            if cell.formula is not None
        )
        all_formula_patterns.update(formula_patterns)
        regions = [
            {
                "orientation": region.orientation,
                "period_axis": region.period_axis,
                "rows": region.max_row - region.min_row + 1,
                "columns": region.max_col - region.min_col + 1,
                "start_row": region.min_row,
                "start_column": region.min_col,
            }
            for region in detect_regions(sheet)
        ]
        sheets.append(
            {
                "sheet": f"sheet_{sheet_index:03d}",
                "visibility": sheet.visibility,
                "rows": sheet.max_row,
                "columns": sheet.max_column,
                "populated_cells": len(sheet.cells),
                "cell_types": dict(sorted(kinds.items())),
                "period_grammars": dict(sorted(periods.items())),
                "hidden_rows": len(sheet.hidden_rows),
                "hidden_columns": len(sheet.hidden_columns),
                "regions": regions,
                "formula_patterns": [
                    {"digest": digest, "count": count}
                    for digest, count in sorted(formula_patterns.items())
                ],
            }
        )
    charts = [
        {
            "chart_type": chart.chart_type,
            "plot_types": [plot.chart_type for plot in chart.plots],
            "plot_axis_groups": [plot.axis_group for plot in chart.plots],
            "series_count": len(chart.series),
            "series": [
                {
                    "values_shape": _range_shape(series.values_ref),
                    "categories_shape": _range_shape(series.categories_ref),
                }
                for series in chart.series
            ],
            "axis_count": len(chart.axes),
            "axis_types": [axis.axis_type for axis in chart.axes],
            "legend": chart.legend is not None,
            "data_label_scopes": sum(
                plot.data_labels is not None for plot in chart.plots
            )
            + sum(series.data_labels is not None for series in chart.series),
            "anchor_type": chart.anchor.anchor_type if chart.anchor else None,
        }
        for chart in workbook.charts
    ]
    pivots = [
        {
            "location_shape": _range_shape(pivot.location_ref),
            "source_shape": _range_shape(pivot.source_ref),
        }
        for pivot in workbook.pivots
    ]
    tables = [
        {
            "range_shape": _range_shape(table.cell_range),
            "column_count": len(table.columns),
            "header_rows": table.header_row_count,
            "totals_rows": table.totals_row_count,
        }
        for table in workbook.tables
    ]
    data_validations = [
        {
            "type": rule.validation_type,
            "operator": rule.operator,
            "target_range_count": len(rule.target_ranges),
            "has_formula1": rule.formula1 is not None,
            "has_formula2": rule.formula2 is not None,
            "allow_blank": rule.allow_blank,
            "has_prompt": bool(rule.prompt or rule.prompt_title),
            "has_error_message": bool(rule.error or rule.error_title),
        }
        for rule in workbook.data_validations
    ]
    conditional_formats = [
        {
            "rule_type": rule.rule_type,
            "operator": rule.operator,
            "target_range_count": len(rule.target_ranges),
            "formula_count": len(rule.formulas),
            "priority": rule.priority,
            "stop_if_true": rule.stop_if_true,
            "semantic_supported": rule.semantic_supported,
            "style_supported": rule.style_supported,
            "has_stable_style": rule.style_key is not None,
        }
        for rule in workbook.conditional_formats
    ]
    return {
        "schema_version": 1,
        "schema": "https://cadence-diff.local/schema/fingerprint-v1.json",
        "artifact": "excel",
        "format": workbook.file_format,
        "capabilities": {
            "formulas": workbook.formulas_available,
            "styles": workbook.styles_available,
            "tables": workbook.tables_available,
            "charts": workbook.charts_available,
            "interaction_rules": workbook.interaction_rules_available,
            "interaction_semantics": workbook.interaction_rules_supported,
            "conditional_format_styles": (
                workbook.conditional_format_styles_supported
            ),
        },
        "sheet_count": len(workbook.sheets),
        "named_range_count": len(workbook.named_ranges),
        "chart_count": len(workbook.charts),
        "pivot_count": len(workbook.pivots),
        "table_count": len(workbook.tables),
        "data_validation_count": len(workbook.data_validations),
        "conditional_format_count": len(workbook.conditional_formats),
        "sheets": sheets,
        "charts": charts,
        "pivots": pivots,
        "tables": tables,
        "data_validations": data_validations,
        "conditional_formats": conditional_formats,
        "formula_patterns": [
            {"digest": digest, "count": count}
            for digest, count in sorted(all_formula_patterns.items())
        ],
    }


def _length_bucket(length: int) -> str:
    if length == 0:
        return "0"
    if length <= 20:
        return "1-20"
    if length <= 80:
        return "21-80"
    if length <= 200:
        return "81-200"
    return "201+"


def deck_fingerprint(path: Path, *, password: str | None = None) -> dict:
    deck: DeckSnapshot = load_deck_snapshot(path, password=password)
    slides: list[dict] = []
    deck_periods: Counter[str] = Counter()
    for slide_index, slide in enumerate(deck.slides, start=1):
        visible_text = [slide.title or "", *slide.texts]
        visible_text.extend(
            cell for table in slide.tables for row in table.rows for cell in row
        )
        periods = Counter(
            grammar
            for text in visible_text
            if (grammar := _period_grammar(text)) is not None
        )
        for chart in slide.charts:
            chart_grammars: list[str] = []
            categories = (
                [
                    category
                    for series in chart.all_series
                    for category in series.categories
                ]
                if chart.all_series
                else chart.categories
            )
            for category in categories:
                grammar = _period_grammar(category)
                if grammar is not None:
                    chart_grammars.append(grammar)
            chart_periods: Counter[str] = Counter(chart_grammars)
            periods.update(chart_periods)
        deck_periods.update(periods)
        slides.append(
            {
                "slide": f"slide_{slide_index:03d}",
                "has_title": slide.title is not None,
                "shape_count": slide.shape_count,
                "shape_types": dict(
                    sorted(Counter(shape.shape_type for shape in slide.shapes).items())
                ),
                "has_notes": bool(slide.notes),
                "note_paragraph_count": len(slide.notes),
                "text_block_count": len(slide.texts),
                "text_length_buckets": dict(
                    sorted(Counter(_length_bucket(len(text)) for text in slide.texts).items())
                ),
                "figure_count": sum(len(extract_figures(text)) for text in visible_text),
                "period_grammars": dict(sorted(periods.items())),
                "tables": [
                    {
                        "rows": len(table.rows),
                        "columns": max((len(row) for row in table.rows), default=0),
                    }
                    for table in slide.tables
                ],
                "charts": [
                    {
                        "chart_type": chart.chart_type,
                        "plot_types": [plot.chart_type for plot in chart.plots],
                        "plot_axis_groups": [plot.axis_group for plot in chart.plots],
                        "category_count": sum(
                            len(series.categories) for series in chart.all_series
                        ),
                        "series_count": len(chart.all_series),
                        "series_lengths": [
                            len(series.values) for series in chart.all_series
                        ],
                        "axis_count": len(chart.axes),
                        "axis_types": [axis.axis_type for axis in chart.axes],
                        "legend": chart.legend is not None,
                        "visible_label_count": sum(
                            len(plot.visible_label_values) for plot in chart.plots
                        ),
                    }
                    for chart in slide.charts
                ],
            }
        )
    return {
        "schema_version": 1,
        "schema": "https://cadence-diff.local/schema/fingerprint-v1.json",
        "artifact": "powerpoint",
        "format": "pptx",
        "capabilities": {
            "charts": deck.charts_available,
            "notes": deck.notes_available,
        },
        "slide_count": len(deck.slides),
        "period_grammars": dict(sorted(deck_periods.items())),
        "slides": slides,
    }


def fingerprint_file(path: Path, *, password: str | None = None) -> dict:
    suffix = path.suffix.casefold()
    if suffix in {".xlsx", ".xlsm", ".xlsb"}:
        payload = workbook_fingerprint(path, password=password)
    elif suffix == ".pptx":
        payload = deck_fingerprint(path, password=password)
    else:
        raise ValueError(f"{path.name}: fingerprint supports .xlsx/.xlsm/.xlsb/.pptx")
    canonical = json.dumps(payload, sort_keys=True, separators=(",", ":"))
    return {
        **payload,
        "fingerprint_id": hashlib.sha256(canonical.encode()).hexdigest(),
    }


def package_fingerprint(
    package_files: Mapping[str, Path],
    package_manifest: PackageManifest,
    *,
    passwords: Mapping[str, str] | None = None,
) -> dict:
    """Fingerprint an exact package topology without source member identifiers."""
    paths_by_member(package_files, package_manifest)
    aliases = structural_member_aliases(package_manifest)

    redacted_members: list[PackageMember] = []
    fingerprints: list[dict[str, object]] = []
    credentials = passwords or {}
    for member in package_manifest.members:
        alias = aliases[(member.artifact, member.member_id)]
        redacted = PackageMember(
            member_id=alias,
            side=member.side,
            artifact=member.artifact,
            display_name=f"{alias}.office",
        )
        redacted_members.append(redacted)
        fingerprints.append(
            {
                "role": redacted.role_key,
                "fingerprint": fingerprint_file(
                    package_files[member.role_key],
                    password=credentials.get(member.role_key),
                ),
            }
        )
    redacted_manifest = PackageManifest(members=tuple(redacted_members))
    payload = {
        "schema_version": 2,
        "schema": "https://cadence-diff.local/schema/package-fingerprint-v2.json",
        "artifact": "package",
        "format": "package",
        "package_manifest": redacted_manifest.model_dump(mode="json"),
        "members": fingerprints,
    }
    canonical = json.dumps(payload, sort_keys=True, separators=(",", ":"))
    return {
        **payload,
        "fingerprint_id": hashlib.sha256(canonical.encode()).hexdigest(),
    }


def write_fingerprint(path: Path, output: Path, *, password: str | None = None) -> dict:
    payload = fingerprint_file(path, password=password)
    private_directory(output.parent)
    output.write_text(json.dumps(payload, indent=2), encoding="utf-8")
    private_file(output)
    return payload


def write_package_fingerprint(
    package_files: Mapping[str, Path],
    package_manifest: PackageManifest,
    output: Path,
    *,
    passwords: Mapping[str, str] | None = None,
) -> dict:
    payload = package_fingerprint(
        package_files,
        package_manifest,
        passwords=passwords,
    )
    private_directory(output.parent)
    output.write_text(json.dumps(payload, indent=2), encoding="utf-8")
    private_file(output)
    return payload
