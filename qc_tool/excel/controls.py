"""Profile-declared current-workbook controls and arithmetic tie-outs."""

from dataclasses import dataclass, field

from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import range_boundaries

from qc_tool.config.profile import ExcelControls
from qc_tool.coverage import CoverageItem, CoverageState
from qc_tool.excel.references import ReferenceResolution, ReferenceStatus, resolve_reference
from qc_tool.findings import Finding, FindingClass
from qc_tool.io.model import (
    CellRecord,
    SheetSnapshot,
    WorkbookSnapshot,
    display_cell_value,
)


@dataclass(slots=True)
class ControlResult:
    findings: list[Finding] = field(default_factory=list)
    coverage: CoverageItem | None = None


def _cell_ref(row: int, col: int) -> str:
    return f"{get_column_letter(col)}{row}"


def _resolve_range(
    workbook: WorkbookSnapshot, sheet_name: str, cell_range: str
) -> tuple[SheetSnapshot, tuple[int, int, int, int]] | None:
    try:
        sheet = workbook.sheet(sheet_name)
        min_col, min_row, max_col, max_row = range_boundaries(cell_range)
    except (KeyError, ValueError):
        return None
    if min_col is None or min_row is None or max_col is None or max_row is None:
        return None
    return sheet, (min_col, min_row, max_col, max_row)


def _control_invalid(name: str, target: str) -> Finding:
    return Finding(
        artifact="excel",
        finding_class=FindingClass.CONTROL_INVALID,
        element=name or target,
        current_value=target,
        message=f"profile control {name or target!r} has an invalid target {target!r}",
    )


def _is_blank(cell: CellRecord | None) -> bool:
    return cell is None or cell.value is None or (
        isinstance(cell.value, str) and not cell.value.strip()
    )


def _numeric_values(
    workbook: WorkbookSnapshot,
    resolution: ReferenceResolution,
) -> list[float] | None:
    values: list[float] = []
    for resolved in resolution.ranges:
        sheet = workbook.sheet(resolved.sheet)
        for row in range(resolved.min_row, resolved.max_row + 1):
            for col in range(resolved.min_col, resolved.max_col + 1):
                cell = sheet.cells.get((row, col))
                value = None if cell is None else cell.value
                if isinstance(value, bool) or not isinstance(value, int | float):
                    return None
                values.append(float(value))
    return values


def _resolve_control_reference(
    workbook: WorkbookSnapshot,
    reference: str,
) -> ReferenceResolution:
    sheet_name, separator, _target = reference.rpartition("!")
    host_sheet = (
        sheet_name.strip("'").replace("''", "'")
        if separator
        else (workbook.sheets[0].name if workbook.sheets else "")
    )
    return resolve_reference(
        workbook,
        reference,
        host_sheet=host_sheet,
        require_within_sheet=True,
    )


def evaluate_controls(
    workbook: WorkbookSnapshot,
    controls: ExcelControls,
    *,
    ignored_sheets: set[str] | None = None,
) -> ControlResult:
    result = ControlResult()
    ignored = ignored_sheets or set()
    required_ranges = [
        control for control in controls.required_ranges if control.sheet not in ignored
    ]
    unique_ranges = [
        control for control in controls.unique_ranges if control.sheet not in ignored
    ]
    numeric_bounds = [
        control for control in controls.numeric_bounds if control.sheet not in ignored
    ]
    configured = (
        len(required_ranges)
        + len(unique_ranges)
        + len(numeric_bounds)
    )

    for control in required_ranges:
        resolved = _resolve_range(workbook, control.sheet, control.cell_range)
        if resolved is None:
            result.findings.append(
                _control_invalid(control.name, f"{control.sheet}!{control.cell_range}")
            )
            continue
        sheet, (min_col, min_row, max_col, max_row) = resolved
        for row in range(min_row, max_row + 1):
            for col in range(min_col, max_col + 1):
                if not _is_blank(sheet.cells.get((row, col))):
                    continue
                location = _cell_ref(row, col)
                result.findings.append(
                    Finding(
                        artifact="excel",
                        finding_class=FindingClass.REQUIRED_VALUE_MISSING,
                        sheet=sheet.name,
                        location=location,
                        element=control.name or control.cell_range,
                        message=f"{sheet.name}!{location}: required value is blank",
                    )
                )

    for control in unique_ranges:
        resolved = _resolve_range(workbook, control.sheet, control.cell_range)
        if resolved is None:
            result.findings.append(
                _control_invalid(control.name, f"{control.sheet}!{control.cell_range}")
            )
            continue
        sheet, (min_col, min_row, max_col, max_row) = resolved
        first_row = min_row + 1 if control.skip_header else min_row
        seen: dict[tuple[object, ...], int] = {}
        for row in range(first_row, max_row + 1):
            row_cells = [
                sheet.cells.get((row, col)) for col in range(min_col, max_col + 1)
            ]
            key = tuple(cell.value if cell is not None else None for cell in row_cells)
            if all(value is None or value == "" for value in key):
                continue
            if key not in seen:
                seen[key] = row
                continue
            location = (
                f"{_cell_ref(row, min_col)}:{_cell_ref(row, max_col)}"
                if min_col != max_col
                else _cell_ref(row, min_col)
            )
            result.findings.append(
                Finding(
                    artifact="excel",
                    finding_class=FindingClass.DUPLICATE_KEY,
                    sheet=sheet.name,
                    location=location,
                    baseline_location=f"row {seen[key]}",
                    current_value=" | ".join(display_cell_value(value) for value in key),
                    element=control.name or control.cell_range,
                    message=f"{sheet.name}!{location}: duplicate key",
                )
            )

    for control in numeric_bounds:
        resolved = _resolve_range(workbook, control.sheet, control.cell_range)
        if resolved is None:
            result.findings.append(
                _control_invalid(control.name, f"{control.sheet}!{control.cell_range}")
            )
            continue
        sheet, (min_col, min_row, max_col, max_row) = resolved
        for row in range(min_row, max_row + 1):
            for col in range(min_col, max_col + 1):
                cell = sheet.cells.get((row, col))
                value = None if cell is None else cell.value
                if isinstance(value, bool) or not isinstance(value, int | float):
                    continue
                below = control.minimum is not None and value < control.minimum
                above = control.maximum is not None and value > control.maximum
                if not below and not above:
                    continue
                location = _cell_ref(row, col)
                result.findings.append(
                    Finding(
                        artifact="excel",
                        finding_class=FindingClass.NUMERIC_BOUND_VIOLATION,
                        sheet=sheet.name,
                        location=location,
                        current_value=display_cell_value(value),
                        element=control.name or control.cell_range,
                        message=(
                            f"{sheet.name}!{location}: value {value} is outside "
                            f"configured bounds [{control.minimum}, {control.maximum}]"
                        ),
                    )
                )

    for control in controls.tie_outs:
        uses_components = bool(control.components)
        uses_terms = bool(control.terms)
        if uses_components == uses_terms:
            configured += 1
            result.findings.append(_control_invalid(control.name, control.target))
            continue
        terms = (
            [(component, 1.0) for component in control.components]
            if uses_components
            else [
                (term.reference, 1.0 if term.operation == "add" else -1.0)
                for term in control.terms
            ]
        )
        target_resolution = _resolve_control_reference(workbook, control.target)
        term_resolutions = [
            (_resolve_control_reference(workbook, reference), sign)
            for reference, sign in terms
        ]
        resolutions = [target_resolution, *[item for item, _sign in term_resolutions]]
        if any(item.status is not ReferenceStatus.RESOLVED for item in resolutions):
            configured += 1
            result.findings.append(_control_invalid(control.name, control.target))
            continue
        if any(
            resolved.sheet in ignored
            for resolution in resolutions
            for resolved in resolution.ranges
        ):
            continue
        configured += 1
        target_values = _numeric_values(workbook, target_resolution)
        term_values = [
            (_numeric_values(workbook, resolution), sign)
            for resolution, sign in term_resolutions
        ]
        if target_values is None or len(target_values) != 1 or any(
            values is None for values, _sign in term_values
        ):
            result.findings.append(_control_invalid(control.name, control.target))
            continue
        target = target_values[0]
        expected = sum(
            sign * value
            for values, sign in term_values
            if values is not None
            for value in values
        )
        delta = abs(target - expected)
        within_absolute = delta <= control.absolute_tolerance
        within_relative = (
            expected != 0 and delta / abs(expected) <= control.relative_tolerance
        )
        if within_absolute or within_relative:
            continue
        target_cell = target_resolution.ranges[0]
        location = _cell_ref(target_cell.min_row, target_cell.min_col)
        result.findings.append(
            Finding(
                artifact="excel",
                finding_class=FindingClass.TIE_OUT_MISMATCH,
                sheet=target_cell.sheet,
                location=location,
                element=control.name,
                baseline_value=str(expected),
                current_value=str(target),
                message=(
                    f"{control.name}: target {target} does not tie to component "
                    f"sum {expected}"
                ),
            )
        )

    result.coverage = CoverageItem(
        check_id="excel-profile-controls",
        label="Profile-required values, uniqueness, bounds, and tie-outs",
        artifact="excel",
        state=CoverageState.CHECKED,
        findings=len(result.findings),
        detail=f"{configured} controls configured",
    )
    return result
