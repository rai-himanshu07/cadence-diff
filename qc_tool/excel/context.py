"""Cell-neighborhood excerpts: let analysts verify findings without Excel.

For every cell-level finding, capture a small grid around the affected
cell from both snapshots at run time. Excerpts persist with the finding
(and therefore with run history), so past runs stay inspectable even if
the source files are later replaced.
"""

import contextlib
import re

from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import coordinate_to_tuple

from qc_tool.findings import Finding, FindingClass, GridExcerpt
from qc_tool.io.model import SheetSnapshot, WorkbookSnapshot, display_cell_value

_A1_RE = re.compile(r"^[A-Z]{1,3}\d+$")

#: Classes where seeing the surrounding cells answers "is this real?".
_CONTEXT_CLASSES = frozenset(
    {
        FindingClass.VALUE_CHANGED,
        FindingClass.FORMULA_ERROR,
        FindingClass.FORMULA_HARDCODED,
        FindingClass.FORMULA_REMOVED,
        FindingClass.FORMULA_MISSING,
        FindingClass.FORMULA_NOT_EXTENDED,
        FindingClass.FORMULA_LOGIC_CHANGED,
        FindingClass.FORMULA_INCONSISTENT,
        FindingClass.NUMBER_FORMAT_CHANGED,
        FindingClass.STYLE_CHANGED,
        FindingClass.REQUIRED_VALUE_MISSING,
        FindingClass.NUMERIC_BOUND_VIOLATION,
        FindingClass.TIE_OUT_MISMATCH,
    }
)

_ROW_RADIUS = 3
_COL_RADIUS = 3


def _display(sheet: SheetSnapshot, row: int, col: int) -> str:
    cell = sheet.cells.get((row, col))
    if cell is None:
        return ""
    if cell.value is None and cell.formula is not None:
        return cell.formula
    return display_cell_value(cell.value)


def build_excerpt(sheet: SheetSnapshot, ref: str) -> GridExcerpt | None:
    """A (2*radius+1)-sized neighborhood of ``ref``, clamped to the sheet."""
    if not _A1_RE.match(ref):
        return None
    row, col = coordinate_to_tuple(ref)
    min_row = max(1, row - _ROW_RADIUS)
    max_row = min(max(sheet.max_row, row), row + _ROW_RADIUS)
    min_col = max(1, col - _COL_RADIUS)
    max_col = min(max(sheet.max_column, col), col + _COL_RADIUS)
    rows = list(range(min_row, max_row + 1))
    cols = list(range(min_col, max_col + 1))
    return GridExcerpt(
        cols=[get_column_letter(c) for c in cols],
        rows=rows,
        cells=[[_display(sheet, r, c) for c in cols] for r in rows],
        hit_row=rows.index(row),
        hit_col=cols.index(col),
    )


def attach_excerpts(
    findings: list[Finding], baseline: WorkbookSnapshot, current: WorkbookSnapshot
) -> None:
    """Attach baseline/current excerpts to cell-level findings (in place)."""
    for finding in findings:
        if finding.finding_class not in _CONTEXT_CLASSES or finding.sheet is None:
            continue
        if finding.location:
            with contextlib.suppress(KeyError):  # sheet may not exist on this side
                finding.current_excerpt = build_excerpt(
                    current.sheet(finding.sheet), finding.location
                )
        baseline_ref = finding.baseline_location or finding.location
        if baseline_ref:
            with contextlib.suppress(KeyError):
                finding.baseline_excerpt = build_excerpt(
                    baseline.sheet(finding.sheet), baseline_ref
                )


def attach_current_excerpts(
    findings: list[Finding], current: WorkbookSnapshot
) -> None:
    """Attach only current-side context for standalone/package findings."""
    for finding in findings:
        if (
            finding.finding_class not in _CONTEXT_CLASSES
            or finding.sheet is None
            or finding.location is None
        ):
            continue
        with contextlib.suppress(KeyError):
            finding.current_excerpt = build_excerpt(
                current.sheet(finding.sheet), finding.location
            )
