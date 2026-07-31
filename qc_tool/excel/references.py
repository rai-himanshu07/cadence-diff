"""Bounded Excel A1 and structured-reference resolution."""

from __future__ import annotations

import re
from dataclasses import dataclass
from enum import StrEnum

from openpyxl.utils.cell import range_boundaries

from qc_tool.excel.formula_tokens import parse_dynamic_reference
from qc_tool.io.model import TableDescriptor, WorkbookSnapshot


class ReferenceStatus(StrEnum):
    RESOLVED = "resolved"
    INVALID = "invalid"
    UNSUPPORTED = "unsupported"


@dataclass(frozen=True, slots=True)
class ResolvedRange:
    sheet: str
    min_row: int
    min_col: int
    max_row: int
    max_col: int

    @property
    def size(self) -> int:
        return (self.max_row - self.min_row + 1) * (self.max_col - self.min_col + 1)


@dataclass(frozen=True, slots=True)
class ReferenceResolution:
    status: ReferenceStatus
    ranges: tuple[ResolvedRange, ...] = ()
    detail: str = ""

    @property
    def size(self) -> int:
        return sum(item.size for item in self.ranges)


@dataclass(frozen=True, slots=True)
class _StructuredSpec:
    table_name: str
    row_selectors: tuple[str, ...]
    columns: tuple[str, ...]
    column_span: bool = False


_ESCAPED_SPECIAL_RE = re.compile(r"'([\[\]#'@])")
_ROW_SELECTORS = {
    "#all": "all",
    "#data": "data",
    "#headers": "headers",
    "#totals": "totals",
    "#this row": "this_row",
}
_RANGE_EXTENSION_RE = re.compile(
    r"^(?P<prefix>.*!\$?[A-Z]{1,3}\$?\d+:\$?)(?P<endcol>[A-Z]{1,3})\$?(?P<endrow>\d+)$"
)


def _failure(status: ReferenceStatus, detail: str) -> ReferenceResolution:
    return ReferenceResolution(status=status, detail=detail)


def _unescape_name(value: str) -> str:
    return _ESCAPED_SPECIAL_RE.sub(r"\1", value.replace("]]", "]"))


def _nested_groups(value: str) -> tuple[list[str], list[str]] | None:
    groups: list[str] = []
    separators: list[str] = []
    position = 0
    while position < len(value):
        if value[position] != "[":
            return None
        position += 1
        text: list[str] = []
        while position < len(value):
            character = value[position]
            if character == "]":
                if position + 1 < len(value) and value[position + 1] == "]":
                    text.append("]")
                    position += 2
                    continue
                position += 1
                break
            text.append(character)
            position += 1
        else:
            return None
        groups.append("".join(text))
        if position == len(value):
            break
        separator = value[position]
        if separator not in {",", ":"}:
            return None
        separators.append(separator)
        position += 1
    if len(separators) != max(0, len(groups) - 1):
        return None
    return groups, separators


def _parse_structured_reference(
    target: str,
) -> _StructuredSpec | ReferenceResolution:
    bracket = target.find("[")
    if bracket < 0 or not target.endswith("]"):
        return _failure(ReferenceStatus.INVALID, "malformed structured reference")
    table_name = target[:bracket]
    inner = target[bracket + 1 : -1]
    if not inner:
        return _failure(ReferenceStatus.INVALID, "empty structured selector")

    if inner.startswith("["):
        parsed = _nested_groups(inner)
        if parsed is None:
            return _failure(ReferenceStatus.INVALID, "malformed nested selector")
        groups, separators = parsed
    else:
        if "[" in inner:
            return _failure(ReferenceStatus.INVALID, "malformed structured selector")
        groups = [inner]
        separators = []

    row_selectors: list[str] = []
    columns: list[str] = []
    for group in groups:
        stripped = group.strip()
        normalized = stripped.casefold()
        if normalized.startswith("@"):
            row_selectors.append("this_row")
            column = _unescape_name(stripped[1:])
            if column:
                columns.append(column)
            continue
        if normalized.startswith("#"):
            selector = _ROW_SELECTORS.get(normalized)
            if selector is None:
                return _failure(ReferenceStatus.INVALID, f"unknown selector {group!r}")
            row_selectors.append(selector)
            continue
        if not group:
            return _failure(ReferenceStatus.INVALID, "empty column selector")
        columns.append(_unescape_name(group))

    if ":" in separators:
        if separators != [":"] or row_selectors or len(columns) != 2:
            return _failure(
                ReferenceStatus.UNSUPPORTED,
                "mixed or multi-part structured column span",
            )
        return _StructuredSpec(table_name, (), tuple(columns), column_span=True)
    if len(columns) > 1:
        return _failure(
            ReferenceStatus.UNSUPPORTED,
            "non-contiguous structured column union",
        )
    if len(set(row_selectors)) != len(row_selectors):
        return _failure(ReferenceStatus.INVALID, "duplicate row selector")
    if "all" in row_selectors and len(row_selectors) > 1:
        return _failure(ReferenceStatus.INVALID, "#All cannot be combined")
    if "this_row" in row_selectors and len(row_selectors) > 1:
        return _failure(ReferenceStatus.INVALID, "#This Row cannot be combined")
    return _StructuredSpec(table_name, tuple(row_selectors), tuple(columns))


def _table_bounds(table: TableDescriptor) -> tuple[int, int, int, int] | None:
    try:
        min_col, min_row, max_col, max_row = range_boundaries(table.cell_range)
    except ValueError:
        return None
    if min_col is None or min_row is None or max_col is None or max_row is None:
        return None
    return min_col, min_row, max_col, max_row


def _containing_table(
    workbook: WorkbookSnapshot,
    host_sheet: str,
    host_cell: tuple[int, int] | None,
) -> TableDescriptor | ReferenceResolution:
    if host_cell is None:
        return _failure(
            ReferenceStatus.UNSUPPORTED,
            "current-row reference requires a host cell",
        )
    host_row, host_col = host_cell
    matches: list[TableDescriptor] = []
    for table in workbook.tables:
        bounds = _table_bounds(table)
        if bounds is None or table.sheet != host_sheet:
            continue
        min_col, min_row, max_col, max_row = bounds
        if min_row <= host_row <= max_row and min_col <= host_col <= max_col:
            matches.append(table)
    if len(matches) != 1:
        return _failure(
            ReferenceStatus.INVALID,
            "host cell does not identify one containing table",
        )
    return matches[0]


def _find_table(
    workbook: WorkbookSnapshot,
    table_name: str,
    host_sheet: str,
    host_cell: tuple[int, int] | None,
) -> TableDescriptor | ReferenceResolution:
    if not table_name:
        return _containing_table(workbook, host_sheet, host_cell)
    normalized = table_name.casefold()
    matches = [
        table
        for table in workbook.tables
        if normalized in {table.name.casefold(), table.display_name.casefold()}
    ]
    if len(matches) != 1:
        return _failure(
            ReferenceStatus.INVALID,
            f"table {table_name!r} was not found uniquely",
        )
    return matches[0]


def _structured_ranges(
    workbook: WorkbookSnapshot,
    spec: _StructuredSpec,
    host_sheet: str,
    host_cell: tuple[int, int] | None,
) -> ReferenceResolution:
    table = _find_table(workbook, spec.table_name, host_sheet, host_cell)
    if isinstance(table, ReferenceResolution):
        return table
    bounds = _table_bounds(table)
    if bounds is None:
        return _failure(ReferenceStatus.INVALID, "table range is malformed")
    min_col, min_row, max_col, max_row = bounds
    if len(table.columns) != max_col - min_col + 1:
        return _failure(
            ReferenceStatus.UNSUPPORTED,
            "table column schema does not match its range",
        )

    column_lookup = {name.casefold(): index for index, name in enumerate(table.columns)}
    if spec.columns:
        indexes: list[int] = []
        for name in spec.columns:
            index = column_lookup.get(name.casefold())
            if index is None:
                return _failure(
                    ReferenceStatus.INVALID,
                    f"table column {name!r} was not found",
                )
            indexes.append(index)
        if spec.column_span and indexes[0] > indexes[1]:
            return _failure(ReferenceStatus.INVALID, "reversed table column span")
        selected_min_col = min_col + indexes[0]
        selected_max_col = min_col + indexes[-1]
    else:
        selected_min_col = min_col
        selected_max_col = max_col

    data_min_row = min_row + table.header_row_count
    data_max_row = max_row - table.totals_row_count
    selectors = spec.row_selectors or ("data",)
    row_ranges: list[tuple[int, int]] = []
    for selector in selectors:
        if selector == "all":
            row_ranges.append((min_row, max_row))
        elif selector == "headers":
            if table.header_row_count <= 0:
                return _failure(ReferenceStatus.INVALID, "table has no header rows")
            row_ranges.append((min_row, data_min_row - 1))
        elif selector == "totals":
            if table.totals_row_count <= 0:
                return _failure(ReferenceStatus.INVALID, "table has no totals rows")
            row_ranges.append((data_max_row + 1, max_row))
        elif selector == "data":
            if data_min_row > data_max_row:
                return _failure(ReferenceStatus.INVALID, "table has no data rows")
            row_ranges.append((data_min_row, data_max_row))
        elif selector == "this_row":
            if host_cell is None:
                return _failure(
                    ReferenceStatus.UNSUPPORTED,
                    "current-row reference requires a host cell",
                )
            host_row, _ = host_cell
            if not data_min_row <= host_row <= data_max_row:
                return _failure(
                    ReferenceStatus.INVALID,
                    "host cell is outside the table data rows",
                )
            row_ranges.append((host_row, host_row))

    return ReferenceResolution(
        status=ReferenceStatus.RESOLVED,
        ranges=tuple(
            ResolvedRange(
                sheet=table.sheet,
                min_row=start_row,
                min_col=selected_min_col,
                max_row=end_row,
                max_col=selected_max_col,
            )
            for start_row, end_row in row_ranges
        ),
    )


def _resolve_reference(
    workbook: WorkbookSnapshot,
    target: str,
    *,
    host_sheet: str,
    host_cell: tuple[int, int] | None,
    require_within_sheet: bool,
    seen_names: frozenset[str],
) -> ReferenceResolution:
    normalized_target = target.strip()
    if not normalized_target or "#REF!" in normalized_target.upper():
        return _failure(ReferenceStatus.INVALID, "empty or broken reference")
    dynamic = parse_dynamic_reference(normalized_target)
    if dynamic is not None:
        kind, inner = dynamic
        resolved = _resolve_reference(
            workbook,
            inner,
            host_sheet=host_sheet,
            host_cell=host_cell,
            require_within_sheet=require_within_sheet,
            seen_names=seen_names,
        )
        if resolved.status is not ReferenceStatus.RESOLVED:
            return resolved
        if len(resolved.ranges) != 1:
            return _failure(
                ReferenceStatus.UNSUPPORTED,
                f"{kind} reference resolves to multiple ranges",
            )
        source = resolved.ranges[0]
        if kind == "spill":
            if source.size != 1:
                return _failure(
                    ReferenceStatus.UNSUPPORTED,
                    "spill anchor must resolve to one cell",
                )
            matches = [
                descriptor
                for descriptor in workbook.formula_ranges
                if descriptor.sheet == source.sheet
                and descriptor.anchor_row == source.min_row
                and descriptor.anchor_column == source.min_col
                and descriptor.formula_type.casefold() == "array"
            ]
            if len(matches) != 1:
                return _failure(
                    ReferenceStatus.UNSUPPORTED,
                    "dynamic-array spill extent is unavailable",
                )
            descriptor = matches[0]
            try:
                min_col, min_row, max_col, max_row = range_boundaries(
                    descriptor.cell_range
                )
            except ValueError:
                return _failure(
                    ReferenceStatus.INVALID,
                    "declared spill extent is malformed",
                )
            if (
                min_col is None
                or min_row is None
                or max_col is None
                or max_row is None
            ):
                return _failure(
                    ReferenceStatus.INVALID,
                    "declared spill extent is unbounded",
                )
            return ReferenceResolution(
                status=ReferenceStatus.RESOLVED,
                ranges=(
                    ResolvedRange(
                        sheet=descriptor.sheet,
                        min_row=min_row,
                        min_col=min_col,
                        max_row=max_row,
                        max_col=max_col,
                    ),
                ),
                detail="dynamic-array spill extent declared by the anchor formula",
            )
        if source.size == 1:
            return resolved
        if host_cell is None:
            return _failure(
                ReferenceStatus.UNSUPPORTED,
                "implicit intersection requires a host cell",
            )
        host_row, host_column = host_cell
        if source.min_col == source.max_col and source.min_row <= host_row <= source.max_row:
            return ReferenceResolution(
                status=ReferenceStatus.RESOLVED,
                ranges=(
                    ResolvedRange(
                        source.sheet,
                        host_row,
                        source.min_col,
                        host_row,
                        source.max_col,
                    ),
                ),
                detail="implicit intersection projected by host row",
            )
        if source.min_row == source.max_row and source.min_col <= host_column <= source.max_col:
            return ReferenceResolution(
                status=ReferenceStatus.RESOLVED,
                ranges=(
                    ResolvedRange(
                        source.sheet,
                        source.min_row,
                        host_column,
                        source.max_row,
                        host_column,
                    ),
                ),
                detail="implicit intersection projected by host column",
            )
        return _failure(
            ReferenceStatus.UNSUPPORTED,
            "implicit intersection is ambiguous for this range and host cell",
        )
    if normalized_target.startswith("[") and "!" in normalized_target:
        return _failure(
            ReferenceStatus.UNSUPPORTED,
            "external workbook reference",
        )
    if "," in normalized_target and "[" not in normalized_target:
        return _failure(ReferenceStatus.UNSUPPORTED, "multi-area A1 reference")
    if "[" in normalized_target:
        parsed = _parse_structured_reference(normalized_target)
        if isinstance(parsed, ReferenceResolution):
            return parsed
        return _structured_ranges(workbook, parsed, host_sheet, host_cell)

    normalized_name = normalized_target.casefold()
    named = next(
        (item for item in workbook.named_ranges if item.name.casefold() == normalized_name),
        None,
    )
    if named is not None:
        if normalized_name in seen_names:
            return _failure(ReferenceStatus.INVALID, "cyclic named range")
        return _resolve_reference(
            workbook,
            named.target,
            host_sheet=host_sheet,
            host_cell=host_cell,
            require_within_sheet=require_within_sheet,
            seen_names=seen_names | {normalized_name},
        )

    sheet_part, separator, cell_range = normalized_target.rpartition("!")
    sheet_name = sheet_part.strip("'").replace("''", "'") if separator else host_sheet
    if not sheet_name:
        return _failure(ReferenceStatus.INVALID, "reference has no worksheet context")
    try:
        sheet = workbook.sheet(sheet_name)
        min_col, min_row, max_col, max_row = range_boundaries(
            cell_range.replace("$", "") if separator else normalized_target.replace("$", "")
        )
    except (KeyError, ValueError):
        return _failure(ReferenceStatus.INVALID, "invalid A1 reference")
    bounded = any(value is None for value in (min_col, min_row, max_col, max_row))
    if bounded and (sheet.max_row <= 0 or sheet.max_column <= 0):
        return _failure(
            ReferenceStatus.INVALID,
            "whole-row/column reference cannot be bounded on an empty sheet",
        )
    resolved_min_col = min_col if min_col is not None else 1
    resolved_min_row = min_row if min_row is not None else 1
    resolved_max_col = max_col if max_col is not None else max(sheet.max_column, 1)
    resolved_max_row = max_row if max_row is not None else max(sheet.max_row, 1)
    if require_within_sheet and (
        resolved_max_row > sheet.max_row or resolved_max_col > sheet.max_column
    ):
        return _failure(ReferenceStatus.INVALID, "reference exceeds worksheet bounds")
    return ReferenceResolution(
        status=ReferenceStatus.RESOLVED,
        ranges=(
            ResolvedRange(
                sheet=sheet_name,
                min_row=resolved_min_row,
                min_col=resolved_min_col,
                max_row=resolved_max_row,
                max_col=resolved_max_col,
            ),
        ),
        detail=(
            "bounded whole-row/column reference to populated sheet dimensions" if bounded else ""
        ),
    )


def resolve_reference(
    workbook: WorkbookSnapshot,
    target: str,
    *,
    host_sheet: str,
    host_cell: tuple[int, int] | None = None,
    require_within_sheet: bool = True,
) -> ReferenceResolution:
    """Resolve one A1, named, or supported structured reference."""
    return _resolve_reference(
        workbook,
        target,
        host_sheet=host_sheet,
        host_cell=host_cell,
        require_within_sheet=require_within_sheet,
        seen_names=frozenset(),
    )


def is_pure_range_extension(baseline: str, current: str) -> bool:
    """Whether one A1 range only grows at its trailing row or column edge."""
    baseline_match = _RANGE_EXTENSION_RE.match(baseline)
    current_match = _RANGE_EXTENSION_RE.match(current)
    if baseline_match is None or current_match is None:
        return False
    if baseline_match["prefix"] != current_match["prefix"]:
        return False
    if baseline_match["endcol"] == current_match["endcol"]:
        return int(current_match["endrow"]) >= int(baseline_match["endrow"])
    if baseline_match["endrow"] == current_match["endrow"]:
        baseline_column = baseline_match["endcol"]
        current_column = current_match["endcol"]
        return len(current_column) > len(baseline_column) or (
            len(current_column) == len(baseline_column) and current_column > baseline_column
        )
    return False
