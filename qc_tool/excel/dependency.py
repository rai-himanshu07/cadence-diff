"""Compact cross-sheet formula dependency index and downstream-impact annotation.

Nodes are cells ``(sheet, row, col)``. A precedent -> dependent relationship
means the dependent's formula references the precedent. Single-cell references
are stored as direct edges; range, table, named, and supported dynamic
references are stored as rectangular descriptors indexed by sheet and bounded
row interval, so a 365-cell range costs one descriptor instead of 365 edges.
Findings on a cell gain ``impacts``: every downstream formula cell that
transitively consumes it.

``tests/dependency_oracle.py`` keeps the historical per-cell NetworkX expansion
as an independent parity oracle for supported synthetic cases. It is test-only
and must never be imported by production code.
"""

import hashlib
import json
import logging
import re
from collections import Counter, OrderedDict
from dataclasses import dataclass

from openpyxl.utils import get_column_letter

from qc_tool.coverage import CoverageItem, CoverageState
from qc_tool.excel.formula_tokens import (
    FormulaPatternKey,
    FormulaPrecedentExtraction,
    FormulaPrecedentKind,
    extract_formula_precedents,
    formula_pattern_key,
)
from qc_tool.excel.references import (
    ReferenceStatus,
    ResolvedRange,
    build_reference_index,
    reference_reason_code,
    resolve_reference,
)
from qc_tool.findings import Finding, FindingClass
from qc_tool.io.model import WorkbookSnapshot
from qc_tool.progress import CancellationToken, check_cancelled

logger = logging.getLogger(__name__)

_CELL_REF_RE = re.compile(r"^[A-Z]{1,3}\d+$")
_MAX_RANGE_CELLS = 50_000
_MAX_FAILURE_SAMPLES = 8
_ROW_BUCKET = 64
_COLUMN_BUCKET = 16
_MAX_CYCLE_EDGES = 2_000_000
_CANDIDATE_CACHE_ENTRY_CAP = 131_072
_CANDIDATE_CACHE_MEMBERSHIP_CAP = 2_000_000
_ANNOTATED_CLASSES = frozenset(
    {
        FindingClass.VALUE_CHANGED,
        FindingClass.FORMULA_ERROR,
        FindingClass.FORMULA_HARDCODED,
        FindingClass.FORMULA_REMOVED,
        FindingClass.FORMULA_MISSING,
        FindingClass.FORMULA_NOT_EXTENDED,
        FindingClass.FORMULA_LOGIC_CHANGED,
        FindingClass.FORMULA_INCONSISTENT,
    }
)

Node = tuple[str, int, int]


@dataclass(frozen=True, slots=True)
class SymbolicReference:
    source: ResolvedRange
    dependent: Node
    formula_reference: str

    def contains(self, node: Node) -> bool:
        sheet, row, column = node
        return (
            sheet == self.source.sheet
            and self.source.min_row <= row <= self.source.max_row
            and self.source.min_col <= column <= self.source.max_col
        )


@dataclass(frozen=True, slots=True)
class CircularReferenceResult:
    findings: tuple[Finding, ...]
    coverage: CoverageItem


class FormulaCycleGraph:
    """Formula-only adjacency; ranges query spatially indexed formula nodes.

    Candidate rectangles are memoized under independent entry and retained-ID
    limits. The optional limits exist for deterministic stress contracts; a
    zero entry limit disables admission, while over-sized entries are used for
    the current call without being retained.
    """

    def __init__(
        self,
        formula_nodes: list[Node],
        *,
        max_edges: int = _MAX_CYCLE_EDGES,
        max_candidate_cache_entries: int = _CANDIDATE_CACHE_ENTRY_CAP,
        max_candidate_cache_memberships: int = _CANDIDATE_CACHE_MEMBERSHIP_CAP,
    ) -> None:
        if (
            max_candidate_cache_entries < 0
            or max_candidate_cache_memberships < 0
        ):
            raise ValueError("candidate cache limits must be non-negative")
        self._nodes = sorted(formula_nodes)
        self._node_ids = {node: index for index, node in enumerate(self._nodes)}
        self._adjacency: list[set[int]] = [set() for _ in self._nodes]
        self._sheet_buckets: dict[str, dict[tuple[int, int], list[int]]] = {}
        for node_id, (sheet, row, column) in enumerate(self._nodes):
            self._sheet_buckets.setdefault(sheet, {}).setdefault(
                (row // _ROW_BUCKET, column // _COLUMN_BUCKET), []
            ).append(node_id)
        self._max_edges = max_edges
        self._edge_count = 0
        self.omitted_range_references = 0
        self._candidate_cache: OrderedDict[
            tuple[str, int, int, int, int], frozenset[int]
        ] = OrderedDict()
        self._candidate_cache_memberships = 0
        self._max_candidate_cache_entries = max_candidate_cache_entries
        self._max_candidate_cache_memberships = max_candidate_cache_memberships
        self.candidate_walks = 0
        self.candidate_hits = 0
        self.candidate_evictions = 0
        self.candidate_admission_skips = 0

    @property
    def node_count(self) -> int:
        return len(self._nodes)

    @property
    def edge_count(self) -> int:
        return self._edge_count

    def _candidates(self, resolved: ResolvedRange) -> frozenset[int]:
        """Formula node ids inside the rectangle; memoized per unique rectangle."""
        key = (
            resolved.sheet,
            resolved.min_row,
            resolved.max_row,
            resolved.min_col,
            resolved.max_col,
        )
        cached = self._candidate_cache.get(key)
        if cached is not None:
            self.candidate_hits += 1
            self._candidate_cache.move_to_end(key)
            return cached
        self.candidate_walks += 1
        buckets = self._sheet_buckets.get(resolved.sheet)
        precedent_ids: set[int] = set()
        if buckets:
            min_rb = resolved.min_row // _ROW_BUCKET
            max_rb = resolved.max_row // _ROW_BUCKET
            min_cb = resolved.min_col // _COLUMN_BUCKET
            max_cb = resolved.max_col // _COLUMN_BUCKET
            span = (max_rb - min_rb + 1) * (max_cb - min_cb + 1)
            if span <= len(buckets):
                bucket_lists = (
                    buckets.get((row_bucket, column_bucket), ())
                    for row_bucket in range(min_rb, max_rb + 1)
                    for column_bucket in range(min_cb, max_cb + 1)
                )
            else:
                bucket_lists = (
                    node_ids
                    for (row_bucket, column_bucket), node_ids in buckets.items()
                    if min_rb <= row_bucket <= max_rb
                    and min_cb <= column_bucket <= max_cb
                )
            for node_ids in bucket_lists:
                for node_id in node_ids:
                    _sheet, row, column = self._nodes[node_id]
                    if (
                        resolved.min_row <= row <= resolved.max_row
                        and resolved.min_col <= column <= resolved.max_col
                    ):
                        precedent_ids.add(node_id)
        frozen = frozenset(precedent_ids)
        if (
            self._max_candidate_cache_entries == 0
            or len(frozen) > self._max_candidate_cache_memberships
        ):
            self.candidate_admission_skips += 1
            return frozen
        while self._candidate_cache and (
            len(self._candidate_cache) >= self._max_candidate_cache_entries
            or self._candidate_cache_memberships + len(frozen)
            > self._max_candidate_cache_memberships
        ):
            _, evicted = self._candidate_cache.popitem(last=False)
            self._candidate_cache_memberships -= len(evicted)
            self.candidate_evictions += 1
        self._candidate_cache[key] = frozen
        self._candidate_cache_memberships += len(frozen)
        return frozen

    def add_range(self, resolved: ResolvedRange, dependent: Node) -> None:
        dependent_id = self._node_ids.get(dependent)
        if dependent_id is None:
            raise ValueError("cycle dependent is not a formula node")
        new_ids = self._candidates(resolved) - self._adjacency[dependent_id]
        if self._edge_count + len(new_ids) > self._max_edges:
            self.omitted_range_references += 1
            return
        self._adjacency[dependent_id].update(new_ids)
        self._edge_count += len(new_ids)

    def strongly_connected_components(
        self,
        cancellation_token: CancellationToken | None = None,
    ) -> tuple[tuple[Node, ...], ...]:
        """Return deterministic cyclic SCCs using iterative Kosaraju passes."""
        adjacency = [tuple(sorted(neighbors)) for neighbors in self._adjacency]
        visited: set[int] = set()
        order: list[int] = []
        operations = 0
        for start in range(len(self._nodes)):
            if start in visited:
                continue
            visited.add(start)
            stack: list[tuple[int, int]] = [(start, 0)]
            while stack:
                operations += 1
                if operations % 10_000 == 0:
                    check_cancelled(cancellation_token)
                node_id, neighbor_index = stack[-1]
                neighbors = adjacency[node_id]
                if neighbor_index < len(neighbors):
                    neighbor = neighbors[neighbor_index]
                    stack[-1] = (node_id, neighbor_index + 1)
                    if neighbor not in visited:
                        visited.add(neighbor)
                        stack.append((neighbor, 0))
                    continue
                stack.pop()
                order.append(node_id)

        reverse: list[list[int]] = [[] for _ in self._nodes]
        for source, neighbors in enumerate(adjacency):
            for target in neighbors:
                reverse[target].append(source)
        for neighbors in reverse:
            neighbors.sort()

        visited.clear()
        components: list[tuple[Node, ...]] = []
        for start in reversed(order):
            if start in visited:
                continue
            visited.add(start)
            pending = [start]
            member_ids: list[int] = []
            while pending:
                operations += 1
                if operations % 10_000 == 0:
                    check_cancelled(cancellation_token)
                node_id = pending.pop()
                member_ids.append(node_id)
                for neighbor in reversed(reverse[node_id]):
                    if neighbor not in visited:
                        visited.add(neighbor)
                        pending.append(neighbor)
            is_self_cycle = (
                len(member_ids) == 1
                and member_ids[0] in self._adjacency[member_ids[0]]
            )
            if len(member_ids) > 1 or is_self_cycle:
                components.append(
                    tuple(sorted(self._nodes[node_id] for node_id in member_ids))
                )
        return tuple(sorted(components, key=lambda component: component[0]))


def _display(node: Node) -> str:
    sheet, row, col = node
    return f"{sheet}!{get_column_letter(col)}{row}"


def _strip_sheet_quotes(name: str) -> str:
    if name.startswith("'") and name.endswith("'"):
        return name[1:-1].replace("''", "'")
    return name


class DependencyGraph:
    """Compact integer-keyed dependency index with rectangular descriptors."""

    def __init__(self, formula_nodes: list[Node] | None = None) -> None:
        self.unsupported_references: list[str] = []
        self.invalid_references: list[str] = []
        self.parse_errors: list[str] = []
        self.reference_status_counts: Counter[ReferenceStatus] = Counter()
        self.unsupported_reason_counts: Counter[str] = Counter()
        self.invalid_reason_counts: Counter[str] = Counter()
        self.parse_error_counts: Counter[str] = Counter()
        self.local_symbol_count = 0
        self._unsupported_samples: set[str] = set()
        self._invalid_samples: set[str] = set()
        self._parse_error_samples: set[str] = set()
        self.symbolic_references: list[SymbolicReference] = []
        self._symbolic_reference_set: set[SymbolicReference] = set()
        self._symbolic_rect_ids: dict[tuple[str, int, int, int, int], int] = {}
        self._symbolic_rect_bounds: list[tuple[str, int, int, int, int]] = []
        self._symbolic_rect_dependents: list[set[int]] = []
        self._symbolic_buckets: dict[tuple[str, int], list[int]] = {}
        self.bounded_references: list[str] = []
        self.bounded_reference_count = 0
        self._bounded_reference_samples: set[str] = set()
        self._node_ids: dict[Node, int] = {}
        self._nodes: list[Node] = []
        self._direct: dict[int, set[int]] = {}
        self._rect_ids: dict[tuple[str, int, int, int, int], int] = {}
        self._rect_bounds: list[tuple[str, int, int, int, int]] = []
        self._rect_dependents: list[set[int]] = []
        self._descriptor_count = 0
        #: Per bucket: rects covering the whole bucket (uniform membership)
        #: versus rects cutting through it (per-cell bounds check required).
        self._bucket_full: dict[tuple[str, int, int], list[int]] = {}
        self._bucket_partial: dict[tuple[str, int, int], list[int]] = {}
        self._closure_cache: dict[Node, frozenset[Node]] = {}
        #: Positional dependents keyed by matched rectangle combination; every
        #: cell in the same rectangle regime shares one materialized set.
        self._positional_cache: OrderedDict[
            tuple[tuple[int, ...], tuple[int, ...]], frozenset[Node]
        ] = OrderedDict()
        self._shared_closure_cache: OrderedDict[
            tuple[tuple[int, ...], tuple[int, ...], tuple[int, ...]],
            frozenset[Node],
        ] = OrderedDict()
        #: Downstream of one rectangle-combination's positional frontier,
        #: in id space. Regimes are few, so this stays unbounded.
        self._regime_downstream_cache: dict[
            tuple[tuple[int, ...], tuple[int, ...]], frozenset[int]
        ] = {}
        #: (regime ref, private delta ids) per signature; entries are small
        #: because the regime component is a shared reference.
        self._parts_cache: OrderedDict[
            tuple[tuple[int, ...], tuple[int, ...], tuple[int, ...]],
            tuple[frozenset[int], frozenset[int]],
        ] = OrderedDict()
        #: Smallest display strings per regime, grown on demand.
        self._regime_heads: dict[
            tuple[tuple[int, ...], tuple[int, ...]], tuple[int, tuple[str, ...]]
        ] = {}
        self._closure_summaries: OrderedDict[
            frozenset[Node], tuple[tuple[str, ...], int]
        ] = OrderedDict()
        self.formula_cycles = FormulaCycleGraph(formula_nodes or [])

    # --- construction -------------------------------------------------------

    def _node_id(self, node: Node) -> int:
        node_id = self._node_ids.get(node)
        if node_id is None:
            node_id = len(self._nodes)
            self._node_ids[node] = node_id
            self._nodes.append(node)
        return node_id

    def add_direct(self, precedent: Node, dependent: Node) -> None:
        if precedent == dependent:
            return
        self._direct.setdefault(self._node_id(precedent), set()).add(
            self._node_id(dependent)
        )

    def add_range(self, resolved: ResolvedRange, dependent: Node) -> None:
        """Store one rectangular precedent without materializing its members."""
        if resolved.size == 1:
            self.add_direct(
                (resolved.sheet, resolved.min_row, resolved.min_col), dependent
            )
            return
        key = (
            resolved.sheet,
            resolved.min_row,
            resolved.max_row,
            resolved.min_col,
            resolved.max_col,
        )
        dependent_id = self._node_id(dependent)
        rect_id = self._rect_ids.get(key)
        if rect_id is None:
            rect_id = len(self._rect_bounds)
            self._rect_ids[key] = rect_id
            self._rect_bounds.append(key)
            self._rect_dependents.append(set())
            for row_bucket in range(
                resolved.min_row // _ROW_BUCKET, resolved.max_row // _ROW_BUCKET + 1
            ):
                for column_bucket in range(
                    resolved.min_col // _COLUMN_BUCKET,
                    resolved.max_col // _COLUMN_BUCKET + 1,
                ):
                    bucket = (resolved.sheet, row_bucket, column_bucket)
                    covers_bucket = (
                        resolved.min_row <= row_bucket * _ROW_BUCKET
                        and resolved.max_row >= row_bucket * _ROW_BUCKET + _ROW_BUCKET - 1
                        and resolved.min_col <= column_bucket * _COLUMN_BUCKET
                        and resolved.max_col
                        >= column_bucket * _COLUMN_BUCKET + _COLUMN_BUCKET - 1
                    )
                    target = self._bucket_full if covers_bucket else self._bucket_partial
                    target.setdefault(bucket, []).append(rect_id)
        if dependent_id not in self._rect_dependents[rect_id]:
            self._rect_dependents[rect_id].add(dependent_id)
            self._descriptor_count += 1

    # --- queries ------------------------------------------------------------

    @property
    def direct_edge_count(self) -> int:
        return sum(len(dependents) for dependents in self._direct.values())

    @property
    def range_descriptor_count(self) -> int:
        return self._descriptor_count

    def _match_signature(
        self, node: Node
    ) -> tuple[tuple[int, ...], tuple[int, ...], tuple[int, ...]]:
        """Identity of everything that determines the node's dependent set.

        Components are sorted so equal match sets compare equal regardless of
        bucket iteration order, and the direct component is the dependent-id
        set itself so distinct cells with identical dependents share entries.
        """
        sheet, row, column = node
        node_id = self._node_ids.get(node)
        direct_ids: tuple[int, ...] = ()
        if node_id is not None:
            direct = self._direct.get(node_id)
            if direct:
                direct_ids = tuple(sorted(direct))
        rect_ids: list[int] = []
        bucket = (sheet, row // _ROW_BUCKET, column // _COLUMN_BUCKET)
        rect_ids.extend(self._bucket_full.get(bucket, ()))
        for rect_id in self._bucket_partial.get(bucket, ()):
            _sheet, min_row, max_row, min_col, max_col = self._rect_bounds[rect_id]
            if min_row <= row <= max_row and min_col <= column <= max_col:
                rect_ids.append(rect_id)
        rect_ids.sort()
        symbolic_ids: list[int] = []
        for rect_id in self._symbolic_buckets.get(
            (sheet, column // _COLUMN_BUCKET), ()
        ):
            _sheet, min_row, max_row, min_col, max_col = self._symbolic_rect_bounds[
                rect_id
            ]
            if min_row <= row <= max_row and min_col <= column <= max_col:
                symbolic_ids.append(rect_id)
        symbolic_ids.sort()
        return (direct_ids, tuple(rect_ids), tuple(symbolic_ids))

    def _positional_dependents(
        self, positional_key: tuple[tuple[int, ...], tuple[int, ...]]
    ) -> frozenset[Node]:
        positional = self._positional_cache.get(positional_key)
        if positional is None:
            rect_ids, symbolic_ids = positional_key
            dependent_ids: set[int] = set()
            for rect_id in rect_ids:
                dependent_ids.update(self._rect_dependents[rect_id])
            for rect_id in symbolic_ids:
                dependent_ids.update(self._symbolic_rect_dependents[rect_id])
            positional = frozenset(
                self._nodes[dependent_id] for dependent_id in dependent_ids
            )
            self._positional_cache[positional_key] = positional
            if len(self._positional_cache) > 512:
                self._positional_cache.popitem(last=False)
        else:
            self._positional_cache.move_to_end(positional_key)
        return positional

    def _materialized_dependents(
        self, signature: tuple[tuple[int, ...], tuple[int, ...], tuple[int, ...]]
    ) -> frozenset[Node]:
        direct_ids, rect_ids, symbolic_ids = signature
        positional = self._positional_dependents((rect_ids, symbolic_ids))
        if not direct_ids:
            return positional
        return positional | frozenset(
            self._nodes[dependent_id] for dependent_id in direct_ids
        )

    def direct_dependents(self, node: Node) -> set[Node]:
        """Cells whose formula references ``node`` without transitive closure."""
        dependents = set(self._materialized_dependents(self._match_signature(node)))
        dependents.discard(node)
        return dependents

    # --- telemetry ----------------------------------------------------------

    @staticmethod
    def _sample(target: list[str], seen: set[str], reason: str) -> None:
        if reason in seen:
            return
        seen.add(reason)
        if len(target) < _MAX_FAILURE_SAMPLES:
            target.append(reason)

    def record_reference(self, status: ReferenceStatus, detail: str) -> None:
        self.reference_status_counts[status] += 1
        if status is ReferenceStatus.INVALID:
            reason = reference_reason_code(status, detail)
            self.invalid_reason_counts[reason] += 1
            self._sample(self.invalid_references, self._invalid_samples, reason)
        elif status is ReferenceStatus.UNSUPPORTED:
            reason = reference_reason_code(status, detail)
            self.unsupported_reason_counts[reason] += 1
            self._sample(
                self.unsupported_references,
                self._unsupported_samples,
                reason,
            )

    def record_parse_error(self) -> None:
        reason = "formula_parse_error"
        self.parse_error_counts[reason] += 1
        self._sample(self.parse_errors, self._parse_error_samples, reason)

    def record_extraction_unsupported(self, reason: str) -> None:
        self.reference_status_counts[ReferenceStatus.UNSUPPORTED] += 1
        self.unsupported_reason_counts[reason] += 1
        self._sample(
            self.unsupported_references,
            self._unsupported_samples,
            reason,
        )

    def record_bounded_reference(self) -> None:
        reason = "bounded_whole_row_or_column"
        self.bounded_reference_count += 1
        self._sample(
            self.bounded_references,
            self._bounded_reference_samples,
            reason,
        )

    def record_symbolic_reference(
        self, resolved: ResolvedRange, dependent: Node, formula_reference: str
    ) -> None:
        symbolic = SymbolicReference(resolved, dependent, formula_reference)
        if symbolic in self._symbolic_reference_set:
            return
        self._symbolic_reference_set.add(symbolic)
        self.symbolic_references.append(symbolic)
        key = (
            resolved.sheet,
            resolved.min_row,
            resolved.max_row,
            resolved.min_col,
            resolved.max_col,
        )
        rect_id = self._symbolic_rect_ids.get(key)
        if rect_id is None:
            rect_id = len(self._symbolic_rect_bounds)
            self._symbolic_rect_ids[key] = rect_id
            self._symbolic_rect_bounds.append(key)
            self._symbolic_rect_dependents.append(set())
            for column_bucket in range(
                resolved.min_col // _COLUMN_BUCKET,
                resolved.max_col // _COLUMN_BUCKET + 1,
            ):
                self._symbolic_buckets.setdefault(
                    (resolved.sheet, column_bucket), []
                ).append(rect_id)
        self._symbolic_rect_dependents[rect_id].add(self._node_id(dependent))

    @property
    def coverage_state(self) -> CoverageState:
        return (
            CoverageState.DEGRADED
            if (
                self.reference_status_counts[ReferenceStatus.UNSUPPORTED]
                or self.reference_status_counts[ReferenceStatus.INVALID]
                or self.parse_error_counts
            )
            else CoverageState.CHECKED
        )

    @property
    def coverage_detail(self) -> str:
        details: list[str] = []
        if self.symbolic_references:
            details.append(
                f"{len(self.symbolic_references)} symbolic aggregate references"
            )
        if self.bounded_reference_count:
            details.append(
                f"{self.bounded_reference_count} bounded whole-row/column references"
            )
        unsupported = self.reference_status_counts[ReferenceStatus.UNSUPPORTED]
        if unsupported:
            details.append(
                f"{unsupported} unsupported references"
            )
        invalid = self.reference_status_counts[ReferenceStatus.INVALID]
        if invalid:
            details.append(f"{invalid} invalid references")
        parse_errors = sum(self.parse_error_counts.values())
        if parse_errors:
            details.append(f"{parse_errors} unparseable formulas")
        return "; ".join(details) or "All parsed formula references were resolved"


def build_dependency_graph(
    workbook: WorkbookSnapshot,
    *,
    max_range_cells: int = _MAX_RANGE_CELLS,
    cancellation_token: CancellationToken | None = None,
) -> DependencyGraph:
    formula_nodes = [
        (sheet.name, row, column)
        for sheet in workbook.sheets
        for (row, column), cell in sheet.cells.items()
        if cell.formula is not None
    ]
    graph = DependencyGraph(formula_nodes)
    reference_index = build_reference_index(workbook)
    formula_templates: dict[FormulaPatternKey, FormulaPrecedentExtraction] = {}
    oversized_symbolic: Counter[str] = Counter()
    for sheet in workbook.sheets:
        check_cancelled(cancellation_token)
        for index, ((row, col), cell) in enumerate(sheet.cells.items(), start=1):
            if index % 10_000 == 0:
                check_cancelled(cancellation_token)
            if cell.formula is None:
                continue
            dependent: Node = (sheet.name, row, col)
            try:
                pattern = formula_pattern_key(cell.formula)
                extraction = formula_templates.get(pattern)
                if extraction is None:
                    extraction = extract_formula_precedents(cell.formula)
                    formula_templates[pattern] = extraction
            except Exception:  # malformed formulas must not kill a run
                graph.record_parse_error()
                logger.warning("unparseable formula at %s", _display(dependent))
                continue
            for precedent in extraction.precedents:
                if precedent.kind is FormulaPrecedentKind.LOCAL_SYMBOL:
                    graph.local_symbol_count += 1
                    continue
                if precedent.kind is FormulaPrecedentKind.UNSUPPORTED:
                    graph.record_extraction_unsupported(precedent.reason)
                    continue
                value = precedent.value
                resolution = resolve_reference(
                    workbook,
                    value,
                    host_sheet=sheet.name,
                    host_cell=(row, col),
                    require_within_sheet=False,
                    index=reference_index,
                )
                graph.record_reference(resolution.status, resolution.detail)
                if resolution.status is not ReferenceStatus.RESOLVED:
                    continue
                if resolution.detail and "bounded whole-row/column" in resolution.detail:
                    graph.record_bounded_reference()
                for resolved in resolution.ranges:
                    graph.formula_cycles.add_range(resolved, dependent)
                    if resolved.size > max_range_cells:
                        graph.record_symbolic_reference(resolved, dependent, value)
                        oversized_symbolic[value] += 1
                        continue
                    graph.add_range(resolved, dependent)
    if oversized_symbolic:
        logger.warning(
            "represented %d oversized range references symbolically "
            "(%d unique references)",
            sum(oversized_symbolic.values()),
            len(oversized_symbolic),
        )
    return graph


def detect_circular_references(
    graph: DependencyGraph,
    *,
    cancellation_token: CancellationToken | None = None,
) -> CircularReferenceResult:
    """Report each exact formula SCC once; disclose every incomplete edge source."""
    components = graph.formula_cycles.strongly_connected_components(
        cancellation_token
    )
    findings: list[Finding] = []
    for component in components:
        displays = [_display(node) for node in component]
        canonical = json.dumps(displays, separators=(",", ":"))
        digest = hashlib.sha256(canonical.encode("utf-8")).hexdigest()[:12]
        event_key = f"excel-cycle:{digest}"
        sheet, row, column = component[0]
        sample = displays[:8]
        omitted = len(displays) - len(sample)
        element = "; ".join(sample)
        if omitted:
            element += f"; +{omitted} more cells"
        findings.append(
            Finding(
                artifact="excel",
                finding_class=FindingClass.CIRCULAR_REFERENCE,
                sheet=sheet,
                location=f"{get_column_letter(column)}{row}",
                element=element,
                current_value=f"{len(component)} formula cells",
                event_key=event_key,
                root_cause_key=event_key,
                message=(
                    f"Circular formula dependency contains {len(component)} "
                    "formula cells"
                ),
            )
        )

    unsupported = graph.reference_status_counts[ReferenceStatus.UNSUPPORTED]
    invalid = graph.reference_status_counts[ReferenceStatus.INVALID]
    parse_errors = sum(graph.parse_error_counts.values())
    omitted_ranges = graph.formula_cycles.omitted_range_references
    incomplete = bool(unsupported or invalid or parse_errors or omitted_ranges)
    details = [
        f"{graph.formula_cycles.node_count} formula nodes",
        f"{graph.formula_cycles.edge_count} exact formula edges",
        f"{len(components)} circular components",
    ]
    if unsupported:
        details.append(f"{unsupported} unsupported references")
    if invalid:
        details.append(f"{invalid} invalid references")
    if parse_errors:
        details.append(f"{parse_errors} unparseable formulas")
    if omitted_ranges:
        details.append(
            f"{omitted_ranges} formula ranges exceeded the cycle-edge budget"
        )
    coverage = CoverageItem(
        check_id="excel-circular-references",
        label="Circular formula references",
        artifact="excel",
        state=CoverageState.DEGRADED if incomplete else CoverageState.CHECKED,
        findings=len(findings),
        detail="; ".join(details),
    )
    return CircularReferenceResult(findings=tuple(findings), coverage=coverage)


def _expand_ids(
    graph: DependencyGraph,
    discovered: set[int],
    pending: list[Node],
) -> None:
    """Expand ``pending`` transitively into ``discovered`` (id space).

    Pre-seeded members of ``discovered`` are never re-expanded: a closure is
    transitive, so anything already known complete contributes nothing new.
    """
    absorbed_rects: set[int] = set()
    absorbed_symbolic: set[int] = set()
    absorbed_full_buckets: set[tuple[str, int, int]] = set()
    while pending:
        current = pending.pop()
        complete = graph._closure_cache.get(current)
        if complete is not None:
            for dependent in complete:
                discovered.add(graph._node_ids[dependent])
            continue
        sheet, row, column = current
        current_id = graph._node_ids.get(current)
        if current_id is not None:
            direct = graph._direct.get(current_id)
            if direct:
                for dependent_id in direct:
                    if dependent_id in discovered:
                        continue
                    discovered.add(dependent_id)
                    pending.append(graph._nodes[dependent_id])
        bucket = (sheet, row // _ROW_BUCKET, column // _COLUMN_BUCKET)
        if bucket not in absorbed_full_buckets:
            absorbed_full_buckets.add(bucket)
            for rect_id in graph._bucket_full.get(bucket, ()):
                if rect_id in absorbed_rects:
                    continue
                absorbed_rects.add(rect_id)
                for dependent_id in graph._rect_dependents[rect_id]:
                    if dependent_id in discovered:
                        continue
                    discovered.add(dependent_id)
                    pending.append(graph._nodes[dependent_id])
        for rect_id in graph._bucket_partial.get(bucket, ()):
            if rect_id in absorbed_rects:
                continue
            _sheet, min_row, max_row, min_col, max_col = graph._rect_bounds[rect_id]
            if not (min_row <= row <= max_row and min_col <= column <= max_col):
                continue
            absorbed_rects.add(rect_id)
            for dependent_id in graph._rect_dependents[rect_id]:
                if dependent_id in discovered:
                    continue
                discovered.add(dependent_id)
                pending.append(graph._nodes[dependent_id])
        for rect_id in graph._symbolic_buckets.get(
            (sheet, column // _COLUMN_BUCKET), ()
        ):
            if rect_id in absorbed_symbolic:
                continue
            _sheet, min_row, max_row, min_col, max_col = graph._symbolic_rect_bounds[
                rect_id
            ]
            if not (min_row <= row <= max_row and min_col <= column <= max_col):
                continue
            absorbed_symbolic.add(rect_id)
            for dependent_id in graph._symbolic_rect_dependents[rect_id]:
                if dependent_id in discovered:
                    continue
                discovered.add(dependent_id)
                pending.append(graph._nodes[dependent_id])


def _regime_downstream(
    graph: DependencyGraph,
    positional_key: tuple[tuple[int, ...], tuple[int, ...]],
) -> frozenset[int]:
    """Frontier-plus-downstream of one rectangle combination, computed once."""
    cached = graph._regime_downstream_cache.get(positional_key)
    if cached is not None:
        return cached
    rect_ids, symbolic_ids = positional_key
    seed_ids: set[int] = set()
    for rect_id in rect_ids:
        seed_ids.update(graph._rect_dependents[rect_id])
    for rect_id in symbolic_ids:
        seed_ids.update(graph._symbolic_rect_dependents[rect_id])
    discovered = set(seed_ids)
    pending = [graph._nodes[dependent_id] for dependent_id in seed_ids]
    _expand_ids(graph, discovered, pending)
    frozen = frozenset(discovered)
    graph._regime_downstream_cache[positional_key] = frozen
    return frozen


def _delta_ids(
    graph: DependencyGraph,
    regime: frozenset[int],
    direct_ids: tuple[int, ...],
) -> frozenset[int]:
    """The source's private downstream beyond the shared regime component."""
    if not direct_ids:
        return frozenset()
    discovered = set(regime)
    pending: list[Node] = []
    for dependent_id in direct_ids:
        if dependent_id in discovered:
            continue
        discovered.add(dependent_id)
        pending.append(graph._nodes[dependent_id])
    if pending:
        _expand_ids(graph, discovered, pending)
    discovered.difference_update(regime)
    return frozenset(discovered)


def _closure_parts(
    graph: DependencyGraph,
    signature: tuple[tuple[int, ...], tuple[int, ...], tuple[int, ...]],
) -> tuple[frozenset[int], frozenset[int]]:
    """Signature-determined closure as (shared regime ids, private delta ids).

    Valid for non-formula sources only: dependents are always formula cells,
    so a non-formula source can never appear inside its own closure and no
    source exclusion is needed.
    """
    cached = graph._parts_cache.get(signature)
    if cached is not None:
        graph._parts_cache.move_to_end(signature)
        return cached
    direct_ids, rect_ids, symbolic_ids = signature
    regime = _regime_downstream(graph, (rect_ids, symbolic_ids))
    delta = _delta_ids(graph, regime, direct_ids)
    parts = (regime, delta)
    graph._parts_cache[signature] = parts
    if len(graph._parts_cache) > 65536:
        graph._parts_cache.popitem(last=False)
    return parts


def _regime_head(
    graph: DependencyGraph,
    positional_key: tuple[tuple[int, ...], tuple[int, ...]],
    count: int,
) -> tuple[str, ...]:
    """Smallest ``count`` display strings of the regime, cached and grown."""
    import heapq

    regime = _regime_downstream(graph, positional_key)
    wanted = min(count, len(regime))
    cached = graph._regime_heads.get(positional_key)
    if cached is not None and cached[0] >= wanted:
        return cached[1]
    size = max(count, 64)
    head = tuple(
        heapq.nsmallest(
            size, (_display(graph._nodes[member]) for member in regime)
        )
    )
    graph._regime_heads[positional_key] = (min(size, len(regime)), head)
    return head


def _frozen_closure(graph: DependencyGraph, node: Node) -> frozenset[Node]:
    """Cached transitive closure; the caller must not mutate the result."""
    cached = graph._closure_cache.get(node)
    if cached is not None:
        return cached
    signature = graph._match_signature(node)
    if node not in graph.formula_cycles._node_ids:
        shared = graph._shared_closure_cache.get(signature)
        if shared is not None:
            graph._shared_closure_cache.move_to_end(signature)
            return shared
        regime, delta = _closure_parts(graph, signature)
        frozen = frozenset(
            graph._nodes[dependent_id] for dependent_id in regime
        ) | frozenset(graph._nodes[dependent_id] for dependent_id in delta)
        graph._shared_closure_cache[signature] = frozen
        if len(graph._shared_closure_cache) > 1024:
            graph._shared_closure_cache.popitem(last=False)
        return frozen
    # Formula sources can sit inside their own regime, so they keep the
    # explicit walk with source exclusion and per-node caching.
    direct_ids, rect_ids, symbolic_ids = signature
    discovered = set(_regime_downstream(graph, (rect_ids, symbolic_ids)))
    pending: list[Node] = []
    for dependent_id in direct_ids:
        if dependent_id in discovered:
            continue
        discovered.add(dependent_id)
        pending.append(graph._nodes[dependent_id])
    if pending:
        _expand_ids(graph, discovered, pending)
    source_id = graph._node_ids.get(node)
    if source_id is not None:
        discovered.discard(source_id)
    frozen = frozenset(graph._nodes[dependent_id] for dependent_id in discovered)
    graph._closure_cache[node] = frozen
    return frozen


def frozen_closure(graph: DependencyGraph, node: Node) -> frozenset[Node]:
    """Cached transitive closure; callers must treat the result as immutable."""
    return _frozen_closure(graph, node)


def closure_parts_for(
    graph: DependencyGraph, node: Node
) -> tuple[
    tuple[tuple[int, ...], tuple[int, ...]], frozenset[int], frozenset[int]
] | None:
    """(regime key, regime ids, delta ids) for non-formula sources, else None."""
    if node in graph.formula_cycles._node_ids:
        return None
    direct_ids, rect_ids, symbolic_ids = graph._match_signature(node)
    regime, delta = _closure_parts(graph, (direct_ids, rect_ids, symbolic_ids))
    return ((rect_ids, symbolic_ids), regime, delta)


_RECT_SCAN_CAP = 4096


def rect_intersects_members(
    graph: DependencyGraph,
    members: frozenset[int],
    sheet: str,
    min_row: int,
    max_row: int,
    min_col: int,
    max_col: int,
) -> bool:
    """Whether any member node lies inside the rectangle (smaller-side scan)."""
    if not members:
        return False
    span = (max_row - min_row + 1) * (max_col - min_col + 1)
    if span <= _RECT_SCAN_CAP and span <= len(members):
        node_ids = graph._node_ids
        for row in range(min_row, max_row + 1):
            for column in range(min_col, max_col + 1):
                member = node_ids.get((sheet, row, column))
                if member is not None and member in members:
                    return True
        return False
    nodes = graph._nodes
    for member in members:
        node_sheet, node_row, node_column = nodes[member]
        if (
            node_sheet == sheet
            and min_row <= node_row <= max_row
            and min_col <= node_column <= max_col
        ):
            return True
    return False


def dependent_nodes_of(graph: DependencyGraph, node: Node) -> set[Node]:
    """Transitive dependents across direct edges, range descriptors, aggregates.

    Closures are memoized per requested source node, and additionally shared
    across non-formula sources with identical match signatures: dependents are
    always formula cells, so a non-formula source can never appear inside its
    own closure and equal signatures walk identical frontiers.
    """
    return set(_frozen_closure(graph, node))


def dependents_of(graph: DependencyGraph, sheet: str, ref: str) -> list[str]:
    """Transitive downstream dependents of ``sheet!ref``, as display strings."""
    from openpyxl.utils.cell import coordinate_to_tuple

    row, col = coordinate_to_tuple(ref)
    node: Node = (sheet, row, col)
    return sorted(_display(dependent) for dependent in dependent_nodes_of(graph, node))


def annotate_impacts(findings: list[Finding], graph: DependencyGraph) -> None:
    """Attach downstream impact lists to cell-level findings (in place)."""
    from openpyxl.utils.cell import coordinate_to_tuple

    for finding in findings:
        if finding.finding_class not in _ANNOTATED_CLASSES:
            continue
        if finding.sheet is None or finding.location is None:
            continue
        if not _CELL_REF_RE.match(finding.location):
            continue
        row, col = coordinate_to_tuple(finding.location)
        frozen = _frozen_closure(graph, (finding.sheet, row, col))
        display = sorted(_display(dependent) for dependent in frozen)
        finding.impacts = sorted({*finding.impacts, *display})


_DISPLAY_RE = re.compile(r"^(?P<column>[A-Z]{1,3})(?P<row>[1-9][0-9]*)$")


def _parse_display(display: str) -> Node | None:
    """Inverse of ``_display``; ``None`` when the string is not a cell display."""
    from openpyxl.utils.cell import column_index_from_string

    sheet, separator, ref = display.rpartition("!")
    if not separator or not sheet:
        return None
    match = _DISPLAY_RE.match(ref)
    if match is None:
        return None
    return (sheet, int(match.group("row")), column_index_from_string(match.group("column")))


class ImpactAccumulator:
    """Impact annotation that never materializes a full closure per finding.

    ``annotate`` records each eligible finding's closure as signature parts
    (a shared regime reference plus a private delta); other impact sources
    keep appending their small display lists to ``finding.impacts``.
    ``finalize`` then produces exactly what ``annotate_impacts`` followed by
    ``limit_impacts`` produce: the first ``max_impacts`` of the sorted union
    plus the overflow sentinel — merging a cached regime display head with
    the tiny delta instead of sorting a 10k-string list per finding.
    """

    def __init__(self, graph: DependencyGraph) -> None:
        self._graph = graph
        self._parts: dict[
            int,
            tuple[
                tuple[tuple[int, ...], tuple[int, ...]],
                frozenset[int],
                frozenset[int],
            ],
        ] = {}
        self._formula_closures: dict[int, frozenset[Node]] = {}

    def annotate(self, findings: list[Finding]) -> None:
        from openpyxl.utils.cell import coordinate_to_tuple

        graph = self._graph
        formula_nodes = graph.formula_cycles._node_ids
        for finding in findings:
            if finding.finding_class not in _ANNOTATED_CLASSES:
                continue
            if finding.sheet is None or finding.location is None:
                continue
            if not _CELL_REF_RE.match(finding.location):
                continue
            row, col = coordinate_to_tuple(finding.location)
            node = (finding.sheet, row, col)
            if node in formula_nodes:
                # Formula sources need source exclusion; keep the exact path.
                self._formula_closures[id(finding)] = _frozen_closure(graph, node)
                continue
            direct_ids, rect_ids, symbolic_ids = graph._match_signature(node)
            regime, delta = _closure_parts(
                graph, (direct_ids, rect_ids, symbolic_ids)
            )
            self._parts[id(finding)] = ((rect_ids, symbolic_ids), regime, delta)

    def _summary(
        self, frozen: frozenset[Node], head: int
    ) -> tuple[tuple[str, ...], int]:
        cached = self._graph._closure_summaries.get(frozen)
        if cached is not None and len(cached[0]) >= min(head, cached[1]):
            self._graph._closure_summaries.move_to_end(frozen)
            return cached
        import heapq

        smallest = tuple(heapq.nsmallest(head, map(_display, frozen)))
        summary = (smallest, len(frozen))
        self._graph._closure_summaries[frozen] = summary
        if len(self._graph._closure_summaries) > 1024:
            self._graph._closure_summaries.popitem(last=False)
        return summary

    def _finalize_parts(
        self,
        finding: Finding,
        parts: tuple[
            tuple[tuple[int, ...], tuple[int, ...]],
            frozenset[int],
            frozenset[int],
        ],
        extras: list[str],
        max_impacts: int,
    ) -> None:
        import heapq

        graph = self._graph
        positional_key, regime, delta = parts
        unique_extras = set(extras)
        outside = 0
        for extra in unique_extras:
            extra_node = _parse_display(extra)
            if extra_node is None:
                outside += 1
                continue
            extra_id = graph._node_ids.get(extra_node)
            if extra_id is None or (
                extra_id not in regime and extra_id not in delta
            ):
                outside += 1
        total = len(regime) + len(delta) + outside
        if total <= max_impacts:
            merged = sorted(
                unique_extras
                | {_display(graph._nodes[member]) for member in regime}
                | {_display(graph._nodes[member]) for member in delta}
            )
            finding.impacts = merged
            return
        head_size = max_impacts + len(unique_extras)
        regime_head = _regime_head(graph, positional_key, head_size)
        delta_displays = sorted(
            _display(graph._nodes[member]) for member in delta
        )
        merged_head = heapq.nsmallest(
            head_size,
            set(regime_head) | set(delta_displays) | unique_extras,
        )
        finding.impacts = [
            *merged_head[:max_impacts],
            f"... {total - max_impacts} more",
        ]

    def finalize(
        self, findings: list[Finding], *, max_impacts: int = 25
    ) -> None:
        """Apply the exact merged display cap across every impact source."""
        for finding in findings:
            extras = [
                impact
                for impact in finding.impacts
                if not impact.startswith("... ")
            ]
            parts = self._parts.get(id(finding))
            if parts is not None:
                self._finalize_parts(finding, parts, extras, max_impacts)
                continue
            frozen = self._formula_closures.get(id(finding))
            if frozen is None:
                # Mirror ``limit_impacts`` exactly: it sorts without deduping.
                impacts = sorted(extras)
                if len(impacts) > max_impacts:
                    finding.impacts = [
                        *impacts[:max_impacts],
                        f"... {len(impacts) - max_impacts} more",
                    ]
                else:
                    finding.impacts = impacts
                continue
            unique_extras = set(extras)
            outside = [
                extra
                for extra in unique_extras
                if (node := _parse_display(extra)) is None or node not in frozen
            ]
            total = len(frozen) + len(outside)
            if total <= max_impacts:
                merged = sorted(
                    unique_extras | {_display(node) for node in frozen}
                )
                finding.impacts = merged
                continue
            head, _count = self._summary(frozen, max_impacts + len(unique_extras))
            merged = sorted(unique_extras | set(head))
            finding.impacts = [
                *merged[:max_impacts],
                f"... {total - max_impacts} more",
            ]
        self._parts.clear()
        self._formula_closures.clear()


def limit_impacts(findings: list[Finding], *, max_impacts: int = 25) -> None:
    """Apply one accurate display cap after every impact source has merged."""
    for finding in findings:
        impacts = sorted(
            impact
            for impact in finding.impacts
            if not impact.startswith("... ")
        )
        if len(impacts) > max_impacts:
            finding.impacts = [
                *impacts[:max_impacts],
                f"... {len(impacts) - max_impacts} more",
            ]
        else:
            finding.impacts = impacts
