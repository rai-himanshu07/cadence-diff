"""Formula QC: errors, hardcoding, logic changes, extension, consistency.

Comparisons are shift-invariant: every formula is normalized to R1C1
relative to its host cell, so a formula that merely moved with inserted
cadence data compares equal. Where normalized forms differ only in range
*end* references that grow (same sheet, same anchor, larger end), the
change is classified as an expected range extension, not a logic change.

Checks (paired regions come from the alignment engine):

- error values anywhere in the current workbook, and error literals
  inside formula text;
- baseline formula replaced by a constant (hardcode);
- normalized formula changed vs baseline (minus expected extensions);
- growth rows/columns missing the formula pattern of their axis run
  (not extended);
- formula cells deviating from the dominant R1C1 pattern of their run
  (inconsistent within range).

In a cycle comparison, error and consistency findings also carry additive
provenance so an inherited current-state risk is distinguishable from a new
regression. Provenance is assigned only where a confidently aligned baseline
cell proves it.

xlsb workbooks may expose formula-record presence without formula text. That
presence safely enables hardcode, removal, and missing fill checks; semantic
logic and consistency checks require compatible decoded text on both sides.
"""

import difflib
import hashlib
import logging
import re
import time
from collections import Counter
from dataclasses import dataclass, field

from openpyxl.formula.tokenizer import TokenizerError
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import coordinate_to_tuple

from qc_tool.availability import cell_in_ranges, excel_blank_allowed
from qc_tool.config.profile import DeliverableProfile, SheetProfile
from qc_tool.excel.align import RegionAlignment, WorkbookAlignment
from qc_tool.excel.formula_tokens import (
    FormulaDiffKind,
    FormulaDiffSegment,
    formula_reference_operands,
    tokenize_formula,
)
from qc_tool.excel.population import CandidateSpill
from qc_tool.findings import (
    Finding,
    FindingClass,
    FindingEvidenceTag,
    FindingExpectedReason,
    FindingProvenance,
    FindingSubtype,
)
from qc_tool.io.model import (
    ERROR_LITERALS,
    CellRecord,
    SheetSnapshot,
    WorkbookSnapshot,
    display_cell_value,
)
from qc_tool.progress import CancellationToken, check_cancelled

logger = logging.getLogger(__name__)

_ENDPOINT_RE = re.compile(r"^(\$?)([A-Z]{1,3})(\$?)(\d+)$")
_RUN_MIN_CELLS = 3
_RUN_DOMINANCE = 0.5
_FORMULAIC_SHARE = 0.5


def _ref(row: int, col: int) -> str:
    return f"{get_column_letter(col)}{row}"


# --- R1C1 normalization ---------------------------------------------------


def _endpoint_to_r1c1(endpoint: str, host_row: int, host_col: int) -> str | None:
    match = _ENDPOINT_RE.match(endpoint)
    if match is None:
        return None
    col_abs, col_letters, row_abs, row_digits = match.groups()
    from openpyxl.utils.cell import column_index_from_string

    col = column_index_from_string(col_letters)
    row = int(row_digits)
    row_part = f"R{row}" if row_abs else ("R" if row == host_row else f"R[{row - host_row}]")
    col_part = f"C{col}" if col_abs else ("C" if col == host_col else f"C[{col - host_col}]")
    return row_part + col_part


def _range_token_to_r1c1(token: str, host_row: int, host_col: int) -> str:
    sheet_prefix, sep, ref = token.rpartition("!")
    parts = ref.split(":")
    converted: list[str] = []
    for part in parts:
        r1c1 = _endpoint_to_r1c1(part, host_row, host_col)
        if r1c1 is None:
            return token  # named range, whole-row/col ref, etc. — keep verbatim
        converted.append(r1c1)
    return sheet_prefix + sep + ":".join(converted)


def to_r1c1(formula: str, host_row: int, host_col: int) -> str:
    """Normalize an A1-style formula to R1C1 relative to its host cell."""
    try:
        tokens = tokenize_formula(formula)
    except Exception:  # malformed formulas must not kill a run
        logger.warning("unparseable formula at %s: %r", _ref(host_row, host_col), formula)
        return formula
    rendered = [
        _range_token_to_r1c1(t.value, host_row, host_col)
        if t.type == "OPERAND" and t.subtype == "RANGE"
        else t.value
        for t in tokens
    ]
    return "=" + "".join(rendered)


def _normalize_formula(cell: CellRecord, row: int, col: int) -> str:
    """R1C1 text for a cell: the adapter-supplied value when present
    (currently only the native XLSB kernel, already validated per definition
    -- see `qc_tool.io.native_formula._validated_r1c1_cells`), else computed
    here from `cell.formula`.
    """
    if cell.formula_r1c1 is not None:
        return cell.formula_r1c1
    return to_r1c1(cell.formula or "", row, col)


# --- wrapper detection ------------------------------------------------------


@dataclass(frozen=True, slots=True)
class WrapperMatch:
    """One formula preserved as an argument-level subtree of the other."""

    kind: str  # "wrapped" (baseline inside current) | "unwrapped"
    exact: bool  # False when only the abstracted shape is preserved
    skeleton_key: str  # stable hash of the outer shape with the core removed


def _wrapper_tokens(formula: str) -> list[tuple[str, str, str]] | None:
    try:
        tokens = tokenize_formula(formula)
    except Exception:
        return None
    return [(t.type, t.subtype, t.value) for t in tokens]


def _shape_label(token: tuple[str, str, str]) -> str:
    token_type, subtype, value = token
    if token_type == "OPERAND" and subtype == "RANGE":
        return "REF"
    if token_type == "OPERAND" and subtype == "NUMBER":
        return "NUM"
    if token_type == "OPERAND" and subtype == "TEXT":
        return "TEXT"
    return value.casefold()


def _is_boundary_open(token: tuple[str, str, str] | None) -> bool:
    if token is None:
        return True  # sequence start
    token_type, subtype, _ = token
    return (token_type in ("FUNC", "PAREN") and subtype == "OPEN") or (
        token_type == "SEP" and subtype == "ARG"
    )


def _is_boundary_close(token: tuple[str, str, str] | None) -> bool:
    if token is None:
        return True  # sequence end
    token_type, subtype, _ = token
    return (token_type in ("FUNC", "PAREN") and subtype == "CLOSE") or (
        token_type == "SEP" and subtype == "ARG"
    )


def _find_argument_occurrence(
    outer: list[tuple[str, str, str]],
    inner_keys: list[str],
    key_of,
) -> int | None:
    """Index where `inner_keys` occurs in `outer` at argument boundaries."""
    outer_keys = [key_of(token) for token in outer]
    span = len(inner_keys)
    if span == 0 or span >= len(outer_keys):
        return None
    for start in range(len(outer_keys) - span + 1):
        if outer_keys[start : start + span] != inner_keys:
            continue
        before = outer[start - 1] if start > 0 else None
        after_index = start + span
        after = outer[after_index] if after_index < len(outer) else None
        if _is_boundary_open(before) and _is_boundary_close(after):
            return start
    return None


def _is_meaningful_core(tokens: list[tuple[str, str, str]]) -> bool:
    """Bare literals never classify as a preserved core."""
    return any(
        (token_type == "OPERAND" and subtype == "RANGE")
        or (token_type == "FUNC" and subtype == "OPEN")
        for token_type, subtype, _ in tokens
    )


def _is_structural_core(tokens: list[tuple[str, str, str]]) -> bool:
    """Shape-level matches need real structure: a lone REF would match any
    reference argument once abstracted."""
    return any(
        token_type == "FUNC" and subtype == "OPEN" for token_type, subtype, _ in tokens
    ) or sum(1 for token_type, _, _ in tokens if token_type == "OPERAND") >= 2


def _skeleton_key(outer: list[tuple[str, str, str]], start: int, span: int) -> str:
    labels = [_shape_label(token) for token in outer[:start]]
    labels.append("<CORE>")
    labels.extend(_shape_label(token) for token in outer[start + span :])
    digest = hashlib.sha256("|".join(labels).encode("utf-8")).hexdigest()
    return digest[:12]


def detect_formula_wrapper(base_norm: str, curr_norm: str) -> WrapperMatch | None:
    """Detect a formula preserved as an argument of the other formula.

    Containment is token-exact first (identical R1C1 logic inside a new
    guard), then shape-abstracted (same structure with shifted references or
    changed literals — the rolling-window case). Matches must sit at argument
    boundaries, so `A1+B1` inside `A1+B1*2` never counts, and string literals
    cannot fake a match because they tokenize as single TEXT operands.
    """
    base_tokens = _wrapper_tokens(base_norm)
    curr_tokens = _wrapper_tokens(curr_norm)
    if not base_tokens or not curr_tokens:
        return None

    def exact_key(token: tuple[str, str, str]) -> str:
        return token[2]

    candidates: list[tuple[str, list[tuple[str, str, str]], list[tuple[str, str, str]]]] = [
        ("wrapped", curr_tokens, base_tokens),
        ("unwrapped", base_tokens, curr_tokens),
    ]
    for kind, outer, inner in candidates:
        if not _is_meaningful_core(inner):
            continue
        start = _find_argument_occurrence(outer, [exact_key(t) for t in inner], exact_key)
        if start is not None:
            return WrapperMatch(
                kind=kind,
                exact=True,
                skeleton_key=_skeleton_key(outer, start, len(inner)),
            )
    for kind, outer, inner in candidates:
        if not _is_meaningful_core(inner) or not _is_structural_core(inner):
            continue
        start = _find_argument_occurrence(
            outer, [_shape_label(t) for t in inner], _shape_label
        )
        if start is not None:
            return WrapperMatch(
                kind=kind,
                exact=False,
                skeleton_key=_skeleton_key(outer, start, len(inner)),
            )
    return None


_WRAPPER_NOTES = {
    ("wrapped", True): " (existing logic preserved inside a new wrapper)",
    ("wrapped", False): (
        " (logic shape preserved inside a new wrapper; inner references shifted)"
    ),
    ("unwrapped", True): " (wrapper removed; inner logic preserved)",
    ("unwrapped", False): (
        " (wrapper removed; inner logic shape preserved with shifted references)"
    ),
}

_WRAPPER_SUBTYPES = {
    "wrapped": FindingSubtype.FORMULA_WRAPPED,
    "unwrapped": FindingSubtype.FORMULA_UNWRAPPED,
}


# --- expected range extension ---------------------------------------------


def _split_range(token: str) -> tuple[str, str, str] | None:
    sheet, _, ref = token.rpartition("!")
    start, sep, end = ref.partition(":")
    if not sep:
        return None
    return sheet, start, end


def _is_range_extension(base_token: str, curr_token: str) -> bool:
    base_parts = _split_range(base_token)
    curr_parts = _split_range(curr_token)
    if base_parts is None or curr_parts is None:
        return False
    if base_parts[0] != curr_parts[0] or base_parts[1] != curr_parts[1]:
        return False
    base_end = _ENDPOINT_RE.match(base_parts[2])
    curr_end = _ENDPOINT_RE.match(curr_parts[2])
    if base_end is None or curr_end is None:
        return False
    same_col = base_end[2] == curr_end[2]
    same_row = base_end[4] == curr_end[4]
    if same_col and int(curr_end[4]) >= int(base_end[4]):
        return True
    return bool(
        same_row
        and (
            len(curr_end[2]) > len(base_end[2])
            or (len(curr_end[2]) == len(base_end[2]) and curr_end[2] >= base_end[2])
        )
    )


def _differs_only_by_extension(base_formula: str, curr_formula: str) -> bool:
    try:
        base_tokens = tokenize_formula(base_formula)
        curr_tokens = tokenize_formula(curr_formula)
    except Exception:  # malformed formulas cannot be extension-classified
        return False
    if len(base_tokens) != len(curr_tokens):
        return False
    extension_seen = False
    for base_tok, curr_tok in zip(base_tokens, curr_tokens, strict=True):
        if (base_tok.type, base_tok.subtype) != (curr_tok.type, curr_tok.subtype):
            return False
        if base_tok.value == curr_tok.value:
            continue
        if base_tok.type == "OPERAND" and base_tok.subtype == "RANGE":
            if not _is_range_extension(base_tok.value, curr_tok.value):
                return False
            extension_seen = True
        else:
            return False
    return extension_seen


# --- telemetry ---------------------------------------------------------------


@dataclass(slots=True)
class FormulaComparisonTelemetry:
    """Aggregate-only counters/timers for one ``diff_workbook_formulas()`` call.

    Every field is a count or a duration in seconds; no formula, sheet name,
    coordinate, or defined name is ever recorded here. Passing an instance
    changes no finding, ordering, or evidence -- it only accumulates these
    counters as a side effect. The whole-function fields
    (``error_scan_seconds``, ``paired_traversal_seconds``,
    ``consistency_seconds``, ``extension_findings_seconds``) are INCLUSIVE of
    every sub-step measured below them; they are wall-clock totals for that
    function, not an exclusive residual.

    ``complexity_assessment_seconds`` is the one exception to the "one
    ``diff_workbook_formulas()`` call" framing above: ``run_qc()`` measures it
    directly around its own ``assess_workbook_complexity()`` call, which runs
    inside the same ``RunPhase.COMPARING_FORMULAS`` boundary but is not part
    of ``diff_workbook_formulas()`` itself. It is included here so a
    ``comparing_formulas`` phase-duration reconciliation has somewhere to put
    that cost instead of leaving it as an unexplained residual
    (plan-20260908-phase-b-guest-performance-followup.md).
    """

    error_scan_seconds: float = 0.0
    error_scan_cells: int = 0

    paired_traversal_seconds: float = 0.0
    paired_pairs_considered: int = 0
    #: Pairs whose formula text and host coordinates are both identical --
    #: candidates for a same-position exact-text shortcut that would skip
    #: normalization entirely.
    exact_text_same_host_pairs: int = 0

    normalization_calls: int = 0
    normalization_seconds: float = 0.0
    #: Hits/misses against the region-scoped current-side normalization memo
    #: shared by paired and consistency checks (Step 2). A hit means
    #: `normalization_calls`/`consistency_normalization_calls` counted only
    #: the baseline-side (or first-seen) `to_r1c1` call, not a repeat.
    memo_hits: int = 0
    memo_misses: int = 0

    consistency_seconds: float = 0.0
    consistency_normalization_calls: int = 0
    consistency_normalization_seconds: float = 0.0

    wrapper_reference_seconds: float = 0.0
    extension_seconds: float = 0.0
    extension_findings_seconds: float = 0.0

    finding_construction_seconds: float = 0.0

    #: Populated by ``run_qc()``, not by this module -- see the class
    #: docstring.
    complexity_assessment_seconds: float = 0.0

    #: Hits/misses against an optional ``FormulaPairAnalysisMemo`` (Step 5 of
    #: plan-20260910). A hit means the extension/wrapper/reference-delta
    #: analysis for this changed pair's canonical key was reused, not
    #: recomputed -- so a growing hit share directly explains a shrinking
    #: ``extension_seconds``/``wrapper_reference_seconds``.
    pair_analysis_memo_hits: int = 0
    pair_analysis_memo_misses: int = 0


@dataclass(slots=True)
class PairKeyTelemetry:
    """Aggregate-only diagnostic for the changed-formula-pair classification
    hot path (plan-20260910): whether a candidate memoization key -- the
    canonical ``(baseline_r1c1, current_r1c1)`` pair -- would be safe to
    cache. A private, in-memory map from that key to the single
    classification signature seen so far lives only for this object's
    lifetime; the public surface (``changed_pairs``, ``distinct_pair_keys``,
    ``classification_conflicts``, ``frequency_histogram()``) never exposes a
    key, a formula, a sheet name, or a coordinate -- only bounded counts.
    Passing an instance changes no finding, ordering, or evidence.
    """

    changed_pairs: int = 0
    _signatures: dict[tuple[str, str], tuple[object, ...]] = field(
        default_factory=dict, repr=False
    )
    _occurrences: dict[tuple[str, str], int] = field(default_factory=dict, repr=False)
    _conflicted_keys: set[tuple[str, str]] = field(default_factory=set, repr=False)

    def observe(
        self, base_norm: str, curr_norm: str, signature: tuple[object, ...]
    ) -> None:
        """Record one changed pair's canonical key and classification.

        ``signature`` must be a hashable tuple built only from bounded,
        already-computed classification facts (extension flag, wrapper kind/
        exactness, sorted evidence-tag values, event skeleton) -- never raw
        formula text.
        """
        self.changed_pairs += 1
        key = (base_norm, curr_norm)
        self._occurrences[key] = self._occurrences.get(key, 0) + 1
        existing = self._signatures.get(key)
        if existing is None:
            self._signatures[key] = signature
        elif existing != signature:
            self._conflicted_keys.add(key)

    @property
    def distinct_pair_keys(self) -> int:
        return len(self._signatures)

    @property
    def classification_conflicts(self) -> int:
        """Distinct pair keys that produced 2+ different classification
        signatures across their occurrences -- the exact condition that
        would make memoizing on this key unsafe (plan Criterion 15).
        """
        return len(self._conflicted_keys)

    def frequency_histogram(self) -> dict[str, int]:
        """Distinct-key counts bucketed by occurrence count -- never the
        keys, a formula, or a coordinate themselves.
        """
        buckets = {"1": 0, "2-4": 0, "5-9": 0, "10-24": 0, "25+": 0}
        for count in self._occurrences.values():
            if count == 1:
                buckets["1"] += 1
            elif count <= 4:
                buckets["2-4"] += 1
            elif count <= 9:
                buckets["5-9"] += 1
            elif count <= 24:
                buckets["10-24"] += 1
            else:
                buckets["25+"] += 1
        return buckets


@dataclass(frozen=True, slots=True)
class FormulaPairAnalysis:
    """Immutable wrapper-shape classification of one changed aligned formula
    pair, keyed by the canonical ``(baseline_r1c1, current_r1c1)`` pair.

    Every field here is derived ONLY from ``detect_formula_wrapper(base_norm,
    curr_norm)`` -- a pure function of the normalized pair itself, nothing
    else -- which is what makes reusing one analysis across every occurrence
    of the same canonical pair safe (plan-20260910, Step 5).

    Deliberately EXCLUDED: the "expected" (range-extension) flag and the
    ADDED_REFERENCE evidence tag. Both are computed from the pair's RAW
    (pre-normalization) formula text (``_differs_only_by_extension`` and
    ``formula_reference_operands`` respectively), so neither is a pure
    function of ``(base_norm, curr_norm)`` alone -- two different raw-text
    occurrences of the identical R1C1 shape can legitimately disagree on
    either. Caching them here was a real, confirmed bug (found during
    plan-20260910 Step 8 guest validation: a real-data conflict probe using
    ``PairKeyTelemetry`` found exactly one canonical key whose occurrences
    disagreed on ADDED_REFERENCE). Both are now always recomputed fresh, on
    every occurrence, regardless of this cache's hit/miss state -- see
    ``_paired_cell_findings``'s own control flow.
    """

    wrapper_kind: str | None
    wrapper_exact: bool | None
    event_key: str


#: Bounded so the memo's added RSS stays within the plan's 128 MiB budget.
#: Measured empirically (not guessed): a representative dict entry -- a
#: ``(base_r1c1, curr_r1c1)`` string-pair key plus its ``FormulaPairAnalysis``
#: -- costs ~330 bytes deep (CPython 3.11, `sys.getsizeof` recursive
#: measurement over a 5,000-entry synthetic batch with realistic R1C1 text).
#: 128 MiB / 330 bytes is ~405K entries; this cap keeps a ~2.7x safety margin
#: for dict growth/resize overhead and longer real-world R1C1 strings than
#: the synthetic sample used to measure it.
_FORMULA_PAIR_ANALYSIS_MEMO_CAP = 150_000


class FormulaPairAnalysisMemo:
    """Bounded cache from a canonical ``(baseline_r1c1, current_r1c1)`` pair
    to its ``FormulaPairAnalysis`` -- skips re-running wrapper detection
    (full tokenization plus structural argument-boundary matching) for a
    pair whose normalized text was already classified once.

    Only the wrapper-shape result is cached here: it is a pure function of
    ``(base_norm, curr_norm)`` alone (``detect_formula_wrapper`` takes only
    the normalized strings as input), so reusing it across every occurrence
    of the same canonical pair is safe regardless of the pairs' raw text.
    The "expected" (range-extension) flag and the ADDED_REFERENCE evidence
    tag are NOT cached here -- see ``FormulaPairAnalysis``'s own docstring
    for why they are unsafe to memoize on this key and are always
    recomputed fresh instead. Passing an instance changes no finding,
    ordering, or evidence, only whether the wrapper-detection cost is
    recomputed or reused.
    """

    __slots__ = ("_cap", "_values")

    def __init__(self, cap: int = _FORMULA_PAIR_ANALYSIS_MEMO_CAP) -> None:
        self._values: dict[tuple[str, str], FormulaPairAnalysis] = {}
        self._cap = cap

    def get(self, base_norm: str, curr_norm: str) -> FormulaPairAnalysis | None:
        return self._values.get((base_norm, curr_norm))

    def put(
        self, base_norm: str, curr_norm: str, analysis: FormulaPairAnalysis
    ) -> None:
        if len(self._values) < self._cap:
            self._values[(base_norm, curr_norm)] = analysis

    def __len__(self) -> int:
        return len(self._values)


#: Bounded per-region so the memo's added RSS is small and predictable; a
#: region larger than this recomputes past the cap instead of growing further.
_CURRENT_NORMALIZATION_MEMO_CAP = 200_000


class _CurrentNormalizationMemo:
    """Region-scoped current-side R1C1 cache shared by paired and consistency
    checks (Step 2 of plan-20260904-large_workbook-load-and-formula-compare.md).

    Keyed by current coordinate only: within one `diff_workbook_formulas()`
    call the current `WorkbookSnapshot` is read-only, so a given (row, column)
    always holds the same formula text, and `to_r1c1` is a pure function of
    (formula text, host row, host column) -- the cached value stays valid for
    the memo's whole lifetime. Bounded; past the cap, `normalize()` still
    returns the correct value by recomputing rather than ever skipping
    analysis or growing unbounded.
    """

    __slots__ = ("_cap", "_values")

    def __init__(self, cap: int = _CURRENT_NORMALIZATION_MEMO_CAP) -> None:
        self._values: dict[tuple[int, int], str] = {}
        self._cap = cap

    def normalize(
        self,
        row: int,
        column: int,
        formula: str,
        precomputed_r1c1: str | None = None,
    ) -> tuple[str, bool]:
        """Return ``(normalized, was_cache_hit)``.

        ``precomputed_r1c1`` -- when given (an adapter-supplied
        ``CellRecord.formula_r1c1``) -- is used directly on a cache miss
        instead of calling `to_r1c1`.
        """
        key = (row, column)
        cached = self._values.get(key)
        if cached is not None:
            return cached, True
        value = precomputed_r1c1 if precomputed_r1c1 is not None else to_r1c1(
            formula, row, column
        )
        if len(self._values) < self._cap:
            self._values[key] = value
        return value, False


# --- error scan -------------------------------------------------------------


@dataclass(frozen=True, slots=True)
class _RegionPairs:
    """Current -> baseline index maps for one confidently aligned region."""

    rows: dict[int, int]
    columns: dict[int, int]

    def baseline_cell(self, row: int, column: int) -> tuple[int, int] | None:
        base_row = self.rows.get(row)
        base_col = self.columns.get(column)
        if base_row is None or base_col is None:
            return None
        return base_row, base_col


def _region_pairs(region: RegionAlignment) -> _RegionPairs:
    return _RegionPairs(
        rows={current: baseline for baseline, current in region.rows.pairs},
        columns={current: baseline for baseline, current in region.columns.pairs},
    )


def _aligned_regions(alignment: WorkbookAlignment) -> dict[str, list[_RegionPairs]]:
    """Per-sheet baseline correspondence, excluding low-confidence regions."""
    return {
        sheet_name: [
            _region_pairs(region) for region in regions if not region.low_confidence
        ]
        for sheet_name, regions in alignment.regions.items()
    }


def _baseline_cell(
    regions: list[_RegionPairs], row: int, column: int
) -> tuple[int, int] | None:
    for region in regions:
        pair = region.baseline_cell(row, column)
        if pair is not None:
            return pair
    return None


_PROVENANCE_NOTES: dict[FindingProvenance | None, str] = {
    FindingProvenance.INHERITED: " (inherited from the baseline)",
    FindingProvenance.CHANGED: " (error state changed since the baseline)",
    FindingProvenance.NEW: " (new in this cycle)",
}

#: A deliberate `NA()` call: the standard "skip this chart point" idiom.
_NA_CALL_RE = re.compile(r"(?<![A-Za-z0-9_.])NA\(\s*\)", re.IGNORECASE)


def _deliberate_na_note(cell: CellRecord) -> str:
    if (
        cell.value == "#N/A"
        and cell.formula is not None
        and _NA_CALL_RE.search(cell.formula)
    ):
        return " [formula deliberately produces #N/A via NA()]"
    return ""


def _error_evidence(cell: CellRecord, literal: str) -> set[FindingEvidenceTag]:
    evidence: set[FindingEvidenceTag] = set()
    if cell.has_formula:
        evidence.add(FindingEvidenceTag.FORMULA_PRESENCE)
    else:
        evidence.add(FindingEvidenceTag.CACHED_VALUE_ONLY)
    if cell.formula is not None:
        evidence.add(FindingEvidenceTag.FORMULA_TEXT)
    if _deliberate_na_note(cell):
        evidence.add(FindingEvidenceTag.EXPLICIT_NA)
    if literal in _STRUCTURAL_ERROR_LITERALS:
        evidence.add(FindingEvidenceTag.STRUCTURAL_ERROR)
    return evidence
_OUTLIER_NOTES: dict[FindingProvenance | None, str] = {
    FindingProvenance.INHERITED: " (the baseline cell already deviated this way)",
    FindingProvenance.HISTORICAL_PATTERN: (
        " (new-cycle cell reusing a pattern already present in the historical range)"
    ),
    FindingProvenance.NEW: " (novel deviation in this cycle)",
}


def _outlier_provenance(
    base_sheet: SheetSnapshot,
    baseline_cell: tuple[int, int] | None,
    pattern: str,
    historical_patterns: set[str],
) -> FindingProvenance:
    if baseline_cell is None:
        return (
            FindingProvenance.HISTORICAL_PATTERN
            if pattern in historical_patterns
            else FindingProvenance.NEW
        )
    cell = base_sheet.cells.get(baseline_cell)
    if cell is None or not cell.has_formula:
        return FindingProvenance.NEW
    base_row, base_col = baseline_cell
    return (
        FindingProvenance.INHERITED
        if _normalize_formula(cell, base_row, base_col) == pattern
        else FindingProvenance.NEW
    )


def _error_literals(formula: str | None) -> tuple[str, ...]:
    if formula is None:
        return ()
    return tuple(sorted(err for err in ERROR_LITERALS if err in formula))


def _provenance(baseline_state: object, current_state: object) -> FindingProvenance:
    if not baseline_state:
        return FindingProvenance.NEW
    return (
        FindingProvenance.INHERITED
        if baseline_state == current_state
        else FindingProvenance.CHANGED
    )


def _ignored(profile: SheetProfile | None, row: int, column: int) -> bool:
    return bool(
        profile is not None
        and cell_in_ranges(row, column, profile.ignore_ranges)
    )


#: A single error literal dominating one column is one systemic population.
_COLUMNAR_ERROR_MIN_CELLS = 20
_COLUMNAR_ERROR_MIN_SHARE = 0.5
#: In huge columns the share test dilutes; this many identical literals in one
#: column is systemic evidence on its own (e.g. a lookup with unmatched rows).
_COLUMNAR_ERROR_MASS_CELLS = 200
#: Structural breakage is never a tolerated data state: a column full of these
#: is one incident, but a broken one — it groups without leaving CRITICAL.
_STRUCTURAL_ERROR_LITERALS = frozenset({"#REF!", "#NAME?", "#NULL!"})


def _error_population_event_key(
    artifact: str, sheet: str, column: int, literal: str
) -> str:
    payload = "\0".join(
        (artifact, sheet, get_column_letter(column), literal)
    ).encode("utf-8")
    return f"error-population:{hashlib.sha256(payload).hexdigest()[:12]}"


def _run_evidence(rows: list[int]) -> tuple[int, int]:
    run_count = 0
    longest_run = 0
    previous: int | None = None
    current_run = 0
    for row in sorted(rows):
        if previous is None or row != previous + 1:
            run_count += 1
            current_run = 1
        else:
            current_run += 1
        longest_run = max(longest_run, current_run)
        previous = row
    return run_count, longest_run


def _mark_columnar_error_populations(
    sheet: SheetSnapshot,
    saved_error_findings: list[tuple[Finding, int, int, str, CellRecord]],
) -> None:
    """Group proven same-column/literal populations without assuming benignity."""
    if not saved_error_findings:
        return
    populated: Counter[int] = Counter()
    for (_, col), cell in sheet.cells.items():
        if cell.value is not None:
            populated[col] += 1
    groups: dict[tuple[int, str], list[tuple[Finding, int, CellRecord]]] = {}
    for finding, row, col, literal, cell in saved_error_findings:
        groups.setdefault((col, literal), []).append((finding, row, cell))
    for (col, literal), records in groups.items():
        members = [finding for finding, _, _ in records]
        count = len(members)
        column_population = populated.get(col, 0)
        share = count / column_population if column_population else 0.0
        concentrated = (
            count >= _COLUMNAR_ERROR_MIN_CELLS
            and column_population > 0
            and share >= _COLUMNAR_ERROR_MIN_SHARE
        )
        run_count, longest_run = _run_evidence([row for _, row, _ in records])
        contiguous = (
            count >= _COLUMNAR_ERROR_MIN_CELLS and longest_run / count >= 0.8
        )
        mass = count >= _COLUMNAR_ERROR_MASS_CELLS
        if not (concentrated or contiguous or mass):
            continue
        population_evidence: set[FindingEvidenceTag] = set()
        if concentrated:
            population_evidence.add(FindingEvidenceTag.CONCENTRATED_POPULATION)
        if contiguous:
            population_evidence.add(FindingEvidenceTag.CONTIGUOUS_POPULATION)
        if mass and not concentrated:
            population_evidence.add(FindingEvidenceTag.SPARSE_MASS_POPULATION)
        formula_count = sum(1 for _, _, cell in records if cell.has_formula)
        note = (
            f" [error population: {count} of {column_population} populated cells "
            f"({share:.1%}); formula presence {formula_count}/{count}; "
            f"runs {run_count}, longest {longest_run}]"
        )
        event_key = _error_population_event_key("excel", sheet.name, col, literal)
        for finding in members:
            finding.subtype = FindingSubtype.COLUMNAR_ERROR_POPULATION
            finding.event_key = event_key
            finding.evidence_tags.update(population_evidence)
            finding.message += note


def _error_findings(
    baseline: WorkbookSnapshot,
    current: WorkbookSnapshot,
    alignment: WorkbookAlignment,
    profile: DeliverableProfile | None,
    *,
    cycle: bool,
    telemetry: FormulaComparisonTelemetry | None = None,
) -> list[Finding]:
    findings = []
    scan_start = time.perf_counter()
    scanned_cells = 0
    aligned = _aligned_regions(alignment) if cycle else {}
    for sheet in current.sheets:
        sheet_profile = profile.sheet_profile(sheet.name) if profile is not None else None
        regions = aligned.get(sheet.name, [])
        base_sheet = baseline.sheet(sheet.name) if regions else None
        saved_error_findings: list[tuple[Finding, int, int, str, CellRecord]] = []
        for (row, col), cell in sorted(sheet.cells.items()):
            if _ignored(sheet_profile, row, col):
                continue
            scanned_cells += 1
            location = _ref(row, col)
            base_cell: CellRecord | None = None
            aligned_baseline = False
            if base_sheet is not None:
                pair = _baseline_cell(regions, row, col)
                aligned_baseline = pair is not None
                if pair is not None:
                    base_cell = base_sheet.cells.get(pair)
            if cell.is_error:
                provenance = (
                    _provenance(
                        base_cell.value if base_cell is not None and base_cell.is_error
                        else None,
                        cell.value,
                    )
                    if aligned_baseline
                    else None
                )
                finding = Finding(
                    artifact="excel",
                    finding_class=FindingClass.FORMULA_ERROR,
                    provenance=provenance,
                    evidence_tags=_error_evidence(cell, str(cell.value)),
                    sheet=sheet.name,
                    location=location,
                    current_value=display_cell_value(cell.value),
                    message=(
                        f"{sheet.name}!{location}: error value {cell.value}"
                        + _PROVENANCE_NOTES.get(provenance, "")
                        + _deliberate_na_note(cell)
                    ),
                )
                findings.append(finding)
                saved_error_findings.append((finding, row, col, str(cell.value), cell))
            current_literals = _error_literals(cell.formula)
            if cell.formula and current_literals:
                provenance = (
                    _provenance(
                        _error_literals(
                            base_cell.formula if base_cell is not None else None
                        ),
                        current_literals,
                    )
                    if aligned_baseline
                    else None
                )
                findings.append(
                    Finding(
                        artifact="excel",
                        finding_class=FindingClass.FORMULA_ERROR,
                        provenance=provenance,
                        evidence_tags={
                            FindingEvidenceTag.FORMULA_TEXT,
                            FindingEvidenceTag.FORMULA_PRESENCE,
                            *(
                                {FindingEvidenceTag.STRUCTURAL_ERROR}
                                if _STRUCTURAL_ERROR_LITERALS.intersection(
                                    current_literals
                                )
                                else set()
                            ),
                        },
                        sheet=sheet.name,
                        location=location,
                        current_value=cell.formula,
                        message=(
                            f"{sheet.name}!{location}: formula contains an error "
                            f"reference ({cell.formula})"
                            + _PROVENANCE_NOTES.get(provenance, "")
                        ),
                    )
                )
        _mark_columnar_error_populations(sheet, saved_error_findings)
    if telemetry is not None:
        telemetry.error_scan_seconds += time.perf_counter() - scan_start
        telemetry.error_scan_cells += scanned_cells
    return findings


# --- paired-cell checks ------------------------------------------------------


def _paired_cell_findings(
    base_sheet: SheetSnapshot,
    curr_sheet: SheetSnapshot,
    region: RegionAlignment,
    *,
    compare_text: bool,
    strict_text: bool,
    profile: SheetProfile | None,
    telemetry: FormulaComparisonTelemetry | None = None,
    current_memo: _CurrentNormalizationMemo | None = None,
    candidate_sink: CandidateSpill | None = None,
    pair_key_telemetry: PairKeyTelemetry | None = None,
    pair_analysis_memo: FormulaPairAnalysisMemo | None = None,
) -> list[Finding]:
    findings = []
    sheet_name = curr_sheet.name
    traversal_start = time.perf_counter()
    for (base_row, base_col), (curr_row, curr_col) in region.cell_pairs():
        base_cell = base_sheet.cells.get((base_row, base_col))
        curr_cell = curr_sheet.cells.get((curr_row, curr_col))
        if _ignored(profile, curr_row, curr_col):
            continue
        base_has_formula = base_cell is not None and base_cell.has_formula
        curr_has_formula = curr_cell is not None and curr_cell.has_formula
        if not base_has_formula and not curr_has_formula:
            continue
        location = _ref(curr_row, curr_col)

        if base_has_formula and not curr_has_formula:
            baseline_value = (
                base_cell.formula
                if base_cell is not None and base_cell.formula is not None
                else "formula record"
            )
            construct_start = time.perf_counter()
            if curr_cell is not None and curr_cell.value is not None:
                findings.append(
                    Finding(
                        artifact="excel",
                        finding_class=FindingClass.FORMULA_HARDCODED,
                        sheet=sheet_name,
                        location=location,
                        baseline_location=_ref(base_row, base_col),
                        baseline_value=baseline_value,
                        current_value=display_cell_value(curr_cell.value),
                        message=(
                            f"{sheet_name}!{location}: formula replaced by "
                            f"hardcoded value {curr_cell.value}"
                        ),
                    )
                )
            else:
                if excel_blank_allowed(
                    curr_sheet,
                    profile,
                    curr_row,
                    curr_col,
                ):
                    continue
                findings.append(
                    Finding(
                        artifact="excel",
                        finding_class=FindingClass.FORMULA_REMOVED,
                        sheet=sheet_name,
                        location=location,
                        baseline_location=_ref(base_row, base_col),
                        baseline_value=baseline_value,
                        message=f"{sheet_name}!{location}: formula cleared or removed",
                    )
                )
            if telemetry is not None:
                telemetry.finding_construction_seconds += (
                    time.perf_counter() - construct_start
                )
            continue

        if base_has_formula and curr_has_formula and compare_text:
            base_formula = base_cell.formula if base_cell is not None else None
            curr_formula = curr_cell.formula if curr_cell is not None else None
            if base_formula is None or curr_formula is None:
                if strict_text:
                    logger.warning(
                        "%s!%s: formula text capability contradicted cell data",
                        sheet_name,
                        location,
                    )
                # Under partial XLSB coverage, a has_formula cell with no merged
                # text is an expected, already-counted coverage gap, not a bug.
                continue
            if telemetry is not None:
                telemetry.paired_pairs_considered += 1
            if (
                base_row == curr_row
                and base_col == curr_col
                and base_formula == curr_formula
            ):
                # Identical text at the identical host position always
                # normalizes identically -- to_r1c1 is a pure function of
                # (formula text, host row, host column) -- so this pair can
                # never produce a finding. Skip normalization entirely rather
                # than normalizing both sides only to discover they match.
                if telemetry is not None:
                    telemetry.exact_text_same_host_pairs += 1
                continue
            norm_start = time.perf_counter()
            base_r1c1 = base_cell.formula_r1c1 if base_cell is not None else None
            base_norm = (
                base_r1c1
                if base_r1c1 is not None
                else to_r1c1(base_formula, base_row, base_col)
            )
            calls = 0 if base_r1c1 is not None else 1
            if current_memo is not None:
                curr_r1c1 = curr_cell.formula_r1c1 if curr_cell is not None else None
                curr_norm, memo_hit = current_memo.normalize(
                    curr_row, curr_col, curr_formula, curr_r1c1
                )
                if telemetry is not None:
                    if memo_hit:
                        telemetry.memo_hits += 1
                    else:
                        telemetry.memo_misses += 1
                calls += 0 if memo_hit else 1
            else:
                curr_r1c1 = curr_cell.formula_r1c1 if curr_cell is not None else None
                curr_norm = (
                    curr_r1c1
                    if curr_r1c1 is not None
                    else to_r1c1(curr_formula, curr_row, curr_col)
                )
                calls += 0 if curr_r1c1 is not None else 1
            if telemetry is not None:
                telemetry.normalization_seconds += time.perf_counter() - norm_start
                telemetry.normalization_calls += calls
            if base_norm == curr_norm:
                continue
            # `expected` and the ADDED_REFERENCE evidence tag are both
            # computed from RAW (pre-normalization) formula text, so they
            # are NOT a pure function of the (base_norm, curr_norm) cache
            # key -- two different occurrences of the identical R1C1 shape
            # can legitimately disagree on either (confirmed empirically on
            # real data: plan-20260910 Step 8 guest validation found exactly
            # one such key with a genuine ADDED_REFERENCE conflict across
            # its occurrences). Only wrapper_kind/wrapper_exact/event_key
            # are a pure function of the key (detect_formula_wrapper takes
            # only base_norm/curr_norm as input) and are safe to memoize;
            # `expected` and the reference-added check are always
            # recomputed fresh below, on every occurrence, regardless of
            # cache hit/miss.
            ext_start = time.perf_counter()
            expected = _differs_only_by_extension(base_formula, curr_formula)
            if telemetry is not None:
                telemetry.extension_seconds += time.perf_counter() - ext_start
            evidence_tags: set[FindingEvidenceTag] = set()
            wrapper_start = time.perf_counter()
            if expected:
                wrapper_kind = None
                wrapper_exact = None
                event_key = ""
            else:
                cached_analysis = (
                    pair_analysis_memo.get(base_norm, curr_norm)
                    if pair_analysis_memo is not None
                    else None
                )
                if cached_analysis is not None:
                    if telemetry is not None:
                        telemetry.pair_analysis_memo_hits += 1
                    wrapper_kind = cached_analysis.wrapper_kind
                    wrapper_exact = cached_analysis.wrapper_exact
                    event_key = cached_analysis.event_key
                else:
                    if pair_analysis_memo is not None and telemetry is not None:
                        telemetry.pair_analysis_memo_misses += 1
                    wrapper = detect_formula_wrapper(base_norm, curr_norm)
                    wrapper_kind = wrapper.kind if wrapper is not None else None
                    wrapper_exact = wrapper.exact if wrapper is not None else None
                    event_key = (
                        f"formula-wrapper:{wrapper.kind}:{wrapper.skeleton_key}"
                        if wrapper is not None
                        else ""
                    )
                    if pair_analysis_memo is not None:
                        pair_analysis_memo.put(
                            base_norm,
                            curr_norm,
                            FormulaPairAnalysis(
                                wrapper_kind=wrapper_kind,
                                wrapper_exact=wrapper_exact,
                                event_key=event_key,
                            ),
                        )
                if wrapper_kind is not None:
                    assert wrapper_exact is not None
                    evidence_tags.add(
                        FindingEvidenceTag.EXACT_WRAPPER
                        if wrapper_exact
                        else FindingEvidenceTag.SHAPE_WRAPPER
                    )
            try:
                baseline_references = {
                    operand.value.casefold()
                    for operand in formula_reference_operands(base_formula)
                }
                current_references = {
                    operand.value.casefold()
                    for operand in formula_reference_operands(curr_formula)
                }
            except (TokenizerError, TypeError, ValueError) as exc:
                logger.warning(
                    "formula-reference-tag-unavailable %s!%s: %s",
                    sheet_name,
                    location,
                    type(exc).__name__,
                )
            else:
                if current_references - baseline_references:
                    evidence_tags.add(FindingEvidenceTag.ADDED_REFERENCE)
            if telemetry is not None:
                telemetry.wrapper_reference_seconds += (
                    time.perf_counter() - wrapper_start
                )
            if expected:
                wording = "formula range extended with new-cycle data"
            elif wrapper_kind is not None:
                assert wrapper_exact is not None
                wording = "formula logic changed" + _WRAPPER_NOTES[
                    (wrapper_kind, wrapper_exact)
                ]
            else:
                wording = "formula logic changed"
            expected_reason = (
                FindingExpectedReason.CADENCE_EXTENSION if expected else None
            )
            subtype = (
                _WRAPPER_SUBTYPES[wrapper_kind] if wrapper_kind is not None else None
            )
            if pair_key_telemetry is not None:
                # Observe only the portion FormulaPairAnalysisMemo actually
                # caches (wrapper_kind/wrapper_exact/event_key) -- `expected`
                # and the ADDED_REFERENCE tag are always freshly computed
                # now (see the comment above), so including them here would
                # make every real workload show spurious "conflicts" for
                # values that were never claimed to be cacheable.
                pair_key_telemetry.observe(
                    base_norm,
                    curr_norm,
                    (wrapper_kind, wrapper_exact, event_key),
                )
            message = f"{sheet_name}!{location}: {wording}"
            construct_start = time.perf_counter()
            if candidate_sink is not None:
                candidate = Finding(
                    artifact="excel",
                    finding_class=FindingClass.FORMULA_LOGIC_CHANGED,
                    expected_reason=expected_reason,
                    subtype=subtype,
                    event_key=event_key,
                    evidence_tags=evidence_tags,
                    sheet=sheet_name,
                    location=location,
                    baseline_location=_ref(base_row, base_col),
                    baseline_value=base_formula,
                    current_value=curr_formula,
                    message=message,
                )
                candidate_sink.add(
                    candidate,
                    shape_before=hashlib.sha256(
                        base_norm.encode("utf-8")
                    ).hexdigest(),
                    shape_after=hashlib.sha256(
                        curr_norm.encode("utf-8")
                    ).hexdigest(),
                )
            else:
                findings.append(
                    Finding(
                        artifact="excel",
                        finding_class=FindingClass.FORMULA_LOGIC_CHANGED,
                        expected_reason=expected_reason,
                        subtype=subtype,
                        event_key=event_key,
                        evidence_tags=evidence_tags,
                        sheet=sheet_name,
                        location=location,
                        baseline_location=_ref(base_row, base_col),
                        baseline_value=base_formula,
                        current_value=curr_formula,
                        message=message,
                    )
                )
            if telemetry is not None:
                telemetry.finding_construction_seconds += (
                    time.perf_counter() - construct_start
                )
    if telemetry is not None:
        telemetry.paired_traversal_seconds += time.perf_counter() - traversal_start
    return findings


# --- run helpers (formula pattern along the data axis) -----------------------


def _run_axes(
    region: RegionAlignment,
) -> tuple[list[int], list[int]] | None:
    """(run positions, cell positions) as (fixed axis indices, moving indices)."""
    current = region.current
    if current.orientation == "long":
        data_rows = sorted([c for _, c in region.rows.pairs] + region.rows.growth)
        columns = sorted([c for _, c in region.columns.pairs] + region.columns.growth)
        return columns, data_rows
    if current.orientation == "wide":
        header = current.header_row or current.min_row
        data_rows = sorted(
            [c for _, c in region.rows.pairs if c != header] + region.rows.growth
        )
        columns = sorted([c for _, c in region.columns.pairs] + region.columns.growth)
        return data_rows, columns
    return None


def _run_cells(
    sheet: SheetSnapshot, region: RegionAlignment, fixed: int, moving: list[int]
) -> list[tuple[int, int, CellRecord]]:
    is_long = region.current.orientation == "long"
    out = []
    for position in moving:
        key = (position, fixed) if is_long else (fixed, position)
        cell = sheet.cells.get(key)
        if cell is not None:
            out.append((key[0], key[1], cell))
    return out


def _consistency_findings(
    base_sheet: SheetSnapshot,
    curr_sheet: SheetSnapshot,
    region: RegionAlignment,
    profile: SheetProfile | None,
    *,
    cycle: bool,
    telemetry: FormulaComparisonTelemetry | None = None,
    current_memo: _CurrentNormalizationMemo | None = None,
) -> list[Finding]:
    axes = _run_axes(region)
    if axes is None:
        return []
    consistency_start = time.perf_counter()
    run_positions, cell_positions = axes
    pairs = _region_pairs(region)
    findings = []
    for fixed in run_positions:
        cells = _run_cells(curr_sheet, region, fixed, cell_positions)
        formula_cells = [
            (row, col, cell)
            for row, col, cell in cells
            if cell.has_formula and not _ignored(profile, row, col)
        ]
        if len(formula_cells) < _RUN_MIN_CELLS:
            continue
        if any(cell.formula is None for _, _, cell in formula_cells):
            # Partial XLSB coverage cannot prove this run's dominant pattern;
            # a complete run always has merged text for every formula cell,
            # so this is a no-op once coverage is complete.
            continue
        norm_start = time.perf_counter()
        normalized: dict[tuple[int, int], str] = {}
        calls = 0
        for row, col, cell in formula_cells:
            formula = cell.formula or ""
            if current_memo is not None:
                value, memo_hit = current_memo.normalize(
                    row, col, formula, cell.formula_r1c1
                )
                if telemetry is not None:
                    if memo_hit:
                        telemetry.memo_hits += 1
                    else:
                        telemetry.memo_misses += 1
                calls += 0 if memo_hit else 1
            else:
                value = (
                    cell.formula_r1c1
                    if cell.formula_r1c1 is not None
                    else to_r1c1(formula, row, col)
                )
                calls += 0 if cell.formula_r1c1 is not None else 1
            normalized[(row, col)] = value
        if telemetry is not None:
            telemetry.consistency_normalization_seconds += (
                time.perf_counter() - norm_start
            )
            telemetry.consistency_normalization_calls += calls
        patterns = Counter(normalized.values())
        dominant, dominant_count = patterns.most_common(1)[0]
        if dominant_count / len(formula_cells) <= _RUN_DOMINANCE:
            continue
        historical_patterns = {
            pattern
            for (row, col), pattern in normalized.items()
            if pairs.baseline_cell(row, col) is not None
        }
        for row, col, cell in formula_cells:
            pattern = normalized[(row, col)]
            if pattern == dominant:
                continue
            location = _ref(row, col)
            provenance = (
                _outlier_provenance(
                    base_sheet,
                    pairs.baseline_cell(row, col),
                    pattern,
                    historical_patterns,
                )
                if cycle
                else None
            )
            construct_start = time.perf_counter()
            findings.append(
                Finding(
                    artifact="excel",
                    finding_class=FindingClass.FORMULA_INCONSISTENT,
                    provenance=provenance,
                    sheet=curr_sheet.name,
                    location=location,
                    baseline_value=dominant,
                    current_value=cell.formula,
                    message=(
                        f"{curr_sheet.name}!{location}: formula deviates from the "
                        f"dominant pattern of its range ({dominant})"
                        + _OUTLIER_NOTES.get(provenance, "")
                    ),
                )
            )
            if telemetry is not None:
                telemetry.finding_construction_seconds += (
                    time.perf_counter() - construct_start
                )
    if telemetry is not None:
        telemetry.consistency_seconds += time.perf_counter() - consistency_start
    return findings


def _extension_findings(
    curr_sheet: SheetSnapshot,
    region: RegionAlignment,
    profile: SheetProfile | None,
    *,
    telemetry: FormulaComparisonTelemetry | None = None,
) -> list[Finding]:
    """Growth rows/columns must carry the formula pattern of their run."""
    current = region.current
    findings = []
    ext_start = time.perf_counter()

    def check(
        growth: list[int], run_positions: list[int], paired: list[int], *, rows_grow: bool
    ) -> None:
        for fixed in run_positions:
            paired_cells = [
                curr_sheet.cells.get((p, fixed) if rows_grow else (fixed, p)) for p in paired
            ]
            populated = [c for c in paired_cells if c is not None]
            with_formula = [c for c in populated if c.has_formula]
            if len(populated) < 2 or len(with_formula) / len(populated) < _FORMULAIC_SHARE:
                continue
            for g in growth:
                key = (g, fixed) if rows_grow else (fixed, g)
                if _ignored(profile, *key):
                    continue
                cell = curr_sheet.cells.get(key)
                if cell is not None and cell.has_formula:
                    continue
                if excel_blank_allowed(curr_sheet, profile, *key):
                    continue
                location = _ref(*key)
                findings.append(
                    Finding(
                        artifact="excel",
                        finding_class=FindingClass.FORMULA_NOT_EXTENDED,
                        sheet=curr_sheet.name,
                        location=location,
                        current_value=None if cell is None else display_cell_value(cell.value),
                        message=(
                            f"{curr_sheet.name}!{location}: new-cycle cell is missing "
                            "the formula used by its historical range"
                        ),
                    )
                )

    if current.orientation == "long" and region.rows.growth:
        columns = sorted([c for _, c in region.columns.pairs] + region.columns.growth)
        paired_rows = [c for _, c in region.rows.pairs]
        check(region.rows.growth, columns, paired_rows, rows_grow=True)
    elif current.orientation == "wide" and region.columns.growth:
        header = current.header_row or current.min_row
        data_rows = [c for _, c in region.rows.pairs if c != header] + region.rows.growth
        paired_cols = [c for _, c in region.columns.pairs]
        check(region.columns.growth, sorted(data_rows), paired_cols, rows_grow=False)
    if telemetry is not None:
        telemetry.extension_findings_seconds += time.perf_counter() - ext_start
    return findings


# --- entry point --------------------------------------------------------------


def formula_text_compatible(
    baseline: WorkbookSnapshot, current: WorkbookSnapshot
) -> bool:
    """Whether formula text can be compared without crossing engine dialects."""
    return bool(
        baseline.formulas_available
        and current.formulas_available
        and baseline.formula_source == current.formula_source
    )


def formula_text_comparable(
    baseline: WorkbookSnapshot, current: WorkbookSnapshot
) -> bool:
    """Whether any paired cell's formula text may be compared at all.

    Complete coverage on both sides (the historical, strict rule above) is
    always comparable. Partial XLSB coverage is comparable too, provided
    both sides still share the same non-empty formula-text provenance;
    `_paired_cell_findings` skips any individual pair missing merged text on
    either side, and `_consistency_findings` skips any run containing one.
    OOXML sources never populate `formula_text_coverage` (it stays at its
    default ``"none"``); they are already covered by the first branch below
    because `formulas_available` is always True for them.
    """
    if formula_text_compatible(baseline, current):
        return True
    return bool(
        baseline.formula_source
        and baseline.formula_source == current.formula_source
        and baseline.formula_text_coverage.state != "none"
        and current.formula_text_coverage.state != "none"
    )


def diff_workbook_formulas(
    baseline: WorkbookSnapshot,
    current: WorkbookSnapshot,
    alignment: WorkbookAlignment,
    profile: DeliverableProfile | None = None,
    *,
    cycle: bool = True,
    cancellation_token: CancellationToken | None = None,
    telemetry: FormulaComparisonTelemetry | None = None,
    candidate_sink: CandidateSpill | None = None,
    pair_key_telemetry: PairKeyTelemetry | None = None,
    pair_analysis_memo: FormulaPairAnalysisMemo | None = None,
) -> list[Finding]:
    findings = _error_findings(
        baseline, current, alignment, profile, cycle=cycle, telemetry=telemetry
    )
    presence_pair = bool(
        baseline.formula_presence_available and current.formula_presence_available
    )
    compare_text = formula_text_comparable(baseline, current)
    strict_text = formula_text_compatible(baseline, current)
    if not presence_pair:
        logger.info(
            "formula presence unavailable (%s/%s): formula QC limited to error values",
            baseline.file_format,
            current.file_format,
        )
        return findings
    if not compare_text:
        logger.info(
            "compatible formula text unavailable (%s/%s): semantic formula QC skipped",
            baseline.formula_source,
            current.formula_source,
        )
    for sheet_name, regions in alignment.regions.items():
        check_cancelled(cancellation_token)
        base_sheet = baseline.sheet(sheet_name)
        curr_sheet = current.sheet(sheet_name)
        sheet_profile = profile.sheet_profile(sheet_name) if profile is not None else None
        for region in regions:
            check_cancelled(cancellation_token)
            if region.low_confidence:
                continue
            # Fresh per region and discarded at the end of its iteration --
            # bounded scope keeps its added RSS small and predictable.
            current_memo = _CurrentNormalizationMemo()
            findings.extend(
                _paired_cell_findings(
                    base_sheet,
                    curr_sheet,
                    region,
                    compare_text=compare_text,
                    strict_text=strict_text,
                    profile=sheet_profile,
                    telemetry=telemetry,
                    current_memo=current_memo,
                    candidate_sink=candidate_sink,
                    pair_key_telemetry=pair_key_telemetry,
                    pair_analysis_memo=pair_analysis_memo,
                )
            )
            if compare_text:
                findings.extend(
                    _consistency_findings(
                        base_sheet,
                        curr_sheet,
                        region,
                        sheet_profile,
                        cycle=cycle,
                        telemetry=telemetry,
                        current_memo=current_memo,
                    )
                )
            findings.extend(
                _extension_findings(
                    curr_sheet, region, sheet_profile, telemetry=telemetry
                )
            )
    return findings


def formula_token_diff(
    baseline_formula: str | None,
    current_formula: str | None,
    baseline_location: str,
    current_location: str,
) -> tuple[FormulaDiffSegment, ...]:
    """Compute a token-level diff of two original formula texts.

    Returns an empty tuple on any parse/validation problem. The returned
    segments are coalesced by kind and always prefixed by a single '=' equal
    segment when non-empty.
    """
    try:
        if not baseline_formula or not current_formula:
            return ()
        if not (baseline_formula.startswith("=") and current_formula.startswith("=")):
            return ()
        try:
            base_row, base_col = coordinate_to_tuple(baseline_location)
            curr_row, curr_col = coordinate_to_tuple(current_location)
        except Exception:
            return ()
        try:
            base_tokens = tokenize_formula(baseline_formula)
            curr_tokens = tokenize_formula(current_formula)
        except Exception:
            return ()

        def _keys_texts(tokens, host_row, host_col):
            keys: list[tuple[str, str, str]] = []
            texts: list[str] = []
            for t in tokens:
                if t.type == "OPERAND" and t.subtype == "RANGE":
                    normalized = _range_token_to_r1c1(t.value, host_row, host_col)
                    key = normalized.casefold()
                    keys.append((t.type, t.subtype, key))
                    texts.append(normalized)
                else:
                    text = t.value
                    key = (
                        text
                        if t.type == "OPERAND" and t.subtype == "TEXT"
                        else text.casefold()
                    )
                    keys.append((t.type, t.subtype, key))
                    texts.append(text)
            return keys, texts

        base_keys, base_texts = _keys_texts(base_tokens, base_row, base_col)
        curr_keys, curr_texts = _keys_texts(curr_tokens, curr_row, curr_col)

        matcher = difflib.SequenceMatcher(a=base_keys, b=curr_keys, autojunk=False)
        segments: list[FormulaDiffSegment] = []
        # prefix one leading '=' equal segment
        segments.append(FormulaDiffSegment("=", FormulaDiffKind.EQUAL))
        for tag, i1, i2, j1, j2 in matcher.get_opcodes():
            if tag == "equal":
                text = "".join(curr_texts[j1:j2])
                if text:
                    segments.append(FormulaDiffSegment(text, FormulaDiffKind.EQUAL))
            elif tag == "delete":
                text = "".join(base_texts[i1:i2])
                if text:
                    segments.append(FormulaDiffSegment(text, FormulaDiffKind.REMOVED))
            elif tag == "insert":
                text = "".join(curr_texts[j1:j2])
                if text:
                    segments.append(FormulaDiffSegment(text, FormulaDiffKind.ADDED))
            elif tag == "replace":
                text_b = "".join(base_texts[i1:i2])
                text_c = "".join(curr_texts[j1:j2])
                if text_b:
                    segments.append(FormulaDiffSegment(text_b, FormulaDiffKind.REMOVED))
                if text_c:
                    segments.append(FormulaDiffSegment(text_c, FormulaDiffKind.ADDED))

        # coalesce adjacent segments of same kind
        coalesced: list[FormulaDiffSegment] = []
        for seg in segments:
            if coalesced and coalesced[-1].kind is seg.kind:
                coalesced[-1] = FormulaDiffSegment(
                    coalesced[-1].text + seg.text,
                    coalesced[-1].kind,
                )
            else:
                coalesced.append(seg)
        return tuple(coalesced)
    except Exception:
        return ()
