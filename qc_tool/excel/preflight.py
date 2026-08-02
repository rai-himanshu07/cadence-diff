"""Intrinsic QC for one current workbook, without a historical baseline."""

from collections import Counter
from dataclasses import dataclass, field
from itertools import pairwise

from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import range_boundaries

from qc_tool.availability import (
    availability_coverage,
    cell_in_ranges,
    excel_availability_issues,
    excel_blank_allowed,
)
from qc_tool.config.profile import CadenceBand, DeliverableProfile, SheetProfile
from qc_tool.coverage import CoverageItem, CoverageState
from qc_tool.excel.align import align_workbooks
from qc_tool.excel.charts import annotate_chart_impacts
from qc_tool.excel.controls import evaluate_controls
from qc_tool.excel.dependency import (
    DependencyGraph,
    annotate_impacts,
    build_dependency_graph,
    limit_impacts,
)
from qc_tool.excel.formulas import diff_workbook_formulas
from qc_tool.excel.interaction import (
    conditional_style_coverage,
    interaction_rule_coverage,
)
from qc_tool.excel.periods import Period, is_period_after, parse_period
from qc_tool.excel.references import ReferenceStatus, resolve_reference
from qc_tool.excel.regions import TableRegion, detect_regions
from qc_tool.excel.workbook_risks import workbook_risk_findings
from qc_tool.findings import Finding, FindingClass
from qc_tool.io.model import SheetSnapshot, WorkbookSnapshot, display_cell_value
from qc_tool.progress import CancellationToken, check_cancelled

_FORMULA_RUN_MIN = 4
_FORMULA_SHARE = 0.60


@dataclass(slots=True)
class ExcelPreflightResult:
    findings: list[Finding] = field(default_factory=list)
    coverage: list[CoverageItem] = field(default_factory=list)
    dependency_graph: DependencyGraph | None = None


def _ref(row: int, col: int) -> str:
    return f"{get_column_letter(col)}{row}"


def _formula_gap_findings(
    sheet: SheetSnapshot,
    region: TableRegion,
    profile: SheetProfile | None,
) -> list[Finding]:
    if region.orientation == "long":
        moving = list(range((region.header_row or region.min_row) + 1, region.max_row + 1))
        fixed = list(range(region.min_col, region.max_col + 1))

        def keys(fixed_pos: int, moving_pos: int) -> tuple[int, int]:
            return moving_pos, fixed_pos

    elif region.orientation == "wide":
        moving = list(range(region.min_col + 1, region.max_col + 1))
        fixed = list(range((region.header_row or region.min_row) + 1, region.max_row + 1))

        def keys(fixed_pos: int, moving_pos: int) -> tuple[int, int]:
            return fixed_pos, moving_pos

    else:
        return []

    findings: list[Finding] = []
    for fixed_pos in fixed:
        cells = [sheet.cells.get(keys(fixed_pos, moving_pos)) for moving_pos in moving]
        formula_count = sum(cell is not None and cell.has_formula for cell in cells)
        if formula_count < _FORMULA_RUN_MIN or formula_count / max(len(cells), 1) < _FORMULA_SHARE:
            continue
        for moving_pos, cell in zip(moving, cells, strict=True):
            if cell is not None and cell.has_formula:
                continue
            row, col = keys(fixed_pos, moving_pos)
            if profile is not None and cell_in_ranges(
                row,
                col,
                profile.ignore_ranges,
            ):
                continue
            if excel_blank_allowed(sheet, profile, row, col):
                continue
            findings.append(
                Finding(
                    artifact="excel",
                    finding_class=FindingClass.FORMULA_MISSING,
                    sheet=sheet.name,
                    location=_ref(row, col),
                    current_value=None if cell is None else display_cell_value(cell.value),
                    message=(
                        f"{sheet.name}!{_ref(row, col)}: formula missing inside "
                        "a formula-dense range"
                    ),
                )
            )
    return findings


PeriodAxisItem = tuple[str, Period, int, int]


def _period_axis(sheet: SheetSnapshot, region: TableRegion) -> list[PeriodAxisItem]:
    raw: list[tuple[object, int, int]] = []
    if region.orientation == "wide":
        row = region.header_row or region.min_row
        raw = [
            (sheet.cells[(row, col)].value, row, col)
            for col in range(region.min_col + 1, region.max_col + 1)
            if (row, col) in sheet.cells
        ]
    elif region.orientation == "long":
        period_col = region.key_col or region.min_col
        raw = [
            (sheet.cells[(row, period_col)].value, row, period_col)
            for row in range((region.header_row or region.min_row) + 1, region.max_row + 1)
            if (row, period_col) in sheet.cells
        ]
    result: list[PeriodAxisItem] = []
    for value, row, column in raw:
        period = parse_period(value)
        if period is not None:
            result.append((display_cell_value(value), period, row, column))
    return result


def _auto_period_bands(axis: list[PeriodAxisItem]) -> list[list[PeriodAxisItem]]:
    bands: list[list[PeriodAxisItem]] = []
    for item in axis:
        if not bands:
            bands.append([item])
            continue
        previous = bands[-1][-1]
        same_kind = previous[1].kind == item[1].kind
        adjacent = abs(previous[2] - item[2]) + abs(previous[3] - item[3]) == 1
        if same_kind and adjacent:
            bands[-1].append(item)
        else:
            bands.append([item])
    return bands


def _configured_period_bands(
    axis: list[PeriodAxisItem], configured: list[CadenceBand]
) -> list[list[PeriodAxisItem]]:
    if not configured:
        return _auto_period_bands(axis)
    bands: list[list[PeriodAxisItem]] = []
    consumed: set[tuple[int, int]] = set()
    for band in configured:
        min_col, min_row, max_col, max_row = range_boundaries(band.cell_range)
        selected = [
            item
            for item in axis
            if (min_row or 1) <= item[2] <= (max_row or item[2])
            and (min_col or 1) <= item[3] <= (max_col or item[3])
            and item[1].kind == band.kind
        ]
        if selected:
            bands.append(selected)
            consumed.update((item[2], item[3]) for item in selected)
    remainder = [item for item in axis if (item[2], item[3]) not in consumed]
    bands.extend(_auto_period_bands(remainder))
    return bands


def _period_distance(previous: Period, current: Period) -> int | None:
    if previous.kind != current.kind:
        return None
    if previous.kind == "month" and previous.sort_key[0] and current.sort_key[0]:
        previous_ordinal = previous.sort_key[0] * 12 + previous.sort_key[1]
        current_ordinal = current.sort_key[0] * 12 + current.sort_key[1]
        return current_ordinal - previous_ordinal
    if previous.kind == "quarter" and previous.sort_key[0] and current.sort_key[0]:
        previous_ordinal = previous.sort_key[0] * 4 + previous.sort_key[1]
        current_ordinal = current.sort_key[0] * 4 + current.sort_key[1]
        return current_ordinal - previous_ordinal
    if previous.kind == "week":
        previous_year, _, previous_week = previous.sort_key
        current_year, _, current_week = current.sort_key
        if previous_year == current_year:
            return current_week - previous_week
        if previous_year == current_year == 0 and previous_week >= 40 and current_week <= 13:
            return 53 - previous_week + current_week
    return None


def _period_findings(
    sheet: SheetSnapshot,
    region: TableRegion,
    cadence_bands: list[CadenceBand] | None = None,
) -> list[Finding]:
    axis = _period_axis(sheet, region)
    if not axis:
        return []
    findings: list[Finding] = []
    bands = _configured_period_bands(axis, cadence_bands or [])
    if region.orientation == "wide":
        for band in bands:
            duplicates = [
                label
                for label, count in Counter(item[0] for item in band).items()
                if count > 1
            ]
            for label in duplicates:
                findings.append(
                    Finding(
                        artifact="excel",
                        finding_class=FindingClass.PERIOD_DUPLICATE,
                        sheet=sheet.name,
                        element=region.region_id,
                        current_value=label,
                        message=(
                            f"{sheet.name} ({region.cell_range}): "
                            f"duplicate period {label!r}"
                        ),
                    )
                )
    for band in bands:
        axis_pairs = [(label, period) for label, period, _, _ in band]
        unique: list[tuple[str, Period]] = []
        for label, period in axis_pairs:
            if not unique or period.sort_key != unique[-1][1].sort_key:
                unique.append((label, period))
        for (previous_label, previous), (current_label, current) in pairwise(unique):
            if not is_period_after(current, previous):
                findings.append(
                    Finding(
                        artifact="excel",
                        finding_class=FindingClass.PERIOD_OUT_OF_ORDER,
                        sheet=sheet.name,
                        element=region.region_id,
                        baseline_value=previous_label,
                        current_value=current_label,
                        message=(
                            f"{sheet.name} ({region.cell_range}): period "
                            f"{current_label!r} does not follow {previous_label!r}"
                        ),
                    )
                )
                continue
            distance = _period_distance(previous, current)
            if distance is not None and distance > 1:
                findings.append(
                    Finding(
                        artifact="excel",
                        finding_class=FindingClass.PERIOD_GAP,
                        sheet=sheet.name,
                        element=region.region_id,
                        baseline_value=previous_label,
                        current_value=current_label,
                        message=(
                            f"{sheet.name} ({region.cell_range}): period gap between "
                            f"{previous_label!r} and {current_label!r}"
                        ),
                    )
                )
    return findings


def _structure_findings(
    workbook: WorkbookSnapshot,
    ignored_sheets: set[str],
) -> tuple[list[Finding], int]:
    findings: list[Finding] = []
    unsupported = 0
    for named in workbook.named_ranges:
        resolution = resolve_reference(workbook, named.target, host_sheet="")
        if resolution.ranges and all(
            resolved.sheet in ignored_sheets for resolved in resolution.ranges
        ):
            continue
        if resolution.status is ReferenceStatus.UNSUPPORTED:
            unsupported += 1
        elif resolution.status is ReferenceStatus.INVALID:
            findings.append(
                Finding(
                    artifact="excel",
                    finding_class=FindingClass.NAMED_RANGE_INVALID,
                    element=named.name,
                    current_value=named.target,
                    message=f"named range {named.name!r} points to an invalid target",
                )
            )
    for chart in workbook.charts:
        if chart.sheet in ignored_sheets:
            continue
        for series in chart.series:
            sizes: dict[str, int] = {}
            for kind, target in (
                ("values", series.values_ref),
                ("categories", series.categories_ref),
            ):
                if target is None:
                    continue
                resolution = resolve_reference(
                    workbook,
                    target,
                    host_sheet=chart.sheet,
                )
                if resolution.status is ReferenceStatus.UNSUPPORTED:
                    unsupported += 1
                elif resolution.status is ReferenceStatus.INVALID:
                    findings.append(
                        Finding(
                            artifact="excel",
                            finding_class=FindingClass.CHART_REFERENCE_INVALID,
                            sheet=chart.sheet,
                            element=chart.title or chart.chart_type,
                            current_value=target,
                            message=(
                                f"chart {chart.title or chart.chart_type!r} series "
                                f"{series.index} has an invalid {kind} reference"
                            ),
                        )
                    )
                else:
                    sizes[kind] = resolution.size
            if len(sizes) == 2 and sizes["values"] != sizes["categories"]:
                findings.append(
                    Finding(
                        artifact="excel",
                        finding_class=FindingClass.CHART_LENGTH_MISMATCH,
                        sheet=chart.sheet,
                        element=chart.title or chart.chart_type,
                        baseline_value=str(sizes["categories"]),
                        current_value=str(sizes["values"]),
                        message=(
                            f"chart {chart.title or chart.chart_type!r} series "
                            f"{series.index} has different category and value lengths"
                        ),
                    )
                )
    for pivot in workbook.pivots:
        target = (
            f"{pivot.source_sheet}!{pivot.source_ref}"
            if pivot.source_sheet and pivot.source_ref
            else ""
        )
        resolution = resolve_reference(workbook, target, host_sheet="")
        if resolution.status is ReferenceStatus.UNSUPPORTED:
            unsupported += 1
        elif resolution.status is ReferenceStatus.INVALID:
            findings.append(
                Finding(
                    artifact="excel",
                    finding_class=FindingClass.PIVOT_SOURCE_INVALID,
                    element=pivot.name,
                    current_value=target,
                    message=f"pivot table {pivot.name!r} has an invalid source",
                )
            )
    return findings, unsupported


def preflight_workbook(
    workbook: WorkbookSnapshot,
    profile: DeliverableProfile,
    *,
    defer_impacts: bool = False,
    cancellation_token: CancellationToken | None = None,
) -> ExcelPreflightResult:
    check_cancelled(cancellation_token)
    result = ExcelPreflightResult()
    ignored_sheets = set(profile.excel.ignore_sheets)
    ignored_sheets.update(
        sheet_name
        for sheet_name, sheet_profile in profile.excel.sheets.items()
        if sheet_profile.ignore
    )

    if workbook.calculation_mode == "manual":
        result.findings.append(
            Finding(
                artifact="excel",
                finding_class=FindingClass.CALCULATION_MODE,
                element="workbook calculation",
                current_value="manual",
                message="workbook calculation mode is manual; saved values may be stale",
            )
        )
    result.findings.extend(workbook_risk_findings(workbook))
    result.coverage.append(
        CoverageItem(
            check_id="excel-workbook-settings",
            label="Workbook calculation and package risks",
            artifact="excel",
            state=CoverageState.CHECKED,
            findings=len(result.findings),
        )
    )

    formula_start = len(result.findings)
    if workbook.formula_presence_available:
        self_alignment = align_workbooks(
            workbook,
            workbook,
            profile,
            cancellation_token=cancellation_token,
        )
        result.findings.extend(
            diff_workbook_formulas(
                workbook,
                workbook,
                self_alignment,
                profile,
                cycle=False,
                cancellation_token=cancellation_token,
            )
        )
        for sheet_name, regions in self_alignment.regions.items():
            sheet = workbook.sheet(sheet_name)
            sheet_profile = profile.sheet_profile(sheet_name)
            for region_alignment in regions:
                result.findings.extend(
                    _formula_gap_findings(
                        sheet,
                        region_alignment.current,
                        sheet_profile,
                    )
                )
        for sheet in workbook.sheets:
            check_cancelled(cancellation_token)
            if sheet.name in ignored_sheets:
                continue
            missing_cache = sum(
                cell.has_formula and cell.value is None
                for cell in sheet.cells.values()
            )
            if missing_cache:
                result.findings.append(
                    Finding(
                        artifact="excel",
                        finding_class=FindingClass.FORMULA_CACHE_MISSING,
                        sheet=sheet.name,
                        current_value=str(missing_cache),
                        message=(
                            f"{sheet.name}: {missing_cache} formula cells have no saved "
                            "calculated value"
                        ),
                    )
                )
    formula_state = (
        CoverageState.CHECKED if workbook.formulas_available else CoverageState.DEGRADED
    )
    result.coverage.append(
        CoverageItem(
            check_id="excel-intrinsic-formulas",
            label="Formula errors, consistency, gaps, and cached values",
            artifact="excel",
            state=formula_state,
            findings=len(result.findings) - formula_start,
            detail=(
                workbook.formula_detail
                if not workbook.formulas_available
                else f"Formula text source: {workbook.formula_source}"
            ),
        )
    )
    dependency_graph: DependencyGraph | None = None
    if workbook.formulas_available:
        dependency_graph = build_dependency_graph(
            workbook,
            cancellation_token=cancellation_token,
        )
        result.dependency_graph = dependency_graph
        if not defer_impacts:
            annotate_impacts(result.findings, dependency_graph)
        dependency_state = dependency_graph.coverage_state
        dependency_detail = dependency_graph.coverage_detail
    else:
        dependency_state = CoverageState.UNAVAILABLE
        dependency_detail = "Formula text is unavailable for dependency extraction"
    result.coverage.append(
        CoverageItem(
            check_id="excel-dependencies",
            label="Formula dependency impact tracing",
            artifact="excel",
            state=dependency_state,
            detail=dependency_detail,
        )
    )

    period_start = len(result.findings)
    for sheet in workbook.sheets:
        check_cancelled(cancellation_token)
        if sheet.name in ignored_sheets:
            continue
        sheet_profile = profile.sheet_profile(sheet.name)
        for region in detect_regions(sheet, sheet_profile):
            result.findings.extend(
                _period_findings(
                    sheet,
                    region,
                    None if sheet_profile is None else sheet_profile.cadence_bands,
                )
            )
    result.coverage.append(
        CoverageItem(
            check_id="excel-period-integrity",
            label="Period ordering, duplicates, and detectable gaps",
            artifact="excel",
            state=CoverageState.CHECKED,
            findings=len(result.findings) - period_start,
        )
    )

    structure_start = len(result.findings)
    check_cancelled(cancellation_token)
    structure_findings, unsupported_references = _structure_findings(
        workbook,
        ignored_sheets,
    )
    result.findings.extend(structure_findings)
    structure_details: list[str] = []
    if not workbook.tables_available:
        structure_details.append("Excel table metadata is unavailable for this format")
    if not workbook.charts_available:
        structure_details.append("Complete Excel chart metadata is unavailable")
    if unsupported_references:
        structure_details.append(
            f"{unsupported_references} unsupported references were not validated"
        )
    for sheet in workbook.sheets:
        if sheet.name in ignored_sheets:
            continue
        if sheet.visibility != "visible" or sheet.hidden_rows or sheet.hidden_columns:
            detail = []
            if sheet.visibility != "visible":
                detail.append(f"sheet is {sheet.visibility}")
            if sheet.hidden_rows:
                detail.append(f"{len(sheet.hidden_rows)} hidden rows")
            if sheet.hidden_columns:
                detail.append(f"{len(sheet.hidden_columns)} hidden columns")
            result.findings.append(
                Finding(
                    artifact="excel",
                    finding_class=FindingClass.HIDDEN_CONTENT,
                    sheet=sheet.name,
                    current_value=", ".join(detail),
                    message=f"{sheet.name}: " + ", ".join(detail),
                )
            )
    result.coverage.append(
        CoverageItem(
            check_id="excel-intrinsic-structure",
            label="Tables, named ranges, charts, pivots, and hidden content",
            artifact="excel",
            state=(
                CoverageState.DEGRADED
                if structure_details
                else CoverageState.CHECKED
            ),
            findings=len(result.findings) - structure_start,
            detail="; ".join(structure_details),
        )
    )
    interaction_state, interaction_detail = interaction_rule_coverage(workbook)
    result.coverage.append(
        CoverageItem(
            check_id="excel-interaction-rules",
            label="Data validation and conditional-format rules",
            artifact="excel",
            state=interaction_state,
            detail=interaction_detail,
        )
    )
    style_state, style_detail = conditional_style_coverage(workbook)
    result.coverage.append(
        CoverageItem(
            check_id="excel-conditional-format-styles",
            label="Conditional-format style capture",
            artifact="excel",
            state=style_state,
            detail=style_detail,
        )
    )
    result.coverage.append(
        availability_coverage(
            artifact="excel",
            rule_count=sum(
                len(sheet.availability_rules)
                for sheet_name, sheet in profile.excel.sheets.items()
                if sheet_name not in ignored_sheets
            ),
            issues=excel_availability_issues(workbook, profile),
        )
    )
    controls = evaluate_controls(
        workbook,
        profile.excel.controls,
        ignored_sheets=ignored_sheets,
    )
    result.findings.extend(controls.findings)
    if controls.coverage is not None:
        result.coverage.append(controls.coverage)
    annotate_chart_impacts(result.findings, workbook, dependency_graph)
    limit_impacts(result.findings)
    return result
