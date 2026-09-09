"""Table-region detection: find data blocks and infer their orientation.

Sheets may hold one clean table or several blocks (dashboard layouts).
Detection flood-fills populated cells into 4-connected components, merges
overlapping bounding boxes, and infers each region's orientation from
where period labels sit:

- period labels across the top row      -> ``wide``  (periods on columns)
- period labels down an early column    -> ``long``  (periods on rows)
- no period axis                        -> ``block`` (label/value layout)

A profile's `RegionOverride` list for a sheet replaces detection entirely
for that sheet — analyst pins always win.
"""

from collections import deque
from collections.abc import Sequence
from dataclasses import dataclass
from itertools import pairwise

from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import column_index_from_string, range_boundaries

from qc_tool.config.profile import Orientation, RegionOverride, SheetProfile
from qc_tool.excel.periods import Period, is_period_label, parse_period
from qc_tool.io.model import SheetSnapshot

#: Minimum count and share of period labels required to call an axis.
_MIN_PERIODS = 2
_PERIOD_SHARE = 0.6

#: Evidence-scored period-band thresholds (Step 3, client feedback: a
#: percent-formatted row sitting beside a date header row must never win by
#: accident). Pinned by synthetic counterexamples; do not retune against
#: external data without a plan amendment.
_BAND_MIN_PERIODS = 4
_BAND_MIN_SHARE = 0.8
_BAND_MIN_CONTIGUOUS_SHARE = 0.9
_BAND_AUTOSELECT_MARGIN = 1.5
_BAND_OPPOSING_MARGIN = 0.2


@dataclass(frozen=True, slots=True)
class PeriodBandCandidate:
    """One scored row/column period-axis candidate."""

    axis: str  # "rows" | "columns"
    anchor: int  # key_col for rows, header_row for columns
    positions: dict[int, Period]
    score: float


@dataclass(frozen=True, slots=True)
class PeriodBandSuggestion:
    """A strong internal axis that needs an explicit profile region pin."""

    sheet: str
    region_range: str
    axis: str
    anchor: int
    period_count: int


def _longest_contiguous_run(positions: Sequence[int]) -> int:
    if not positions:
        return 0
    ordered = sorted(set(positions))
    best = current = 1
    for previous, current_position in pairwise(ordered):
        if current_position == previous + 1:
            current += 1
            best = max(best, current)
        else:
            current = 1
    return best


def _score_period_band(
    axis: str, anchor: int, values: Sequence[tuple[int, object]]
) -> PeriodBandCandidate | None:
    """An evidence-scored candidate, or ``None`` if it fails a hard gate."""
    populated = [(position, value) for position, value in values if value is not None]
    if not populated:
        return None
    parsed = {
        position: period
        for position, value in populated
        if (period := parse_period(value)) is not None
    }
    if len(parsed) < _BAND_MIN_PERIODS:
        return None
    if len(parsed) / len(populated) < _BAND_MIN_SHARE:
        return None
    ordered_positions = sorted(parsed)
    periods_in_order = [parsed[position] for position in ordered_positions]
    # monotonic cadence: allow ties, reject any backward step
    if any(
        later.sort_key < earlier.sort_key
        for earlier, later in pairwise(periods_in_order)
    ):
        return None
    longest_run = _longest_contiguous_run(ordered_positions)
    if longest_run / len(parsed) < _BAND_MIN_CONTIGUOUS_SHARE:
        return None
    return PeriodBandCandidate(
        axis=axis, anchor=anchor, positions=parsed, score=float(len(parsed))
    )


def _select_period_band(
    candidates: Sequence[PeriodBandCandidate],
) -> PeriodBandCandidate | None:
    """The clear winner, or ``None`` when competing bands make it ambiguous.

    Never silently reclassifies a block: a tie or a near-tie between two
    candidates (same axis or opposing axes) means "do not guess" rather than
    "pick one". Callers fall back to `block` in that case.
    """
    if not candidates:
        return None
    best = max(candidates, key=lambda candidate: candidate.score)
    same_axis_runner_up = max(
        (c for c in candidates if c.axis == best.axis and c is not best),
        key=lambda candidate: candidate.score,
        default=None,
    )
    if same_axis_runner_up is not None and best.score < (
        same_axis_runner_up.score * _BAND_AUTOSELECT_MARGIN
    ):
        return None
    opposing_best = max(
        (c for c in candidates if c.axis != best.axis),
        key=lambda candidate: candidate.score,
        default=None,
    )
    if opposing_best is not None and opposing_best.score >= best.score * (
        1 - _BAND_OPPOSING_MARGIN
    ):
        return None
    return best


@dataclass(frozen=True, slots=True)
class TableRegion:
    sheet: str
    min_row: int
    min_col: int
    max_row: int
    max_col: int
    orientation: Orientation
    header_row: int | None
    key_col: int | None  # column holding row identity (long/block layouts)
    period_axis: str  # "rows" | "columns" | "none"

    @property
    def cell_range(self) -> str:
        return (
            f"{get_column_letter(self.min_col)}{self.min_row}:"
            f"{get_column_letter(self.max_col)}{self.max_row}"
        )

    @property
    def region_id(self) -> str:
        return f"{self.sheet}!{self.cell_range}"


def _components(sheet: SheetSnapshot) -> list[tuple[int, int, int, int]]:
    """4-connected components of populated cells -> bounding boxes."""
    remaining = set(sheet.cells)
    boxes: list[tuple[int, int, int, int]] = []
    while remaining:
        seed = remaining.pop()
        queue = deque([seed])
        min_row = max_row = seed[0]
        min_col = max_col = seed[1]
        while queue:
            row, col = queue.popleft()
            min_row, max_row = min(min_row, row), max(max_row, row)
            min_col, max_col = min(min_col, col), max(max_col, col)
            for neighbor in ((row - 1, col), (row + 1, col), (row, col - 1), (row, col + 1)):
                if neighbor in remaining:
                    remaining.discard(neighbor)
                    queue.append(neighbor)
        boxes.append((min_row, min_col, max_row, max_col))
    return _merge_overlaps(boxes)


def _overlaps(a: tuple[int, int, int, int], b: tuple[int, int, int, int]) -> bool:
    return not (a[2] < b[0] or b[2] < a[0] or a[3] < b[1] or b[3] < a[1])


def _merge_overlaps(boxes: list[tuple[int, int, int, int]]) -> list[tuple[int, int, int, int]]:
    merged = True
    while merged:
        merged = False
        result: list[tuple[int, int, int, int]] = []
        for box in boxes:
            for index, existing in enumerate(result):
                if _overlaps(box, existing):
                    result[index] = (
                        min(existing[0], box[0]),
                        min(existing[1], box[1]),
                        max(existing[2], box[2]),
                        max(existing[3], box[3]),
                    )
                    merged = True
                    break
            else:
                result.append(box)
        boxes = result
    return sorted(boxes)


def _axis_is_periodic(values: Sequence[object]) -> bool:
    present = [v for v in values if v is not None]
    if not present:
        return False
    hits = sum(1 for v in present if is_period_label(v))
    return hits >= _MIN_PERIODS and hits / len(present) >= _PERIOD_SHARE


def period_positions(sheet: SheetSnapshot, region: TableRegion) -> dict[int, Period]:
    """Map each period-axis position (row or column number) to its `Period`.

    For ``rows`` the detected `key_col` is tried first, then every region
    column, so profile-pinned regions with a non-period key column still
    resolve. Positions whose label does not parse are absent — callers must
    treat them as non-recent. Empty for ``none`` axes.
    """
    if region.period_axis == "rows":
        candidates = [region.key_col] if region.key_col is not None else []
        candidates += [
            col
            for col in range(region.min_col, region.max_col + 1)
            if col not in candidates
        ]
        for col in candidates:
            start = region.min_row + (1 if region.header_row == region.min_row else 0)
            values = [
                (row, sheet.cells[(row, col)].value)
                for row in range(start, region.max_row + 1)
                if (row, col) in sheet.cells
            ]
            if not _axis_is_periodic([value for _, value in values]):
                continue
            positions = {
                row: period
                for row, value in values
                if (period := parse_period(value)) is not None
            }
            if positions:
                return positions
        return {}
    if region.period_axis == "columns":
        rows = [region.header_row] if region.header_row is not None else []
        rows += [row for row in range(region.min_row, region.max_row + 1) if row not in rows]
        for row in rows:
            values = [
                (col, sheet.cells[(row, col)].value)
                for col in range(region.min_col + 1, region.max_col + 1)
                if (row, col) in sheet.cells
            ]
            if not _axis_is_periodic([value for _, value in values]):
                continue
            positions = {
                col: period
                for col, value in values
                if (period := parse_period(value)) is not None
            }
            if positions:
                return positions
        return {}
    return {}


def _infer_region(
    sheet: SheetSnapshot, box: tuple[int, int, int, int]
) -> TableRegion:
    min_row, min_col, max_row, max_col = box
    cells = sheet.cells

    def cell_value(row: int, col: int) -> object:
        cell = cells.get((row, col))
        return cell.value if cell is not None else None

    top_row_values = [
        (col, cell_value(min_row, col)) for col in range(min_col + 1, max_col + 1)
    ]
    candidates: list[PeriodBandCandidate] = []
    wide_candidate = _score_period_band("columns", min_row, top_row_values)
    if wide_candidate is not None:
        candidates.append(wide_candidate)
    long_candidate_columns = range(min_col, min(min_col + 3, max_col + 1))
    for col in long_candidate_columns:
        column_values = [
            (row, cell_value(row, col)) for row in range(min_row + 1, max_row + 1)
        ]
        candidate = _score_period_band("rows", col, column_values)
        if candidate is not None:
            candidates.append(candidate)

    # Evidence-scored path: a clear winner (>=4 periods, >=80% share, monotonic,
    # >=90% contiguous, and no near-tied competing band) wins outright. A
    # genuine tie between competing bands (e.g. a percent-formatted row beside
    # a date header) never guesses; it falls through to `block` below instead
    # of silently reclassifying on whichever axis happened to be checked first.
    if candidates:
        winner = _select_period_band(candidates)
        if winner is not None:
            if winner.axis == "columns":
                return TableRegion(
                    sheet=sheet.name,
                    min_row=min_row,
                    min_col=min_col,
                    max_row=max_row,
                    max_col=max_col,
                    orientation="wide",
                    header_row=min_row,
                    key_col=min_col,
                    period_axis="columns",
                )
            return TableRegion(
                sheet=sheet.name,
                min_row=min_row,
                min_col=min_col,
                max_row=max_row,
                max_col=max_col,
                orientation="long",
                header_row=min_row,
                key_col=winner.anchor,
                period_axis="rows",
            )
        return TableRegion(
            sheet=sheet.name,
            min_row=min_row,
            min_col=min_col,
            max_row=max_row,
            max_col=max_col,
            orientation="block",
            header_row=None,
            key_col=min_col,
            period_axis="none",
        )

    # No candidate cleared the strict evidence bar (e.g. a short fixture with
    # only 2-3 periods): fall back to the original lenient detection so small,
    # already-relied-upon layouts keep working exactly as before.
    top_row = [value for _, value in top_row_values]
    if _axis_is_periodic(top_row):
        return TableRegion(
            sheet=sheet.name,
            min_row=min_row,
            min_col=min_col,
            max_row=max_row,
            max_col=max_col,
            orientation="wide",
            header_row=min_row,
            key_col=min_col,
            period_axis="columns",
        )

    for col in range(min_col, min(min_col + 3, max_col + 1)):
        column_values = [
            cells[(row, col)].value
            for row in range(min_row + 1, max_row + 1)
            if (row, col) in cells
        ]
        if _axis_is_periodic(column_values):
            return TableRegion(
                sheet=sheet.name,
                min_row=min_row,
                min_col=min_col,
                max_row=max_row,
                max_col=max_col,
                orientation="long",
                header_row=min_row,
                key_col=col,
                period_axis="rows",
            )

    return TableRegion(
        sheet=sheet.name,
        min_row=min_row,
        min_col=min_col,
        max_row=max_row,
        max_col=max_col,
        orientation="block",
        header_row=None,
        key_col=min_col,
        period_axis="none",
    )


def _region_from_override(sheet_name: str, override: RegionOverride) -> TableRegion:
    min_col, min_row, max_col, max_row = range_boundaries(override.cell_range)
    orientation = override.orientation
    period_axis = {"long": "rows", "wide": "columns", "block": "none"}[orientation]
    key_col = (
        column_index_from_string(override.key_column) if override.key_column else min_col
    )
    header_row = override.header_row
    if header_row is None and orientation != "block":
        header_row = min_row
    return TableRegion(
        sheet=sheet_name,
        min_row=int(min_row or 1),
        min_col=int(min_col or 1),
        max_row=int(max_row or 1),
        max_col=int(max_col or 1),
        orientation=orientation,
        header_row=header_row,
        key_col=key_col,
        period_axis=period_axis,
    )


def detect_regions(
    sheet: SheetSnapshot, profile: SheetProfile | None = None
) -> list[TableRegion]:
    """Detect table regions, honoring profile pins and ignores."""
    if profile is not None:
        if profile.ignore:
            return []
        if profile.regions:
            return [_region_from_override(sheet.name, o) for o in profile.regions]
    if not sheet.cells:
        return []
    return [_infer_region(sheet, box) for box in _components(sheet)]


def internal_period_band_suggestions(
    sheet: SheetSnapshot, regions: Sequence[TableRegion]
) -> list[PeriodBandSuggestion]:
    """Strong axes inside blocks that are unsafe to split automatically."""
    suggestions: list[PeriodBandSuggestion] = []
    for region in regions:
        if region.orientation != "block":
            continue
        row_values: dict[int, list[tuple[int, object]]] = {}
        column_values: dict[int, list[tuple[int, object]]] = {}
        for (row, column), cell in sheet.cells.items():
            if not (
                region.min_row <= row <= region.max_row
                and region.min_col <= column <= region.max_col
            ):
                continue
            row_values.setdefault(row, []).append((column, cell.value))
            column_values.setdefault(column, []).append((row, cell.value))
        candidates = [
            candidate
            for row, values in row_values.items()
            if row != region.min_row
            and (candidate := _score_period_band("columns", row, values)) is not None
        ]
        candidates.extend(
            candidate
            for column, values in column_values.items()
            if (candidate := _score_period_band("rows", column, values)) is not None
        )
        for candidate in sorted(candidates, key=lambda item: (-item.score, item.anchor)):
            suggestions.append(
                PeriodBandSuggestion(
                    sheet=sheet.name,
                    region_range=region.cell_range,
                    axis=candidate.axis,
                    anchor=candidate.anchor,
                    period_count=len(candidate.positions),
                )
            )
    return suggestions
