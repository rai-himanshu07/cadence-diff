"""Cycle-comparison prerequisite checks.

A prerequisite pins one profile-configured selector/scenario cell that must
be present, non-blank, and exactly equal between baseline and current before
any alignment, diffing, reports, or history proceed. A mismatch means the
two files were not prepared under the same configuration (for example, a
different dropdown selection), so a value-level comparison would prove
nothing until the analyst aligns them.
"""

from __future__ import annotations

from typing import TYPE_CHECKING

from openpyxl.utils.cell import coordinate_to_tuple

from qc_tool.config.profile import ComparisonPrerequisite
from qc_tool.io.model import WorkbookSnapshot
from qc_tool.run_action import RunActionItem

if TYPE_CHECKING:
    from qc_tool.config.resolved_input import ResolvedInputConfigurationV1, ResolvedRegion
    from qc_tool.excel.align import WorkbookAlignment
    from qc_tool.excel.regions import TableRegion


def _is_blank(value: object) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def check_comparison_prerequisites(
    baseline: WorkbookSnapshot,
    current: WorkbookSnapshot,
    prerequisites: list[ComparisonPrerequisite],
    *,
    member_id: str = "primary",
) -> list[RunActionItem]:
    """Return one bounded mismatch item per failing prerequisite.

    An empty result means every configured prerequisite matched. Only a
    bounded reason category is returned -- never the cell's actual value.
    """
    mismatches: list[RunActionItem] = []
    for prerequisite in prerequisites:
        base_cell = None
        curr_cell = None
        try:
            base_cell = baseline.sheet(prerequisite.sheet).cell(prerequisite.cell)
        except (KeyError, ValueError):
            base_cell = None
        try:
            curr_cell = current.sheet(prerequisite.sheet).cell(prerequisite.cell)
        except (KeyError, ValueError):
            curr_cell = None
        base_value = None if base_cell is None else base_cell.value
        curr_value = None if curr_cell is None else curr_cell.value
        if base_cell is None or curr_cell is None:
            reason = "the configured sheet or cell is missing from one or both files"
        elif _is_blank(base_value) or _is_blank(curr_value):
            reason = "the prerequisite cell is blank in one or both files"
        elif base_value != curr_value:
            reason = "baseline and current do not have the same prerequisite value"
        else:
            continue
        mismatches.append(
            RunActionItem(
                member_id=member_id,
                sheet=prerequisite.sheet,
                cell=prerequisite.cell,
                label=prerequisite.name,
                detail=reason,
            )
        )
    return mismatches


def check_resolved_selectors(
    baseline: WorkbookSnapshot,
    current: WorkbookSnapshot,
    resolved: ResolvedInputConfigurationV1 | None,
    *,
    member_id: str = "primary",
) -> list[RunActionItem]:
    """Run-time gate for selector prerequisites declared through the
    mode-aware configuration workspace (plan-20260913, Step 8) -- a
    resolved-configuration analog of ``check_comparison_prerequisites``
    with the identical bounded, value-free mismatch contract. Selector
    VALUES are read only to compare them; neither side's value is ever
    returned or persisted.
    """
    if resolved is None:
        return []
    member = next((m for m in resolved.members if m.member_id == member_id), None)
    if member is None:
        return []
    mismatches: list[RunActionItem] = []
    for sheet in member.sheets:
        for selector in sheet.selectors:
            if not sheet.current_sheet_name or not selector.current_cell:
                continue
            base_cell = None
            curr_cell = None
            if sheet.baseline_sheet_name and selector.baseline_cell:
                try:
                    base_cell = baseline.sheet(sheet.baseline_sheet_name).cell(
                        selector.baseline_cell
                    )
                except (KeyError, ValueError):
                    base_cell = None
            try:
                curr_cell = current.sheet(sheet.current_sheet_name).cell(selector.current_cell)
            except (KeyError, ValueError):
                curr_cell = None
            base_value = None if base_cell is None else base_cell.value
            curr_value = None if curr_cell is None else curr_cell.value
            if base_cell is None or curr_cell is None:
                reason = "the configured sheet or cell is missing from one or both files"
            elif _is_blank(base_value) or _is_blank(curr_value):
                reason = "the selector cell is blank in one or both files"
            elif base_value != curr_value:
                reason = "baseline and current do not have the same selector value"
            else:
                continue
            mismatches.append(
                RunActionItem(
                    member_id=member_id,
                    sheet=sheet.current_sheet_name,
                    cell=selector.current_cell,
                    label=selector.selector_id,
                    detail=reason,
                )
            )
    return mismatches


def _region_anchor_containing(
    regions: tuple[ResolvedRegion, ...], current_region: TableRegion
) -> ResolvedRegion | None:
    """The first region whose ``current_outer_range`` top-left anchor falls
    inside ``current_region`` -- mirrors ``qc_tool.excel.align.
    _matching_row_identity_rule``'s anchor-containment matching exactly, so
    a region keeps resolving across ordinary row/column growth between
    cycles instead of needing an unstable exact-range match.
    """
    for region in regions:
        if not region.current_outer_range:
            continue
        try:
            anchor_row, anchor_col = coordinate_to_tuple(region.current_outer_range.split(":")[0])
        except ValueError:
            continue
        if (
            current_region.min_row <= anchor_row <= current_region.max_row
            and current_region.min_col <= anchor_col <= current_region.max_col
        ):
            return region
    return None


def check_blank_identity_keys(
    alignment: WorkbookAlignment,
    resolved: ResolvedInputConfigurationV1 | None,
    *,
    member_id: str = "primary",
) -> list[RunActionItem]:
    """Run-time gate for a keyed region's ``blank_key_policy == "block"``
    (plan-20260913, Step 8): refuses the run when the alignment already
    computed by ``align_workbooks()`` found any row excluded from key
    matching because an identity-column component was blank, rather than
    silently letting it surface as an ordinary insert/delete (the
    ``"system_default"``/``"tolerate"`` behavior). Only a bounded row COUNT
    is ever returned -- never a key value.

    Deliberately a POST-alignment check, not a change to
    ``_align_rows_by_identity`` itself: that function already computes and
    discloses ``AxisAlignment.blank_key_rows` unconditionally; this
    function only decides whether a nonzero count should block, keeping the
    heavily-tested core alignment function free of any new control-flow
    branch or exception path.
    """
    if resolved is None:
        return []
    member = next((m for m in resolved.members if m.member_id == member_id), None)
    if member is None:
        return []
    blockers: list[RunActionItem] = []
    for sheet in member.sheets:
        sheet_name = sheet.current_sheet_name
        if not sheet_name:
            continue
        for region_alignment in alignment.regions.get(sheet_name, []):
            if region_alignment.rows.method != "keys" or region_alignment.rows.blank_key_rows <= 0:
                continue
            region = _region_anchor_containing(sheet.regions, region_alignment.current)
            if region is None or region.blank_key_policy != "block":
                continue
            blockers.append(
                RunActionItem(
                    member_id=member_id,
                    sheet=sheet_name,
                    label=region.region_id,
                    detail=(
                        f"{region_alignment.rows.blank_key_rows} row(s) have a "
                        "blank identity-key column; this region's policy blocks "
                        "on blank keys"
                    ),
                )
            )
    return blockers
