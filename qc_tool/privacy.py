"""Fail-closed privacy verification for strictly sanitized Office packages."""

import hashlib
import io
import re
import zipfile
from pathlib import Path
from typing import Any, cast
from xml.etree import ElementTree

from openpyxl import load_workbook
from pptx import Presentation
from pydantic import BaseModel, Field

from qc_tool.excel.periods import parse_period

_PSEUDONYM_RE = re.compile(r"TXT_[0-9a-f]{6}|(?:Sheet|Table|Chart|Shape|Name)_\d+", re.I)
_DISPLAY_NUMBER_RE = re.compile(
    r"(?<!\w)\(?\s*[$€£₹]?\s*[-+]?\d[\d,]*(?:\.\d+)?\s*"
    r"(?:%|bn|mn|k|m|b)?\s*\)?(?!\w)",
    re.IGNORECASE,
)
_PERIOD_TOKEN_RE = re.compile(
    r"\b(?:"
    r"(?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)[A-Za-z]*[-_ ]?'?\d{2,4}"
    r"|\d{4}-\d{2}(?:-\d{2})?"
    r"|(?:CW|WK|W)[-_ ]?\d{1,2}(?:[-_ ]?'?\d{2,4})?"
    r"|Q[1-4](?:[-_ ]?(?:FY)?[-_ ]?'?\d{2,4})?"
    r")\b",
    re.IGNORECASE,
)
_FORMULA_STRING_RE = re.compile(r'"((?:[^"]|"")*)"')
_GENERIC_SHAPE_RE = re.compile(
    r"^(?:Title|Subtitle|Content Placeholder|Text Placeholder|TextBox|Chart|Table|"
    r"Shape|Picture|Object|Media|Slide Number Placeholder|Date Placeholder|"
    r"Footer Placeholder|Shape_)\s*_?\d+(?:_\d+)?$",
    re.IGNORECASE,
)
_FORBIDDEN_PART_MARKERS = (
    "vbaproject",
    "/externallinks/",
    "/comments",
    "/threadedcomments/",
    "/persons/",
    "/customxml/",
    "/activex/",
    "/ctrlprops/",
    "/oleobject",
    "/media/",
    "/connections",
    "docprops/custom.xml",
)
_SENSITIVE_PROPERTY_NAMES = {
    "creator",
    "lastmodifiedby",
    "title",
    "subject",
    "description",
    "comments",
    "keywords",
    "category",
    "contentstatus",
    "identifier",
    "company",
    "manager",
    "hyperlinkbase",
}


class PrivacyIssue(BaseModel):
    code: str
    location: str
    message: str


class PrivacyReport(BaseModel):
    schema_version: int = 1
    file_type: str
    sha256: str
    safe: bool = True
    checks: list[str] = Field(default_factory=list)
    issues: list[PrivacyIssue] = Field(default_factory=list)

    def add(self, code: str, location: str, message: str) -> None:
        self.safe = False
        self.issues.append(PrivacyIssue(code=code, location=location, message=message))


def _local_name(tag: str) -> str:
    return tag.rsplit("}", 1)[-1].casefold()


def _visible_text_safe(text: str) -> bool:
    stripped = text.strip()
    if not stripped or parse_period(stripped) is not None:
        return True
    remainder = _PSEUDONYM_RE.sub("", stripped)
    remainder = _PERIOD_TOKEN_RE.sub("", remainder)
    remainder = _DISPLAY_NUMBER_RE.sub("", remainder)
    return not re.search(r"[A-Za-z0-9]", remainder)


def _scan_package(
    data: bytes,
    report: PrivacyReport,
    *,
    forbidden_tokens: list[str],
    allow_chart_embeddings: bool,
) -> list[tuple[str, bytes]]:
    embeddings: list[tuple[str, bytes]] = []
    try:
        archive = zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile:
        report.add("not-ooxml", "package", "file is not a valid OOXML ZIP package")
        return embeddings
    with archive:
        for name in archive.namelist():
            lowered = name.casefold()
            part = archive.read(name)
            if lowered.startswith("ppt/embeddings/"):
                if allow_chart_embeddings and lowered.endswith(".xlsx"):
                    embeddings.append((name, part))
                else:
                    report.add("embedding", name, "unverifiable embedded object remains")
            if any(marker in lowered for marker in _FORBIDDEN_PART_MARKERS):
                report.add("forbidden-part", name, "sensitive or active package part remains")
            if lowered.endswith(".rels"):
                try:
                    root = ElementTree.fromstring(part)
                except ElementTree.ParseError:
                    report.add("invalid-xml", name, "relationship part is not valid XML")
                else:
                    for relationship in root:
                        if relationship.get("TargetMode") == "External":
                            report.add(
                                "external-relationship",
                                name,
                                f"external target remains: {relationship.get('Target', '')}",
                            )
            if lowered in {"docprops/core.xml", "docprops/app.xml"}:
                try:
                    root = ElementTree.fromstring(part)
                except ElementTree.ParseError:
                    report.add("invalid-metadata", name, "metadata part is not valid XML")
                else:
                    for element in root.iter():
                        if (
                            _local_name(element.tag) in _SENSITIVE_PROPERTY_NAMES
                            and (element.text or "").strip()
                        ):
                            report.add(
                                "metadata",
                                name,
                                f"non-empty {_local_name(element.tag)} property remains",
                            )
            if lowered.startswith(
                ("ppt/slidemasters/", "ppt/slidelayouts/")
            ) and lowered.endswith(".xml"):
                try:
                    root = ElementTree.fromstring(part)
                except ElementTree.ParseError:
                    report.add("invalid-xml", name, "master/layout part is not valid XML")
                else:
                    for element in root.iter():
                        if _local_name(element.tag) == "t" and not _visible_text_safe(
                            element.text or ""
                        ):
                            report.add(
                                "master-layout-text",
                                name,
                                "non-pseudonymous master/layout text remains",
                            )
            folded = part.lower()
            for token in forbidden_tokens:
                encoded = token.casefold().encode("utf-8")
                encoded_utf16 = token.casefold().encode("utf-16le")
                if encoded in folded or encoded_utf16 in folded:
                    report.add("forbidden-token", name, f"token {token!r} remains")
    report.checks.extend(
        ["package-parts", "external-relationships", "document-metadata", "token-scan"]
    )
    return embeddings


def _verify_workbook_data(
    data: bytes,
    report: PrivacyReport,
    *,
    embedded: bool = False,
) -> None:
    try:
        workbook = load_workbook(io.BytesIO(data), data_only=False, keep_links=False)
    except Exception as exc:
        report.add("workbook-open", "workbook", f"cannot inspect workbook: {exc}")
        return
    dynamic_workbook = cast(Any, workbook)
    allowed_sheet = re.compile(r"^Sheet_?\d+$", re.I)
    if len(dynamic_workbook.custom_doc_props):
        report.add("custom-properties", "workbook", "custom document properties remain")
    for sheet in workbook.worksheets:
        if not allowed_sheet.fullmatch(sheet.title):
            report.add("sheet-name", sheet.title, "worksheet name is not anonymized")
        for row in sheet.iter_rows():
            for cell in row:
                if cell.comment is not None:
                    report.add("comment", f"{sheet.title}!{cell.coordinate}", "comment remains")
                if cell.hyperlink is not None:
                    report.add("hyperlink", f"{sheet.title}!{cell.coordinate}", "hyperlink remains")
                if not isinstance(cell.value, str):
                    continue
                if cell.data_type == "f":
                    if "[" in cell.value and "]" in cell.value:
                        report.add(
                            "external-formula",
                            f"{sheet.title}!{cell.coordinate}",
                            "external workbook reference remains",
                        )
                    for value in _FORMULA_STRING_RE.findall(cell.value):
                        if not _visible_text_safe(value):
                            report.add(
                                "formula-string",
                                f"{sheet.title}!{cell.coordinate}",
                                "literal text remains in formula",
                            )
                elif not _visible_text_safe(cell.value):
                    report.add(
                        "visible-text",
                        f"{sheet.title}!{cell.coordinate}",
                        "non-pseudonymous visible text remains",
                    )
        for table in sheet.tables.values():
            if not re.fullmatch(r"Table_\d+", table.name, re.I):
                report.add("table-name", table.name, "table name is not anonymized")
            for column in table.tableColumns:
                if not re.fullmatch(r"Column_\d+", column.name, re.I):
                    report.add(
                        "table-column-name",
                        table.name,
                        "table column name is not anonymized",
                    )
        for header_footer_name in (
            "oddHeader",
            "evenHeader",
            "firstHeader",
            "oddFooter",
            "evenFooter",
            "firstFooter",
        ):
            header_footer = getattr(sheet, header_footer_name)
            for section_name in ("left", "center", "right"):
                if getattr(header_footer, section_name).text:
                    report.add(
                        "header-footer",
                        sheet.title,
                        "header or footer text remains",
                    )
        if sheet.data_validations.dataValidation:
            report.add("data-validation", sheet.title, "data validation metadata remains")
        if getattr(sheet.conditional_formatting, "_cf_rules", {}):
            report.add(
                "conditional-formatting",
                sheet.title,
                "conditional formatting rules remain",
            )
        for chart_index, chart in enumerate(getattr(sheet, "_charts", []), start=1):
            for series_index, series in enumerate(chart.series):
                for attribute in ("val", "cat", "xVal", "yVal", "bubbleSize"):
                    source = getattr(series, attribute, None)
                    if source is None:
                        continue
                    for ref_attribute in ("numRef", "strRef"):
                        reference = getattr(source, ref_attribute, None)
                        formula = getattr(reference, "f", None)
                        if not formula or "!" not in formula:
                            continue
                        sheet_name = formula.rpartition("!")[0].strip("'").replace("''", "'")
                        if not allowed_sheet.fullmatch(sheet_name):
                            report.add(
                                "chart-reference",
                                f"{sheet.title} chart {chart_index} series {series_index}",
                                "chart references a non-anonymized worksheet",
                            )
    if workbook.defined_names and not embedded:
        report.add("defined-names", "workbook", "defined names remain")
    if not embedded:
        for sheet in workbook.worksheets:
            if getattr(sheet, "defined_names", None):
                report.add(
                    "defined-names",
                    "worksheet",
                    "sheet-scoped defined names remain",
                )
                break
    for style in dynamic_workbook._named_styles:
        if style.name != "Normal" and not re.fullmatch(r"Style_\d+", style.name, re.I):
            report.add("named-style", style.name, "named style is not anonymized")
    report.checks.extend(["workbook-identifiers", "workbook-visible-text"])


def _verify_presentation_data(data: bytes, report: PrivacyReport) -> None:
    try:
        presentation = Presentation(io.BytesIO(data))
    except Exception as exc:
        report.add("presentation-open", "presentation", f"cannot inspect deck: {exc}")
        return
    for slide_index, slide in enumerate(presentation.slides, start=1):
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame
            if notes is not None and notes.text.strip():
                report.add("speaker-notes", f"slide {slide_index}", "speaker notes remain")
        for shape_index, base_shape in enumerate(slide.shapes, start=1):
            shape = cast(Any, base_shape)
            if not _GENERIC_SHAPE_RE.fullmatch(base_shape.name):
                report.add(
                    "shape-name",
                    f"slide {slide_index} shape {shape_index}",
                    f"shape name {base_shape.name!r} is not anonymized",
                )
            for element in base_shape._element.iter():
                if _local_name(element.tag) == "cnvpr" and (
                    element.get("descr") or element.get("title")
                ):
                    report.add(
                        "shape-alt-text",
                        f"slide {slide_index} shape {shape_index}",
                        "shape alt text remains",
                    )
            if base_shape.has_text_frame and not _visible_text_safe(shape.text_frame.text):
                report.add(
                    "visible-text",
                    f"slide {slide_index} shape {shape_index}",
                    "non-pseudonymous visible text remains",
                )
            if base_shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        if not _visible_text_safe(cell.text):
                            report.add(
                                "table-text",
                                f"slide {slide_index} shape {shape_index}",
                                "non-pseudonymous table text remains",
                            )
            if base_shape.has_chart:
                for plot in shape.chart.plots:
                    for category in plot.categories:
                        if not _visible_text_safe(str(category)):
                            report.add(
                                "chart-category",
                                f"slide {slide_index} shape {shape_index}",
                                "non-pseudonymous chart category remains",
                            )
                    for series in plot.series:
                        if series.name and not _visible_text_safe(series.name):
                            report.add(
                                "chart-series",
                                f"slide {slide_index} shape {shape_index}",
                                "non-pseudonymous chart series name remains",
                            )
    report.checks.extend(["presentation-identifiers", "presentation-visible-text"])


def verify_sanitized(
    path: Path, *, forbidden_tokens: list[str] | None = None
) -> PrivacyReport:
    """Inspect strict sanitizer output; any unsupported residue fails closed."""
    data = path.read_bytes()
    suffix = path.suffix.casefold()
    report = PrivacyReport(
        file_type=suffix.lstrip("."),
        sha256=hashlib.sha256(data).hexdigest(),
    )
    tokens = [token for token in (forbidden_tokens or []) if token]
    embeddings = _scan_package(
        data,
        report,
        forbidden_tokens=tokens,
        allow_chart_embeddings=suffix == ".pptx",
    )
    if suffix == ".xlsx":
        _verify_workbook_data(data, report)
    elif suffix == ".pptx":
        _verify_presentation_data(data, report)
        for name, embedded_data in embeddings:
            embedded_report = PrivacyReport(
                file_type="embedded-xlsx",
                sha256=hashlib.sha256(embedded_data).hexdigest(),
            )
            _scan_package(
                embedded_data,
                embedded_report,
                forbidden_tokens=tokens,
                allow_chart_embeddings=False,
            )
            _verify_workbook_data(embedded_data, embedded_report, embedded=True)
            for issue in embedded_report.issues:
                report.add(issue.code, f"{name}:{issue.location}", issue.message)
            report.checks.append(f"embedded-chart-workbook:{name}")
    else:
        report.add("unsupported-format", path.name, "verification supports .xlsx/.pptx")
    report.safe = not report.issues
    return report
